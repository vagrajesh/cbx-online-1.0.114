"""1.0.99: Entity Compliance → Export (GET /api/ec/export): one workbook of the companies the person sees in the tool
folder's tables, coloured like the grid. The rows are added straight into the tool's own file and taken out again."""
import io
import shutil
from datetime import datetime, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from cbx.docwatch import ec_export as EX
from cbx.engine import files as FILES
from cbx.engine import loader
from cbx.engine.store import Store
from cbx.main import app
from cbx.models import AuditLog, Tool, ToolAccess, User
from cbx.util import local_today
from tests._pdfs import certificate
from tests.conftest import _make, login

SLUG = "entity-compliance"
OPEN, CLOSED = "Everyone with the tool", "Only its people"


@pytest.fixture()
def ec(owner, db):
    """A way to add rows to the tool's tables as the owner; whatever it added is removed afterwards."""
    t = loader.get(SLUG)
    st = Store(SLUG)
    me = db.query(User).filter(User.email == "owner@test.local").one()
    added: list[tuple[str, int]] = []

    def add(block, **values):
        row = st.insert(block, t.data[block]["columns"], values, me.id)
        added.append((block, row["id"]))
        return row

    def types():
        rows = st.rows("doc_types", t.data["doc_types"]["columns"], limit=2000)["rows"]
        rows = [r for r in rows if r.get("shown") is not False]
        return sorted(rows, key=lambda x: (x.get("order") if isinstance(x.get("order"), int) else 100, str(x.get("name") or "").lower()))

    add.types = types
    add.tool = t
    yield add
    for block, rid in reversed(added):
        st.delete(block, t.data[block]["columns"], rid)
        shutil.rmtree(FILES.row_dir(SLUG, block, rid), ignore_errors=True)


def _client(owner, email, name, *, only_theirs=False, grant=True):
    tools = [x["id"] for x in owner.get("/api/users").json()["tools"] if x["slug"] == "doc-watch"] if grant else []
    pw = _make(owner, email, name, "user", tools=tools, team="EC Export Team")
    if only_theirs:
        from cbx.db import SessionLocal

        with SessionLocal() as s:
            u = s.query(User).filter(User.email == email).one()
            ta = s.query(ToolAccess).join(Tool, Tool.id == ToolAccess.tool_id).filter(Tool.slug == "doc-watch", ToolAccess.user_id == u.id).one()
            ta.scope = "granted"  # "only theirs": the companies open to everyone are not open to them
            s.commit()
    c = TestClient(app)
    login(c, email, pw)
    return c


def _book(client):
    r = client.get("/api/ec/export")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert f'filename="Entity Compliance {local_today().isoformat()}.xlsx"' in r.headers["content-disposition"]
    return load_workbook(io.BytesIO(r.content))


def _rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _row(ws, first):
    for r in ws.iter_rows():
        if r[0].value == first:
            return r
    raise AssertionError(f"{first} is not on {ws.title}")


def _ensure_types(ec, n=4):
    have = ec.types()
    for i in range(len(have), n):
        ec(f"doc_types", name=f"EXP Type {i}", short=f"EX{i}", order=900 + i, warn_days=90, shown=True)
    return [(x["id"], x["name"]) for x in ec.types()]


def test_the_page_has_the_button_and_the_folder_is_sound(owner):
    t = loader.get(SLUG)
    assert t.ok, t.errors
    w = next(x for x in loader.widgets_of("grid", t.pages["grid"]) if x["id"] == "w_export")
    assert w["type"] == "links" and w["items"][0]["url"] == "/api/ec/export"
    p = owner.get(f"/api/engine/tools/{SLUG}/pages/grid").json()
    assert p["strings"]["act.export"] == "Export"
    assert not loader.DOWNLOAD_RE.match("/api/../secret") and not loader.DOWNLOAD_RE.match("/api/ec/export?x=1") and loader.DOWNLOAD_RE.match("/api/ec/export")


def test_sheets_headers_and_the_header_style(owner, ec):
    names = [n for _, n in _ensure_types(ec)]
    ec("companies", name="EXP Shape Co", visible=OPEN)
    wb = _book(owner)
    assert wb.sheetnames[:3] == ["Overview", "Documents", "Expiring and expired"] and wb.sheetnames[-2:] == ["Document types", "About"]
    ov = wb["Overview"]
    head = [c.value for c in ov[1]]
    assert head[:9] == ["Entity", "Responsible person", "Responsible email", "Valid", "Expiring soon", "Expired", "Missing", "Not required", "Overall"]
    assert head[9:] == names
    for c in ov[1]:
        assert c.fill.fgColor.rgb.endswith("003F2D") and c.font.bold and c.font.color.rgb.endswith("FFFFFF")
    assert ov.freeze_panes == "B2" and ov.auto_filter.ref.startswith("A1:")
    assert [c.value for c in wb["Documents"][1]] == ["Entity", "Document type", "Title", "Expiry date", "Days left", "State", "Required", "No expiry",
                                                     "File name", "Uploaded by", "Uploaded at", "Notes"]
    assert [c.value for c in wb["Expiring and expired"][1]] == ["Entity", "Document type", "Title", "Expiry date", "Days left", "State",
                                                                "Responsible person", "Responsible email"]
    assert [c.value for c in wb["Entities"][1]] == ["Entity", "Responsible person", "Responsible account", "Responsible email", "Who sees it",
                                                    "Given access", "Notes", "Created"]
    assert [c.value for c in wb["Document types"][1]] == ["Name", "Short name", "Warn (days before)", "Order", "Active"]
    about = {r[0]: r[1] for r in _rows(wb["About"]) if r and r[0]}
    assert about["Exported by"].endswith("(owner@test.local)") and isinstance(about["Exported at"], datetime) and about["Time zone"] == "Africa/Johannesburg"


def test_a_person_given_two_companies_gets_only_those_two(owner, ec, db):
    a = ec("companies", name="EXP Mozambique Lda", visible=CLOSED)
    b = ec("companies", name="EXP Zambia Ltd", visible=CLOSED)
    ec("companies", name="EXP Botswana (Pty) Ltd", visible=CLOSED)
    ec("companies", name="EXP Open Co", visible=OPEN)
    narrow = _client(owner, "exp-narrow@test.local", "Nandi Narrow", only_theirs=True)
    uid = db.query(User).filter(User.email == "exp-narrow@test.local").one().id
    for c in (a, b):
        ec("company_people", company=c["id"], person=uid)
    wb = _book(narrow)
    names = sorted(r[0] for r in _rows(wb["Overview"])[1:])
    assert names == ["EXP Mozambique Lda", "EXP Zambia Ltd"]
    assert {r[0] for r in _rows(wb["Documents"])[1:]} <= set(names)
    assert sorted(r[0] for r in _rows(wb["Entities"])[1:]) == names
    assert "Given access" not in [c.value for c in wb["Entities"][1]]  # who was given access: admins only
    assert "Archived entities" not in wb.sheetnames
    # without "only theirs", the open company comes too — never the third closed one
    wide = _client(owner, "exp-wide@test.local", "Wanda Wide")
    names = {r[0] for r in _rows(_book(wide)["Overview"])[1:]}
    assert "EXP Open Co" in names and "EXP Botswana (Pty) Ltd" not in names and "EXP Mozambique Lda" not in names


def test_each_state_has_its_word_and_its_colour(owner, ec):
    today = local_today()
    (t1, n1), (t2, n2), (t3, n3), (t4, n4) = _ensure_types(ec)[:4]
    c = ec("companies", name="EXP States Co", responsible="Sipho Mokoena", responsible_email="sipho@client.example", visible=OPEN)
    old = ec("documents", company=c["id"], doc_type=t1, title=n1, expiry=(today - timedelta(days=10)).isoformat(), warn_days=90)
    up = owner.post(f"/api/engine/tools/{SLUG}/data/documents/{old['id']}/files/file", files=[("files", ("fire-2025.pdf", certificate(), "application/pdf"))])
    assert up.status_code == 200, up.text
    ec("documents", company=c["id"], doc_type=t2, title=n2, expiry=(today + timedelta(days=23)).isoformat(), warn_days=90)
    ec("documents", company=c["id"], doc_type=t3, title=n3, expiry=(today + timedelta(days=400)).isoformat(), warn_days=90)
    ec("documents", company=c["id"], title="Lease agreement", no_expiry=True)
    wb = _book(owner)
    ov = wb["Overview"]
    head = [x.value for x in ov[1]]
    row = _row(ov, "EXP States Co")
    cell = {head[i]: row[i] for i in range(len(head))}
    d = lambda x: f"{x.day} {x:%b} {x.year}"  # noqa: E731
    assert cell[n1].value == f"Expired {d(today - timedelta(days=10))}" and cell[n1].fill.fgColor.rgb.endswith(EX.STATE_FILL["expired"][0])
    assert cell[n2].value == "Expires in 23 days" and cell[n2].fill.fgColor.rgb.endswith(EX.STATE_FILL["expiring"][0])
    assert cell[n3].value == f"Valid until {d(today + timedelta(days=400))}" and cell[n3].fill.fgColor.rgb.endswith(EX.STATE_FILL["ok"][0])
    assert cell[n4].value == "Missing" and cell[n4].fill.fgColor.rgb.endswith(EX.STATE_FILL["missing"][0])
    assert cell["Overall"].value == "Expired" and cell["Overall"].fill.fgColor.rgb.endswith(EX.STATE_FILL["expired"][0])
    assert (cell["Valid"].value, cell["Expiring soon"].value, cell["Expired"].value) == (2, 1, 1)
    docs = [r for r in wb["Documents"].iter_rows(min_row=2) if r[0].value == "EXP States Co"]
    by = {r[1].value if r[1].value != "Other" else r[2].value: r for r in docs}
    assert by[n1][3].value.date() == today - timedelta(days=10) and by[n1][3].number_format == "dd mmm yyyy" and by[n1][4].value == -10
    assert by[n1][8].value == "fire-2025.pdf" and by[n1][9].value == "Test Owner" and isinstance(by[n1][10].value, datetime)
    assert by[n2][4].value == 23 and by["Lease agreement"][5].value == "Valid, doesn't expire" and by[n4][5].value == "Missing"
    todo = [r for r in _rows(wb["Expiring and expired"])[1:] if r[0] == "EXP States Co"]
    assert [r[1] for r in todo][:3] == [n1, n2, n4] and n3 not in [r[1] for r in todo]
    assert todo[0][6] == "Sipho Mokoena" and todo[0][7] == "sipho@client.example"


def test_archived_companies_are_on_their_own_sheet_for_admins_only(owner, ec):
    ec("companies", name="EXP Closed Co", visible=OPEN, archived=True, archived_reason="contract ended")
    wb = _book(owner)
    assert "EXP Closed Co" not in {r[0] for r in _rows(wb["Overview"])} | {r[0] for r in _rows(wb["Entities"])}
    assert _row(wb["Archived entities"], "EXP Closed Co")[4].value == "contract ended"
    plain = _client(owner, "exp-plain@test.local", "Pieter Plain")
    wb2 = _book(plain)
    assert "Archived entities" not in wb2.sheetnames
    assert all("EXP Closed Co" not in [r[0] for r in _rows(ws)] for ws in wb2.worksheets)


def test_a_formula_looking_name_stays_text(owner, ec):
    ec("companies", name="=HYPERLINK(\"http://evil.example\",\"x\")", visible=OPEN)
    wb = _book(owner)
    for ws in (wb["Overview"], wb["Entities"]):
        assert _row(ws, "=HYPERLINK(\"http://evil.example\",\"x\")")[0].data_type == "s"


def test_strangers_are_refused_and_the_export_is_audited(owner, ec, db):
    ec("companies", name="EXP Audit Co", visible=OPEN)
    nobody = _client(owner, "exp-nobody@test.local", "No Grant", grant=False)
    assert nobody.get("/api/ec/export").status_code == 403
    _book(owner)
    db.expire_all()
    row = db.query(AuditLog).filter(AuditLog.action == "ec_export").order_by(AuditLog.id.desc()).first()
    assert row is not None and " entities, " in row.detail and "EXP Audit Co" not in row.detail  # counts only
