"""Made-up IFRS 16 source files rich enough to make the engine write import workbooks: a Contract Movement tab
(additions not on LA, a renewal and a removal that are), its Step Payment tab, the LA register and a Sequentra
validation export. Every name and number is invented."""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook


def _save(wb, path):
    wb.save(path)
    return path


def make(out: Path) -> dict:
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    la = Workbook()
    ws = la.active
    ws.title = "Portfolio Transactions"
    ws.append(["Landlord Advantage - Portfolio Transactions"])
    ws.append([])
    ws.append(["Transaction #", "Company Name", "Currency", "Current Payment (Transactional)"])
    for ref, name, pay in (("TXN10001", "Rosebank Branch", 18000), ("TXN10002", "Durban North Branch", 22000), ("TXN10003", "Menlyn Kiosk", 9000)):
        ws.append([ref, name, "ZAR", pay])

    mv = Workbook()
    ws = mv.active
    ws.title = "Contract Movement"
    ws.append(["RR06 Contract Movement export"])
    ws.append([])
    ws.append(["Lease Reference", "Move Type", "DCF date", "Current Start Date", "Lease Expiry Date", "Termination Date", "Property Ref",
               "One View Ref", "Landlord / Tenant Name", "Contract Cost Centre", "Comments"])
    rows = [
        ("TXN20001", "Addition", datetime(2026, 8, 12), datetime(2026, 8, 1), datetime(2031, 7, 31), None, "PR-501", "OV-501", "Invented Properties (Pty) Ltd", "CC501", "New branch"),
        ("TXN20002", "Addition", datetime(2026, 8, 18), datetime(2026, 9, 1), datetime(2029, 8, 31), None, "PR-502", "OV-502", "Made Up Landlords Ltd", "CC502", ""),
        ("TXN10002", "Renewal", datetime(2026, 8, 20), datetime(2026, 8, 1), datetime(2030, 7, 31), None, "PR-202", "OV-202", "Sample Holdings", "CC202", "Renewed for 4 years"),
        ("TXN10003", "Removal", datetime(2026, 8, 25), datetime(2021, 3, 1), datetime(2026, 9, 30), datetime(2026, 8, 31), "PR-303", "OV-303", "Fictional Trust", "CC303", "Closing early"),
        ("TXN20009", "Addition", datetime(2026, 7, 10), datetime(2026, 7, 1), datetime(2028, 6, 30), None, "PR-509", "OV-509", "Last Month Ltd", "CC509", ""),
    ]
    for r in rows:
        ws.append(list(r))
    sp = mv.create_sheet("Step Payment")
    sp.append(["Schedule Number", "Expense Type Name", "StartDate", "EndDate", "Number of Payments", "Payment Amount", "Starting Payment Number", "Payment Frequency"])
    sp.append(["TXN20001", "Base Rent", datetime(2026, 8, 1), datetime(2027, 7, 31), 12, 25000, 1, "Monthly"])
    sp.append(["TXN20001", "Base Rent", datetime(2027, 8, 1), datetime(2031, 7, 31), 48, 27000, 13, "Monthly"])
    sp.append(["TXN20002", "Base Rent", datetime(2026, 9, 1), datetime(2029, 8, 31), 36, 14000, 1, "Monthly"])
    sp.append(["TXN10002", "Base Rent", datetime(2026, 8, 1), datetime(2030, 7, 31), 48, 23500, 1, "Monthly"])
    sp.append(["TXN20009", "Base Rent", datetime(2026, 7, 1), datetime(2028, 6, 30), 24, 8000, 1, "Monthly"])

    sq = Workbook()
    ws = sq.active
    ws.title = "Validation - Record Info"
    ws.append(["Record ID", "Record Name", "Country", "Currency", "City", "Address 1", "Address 2", "Postal Code", "Current Term Start Date",
               "Current Lease End Date", "Move-In Date", "Discount Rate", "New Discount Rate", "External Reference", "Contract Cost Centre", "Title Description"])
    for ref, name, city, start, end, rate in (("TXN20001", "Parkhurst Branch", "Johannesburg", datetime(2026, 8, 1), datetime(2031, 7, 31), 10.75),
                                              ("TXN20002", "Umhlanga Kiosk", "Durban", datetime(2026, 9, 1), datetime(2029, 8, 31), 10.5),
                                              ("TXN10002", "Durban North Branch", "Durban", datetime(2026, 8, 1), datetime(2030, 7, 31), 10.25),
                                              ("TXN10003", "Menlyn Kiosk", "Pretoria", datetime(2021, 3, 1), datetime(2026, 9, 30), 9.5),
                                              ("TXN20009", "Last Month Branch", "Cape Town", datetime(2026, 7, 1), datetime(2028, 6, 30), 10.0)):
        ws.append([ref, name, "South Africa", "ZAR", city, "1 Invented Road", "Suite 2", "2000", start, end, start, rate, None, "EXT-" + ref, "CC-" + ref[-3:], "Lease"])
    return {"la": _save(la, out / "la_register_rich.xlsx"), "movements": _save(mv, out / "rr06_contract_movement.xlsx"), "sequentra": _save(sq, out / "sequentra_validation.xlsx")}


if __name__ == "__main__":
    import sys

    print(make(Path(sys.argv[1] if len(sys.argv) > 1 else ".")))
