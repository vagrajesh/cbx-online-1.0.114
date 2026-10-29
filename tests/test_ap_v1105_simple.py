"""1.0.105 (SPEC-1105, Builder S): the AP Tool's one screen, "Make SAP upload".

A run of mixed invoices ends with one download holding the green and yellow ones and leaving out the duplicate, the
invalid tax invoice and the one waiting for a second person's check; the download no longer refuses while a check
waits (the waiting ones are simply not in the file); Claim / Don't claim and "Check it" from the screen; "Download again
(n more)" once they are settled; the recent runs' record of what was sent (never a second uploadable copy); the company
picker for "company unknown"; the status in a few reads however many wait; the More link only for admins and role holders; the guessed cells
yellow with their note in the file; the AI off. Invoices, vendors and people are made up."""
from __future__ import annotations

import io
import random
import uuid
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from cbx.ap import capture as C
from cbx.ap import masters as M
from cbx.ap import portfolio as PF
from cbx.ap import sap as P
from cbx.ap import simple as SIMPLE
from cbx.ap import history as H
from cbx.db import SessionLocal
from cbx.models import AuditLog
from cbx.models.ap import ApCapture, ApEntity, ApPosted, ApRole, ApVendorCode

from .fixtures import make_ap
from .test_ap_capture_1098 import _admin, _app, _chart_of_accounts, _invoice_answer, _user, _vision, signed_in  # noqa: F401

VAT1, VENDOR1, SUP1 = "4555000222", "300561", "Simpleco Services (Pty) Ltd"
VAT2, VENDOR2, SUP2 = "4555000333", "300562", "Learnt Supplies (Pty) Ltd"


@pytest.fixture(autouse=True)
def _clean():
    yield
    with SessionLocal() as s:
        s.query(ApRole).delete()
        s.query(ApPosted).filter(ApPosted.doc_number.like("S5TEST%")).delete(synchronize_session=False)
        s.commit()


def _vendors():
    with SessionLocal() as s:
        s.query(ApVendorCode).filter(ApVendorCode.vat_no.in_((VAT1, VAT2))).delete(synchronize_session=False)
        s.add(ApVendorCode(vendor_no=VENDOR1, name=SUP1, name_key=M.name_key(SUP1), vat_no=VAT1, gl="6100100", cost_centre="CC77", tax_code="I9", source="excel"))
        # a coding a person added and nobody confirmed yet: its invoices are yellow
        s.add(ApVendorCode(vendor_no=VENDOR2, name=SUP2, name_key=M.name_key(SUP2), vat_no=VAT2, gl="6100100", cost_centre="CC77", tax_code="I9", source="person",
                           learned=True, learned_by="Someone Else"))
        x143 = s.query(ApEntity).filter(ApEntity.code == "X143").one()
        if "CBRE Excellerate Facilities Management" not in x143.aliases:
            x143.aliases = (x143.aliases + "\n" if x143.aliases else "") + "CBRE Excellerate Facilities Management"
        s.commit()


def _posted(vendor: str, ref: str = "OLD-1", amount: float = 100.0) -> None:
    with SessionLocal() as s:
        s.add(ApPosted(company_code="X143", vendor=vendor, reference=ref, ref_key=H.ref_key(ref), doc_type="KR", doc_number=f"S5TEST{random.randint(10**8, 10**9)}",
                       posting_date="2026-08-17", amount=amount, source="excel"))
        s.commit()


def _answer(no, supplier=SUP1, vat_no=VAT1, excl="1000.00", vat="150.00", incl="1150.00", bill_to_vat="4999888777"):
    a = _invoice_answer(no, supplier=supplier, vat_no=vat_no, excl=excl, vat=vat, incl=incl)
    a["bill_to_vat"], a["tax_invoice_words"] = bill_to_vat, True
    return a


def _rhea():
    return signed_in("v1105.rev@test.local", "user", name="Rhea Reviewer", tools=["ap"])


def _start(c, files: dict) -> dict:
    r = c.post("/api/ap/simple/runs", files=[("files", (n, d, "application/pdf")) for n, d in files.items()])
    assert r.status_code == 200, r.text
    return r.json()


def _status(c, bid) -> dict:
    r = c.get(f"/api/ap/simple/runs/{bid}")
    assert r.status_code == 200, r.text
    return r.json()


def _refs(content: bytes) -> set:
    ws = load_workbook(io.BytesIO(content)).active
    heads = [x.value for x in ws[1]]
    return {row[heads.index("refdocnr")] for row in ws.iter_rows(min_row=2, values_only=True) if row[heads.index("pkey")] in ("31", "21")}


def test_a_run_downloads_green_and_yellow_and_leaves_out_the_rest(monkeypatch, ap_review_gates):
    _vendors()
    _posted(VENDOR1)  # SAP posted this vendor before: not a first invoice
    _posted(VENDOR2)
    monkeypatch.setitem(PF.DEFAULTS, "review_sample_pct", 0)
    run = uuid.uuid4().hex[:6].upper()
    no = {k: f"S5-{run}-{i}" for i, k in enumerate(("green", "yellow", "dup", "badvat", "big"))}
    _posted(VENDOR1, no["dup"], 1150.0)  # the duplicate: SAP posted this very invoice
    answers = {"green.pdf": _answer(no["green"]), "yellow.pdf": _answer(no["yellow"], supplier=SUP2, vat_no=VAT2), "dup.pdf": _answer(no["dup"]),
               "badvat.pdf": _answer(no["badvat"], excl="10000.00", vat="1500.00", incl="11500.00", bill_to_vat=None),
               "big.pdf": _answer(no["big"], excl="60000.00", vat="9000.00", incl="69000.00")}
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision(answers))
    cara, rhea = _user(), _rhea()
    pdfs = {n: make_ap.invoice(no=a["invoice_no"], supplier=a["supplier"], vat_no=a["supplier_vat"]) for n, a in answers.items()}
    st = _start(cara, pdfs)
    bid = st["run"]["id"]
    assert st["run"]["name"].startswith("Run ") and st["phase"] == "done", st
    st = _status(cara, bid)
    left = {x["file"]: x for x in st["left_out"]}
    assert set(left) == {"yellow.pdf", "dup.pdf", "badvat.pdf", "big.pdf"}, st["docs"]
    # one plain line each; the person who dropped them may not decide or check: told who it waits for
    assert left["dup.pdf"]["kind"] == "duplicate" and "already posted in SAP" in left["dup.pdf"]["note"]
    assert left["badvat.pdf"]["kind"] == "vat" and "not a valid tax invoice" in left["badvat.pdf"]["note"] and not left["badvat.pdf"]["may"]
    assert left["badvat.pdf"]["waits"].startswith("waits for ") and left["badvat.pdf"]["waits"].endswith(" to decide on the VAT") and "Cara Person" not in left["badvat.pdf"]["waits"]
    assert left["big.pdf"]["kind"] == "review" and "second person" in left["big.pdf"]["note"] and not left["big.pdf"]["may"]
    assert left["big.pdf"]["waits"].startswith("waits for ") and left["big.pdf"]["waits"].endswith(" to check it") and "Cara Person" not in left["big.pdf"]["waits"]
    with SessionLocal() as s:  # who may check it: reviewers in its portfolio, never the person who dropped it
        names = SIMPLE.reviewers_for(s, s.get(ApCapture, left["big.pdf"]["id"]), limit=1000)
        assert "Rhea Reviewer" in names and "Cara Person" not in names
    assert left["yellow.pdf"]["kind"] == "review"  # 1.0.104 V05 picks every yellow: it waits for its check too
    assert st["summary"]["line"] == "5 invoices · 1 ready · 4 left out" and st["download"]["label"] == "Download SAP upload"
    assert {d["file"]: d["word"] for d in st["docs"]} == {"green.pdf": "Ready", "yellow.pdf": "Left out", "dup.pdf": "Left out", "badvat.pdf": "Left out", "big.pdf": "Left out"}
    # the reviewer sees the same run with her buttons
    mine = {x["file"]: x for x in _status(rhea, bid)["left_out"]}
    assert mine["big.pdf"]["may"] and mine["yellow.pdf"]["may"] and mine["badvat.pdf"]["may"]
    # "Check it" on the yellow one: the page beside the fields; Agree
    one = rhea.get(f"/api/ap/review/{left['yellow.pdf']['id']}").json()
    assert one["pages"] >= 1 and one["editable"]["invoice_no"] == no["yellow"]
    assert rhea.post(f"/api/ap/review/{left['yellow.pdf']['id']}", json={"action": "agree"}).status_code == 200
    st = _status(cara, bid)
    assert st["summary"]["line"] == "5 invoices · 1 ready · 1 with guesses (yellow) · 3 left out"
    # the download: green and yellow, never the duplicate, the invalid tax invoice nor the one still waiting for its check —
    # and no refusal while it waits (1.0.104 refused the whole download)
    r = cara.post(f"/api/ap/simple/runs/{bid}/download")
    assert r.status_code == 200, r.text
    assert _refs(r.content) == {no["green"], no["yellow"]}
    with SessionLocal() as s:
        log = s.query(AuditLog).filter(AuditLog.action == "ap_sap_upload", AuditLog.target == f"ap_batch:{bid}").one()
        assert "left out: waiting for review" in log.detail
        assert s.get(ApCapture, left["big.pdf"]["id"]).uploaded_at is None
    st = _status(cara, bid)
    assert not st["download"]["can"] and st["download"]["sent"] == 2
    # Don't claim from the screen (a reason is needed; the dropper can't), then the second person's checks
    cid = left["badvat.pdf"]["id"]
    assert cara.post(f"/api/ap/captures/{cid}/vat", json={"decision": "no_claim", "reason": "supplier to send a valid one"}).status_code == 403
    assert rhea.post(f"/api/ap/captures/{cid}/vat", json={"decision": "no_claim", "reason": ""}).status_code == 400
    assert rhea.post(f"/api/ap/captures/{cid}/vat", json={"decision": "no_claim", "reason": "supplier to send a valid one"}).status_code == 200
    mine = {x["file"]: x for x in _status(rhea, bid)["left_out"]}
    assert mine["badvat.pdf"]["kind"] == "review" and "not a valid tax invoice" in mine["badvat.pdf"]["note"]  # V05 picks it now
    for f in ("badvat.pdf", "big.pdf"):
        assert rhea.post(f"/api/ap/review/{mine[f]['id']}", json={"action": "agree"}).status_code == 200
    st = _status(cara, bid)
    assert st["download"]["label"] == "Download again (2 more)" and [x["file"] for x in st["left_out"]] == ["dup.pdf"]
    r = cara.post(f"/api/ap/simple/runs/{bid}/download")
    assert r.status_code == 200 and _refs(r.content) == {no["badvat"], no["big"]}
    assert cara.post(f"/api/ap/simple/runs/{bid}/download").status_code == 400  # nothing more: never twice
    # recent runs: the run, and the record of what was sent — never a file SAP could take again (1.0.105 review: a second
    # uploadable copy is a double posting one click away)
    home = cara.get("/api/ap/simple").json()
    assert any(x["id"] == bid and x["files"] == 5 and x["sent"] == 4 for x in home["recent"])
    assert cara.get(f"/api/ap/simple/runs/{bid}/copy").status_code in (404, 405)
    rec = cara.get(f"/api/ap/simple/runs/{bid}/sent")
    assert rec.status_code == 200 and "_record" in rec.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(rec.content))
    first = wb.worksheets[0]
    assert wb.sheetnames[0] == SIMPLE.SENT_SHEET and wb.active.title == SIMPLE.SENT_SHEET
    assert "TransCode" not in [x.value for x in first[1]] and "posts them twice" in first["A1"].value  # SAP's uploader finds no SAP columns on it
    sent_rows = [r for r in first.iter_rows(min_row=4, values_only=True) if r[0]]
    assert {r[2] for r in sent_rows} == {no["green"], no["yellow"], no["badvat"], no["big"]} and {r[1] for r in sent_rows} == {"Cara Person"}
    assert "SAP rows" in wb.sheetnames  # the review workbook behind it
    wb = cara.get(f"/api/ap/simple/runs/{bid}/review")
    assert wb.status_code == 200 and wb.content[:2] == b"PK"


def test_the_old_download_leaves_the_waiting_ones_out_too(monkeypatch, ap_review_gates):
    """/sap/upload (the More → Inbox page) behaves the same: the documents waiting for review stay out, the rest go."""
    _vendors()
    _posted(VENDOR1)
    monkeypatch.setitem(PF.DEFAULTS, "review_sample_pct", 0)
    run = uuid.uuid4().hex[:6].upper()
    small, big = f"S5O-{run}-1", f"S5O-{run}-2"
    answers = {"small.pdf": _answer(small), "big.pdf": _answer(big, excl="60000.00", vat="9000.00", incl="69000.00")}
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision(answers))
    cara = _user()
    bid = _start(cara, {n: make_ap.invoice(no=a["invoice_no"], supplier=SUP1, vat_no=VAT1) for n, a in answers.items()})["run"]["id"]
    r = cara.post(f"/api/ap/batches/{bid}/sap/upload", json={"include_yellow": True})
    assert r.status_code == 200 and _refs(r.content) == {small}
    r = cara.post(f"/api/ap/batches/{bid}/sap/upload", json={"include_yellow": True})
    assert r.status_code == 400 and r.json()["detail"].startswith("1 document waits for review")  # only the waiting one is left: the 1.0.104 words


def test_more_link_only_for_admins_and_role_holders():
    cara, adm, rhea = _user(), _admin(), _rhea()
    assert cara.get("/api/ap/simple").json()["more"] is False  # an everyday user: one screen
    assert adm.get("/api/ap/simple").json()["more"] is True
    assert rhea.get("/api/ap/simple").json()["more"] is False
    with SessionLocal() as s:
        from cbx.models import User

        rid = s.query(User).filter(User.email == "v1105.rev@test.local").one().id
    assert adm.put(f"/api/ap/roles/{rid}", json={"portfolio": "", "roles": ["reviewer"]}).status_code == 200
    assert rhea.get("/api/ap/simple").json()["more"] is True
    viewer = signed_in("v1105.viewer@test.local", "viewer", name="Vee Viewer", tools=["ap"])
    v = viewer.get("/api/ap/simple").json()
    assert v["can_write"] is False and v["more"] is False
    assert viewer.post("/api/ap/simple/runs", files=[("files", ("a.pdf", make_ap.invoice(), "application/pdf"))]).status_code == 403


def test_with_the_ai_off_the_run_finishes_and_says_so():
    cara = _user()
    st = _start(cara, {"later.pdf": make_ap.invoice(no=f"S5AI-{uuid.uuid4().hex[:6]}", supplier=SUP1, vat_no=VAT1)})
    if st["phase"] == "reading":
        st = _status(cara, st["run"]["id"])
    assert st["phase"] == "done" and st["ai_off"] == 1 and st["ai_line"] == "The AI is off: 1 invoice will be read when it is back."
    assert st["left_out"] == [] and not st["download"]["can"] and st["summary"]["line"].startswith("1 invoice · 0 ready")


def test_guessed_cells_are_yellow_with_a_note_in_the_upload():
    c = SimpleNamespace(guesses={"gl": "the model: cleaning services → 6100100", "due_date": {"how": "invoice date + 30 days", "sure": False},
                                 "company": {"how": "VAT number", "sure": True}})
    rows = [{"TransCode": "FB01", "pkey": "31", "account": "300561", "amount": "1150.00", "comp": "Z999", "DueDate": "18112026", "refdocnr": "X1"},
            {"TransCode": "ZB01", "pkey": "40", "account": "6100100", "amount": "1000.00", "tc": "I9", "costcentre": "NV0001"},
            {"TransCode": "ZB01", "pkey": "40", "account": "JA600", "amount": "150.00", "dt": "X"}]
    marks = [SIMPLE.cell_marks(c, r) for r in rows]
    assert marks == [{"_guess": {"DueDate": "invoice date + 30 days"}}, {"_guess": {"account": "the model: cleaning services → 6100100"}}, {}]
    ws = load_workbook(io.BytesIO(P.upload_file([{**r, **m, "_tess": "yellow"} for r, m in zip(rows, marks)]))).active
    heads = [x.value for x in ws[1]]
    due, gl, vat_gl = ws.cell(row=2, column=heads.index("DueDate") + 1), ws.cell(row=3, column=heads.index("account") + 1), ws.cell(row=4, column=heads.index("account") + 1)
    assert due.comment and "invoice date + 30 days" in due.comment.text and due.fill.fgColor.rgb.endswith(P.GUESS_FILL)
    assert gl.comment and "the model" in gl.comment.text and vat_gl.comment is None
    assert ws.cell(row=2, column=1).fill.fgColor.rgb.endswith(P.STATUS_FILL["yellow"])  # Tess's colour stays on the TransCode cell
    assert SIMPLE.guesses_of(SimpleNamespace(guesses=[{"field": "vendor", "how": "most recent of 6", "sure": False}, {"field": "company", "how": "VAT", "sure": True}])) == {"vendor": "most recent of 6"}
    assert SIMPLE.cell_marks(SimpleNamespace(), rows[0]) == {}
    # Builder G's autofill keeps every guess in match["autofill"] and says which cells carry it (row_guesses)
    g = SimpleNamespace(match={"autofill": {"gl": {"value": "6100100", "how": "the model", "sure": False}, "company": {"value": "Z999", "how": "VAT number", "sure": True}}},
                        guesses={"gl": "the model"})
    assert SIMPLE.cell_marks(g, rows[1])["_guess"]["account"] == "the model" and SIMPLE.cell_marks(g, rows[0]) == {} and SIMPLE.cell_marks(g, rows[2]) == {}


def test_a_run_is_named_by_the_day_and_time():
    from datetime import datetime, timezone

    assert SIMPLE.run_name(datetime(2026, 10, 19, 12, 5, tzinfo=timezone.utc)) == "Run 19 Oct, 14:05"


def test_the_corrected_excel_goes_to_builder_g_s_learn_from_upload(monkeypatch):
    """"Drop the corrected Excel here to teach CBX": the file reaches autofill.learn_from_upload with the person; its words
    come back; its refusal is a 400; a file that is not Excel is refused before."""
    import importlib
    import sys
    import types

    try:
        AF = importlib.import_module("cbx.ap.autofill")
    except ImportError:  # Builder G's module not in this tree yet: a stand-in with the agreed shape
        AF = types.ModuleType("cbx.ap.autofill")

        class Refused(Exception):
            pass

        AF.Refused = Refused
        monkeypatch.setitem(sys.modules, "cbx.ap.autofill", AF)
    seen = {}

    def learn(db, user, data):
        seen["who"], seen["n"] = user.name, len(data)
        if data == b"PKbad":
            raise AF.Refused("This is not the SAP upload: no TransCode column.")
        return {"learnt": 7, "documents": 7, "updated": 2, "text": "Learnt 7 codings from your file."}

    monkeypatch.setattr(AF, "learn_from_upload", learn, raising=False)
    cara = _user()
    r = cara.post("/api/ap/simple/learn", files={"file": ("SAP_upload_fixed.xlsx", b"PKgood", "application/octet-stream")})
    assert r.status_code == 200 and r.json()["text"] == "Learnt 7 codings from your file." and seen == {"who": "Cara Person", "n": 6}
    r = cara.post("/api/ap/simple/learn", files={"file": ("SAP_upload_fixed.xlsx", b"PKbad", "application/octet-stream")})
    assert r.status_code == 400 and "no TransCode column" in r.json()["detail"]
    assert cara.post("/api/ap/simple/learn", files={"file": ("notes.pdf", b"%PDF", "application/pdf")}).status_code == 400


def test_the_status_reads_the_same_few_times_however_many_wait(monkeypatch, ap_review_gates):
    """1.0.105 review (perf): who each waiting invoice waits for is worked out in bulk once the run is read — not ~16 reads
    per invoice on every poll."""
    from sqlalchemy import event

    from cbx.db import engine

    _vendors()
    _posted(VENDOR1)
    monkeypatch.setitem(PF.DEFAULTS, "review_sample_pct", 0)
    monkeypatch.setattr(C, "model_on", lambda: True)
    cara = _user()
    _rhea()
    counts = {}
    for n in (1, 6):
        run = uuid.uuid4().hex[:6].upper()
        answers = {f"big{i}.pdf": _answer(f"S5P-{run}-{i}", excl="60000.00", vat="9000.00", incl="69000.00") for i in range(n)}
        monkeypatch.setattr(C, "ask_vision", _vision(answers))
        bid = _start(cara, {k: make_ap.invoice(no=a["invoice_no"], supplier=SUP1, vat_no=VAT1) for k, a in answers.items()})["run"]["id"]
        seen = {"q": 0}

        def count(*_a, **_k):
            seen["q"] += 1

        event.listen(engine, "before_cursor_execute", count)
        try:
            st = _status(cara, bid)
        finally:
            event.remove(engine, "before_cursor_execute", count)
        assert len(st["left_out"]) == n and all(x["kind"] == "review" and x["waits"].startswith("waits for ") for x in st["left_out"]), st["left_out"]
        counts[n] = seen["q"]
    assert counts[6] - counts[1] <= 3 and counts[6] < 60, counts


def test_company_unknown_is_picked_on_the_screen(monkeypatch):
    """Builder G's "company unknown" (red, needs "entity"): a capturer picks the company in Left out (the group's own
    companies first); it goes through capture.decide "manual" and Mia matches it again. A viewer can't."""
    from cbx.ap import autofill as AF

    _vendors()
    _posted(VENDOR1)
    monkeypatch.setattr(C, "model_on", lambda: True)
    run = uuid.uuid4().hex[:6].upper()
    answers = {"lost.pdf": _answer(f"S5C-{run}")}
    monkeypatch.setattr(C, "ask_vision", _vision(answers))
    cara = _user()
    bid = _start(cara, {"lost.pdf": make_ap.invoice(no=f"S5C-{run}", supplier=SUP1, vat_no=VAT1)})["run"]["id"]
    with SessionLocal() as s:  # as Mia leaves a bill-to no company could be guessed for (cbx/ap/capture.py match_one)
        c = s.query(ApCapture).filter(ApCapture.batch_id == bid).one()
        c.fields = {**(c.fields or {}), "bill_to": "Somebody Unknown Trading", "bill_to_vat": ""}
        c.match = {"entity": {"found": False, "bill_to": "Somebody Unknown Trading"}}
        c.bucket, c.stage, c.needs, c.sap_status, c.sap_issues = "other", "needs", "entity", "red", [{"level": "red", "text": AF.UNKNOWN_COMPANY}]
        s.commit()
        cid = c.id
    left = _status(cara, bid)["left_out"]
    assert [(x["kind"], x["may"]) for x in left] == [("company", True)] and "company unknown" in left[0]["note"] and "Somebody Unknown Trading" in left[0]["note"]
    viewer = signed_in("v1105.viewer@test.local", "viewer", name="Vee Viewer", tools=["ap"])
    v = _status(viewer, bid)["left_out"][0]
    assert v["may"] is False and v["waits"].startswith("waits for a capturer")
    cos = cara.get("/api/ap/simple/companies").json()
    assert {"Z999", "Z102"} <= {x["code"] for x in cos["group"]} and "X143" in {x["code"] for x in cos["other"]}
    assert not {x["code"] for x in cos["group"]} & {x["code"] for x in cos["other"]}
    r = cara.post(f"/api/ap/captures/{cid}/decide", json={"action": "manual", "data": {"company_code": "Z999"}})
    assert r.status_code == 200, r.text
    with SessionLocal() as s:
        c = s.get(ApCapture, cid)
        assert c.company_code == "Z999" and (c.match or {}).get("entity", {}).get("how") == "a person" and c.needs == "" and c.stage != "needs"
        assert s.query(AuditLog).filter(AuditLog.action == "ap_capture_manual", AuditLog.target == f"ap_item:{c.item_id}").count() == 1
    assert not [x for x in _status(cara, bid)["left_out"] if x["kind"] == "company"]
