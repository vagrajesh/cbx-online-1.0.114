"""1.0.104 (SPEC-1104 Phase 2, Builder C): V10 account status, masterlists and the completeness view; SAP payments kept apart
from the posting history; V11 the reconciliation population, the municipal roll-forward, the register with its items,
prepare → review by a different person, the report and Needs you's overdue line. Accounts, vendors and amounts are made up;
everything happens in a test portfolio (CTEST, company code Q900) so the seeded utilities list stays out of it."""
from __future__ import annotations

import io
import random
from datetime import date, timedelta

import pytest
from openpyxl import Workbook, load_workbook

from cbx.ap import completeness as CP
from cbx.ap import history as H
from cbx.ap import payments as PAY
from cbx.ap import portfolio as PF
from cbx.ap import recons as RC
from cbx.db import SessionLocal
from cbx.models import AuditLog, User
from cbx.models.ap import (ApBatch, ApBillHistory, ApCapture, ApEstimate, ApItem, ApMasterlistException, ApMasterlistImport, ApPayment, ApPortfolio,
                           ApPosted, ApRecon, ApReconItem, ApRole, ApUtilityAccount)

from .test_ap_capture_1098 import _admin, _app, _chart_of_accounts, _user, signed_in  # noqa: F401

PF_CODE = "CTEST"
CO = "Q900"
TAG = "v1104c test"


@pytest.fixture(autouse=True)
def _clean():
    with SessionLocal() as s:
        if not PF.get(s, PF_CODE):
            s.add(ApPortfolio(code=PF_CODE, name="C test", company_codes=CO, settings={}, fixed={}, gl={}, notes={}, updated_by=TAG))
            s.commit()
    yield
    with SessionLocal() as s:
        s.query(ApRole).delete()
        keys = [k for (k,) in s.query(ApUtilityAccount.account_key).filter(ApUtilityAccount.company_code == CO).all()]
        for b in s.query(ApBatch).filter(ApBatch.name == "C test batch").all():
            s.query(ApCapture).filter(ApCapture.batch_id == b.id).delete(synchronize_session=False)
            s.query(ApItem).filter(ApItem.batch_id == b.id).delete(synchronize_session=False)
            s.delete(b)
        rids = [i for (i,) in s.query(ApRecon.id).filter(ApRecon.portfolio == PF_CODE).all()]
        if rids:
            s.query(ApReconItem).filter(ApReconItem.recon_id.in_(rids)).delete(synchronize_session=False)
        s.query(ApRecon).filter(ApRecon.portfolio == PF_CODE).delete(synchronize_session=False)
        s.query(ApMasterlistException).filter(ApMasterlistException.portfolio == PF_CODE).delete(synchronize_session=False)
        s.query(ApMasterlistImport).filter(ApMasterlistImport.portfolio == PF_CODE).delete(synchronize_session=False)
        s.query(ApPayment).filter(ApPayment.company_code == CO).delete(synchronize_session=False)
        s.query(ApPosted).filter(ApPosted.company_code == CO).delete(synchronize_session=False)
        if keys:
            s.query(ApBillHistory).filter(ApBillHistory.account_key.in_(keys)).delete(synchronize_session=False)
            s.query(ApEstimate).filter(ApEstimate.account_key.in_(keys)).delete(synchronize_session=False)
        s.query(ApUtilityAccount).filter(ApUtilityAccount.company_code == CO).delete(synchronize_session=False)
        s.query(ApUtilityAccount).filter(ApUtilityAccount.updated_by == TAG).delete(synchronize_session=False)
        p = PF.get(s, PF_CODE)
        if p is not None:
            s.delete(p)
        s.commit()


def _acct(site: str = "Site", vendor: str = "700777", **kw) -> ApUtilityAccount:
    no = "77" + "".join(random.choice("0123456789") for _ in range(8))
    with SessionLocal() as s:
        a = ApUtilityAccount(account_key=no, account=no, vendor=vendor, cost_centre="01306", site=f"{site} {no[-4:]}", municipality="TEST MUNI", company_code=CO,
                             source="excel", updated_by=TAG, **kw)
        s.add(a)
        s.commit()
        s.refresh(a)
        s.expunge(a)
        return a


def _bill(key: str, period: str, *, bill_day: int = 10, opening=None, closing=None, total=1000.0, capture_id=None, reading="A", invoice=None) -> None:
    with SessionLocal() as s:
        s.add(ApBillHistory(account_key=key, portfolio=PF_CODE, municipality="TEST MUNI", period=period, bill_date=f"{period}-{bill_day:02d}",
                            invoice_no=invoice or f"INV{key[-4:]}{period.replace('-', '')}", read_start=f"{period}-01", read_end=f"{period}-{bill_day:02d}",
                            total_excl=total, opening=opening, closing=closing, reading_type=reading, capture_id=capture_id, source="import", findings=[]))
        s.commit()


def _uploaded_capture() -> int:
    with SessionLocal() as s:
        uid = s.query(User.id).order_by(User.id).first()[0]
        b = ApBatch(user_id=uid, name="C test batch", status="ready")
        s.add(b)
        s.flush()
        it = ApItem(batch_id=b.id, name="bill.pdf", kind="utility")
        s.add(it)
        s.flush()
        from cbx.db import utcnow

        c = ApCapture(item_id=it.id, batch_id=b.id, bucket="utility", stage="ready", uploaded_at=utcnow())
        s.add(c)
        s.commit()
        return c.id


def _user_obj(email: str) -> User:
    s = SessionLocal()
    u = s.query(User).filter(User.email == email).one()
    s.close()
    return u


def _xlsx(rows: list) -> bytes:
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _iso(d: date) -> str:
    return d.isoformat()


# ---- SAP payments: from the same FBL1N upload, kept apart ---------------------------------------------------------------------
def test_payments_from_the_vendor_line_items_are_kept_apart_and_tied_to_accounts():
    a1, a2 = _acct("Pay one"), _acct("Pay two")
    heads = ["Company Code", "Account", "Reference", "Document Type", "Posting Date", "Document Date", "Clearing Date", "Clearing Document", "Document Number",
             "Local Currency", "Tax Code", "Amount in Local Currency", "Text", "Assignment"]
    data = _xlsx([heads,
                  [CO, "700777", a1.account, "KR", "2026-08-05", "2026-08-05", "2026-08-20", "1500009901", "1900009901", "ZAR", "", -1150.00, "bill", ""],
                  [CO, "700777", a2.account, "KR", "2026-08-06", "2026-08-06", "2026-08-20", "1500009901", "1900009902", "ZAR", "", -230.00, "bill", ""],
                  [CO, "700777", "", "KZ", "2026-08-20", "2026-08-20", "2026-08-20", "1500009901", "1500009901", "ZAR", "", 1380.00, "Payment - EFT", "00000"],
                  [CO, "700777", f"ACC {a1.account}", "KZ", "2026-08-25", "2026-08-25", "", "", "1500009902", "ZAR", "", 99.00, "Payment - EFT", ""]])
    rows = H.read_items("vendor", data, "fbl1n.xlsx")
    assert rows[0]["clearing_doc"] == "1500009901" and rows[0]["clearing_date"] == "2026-08-20"
    admin = _admin()  # noqa: F841 — the admin person exists
    with SessionLocal() as db:
        u = db.query(User).filter(User.email == "cap.admin@test.local").one()
        prev = H.save(db, u, rows, None, confirm=False)
        assert prev["payments"] == 2 and prev["documents"] == 2  # the duplicate check's documents: the two invoices only
        out = H.save(db, u, rows, None, confirm=True)
        db.commit()
        assert out["payments_saved"] == {"payments": 2, "added": 2, "replaced": 0, "allocated": 2}
        assert not db.query(ApPosted).filter(ApPosted.company_code == CO, ApPosted.doc_type == "KZ").first()  # never in ap_posted
        p1 = db.query(ApPayment).filter(ApPayment.company_code == CO, ApPayment.doc_number == "1500009901").one()
        assert p1.amount == 1380.0 and p1.allocations == {a1.account_key: 1150.0, a2.account_key: 230.0} and p1.account_key == ""
        p2 = db.query(ApPayment).filter(ApPayment.company_code == CO, ApPayment.doc_number == "1500009902").one()
        assert p2.allocations == {a1.account_key: 99.0} and p2.account_key == a1.account_key  # the account number in its reference
        assert PAY.paid(db, a1.account_key, "2026-08-01", "2026-08-31") == 1249.0
        assert PAY.paid(db, a1.account_key, "2026-08-21", "2026-08-31") == 99.0
        assert PAY.paid_in_month(db, [a1.account_key, a2.account_key], "2026-08") == {a1.account_key: 1249.0, a2.account_key: 230.0}
        assert db.query(AuditLog).filter(AuditLog.action == "ap_payments_loaded").count() >= 1
        # the same upload again replaces, never doubles
        again = H.save(db, u, rows, None, confirm=True)
        db.commit()
        assert again["payments_saved"]["replaced"] == 2 and db.query(ApPayment).filter(ApPayment.company_code == CO).count() == 2


# ---- V10: account fields ------------------------------------------------------------------------------------------------------
def test_an_accounts_status_and_fields_are_set_by_a_reviewer_and_audited():
    a = _acct("Status")
    adm = _admin()
    r = adm.put(f"/api/ap/completeness/accounts/{a.id}", json={"status": "terminated"})
    assert r.status_code == 400 and "day" in r.json()["detail"]
    r = adm.put(f"/api/ap/completeness/accounts/{a.id}", json={"status": "terminated", "status_from": "2026-09-30", "co_do": "do", "cluster": "2", "check_meter": True, "bill_day": 12})
    assert r.status_code == 200, r.text
    got = r.json()
    assert (got["status"], got["status_from"], got["co_do"], got["cluster"], got["check_meter"], got["bill_day"], got["portfolio"]) == ("terminated", "2026-09-30", "DO", "2", True, 12, PF_CODE)
    with SessionLocal() as db:
        row = db.query(AuditLog).filter(AuditLog.action == "ap_account_status", AuditLog.target == f"ap_account:{a.account_key}").order_by(AuditLog.id.desc()).first()
        assert row is not None and "terminated" in row.detail and "CO/DO - → DO" in row.detail
    # once the portfolio has AP roles, a person without the reviewer role there can't change it
    uid = _user_obj("cap.admin@test.local").id
    with SessionLocal() as db:
        db.add(ApRole(user_id=uid, portfolio=PF_CODE, role="reviewer", granted_by=TAG))
        db.commit()
    u = _user()
    r = u.put(f"/api/ap/completeness/accounts/{a.id}", json={"cluster": "3"})
    assert r.status_code == 403 and "reviewer" in r.json()["detail"]
    assert u.get(f"/api/ap/completeness?portfolio={PF_CODE}").status_code == 403  # the register is for reviewers and admins


# ---- V10: masterlist import -----------------------------------------------------------------------------------------------------
def test_a_masterlist_is_compared_taken_and_its_exceptions_resolved_with_a_note():
    on_both = _acct("Both", co_do="CO")
    gone = _acct("Gone")
    not_listed = _acct("Not listed")
    _bill(gone.account_key, "2026-10", bill_day=15)  # a bill after the termination date
    new_no = "7799" + "".join(random.choice("0123456789") for _ in range(6))
    data = _xlsx([["Masterlist October 2026"], [],
                  ["Account Number", "Vendor", "Site Name", "Status", "Termination Date", "CO/DO", "Cluster", "Check Meter", "Bill Day"],
                  [on_both.account, "700777", on_both.site, "Active", "", "DO", "1", "Y", "8"],
                  [gone.account, "700777", gone.site, "Terminated", "2026-09-30", "", "1", "", ""],
                  [new_no, "9999999999", "Brand new site", "Active", "", "CO", "2", "", ""]])
    rows = CP.read_list(data, "masterlist.xlsx")
    assert [r["account_key"] for r in rows] == [on_both.account_key, gone.account_key, new_no]
    adm = _admin()
    files = {"file": ("masterlist.xlsx", data, "application/octet-stream")}
    form = {"portfolio": PF_CODE, "kind": "masterlist", "source": "Test masterlist Oct 2026", "confirm": "false"}
    prev = adm.post("/api/ap/completeness/import", files=files, data=form)
    assert prev.status_code == 200, prev.text
    c = prev.json()["counts"]
    assert (c["rows"], c["matched"], c["changed"], c["not_in_cbx"], c["not_on_list"], c["terminated_active"], c["terminated_billed"], c["co_do"]) == (3, 2, 2, 1, 1, 1, 1, 1)
    with SessionLocal() as db:
        assert CP.status_of(db.get(ApUtilityAccount, gone.id)) == "active"  # a preview writes nothing
        sap_loaded = CP._sap_vendors(db) is not None
    assert c["not_in_sap"] == (3 if sap_loaded else 0)  # 700777 (two rows) and 9999999999 are made up
    r = adm.post("/api/ap/completeness/import", files=files, data={**form, "confirm": "true"})
    assert r.status_code == 200 and r.json()["saved"]
    with SessionLocal() as db:
        a = db.get(ApUtilityAccount, on_both.id)
        assert (a.co_do, a.cluster, a.check_meter, a.bill_day) == ("DO", "1", True, 8)
        g = db.get(ApUtilityAccount, gone.id)
        assert (g.status, g.status_from) == ("terminated", "2026-09-30") and g.status_source.startswith("Test masterlist Oct 2026")
        assert CP.status_of(db.get(ApUtilityAccount, not_listed.id)) == "active"  # not on the list: an exception, never a change
        imp = db.query(ApMasterlistImport).filter(ApMasterlistImport.portfolio == PF_CODE).one()
        assert imp.source == "Test masterlist Oct 2026" and imp.rows == 3
        log = db.query(AuditLog).filter(AuditLog.action == "ap_masterlist_import", AuditLog.target == f"ap_portfolio:{PF_CODE}").order_by(AuditLog.id.desc()).first()
        assert "Test masterlist Oct 2026" in log.detail and date.today().isoformat() in log.detail
    ex = adm.get(f"/api/ap/completeness/exceptions?portfolio={PF_CODE}").json()
    kinds = sorted(e["kind"] for e in ex["rows"])
    assert {"not_in_cbx", "not_on_list", "terminated_active", "terminated_billed", "co_do"} <= set(kinds)
    billed = next(e for e in ex["rows"] if e["kind"] == "terminated_billed")
    assert "still billed" in billed["text"]
    assert adm.post(f"/api/ap/completeness/exceptions/{billed['id']}/resolve", json={"note": ""}).status_code == 400
    r = adm.post(f"/api/ap/completeness/exceptions/{billed['id']}/resolve", json={"note": "Final bill; council confirmed closure"})
    assert r.status_code == 200 and r.json()["status"] == "resolved"
    assert adm.get(f"/api/ap/completeness/exceptions?portfolio={PF_CODE}").json()["total"] == len(kinds) - 1
    with SessionLocal() as db:
        assert db.query(AuditLog).filter(AuditLog.action == "ap_masterlist_exception").order_by(AuditLog.id.desc()).first().detail.endswith("council confirmed closure")
    # a Termination report: every listed account is terminated
    t = _xlsx([["Account", "Date terminated"], [on_both.account, "15.10.2026"]])
    r = adm.post("/api/ap/completeness/import", files={"file": ("terminations.xlsx", t, "application/octet-stream")},
                 data={"portfolio": PF_CODE, "kind": "terminations", "source": "Termination report", "confirm": "true"})
    assert r.status_code == 200 and r.json()["counts"]["terminated_active"] == 1
    with SessionLocal() as db:
        assert (db.get(ApUtilityAccount, on_both.id).status, db.get(ApUtilityAccount, on_both.id).status_from) == ("terminated", "2026-10-15")
    bad = adm.post("/api/ap/completeness/import", files={"file": ("x.xlsx", _xlsx([["Nothing", "here"], [1, 2]]), "application/octet-stream")}, data={"portfolio": PF_CODE})
    assert bad.status_code == 400 and "account column" in bad.json()["detail"]


# ---- V10: the completeness view -----------------------------------------------------------------------------------------------
def test_the_completeness_view_says_posted_read_estimated_missing_and_late():
    period = "2026-09"
    posted, read, est, late, early, gone = (_acct(n, cluster=c) for n, c in (("Posted", "1"), ("Read", "1"), ("Est", "2"), ("Late", "2"), ("Early", "2"), ("Gone", "1")))
    _bill(posted.account_key, period, capture_id=_uploaded_capture())
    _bill(read.account_key, period)
    with SessionLocal() as db:
        db.add(ApEstimate(account_key=est.account_key, account=est.account, portfolio=PF_CODE, period=period, value=812.5, status="approved", prepared_by="x", approved_by="y"))
        a = db.get(ApUtilityAccount, early.id)
        a.bill_day = 28
        g = db.get(ApUtilityAccount, gone.id)
        g.status, g.status_from = "terminated", "2026-08-31"  # closed before the month: not expected
        db.add(ApPayment(company_code=CO, vendor="700777", doc_number="1500008801", doc_type="KZ", posting_date="2026-09-21", amount=500.0, allocations={read.account_key: 500.0}))
        db.commit()
    _bill(late.account_key, "2026-08", bill_day=5)  # its usual day is learned: the 5th
    with SessionLocal() as db:
        v = CP.month_view(db, PF_CODE, period, today=date(2026, 9, 20))
    by = {r["account_key"]: r for r in v["rows"]}
    assert gone.account_key not in by and v["summary"]["accounts"] == 5
    assert by[posted.account_key]["state"] == "posted" and by[read.account_key]["state"] == "read" and by[est.account_key]["state"] == "estimated"
    assert by[read.account_key]["paid"] == 500.0 and v["payments_loaded"]
    lt = by[late.account_key]
    assert (lt["state"], lt["bill_day"], lt["bill_day_how"], lt["expected"], lt["late_from"], lt["age_days"], lt["late"]) == ("missing", 5, "learned", "2026-09-05", "2026-09-15", 15, True)
    er = by[early.account_key]
    assert (er["state"], er["late"], er["expected"]) == ("missing", False, "2026-09-28")
    s = v["summary"]
    assert (s["posted"], s["read"], s["estimated"], s["missing"], s["late"], s["complete_pct"]) == (1, 1, 1, 2, 1, 60.0)
    assert {c["cluster"]: c["missing"] for c in v["by_cluster"]} == {"1": 0, "2": 2}
    with SessionLocal() as db:
        li = CP.liaison(db, PF_CODE, period, today=date(2026, 9, 20))
        assert [r["account_key"] for r in li] == [late.account_key]
        assert [r["account_key"] for r in CP.month_view(db, PF_CODE, period, cluster="2", state="estimated", today=date(2026, 9, 20))["rows"]] == [est.account_key]
    adm = _admin()
    r = adm.get(f"/api/ap/completeness?portfolio={PF_CODE}&period={period}")
    assert r.status_code == 200 and r.json()["summary"]["accounts"] == 5 and r.json()["manage"]
    x = adm.get(f"/api/ap/completeness/export?portfolio={PF_CODE}&period={period}")
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content), read_only=True)
    assert wb.sheetnames == ["Summary", "Accounts", "Council Liaison", "Exceptions"]
    assert sum(1 for _ in wb["Accounts"].iter_rows()) == 6
    wb.close()


# ---- V10: Tess on a terminated account --------------------------------------------------------------------------------------
def test_tess_makes_a_bill_on_a_terminated_account_red_and_a_final_bill_yellow():
    a = _acct("Tess", status="terminated", status_from="2026-09-30")
    s_ = _acct("Susp", status="suspended", status_from="2026-09-01")
    with SessionLocal() as db:
        c = ApCapture(bucket="utility", match={"account": {"id": a.id, "found": True}}, fields={"read_end": "2026-10-28"})
        assert CP.tess_note(db, c) == (["account terminated on 30.09.2026"], [])
        c.fields = {"read_end": "2026-09-30"}
        red, yellow = CP.tess_note(db, c)
        assert not red and yellow[0].startswith("account terminated on 30.09.2026: final bill?")
        c2 = ApCapture(bucket="utility", match={"account": {"id": s_.id}}, fields={})
        assert CP.tess_note(db, c2) == ([], ["account suspended since 01.09.2026"])
        assert CP.tess_note(db, ApCapture(bucket="manual", match={}, fields={})) == ([], [])


# ---- V11: roll-forward, population, register ---------------------------------------------------------------------------------------
def _recon_setup():
    """Three accounts: one whose roll-forward breaks, one in credit, one that rolls forward cleanly; and the payments."""
    broke, credit, clean = _acct("Broke"), _acct("Credit"), _acct("Clean")
    _bill(broke.account_key, "2026-08", opening=0, closing=5000.0)
    _bill(broke.account_key, "2026-09", opening=3200.0, closing=4100.0)  # paid 2 000 → should open at 3 000
    _bill(credit.account_key, "2026-08", opening=0, closing=100.0)
    _bill(credit.account_key, "2026-09", opening=100.0, closing=-450.0)
    _bill(clean.account_key, "2026-08", opening=0, closing=800.0)
    _bill(clean.account_key, "2026-09", opening=300.0, closing=900.0)  # paid 500
    with SessionLocal() as db:
        db.add(ApPayment(company_code=CO, vendor="700777", doc_number="1500007701", doc_type="ZP", posting_date="2026-08-25", amount=2500.0,
                         allocations={broke.account_key: 2000.0, clean.account_key: 500.0}))
        db.add(ApPayment(company_code=CO, vendor="700888", doc_number="1500007702", doc_type="KZ", posting_date="2026-09-03", amount=1234.0, allocations={}))
        db.commit()
    return broke, credit, clean


def test_the_roll_forward_finds_what_does_not_agree():
    broke, credit, clean = _recon_setup()
    with SessionLocal() as db:
        rf = RC.roll_forward(db, broke.account_key, "2026-09")
        assert (rf["state"], rf["prev_closing"], rf["paid"], rf["expected"], rf["opening"], rf["difference"]) == ("exception", 5000.0, 2000.0, 3000.0, 3200.0, 200.0)
        assert "this bill opens at R 3 200.00" in rf["text"]
        assert RC.roll_forward(db, clean.account_key, "2026-09")["state"] == "ok"
        assert RC.roll_forward(db, clean.account_key, "2026-10")["state"] == "no_bill"
        db.query(ApPayment).filter(ApPayment.company_code == CO).delete()
        db.flush()
        if not PAY.loaded(db, "2026-08", "2026-09"):  # nothing else in the database for these months: the roll-forward can't judge
            assert RC.roll_forward(db, broke.account_key, "2026-09")["state"] == "no_payments"
        db.rollback()


def test_the_population_register_prepare_review_and_report():
    broke, credit, clean = _recon_setup()
    period = "2026-09"
    with SessionLocal() as db:
        db.query(ApPortfolio).filter(ApPortfolio.code == PF_CODE).one().settings = {"recon_top_n": 1}
        db.commit()
    adm = _admin()
    r = adm.post("/api/ap/recons/build", json={"portfolio": PF_CODE, "period": period})
    assert r.status_code == 200, r.text
    assert r.json()["population"] == 3  # top balance (broke), credit, the vendor's manual KZ payment tied to nothing
    again = adm.post("/api/ap/recons/build", json={"portfolio": PF_CODE, "period": period}).json()
    assert again["added"] == 0 and again["items"] == 0  # run again: nothing doubles
    lst = adm.get(f"/api/ap/recons?portfolio={PF_CODE}&period={period}").json()
    by = {(x["account_key"], x["vendor"]): x for x in lst["rows"]}
    rb = by[(broke.account_key, "700777")]
    assert {w["rule"] for w in rb["why"]} == {"top_balances", "roll_forward"} and rb["statement_balance"] == 4100.0 and rb["roll"]["difference"] == 200.0
    rc = by[(credit.account_key, "700777")]
    assert [w["rule"] for w in rc["why"]] == ["credit_balances"]
    rv = by[("", "700888")]
    assert rv["why"][0]["rule"] == "manual_payments" and rv["items_open"] == 1
    assert (clean.account_key, "700777") not in by
    assert lst["credits_recoverable"] == 450.0 and lst["summary"]["population"] == 3
    d = adm.get(f"/api/ap/recons/{rb['id']}").json()
    assert [i["type"] for i in d["items"]] == ["roll_forward"] and d["items"][0]["amount"] == 200.0 and len(d["bills"]) == 2
    # prepare: the difference must be explained by the open items, or the note says why
    prep = signed_in("recon.prep@test.local", "user", name="Pat Preparer", tools=["ap"])
    rev = signed_in("recon.rev@test.local", "user", name="Rita Reviewer", tools=["ap"])
    p = prep.put(f"/api/ap/recons/{rb['id']}", json={"statement_balance": 4100, "ledger_balance": 3800, "sign": True})
    assert p.status_code == 400 and "not explained" in p.json()["detail"]
    p = prep.put(f"/api/ap/recons/{rb['id']}", json={"statement_balance": 4100, "ledger_balance": 3900, "sign": True})
    assert p.status_code == 200, p.text
    assert p.json()["status"] == "prepared" and p.json()["difference"] == 200.0 and p.json()["preparer"] == "Pat Preparer"
    # the preparer never reviews it — bootstrap: a reviewer counts as a reconciliation reviewer
    assert prep.post(f"/api/ap/recons/{rb['id']}/review", json={"outcome": "reviewed"}).status_code == 403
    assert rev.post(f"/api/ap/recons/{rb['id']}/review", json={"outcome": "returned"}).status_code == 400  # returning needs a note
    ok = rev.post(f"/api/ap/recons/{rb['id']}/review", json={"outcome": "reviewed", "note": "agreed to the statement"})
    assert ok.status_code == 200 and ok.json()["status"] == "reviewed" and ok.json()["reviewer"] == "Rita Reviewer"
    with SessionLocal() as db:
        assert db.query(AuditLog).filter(AuditLog.action == "ap_recon_prepared", AuditLog.target == f"ap_recon:{rb['id']}").count() == 1
        assert db.query(AuditLog).filter(AuditLog.action == "ap_recon_reviewed", AuditLog.target == f"ap_recon:{rb['id']}").count() == 1
    # once the portfolio has AP roles: only the reconciliation-reviewer role reviews — an admin gets no exemption
    p = prep.put(f"/api/ap/recons/{rc['id']}", json={"statement_balance": -450, "ledger_balance": -450, "note": "credit to claim back", "sign": True})
    assert p.status_code == 200 and p.json()["status"] == "prepared"
    with SessionLocal() as db:
        db.add(ApRole(user_id=_user_obj("recon.rev@test.local").id, portfolio=PF_CODE, role="recon_reviewer", granted_by=TAG))
        db.commit()
    assert adm.post(f"/api/ap/recons/{rc['id']}/review", json={"outcome": "reviewed"}).status_code == 403
    assert rev.post(f"/api/ap/recons/{rc['id']}/review", json={"outcome": "reviewed"}).status_code == 200
    # items: a person adds one, owns it, clears it with a note
    it = rev.post(f"/api/ap/recons/{rv['id']}/items", json={"type": "missing_invoice", "amount": 99.5, "text": "invoice 4411 not captured"})
    assert it.status_code == 200, it.text
    new = next(i for i in it.json()["items"] if i["text"] == "invoice 4411 not captured")
    assert new["owner"] == "Rita Reviewer" and new["target_date"] > date.today().isoformat()
    assert rev.post(f"/api/ap/recons/{rv['id']}/items/{new['id']}/clear", json={"note": ""}).status_code == 400
    cl = rev.post(f"/api/ap/recons/{rv['id']}/items/{new['id']}/clear", json={"note": "captured in batch 12"})
    assert cl.status_code == 200 and next(i for i in cl.json()["items"] if i["id"] == new["id"])["cleared"]
    # the report
    x = adm.get(f"/api/ap/recons/report?portfolio={PF_CODE}&period={period}")
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content), read_only=True)
    assert wb.sheetnames == ["Summary", "Register", "Items"]
    summ = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True) if r and r[0]}
    assert summ["In the population"] == 3 and summ["Reviewed"] == 2 and summ["Completion % (reviewed)"] == 66.7 and summ["Credits recoverable (R)"] == 450.0
    wb.close()


def test_needs_you_shows_only_overdue_items():
    broke, _credit, _clean = _recon_setup()
    with SessionLocal() as db:
        RC.build(db, None, PF_CODE, "2026-09")
        r = db.query(ApRecon).filter(ApRecon.portfolio == PF_CODE, ApRecon.account_key == broke.account_key).one()
        db.commit()
    u = _user()
    uid = _user_obj("cap.user@test.local").id
    assert not any(x["key"] == "recons_overdue" for x in u.get("/api/ap/needs").json().get("controls", []))
    with SessionLocal() as db:
        it = db.query(ApReconItem).filter(ApReconItem.recon_id == r.id).first()
        it.owner_id, it.owner, it.target_date = uid, "Cara Person", (date.today() - timedelta(days=2)).isoformat()
        db.commit()
    lines = [x for x in u.get("/api/ap/needs").json().get("controls", []) if x["key"] == "recons_overdue"]
    assert lines == [{"key": "recons_overdue", "sub": "reconciliations", "params": {"overdue": "1"}, "text": "1 reconciliation item is past the target date"}]
    other = signed_in("recon.other@test.local", "user", name="Olga Other", tools=["ap"])
    assert not any(x["key"] == "recons_overdue" for x in other.get("/api/ap/needs").json().get("controls", []))


def test_the_settings_keys_are_offered_per_portfolio():
    keys = {k["key"]: k for k in PF.keys_out()}
    for k in ("default_bill_day", "recon_rules", "recon_top_n", "recon_bf_min_r", "recon_tolerance_r", "recon_manual_payment_types", "recon_item_days"):
        assert k in keys, k
    assert keys["recon_rules"]["kind"] == "lines" and keys["recon_top_n"]["kind"] == "number" and keys["recon_top_n"]["label"] == "Top balances to reconcile"
    with SessionLocal() as db:
        assert PF.setting(db, "SHELL", "recon_manual_payment_types", ["KZ"]) == ["KZ"]


def test_a_prepared_reconciliation_is_not_changed_by_its_reviewer_and_no_editor_reviews_it():
    """Review finding: the reviewer could change the figures of a prepared reconciliation and then review them."""
    broke, _credit, _clean = _recon_setup()
    with SessionLocal() as db:
        RC.build(db, None, PF_CODE, "2026-09")
        rid = db.query(ApRecon.id).filter(ApRecon.portfolio == PF_CODE, ApRecon.account_key == broke.account_key).scalar()
        db.commit()
    prep = signed_in("recon.prep@test.local", "user", name="Pat Preparer", tools=["ap"])
    rev = signed_in("recon.rev@test.local", "user", name="Rita Reviewer", tools=["ap"])
    third = signed_in("recon.third@test.local", "user", name="Tom Third", tools=["ap"])
    assert prep.put(f"/api/ap/recons/{rid}", json={"statement_balance": 4100, "ledger_balance": 3900, "sign": True}).json()["status"] == "prepared"
    # while it waits for review nobody changes it — not its balances, not its items
    r = rev.put(f"/api/ap/recons/{rid}", json={"ledger_balance": 4100})
    assert r.status_code == 400 and "send it back" in r.json()["detail"]
    assert rev.post(f"/api/ap/recons/{rid}/items", json={"type": "other", "text": "a quiet change"}).status_code == 400
    item = rev.get(f"/api/ap/recons/{rid}").json()["items"][0]
    assert rev.post(f"/api/ap/recons/{rid}/items/{item['id']}/clear", json={"note": "cleared by the reviewer"}).status_code == 400
    with SessionLocal() as db:
        assert db.get(ApRecon, rid).ledger_balance == 3900.0
    # the reviewer sends it back (a note is needed), changes it, and so may no longer review it
    assert rev.post(f"/api/ap/recons/{rid}/send-back", json={"note": ""}).status_code == 400
    sb = rev.post(f"/api/ap/recons/{rid}/send-back", json={"note": "ledger looks wrong"})
    assert sb.status_code == 200 and sb.json()["status"] == "open"
    ed = rev.put(f"/api/ap/recons/{rid}", json={"ledger_balance": 3900, "note": "checked the ledger"})
    assert ed.status_code == 200 and ed.json()["preparer"] == "Rita Reviewer"  # whoever edits becomes the preparer
    assert prep.put(f"/api/ap/recons/{rid}", json={"sign": True}).json()["preparer"] == "Pat Preparer"
    d = rev.get(f"/api/ap/recons/{rid}").json()
    assert not d["can_review"] and d["can_send_back"]
    r = rev.post(f"/api/ap/recons/{rid}/review", json={"outcome": "reviewed"})
    assert r.status_code == 403 and "changed" in r.json()["detail"]
    ok = third.post(f"/api/ap/recons/{rid}/review", json={"outcome": "reviewed"})
    assert ok.status_code == 200 and ok.json()["reviewer"] == "Tom Third"
    with SessionLocal() as db:
        assert db.query(AuditLog).filter(AuditLog.action == "ap_recon_sent_back", AuditLog.target == f"ap_recon:{rid}").count() == 1
