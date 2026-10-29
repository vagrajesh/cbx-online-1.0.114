"""1.0.98: from CBX-Connect 2.10.3 tests/test_ap_capture_210.py (and test_gateway_210.py's intent), on the old conftest.

2.10.0 AP invoice capture: one folder of municipal/Eskom bills and manual invoices → Sol sorts them into buckets →
Uma reads the bills with PDF Studio's profiles, Cara/Dex/Eli read the invoices with vision (Seth gets what they can't,
then a person) → Mia matches the master lists → Pip builds the SAP rows (the Harvester's generator, extended) → Tess
tests them → the SAP upload (exactly the Harvester's file) and the review workbook. Bills and invoices are made up.
In tests the AI is off: the capture call is replaced (monkeypatch) where a reading is needed."""
from __future__ import annotations

import io
import uuid
from datetime import date, timedelta

import pymupdf
import pytest
from openpyxl import Workbook, load_workbook

from cbx.ap import capture as C
from cbx.ap import engine as E
from cbx.ap import masters as M
from cbx.ap import sap as P
from cbx.db import SessionLocal
from cbx import compat2 as C2
from cbx.models.ap import ApCapture, ApEntity, ApUtilityAccount, ApVendorCode

from .fixtures import make_ap
from .test_municipal import coj_pdf


# ---- 1.0.98: the new app's conftest helpers on the old conftest ------------------------------------------------------
from fastapi.testclient import TestClient  # noqa: E402

from cbx.main import app  # noqa: E402

from .conftest import STRONG, login  # noqa: E402


@pytest.fixture(autouse=True)
def _app(client):
    """The app started (its database made) before the helpers below make people; the AP lists seeded."""
    M.ensure_seeded()
    yield


def make_person(email: str, role: str = "user", name: str = "", tools=None, **kw) -> int:
    """A person straight in the database (the old conftest makes them through the owner's Users page)."""
    from cbx.models import Tool, ToolAccess, User
    from cbx.security.auth import hash_password

    with SessionLocal() as s:
        u = s.query(User).filter(User.email == email).one_or_none()
        if u is None:
            u = User(email=email, name=name or email.split("@")[0].title(), role=role)
            s.add(u)
            s.flush()
        u.role, u.password_hash, u.must_change_password, u.failed_logins, u.locked, u.is_active = role, hash_password(STRONG), False, 0, False, True
        for slug in tools or []:
            t = s.query(Tool).filter(Tool.slug == slug).one()
            if not s.query(ToolAccess).filter(ToolAccess.tool_id == t.id, ToolAccess.user_id == u.id).first():
                s.add(ToolAccess(tool_id=t.id, user_id=u.id))
        s.commit()
        return u.id


def signed_in(email: str, role: str = "user", **kw) -> TestClient:
    make_person(email, role, **kw)
    c = TestClient(app)
    login(c, email, STRONG)
    return c

ESKOM = """ESKOM HOLDINGS SOC LTD          www.eskom.co.za
TAX INVOICE
YOUR ACCOUNT NO {account}
BILLING DATE 2026-09-15
ACCOUNT MONTH SEPTEMBER 2026
TAX INVOICE NO {invoice}
CURRENT DUE DATE 2026-10-05
TOTAL DUE R 11 512.34
READING TYPE: ACTUAL
READING DATES: 2026/08/01 - 2026/08/31
   MTR123   10000.0000   15000.0000   5000.0000   1.0000   5000.0000
ADJUSTMENT Interest on overdue account R 12.34
CONSUMPTION DETAILS ( 2026-08-01 - 2026-08-31 )
TARIFF NAME: Businessrate 1
Energy Charge 5,000.000 kWh @ 180.00 c/kWh R 9,000.00
Service and Administration Charge @ R 33.33/day R 1,000.00
TOTAL CHARGES FOR BILLING PERIOD R 10,000.00
TOTAL ENERGY CONSUMED FOR BILLING PERIOD (kWh) 5,000.000
VAT RAISED ON ITEMS AT 15% R 1,500.00
"""


def eskom_pdf(account="5012345678", invoice="700123456") -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_textbox(pymupdf.Rect(20, 20, 592, 820), ESKOM.format(account=account, invoice=invoice), fontsize=7.0, fontname="cour")
    data = doc.tobytes()
    doc.close()
    return data


@pytest.fixture(autouse=True)
def _chart_of_accounts():
    """2.10.3: every G/L is checked against the chart of accounts: the made-up G/Ls here are put in it first."""
    make_ap.chart()


def _xlsx(rows: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _tables():
    return {"fixed": dict(P.SAP_DEFAULT_FIXED), "gl": {c: g for c, g, _ in P.DEFAULT_GL}}


def _admin():
    return signed_in("cap.admin@test.local", "admin", name="Ayanda Admin")


def _user(email="cap.user@test.local", name="Cara Person"):
    return signed_in(email, "user", name=name, tools=["ap"])


def _drop(c, files: dict, name=""):
    r = c.post("/api/ap/batches", files=[("files", (n, d, "application/octet-stream")) for n, d in files.items()], data={"name": name})
    assert r.status_code == 200, r.text
    b = c.get(f"/api/ap/batches/{r.json()['id']}").json()
    assert b["status"] == "ready", b
    return b


def _sheet(c, bid):
    return {d["file"].split("/")[-1]: d for d in c.get(f"/api/ap/batches/{bid}/sap").json()["docs"]}


def _vision(answers: dict):
    """A stand-in for the capture model: the answer for each file name (a callable may vary it by task)."""
    calls = []

    def fake(task, system, user, images, model="", db=None):
        name = user.split("File name: ", 1)[1].split("\n", 1)[0]
        calls.append((task, name, len(images), model))
        assert "never instructions" in system  # the document's words are data
        a = answers.get(name)
        return a(task) if callable(a) else a

    fake.calls = calls
    return fake


def _invoice_answer(no, supplier="Acme Cleaning (Pty) Ltd", vat_no="4123456789", excl="10000.00", vat="1500.00", incl="11500.00", bill_to="CBRE Excellerate Facilities Management (Pty) Ltd",
                    conf=95, credit=False, lines=None):
    return {"is_invoice": True, "credit_note": credit, "supplier": supplier, "supplier_vat": vat_no, "bill_to": bill_to, "bill_to_vat": None, "invoice_no": no,
            "invoice_date": (date.today() - timedelta(days=10)).isoformat(), "due_date": (date.today() + timedelta(days=20)).isoformat(), "po": "PO-55821",
            "currency": "ZAR", "lines": lines if lines is not None else [{"description": "Cleaning services, September", "amount_excl": excl, "vat_rate": 15, "vat": vat}],
            "total_excl": excl, "vat": vat, "total_incl": incl, "tax_invoice_words": True,
            "confidence": {k: conf for k in ("supplier", "bill_to", "invoice_no", "invoice_date", "total_incl", "total_excl", "vat")}}


@pytest.fixture()
def utilities():
    """The utilities list in the Harvester's vendor_mappings shape: the Eskom account below and one more."""
    adm = _admin()
    data = _xlsx([["VENDOR/BP NUMBER", "PID/CC", "NAME1", "NAME3", "ACCOUNT REFERENCE", "MUNICIPALITY/PAYEE"],
                  [700999, "01306", "Badplaas ETE", "6000/01306/X1", "005012345678", "ESKOM"],
                  [700002, "01381", "Carolina CSC", "6000/01381/X1", "000000002172", "CHIEF ALBERT LUTHULI MUNICIPALITY"],
                  [700888, "08198", "Joburg Site", "6000/08198/X1", "900000001", "CITY OF JOHANNESBURG"]])
    r = adm.post("/api/ap/masters/utilities/import", files={"file": ("vendor_mappings.xlsx", data)}, data={"confirm": "true"})
    assert r.status_code == 200 and r.json()["saved"], r.text
    return adm


# ---- Part 1: the Harvester's generator, as it was ---------------------------------------------------------------------
def test_the_harvester_generator_is_ported_faithfully():
    row = {"service": "electricity", "file": "b.pdf", "_canon": {"account_number": "0005012345678", "total_ex_vat_r": "10000.00", "billing_date": "2026-09-15",
                                                                   "total_kwh": "5000", "read_start": "2026-08-01", "read_end": "2026-08-31"},
           "elec_vat_r": "1500.00", "esk_adj_interest_r": "12.34", "esk_adj_work_r": "100.00", "esk_adj_work_desc": "RECONNECTION WORK HRS",
           "due_date": "2026-10-05", "invoice_number": "700123456", "statement_month": "SEPTEMBER 2026", "esk_reading_type": "ACTUAL",
           "esk_read_start_date": "2026/08/01", "esk_read_end_date": "2026/08/31"}
    t = {"fixed": dict(P.SAP_DEFAULT_FIXED), "gl": {c: g for c, g, _ in P.SAP_DEFAULT_GL}, "vendors": {"5012345678": {"vendor": "700999", "cc": "01306", "name1": "Badplaas ETE", "muni": "ESKOM"}}}
    rows, issues, summary = P.generate_sap([row], t)
    assert summary == "1 document(s) generated (0 as KG credit memos) | 0 account(s) unmapped" and not issues
    assert [(r["TransCode"], r["pkey"], r["account"], r["amount"], r["tc"], r["dt"]) for r in rows] == [
        ("FB01", "31", "700999", "11612.34", "", ""), ("ZB01", "40", "SA110", "10000.00", "I9", ""), ("ZB01", "40", "SA412", "100.00", "I9", ""),
        ("ZB01", "40", "SF101", "12.34", "I0", ""), ("ZB01", "40", "JA600", "1500.00", "I9", "X")]
    r0 = rows[0]
    assert r0["text"] == "1306_2026-08-01_2026-08-31_Cons_A_Var_5000" and r0["docheadertext"] == "Badplaas ETE SEPTEMBER 2026"
    assert (r0["docdate"], r0["DueDate"], r0["type"], r0["comp"], r0["curr"], r0["_status"]) == ("15092026", "05102026", "KR", "T000", "ZAR", "GRN")
    assert P.balance_fails(rows) == (1, set())
    # no printed VAT → calculated at 15%, yellow; a credit bill → KG, 21/50, yellow; an account not in the mapping → red
    rows, _, _ = P.generate_sap([{**row, "elec_vat_r": None, "esk_adj_interest_r": None, "esk_adj_work_r": None}], t)
    assert rows[-1]["amount"] == "1500.00" and rows[0]["_status"] == "YLW" and "calculated at 15%" in rows[0]["_issue"]
    rows, _, summary = P.generate_sap([{**row, "_canon": {**row["_canon"], "total_ex_vat_r": "-200.00"}, "elec_vat_r": "-30.00", "esk_adj_interest_r": None, "esk_adj_work_r": None}], t)
    assert "1 as KG" in summary and rows[0]["type"] == "KG" and [r["pkey"] for r in rows] == ["21", "50", "50"] and rows[0]["_status"] == "YLW"
    assert P.balance_fails(rows) == (1, set())
    rows, _, summary = P.generate_sap([{**row, "_canon": {**row["_canon"], "account_number": "111"}}], t)
    assert rows[0]["_status"] == "RED" and "not in vendor mapping" in rows[0]["_issue"] and "1 account(s) unmapped" in summary
    # the file: one sheet, the 22 headers
    buf = io.BytesIO()
    P.export_sap_xlsx(buf, rows)
    ws = load_workbook(io.BytesIO(buf.getvalue())).active
    assert ws.title == "Sheet1" and [c.value for c in ws[1]] == P.SAP_HEADERS and len(P.SAP_HEADERS) == 22 and ws.max_row == 1 + len(rows)


def test_an_eskom_bill_through_the_profile_posts_exactly_as_the_harvester_did():
    from cbx.pdfstudio import municipal as MU

    got = MU.read(eskom_pdf(), "eskom.pdf")
    assert got["profile"]["profile_id"] == "eskom"
    elec = [r for r in got["rows"] if r["service"] == "electricity"]
    acct = {"vendor": "700999", "cc": "01306", "name1": "Badplaas ETE", "muni": "ESKOM"}
    want, _, _ = P.generate_sap(elec, {**_tables(), "vendors": {"5012345678": acct}})
    f = C.utility_fields(got)
    assert f["invoice_no"] == "700123456" and f["services"] == [{"service": "electricity", "ex": "10000.00", "vat": "1500.00", "how": "the profile's total excl. VAT", "sure": True}]
    rows, status, issues = P.utility_doc(f, C._jsonable(elec[0]), acct, _tables(), file="eskom.pdf")
    assert rows == want and status == "GRN" and not issues and rows[0]["amount"] == "11512.34"


# ---- Part 2: the other services, and manual invoices ------------------------------------------------------------------
def test_other_services_post_to_their_gl_and_a_blank_gl_is_red():
    f = {"account": "5012345678", "billing_date": "2026-09-15", "due_date": "2026-10-05", "invoice_no": "INV-77", "statement_month": "SEPTEMBER 2026",
         "read_start": "2026-08-01", "read_end": "2026-08-31", "reading_type": "Actual", "kwh": "100",
         "services": [{"service": "electricity", "ex": "1000.00", "vat": "150.00"}, {"service": "water", "ex": "200.00", "vat": "30.00"},
                      {"service": "rates", "ex": "500.00", "vat": None}, {"service": "sewerage", "ex": "-50.00", "vat": "-7.50"}],
         "adjustments": {"interest": "10.00"}}
    acct = {"vendor": "700999", "cc": "01306", "name1": "Badplaas ETE", "muni": "X"}
    rows, status, issues = P.utility_doc(f, None, acct, _tables(), file="m.pdf")
    assert status == "GRN" and not issues, issues
    gl = {(r["account"], r["pkey"]): r for r in rows}
    assert rows[0]["pkey"] == "31" and rows[0]["amount"] == "1832.50"  # 1150 + 230 + 500 + (-57.50) + 10
    assert gl[("SA210", "40")]["amount"] == "200.00" and gl[("SA010", "40")]["tc"] == "I0" and gl[("SA312", "50")]["amount"] == "50.00"
    assert gl[("JA600", "40")]["amount"] == "172.50" and gl[("JA600", "40")]["dt"] == "X" and gl[("SF101", "40")]["tc"] == "I0"
    assert P.balance_fails(rows) == (1, set())
    # sundry charges have no GL until an admin sets one: red, never a guess
    f2 = {**f, "services": f["services"] + [{"service": "sundry", "ex": "25.00", "vat": "3.75"}]}
    rows, status, issues = P.utility_doc(f2, None, acct, _tables(), file="m.pdf")
    assert status == "RED" and any("no GL for sundry" in i for i in issues)


def test_a_manual_invoice_and_a_credit_note_become_documents():
    coding = {"vendor_no": "300100", "name": "Acme Cleaning", "gl": "6100100", "cost_centre": "CC77", "tax_code": "I9"}
    f = {"supplier": "Acme Cleaning (Pty) Ltd", "invoice_no": "AC-1", "invoice_date": "2026-10-01", "due_date": "2026-10-31", "currency": "ZAR",
         "lines": [{"description": "Cleaning", "excl": "800.00", "vat_rate": 15, "vat": "120.00"}, {"description": "Levy (zero rated)", "excl": "200.00", "vat_rate": 0, "vat": "0.00"}],
         "excl": "1000.00", "vat": "120.00", "incl": "1120.00", "credit": False}
    rows, status, issues = P.manual_doc(f, coding, "X143", _tables(), file="ac.pdf")
    assert status == "GRN" and not issues, issues
    assert [(r["pkey"], r["account"], r["amount"], r["tc"]) for r in rows] == [("31", "300100", "1120.00", ""), ("40", "6100100", "800.00", "I9"),
                                                                              ("40", "6100100", "200.00", "I0"), ("40", "JA600", "120.00", "I9")]
    assert rows[0]["comp"] == "X143" and rows[0]["docdate"] == "01102026" and rows[0]["DueDate"] == "31102026" and rows[3]["dt"] == "X"
    assert P.balance_fails(rows) == (1, set())
    rows, status, issues = P.manual_doc({**f, "credit": True}, coding, "X143", _tables(), file="cn.pdf", doctypes=("YR", "YG"))
    assert status == "YLW" and rows[0]["type"] == "YG" and [r["pkey"] for r in rows] == ["21", "50", "50", "50"] and P.balance_fails(rows) == (1, set())
    rows, status, issues = P.manual_doc({**f, "due_date": ""}, {**coding, "gl": ""}, "X143", _tables(), file="x.pdf")
    assert status == "RED" and "no GL for this vendor" in issues and "no due date on the invoice" in issues


def test_tess_tests_every_document():
    rows, _, _ = P.manual_doc({"supplier": "Acme", "invoice_no": "T-1", "invoice_date": "2026-10-01", "due_date": "2026-10-31", "excl": "100.00", "vat": "15.00", "incl": "115.00"},
                              {"vendor_no": "300100", "gl": "6100100", "cost_centre": "CC1"}, "X143", _tables(), file="t.pdf")
    known = {"vendors": {"300100"}, "gls": {"6100100", "JA600"}, "companies": {"X143"}, "doc_types": {"KR", "KG"}, "tax_codes": {"I9", "I0"}}
    assert P.test_doc(rows, known) == ("green", [])
    st, why = P.test_doc(rows, {**known, "vendors": set(), "gls": {"JA600"}})
    assert st == "red" and {"vendor 300100 is not in the vendor tables", "GL 6100100 is not in the GL tables"} <= {x["text"] for x in why}
    bad = [dict(r) for r in rows]
    bad[1]["amount"] = "99.00"
    bad[0]["DueDate"] = "31022026"
    st, why = P.test_doc(bad, known)
    assert st == "red" and any("do not add up" in x["text"] for x in why) and any("not a real date" in x["text"] for x in why)
    st, why = P.test_doc(rows, known, pip_status=P.YLW, pip_issues=["no due date on the invoice"])
    assert st == "yellow" and why == [{"level": "yellow", "text": "no due date on the invoice"}]
    assert P.test_doc(rows, known, dup="same invoice number and vendor as a.pdf in this batch")[0] == "red"


def test_the_capture_checks_send_a_doubtful_reading_up():
    ok, f, conf, probs = C.parse(_invoice_answer("Q-1"), 80)
    assert ok and not probs and conf == 95 and f["incl"] == "11500.00" and f["lines"][0]["excl"] == "10000.00"
    _, _, _, probs = C.parse(_invoice_answer("Q-1", conf=50), 80)
    assert probs == ["not sure enough (50%)"]
    _, _, _, probs = C.parse(_invoice_answer("Q-1", incl="11600.00"), 80)
    assert any("is not the total" in p for p in probs)
    _, _, _, probs = C.parse(_invoice_answer("Q-1", lines=[{"description": "x", "amount_excl": "9000.00", "vat_rate": 15, "vat": "1350"}]), 80)
    assert "the lines do not add up to the excl. total" in probs
    _, _, _, probs = C.parse(_invoice_answer("Q-1", vat="900.00", incl="10900.00", lines=[]), 80)
    assert any("VAT 900.00 is not 15%" in p for p in probs)
    assert C.parse({"is_invoice": False}, 80)[0] is False
    _, _, _, probs = C.parse({**_invoice_answer("Q-1"), "invoice_no": None, "invoice_date": "2031-01-01"}, 80)
    assert "no invoice number" in probs and any("not plausible" in p for p in probs)


# ---- the master lists ----------------------------------------------------------------------------------------------------
def test_the_lists_are_seeded_and_admins_replace_them_with_a_preview():
    adm = _admin()
    m = adm.get("/api/ap/masters").json()
    assert m["counts"]["entities"] >= 30 and {d["code"] for d in m["doc_types"]} >= {"KR", "YR", "KG", "YG", "KZ", "ZV", "AB", "KA", "SA", "UE", "DR"}
    assert {r["key"]: r["value"] for r in m["sap"]["fixed"]} == {**dict(P.SAP_DEFAULT_FIXED), "TaxCodeExempt": ""}  # 1.0.102: + TaxCodeExempt (empty)
    assert {r["key"]: r["value"] for r in m["sap"]["gl"]}["sanitation"] == "SA312" and m["models"]["capture"]
    ents = adm.get("/api/ap/masters/entities?q=X143").json()["rows"]
    assert ents[0]["name"] == "CBRE Excellerate Faciliti" and ents[0]["entity"] == "EFM SA"
    # the utilities list in the master-list shape (Supplier, PID, Reference, Site Name, PID, Vendor Name …): preview first
    master = _xlsx([["Supplier", "PID", "Reference", "Site Name", "PID", "Vendor Name", "District"],
                    [700002, 1306, 6604, "Badplaas ETE", 1306, "CHIEF ALBERT LUTHULI MUNICIPALITY", "North Eastern"],
                    [700006, 8198, 2100090266, "Andries Pretorius  Cabinet", 8198, "EKURHULENI METRO (ALBERTON)", "Gauteng"],
                    [700007, 8199, None, "No account", 8199, "X", "Y"]])
    r = adm.post("/api/ap/masters/utilities/import", files={"file": ("Utilities_Master_List.xlsx", master)}, data={"confirm": "false"})
    assert r.status_code == 200 and r.json()["saved"] is False and r.json()["rows"] == 2 and r.json()["skipped"] == 1
    before = adm.get("/api/ap/masters").json()["counts"]["utilities"]
    assert before == adm.get("/api/ap/masters").json()["counts"]["utilities"]  # nothing saved by a preview
    r = adm.post("/api/ap/masters/utilities/import", files={"file": ("Utilities_Master_List.xlsx", master)}, data={"confirm": "true"})
    assert r.json()["saved"] and adm.get("/api/ap/masters").json()["counts"]["utilities"] == 2
    row = adm.get("/api/ap/masters/utilities?q=2100090266").json()["rows"][0]
    assert row["vendor"] == "700006" and row["cost_centre"] == "8198" and row["site"] == "Andries Pretorius  Cabinet" and row["municipality"].startswith("EKURHULENI")
    with SessionLocal() as s:
        assert s.query(ApUtilityAccount).filter(ApUtilityAccount.account_key == "6604").one().extra == {"District": "North Eastern"}
    # an entities sheet: what would change, then saved
    sheet = _xlsx([["Company Code", "Company Name", "Entity Name"], ["X143", "CBRE Excellerate Facilities Management", "EFM SA"], ["X999", "Test Co", "Test"]])
    r = adm.post("/api/ap/masters/entities/import", files={"file": ("codes.xlsx", sheet)}, data={"confirm": "false"}).json()
    assert r["added"] == 1 and r["changed"] == 1 and r["removed"] >= 28 and not r["saved"]
    # SAP values: the Harvester's sap_config.xlsx out and back in
    cfg = adm.get("/api/ap/sap-values/file")
    wb = load_workbook(io.BytesIO(cfg.content))
    assert wb.sheetnames == ["Fixed", "GL Codes"] and [c.value for c in wb["GL Codes"][1]] == ["Charge", "GL", "Description"]
    wb["Fixed"]["B4"] = "KR"
    buf = io.BytesIO()
    wb.save(buf)
    r = adm.post("/api/ap/sap-values/import", files={"file": ("sap_config.xlsx", buf.getvalue())}, data={"confirm": "true"}).json()
    assert r["saved"] and r["removed"] == 0
    # people without admin read the lists but never change them; pulling from SAP needs a connection ticked for the AP Tool
    u = _user()
    assert u.get("/api/ap/masters").status_code == 200 and "connections" not in u.get("/api/ap/masters").json()
    assert u.post("/api/ap/masters/entities/import", files={"file": ("c.xlsx", sheet)}, data={"confirm": "true"}).status_code == 403
    assert u.put("/api/ap/sap-values", json={"fixed": [], "gl": []}).status_code == 403
    r = adm.post("/api/ap/masters/utilities/pull", json={"connection_id": 999, "path": "/x", "columns": {"account": "A", "vendor": "B"}})
    assert r.status_code == 400 and "Admin → Connections" in r.json()["detail"]


def test_learned_vendors_stay_when_the_list_is_replaced():
    adm = _admin()
    with SessionLocal() as s:
        from cbx.models import User

        me = s.query(User).filter(User.email == "cap.admin@test.local").one()
        M.learn_vendor(s, me, {"vendor_no": "900001", "name": "Learned Supplies", "vat_no": "4555555555", "gl": "6200100"})
        s.commit()
    sheet = _xlsx([["Vendor", "Name", "VAT No", "GL", "Cost centre", "Tax code"], ["300100", "Acme Cleaning (Pty) Ltd", "4123456789", "6100100", "CC77", "I9"]])
    r = adm.post("/api/ap/masters/vendors/import", files={"file": ("vendors.xlsx", sheet)}, data={"confirm": "true"}).json()
    assert r["saved"] and r["kept_learned"] >= 1
    rows = adm.get("/api/ap/masters/vendors?q=Learned").json()["rows"]
    assert rows and rows[0]["learned"] and rows[0]["learned_by"] == "Ayanda Admin"
    wb = load_workbook(io.BytesIO(adm.get("/api/ap/masters/vendors/file").content))
    assert any(r[0].value == "900001" and r[8].value == "Ayanda Admin" for r in wb.active.iter_rows(min_row=2))


# ---- the flow, through the page's API ---------------------------------------------------------------------------------------
def test_a_mixed_folder_is_sorted_read_matched_built_tested_and_downloaded(monkeypatch, utilities):
    adm = utilities
    run = uuid.uuid4().hex[:6].upper()
    good, hard, newv = f"CAP-{run}-1", f"CAP-{run}-2", f"CAP-{run}-3"
    with SessionLocal() as s:
        s.query(ApVendorCode).filter(ApVendorCode.vat_no == "4123456789").delete()
        s.add(ApVendorCode(vendor_no="300100", name="Acme Cleaning (Pty) Ltd", name_key=M.name_key("Acme Cleaning (Pty) Ltd"), vat_no="4123456789", gl="6100100",
                           cost_centre="CC77", tax_code="I9", source="excel"))
        x143 = s.query(ApEntity).filter(ApEntity.code == "X143").one()  # 2.10.0: SAP's 25-character name is only the start of the bill-to (yellow): the full name is taught
        if "CBRE Excellerate Facilities Management" not in x143.aliases:
            x143.aliases = (x143.aliases + "\n" if x143.aliases else "") + "CBRE Excellerate Facilities Management"
        s.commit()
    seth = {"n": 0}

    def hard_one(task):
        seth["n"] += task == "ap_capture_senior"
        return _invoice_answer(hard, conf=40)  # neither the capture bot nor Seth is sure

    fake = _vision({"acme_good.pdf": _invoice_answer(good), "acme_hard.pdf": hard_one,
                    "newco.pdf": _invoice_answer(newv, supplier="New Co Trading", vat_no="4777777777", excl="200.00", vat="30.00", incl="230.00")})
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", fake)
    u = _user()
    folder = {"eskom_sep.pdf": eskom_pdf(), "eskom_unknown.pdf": eskom_pdf(account="5099999999", invoice="700555"), "coj_mar.pdf": coj_pdf(),
              "acme_good.pdf": make_ap.invoice(no=good), "acme_hard.pdf": make_ap.invoice(no=hard), "newco.pdf": make_ap.invoice(no=newv, supplier="New Co Trading"),
              "canteen_menu.txt": b"Canteen: Monday soup.", "run.xlsx": make_ap.payment_run()}
    b = _drop(u, folder, f"Capture {run}")
    cap = b["capture"]
    assert cap["files"] == 8 and cap["sorted"] == 8
    # two Eskom bills and the CoJ bill are utilities, but one Eskom account is not in the list: Other; the new vendor waits
    # in manual capture for a person to code it
    assert {k: v["n"] for k, v in cap["buckets"].items() if v["n"]} == {"utility": 2, "manual": 3, "other": 1, "odd": 1, "ap": 1}
    kinds = {i["name"]: (i["kind"], i["bot_name"]) for i in b["items"]}
    assert kinds["eskom_sep.pdf"] == ("utility", "Uma") and kinds["run.xlsx"][0] == "payment_run"
    # the three capture bots read; the hard one went to Seth with every page, then to a person
    tasks = [(t, n) for t, n, _, _ in fake.calls]
    assert ("ap_capture", "acme_good.pdf") in tasks and ("ap_capture_senior", "acme_hard.pdf") in tasks and seth["n"] == 1
    docs = _sheet(u, b["id"])
    assert docs["eskom_sep.pdf"]["status"] == "green" and docs["eskom_sep.pdf"]["vendor"] == "700999" and docs["eskom_sep.pdf"]["amount"] == "11512.34"
    assert docs["acme_good.pdf"]["status"] == "green" and docs["acme_good.pdf"]["entity"] == "X143" and docs["acme_good.pdf"]["gl"] == "6100100, JA600"
    assert docs["coj_mar.pdf"]["status"] == "red" and any("no invoice number" in x["text"] for x in docs["coj_mar.pdf"]["issues"])
    assert docs["eskom_unknown.pdf"]["bucket"] == "other" and docs["eskom_unknown.pdf"]["needs"] == "entity"
    hist = " ".join(h["who"] + ": " + h["text"] for h in u.get(f"/api/ap/items/{docs['acme_good.pdf']['item_id']}").json()["history"])
    assert "Sol:" in hist and "Mia:" in hist and "Pip:" in hist and "Tess:" in hist and any(w in hist for w in ("Cara:", "Dex:", "Eli:"))
    item = u.get(f"/api/ap/items/{docs['acme_good.pdf']['item_id']}").json()
    assert item["fields"]["invoice_no"] == good and item["fields"]["total"] == 11500.0  # what the bot read is on the AP item, and Ivy checked it
    # Needs you: each section has what it should
    n = u.get("/api/ap/needs").json()
    ids = lambda sec: {x["file"].split("/")[-1] for x in n[sec]}  # noqa: E731
    assert "eskom_unknown.pdf" in ids("other") and "canteen_menu.txt" in ids("odd") and "acme_hard.pdf" in ids("unread") and "coj_mar.pdf" in ids("red")
    vend = next(g for g in n["vendors"] if g["supplier"] == "New Co Trading")
    # a person codes the vendor: learned, and the invoice moves on
    r = u.post("/api/ap/vendors/code", json={"ids": vend["ids"], "vendor_no": "300777", "gl": "6100100", "cost_centre": "CC12", "tax_code": "I9"})
    assert r.status_code == 200 and r.json()["vendor"]["learned"] and r.json()["vendor"]["learned_by"] == "Cara Person"
    # 2.10.0: a person's coding is yellow until an admin confirms it (Needs you, for admins)
    assert _sheet(u, b["id"])["newco.pdf"]["status"] == "yellow"
    learned = next(x for x in adm.get("/api/ap/needs").json()["confirm"] if x["kind"] == "vendors" and x["id"] == r.json()["vendor"]["id"])
    assert adm.post(f"/api/ap/masters/vendors/{learned['id']}/confirm").status_code == 200
    assert _sheet(u, b["id"])["newco.pdf"]["status"] == "green"
    # a person types the one the senior couldn't read
    hard_c = next(x for x in n["unread"] if x["file"].endswith("acme_hard.pdf"))
    png = u.get(f"/api/ap/captures/{hard_c['id']}/page/1")
    assert png.status_code == 200 and png.headers["content-type"] == "image/png"
    r = u.post(f"/api/ap/captures/{hard_c['id']}/decide", json={"action": "save", "data": {"fields": {"supplier": "Acme Cleaning (Pty) Ltd", "supplier_vat": "4123456789",
               "bill_to": "CBRE Excellerate Facilities Management", "invoice_no": hard, "invoice_date": (date.today() - timedelta(days=3)).isoformat(),
               "due_date": (date.today() + timedelta(days=27)).isoformat(), "excl": "10000.00", "vat": "1500.00", "incl": "11500.00"}}})
    assert r.status_code == 200, r.text
    assert r.json()["read_by"] == "person" and r.json()["sap_status"] == "green"
    # CoJ had no invoice number: a person corrects it, it goes green
    coj = docs["coj_mar.pdf"]
    r = u.post(f"/api/ap/captures/{coj['id']}/decide", json={"action": "save", "data": {"fields": {"invoice_no": "COJ-123"}}})
    assert r.json()["sap_status"] == "green", r.json()["sap_issues"]
    # Other: the unknown Eskom account is added to the list and continues as a utility — 1.0.101: by an admin (Add to the
    # utilities list); a person without admin can't. 1.0.104 (Valuence F0b): an admin's own addition is learned (yellow) too,
    # she can't confirm it herself, a second admin does
    other = next(x for x in n["other"] if x["file"].endswith("eskom_unknown.pdf"))
    assert u.post(f"/api/ap/captures/{other['id']}/decide", json={"action": "manual", "data": {"vendor": "700123", "cost_centre": "01400", "site": "New Site"}}).status_code == 403
    r = adm.post(f"/api/ap/captures/{other['id']}/decide", json={"action": "manual", "data": {"vendor": "700123", "cost_centre": "01400", "site": "New Site"}})
    assert r.status_code == 200 and r.json()["bucket"] == "utility" and r.json()["sap_status"] == "yellow"
    row = next(x for x in adm.get("/api/ap/needs").json()["confirm"] if x["kind"] == "utilities" and x["vendor"] == "700123")
    assert adm.post(f"/api/ap/masters/utilities/{row['id']}/confirm").status_code == 403
    assert signed_in("cap.admin2@test.local", "admin", name="Bongani Admin").post(f"/api/ap/masters/utilities/{row['id']}/confirm").status_code == 200
    assert u.get(f"/api/ap/captures/{other['id']}").json()["sap_status"] == "green"
    # Odd: skipped
    odd = next(x for x in n["odd"] if x["file"].endswith("canteen_menu.txt"))
    assert u.post(f"/api/ap/captures/{odd['id']}/decide", json={"action": "skip", "data": {}}).json()["bucket"] == "skipped"
    # the SAP upload: exactly the Harvester's file, every green document, each only once
    r = u.post(f"/api/ap/batches/{b['id']}/sap/upload", json={"include_yellow": False})
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.title == "Sheet1" and [c.value for c in ws[1]] == P.SAP_HEADERS
    refs = {row[6].value for row in ws.iter_rows(min_row=2)}
    assert {good, hard, newv, "700123456", "COJ-123", "700555"} == refs
    assert sum(1 for row in ws.iter_rows(min_row=2) if row[8].value == "31") == 6
    r = u.post(f"/api/ap/batches/{b['id']}/sap/upload", json={"include_yellow": True})
    assert r.status_code == 400 and "Nothing here is ready" in r.json()["detail"]
    assert all(d["uploaded"] for d in _sheet(u, b["id"]).values() if d["status"] == "green")
    # the same invoice in a later folder is red: it already went out
    monkeypatch.setattr(C, "ask_vision", _vision({"acme_again.pdf": _invoice_answer(good)}))
    b2 = _drop(u, {"acme_again.pdf": make_ap.invoice(no=good)}, f"Again {run}")
    again = _sheet(u, b2["id"])["acme_again.pdf"]
    assert again["status"] == "red" and any("already went out" in x["text"] for x in again["issues"])
    # the review workbook has every row with its status and the source file
    wb = load_workbook(io.BytesIO(u.get(f"/api/ap/batches/{b['id']}/sap/review").content))
    heads = [c.value for c in wb.active[1]]
    n = 4 + len(P.SAP_HEADERS)  # 1.0.104 (Builder R): the review's four columns come after SAP's names
    assert heads[:4] == ["Status", "Why", "Source file", "Uploaded"] and heads[4:n] == P.SAP_HEADERS and heads[n:n + 2] == ["G/L name", "Cost centre name"]  # 2.10.3: + SAP's names
    assert heads[n + 2:] == ["Selected why", "Reviewer", "Reviewed", "Outcome"]
    # only an admin releases the upload mark
    assert u.post(f"/api/ap/batches/{b['id']}/sap/release").status_code == 403
    assert adm.post(f"/api/ap/batches/{b['id']}/sap/release").json()["released"] == 6


def test_other_odd_and_a_new_profile(monkeypatch, utilities):
    run = uuid.uuid4().hex[:6].upper()
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision({
        "faraway.pdf": _invoice_answer(f"FA-{run}", supplier="Acme Cleaning (Pty) Ltd", bill_to="Somebody Else Holdings"),
        "not_an_invoice.pdf": {"is_invoice": False},
        "odd_scan.pdf": _invoice_answer(f"OS-{run}"), "unknown_billto.pdf": _invoice_answer(f"UB-{run}", supplier="Acme Cleaning (Pty) Ltd", bill_to="Nobody We Know Ltd")}))
    u = _user("cap.other@test.local", "Olu Other")
    b = _drop(u, {"faraway.pdf": make_ap.invoice(no=f"FA-{run}"), "not_an_invoice.pdf": make_ap.invoice(no=f"NI-{run}"), "odd.txt": b"Hello there"}, f"Other {run}")
    n = u.get("/api/ap/needs").json()
    # 1.0.105: a bill-to that is no company in the list no longer waits in Other when it shares a word with one of the
    # group's own entities ("Holdings" — Excellerate Holdings, Z102): that one, as a yellow guess (cbx/ap/autofill.py)
    assert not [x for x in n["other"] if x["batch_id"] == b["id"]]
    far = _sheet(u, b["id"])["faraway.pdf"]
    ent = u.get(f"/api/ap/captures/{far['id']}").json()["match"]["entity"]
    assert far["entity"] == "Z102" and (ent["how"], ent["unsure"]) == ("a word of the bill-to names the group's Z102 (holdings)", True)
    assert {x["file"] for x in n["odd"] if x["batch_id"] == b["id"]} == {"not_an_invoice.pdf", "odd.txt"}
    # Odd → "not an invoice" goes back to the AP team's flow; another Odd → manual capture, read by the bots
    odd = {x["file"]: x for x in n["odd"] if x["batch_id"] == b["id"]}
    assert u.post(f"/api/ap/captures/{odd['odd.txt']['id']}/decide", json={"action": "ap"}).json()["bucket"] == "ap"
    r = u.post(f"/api/ap/captures/{odd['not_an_invoice.pdf']['id']}/decide", json={"action": "profile"})
    assert r.status_code == 400  # only an Other file waits for a profile
    # 1.0.101: a bill PDF Studio already reads (Eskom) never waits for a new profile: refused in plain words
    b2 = _drop(u, {"unknown_munic.pdf": eskom_pdf(account="5088888888", invoice="700777")}, f"Profile {run}")
    o = next(x for x in u.get("/api/ap/needs").json()["other"] if x["batch_id"] == b2["id"])
    r = u.post(f"/api/ap/captures/{o['id']}/decide", json={"action": "profile"})
    assert r.status_code == 400 and r.json()["detail"].startswith("PDF Studio already reads these bills (profile Eskom).")
    # New profile (a bill no profile knows): PDF Studio's profile keeper hears about it; Read again routes it once a profile knows the bill
    b3 = _drop(u, {"unknown_billto.pdf": make_ap.invoice(no=f"UB-{run}")}, f"Profile2 {run}")
    with SessionLocal() as s:  # 1.0.105: Mia no longer sends an unknown bill-to to Other (the default company, yellow): one left there by 1.0.104
        x = s.query(ApCapture).filter(ApCapture.batch_id == b3["id"]).one()
        x.bucket, x.stage, x.needs, x.match = "other", "needs", "entity", {"entity": {"found": False, "bill_to": "Nobody We Know Ltd"}}
        s.commit()
    o = next(x for x in u.get("/api/ap/needs").json()["other"] if x["batch_id"] == b3["id"])
    r = u.post(f"/api/ap/captures/{o['id']}/decide", json={"action": "profile"})
    assert r.status_code == 200 and r.json()["needs"] == "profile"
    with SessionLocal() as s:
        assert sum(1 for x in C2.notes(s, "ap") if x["ref"] == f"ap_item:{o['item_id']}" and "needs a PDF Studio profile" in x["text"]) == 1  # 1.0.98: the AP team view
    r = u.post(f"/api/ap/captures/{o['id']}/decide", json={"action": "again"})
    assert r.status_code == 200
    again = u.get(f"/api/ap/captures/{o['id']}").json()
    # read again: still no profile, and the bill-to is still no company code; 1.0.105: it names no group entity either, so
    # never the default X143 — no company, red, a person picks it
    assert again["bucket"] == "other" and again["needs"] == "entity" and again["company_code"] == "" and again["sap_status"] == "red"


def test_access_and_the_team_view(monkeypatch):
    make_person("cap.none@test.local", name="No Ap")
    assert signed_in("cap.none@test.local").get("/api/ap/live").status_code == 403
    viewer = signed_in("cap.viewer@test.local", "viewer", name="Vee Viewer", tools=["ap"])
    assert viewer.get("/api/ap/needs").status_code == 200
    with SessionLocal() as s:
        cid = s.query(ApCapture.id).order_by(ApCapture.id.desc()).first()
    if cid:
        assert viewer.post(f"/api/ap/captures/{cid[0]}/decide", json={"action": "skip"}).status_code == 403
    assert viewer.post("/api/ap/vendors/code", json={"ids": [1], "vendor_no": "1", "gl": "1"}).status_code == 403
    u = _user()
    live = u.get("/api/ap/live").json()
    assert [x["name"] for x in live["bots"]] == ["Sol", "Uma", "Cara", "Dex", "Eli", "Seth", "Mia", "Pip", "Tess"]
    assert all("model" not in x for x in live["bots"]) and all(x["status"] in ("idle", "working", "waiting", "paused") for x in live["bots"])
    adm = _admin()
    assert all(x.get("model") for x in adm.get("/api/ap/live").json()["bots"])
    assert u.put("/api/ap/capture-settings", json={"capture_model": "x"}).status_code == 403
    r = adm.put("/api/ap/capture-settings", json={"capture_model": "qwen3-vl:32b", "min_confidence": 85})
    assert r.status_code == 200 and r.json()["capture_model"] == "qwen3-vl:32b" and r.json()["min_confidence"] == 85
    assert adm.put("/api/ap/capture-settings", json={"capture_model": "bad model; rm"}).status_code == 400
    adm.put("/api/ap/capture-settings", json={"capture_model": "", "min_confidence": 80})
    page = u.get("/static/views/ap.js")  # 1.0.98: the old SPA's AP view
    assert page.status_code == 200 and "The team" in page.text


def test_rae_tells_document_types_apart():
    rows = [{"vendor": "100", "name": "Acme", "ref": "INV-1", "amount": 1000.0, "date": date(2026, 9, 1), "due": None, "doc_type": "KR"},
            {"vendor": "100", "name": "Acme", "ref": "CN-1", "amount": -200.0, "date": date(2026, 9, 2), "due": None, "doc_type": "KG"},
            {"vendor": "100", "name": "Acme", "ref": "PAY-1", "amount": -500.0, "date": date(2026, 9, 3), "due": None, "doc_type": "KZ"},
            {"vendor": "100", "name": "Acme", "ref": "J-1", "amount": 50.0, "date": date(2026, 9, 4), "due": None, "doc_type": "ZZ"}]
    with SessionLocal() as s:
        table = M.doc_type_map(s)
    found = {f["code"]: f for f in E.check("ledger", {"rows": rows}, settings={"doc_types": table}, today=date(2026, 10, 15))}
    assert "invoices 1 (R1,000.00)" in found["by_doc_type"]["text"] and "credit memos 1" in found["by_doc_type"]["text"]
    assert found["payments_open"]["severity"] == "check" and "ZZ" in found["unknown_doc_type"]["text"]
    assert E.find_col(["Document Type", "Amount"], "doc_type") == "Document Type"
    assert "by_doc_type" not in {f["code"] for f in E.check("ledger", {"rows": rows}, today=date(2026, 10, 15))}  # without the table: as before


# ---- 1.0.98: test_gateway_210.py's intent on the old gateway: a tool's own model reaches the call ---------------------------
def test_a_tools_own_model_reaches_the_call_and_the_capture_tasks_have_their_classes(monkeypatch, db):
    from cbx.cexa import gateway as G

    sent = []
    g = G.Gateway(db)
    g.mode = "live"
    monkeypatch.setattr(g, "_ollama", lambda model, *a, **k: sent.append(model) or '{"x": 1}')
    assert g.json("ap_capture", "s", "u", {"x": "number"}, model="my-vision-model") == {"x": 1}
    g.json("ap_capture", "s", "u", {"x": "number"})
    g.text("ap_capture_senior", "s", "u", model="big-vision")
    assert sent == ["my-vision-model", g.model_for("ap_capture"), "big-vision"]
    assert G.TASK_MODEL["ap_capture"] == "designer" and G.TASK_MODEL["ap_capture_senior"] == "designer" and G.TASK_MODEL["ap_match"] == "helper"
    # the capture settings' empty models are the gateway's choice (the senior: C-EXA's own); the page shows them to admins
    models = C.default_models()
    assert models["capture"] == g.model_for("ap_capture") and models["senior"] == g.model_for("ap_capture_senior")
    seen = []
    monkeypatch.setattr(G.Gateway, "json", lambda self, task, system, user, schema, images=None, model="", **k: seen.append((task, model, len(images or []))) or {})
    C.ask_vision("ap_capture_senior", C.SENIOR, "File name: x.pdf", [b"png"], model="")
    assert seen == [("ap_capture_senior", "", 1)]  # "" = the gateway picks (the eyes model for the senior, 1.0.110)


# ---- 1.0.98: the old CBX's side of it (cbx/compat2.py): the Office, the page's live view, jobs, the seed ---------------------
def test_the_bots_show_in_the_office_as_ap_and_on_the_page_by_their_short_name(monkeypatch, owner):
    from cbx.cexa.teams import ensure_teams
    from cbx.models import AgentWork, TeamMember

    with SessionLocal() as s:
        ensure_teams(s)
        s.commit()
    run = uuid.uuid4().hex[:6].upper()
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision({f"office_{run}.pdf": _invoice_answer(f"OF-{run}")}))
    u = _user("cap.office@test.local", "Otto Office")
    seen = {}
    real = C.step

    def step(worker):  # while Cara reads, she is "working" on the page's live view and in the Office
        if worker in C.WORKERS and worker not in seen:
            orig = C.ask_vision

            def look(*a, **k):
                seen[worker] = C2.status_of([C.office_name(worker)])[C.office_name(worker)]
                with SessionLocal() as s:
                    seen["office"] = [(w.title, w.status) for w, m in s.query(AgentWork, TeamMember).join(TeamMember, TeamMember.id == AgentWork.to_member_id)
                                      .filter(TeamMember.name == C.office_name(worker), AgentWork.status == "working").all()]
                return orig(*a, **k)

            monkeypatch.setattr(C, "ask_vision", look)
        return real(worker)

    monkeypatch.setattr(C, "step", step)
    _drop(u, {f"office_{run}.pdf": make_ap.invoice(no=f"OF-{run}")}, f"Office {run}")
    reader = next(k for k in seen if k in C.WORKERS)
    assert seen[reader]["state"] == "working" and seen[reader]["task"] == f"Reading office_{run}.pdf"
    assert seen["office"] == [(f"Reading office_{run}.pdf", "working")]
    with SessionLocal() as s:
        names = {m.name: m for m in s.query(TeamMember).filter(TeamMember.name.like("% (AP)")).all()}
        assert {"Sol (AP)", "Mia (AP)", C.office_name(reader)} <= set(names)  # (Pip and Tess wait: the vendor is not coded)
        assert all(m.role == "bot" for m in names.values())  # no hands: an unknown role only looks
        # the Office row goes when the work is over (never a "piece of work finished" in the owner's digest or on a plant)
        assert not s.query(AgentWork).filter(AgentWork.source == "tool").count()
        assert C2.status_of(["Sol (AP)"], s)["Sol (AP)"]["done_today"] >= 1  # counted in app_state (it survives a restart)
    live = u.get("/api/ap/live").json()
    by = {b["key"]: b for b in live["bots"]}
    assert by["sol"]["name"] == "Sol" and by["sol"]["status"] in ("idle", "waiting") and by[reader]["done_today"] >= 1
    ap = next(t for t in owner.get("/api/cexa/teams").json()["teams"] if t["key"] == "ap")  # their desks in the Office
    assert {"Sol (AP)", "Mia (AP)"} <= {m["name"] for m in ap["members"]}


def test_jobs_dedupe_by_ref_and_a_rolled_back_ref_does_not_block(db):
    ran = []
    C2.handler("t1098")(lambda j, p: ran.append(p["n"]))
    assert not C2.enqueue(db, tool="ap", kind="t1098", title="t", ref="t1098:a", payload={"n": 1}).dropped
    assert C2.enqueue(db, tool="ap", kind="t1098", title="t", ref="t1098:a", payload={"n": 2}).dropped  # one is waiting already
    db.commit()
    assert ran == [1] and not C2.queued("t1098:a")
    C2.enqueue(db, tool="ap", kind="t1098", title="t", ref="t1098:b", payload={"n": 3})
    assert C2.queued("t1098:b")
    db.rollback()  # the request failed: the job never starts, and the ref is free again
    assert not C2.queued("t1098:b")
    db.info.pop("after_commit", None)
    C2.enqueue(db, tool="ap", kind="t1098", title="t", ref="t1098:b", payload={"n": 4})
    db.commit()
    assert ran == [1, 4]


def test_a_reader_with_an_empty_queue_stops_and_the_tidy_closes_what_a_restart_left(db):
    from cbx.cexa.teams import ensure_teams
    from cbx.models import AgentWork, TeamMember

    assert C.capture_job(0, {"worker": "cara"}) == {"read": 0}  # inline (tests): it ends when nothing waits
    ensure_teams(db)
    db.commit()
    with C2.working(None, "Dex (AP)", "Reading x.pdf"):
        pass
    mid = db.query(TeamMember.id).filter(TeamMember.name == "Dex (AP)").scalar()
    stale = AgentWork(kind="task", title="Reading y.pdf", status="working", to_member_id=mid, source="tool")  # left by a restart
    handed = AgentWork(kind="task", title="Please pay supplier X", status="open", to_member_id=mid, source="agent")  # another agent's work
    db.add_all([stale, handed])
    db.commit()
    assert C2.tidy(db) >= 2
    db.commit()
    db.expire_all()
    assert db.get(AgentWork, stale.id) is None and db.get(AgentWork, handed.id).status == "done"
    assert "only inside my own tool" in db.get(AgentWork, handed.id).result


def test_the_seed_runs_once_and_fills_sap_s_lists(db):
    import time

    from cbx.models.ap import ApCostCentre, ApDocType, ApGlAccount, ApSapValue, ApVendorCompany

    t0 = time.time()
    M.ensure_seeded()  # done already (the first AP request): nothing to do
    assert time.time() - t0 < 1
    assert db.query(ApEntity).filter(ApEntity.code == "X143").count() == 1 and db.query(ApDocType).filter(ApDocType.code == "KR").count() == 1
    assert db.query(ApSapValue).count() and db.query(ApGlAccount).count() > 1000 and db.query(ApVendorCompany).count() > 100000 and db.query(ApCostCentre).count() > 20000
    assert C2.state_get(db, "ap.seeded.entities", None)  # JSON over the old text column


def test_users_never_see_model_names(monkeypatch):
    u, adm = _user(), _admin()
    m = u.get("/api/ap/masters").json()
    assert m["models"] == {} and not any(k.endswith("_model") for k in m["capture"])
    am = adm.get("/api/ap/masters").json()
    assert am["models"]["capture"] and "capture_model" in am["capture"]
