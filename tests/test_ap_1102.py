"""1.0.102: from CBX-Connect 2.10.7 tests/test_ap_2107.py, on the old conftest.
(Edwin, 18 Oct): SAP rows for manual invoices — the batch's posting date (today unless a person sets another;
"today" is the day of the download), the cost centre from the bill-to's site, a G/L per line from the vendor coding's
line rules, the due date from the vendor's payment days, tax codes per line, SAP's vendor general data (LFA1) and a
header text that fits SAP's 25 characters. Suppliers, VAT numbers and amounts are made up; the AI is off in tests (the
capture call returns what a vision model would read)."""
from __future__ import annotations

import csv
import io
import uuid
from datetime import date, timedelta

import pytest
from openpyxl import Workbook, load_workbook

from cbx import knowledge as K
from cbx.ap import capture as C
from cbx.ap import masters as M
from cbx.ap import sap as P
from cbx.db import SessionLocal
from cbx.models import AuditLog
from cbx.models.ap import ApCapture, ApEntity, ApItem, ApSapValue, ApSapVendor, ApSite, ApVendorCode

from .test_ap_capture_1098 import _admin, _app, _drop, _sheet, _user, _vision, signed_in  # noqa: F401
from .test_ap_scans_1098 import brightline_pdf

FULL = "CBRE Excellerate Facilities Management"
BILL_VAT = "4999000011"
INV_DAY = date.today() - timedelta(days=12)
DUE_DAY = date.today() + timedelta(days=18)


@pytest.fixture(autouse=True)
def _x143():
    """X143's full name taught (a sure bill-to match), as the scans' tests do; nothing of SAP's vendor list or the sites
    is left behind for other tests."""
    with SessionLocal() as s:
        e = s.query(ApEntity).filter(ApEntity.code == "X143").one()
        if FULL not in (e.aliases or ""):
            e.aliases = (e.aliases + "\n" if e.aliases else "") + FULL
        s.commit()
    K.forget()
    yield
    with SessionLocal() as s:
        s.query(ApSapVendor).delete()
        s.query(ApSite).delete()
        from cbx.models.connections import ApiConnection

        s.query(ApiConnection).filter(ApiConnection.slug.like("lfa1-sap-%")).delete(synchronize_session=False)  # 1.0.102: test_connections counts them
        s.commit()
    K.forget()


def _vat() -> str:
    return "47" + str(uuid.uuid4().int)[:8]


def _code(name: str, vat: str, **kw) -> str:
    """An admin's vendor coding (not learned)."""
    no = kw.pop("vendor_no", "99" + str(uuid.uuid4().int)[:5])
    row = {"gl": "SC600", "cost_centre": "AC0328", "tax_code": "I9", "item_text": "", **kw}
    with SessionLocal() as s:
        s.query(ApVendorCode).filter(ApVendorCode.vat_no == vat).delete()
        s.add(ApVendorCode(vendor_no=no, name=name, name_key=M.name_key(name), vat_no=vat, source="excel", **row))
        s.commit()
    return no


def _answer(no: str, supplier: str, vat: str, *, bill_to: str = FULL, lines=None, excl="1000.00", vat_amt="150.00", incl="1150.00", due=True) -> dict:
    conf = {k: 93 for k in ("supplier", "bill_to", "invoice_no", "invoice_date", "total_incl", "total_excl", "vat")}
    return {"is_invoice": True, "credit_note": False, "supplier": supplier, "supplier_vat": vat, "bill_to": bill_to, "bill_to_vat": BILL_VAT, "invoice_no": no,
            "invoice_date": INV_DAY.isoformat(), "due_date": DUE_DAY.isoformat() if due else None, "po": "PO-1", "our_account": "", "currency": "ZAR",
            "lines": lines if lines is not None else [{"description": "Services", "amount_excl": excl, "vat_rate": 15, "vat": vat_amt}],
            "total_excl": excl, "vat": vat_amt, "total_incl": incl, "tax_invoice_words": True, "confidence": conf}


def _run(monkeypatch, u, answers: dict, name: str) -> dict:
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision(answers))
    return _drop(u, {n: brightline_pdf(n) for n in answers}, name)


def _cap(u, cid: int) -> dict:
    return u.get(f"/api/ap/captures/{cid}").json()


def _ccs(rows: list[dict]) -> set[str]:
    """The cost centres of a document's G/L lines (not the VAT line)."""
    return {r["costcentre"] for r in rows if r["pkey"] in ("40", "50") and r["account"] != "JA600"}


def _hist(item_id: int) -> list[str]:
    with SessionLocal() as s:
        return [h["text"] for h in s.get(ApItem, item_id).history or []]


def _tables(**fixed):
    return {"fixed": {**dict(P.SAP_DEFAULT_FIXED), "TaxCodeExempt": "", **fixed}, "gl": {c: g for c, g, _ in P.DEFAULT_GL}}


# ---- 1. the posting date -----------------------------------------------------------------------------------------------
def test_the_posting_date_is_today_a_day_a_person_sets_or_the_day_of_the_download(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    fri, mon = date(2026, 10, 16), date(2026, 10, 19)
    monkeypatch.setattr(P, "sap_today", lambda: fri)
    vat = _vat()
    _code("Posting Test (Pty) Ltd", vat)
    u, viewer = _user(), signed_in("pd.viewer@test.local", "viewer", name="Vera Viewer", tools=["ap"])
    b = _run(monkeypatch, u, {f"pd1_{run}.pdf": _answer(f"PD1-{run}", "Posting Test (Pty) Ltd", vat), f"pd2_{run}.pdf": _answer(f"PD2-{run}", "Posting Test (Pty) Ltd", vat)}, f"PD {run}")
    docs = _sheet(u, b["id"])
    assert {d["status"] for d in docs.values()} == {"green"}, [d["issues"] for d in docs.values()]
    # by default: today; the document's own date stays the invoice date
    rows = _cap(u, docs[f"pd1_{run}.pdf"]["id"])["sap_rows"]
    assert {r["postdate"] for r in rows} == {"16102026"} and {r["docdate"] for r in rows} == {f"{INV_DAY:%d%m%Y}"}
    sheet = u.get(f"/api/ap/batches/{b['id']}/sap").json()
    assert sheet["posting"] == {"date": "2026-10-16", "set": False, "by": "", "today": "2026-10-16", "far": False}
    # a viewer can't set it; a person who can download can, and every row carries it (the review workbook too)
    assert viewer.post(f"/api/ap/batches/{b['id']}/sap/posting-date", json={"date": "2026-10-14"}).status_code == 403
    r = u.post(f"/api/ap/batches/{b['id']}/sap/posting-date", json={"date": "2026-10-14"})
    assert r.status_code == 200 and r.json()["posting"]["set"] and r.json()["posting"]["date"] == "2026-10-14" and r.json()["posting"]["by"] == "Cara Person"
    for d in docs.values():
        assert {x["postdate"] for x in _cap(u, d["id"])["sap_rows"]} == {"14102026"}
    ws = load_workbook(io.BytesIO(u.get(f"/api/ap/batches/{b['id']}/sap/review").content)).active
    heads = [c.value for c in ws[1]]
    assert {r[heads.index("postdate")] for r in ws.iter_rows(min_row=2, values_only=True)} == {"14102026"}
    with SessionLocal() as s:
        a = s.query(AuditLog).filter(AuditLog.action == "ap_posting_date", AuditLog.target == f"ap_batch:{b['id']}").order_by(AuditLog.id.desc()).first()
        assert a is not None and a.detail.endswith("posting date today → 2026-10-14")
    # far from today: yellow, never red (before the 1st of last month, or more than 7 days ahead)
    for day, far in (("2026-08-31", True), ("2026-09-01", False), ("2026-10-23", False), ("2026-10-24", True)):
        r = u.post(f"/api/ap/batches/{b['id']}/sap/posting-date", json={"date": day}).json()
        st = {d["status"] for d in r["docs"]}
        want = f"posting date {date.fromisoformat(day):%d.%m.%Y} is far from today"
        assert st == ({"yellow"} if far else {"green"}) and r["posting"]["far"] is far, (day, r["docs"][0]["issues"])
        assert any(x["text"] == want for x in r["docs"][0]["issues"]) is far
    # "Today" resets it
    r = u.post(f"/api/ap/batches/{b['id']}/sap/posting-date", json={"date": ""}).json()
    assert r["posting"]["set"] is False and {d["status"] for d in r["docs"]} == {"green"}
    assert {x["postdate"] for x in _cap(u, docs[f"pd1_{run}.pdf"]["id"])["sap_rows"]} == {"16102026"}
    # dropped on Friday, left at today, downloaded on Monday: it posts on Monday
    monkeypatch.setattr(P, "sap_today", lambda: mon)
    up = u.post(f"/api/ap/batches/{b['id']}/sap/upload", json={"include_yellow": False})
    assert up.status_code == 200, up.text
    ws = load_workbook(io.BytesIO(up.content)).active
    assert {r[2] for r in ws.iter_rows(min_row=2, values_only=True)} == {"19102026"} and ws.cell(1, 3).value == "postdate"
    # what went out keeps its date
    monkeypatch.setattr(P, "sap_today", lambda: mon + timedelta(days=1))
    assert {x["postdate"] for x in _cap(u, docs[f"pd1_{run}.pdf"]["id"])["sap_rows"]} == {"19102026"}


def test_a_posting_date_a_person_set_stays_when_the_download_is_later(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    monkeypatch.setattr(P, "sap_today", lambda: date(2026, 10, 16))
    vat = _vat()
    _code("Kept Day Supplies", vat)
    u = _user()
    b = _run(monkeypatch, u, {f"kd_{run}.pdf": _answer(f"KD-{run}", "Kept Day Supplies", vat)}, f"KD {run}")
    assert u.post(f"/api/ap/batches/{b['id']}/sap/posting-date", json={"date": "2026-10-15"}).status_code == 200
    monkeypatch.setattr(P, "sap_today", lambda: date(2026, 10, 19))
    ws = load_workbook(io.BytesIO(u.post(f"/api/ap/batches/{b['id']}/sap/upload", json={"include_yellow": True}).content)).active
    assert {r[2] for r in ws.iter_rows(min_row=2, values_only=True)} == {"15102026"}
    assert u.post(f"/api/ap/batches/{b['id']}/sap/posting-date", json={"date": "15/10/2026"}).status_code == 400  # not an ISO date


def test_utility_bills_take_the_posting_date_too():
    f = {"account": "5012345678", "billing_date": "2026-09-15", "due_date": "2026-10-05", "invoice_no": "INV-77", "statement_month": "SEPTEMBER 2026",
         "read_start": "2026-08-01", "read_end": "2026-08-31", "reading_type": "Actual", "kwh": "100", "services": [{"service": "electricity", "ex": "1000.00", "vat": "150.00"}]}
    acct = {"vendor": "700999", "cc": "01306", "name1": "Badplaas ETE", "muni": "X"}
    rows, st, _ = P.utility_doc(f, None, acct, _tables(), file="u.pdf")
    assert {r["postdate"] for r in rows} == {"15092026"} and st == "GRN"  # no posting date given: the Harvester's (the bill's date)
    rows, st, _ = P.utility_doc(f, None, acct, _tables(), file="u.pdf", postdate="16102026")
    assert {r["postdate"] for r in rows} == {"16102026"} and {r["docdate"] for r in rows} == {"15092026"}
    assert rows[0]["docheadertext"] == "Badplaas ETE SEPTEMBER 2026"  # utility bills keep the Harvester's header
    known = {"vendors": {"700999"}, "gls": {"SA110", "JA600"}, "companies": {"T000"}, "doc_types": {"KR", "KG"}, "tax_codes": {"I9", "I0"}, "today": date(2026, 10, 16)}
    assert P.test_doc(rows, known) == ("green", [])
    far = [{**r, "postdate": "16122026"} for r in rows]
    assert P.test_doc(far, known) == ("yellow", [{"level": "yellow", "text": "posting date 16.12.2026 is far from today"}])


# ---- 2. the cost centre per site -------------------------------------------------------------------------------------------
def test_a_site_gives_the_cost_centre_and_a_site_with_no_row_waits_in_sites_to_map(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    site = f"ZS{run[:4]}"
    vat = _vat()
    _code("Site Test Services (Pty) Ltd", vat, cost_centre="AC0328")
    u, adm = _user(), _admin()
    b = _run(monkeypatch, u, {f"st_{run}.pdf": _answer(f"ST-{run}", "Site Test Services (Pty) Ltd", vat, bill_to=f"{FULL} (Pty) Ltd - {site}")}, f"Site {run}")
    d = _sheet(u, b["id"])[f"st_{run}.pdf"]
    c = _cap(u, d["id"])
    assert c["bill_to_site"] == site and _ccs(c["sap_rows"]) == {"AC0328"}  # no row yet: the vendor coding's
    assert c["sap_rows"][0]["docheadertext"] == P.header_text("Site Test Services (Pty) Ltd", INV_DAY.isoformat(), site) == f"{site} Site Test {P.mmm_yy(INV_DAY.isoformat())}"
    waiting = next(g for g in u.get("/api/ap/needs").json()["sites"] if g["site"] == site)
    assert waiting["n"] == 1 and waiting["companies"] == ["X143"] and d["status"] == "green"  # it does not hold the invoice up
    # only admins map a site; the cost centre must exist in the company code in SAP
    assert u.post("/api/ap/sites/map", json={"site_code": site, "company_code": "X143", "cost_centre": "AC0326"}).status_code == 403
    r = adm.post("/api/ap/sites/map", json={"site_code": site, "company_code": "X143", "cost_centre": "2822"})
    assert r.status_code == 400 and r.json()["detail"] == "Cost centre 2822 doesn't exist in company code X143 in SAP."
    assert adm.post("/api/ap/sites/map", json={"site_code": site, "cost_centre": "SAP-DUMMY"}).status_code == 400
    # the site's row for any company: Pip builds the invoice again with it
    r = adm.post("/api/ap/sites/map", json={"site_code": site.lower(), "company_code": "", "cost_centre": "AC0326", "note": "made up"})
    assert r.status_code == 200 and r.json()["rebuilt"] == 1 and r.json()["site"]["site_code"] == site
    c = _cap(u, d["id"])
    assert _ccs(c["sap_rows"]) == {"AC0326"} and c["sap_status"] == "green"
    assert any(f"cost centre AC0326 from the site {site} (any company)" in h for h in _hist(c["item_id"]))
    assert all(g["site"] != site for g in u.get("/api/ap/needs").json()["sites"])
    # the site's row for this company wins over its row for any company (Settings → Sites)
    r = adm.post("/api/ap/masters/sites", json={"site_code": site, "company_code": "X143", "cost_centre": "AC0319"})
    assert r.status_code == 200
    assert _ccs(_cap(u, d["id"])["sap_rows"]) == {"AC0319"}
    # a person's cost centre for this invoice wins over the site
    r = u.post(f"/api/ap/captures/{d['id']}/decide", json={"action": "save", "data": {"fields": {"cost_centre": "AC0328"}}})
    assert _ccs(r.json()["sap_rows"]) == {"AC0328"}
    # Settings → Sites lists them; viewers read, only admins change
    rows = u.get(f"/api/ap/masters/sites?q={site}").json()["rows"]
    assert sorted((x["company_code"], x["cost_centre"]) for x in rows) == [("", "AC0326"), ("X143", "AC0319")]
    assert u.post("/api/ap/masters/sites", json={"site_code": site, "cost_centre": "AC0326"}).status_code == 403
    with SessionLocal() as s:
        assert s.query(AuditLog).filter(AuditLog.action == "ap_site", AuditLog.target == f"ap_site:{site}").count() == 2


def test_the_sites_list_uploads_with_a_preview_and_pip_s_order():
    adm = _admin()
    wb = Workbook()
    ws = wb.active
    for r in (["Site code", "Company code", "Cost centre", "Note"], ["RIO2X21", "X143", "AC0326", "mine client"], ["dell", "", "AC0319", ""], ["", "X143", "AC0328", "no site"]):
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    r = adm.post("/api/ap/masters/sites/import", files={"file": ("sites.xlsx", buf.getvalue())}, data={"confirm": "false"}).json()
    assert (r["added"], r["skipped"], r["saved"]) == (2, 1, False) and r["mapping"] == {"site_code": "Site code", "company_code": "Company code", "cost_centre": "Cost centre", "note": "Note"}
    r = adm.post("/api/ap/masters/sites/import", files={"file": ("sites.xlsx", buf.getvalue())}, data={"confirm": "true"}).json()
    assert r["saved"] and r["rebuilt"] >= 0
    with SessionLocal() as s:
        assert M.site_cost_centre(s, "RIO2X21", "X143") == ("AC0326", "company")
        assert M.site_cost_centre(s, "Dell", "X120") == ("AC0319", "any")
        assert M.site_cost_centre(s, "RIO2X21", "X120") == ("", "")
    out = load_workbook(io.BytesIO(adm.get("/api/ap/masters/sites/file").content)).active
    assert [c.value for c in out[1]] == ["Site code", "Company code", "Cost centre", "Note"] and out.max_row == 3
    # Pip: the site's cost centre, else the vendor coding's, else none (yellow, as before)
    f = {"supplier": "Acme", "invoice_no": "S-1", "invoice_date": "2026-10-01", "due_date": "2026-10-31", "excl": "100.00", "vat": "15.00", "incl": "115.00"}
    rows, st, issues = P.manual_doc(f, {"vendor_no": "300100", "gl": "SC600", "cost_centre": ""}, "X143", _tables(), file="s.pdf")
    assert st == "YLW" and "no cost centre for this vendor" in issues


# ---- 3. G/L per line, 5. tax codes per line ----------------------------------------------------------------------------------
RULES = [{"words": "biokinetic", "gl": "SC109", "cost_centre": "AC0326"}, {"words": "equipment, tools", "gl": "SB006"}]
LINES = [{"description": "Biokinetic assessment", "excl": "1000.00", "vat_rate": 15, "vat": "150.00"},
         {"description": "Equipment service", "excl": "500.00", "vat_rate": 15, "vat": "75.00"},
         {"description": "Management fee", "excl": "2000.00", "vat_rate": 15, "vat": "300.00"},
         {"description": "BIOKINETIC follow-up", "excl": "300.00", "vat_rate": 15, "vat": "45.00"},
         {"description": "Medical aid levy (exempt)", "excl": "200.00", "vat_rate": 0, "vat": "0.00"}]
F = {"supplier": "Peak Wellness (Pty) Ltd", "invoice_no": "PW-9", "invoice_date": "2026-08-14", "due_date": "2026-09-13", "currency": "ZAR", "lines": LINES,
     "excl": "4000.00", "vat": "570.00", "incl": "4570.00"}


def _lines(rows):
    return [(r["pkey"], r["account"], r["amount"], r["tc"], r["costcentre"]) for r in rows]


def test_line_rules_give_each_line_its_gl_and_lines_group_by_gl_tax_code_and_cost_centre():
    coding = {"vendor_no": "300100", "name": "Peak Wellness (Pty) Ltd", "gl": "SC600", "cost_centre": "AC0328", "tax_code": "I9", "line_rules": RULES}
    rows, st, issues = P.manual_doc(F, coding, "X143", _tables(), file="pw.pdf")
    assert st == "GRN" and not issues, issues
    assert _lines(rows) == [("31", "300100", "4570.00", "", ""), ("40", "SC109", "1300.00", "I9", "AC0326"), ("40", "SB006", "500.00", "I9", "AC0328"),
                            ("40", "SC600", "2000.00", "I9", "AC0328"), ("40", "SC600", "200.00", "I0", "AC0328"), ("40", "JA600", "570.00", "I9", "")]
    assert P.balance_fails([{**r, "_file": "x"} for r in rows]) == (1, set())
    # a person's G/L and cost centre for this invoice: the rules don't change them
    rows, _, _ = P.manual_doc(F, {**coding, "gl": "SC109", "cost_centre": "AC0319", "gl_fixed": True, "cc_fixed": True}, "X143", _tables(), file="pw.pdf")
    assert {(a, cc) for _, a, _, _, cc in _lines(rows)[1:-1]} == {("SC109", "AC0319")}
    assert P.line_rule(RULES, "Hand tools") == RULES[1] and P.line_rule(RULES, "Cleaning") is None


def test_a_mixed_invoice_gets_the_zero_rated_code_on_its_zero_lines_and_a_rule_may_set_an_exempt_code():
    coding = {"vendor_no": "300100", "name": "Peak", "gl": "SC600", "cost_centre": "AC0328", "tax_code": "I9"}
    rows, st, _ = P.manual_doc(F, coding, "X143", _tables(), file="pw.pdf")
    assert {(r["amount"], r["tc"]) for r in rows if r["pkey"] == "40" and r["account"] == "SC600"} == {("3800.00", "I9"), ("200.00", "I0")} and st == "GRN"
    exempt = {**coding, "line_rules": [{"words": "levy", "tax_code": "EXEMPT"}]}
    rows, _, _ = P.manual_doc(F, exempt, "X143", _tables(), file="pw.pdf")  # TaxCodeExempt empty: nothing changes
    assert ("200.00", "I0") in {(r["amount"], r["tc"]) for r in rows}
    rows, _, _ = P.manual_doc(F, exempt, "X143", _tables(TaxCodeExempt="IE"), file="pw.pdf")
    assert ("200.00", "IE") in {(r["amount"], r["tc"]) for r in rows}
    rows, _, _ = P.manual_doc(F, {**coding, "line_rules": [{"words": "levy", "tax_code": "I3"}]}, "X143", _tables(), file="pw.pdf")
    assert ("200.00", "I3") in {(r["amount"], r["tc"]) for r in rows}
    with SessionLocal() as s:  # the Fixed value is there, empty, for an admin to set
        assert s.query(ApSapValue).filter(ApSapValue.kind == "fixed", ApSapValue.key == "TaxCodeExempt").one().value == ""


def test_line_rules_are_kept_on_the_vendor_coding_and_their_gls_checked_against_the_chart(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    vat = _vat()
    adm, u = _admin(), _user()
    body = {"vendor_no": "88" + str(uuid.uuid4().int)[:5], "name": f"Rule Wellness {run}", "vat_no": vat, "gl": "SC600", "cost_centre": "AC0328", "tax_code": "I9",
            "payment_days": "30", "line_rules": [{"words": "biokinetic", "gl": "sc109", "cost_centre": "AC0326"}, {"words": "gym", "gl": "ZZ999"}, {"words": "", "gl": "SB006"}]}
    r = adm.post("/api/ap/masters/vendors", json=body)
    assert r.status_code == 200, r.text
    v = r.json()
    assert v["payment_days"] == 30 and v["line_rules"] == [{"words": "biokinetic", "gl": "SC109", "cost_centre": "AC0326", "tax_code": ""},
                                                           {"words": "gym", "gl": "ZZ999", "cost_centre": "", "tax_code": ""}]
    assert v["warning"] == "G/L ZZ999 is not in the chart of accounts CBX has." and v["missing"] == ["ZZ999"]
    lines = [{"description": "Biokinetic session", "amount_excl": "600.00", "vat_rate": 15, "vat": "90.00"}, {"description": "Gym access", "amount_excl": "400.00", "vat_rate": 15, "vat": "60.00"}]
    b = _run(monkeypatch, u, {f"rw_{run}.pdf": _answer(f"RW-{run}", f"Rule Wellness {run}", vat, lines=lines)}, f"Rules {run}")
    d = _sheet(u, b["id"])[f"rw_{run}.pdf"]
    c = _cap(u, d["id"])
    assert {(r["account"], r["costcentre"]) for r in c["sap_rows"] if r["pkey"] == "40" and r["account"] != "JA600"} == {("SC109", "AC0326"), ("ZZ999", "AC0328")}
    assert d["status"] == "yellow" and {"level": "yellow", "text": "G/L ZZ999 is not in the chart of accounts — check with finance"} in d["issues"]
    assert "ZZ999" in adm.get("/api/ap/masters").json()["gl_missing"]
    wb = load_workbook(io.BytesIO(adm.get("/api/ap/masters/vendors/file").content)).active
    heads = [x.value for x in wb[1]]
    row = next(r for r in wb.iter_rows(min_row=2, values_only=True) if r[0] == body["vendor_no"])
    assert heads[-2:] == ["Payment days", "Line rules"] and row[-2:] == (30, "biokinetic => SC109 / AC0326; gym => ZZ999")
    assert M.clean_rules(row[-1]) == v["line_rules"]  # the upload reads the cell back
    with SessionLocal() as s:
        s.query(ApVendorCode).filter(ApVendorCode.vendor_no == body["vendor_no"]).delete()
        s.commit()


# ---- 4. the due date from payment terms ----------------------------------------------------------------------------------
def test_no_due_date_is_the_invoice_date_plus_the_vendor_s_payment_days(monkeypatch):
    f = {"supplier": "Acme", "invoice_no": "D-1", "invoice_date": "2026-08-14", "due_date": "", "excl": "100.00", "vat": "15.00", "incl": "115.00"}
    coding = {"vendor_no": "300100", "gl": "SC600", "cost_centre": "AC0328"}
    rows, st, issues = P.manual_doc(f, {**coding, "payment_days": 30}, "X143", _tables(), file="d.pdf")
    assert rows[0]["DueDate"] == "13092026" and st == "GRN" and not issues
    rows, st, issues = P.manual_doc(f, coding, "X143", _tables(), file="d.pdf")
    assert rows[0]["DueDate"] == "" and st == "YLW" and "no due date on the invoice" in issues
    rows, _, _ = P.manual_doc({**f, "due_date": "2026-08-20"}, {**coding, "payment_days": 30}, "X143", _tables(), file="d.pdf")
    assert rows[0]["DueDate"] == "20082026"  # a printed due date wins
    # through the line: Code this vendor with its payment days → the history says how the due date came about
    run = uuid.uuid4().hex[:6].upper()
    vat = _vat()
    u = _user()
    b = _run(monkeypatch, u, {f"dd_{run}.pdf": _answer(f"DD-{run}", f"Due Days {run} (Pty) Ltd", vat, due=False)}, f"Due {run}")
    g = next(x for x in u.get("/api/ap/needs").json()["vendors"] if x["vat"] == vat)
    r = u.post("/api/ap/vendors/code", json={"ids": g["ids"], "vendor_no": "99" + str(uuid.uuid4().int)[:5], "gl": "SC600", "cost_centre": "AC0328", "tax_code": "I9",
                                            "name": g["supplier"], "vat_no": vat, "payment_days": 45})
    assert r.status_code == 200 and r.json()["vendor"]["payment_days"] == 45
    d = _sheet(u, b["id"])[f"dd_{run}.pdf"]
    c = _cap(u, d["id"])
    due = INV_DAY + timedelta(days=45)
    assert c["sap_rows"][0]["DueDate"] == f"{due:%d%m%Y}" and "no due date on the invoice" not in {x["text"] for x in d["issues"]}
    assert f"No due date on the invoice: due {due.isoformat()} = the invoice date + 45 days (the vendor's payment terms)." in _hist(c["item_id"])


# ---- 6. SAP's vendor general data (LFA1) ------------------------------------------------------------------------------------
def _lfa1_csv(rows: list[list]) -> bytes:
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\n").writerows(rows)
    return buf.getvalue().encode("utf-8")


def _lfa1_xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_the_lfa1_export_loads_in_either_column_style_with_vat_numbers_normalised():
    adm, u = _admin(), _user()
    tech = [["LIFNR", "NAME1", "NAME2", "STCEG", "STCD1", "ORT01", "LAND1", "SPERR", "LOEVM", "ZTERM"],
            ["0000120083", "Twin Supplies", "(Pty) Ltd", "ZA 4400 000 014", "", "Midrand", "ZA", "", "", "N30"],
            ["120000", "Twin Supplies", "", "", "4400000014", "Midrand", "ZA", "", "", "Z045"],
            ["130001", "Old Vendor CC", "", "4500000013", "", "Durban", "ZA", "X", "", "0001"],
            ["130002", "Gone Vendor", "", "4500000021", "", "Durban", "ZA", "", "X", ""],
            ["", "No number", "", "4500000039", "", "", "", "", "", ""]]
    r = adm.post("/api/ap/masters/sap_vendors/import", files={"file": ("LFA1.csv", _lfa1_csv(tech))}, data={"confirm": "false"}).json()
    assert (r["rows"], r["added"], r["skipped"], r["with_vat"], r["blocked"], r["saved"]) == (4, 4, 0, 4, 2, False)
    assert r["mapping"] == {"vendor": "LIFNR", "name": "NAME1", "name2": "NAME2", "vat": "STCEG", "tax1": "STCD1", "city": "ORT01", "country": "LAND1",
                            "blocked": "SPERR", "deleted": "LOEVM", "payment_terms": "ZTERM"}
    assert adm.post("/api/ap/masters/sap_vendors/import", files={"file": ("LFA1.csv", _lfa1_csv(tech))}, data={"confirm": "true"}).json()["saved"]
    with SessionLocal() as s:
        K.forget()
        assert K.sap_vendor(s, "120083")["vat"] == "4400000014" and K.sap_vendor(s, "120083")["name"] == "Twin Supplies (Pty) Ltd"
        assert K.sap_vendor(s, "0000120000")["vat"] == "4400000014"  # the VAT number in STCD1, leading zeros as SAP may give them
        assert K.sap_vendor(s, "120083")["payment_days"] == 30 and K.sap_vendor(s, "120000")["payment_days"] == 45 and K.sap_vendor(s, "130001")["payment_days"] is None
        assert K.vendor_blocked(s, "130001") is True and K.vendor_blocked(s, "130002") is True and K.vendor_blocked(s, "120083") is False and K.vendor_blocked(s, "9") is None
        # a VAT number fits two vendor numbers: the one opened in the posting company (SAP's vendor / company code list) first
        assert sorted(x["vendor"] for x in K.sap_vendors_by_vat(s, "4400000014")) == ["0000120083", "120000"]
        assert C.sap_vendor_for(s, "Twin", "4400000014", "X120")["vendor"] == "0000120083"
        # 1.0.103: X143 is not in that export and neither vendor has posting history there: no pick — a person picks one
        twin = C.sap_vendor_for(s, "Twin", "4400000014", "X143")
        assert twin["vendor"] == "" and [x["vendor"] for x in twin["several"]] == ["0000120083", "120000"]
        assert C.sap_vendor_for(s, "Twin Supplies (Pty) Ltd", "", "X143")["how"] == "name"  # a name only when the invoice has no VAT number
        assert C.sap_vendor_for(s, "Twin Supplies (Pty) Ltd", "4999999990", "X143") is None
    assert K.sap_vat("ZA4123456789") == "4123456789" and K.terms_days("N30") == 30 and K.terms_days("Z007") == 7 and K.terms_days("EOM") is None
    # the English labels of a report: the same table
    labels = [["Vendor", "Name 1", "VAT Registration No.", "Tax Number 1", "City", "Country", "Central posting block", "Terms of Payment"],
              ["120083", "Twin Supplies (Pty) Ltd", "4400000014", "", "Midrand", "ZA", "", "N30"], ["140001", "Label Vendor", "", "4600000012", "Pretoria", "ZA", "X", "N14"]]
    r = adm.post("/api/ap/masters/sap_vendors/import", files={"file": ("vendors.xlsx", _lfa1_xlsx(labels))}, data={"confirm": "false"}).json()
    assert r["mapping"]["vat"] == "VAT Registration No." and r["mapping"]["payment_terms"] == "Terms of Payment" and r["mapping"]["blocked"] == "Central posting block"
    assert (r["added"], r["removed"]) == (2, 4)  # "120083" is not "0000120083": SAP's own export decides the spelling
    # only admins load it; everyone with the tool reads it
    assert u.post("/api/ap/masters/sap_vendors/import", files={"file": ("LFA1.csv", _lfa1_csv(tech))}, data={"confirm": "true"}).status_code == 403
    assert u.get("/api/ap/masters/sap_vendors?q=4400000014").json()["total"] == 4
    lf = load_workbook(io.BytesIO(u.get("/api/ap/masters/sap_vendors/file").content)).active
    assert [c.value for c in lf[1]] == ["LIFNR", "NAME1", "STCEG", "STCD1", "ORT01", "LAND1", "SPERR", "ZTERM", "KTOKK"]  # 1.0.103: + the account group


def test_mia_finds_an_uncoded_supplier_in_sap_s_list_and_a_blocked_vendor_is_red(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    adm, u = _admin(), _user()
    new_vat, coded_vat = _vat(), _vat()
    blocked_no = _code(f"Blocked Coded {run}", coded_vat)
    lfa1 = [["LIFNR", "NAME1", "STCEG", "SPERR", "ZTERM"], ["77" + str(uuid.uuid4().int)[:5], f"Fresh Vendor {run} (Pty) Ltd", new_vat, "", "N30"],
            [blocked_no, f"Blocked Coded {run}", coded_vat, "X", "N30"], ["7799991", f"Named Only {run}", "", "", "N14"]]
    assert adm.post("/api/ap/masters/sap_vendors/import", files={"file": ("LFA1.csv", _lfa1_csv(lfa1))}, data={"confirm": "true"}).json()["saved"]
    K.forget()
    b = _run(monkeypatch, u, {f"fresh_{run}.pdf": _answer(f"FV-{run}", f"Fresh Vendor {run}", new_vat, due=False),
                              f"blk_{run}.pdf": _answer(f"BK-{run}", f"Blocked Coded {run}", coded_vat),
                              f"named_{run}.pdf": _answer(f"NO-{run}", f"Named Only {run}", "", due=False)}, f"LFA1 {run}")
    # 1.0.105: no longer Needs you ("Code this vendor" prefilled) — CBX takes SAP's vendor itself (cbx/ap/autofill.py): by the
    # VAT number (sure), by the name (a guess: yellow), its payment terms (ZTERM) give the due date (a guess: yellow)
    docs = _sheet(u, b["id"])
    fresh_doc, named_doc = docs[f"fresh_{run}.pdf"], docs[f"named_{run}.pdf"]
    fresh = _cap(u, fresh_doc["id"])["match"]["sap_vendor"]
    assert (fresh["vendor"], fresh["name"], fresh["payment_days"], fresh["how"]) == (lfa1[1][0], f"Fresh Vendor {run} (Pty) Ltd", 30, "VAT number")
    named = _cap(u, named_doc["id"])["match"]["sap_vendor"]
    assert (named["vendor"], named["how"], named["payment_days"]) == ("7799991", "name", 14)
    assert fresh_doc["vendor"] == lfa1[1][0] and named_doc["vendor"] == "7799991"
    assert any(f"vendor {fresh['vendor']} Fresh Vendor {run} (Pty) Ltd (SAP's vendor list, by its VAT number)" in h for h in _hist(fresh_doc["item_id"]))
    assert _cap(u, fresh_doc["id"])["sap_rows"][0]["DueDate"] == f"{INV_DAY + timedelta(days=30):%d%m%Y}"  # N30
    # the vendor SAP has blocked: red, never in the upload
    blk = docs[f"blk_{run}.pdf"]
    assert blk["status"] == "red" and f"vendor {blocked_no} is blocked in SAP" in {x["text"] for x in blk["issues"]}
    # found by name: yellow (a guess); SAP's payment terms (N14) give the due date
    texts = {x["text"] for x in named_doc["issues"]}
    assert named_doc["status"] == "yellow" and "vendor 7799991 guessed: SAP's vendor list, by its name - check it" in texts
    assert _cap(u, named_doc["id"])["sap_rows"][0]["DueDate"] == f"{INV_DAY + timedelta(days=14):%d%m%Y}"
    with SessionLocal() as s:
        s.query(ApVendorCode).filter(ApVendorCode.vendor_no.in_([blocked_no, "7799991"])).delete()
        s.query(ApCapture).filter(ApCapture.batch_id == b["id"]).update({"uploaded_at": None})
        s.commit()


# ---- 7. the header text ---------------------------------------------------------------------------------------------------------
def test_the_header_text_puts_the_useful_part_first_within_25_characters():
    assert P.header_text("Excellerate Services (Pty) Ltd", "2026-08-15") == "Excellerate Serv. AUG 26"
    assert P.header_text("SHINE SERVICES (PTY) LTD", "2026-09-30", "DELL") == "DELL SHINE SERV. SEP 26"
    assert P.header_text("Acme Cleaning", "2026-10-01") == "Acme Cleaning OCT 26"
    assert P.header_text("Brightline Studio (Pty) Ltd", "2026-10-01", "RIO2X21") == "RIO2X21 Brightline OCT 26"
    for name in ("Supercalifragilisticexpialidocious Holdings", "A B C D E F G H I J K L M N O P", "The Very Long Name Of A Facilities Company Limited"):
        for site in ("", "RIO2X21", "A-VERY-LONG-SITE-CODE-X"):
            h = P.header_text(name, "2026-10-01", site)
            assert len(h) <= 25 and h.endswith("OCT 26"), h
    assert P.short_name("Twin Supplies (Pty) Ltd") == "Twin Supplies" and P.short_name("Old Vendor CC") == "Old Vendor"
    f = {"supplier": "Excellerate Services (Pty) Ltd", "invoice_no": "H-1", "invoice_date": "2026-08-15", "due_date": "2026-09-14", "excl": "100.00", "vat": "15.00", "incl": "115.00"}
    rows, _, _ = P.manual_doc(f, {"vendor_no": "1", "name": "", "gl": "SC600", "cost_centre": "AC0328", "item_text": "{site} {description}"}, "X143", _tables(), file="h.pdf", site="RIO2X21")
    assert {r["docheadertext"] for r in rows} == {"Excellerate Serv. AUG 26"}  # a line-text rule: the site goes in the lines, not the header


def test_pull_from_sap_reads_a_supplier_list_preview_then_save(monkeypatch):
    """Pull from SAP (OData A_Supplier through Admin → Connections): preview, then exactly those rows are saved."""
    import json
    import threading
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

    from cbx.config import settings
    from cbx.models.connections import ApiConnection

    class Sap(BaseHTTPRequestHandler):
        def log_message(self, *a):
            pass

        def do_GET(self):
            rows = [{"Supplier": "0000150001", "SupplierName": "Pulled Supplier", "VATRegistration": "ZA4800000010", "PostingIsBlocked": False, "DeletionIndicator": False},
                    {"Supplier": "0000150002", "SupplierName": "Pulled Blocked", "VATRegistration": "", "PostingIsBlocked": True, "DeletionIndicator": False}]
            b = json.dumps({"d": {"results": rows}}).encode()
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(b)))
            self.end_headers()
            self.wfile.write(b)

    monkeypatch.setattr(settings, "api_allow_private", True)
    srv = ThreadingHTTPServer(("127.0.0.1", 0), Sap)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    adm = _admin()
    try:
        with SessionLocal() as s:
            c = ApiConnection(slug="lfa1-sap-" + uuid.uuid4().hex[:6], name="SAP suppliers", base_url=f"http://127.0.0.1:{srv.server_address[1]}", auth_kind="none", tools=["ap"], is_active=True)
            s.add(c)
            s.commit()
            cid = c.id
        cols = {"vendor": "Supplier", "name": "SupplierName", "vat": "VATRegistration", "blocked": "PostingIsBlocked", "deleted": "DeletionIndicator"}
        pre = adm.post("/api/ap/masters/sap_vendors/pull", json={"connection_id": cid, "path": "/A_Supplier", "columns": cols}).json()
        assert (pre["added"], pre["blocked"], pre["with_vat"], pre["saved"]) == (2, 1, 1, False) and pre["preview_id"]
        assert adm.post("/api/ap/masters/sap_vendors/pull", json={"confirm": True, "preview_id": pre["preview_id"]}).json()["saved"]
        with SessionLocal() as s:
            K.forget()
            assert K.sap_vendors_by_vat(s, "4800000010")[0]["vendor"] == "0000150001" and K.vendor_blocked(s, "150002") is True
            assert s.query(ApSapVendor).filter(ApSapVendor.vendor == "0000150001").one().source == "sap"
    finally:
        srv.shutdown()
