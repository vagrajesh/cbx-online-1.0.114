"""1.0.98: from CBX-Connect 2.10.3 tests/test_ap_scans_2103.py, on the old conftest.

2.10.3: scanned supplier invoices (image only, no text layer) through AP capture, with SAP's knowledge.

Three look-alikes of the layouts Edwin sent (a studio's one-line project invoice billed to "… (Pty) Ltd - RIO2X21", a
Sage "Copy Tax Invoice" with an account/reference strip, and a services company's long item list with lines that have
no amount and a "Rounding -1" line, billed to "… - DELL"). Suppliers, VAT numbers, bank details and amounts are made
up; each page is drawn with reportlab, turned into a picture with PyMuPDF and put in a new PDF as that picture only.
The AI is off in tests: the capture call returns what a vision model would read off each page."""
from __future__ import annotations

import io
import uuid
from datetime import date, timedelta

import pymupdf
import pytest
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from cbx import knowledge as K
from cbx.ap import capture as C
from cbx.ap import sap as P
from cbx.db import SessionLocal
from cbx.models.ap import ApCapture, ApEntity, ApVendorCode

from .test_ap_capture_1098 import _admin, _app, _drop, _sheet, _user, _vision  # noqa: F401

FULL = "CBRE Excellerate Facilities Management"
BILL_VAT = "4999000011"  # made up
TODAY = date.today()
INV_DAY, DUE_DAY = TODAY - timedelta(days=12), TODAY + timedelta(days=18)
VENDORS = {  # made-up suppliers, coded to real accounts of the CBX chart
    "brightline": {"vendor_no": "9910001", "name": "Brightline Studio (Pty) Ltd", "vat_no": "4100000017", "gl": "SC109", "cost_centre": "AC0326", "tax_code": "I9",
                   "item_text": "{site} {description}"},
    "peak": {"vendor_no": "9910002", "name": "Peak Wellness (Pty) Ltd", "vat_no": "4200000016", "gl": "SC600", "cost_centre": "AC0328", "tax_code": "I9", "item_text": ""},
    "shine": {"vendor_no": "9910003", "name": "Shine Services (Pty) Ltd", "vat_no": "4300000015", "gl": "SB006", "cost_centre": "AC0319", "tax_code": "I9", "item_text": ""},
}


@pytest.fixture(autouse=True)
def _setup():
    """X143's full name taught (a sure bill-to match) and the three suppliers coded, as an admin keeps them."""
    with SessionLocal() as s:
        e = s.query(ApEntity).filter(ApEntity.code == "X143").one()
        if FULL not in (e.aliases or ""):
            e.aliases = (e.aliases + "\n" if e.aliases else "") + FULL
        for v in VENDORS.values():
            s.query(ApVendorCode).filter(ApVendorCode.vat_no == v["vat_no"]).delete()
            s.add(ApVendorCode(**v, name_key=C.M.name_key(v["name"]), source="excel"))
        s.commit()
    K.forget()


# ---- the look-alikes -----------------------------------------------------------------------------------------------
def _scan(draw) -> bytes:
    """reportlab page → a picture (PyMuPDF, 150 dpi) → a new PDF holding only that picture: no text layer, like a scan."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    draw(c)
    c.showPage()
    c.save()
    with pymupdf.open(stream=buf.getvalue(), filetype="pdf") as src:
        png = src[0].get_pixmap(dpi=150).tobytes("png")
        rect = src[0].rect
    with pymupdf.open() as out:
        page = out.new_page(width=rect.width, height=rect.height)
        page.insert_image(page.rect, stream=png)
        return out.tobytes()


def _d(day: date) -> str:
    return f"{day.day} {day:%b %Y}"


def brightline_pdf(no: str) -> bytes:
    def draw(c):
        c.setFont("Helvetica-Bold", 22); c.drawString(380, 740, "brightline")
        c.setFont("Helvetica", 24); c.drawString(50, 660, "TAX INVOICE")
        c.setFont("Helvetica", 9)
        for i, t in enumerate([f"{FULL} (Pty) Ltd - RIO2X21", "1 Example Road", "Dunkeld West", "2196", f"VAT Number: {BILL_VAT}"]):
            c.drawString(80, 630 - 13 * i, t)
        for i, (k, v) in enumerate([("Invoice Date", _d(INV_DAY)), ("Account Number", "RIO2X21"), ("Invoice Number", no), ("Reference", "PO7700112026"), ("VAT Number", "4100000017")]):
            c.setFont("Helvetica-Bold", 9); c.drawString(360, 680 - 28 * i, k)
            c.setFont("Helvetica", 9); c.drawString(360, 668 - 28 * i, v)
        c.drawString(470, 680, "Brightline Studio (Pty) Ltd"); c.drawString(470, 668, "PO Box 100, Cape Town")
        c.setFont("Helvetica-Bold", 9)
        for x, h in ((55, "Description"), (330, "Quantity"), (400, "Unit Price"), (470, "VAT"), (520, "Amount ZAR")):
            c.drawString(x, 520, h)
        c.setFont("Helvetica", 9)
        c.drawString(55, 500, "260922 - RIO2X21 - MINING CLIENT - EXTRAS - 100% Invoice")
        c.drawString(55, 480, "Project Value"); c.drawString(340, 480, "1.00"); c.drawString(400, 480, "21,400.00"); c.drawString(470, 480, "15%"); c.drawString(520, 480, "21,400.00")
        c.drawString(440, 455, "Subtotal"); c.drawString(520, 455, "21,400.00")
        c.drawString(440, 435, "TOTAL VAT"); c.drawString(520, 435, "3,210.00")
        c.setFont("Helvetica-Bold", 9); c.drawString(440, 412, "TOTAL ZAR"); c.drawString(520, 412, "24,610.00")
        c.setFont("Helvetica", 11); c.drawString(55, 360, f"Due Date: {_d(DUE_DAY)}")
        c.setFont("Helvetica", 9); c.drawString(55, 345, "Banking details: Example Bank 000-111  Acc No: 000 0000 001")
    return _scan(draw)


def peak_pdf(no: str) -> bytes:
    def draw(c):
        c.setFillGray(0.45); c.rect(40, 730, 532, 30, fill=1, stroke=0); c.setFillGray(1)
        c.setFont("Helvetica-Bold", 16); c.drawCentredString(306, 740, "Copy Tax Invoice"); c.setFillGray(0)
        c.setFont("Helvetica", 9)
        for i, (k, v) in enumerate([("Date", INV_DAY.strftime("%d/%m/%y")), ("Due By", DUE_DAY.strftime("%d/%m/%y")), ("Document No", no)]):
            c.drawString(330, 705 - 22 * i, k); c.drawRightString(560, 705 - 22 * i, v)
        c.setFont("Helvetica-Bold", 10); c.drawString(50, 640, "Peak Wellness (Pty) Ltd")
        c.setFont("Helvetica", 8); c.drawString(50, 626, "9 Example Drive, Fourways"); c.drawString(220, 590, "VAT No: 4200000016")
        c.setFont("Helvetica-Bold", 8); c.drawString(330, 640, FULL)
        c.setFont("Helvetica", 8)
        for i, t in enumerate(["Private Bag X1", "Benmore", "2010", "Deliver to", "Excellerate on Summit", f"VAT No: {BILL_VAT}"]):
            c.drawString(330, 628 - 11 * i, t)
        for x, h, v in ((50, "Account", "CBR001"), (130, "Your Reference", "PO4040012026"), (260, "Tax Exempt", "N"), (330, "Tax Reference", BILL_VAT), (520, "Incl/Excl", "Exclusive")):
            c.drawString(x, 540, h); c.drawString(x, 524, v)
        for x, h in ((50, "Code"), (130, "Description"), (330, "Quantity"), (400, "Unit Price"), (470, "Tax"), (520, "Nett Price")):
            c.drawString(x, 495, h)
        c.drawString(130, 480, "Facility management")
        for i, (d, a) in enumerate([("Management fee", "18,250.00"), ("Equipment service", "1,240.50"), ("Assessments", "3,100.00")]):
            y = 468 - 12 * i
            c.drawString(50, y, "001"); c.drawString(130, y, d); c.drawString(340, y, "1.00"); c.drawRightString(450, y, a); c.drawString(470, y, "15.00%"); c.drawRightString(560, y, a)
        c.drawString(50, 160, "Banking Details: Example Bank  Account No: 0000000002  Branch Code: 000000")
        for i, (k, v) in enumerate([("Sub Total", "22,590.50"), ("Discount @ 0.00%", "0.00"), ("Amount Excl Tax", "22,590.50"), ("Tax", "3,388.58"), ("Total", "25,979.08")]):
            c.drawString(400, 170 - 18 * i, k); c.drawRightString(560, 170 - 18 * i, v)
    return _scan(draw)


SHINE_ITEMS = [("Cleaning Services", "1", "12 480.20"), ("Supervisor Mon - Fri 07:00 - 16:00", "1", ""), ("Cleaning Specialist Mon - Fri 07:00 - 16:00", "2", ""),
               ("Chemicals & Small Equipment", "1", "980.40"), ("Wet Floor Sign", "6", ""), ("Microfibre Cloths", "9", ""), ("Hand Gloves", "3", ""),
               ("Cleaning Equipment", "1", "410.15"), ("Mopping System", "3", ""), ("Rounding", "-1", "")]


def shine_pdf(no: str) -> bytes:
    def draw(c):
        c.setFont("Helvetica-Bold", 20); c.drawString(50, 735, "S H I N E   S E R V I C E S")
        c.setFont("Helvetica", 8); c.drawRightString(560, 745, "SHINE SERVICES (PTY) LTD"); c.drawRightString(560, 734, "1 Example Avenue, Midrand")
        c.setFont("Helvetica-Bold", 12); c.drawCentredString(306, 690, f"TAX INVOICE  {no}")
        c.setFont("Helvetica-Bold", 8); c.drawString(50, 668, "To:"); c.drawString(80, 668, f"{FULL.upper()} (PTY) LTD - DELL")
        c.setFont("Helvetica", 8); c.drawString(80, 656, "1 EXAMPLE ROAD"); c.drawString(80, 644, "DUNKELD WEST")
        for i, (k, v) in enumerate([("Date:", _d(INV_DAY)), ("Consultant:", "A. Person"), ("Client Order No:", "2093700001 EXP CLEANING"), ("Client Account No:", "80000999"),
                                    ("Client VAT No:", BILL_VAT)]):
            c.setFont("Helvetica-Bold", 8); c.drawString(360, 668 - 12 * i, k)
            c.setFont("Helvetica", 8); c.drawString(450, 668 - 12 * i, v)
        c.setFont("Helvetica-Bold", 9); c.drawString(55, 590, "Description"); c.drawRightString(560, 590, "Total Ex. VAT")
        c.setFont("Helvetica", 8); c.drawString(55, 576, f"PERIOD  01 {INV_DAY:%B %Y} to 28 {INV_DAY:%B %Y}")
        for i, (d, q, a) in enumerate(SHINE_ITEMS):
            y = 560 - 18 * i
            c.drawString(65, y, d); c.drawString(380, y, q)
            if a:
                c.drawRightString(480, y, a); c.drawRightString(560, y, a)
        for i, (k, v) in enumerate([("Total Ex. VAT:", "13 870.75"), ("VAT:", "2 080.61"), ("Unpaid Amount:", "15 951.36"), ("Total Inc. VAT", "ZAR 15 951.36")]):
            c.drawString(400, 330 - 16 * i, k); c.drawRightString(560, 330 - 16 * i, v)
        c.drawString(55, 250, "Banking Details: Example Bank  Account number: 00000000003  Branch Code: 000000")
    return _scan(draw)


# ---- what a vision model reads off them -------------------------------------------------------------------------------------
def _answer(no, supplier, vat, bill_to, po, our, lines, excl, vat_amt, incl):
    conf = {k: 93 for k in ("supplier", "bill_to", "invoice_no", "invoice_date", "total_incl", "total_excl", "vat")}
    return {"is_invoice": True, "credit_note": False, "supplier": supplier, "supplier_vat": vat, "bill_to": bill_to, "bill_to_vat": BILL_VAT, "invoice_no": no,
            "invoice_date": INV_DAY.isoformat(), "due_date": DUE_DAY.isoformat(), "po": po, "our_account": our, "currency": "ZAR", "lines": lines,
            "total_excl": excl, "vat": vat_amt, "total_incl": incl, "tax_invoice_words": True, "confidence": conf}


def answers(run: str) -> dict:
    return {
        f"bright_{run}.pdf": _answer(f"BL-{run}", "Brightline Studio (Pty) Ltd", "4100000017", f"{FULL} (Pty) Ltd - RIO2X21", "PO7700112026", "RIO2X21",
                                     [{"description": "260922 - RIO2X21 - MINING CLIENT - EXTRAS - 100% Invoice", "amount_excl": None, "vat_rate": None, "vat": None},
                                      {"description": "Project Value", "amount_excl": "21400.00", "vat_rate": 15, "vat": "3210.00"}],
                                     "21400.00", "3210.00", "24610.00"),
        f"peak_{run}.pdf": _answer(f"PW-{run}", "Peak Wellness (Pty) Ltd", "4200000016", FULL, "PO4040012026", "CBR001",
                                   [{"description": "Facility management", "amount_excl": "", "vat_rate": None, "vat": None},
                                    {"description": "Management fee", "amount_excl": "18250.00", "vat_rate": 15, "vat": None},
                                    {"description": "Equipment service", "amount_excl": "1240.50", "vat_rate": 15, "vat": None},
                                    {"description": "Assessments", "amount_excl": "3100.00", "vat_rate": 15, "vat": None}],
                                   "22590.50", "3388.58", "25979.08"),
        # several items with no amount and "Rounding -1" (the -1 is a quantity, not an amount): blank, never 0-lines that break the sum
        f"shine_{run}.pdf": _answer(f"SS-{run}", "Shine Services (Pty) Ltd", "4300000015", f"{FULL.upper()} (PTY) LTD - DELL", "2093700001 EXP CLEANING", "80000999",
                                    [{"description": d, "amount_excl": a.replace(" ", "") or None, "vat_rate": None, "vat": None} for d, _q, a in SHINE_ITEMS],
                                    "13870.75", "2080.61", "15951.36"),
    }


# ---- the tests -------------------------------------------------------------------------------------------------------------
def test_the_look_alikes_have_no_text_layer():
    for pdf in (brightline_pdf("X-1"), peak_pdf("X-2"), shine_pdf("X-3")):
        with pymupdf.open(stream=pdf, filetype="pdf") as d:
            assert d.page_count == 1 and d[0].get_text().strip() == "" and d[0].get_images()


def test_the_checks_leave_blank_amounts_out():
    a = answers("CHK")
    ok, f, conf, probs = C.parse(a["shine_CHK.pdf"], 80)
    assert ok and probs == [] and conf == 93, probs
    assert [x["excl"] for x in f["lines"]] == ["12480.20", "980.40", "410.15"] and f["our_account"] == "80000999" and f["po"] == "2093700001 EXP CLEANING"
    ok, f, conf, probs = C.parse(a["peak_CHK.pdf"], 80)
    assert ok and probs == [] and len(f["lines"]) == 3
    # a Rounding line the model read as an amount of -1 does break the sum: that one goes up to Seth, as it should
    bad = dict(a["shine_CHK.pdf"], lines=a["shine_CHK.pdf"]["lines"][:-1] + [{"description": "Rounding", "amount_excl": "-1", "vat_rate": None, "vat": None}])
    assert "the lines do not add up to the excl. total" in C.parse(bad, 80)[3]


def test_scans_go_to_the_capture_bots_and_come_out_as_green_kr_documents(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    fake = _vision(answers(run))
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", fake)
    u, adm = _user("scan.user@test.local", "Sana Scan"), _admin()
    b = _drop(u, {f"bright_{run}.pdf": brightline_pdf(f"BL-{run}"), f"peak_{run}.pdf": peak_pdf(f"PW-{run}"), f"shine_{run}.pdf": shine_pdf(f"SS-{run}")}, f"Scans {run}")
    # Sol: no text layer → manual capture (vision), never Odd
    assert {k: v["n"] for k, v in b["capture"]["buckets"].items() if v["n"]} == {"manual": 3}
    assert sorted(n for t, n, _, _ in fake.calls if t == "ap_capture") == sorted(answers(run)) and all(imgs >= 1 for _, _, imgs, _ in fake.calls)
    docs = _sheet(u, b["id"])
    for name, d in docs.items():
        assert d["status"] == "green" and d["entity"] == "X143", (name, d["issues"])
    caps = {n: u.get(f"/api/ap/captures/{docs[n]['id']}").json() for n in docs}
    # the bill-to "… (Pty) Ltd - RIO2X21": the company by the name before " - ", the site kept and in the line text (the coding's rule)
    bl = caps[f"bright_{run}.pdf"]
    assert bl["company_code"] == "X143" and bl["bill_to_site"] == "RIO2X21" and bl["match"]["entity"]["site"] == "RIO2X21"
    assert bl["fields"]["po"] == "PO7700112026" and bl["fields"]["our_account"] == "RIO2X21"
    rows = bl["sap_rows"]
    assert [(r["type"], r["pkey"], r["account"], r["amount"]) for r in rows] == [("KR", "31", "9910001", "24610.00"), ("KR", "40", "SC109", "21400.00"), ("KR", "40", "JA600", "3210.00")]
    assert rows[1]["text"].startswith("RIO2X21 ") and rows[1]["gl_name"] == "Consultation Fees - General" and rows[2]["gl_name"] == "Input VAT - Expense"
    assert rows[1]["costcentre"] == "AC0326" and rows[1]["cc_name"] == "SA - Operations"  # SAP's Cost Center List, in X143
    # "- DELL" with no line-text rule: the site leads the header text (1.0.102: within the 25 characters SAP keeps)
    sh = caps[f"shine_{run}.pdf"]
    assert sh["bill_to_site"] == "DELL" and sh["sap_rows"][0]["docheadertext"] == f"DELL Shine Serv. {P.mmm_yy(INV_DAY.isoformat())}"
    assert sh["sap_rows"][0]["amount"] == "15951.36" and sh["sap_rows"][1]["amount"] == "13870.75" and sh["sap_rows"][1]["gl_name"] == "Contract Charges - Cleaning Service"
    assert caps[f"peak_{run}.pdf"]["bill_to_site"] == "" and P.balance_fails([{**r, "_file": "x"} for r in caps[f"peak_{run}.pdf"]["sap_rows"]]) == (1, set())
    # the review workbook names every G/L
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(u.get(f"/api/ap/batches/{b['id']}/sap/review").content)).active
    heads = [c.value for c in ws[1]]
    gi = heads.index("G/L name")  # 1.0.104 (Builder R): the review's columns come after SAP's names
    names = {r[heads.index("account")]: (r[gi], r[gi + 1]) for r in ws.iter_rows(min_row=2, values_only=True)}
    assert heads[gi:gi + 2] == ["G/L name", "Cost centre name"] and names["SC109"] == ("Consultation Fees - General", "SA - Operations") and not names["9910001"][0]

    # a G/L the chart doesn't have: yellow in Tess's words (the chart export may be incomplete); 6123 is in no GL table either: red
    pk = docs[f"peak_{run}.pdf"]
    r = u.post(f"/api/ap/captures/{pk['id']}/decide", json={"action": "save", "data": {"fields": {"gl": "6123"}}})
    assert r.status_code == 200 and r.json()["sap_status"] == "red"
    assert {"level": "yellow", "text": "G/L 6123 is not in the chart of accounts — check with finance"} in r.json()["sap_issues"]
    # a cost centre SAP doesn't have in X143 is red; SAP's dummy too
    r = u.post(f"/api/ap/captures/{pk['id']}/decide", json={"action": "save", "data": {"fields": {"gl": "SC600", "cost_centre": "2822"}}})
    assert r.json()["sap_status"] == "red" and "Cost centre 2822 doesn't exist in company code X143 in SAP" in {x["text"] for x in r.json()["sap_issues"]}
    r = u.post(f"/api/ap/captures/{pk['id']}/decide", json={"action": "save", "data": {"fields": {"cost_centre": "SAP-DUMMY"}}})
    assert r.json()["sap_status"] == "red" and "Cost centre SAP-DUMMY is a dummy cost centre" in {x["text"] for x in r.json()["sap_issues"]}
    r = u.post(f"/api/ap/captures/{pk['id']}/decide", json={"action": "save", "data": {"fields": {"cost_centre": "AC0328"}}})
    assert r.json()["sap_status"] == "green", r.json()["sap_issues"]
    # a company set to YR (SAP: RE customer invoice, account types D S) can't take vendor lines: red
    with SessionLocal() as s:
        x143 = s.query(ApEntity).filter(ApEntity.code == "X143").one()
        body = {k: getattr(x143, k) for k in ("code", "name", "entity", "vat_no", "aliases", "doc_type", "doc_type_credit")}
        eid = x143.id
    try:
        assert adm.put(f"/api/ap/masters/entities/{eid}", json={**body, "doc_type": "YR"}).status_code == 200
        sd = docs[f"shine_{run}.pdf"]
        r = u.post(f"/api/ap/captures/{sd['id']}/decide", json={"action": "save", "data": {"fields": {}}})
        assert r.json()["sap_status"] == "red" and r.json()["sap_rows"][0]["type"] == "YR"
        assert "YR doesn't allow vendor lines in SAP (account types D, S)" in {x["text"] for x in r.json()["sap_issues"]}
    finally:
        adm.put(f"/api/ap/masters/entities/{eid}", json=body)
    with SessionLocal() as s:
        assert s.query(ApCapture).filter(ApCapture.batch_id == b["id"], ApCapture.uploaded_at.isnot(None)).count() == 0


def test_mia_prefers_the_vendor_number_opened_in_the_company_code():
    with SessionLocal() as s:
        s.query(ApVendorCode).filter(ApVendorCode.vat_no == "4400000014").delete()
        # two vendor numbers for one supplier: 120083 is opened in X120 in SAP's export, 120000 is not (and is the newer row)
        s.add(ApVendorCode(vendor_no="120083", name="Twin Supplies", name_key=C.M.name_key("Twin Supplies"), vat_no="4400000014", gl="SC600", source="excel"))
        s.flush()
        s.add(ApVendorCode(vendor_no="120000", name="Twin Supplies", name_key=C.M.name_key("Twin Supplies"), vat_no="4400000014", gl="SC600", source="excel"))
        s.commit()
        K.forget()
        assert K.vendor_in_company(s, "120083", "X120") is True and K.vendor_in_company(s, "120000", "X120") is False
        assert C.match_vendor(s, "Twin Supplies", "4400000014", "X120").vendor_no == "120083"
        assert C.match_vendor(s, "Twin Supplies", "4400000014", "X143").vendor_no == "120000"  # X143 is not in the export: the usual order
        s.query(ApVendorCode).filter(ApVendorCode.vat_no == "4400000014").delete()
        s.commit()


def test_mia_says_when_an_invoice_is_not_in_the_company_s_currency(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    a = answers(run)[f"bright_{run}.pdf"]
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision({f"usd_{run}.pdf": {**a, "currency": "USD"}}))
    u = _user("scan.user@test.local", "Sana Scan")
    b = _drop(u, {f"usd_{run}.pdf": brightline_pdf(f"BL-{run}")}, f"USD {run}")
    d = _sheet(u, b["id"])[f"usd_{run}.pdf"]
    assert d["status"] == "yellow" and "Invoice in USD, X143 books in ZAR" in {x["text"] for x in d["issues"]}
    with SessionLocal() as s:
        assert K.company_currency(s, "X143") == "ZAR" and K.company_currency(s, "x120") == "RWF" and K.company_currency(s, "T000") is None
