"""1.0.95, CBX-Connect 2.8.6's tests line for line. The AP Tool: a folder in → each file read, sorted to its bot and checked against the AP harness → the chain for
what nobody knows (bot → Paula → AXL → Needs you) → people fix, teach and sign off (four-eye on payment runs and
bank details) → what the team knows is learned only from signed-off files. Quinn drafts query answers."""
import io
import zipfile
from pathlib import Path
from datetime import date

from openpyxl import load_workbook

from cbx.ap import engine as E
from cbx.ap import service as S

from fastapi.testclient import TestClient

from cbx.main import app

from .conftest import _make, login
from .fixtures import make_ap

PDF = "application/pdf"


def _ap(owner, email, name, role="user", tools=("ap",)):
    ids = [t["id"] for t in owner.get("/api/users").json()["tools"] if t["slug"] in tools]
    pw = _make(owner, email, name, role, tools=ids)
    c = TestClient(app)
    login(c, email, pw)
    return c


def _drop(c, files: dict, name=""):
    r = c.post("/api/ap/batches", files=[("files", (n, d, "application/octet-stream")) for n, d in files.items()], data={"name": name})
    assert r.status_code == 200, r.text
    b = c.get(f"/api/ap/batches/{r.json()['id']}").json()
    assert b["status"] == "ready", b
    return b


def _by(b, name):
    return next(i for i in b["items"] if i["name"].endswith(name))


def _codes(c, iid):
    return {f["code"] for f in c.get(f"/api/ap/items/{iid}").json()["findings"]}


# ---- the engine on its own -----------------------------------------------------------------------------------
def test_engine_reads_sorts_and_checks_a_folder():
    items = []
    for n, d in make_ap.folder().items():
        doc = E.read(n, d)
        kind, conf, why = E.classify(n, doc)
        f = E.extract(kind, doc)
        f["_name"] = n
        items.append({"name": n, "kind": kind, "fields": f})
    kinds = {i["name"]: i["kind"] for i in items}
    assert kinds["acme_inv_10045.pdf"] == "invoice" and kinds["acme_credit.pdf"] == "credit_note"
    assert kinds["payment_run_proposal.xlsx"] == "payment_run" and kinds["acme_bank_letter.pdf"] == "bank_letter"
    assert kinds["acme_open_items.xlsx"] == "ledger" and kinds["acme_statement.pdf"] == "statement"
    assert kinds["query_inv_10045.eml"] == "query" and kinds["canteen_menu.txt"] == "unknown"
    inv = next(i for i in items if i["name"] == "acme_inv_10045.pdf")["fields"]
    assert inv["invoice_no"] == "INV-10045" and inv["supplier_vat"] == "4123456789" and inv["total"] == 11500.0 and inv["excl"] == 10000.0
    chk = {i["name"]: {x["code"] for x in E.check(i["kind"], i["fields"], batch=items)} for i in items}
    assert "duplicate_folder" in chk["acme_inv_10045.pdf"] and "vat_rate" in chk["acme_inv_10046.pdf"]
    assert {"duplicate_in_run", "zero_or_negative", "late", "new_supplier"} <= chk["payment_run_proposal.xlsx"]
    assert "bank_new" in chk["acme_bank_letter.pdf"] and "write_off" in chk["acme_open_items.xlsx"]
    stmt = next(i for i in items if i["kind"] == "statement")["fields"]
    assert stmt["recon"]["difference"] == 1400.0 and [x["ref"] for x in stmt["recon"]["only_statement"]] == ["INV-10050"]
    q = next(i for i in items if i["kind"] == "query")
    reply = E.draft_reply(q["fields"], E.lookup_refs(q["fields"]["refs"], items))
    assert "Good day Jane" in reply and "INV-10045" in reply and "payment run" in reply


def test_engine_money_dates_and_working_days():
    assert E.money("R 1 234,56") == 1234.56 and E.money("(500.00)") == -500.0 and E.money("1,500.00") == 1500.0
    assert E.to_date("05/03/2026") == date(2026, 3, 5) and E.to_date("5 Mar 2026") == date(2026, 3, 5) and E.to_date(46000) is not None
    assert E.working_days(date(2026, 10, 16), 2) == date(2026, 10, 20)  # Friday + 2 working days = Tuesday


# ---- a folder through the tool -----------------------------------------------------------------------------------
def test_a_folder_is_read_sorted_and_reported(owner):
    c = _ap(owner, "ivy.person@test.local", "Ivan Ivy")
    folder = make_ap.folder()
    folder[".DS_Store"] = b"junk"
    b = _drop(c, folder, "Mid-month 15 Oct")
    assert b["name"] == "Mid-month 15 Oct" and len(b["items"]) == 10  # the hidden file is dropped
    bot = {i["name"]: i["bot"] for i in b["items"]}
    assert bot["acme_inv_10045.pdf"] == "ivy" and bot["payment_run_proposal.xlsx"] == "bo" and bot["acme_open_items.xlsx"] == "rae" and bot["query_inv_10045.eml"] == "quinn"
    assert _by(b, "acme_inv_10046.pdf")["status"] == "exception" and "VAT" in _by(b, "acme_inv_10046.pdf")["top"]
    menu = _by(b, "canteen_menu.txt")
    assert menu["status"] == "unknown" and menu["level"] == 3 and menu["with"] == "Needs you"  # the AI is off in tests: straight up
    needs = c.get("/api/ap/items?level=3&open=true").json()["items"]
    assert any(i["id"] == menu["id"] for i in needs)
    assert c.get("/api/ap").json()["board"]["need_you"] >= 1
    # filter, search
    assert {i["kind"] for i in c.get(f"/api/ap/items?batch={b['id']}&bot=ivy").json()["items"]} == {"invoice", "credit_note"}
    assert [i["name"] for i in c.get(f"/api/ap/items?batch={b['id']}&q=statement").json()["items"]] == ["acme_statement.pdf"]
    # the Excel report
    wb = load_workbook(io.BytesIO(c.get(f"/api/ap/files/{b['report']['id']}").content))
    assert {"Summary", "Files", "Findings", "Invoices", "Payment run", "Ageing", "Queries"} <= set(wb.sheetnames)
    assert wb["Files"].max_row == 11


def test_a_person_fixes_and_teaches_and_the_team_learns(owner):
    c = _ap(owner, "rae.person@test.local", "Rina Rae")
    b = _drop(c, {"canteen_menu_week1.txt": b"Canteen: Monday soup, Tuesday curry."})
    it = b["items"][0]
    assert it["kind"] == "unknown"
    r = c.post(f"/api/ap/items/{it['id']}/fix", json={"kind": "other", "teach": "Canteen", "note": "Facilities' menu"})
    assert r.status_code == 200, r.text
    assert r.json()["kind"] == "other" and r.json()["status"] == "check" and r.json()["level"] == 0
    assert any(x["match"] == "Canteen" for x in c.get("/api/ap").json()["rules"])
    b2 = _drop(c, {"canteen_week2.txt": b"Canteen: Monday pap, Tuesday stew."})
    it2 = b2["items"][0]
    assert it2["kind"] == "other" and it2["confidence"] == 95 and it2["level"] == 0  # read by the taught rule
    assert c.post(f"/api/ap/items/{it2['id']}/sign", json={"action": "clear", "note": ""}).json()["status"] == "cleared"
    # fixing a field re-checks: the wrong VAT fixed by a person
    b3 = _drop(c, {"acme_inv_77001.pdf": make_ap.invoice(no="INV-77001", excl=5000.0, vat=900.0)})
    inv = b3["items"][0]
    assert "vat_rate" in _codes(c, inv["id"])
    r = c.post(f"/api/ap/items/{inv['id']}/fix", json={"fields": {"vat": "750.00", "total": "5750"}})
    assert "vat_rate" not in _codes(c, inv["id"]) and r.json()["fields"]["vat"] == 750.0


def test_four_eye_sign_off_and_what_the_team_learns(owner):
    a = _ap(owner, "bo.first@test.local", "Bongi First")
    b_ = _ap(owner, "bo.second@test.local", "Sipho Second")
    run = make_ap.payment_run()
    letter = make_ap.bank_letter()
    b = _drop(a, {"run_oct.xlsx": run, "acme_letter.pdf": letter})
    pr, bl = _by(b, "run_oct.xlsx"), _by(b, "acme_letter.pdf")
    assert pr["four_eye"] and pr["status"] == "exception"
    r = a.post(f"/api/ap/items/{pr['id']}/sign", json={"action": "check"})
    assert r.status_code == 400 and "verified" in r.json()["detail"]  # a stopped run needs a note
    assert a.post(f"/api/ap/items/{pr['id']}/sign", json={"action": "approve"}).status_code == 400  # not checked yet
    assert a.post(f"/api/ap/items/{pr['id']}/sign", json={"action": "check", "note": "Duplicate line removed in SAP; Zero Supplies line pulled."}).json()["status"] == "waiting"
    r = a.post(f"/api/ap/items/{pr['id']}/sign", json={"action": "approve"})
    assert r.status_code == 400 and "second, different person" in r.json()["detail"]
    r = b_.post(f"/api/ap/items/{pr['id']}/sign", json={"action": "approve"})
    assert r.status_code == 200 and r.json()["status"] == "cleared" and r.json()["approved_by"] == "Sipho Second"
    a.post(f"/api/ap/items/{bl['id']}/sign", json={"action": "check", "note": "Called Acme back on 011 555 0100 (on file); Standard Bank verified."})
    assert b_.post(f"/api/ap/items/{bl['id']}/sign", json={"action": "approve"}).json()["status"] == "cleared"
    # the next folder: the same invoice is "paid before", and a changed bank account stops the run
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Vendor", "Vendor name", "Reference", "Amount", "Due date", "Payment date", "Bank account"])
    ws.append(["100200", "Acme Cleaning (Pty) Ltd", "INV-10045", 11500.0, date.today(), date.today(), "9999888877"])
    ws.append(["100200", "Acme Cleaning (Pty) Ltd", "INV-20001", 95000.0, date.today(), date.today(), "9999888877"])
    buf = io.BytesIO()
    wb.save(buf)
    b2 = _drop(a, {"payment run nov.xlsx": buf.getvalue()})
    codes = _codes(a, b2["items"][0]["id"])
    assert {"paid_before", "bank_changed", "big_change"} <= codes and "new_supplier" not in codes


def test_the_chain_paula_then_axl_then_you(owner, monkeypatch):
    c = _ap(owner, "paula.person@test.local", "Pam Paula")
    calls = []

    def fake(task, system, user, schema=None, db=None):
        calls.append(task)
        if task == "ap_read":
            return {"kind": "unknown", "confidence": 20, "fields": {}, "reason": "no idea"}
        return {"kind": "invoice", "confidence": 80, "fields": {"invoice_no": "HW-55", "total": "R 2,300.00", "supplier": "Hardware Hut"}, "reason": "It bills for goods."}

    monkeypatch.setattr(S, "model_on", lambda: True)
    monkeypatch.setattr(S, "ask_model", fake)
    b = _drop(c, {"scan_0001.txt": b"Hardware Hut. 2 x ladders. Pay R2,300.00 please. Ref HW-55"})
    it = c.get(f"/api/ap/items/{b['items'][0]['id']}").json()
    assert calls == ["ap_read", "ap_axl"] and it["kind"] == "invoice" and it["level"] == 2 and it["with"] == "AXL"
    assert it["fields"]["invoice_no"] == "HW-55" and it["fields"]["total"] == 2300.0 and it["suggestion"]["kind"] == "invoice"
    assert "axl_read" in {f["code"] for f in it["findings"]} and it["status"] == "check"
    # nobody knows → you
    monkeypatch.setattr(S, "ask_model", lambda *a, **k: {"kind": "unknown", "confidence": 10, "fields": {}, "reason": "?"})
    b = _drop(c, {"scan_0002.txt": b"zzz qqq"})
    it = b["items"][0]
    assert it["level"] == 3 and it["status"] == "unknown"
    hist = " ".join(h["text"] for h in c.get(f"/api/ap/items/{it['id']}").json()["history"])
    assert "Sent up to AXL" in hist and "Sent to you" in hist


def test_quinn_drafts_and_a_person_answers(owner):
    c = _ap(owner, "quinn.person@test.local", "Quentin Q")
    f = make_ap.folder()
    b = _drop(c, {k: f[k] for k in ("query_inv_10045.eml", "acme_open_items.xlsx")})
    q = _by(b, "query_inv_10045.eml")
    r = c.post(f"/api/ap/items/{q['id']}/draft")
    assert r.status_code == 200 and "INV-10045" in r.json()["reply"] and "Accounts Payable" in r.json()["reply"]
    assert c.post(f"/api/ap/items/{q['id']}/reply", json={"text": "Good day Jane, paid Friday. AP"}).json()["reply"].startswith("Good day Jane")
    assert c.post(f"/api/ap/items/{q['id']}/sign", json={"action": "answered"}).json()["status"] == "cleared"
    assert c.post(f"/api/ap/items/{_by(b, 'acme_open_items.xlsx')['id']}/draft").status_code == 400  # Quinn answers queries only


def test_a_zip_folder_and_the_limits(owner):
    c = _ap(owner, "zip.person@test.local", "Zanele Zip")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("October/inv_a.pdf", make_ap.invoice(no="INV-90001"))
        z.writestr("October/__MACOSX/._inv_a.pdf", b"x")
        z.writestr("October/sub/menu.txt", b"hello")
    b = _drop(c, {"October.zip": buf.getvalue()})
    assert sorted(i["name"] for i in b["items"]) == ["October/October/inv_a.pdf", "October/October/sub/menu.txt"]
    r = c.post("/api/ap/batches", files=[("files", (".hidden", b"x", "text/plain"))])
    assert r.status_code == 400 and "empty" in r.json()["detail"]


def test_access_and_settings(owner):
    assert _ap(owner, "no.ap@test.local", "No Access", tools=()).get("/api/ap").status_code == 403
    c = _ap(owner, "set.user@test.local", "Sam Setter")
    assert c.put("/api/ap/settings", json={"vat_rate": 16}).status_code == 403
    adm = _ap(owner, "ap.admin@test.local", "Ada Admin", role="admin", tools=())
    r = adm.put("/api/ap/settings", json={"vat_rate": 15, "change_pct": 25, "nonsense": 1})
    assert r.status_code == 200 and r.json()["change_pct"] == 25.0 and "nonsense" not in r.json()
    assert adm.put("/api/ap/settings", json={"query_days": -1}).status_code == 400
    view = (Path(__file__).resolve().parents[1] / "web" / "views" / "ap.js").read_text(encoding="utf-8")  # 1.0.100: Windows reads cp1252 without it
    assert "/api/ap" in view and "AP Tool" in view
    team = adm.get("/api/ap").json()["team"]
    assert [b["name"] for b in team["bots"]] == ["Paula", "Ivy", "Bo", "Rae", "Quinn"] and team["chain"] == ["Bot", "Paula", "AXL", "Needs you"]
