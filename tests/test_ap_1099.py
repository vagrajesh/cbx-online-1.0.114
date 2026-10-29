"""1.0.99: from CBX-Connect 2.10.4 tests/test_ap_2104.py, on the old conftest.
1.0.99 AP fixes from the first real run (Edwin's test copy, 17 Oct): the AI off means wait (never people), the one-time
repair of what it sent to people, Needs you following the team, the row's subtitle, and the SAP upload styled like the
Harvester's grid. The AI is off in tests: the capture call is replaced where a reading is needed."""
from __future__ import annotations

import io
import random
import uuid
from datetime import timedelta

import pytest
from openpyxl import load_workbook

from cbx import compat2 as C2
from cbx.ap import capture as C
from cbx.ap import sap as P
from cbx.db import SessionLocal, utcnow
from cbx.models import AuditLog
from cbx.models.ap import ApCapture, ApEntity, ApItem

from .fixtures import make_ap
from .test_ap_capture_1098 import _admin, _app, _drop, _invoice_answer, _sheet, _user, _vision  # noqa: F401

FULL = "CBRE Excellerate Facilities Management"


@pytest.fixture(autouse=True)
def _known():
    make_ap.chart()
    with SessionLocal() as s:
        e = s.query(ApEntity).filter(ApEntity.code == "X143").one()
        if FULL not in (e.aliases or ""):
            e.aliases = (e.aliases + "\n" if e.aliases else "") + FULL
            s.commit()


def _vat() -> str:
    return "4" + "".join(random.choice("0123456789") for _ in range(9))


def _vendor(no, name, vat):
    r = _admin().post("/api/ap/masters/vendors", json={"vendor_no": no, "name": name, "vat_no": vat, "gl": "6100100", "cost_centre": "CC1", "tax_code": "I9"})
    assert r.status_code == 200, r.text


def _caps(bid):
    with SessionLocal() as s:
        return s.query(ApCapture).filter(ApCapture.batch_id == bid).order_by(ApCapture.id).all()


def _hist(item_id):
    with SessionLocal() as s:
        return [h["text"] for h in s.get(ApItem, item_id).history or []]


# ---- 1. AI off = wait, never people ------------------------------------------------------------------------------------
def test_with_the_ai_off_invoices_wait_and_are_read_once_it_is_on(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    vat = _vat()
    _vendor("302104", f"Waits {run}", vat)
    monkeypatch.setattr(C, "model_on", lambda: False)
    asked = []
    monkeypatch.setattr(C, "ask_vision", lambda *a, **k: asked.append(a) or None)
    u = _user(f"w2104.{run.lower()}@test.local", "Wait Person")
    files = {f"w{i}_{run}.pdf": make_ap.invoice(no=f"WT{i}-{run}") for i in range(2)}
    b = _drop(u, files, f"W2104 {run}")
    caps = _caps(b["id"])
    assert len(caps) == 2 and not asked  # no reader asked the AI
    for c in caps:
        assert (c.stage, c.needs, c.tries, c.claimed_at) == ("queued", "", 0, None)
        assert c.read.get("ai_off") and c.read.get("ai_down")
        assert _hist(c.item_id).count(C.AI_OFF_LINE) == 1
    assert not [x for x in u.get("/api/ap/needs").json()["unread"] if x["batch_id"] == b["id"]]  # never to people
    cap = u.get(f"/api/ap/batches/{b['id']}/sap").json()["capture"]
    assert cap["ai_wait"] == 2 and cap["ai_on"] is False
    docs = _sheet(u, b["id"])
    assert all(d["ai_wait"] for d in docs.values())
    live = {x["key"]: x for x in u.get("/api/ap/live").json()["bots"]}
    assert all(live[k]["status"] == "ai_wait" for k in ("cara", "dex", "eli"))
    # the wake routine and a capture job while it is still off: nothing read, no try, no second history line — and a
    # long wait for the switched-off AI is no outage (no "waited too long")
    with SessionLocal() as s:
        for c in s.query(ApCapture).filter(ApCapture.batch_id == b["id"]).all():
            c.read = {**c.read, "ai_down": (utcnow() - C.AI_PATIENCE - timedelta(hours=1)).isoformat()}
        s.commit()
    C.wake_routine()
    C.capture_job(0, {"worker": "cara"})
    C.capture_job(0, {"worker": "seth"})
    for c in _caps(b["id"]):
        assert (c.stage, c.needs, c.tries) == ("queued", "", 0) and _hist(c.item_id).count(C.AI_OFF_LINE) == 1
    # the AI is on again: the next wake round reads them, without "Read again"
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision({n: _invoice_answer(f"WT{i}-{run}", supplier=f"Waits {run}", vat_no=vat) for i, n in enumerate(files)}))
    C.wake_routine()
    for c in _caps(b["id"]):
        assert (c.stage, c.sap_status, c.tries) == ("ready", "green", 1)
        assert "ai_off" not in c.read and "ai_down" not in c.read
    assert u.get(f"/api/ap/batches/{b['id']}/sap").json()["capture"]["ai_wait"] == 0


def test_a_reader_that_takes_one_while_the_ai_is_off_puts_it_back(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    monkeypatch.setattr(C, "model_on", lambda: False)
    u = _user(f"w2104b.{run.lower()}@test.local", "Wait Person B")
    b = _drop(u, {f"s_{run}.pdf": make_ap.invoice(no=f"SN-{run}")}, f"W2104b {run}")
    c = _caps(b["id"])[0]
    with SessionLocal() as s:  # as if it had reached Seth before the AI went off, not marked yet
        x = s.get(ApCapture, c.id)
        x.stage, x.read = "senior", {}
        s.commit()
    monkeypatch.setattr(C, "claim", lambda f, busy, w: c.id if f == "senior" and _set(c.id, busy, w) else None)
    assert C.step("seth") == c.id
    c = _caps(b["id"])[0]
    assert (c.stage, c.needs, c.tries, c.claimed_at) == ("senior", "", 0, None) and c.read.get("ai_off")
    monkeypatch.undo()
    assert u.post(f"/api/ap/captures/{c.id}/decide", json={"action": "skip", "data": {}}).status_code == 200  # tidy for later tests


def _set(cid, stage, worker):
    with SessionLocal() as s:
        x = s.get(ApCapture, cid)
        x.stage, x.worker, x.claimed_at = stage, worker, utcnow()
        s.commit()
    return True


# ---- 2. the one-time repair --------------------------------------------------------------------------------------------
def test_the_repair_sends_invoices_the_ai_off_gave_to_people_back_to_the_queue(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    vat = _vat()
    _vendor("302105", f"Repair {run}", vat)
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision({f"off_{run}.pdf": _invoice_answer(f"RP-{run}", conf=40), f"hard_{run}.pdf": _invoice_answer(f"HD-{run}", conf=40)}))
    u = _user(f"r2104.{run.lower()}@test.local", "Repair Person")
    b = _drop(u, {f"off_{run}.pdf": make_ap.invoice(no=f"RP-{run}"), f"hard_{run}.pdf": make_ap.invoice(no=f"HD-{run}")}, f"R2104 {run}")
    caps = {c.id: c for c in _caps(b["id"])}
    assert all((c.stage, c.needs) == ("needs", "unread") for c in caps.values())
    with SessionLocal() as s:  # the 16 Oct state: one read with the AI off (before 1.0.99), one the senior truly couldn't read
        off = next(c for c in s.query(ApCapture).filter(ApCapture.batch_id == b["id"]).all() if s.get(ApItem, c.item_id).name.startswith("off_"))
        off.read = {"cara": {"problems": ["the AI is off"], "confidence": 0}, "seth": {"problems": ["the AI is off"], "confidence": 0}}
        off.tries, off.fields = 2, {}
        it = s.get(ApItem, off.item_id)
        it.history = [*it.history, {"at": utcnow().isoformat(), "who": "Seth", "text": "Could not read it either (the AI is off). A person types it in (Needs you)."}]
        s.commit()
        off_id = off.id
    monkeypatch.setattr(C, "ask_vision", _vision({f"off_{run}.pdf": _invoice_answer(f"RP-{run}", supplier=f"Repair {run}", vat_no=vat)}))
    n = C.run_requeue_ai_off(force=True)
    assert n == 1
    caps = {c.id: c for c in _caps(b["id"])}
    fixed = caps.pop(off_id)
    assert (fixed.stage, fixed.sap_status, fixed.tries, fixed.needs) == ("ready", "green", 1, "")  # read by itself, fresh
    assert any("Back in the queue" in h for h in _hist(fixed.item_id))
    assert all((c.stage, c.needs) == ("needs", "unread") for c in caps.values())  # the truly unreadable one stays with people
    with SessionLocal() as s:
        row = s.query(AuditLog).filter(AuditLog.action == "ap_requeued_ai_off").order_by(AuditLog.id.desc()).first()
        assert row is not None and "1" in row.detail
        assert C2.state_get(s, C.REPAIR_AI_OFF) is not None
    assert C.run_requeue_ai_off() is None  # once


# ---- 3 and 4. Needs you: follows the team; the row's subtitle -------------------------------------------------------------
def test_needs_you_rows_name_their_batch_and_how_long_they_waited(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    monkeypatch.setattr(C, "model_on", lambda: False)
    u = _user(f"n2104.{run.lower()}@test.local", "Needs Person")
    named = _drop(u, {f"menu_{run}.txt": f"Canteen {run}: Monday soup.".encode()}, f"October run {run}")
    plain = _drop(u, {f"menu2_{run}.txt": f"Canteen {run}: Tuesday soup.".encode()}, "")
    assert C.unnamed(plain["name"]) and not C.unnamed(named["name"])
    n = u.get("/api/ap/needs").json()
    assert "busy" in n
    rows = {x["batch_id"]: x for x in n["odd"] if x["batch_id"] in (named["id"], plain["id"])}
    assert rows[named["id"]]["batch_named"] is True and rows[named["id"]]["batch"] == f"October run {run}"
    assert rows[plain["id"]]["batch_named"] is False and rows[plain["id"]]["batch_at"]
    assert all(r["waiting_since"] for r in rows.values())
    # a capture that leaves Needs you (sent to manual capture: back in the queue, waiting for the AI) is gone from it
    assert u.post(f"/api/ap/captures/{rows[plain['id']]['id']}/decide", json={"action": "manual", "data": {}}).status_code == 200
    n = u.get("/api/ap/needs").json()
    assert not [x for k in ("odd", "unread", "other", "red") for x in n[k] if x["id"] == rows[plain["id"]]["id"]]
    assert C.unnamed("3 files") and C.unnamed("1 file") and not C.unnamed("Sept invoices")
    # tidy: the waiting one is skipped, so later tests' readers never pick it up
    assert u.post(f"/api/ap/captures/{rows[plain['id']]['id']}/decide", json={"action": "skip", "data": {}}).status_code == 200


# ---- 5. the SAP upload, styled like the Harvester's grid -------------------------------------------------------------------
def _doc(no, vendor, comp, status, credit=False, amount="80086.12"):
    gl = P.dec(amount) - P.dec("10446.02")
    base = {"docdate": "15092026", "postdate": "15092026", "type": "KG" if credit else "KR", "comp": comp, "curr": "ZAR", "refdocnr": no,
            "docheadertext": "=Bela Bela ETE", "sgl": "", "compn": "", "txi": "", "ordernr": "", "DueDate": "30092026", "_tess": status}
    return [{**base, "TransCode": "FB01", "pkey": "21" if credit else "31", "account": vendor, "amount": amount, "tc": "", "dt": "", "costcentre": "",
             "allocation": "7267", "text": "7267_12-06", "ProfitCentre": "07267"},
            {**base, "TransCode": "ZB01", "pkey": "40", "account": "SA110", "amount": P.m2(gl), "tc": "I9", "dt": "", "costcentre": "07267",
             "allocation": "7267", "text": "7267_12-06", "ProfitCentre": ""},
            {**base, "TransCode": "ZB01", "pkey": "40", "account": "JA600", "amount": "10446.02", "tc": "I9", "dt": "X", "costcentre": "",
             "allocation": "7267", "text": "", "ProfitCentre": "7267"}]


def _plain(rows):
    buf = io.BytesIO()
    P.export_sap_xlsx(buf, rows)
    return load_workbook(io.BytesIO(buf.getvalue()))


def test_the_sap_upload_is_the_same_file_styled_like_the_harvesters_grid():
    rows = _doc("0001009715202606", "0000709642", "T000", "green") + _doc("0001009715202605", "709642", "0143", "yellow", amount="87091.52") \
        + _doc("0001009715202604", "709642", "T000", "red", credit=True, amount="84418.50")
    wb = load_workbook(io.BytesIO(P.upload_file(rows)))
    old = _plain(rows).active
    assert wb.sheetnames == ["Sheet1"]
    ws = wb["Sheet1"]
    assert [c.value for c in ws[1]] == P.SAP_HEADERS and ws.max_column == 22 and ws.max_row == old.max_row == 10
    a = P.SAP_HEADERS.index("amount")
    for new_row, old_row in zip(ws.iter_rows(min_row=2), old.iter_rows(min_row=2)):
        for i, (n, o) in enumerate(zip(new_row, old_row)):
            if i == a:
                assert isinstance(n.value, float) and n.value == float(o.value) and n.number_format == "0.00"
            else:
                assert n.value == o.value, (P.SAP_HEADERS[i], n.value, o.value)
    for h in P.TEXT_COLS:  # text, leading zeros kept
        col = P.SAP_HEADERS.index(h)
        assert all(isinstance(r[col].value, str) for r in ws.iter_rows(min_row=2) if r[col].value not in (None, ""))
    assert ws.cell(2, P.SAP_HEADERS.index("refdocnr") + 1).value == "0001009715202606"
    assert ws.cell(2, P.SAP_HEADERS.index("account") + 1).value == "0000709642" and ws.cell(2, P.SAP_HEADERS.index("ProfitCentre") + 1).value == "07267"
    assert ws.cell(2, 2).value == "15092026"  # dates stay ddmmyyyy text
    assert ws.cell(2, P.SAP_HEADERS.index("docheadertext") + 1).data_type == "s"  # xlsafe: "=Bela…" is text, not a formula
    # header: dark green, bold white, filters over the whole table, frozen
    for c in ws[1]:
        assert c.fill.fgColor.rgb.endswith("003F2D") and c.font.bold and c.font.color.rgb.endswith("FFFFFF")
    assert ws.auto_filter.ref == "A1:V10" and ws.freeze_panes == "A2"
    # each document's vendor line shaded light green, its TransCode in Tess's colour; the 40/50 lines white
    for r, st in ((2, "63BE7B"), (5, "FFD966"), (8, "F8696B")):
        assert ws.cell(r, 1).fill.fgColor.rgb.endswith(st)
        assert all(ws.cell(r, k).fill.fgColor.rgb.endswith("E3F1E8") for k in range(2, 23))
        for k in range(1, 23):
            assert ws.cell(r + 1, k).fill.fill_type is None and ws.cell(r + 2, k).fill.fill_type is None
    assert all(0 < (ws.column_dimensions[c.column_letter].width or 0) <= 40 for c in ws[1])


def test_the_downloaded_upload_and_the_review_workbook_carry_the_style(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    vat, no = _vat(), "000" + "".join(random.choice("0123456789") for _ in range(9))
    _vendor("0302106", f"Styled {run}", vat)
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision({f"st_{run}.pdf": _invoice_answer(no, supplier=f"Styled {run}", vat_no=vat)}))
    u = _user(f"s2104.{run.lower()}@test.local", "Style Person")
    b = _drop(u, {f"st_{run}.pdf": make_ap.invoice(no=no)}, f"S2104 {run}")
    assert _sheet(u, b["id"])[f"st_{run}.pdf"]["status"] == "green"
    rv = load_workbook(io.BytesIO(u.get(f"/api/ap/batches/{b['id']}/sap/review").content)).active
    assert rv[1][0].fill.fgColor.rgb.endswith("003F2D") and rv[1][0].font.bold and rv.auto_filter.ref.startswith("A1:")
    rows = _caps(b["id"])[0].sap_rows
    r = u.post(f"/api/ap/batches/{b['id']}/sap/upload", json={"include_yellow": False})
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content))["Sheet1"]
    old = _plain(rows).active
    assert [c.value for c in ws[1]] == P.SAP_HEADERS
    a = P.SAP_HEADERS.index("amount")
    for new_row, old_row in zip(ws.iter_rows(min_row=2), old.iter_rows(min_row=2)):
        assert [c.value for i, c in enumerate(new_row) if i != a] == [c.value for i, c in enumerate(old_row) if i != a]
        assert new_row[a].value == float(old_row[a].value)
    assert ws.cell(2, P.SAP_HEADERS.index("refdocnr") + 1).value == no  # leading zeros kept
    assert ws.cell(2, P.SAP_HEADERS.index("account") + 1).value == "0302106"
    assert ws.cell(2, 1).fill.fgColor.rgb.endswith("63BE7B") and ws.cell(2, 2).fill.fgColor.rgb.endswith("E3F1E8")
    assert ws.cell(3, 2).fill.fill_type is None
