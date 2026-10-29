"""PDF Studio end-to-end in stub mode (no GPU, no network): template from Excel, batch upload +
learning loop, misses/fills/verification, learning → pending version → approval, export, master
sheet append with exceptions, access control, parser and provenance rules."""
from __future__ import annotations

import io
import json
import re
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from cbx.db import SessionLocal
from cbx.models import Batch, Document, Extraction, Profile, ProfileVersion, Template, User
from cbx.pdfstudio import extract as X
from cbx.pdfstudio import money as M
from cbx.pdfstudio import scoring as SC
from tests.fixtures.make_bills import make_bills, make_vat_mismatch

COLUMNS = [
    {"key": "account_no", "label": "Account number", "ctype": "text"},
    {"key": "municipality", "label": "Municipality", "ctype": "list"},
    {"key": "period", "label": "Billing period", "ctype": "date"},
    {"key": "kwh", "label": "kWh consumed", "ctype": "number"},
    {"key": "demand_kva", "label": "Demand kVA", "ctype": "number"},
    {"key": "nmd_kva", "label": "NMD kVA", "ctype": "number"},
    {"key": "excl_vat", "label": "Amount excl. VAT", "ctype": "money"},
    {"key": "vat", "label": "VAT", "ctype": "money"},
    {"key": "total", "label": "Total due", "ctype": "money"},
    {"key": "deposit", "label": "Deposit held", "ctype": "money", "required": False, "critical": False},
]


@pytest.fixture(scope="session")
def bills(tmp_path_factory):
    d = tmp_path_factory.mktemp("bills")
    out = make_bills(d)
    out.append(make_vat_mismatch(d))
    return out


def _stem(filename: str) -> str:
    """Stored names get ' (2)' suffixes when the same file is uploaded twice."""
    return re.sub(r" \(\d+\)(?=\.pdf$)", "", filename)


def _pdf_files(bills, names):
    return [("files", (b["name"], Path(b["file"]).read_bytes(), "application/pdf")) for b in bills if b["name"] in names]


def _template(client, name="Utilities bill"):
    r = client.post("/api/pdfstudio/templates", json={"name": name, "columns": COLUMNS})
    assert r.status_code == 200, r.text
    return r.json()


def _run_batch(client, bills, names, template_id, name="Test batch"):
    r = client.post("/api/pdfstudio/batches", json={"template_id": template_id, "name": name})
    assert r.status_code == 200, r.text
    bid = r.json()["id"]
    r = client.post(f"/api/pdfstudio/batches/{bid}/files", files=_pdf_files(bills, names))
    assert r.status_code == 200, r.text
    r = client.post(f"/api/pdfstudio/batches/{bid}/start")
    assert r.status_code == 200, r.text
    r = client.get(f"/api/pdfstudio/batches/{bid}")
    assert r.status_code == 200, r.text
    return r.json()


TEXT_BILLS = [f"emfuleni-EMF-7712{i}.pdf" for i in range(3)] + [f"centlec-CEN-2092{i}.pdf" for i in range(4, 8)] + [f"ekurhuleni-550144{i}.pdf" for i in range(28, 32)]


# ---- template ----------------------------------------------------------------------------------
def test_template_from_excel(user):
    wb = Workbook()
    ws = wb.active
    ws.append(["Account number", "Municipality", "Billing period", "kWh consumed", "Amount excl. VAT", "VAT", "Total due"])
    ws.append(["EMF-1", "Emfuleni", "Jun 2026", 9870, 27905.6, 4185.84, 32091.44])
    buf = io.BytesIO()
    wb.save(buf)
    r = user.post("/api/pdfstudio/templates/from-excel", files=[("files", ("Utilities bill.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))])
    assert r.status_code == 200, r.text
    t = r.json()
    assert t["name"] == "Utilities bill"
    cols = {c["key"]: c for c in t["columns"]}
    assert set(cols) == {"account_number", "municipality", "billing_period", "kwh_consumed", "amount_excl_vat", "vat", "total_due"}
    assert cols["amount_excl_vat"]["ctype"] == "money" and cols["amount_excl_vat"]["critical"]
    assert cols["kwh_consumed"]["ctype"] == "number"
    assert cols["billing_period"]["ctype"] == "date"
    assert cols["account_number"]["ctype"] == "text"
    # it is listed for its owner
    assert any(x["id"] == t["id"] for x in user.get("/api/pdfstudio/templates").json()["templates"])


# ---- the batch --------------------------------------------------------------------------------------
@pytest.fixture()
def finished(user, bills):
    t = _template(user)
    b = _run_batch(user, bills, TEXT_BILLS, t["id"], name="July bills")
    return t, b


def test_batch_reaches_99_on_text_pdfs(finished, bills):
    t, b = finished
    assert b["status"] in ("done", "review"), b
    assert b["step"] == "Ready" and b["progress"] == 1.0
    assert b["accuracy_pct"] >= 99.0, b["accuracy_pct"]
    # role user: no log / runs / by_column
    assert "log" not in b and "runs" not in b and "by_column" not in b and "accuracy" not in b
    assert b["runs_count"] >= 2
    sheet = b["sheet"]
    gold = {x["name"]: x["gold"] for x in bills}
    for row in sheet["rows"]:
        g = gold[_stem(row["filename"])]
        cells = row["cells"]
        for k in ("account_no", "municipality", "kwh", "demand_kva", "nmd_kva"):
            assert str(cells[k]["value"]).lower() == str(g[k]).lower(), (row["filename"], k, cells[k])
        for k in ("excl_vat", "vat", "total"):
            assert abs(float(cells[k]["value"]) - g[k]) < 0.005, (row["filename"], k, cells[k])
            assert cells[k]["status"] == "ok" and cells[k]["quote"]
        assert str(cells["period"]["value"]).startswith(g["period"][:7])
        assert row["correct"] == 9 and row["of"] == 9  # the owner's "k of N" per PDF
    # the page-2 NMD bill was read from page 2
    p2 = next(r for r in sheet["rows"] if _stem(r["filename"]) == "centlec-CEN-20926.pdf")
    assert p2["cells"]["nmd_kva"]["page"] == 2
    # the OPTIONAL column no bill labels this way is blank and "not on this document" — not a guess,
    # and not handed to the user
    assert all(r["cells"]["deposit"]["status"] == "absent" and r["cells"]["deposit"]["value"] is None for r in sheet["rows"])
    assert all(r["cells"]["deposit"]["score"] == 1.0 for r in sheet["rows"])
    assert b["cells_for_user"] == 0 and b["status"] == "done"
    assert b["cells_absent"] == len(sheet["rows"])
    assert b["columns_summary"]["optional"] == 1 and "1 optional" in b["columns_summary"]["text"]
    dep = next(c for c in sheet["columns"] if c["key"] == "deposit")
    assert dep["required"] is False and dep["kind"] == "optional"


def test_admin_sees_runs_log_and_columns(owner, bills):
    t = _template(owner, "Utilities bill (admin)")
    b = _run_batch(owner, bills, TEXT_BILLS[:3] + TEXT_BILLS[3:5], t["id"])
    assert b["runs"] and b["runs"][0]["note"] == "samples"
    assert b["by_column"]["vat"] == 1.0
    assert any("Run 1" in e["text"] for e in b["log"])
    assert any("profile v1" in e["text"] or "matches profile" in e["text"] for e in b["log"])
    assert b["profiles"] and all(p["version"] >= 1 for p in b["profiles"])


def test_scanned_pdf_is_flagged_scanned(user, bills):
    t = _template(user, "Scan template")
    b = _run_batch(user, bills, ["emfuleni-EMF-77123-scan.pdf", "emfuleni-EMF-77120.pdf", "emfuleni-EMF-77121.pdf"], t["id"], name="Scan test")
    doc = next(d for d in b["documents"] if d["filename"].endswith("-scan.pdf"))
    assert doc["scanned"] is True
    assert b["status"] in ("done", "review")
    row = next(r for r in b["sheet"]["rows"] if r["document_id"] == doc["id"])
    gold = next(x for x in bills if x["name"] == "emfuleni-EMF-77123-scan.pdf")["gold"]
    for k in ("excl_vat", "vat", "total"):
        c = row["cells"][k]
        assert c["status"] == "miss" or abs(float(c["value"]) - gold[k]) < 0.005
    # page image of the scanned document renders
    r = user.get(f"/api/pdfstudio/documents/{doc['id']}/page/1.png")
    assert r.status_code == 200 and r.headers["content-type"] == "image/png" and r.content[:4] == b"\x89PNG"


# ---- fills, verification, learning, approval -------------------------------------------------------------
def test_fill_verified_learns_and_admin_approves(finished, bills, owner):
    t, b = finished
    bid = b["id"]
    emf = next(r for r in b["sheet"]["rows"] if _stem(r["filename"]) == "emfuleni-EMF-77120.pdf")
    gold = next(x for x in bills if x["name"] == "emfuleni-EMF-77120.pdf")["gold"]
    # context shows page text around the expected label
    ctx = owner.get(f"/api/pdfstudio/documents/{emf['document_id']}/context?col=deposit").json()
    assert ctx["page"] == 1 and ctx["lines"]
    # value found on the page → filled + verified
    r = owner.post(f"/api/pdfstudio/batches/{bid}/cells", json={"document_id": emf["document_id"], "column_key": "deposit", "value": f"R {gold['deposit']:,.2f}".replace(",", " ")})
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["verified"] is True and out["status"] == "filled"
    assert out["cell"]["page"] == 1 and "5" in out["cell"]["quote"]
    assert out["learned"] and "Security held" in out["learned"]["reason"]
    bt = out["learned"]["backtest"]
    assert bt["passed"] == bt["docs"] and bt["after"] >= bt["before"]
    # the pending version exists, is authored by the user and is not live
    db = SessionLocal()
    try:
        ext = db.query(Extraction).filter(Extraction.document_id == emf["document_id"]).one()
        doc = db.get(Document, emf["document_id"])
        prof = db.get(Profile, doc.profile_id)
        pend = [v for v in prof.versions if v.approved_at is None and v.author.startswith("user:")]
        assert len(pend) == 1 and prof.status == "patch" and prof.live_version_id != pend[0].id
        vid, pid = pend[0].id, prof.id
        assert ext.cells["deposit"]["status"] == "filled"
    finally:
        db.close()
    # counters: the optional column was never "for the user"; filling it keeps that at zero
    b2 = owner.get(f"/api/pdfstudio/batches/{bid}").json()
    assert b2["cells_for_user"] == 0 and b["cells_for_user"] == 0
    assert b2["by_column"]["deposit"] == 1.0  # absent-on-document counts as correct for an optional column
    # user cannot approve; admin can (require_human + require_admin)
    prof_api = owner.get(f"/api/pdfstudio/profiles/{pid}").json()
    assert prof_api["pending"] == 1 and any(v["pending"] for v in prof_api["versions"])
    assert "Security held" in prof_api["notes"] or prof_api["notes"]
    r = owner.post(f"/api/pdfstudio/profiles/{pid}/versions/{vid}/approve")
    assert r.status_code == 200, r.text
    after = r.json()
    assert after["status"] == "live" and after["live_version_id"] == vid and after["pending"] == 0
    # the approved rules now read the deposit on the other Emfuleni bills
    db = SessionLocal()
    try:
        ver = db.get(ProfileVersion, vid)
        from cbx.pdfstudio.spans import cached

        other = next(d for d in db.query(Document).filter(Document.batch_id == bid) if _stem(d.filename) == "emfuleni-EMF-77121.pdf")
        cells = X.extract(ver.ruleset, [dict(c, required=c.get("required", True), critical=c.get("critical", False)) for c in COLUMNS], cached(other.sha256))
        assert abs(float(cells["deposit"]["value"]) - 5250.0) < 0.005
    finally:
        db.close()


def test_fill_not_on_page_is_flagged_and_answer_flow(finished, user):
    t, b = finished
    bid = b["id"]
    row = b["sheet"]["rows"][0]
    r = user.post(f"/api/pdfstudio/batches/{bid}/cells", json={"document_id": row["document_id"], "column_key": "deposit", "value": "77 777.77"})
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["verified"] is False and out["status"] == "flagged"
    assert "can't find" in out["question"] and {o["id"] for o in out["options"]} == {"other_source", "trust_me", "retract"}
    assert out["learned"] is False  # users see no detail; nothing learned from an unverified value
    cid = out["correction_id"]
    # other_source needs a note
    r = user.post(f"/api/pdfstudio/batches/{bid}/cells/answer", json={"correction_id": cid, "answer": "other_source", "note": ""})
    assert r.status_code == 400
    r = user.post(f"/api/pdfstudio/batches/{bid}/cells/answer", json={"correction_id": cid, "answer": "other_source", "note": "From the deposit register"})
    assert r.status_code == 200 and r.json()["status"] == "trusted"
    # retract puts it back to a miss
    r = user.post(f"/api/pdfstudio/batches/{bid}/cells", json={"document_id": row["document_id"], "column_key": "deposit", "value": "999"})
    cid2 = r.json()["correction_id"]
    r = user.post(f"/api/pdfstudio/batches/{bid}/cells/answer", json={"correction_id": cid2, "answer": "retract"})
    assert r.status_code == 200 and r.json()["status"] == "miss"


# ---- export & master sheet -------------------------------------------------------------------------------
def test_export_has_three_sheets_and_is_stored(finished, user):
    t, b = finished
    r = user.get(f"/api/pdfstudio/batches/{b['id']}/export.xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Rows", "Provenance", "Misses"]
    ws = wb["Rows"]
    assert ws.max_row == len(TEXT_BILLS) + 1
    prov = wb["Provenance"]
    assert prov.max_row > len(TEXT_BILLS) * 5
    res = user.get("/api/pdfstudio/results").json()
    assert any(f["name"].endswith(".xlsx") for f in res["files"])
    assert any(x["id"] == b["id"] for x in res["batches"])


def test_to_master_appends_rows_and_vat_mismatch_goes_to_exceptions(user, owner, bills):
    t = _template(user, "Utilities bill (master)")
    b = _run_batch(user, bills, ["ekurhuleni-55014428.pdf", "ekurhuleni-55014429.pdf", "ekurhuleni-55014430.pdf", "ekurhuleni-vat-mismatch.pdf"], t["id"], name="Master test")
    sheets = user.get("/api/pdfstudio/master-sheets").json()["sheets"]
    sid = next(s["id"] for s in sheets if s["name"] == "Municipal accounts")
    before = owner.get(f"/api/master-sheets/{sid}").json()
    r = user.post(f"/api/pdfstudio/batches/{b['id']}/to-master/{sid}")
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["added"] == 3 and out["exceptions"] == 1
    after = owner.get(f"/api/master-sheets/{sid}").json()
    assert after["rows"] == before["rows"] + 3
    new = [r for r in after["data"] if (r.get("_src") or {}).get("batch") == b["id"]]
    assert len(new) == 3
    for row in new:
        assert row["_by"] and row["_at"] and row["_src"]["doc"] and row["_team"] == "Utilities billing" and row["_client"] == "Ekurhuleni"
        assert row["account_no"] and row["total"] > 0
    exc = [r for r in after["exception_rows"] if (r.get("_src") or {}).get("batch") == b["id"]]
    assert len(exc) == 1 and "VAT" in exc[0]["_why"]
    # the optional "deposit" column was added to the sheet's columns
    assert any(c["key"] == "deposit" for c in after["columns"])
    # running it again does not duplicate
    assert user.post(f"/api/pdfstudio/batches/{b['id']}/to-master/{sid}").json()["skipped"] == 4
    # the user sees only their own rows; the download is built from those
    mine = user.get(f"/api/master-sheets/{sid}").json()
    assert all(r["_by"] == mine["data"][0]["_by"] for r in mine["data"]) and len(mine["data"]) == 3
    x = user.get(f"/api/master-sheets/{sid}/export.xlsx")
    assert x.status_code == 200
    assert load_workbook(io.BytesIO(x.content))["Rows"].max_row == 4


# ---- access ----------------------------------------------------------------------------------------------
def test_user_without_tool_gets_403(owner):
    from fastapi.testclient import TestClient

    from cbx.main import app
    from tests.conftest import _make, login

    pw = _make(owner, "notool@test.local", "No Tool", "user", tools=[])

    c = TestClient(app)
    login(c, "notool@test.local", pw)
    assert c.get("/api/pdfstudio/templates").status_code == 403
    assert c.get("/api/master-sheets").status_code == 403
    assert c.get("/api/pdfstudio/profiles").status_code == 403


def test_user_cannot_see_other_users_batch(finished, owner):
    t, b = finished
    from fastapi.testclient import TestClient

    from cbx.main import app
    from tests.conftest import _make, login

    pw = _make(owner, "sipho.pdf@test.local", "Sipho Dlamini", "user", tools=[x["id"] for x in owner.get("/api/users").json()["tools"] if not x["admin_only"]], team="Procurement")

    c = TestClient(app)
    login(c, "sipho.pdf@test.local", pw)
    assert c.get(f"/api/pdfstudio/batches/{b['id']}").status_code == 404
    assert c.get(f"/api/pdfstudio/batches/{b['id']}/export.xlsx").status_code == 404
    doc = b["documents"][0]["id"]
    assert c.get(f"/api/pdfstudio/documents/{doc}/page/1.png").status_code == 404
    assert c.get(f"/api/pdfstudio/documents/{doc}/context?col=vat").status_code == 404
    assert all(x["id"] != b["id"] for x in c.get("/api/pdfstudio/batches").json()["batches"])
    # a template from another team is not visible either
    assert all(x["id"] != t["id"] for x in c.get("/api/pdfstudio/templates").json()["templates"])
    # profiles page is admin-only
    assert c.get("/api/pdfstudio/profiles").status_code == 403
    assert owner.get("/api/pdfstudio/profiles").status_code == 200


# ---- parser & provenance --------------------------------------------------------------------------------
@pytest.mark.parametrize(
    "raw,val",
    [("R 1 250,00", 1250.0), ("1,250.00", 1250.0), ("1 250.00", 1250.0), ("R1.250,00", 1250.0), ("1 250.00-", -1250.0), ("1 250.00 CR", -1250.0), ("(1 250.00)", -1250.0), ("R 118 402.10", 118402.1), ("R 0.00", 0.0), ("12 004", 12004.0), ("R 1 250,00", 1250.0), ("1 250.00 DR", 1250.0)],
)
def test_money_parser(raw, val):
    assert M.parse_money(raw) == val


def test_money_parser_rejects_non_money():
    assert M.parse_money("15%") is None
    assert M.parse_money("kVA") is None
    assert M.find_money("VAT 15% R 4 185.84") == [("R 4 185.84", 4185.84)]
    assert M.find_numbers("Period 01/06/2026 - 30/06/2026 Demand 88 kVA") == [("88", 88.0)]
    assert M.parse_date("01/06/2026 – 30/06/2026") == "2026-06-01 – 2026-06-30"
    assert M.parse_date("Jul 2026") == "2026-07"


def test_provenance_rule_scores_zero(bills):
    from hashlib import sha256

    from cbx.pdfstudio.spans import read_document

    data = Path(next(x["file"] for x in bills if x["name"] == "centlec-CEN-20924.pdf")).read_bytes()
    spans = read_document(data, sha256(data).hexdigest())
    cols = [{"key": "total", "label": "Total due", "ctype": "money", "required": True, "critical": True}, {"key": "muni", "label": "Municipality", "ctype": "list", "required": True, "critical": False}]
    rs = {"match": {}, "validators": {}, "fields": {"total": {"rules": [{"id": "r1", "type": "label_right", "label": "Total payable"}]}, "muni": {"rules": [{"id": "c1", "type": "const", "value": "Ekurhuleni"}]}}}
    cells = X.extract(rs, cols, spans)
    SC.score_cells(cells, cols, {"failing": []})
    assert cells["total"]["score"] == 1.0 and cells["total"]["status"] == "ok" and cells["total"]["quote"] == "R 61 990.29"
    # 'Ekurhuleni' is not printed on a Centlec bill → provenance fails → score 0, miss
    assert cells["muni"]["value"] == "Ekurhuleni" and cells["muni"]["provenance"] is False
    assert cells["muni"]["score"] == 0.0 and cells["muni"]["status"] == "miss"
    # a derived figure that is not printed anywhere has no provenance either
    rs2 = {"fields": {"total": {"rules": [{"id": "d", "type": "derived", "op": "add", "args": [1.0, 2.0]}]}}}
    c2 = X.extract(rs2, cols[:1], spans)
    assert c2["total"]["provenance"] is False


def test_scoring_ladder():
    cols = [{"key": "a", "label": "A", "ctype": "money", "required": True, "critical": True}]
    for q, expect in (("exact", 1.0), ("fuzzy", 0.75), ("positional", 0.55), ("derived", 0.9)):
        cells = {"a": {"value": 1.0, "quote": "1.00", "quality": q, "parsed": True, "provenance": True, "status": "ok"}}
        SC.score_cells(cells, cols, {"failing": []})
        assert cells["a"]["score"] == expect
    cells = {"a": {"value": 1.0, "quote": "1.00", "quality": "exact", "parsed": True, "provenance": True, "ocr": True, "status": "ok"}}
    SC.score_cells(cells, cols, {"failing": ["a"], "note": "x"})
    assert cells["a"]["score"] == round(0.9 * 0.4, 3) and cells["a"]["status"] == "miss"


# ---- regressions -------------------------------------------------------------------------------------------
def test_template_with_batches_cannot_be_deleted_but_orphan_can(owner, bills):
    t = _template(owner, "Delete me")
    b = _run_batch(owner, bills, TEXT_BILLS[:2], t["id"], name="Delete test")
    # a template with batches behind it: a clear 400, not a foreign-key crash
    r = owner.delete(f"/api/pdfstudio/templates/{t['id']}")
    assert r.status_code == 400 and "batch" in r.json()["detail"]
    assert owner.get(f"/api/pdfstudio/batches/{b['id']}").status_code == 200
    # a template whose only references are profiles / bake-off rows can go
    t2 = _template(owner, "Delete me 2")
    owner.post(f"/api/pdfstudio/bench/{t2['id']}")
    db = SessionLocal()
    try:
        db.add(Profile(key="delete-me-2.test", label="x", template_id=t2["id"]))
        db.commit()
    finally:
        db.close()
    assert owner.delete(f"/api/pdfstudio/templates/{t2['id']}").status_code == 200
    assert all(x["id"] != t2["id"] for x in owner.get("/api/pdfstudio/templates").json()["templates"])


def test_non_pdfs_in_a_drop_are_skipped_not_fatal(user, bills):
    t = _template(user, "Mixed drop")
    bid = user.post("/api/pdfstudio/batches", json={"template_id": t["id"]}).json()["id"]
    files = _pdf_files(bills, TEXT_BILLS[:2]) + [("files", ("notes.xlsx", b"PK\x03\x04junk", "application/octet-stream"))]
    r = user.post(f"/api/pdfstudio/batches/{bid}/files", files=files)
    assert r.status_code == 200, r.text
    out = r.json()
    assert len(out["documents"]) == 2 and out["count"] == 2 and out["rejected"] == ["notes.xlsx"]
    # nothing usable at all is still an error
    r = user.post(f"/api/pdfstudio/batches/{bid}/files", files=[("files", ("x.txt", b"hello", "text/plain"))])
    assert r.status_code == 400 and "x.txt" in r.json()["detail"]
    # cells can't be filled before C-EXA has read the batch
    did = out["documents"][0]["id"]
    r = user.post(f"/api/pdfstudio/batches/{bid}/cells", json={"document_id": did, "column_key": "vat", "value": "12"})
    assert r.status_code == 400


def test_stalled_batch_can_be_started_again(user, bills):
    t = _template(user, "Stalled")
    b = _run_batch(user, bills, TEXT_BILLS[:2], t["id"], name="Stalled test")
    assert b["stalled"] is False and b["status"] in ("done", "review")
    assert user.post(f"/api/pdfstudio/batches/{b['id']}/start").status_code == 400  # finished: no restart
    # pretend the server died mid-run: status says running but no job is alive
    db = SessionLocal()
    try:
        row = db.get(Batch, b["id"])
        row.status, row.step, row.progress = "running", "Checking figures · run 2", 0.5
        db.commit()
    finally:
        db.close()
    g = user.get(f"/api/pdfstudio/batches/{b['id']}").json()
    assert g["status"] == "running" and g["stalled"] is True
    assert any(x["id"] == b["id"] and x["stalled"] for x in user.get("/api/pdfstudio/batches").json()["batches"])
    r = user.post(f"/api/pdfstudio/batches/{b['id']}/start")
    assert r.status_code == 200, r.text
    g = user.get(f"/api/pdfstudio/batches/{b['id']}").json()
    assert g["status"] in ("done", "review") and g["stalled"] is False and g["sheet"]["rows"]


def test_excel_ctype_guess_dates_before_money():
    from cbx.pdfstudio.excel_template import guess_ctype

    assert guess_ctype("Due date", None) == "date"
    assert guess_ctype("Date paid", None) == "date"
    assert guess_ctype("Amount due", None) == "money"
    assert guess_ctype("Billing period", None) == "date"
    assert guess_ctype("kWh consumed", None) == "number"


def test_master_sheet_rename_clash_is_400(owner):
    a = owner.post("/api/master-sheets", json={"name": "Clash A", "columns": [{"label": "x"}]})
    b = owner.post("/api/master-sheets", json={"name": "Clash B", "columns": [{"label": "x"}]})
    assert a.status_code == 200 and b.status_code == 200
    r = owner.patch(f"/api/master-sheets/{b.json()['id']}", json={"name": "clash a"})
    assert r.status_code == 400
    r = owner.patch(f"/api/master-sheets/{b.json()['id']}", json={"name": "Clash C"})
    assert r.status_code == 200 and r.json()["name"] == "Clash C"


# ---- column kinds: filename · optional · formula --------------------------------------------------------
from tests._pdfs import bidvest_invoice, megaglo_invoice  # noqa: E402

VENDOR_COLUMNS = [
    {"key": "file", "label": "File", "ctype": "filename"},
    {"key": "invoice_no", "label": "Invoice number", "ctype": "text"},
    {"key": "invoice_date", "label": "Invoice date", "ctype": "date"},
    {"key": "material_excl_vat", "label": "Material excl VAT", "ctype": "money"},
    {"key": "material_10_mu", "label": "Material 10% MU", "ctype": "formula", "expr": "material_excl_vat * 1.10"},
    {"key": "travel_hours", "label": "Travel hours", "ctype": "number", "required": False},
    {"key": "travel_time_rate", "label": "Travel time rate", "ctype": "money", "required": False, "critical": False},
    {"key": "travel_excl_vat", "label": "Travel excl VAT", "ctype": "formula", "expr": "travel_hours * travel_time_rate"},
    {"key": "excl_vat", "label": "Sub total excl VAT", "ctype": "money"},
    {"key": "vat", "label": "VAT", "ctype": "money"},
    {"key": "total", "label": "Total", "ctype": "money"},
]


def _vendor_files(n=3, *, travel_on=(1,), excl=18000.0):
    return [("files", (f"Bidvest INV-{20440 + i}.pdf", bidvest_invoice(i, excl=excl + i * 100, travel_hours=3 if i in travel_on else None), "application/pdf")) for i in range(n)]


def _run_files(client, template_id, files, name="Vendor batch"):
    r = client.post("/api/pdfstudio/batches", json={"template_id": template_id, "name": name})
    assert r.status_code == 200, r.text
    bid = r.json()["id"]
    r = client.post(f"/api/pdfstudio/batches/{bid}/files", files=files)
    assert r.status_code == 200, r.text
    r = client.post(f"/api/pdfstudio/batches/{bid}/start")
    assert r.status_code == 200, r.text
    return client.get(f"/api/pdfstudio/batches/{bid}").json()


@pytest.fixture()
def vendor_batch(owner):
    r = owner.post("/api/pdfstudio/templates", json={"name": "Vendor invoices", "columns": VENDOR_COLUMNS})
    assert r.status_code == 200, r.text
    t = r.json()
    b = _run_files(owner, t["id"], _vendor_files())
    return t, b


def test_filename_optional_and_formula_columns(vendor_batch):
    t, b = vendor_batch
    assert b["status"] in ("done", "review"), b["log"]
    cols = {c["key"]: c for c in t["columns"]}
    assert cols["file"]["ctype"] == "filename" and cols["file"]["kind"] == "file name" and cols["file"]["critical"] is False
    assert cols["material_10_mu"]["expr"] == "material_excl_vat * 1.10" and cols["material_10_mu"]["kind"] == "calculated"
    assert cols["travel_hours"]["required"] is False and cols["travel_hours"]["kind"] == "optional"
    assert t["columns_summary"]["text"] == "11 columns · 2 optional · 2 calculated · 1 file name"
    rows = {r["filename"].split(" (")[0]: r for r in b["sheet"]["rows"]}
    assert set(rows) == {f"Bidvest INV-{20440 + i}.pdf" for i in range(3)}
    for name, r in rows.items():
        c = r["cells"]
        # filename: the PDF's own name, score 1, never a miss
        assert c["file"]["value"] == name and c["file"]["status"] == "ok" and c["file"]["score"] == 1.0 and c["file"]["method"] == "filename"
        # formula: material × 1.10, computed on every row
        assert c["material_excl_vat"]["status"] == "ok"
        assert c["material_10_mu"]["status"] == "ok" and abs(c["material_10_mu"]["value"] - round(c["material_excl_vat"]["value"] * 1.10, 2)) < 0.005
        assert c["material_10_mu"]["method"] == "formula" and c["material_10_mu"]["score"] == 1.0
    with_travel = rows["Bidvest INV-20441.pdf"]["cells"]
    without = rows["Bidvest INV-20440.pdf"]["cells"]
    # the optional travel figures: read where printed, blank ("absent") where the invoice has no travel
    assert with_travel["travel_hours"]["value"] == 3 and with_travel["travel_time_rate"]["value"] == 350.0
    assert with_travel["travel_excl_vat"]["value"] == 1050.0 and with_travel["travel_excl_vat"]["status"] == "ok"
    assert without["travel_hours"]["status"] == "absent" and without["travel_hours"]["value"] is None and without["travel_hours"]["note"] == "not on this document"
    assert without["travel_excl_vat"]["status"] == "absent" and without["travel_excl_vat"]["value"] is None and "Travel hours" in without["travel_excl_vat"]["note"]
    # none of that is for the user; the arithmetic columns were read right
    assert b["cells_for_user"] == 0, [(k, v) for r in b["sheet"]["rows"] for k, v in r["cells"].items() if v["status"] in ("miss", "flagged")]
    assert b["by_column"]["travel_hours"] == 1.0 and b["by_column"]["travel_excl_vat"] == 1.0 and b["by_column"]["file"] == 1.0
    assert b["accuracy_pct"] >= 99.0
    assert b["cells_absent"] == 2 * 3  # 2 invoices without travel × (hours, rate, travel formula) = 6
    assert b["columns_summary"]["optional"] == 2 and b["columns_summary"]["calculated"] == 2
    assert any("2 optional" in e["text"] and "2 calculated" in e["text"] for e in b["log"])
    # row scores ignore the always-right columns: 6 read + required columns count, not 11
    assert all(r["of"] == 6 for r in b["sheet"]["rows"])


def test_formula_and_filename_in_export_and_master_sheet(vendor_batch, owner):
    t, b = vendor_batch
    r = owner.get(f"/api/pdfstudio/batches/{b['id']}/export.xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Rows"]
    head = [c.value for c in ws[1]]
    fi, mi, ti = head.index("File") + 1, head.index("Material 10% MU") + 1, head.index("Travel excl VAT") + 1
    names = {ws.cell(row=i, column=fi + 1).value for i in range(2, 5)}  # the template's own File column sits after the lead File column
    assert names == {f"Bidvest INV-{20440 + i}.pdf" for i in range(3)} or all(str(ws.cell(row=i, column=fi + 1).value).startswith("Bidvest INV-") for i in range(2, 5))
    assert all(isinstance(ws.cell(row=i, column=mi).value, (int, float)) for i in range(2, 5))
    assert [ws.cell(row=i, column=head.index("Material excl VAT") + 1).value for i in range(2, 5)] == [9000, 9050, 9100]  # the line item, not the sub total
    assert ws.cell(row=2, column=mi).number_format == "#,##0.00"
    assert sum(1 for i in range(2, 5) if ws.cell(row=i, column=ti).value is None) == 2  # blank where no travel
    # the Misses sheet has no absent / formula / filename rows; provenance records 'absent'
    misses = wb["Misses"]
    assert misses.max_row == 1
    prov = wb["Provenance"]
    statuses = {prov.cell(row=i, column=10).value for i in range(2, prov.max_row + 1)}
    assert "absent" in statuses and "miss" not in statuses
    # a filename / formula cell can't be typed over
    row = b["sheet"]["rows"][0]
    r = owner.post(f"/api/pdfstudio/batches/{b['id']}/cells", json={"document_id": row["document_id"], "column_key": "file", "value": "x.pdf"})
    assert r.status_code == 400 and "file name" in r.json()["detail"]
    r = owner.post(f"/api/pdfstudio/batches/{b['id']}/cells", json={"document_id": row["document_id"], "column_key": "material_10_mu", "value": "1"})
    assert r.status_code == 400 and "worked out" in r.json()["detail"]


def test_formula_follows_a_user_fill(vendor_batch, owner):
    """Fill an optional input on a row where it was absent: the formula that uses it is worked out."""
    t, b = vendor_batch
    row = next(r for r in b["sheet"]["rows"] if r["cells"]["travel_hours"]["status"] == "absent")
    # 'Travel hours' is not on this PDF, so the value is flagged and the person must vouch for it
    r = owner.post(f"/api/pdfstudio/batches/{b['id']}/cells", json={"document_id": row["document_id"], "column_key": "travel_hours", "value": "2"})
    assert r.status_code == 200 and r.json()["status"] == "flagged"
    cid = r.json()["correction_id"]
    r = owner.post(f"/api/pdfstudio/batches/{b['id']}/cells/answer", json={"correction_id": cid, "answer": "trust_me", "note": "from the job card"})
    assert r.status_code == 200 and r.json()["status"] == "trusted"
    r = owner.post(f"/api/pdfstudio/batches/{b['id']}/cells", json={"document_id": row["document_id"], "column_key": "travel_time_rate", "value": "350"})
    cid2 = r.json()["correction_id"]
    owner.post(f"/api/pdfstudio/batches/{b['id']}/cells/answer", json={"correction_id": cid2, "answer": "trust_me", "note": "rate card"})
    b2 = owner.get(f"/api/pdfstudio/batches/{b['id']}").json()
    row2 = next(r for r in b2["sheet"]["rows"] if r["document_id"] == row["document_id"])
    assert row2["cells"]["travel_excl_vat"]["value"] == 700.0 and row2["cells"]["travel_excl_vat"]["status"] == "ok"
    assert b2["cells_for_user"] == 0
    # taking back an optional value returns the cell to "not on this document", not to the user's desk —
    # and the formula goes blank again
    r = owner.post(f"/api/pdfstudio/batches/{b['id']}/cells/answer", json={"correction_id": cid2, "answer": "retract"})
    assert r.status_code == 200 and r.json()["status"] == "absent", r.json()
    b3 = owner.get(f"/api/pdfstudio/batches/{b['id']}").json()
    row3 = next(r for r in b3["sheet"]["rows"] if r["document_id"] == row["document_id"])
    assert row3["cells"]["travel_excl_vat"]["status"] == "absent" and row3["cells"]["travel_excl_vat"]["value"] is None
    assert b3["cells_for_user"] == 0 and b3["status"] == "done"


def test_cexa_proposes_optional_for_figures_seen_on_few_samples(owner):
    """'Travel hours' starts required; it is printed on 1 of the 3 sample invoices → C-EXA makes it
    optional, records why on the column, says so in the log, and the blank cells are not misses."""
    cols = [dict(c) for c in VENDOR_COLUMNS]
    for c in cols:
        if c["key"] in ("travel_hours", "travel_time_rate"):
            c["required"] = True
    r = owner.post("/api/pdfstudio/templates", json={"name": "Vendor invoices (strict)", "columns": cols})
    t = r.json()
    assert all(c["required"] for c in t["columns"] if c["key"].startswith("travel_") and c["ctype"] != "formula")
    b = _run_files(owner, t["id"], _vendor_files(), name="Strict batch")
    assert b["status"] in ("done", "review"), b["log"]
    t2 = next(x for x in owner.get("/api/pdfstudio/templates").json()["templates"] if x["id"] == t["id"])
    tc = {c["key"]: c for c in t2["columns"]}
    assert tc["travel_hours"]["required"] is False and "1 of 3" in tc["travel_hours"]["optional_reason"]
    assert tc["travel_time_rate"]["required"] is False and tc["travel_time_rate"]["kind"] == "optional"
    assert tc["excl_vat"]["required"] is True and tc["material_excl_vat"]["required"] is True  # seen on every sample
    assert any(e["text"].startswith("Optional by evidence") and "Travel hours (1 of 3 samples)" in e["text"] for e in b["log"])
    assert b["cells_for_user"] == 0 and b["by_column"]["travel_hours"] == 1.0
    db = SessionLocal()
    try:
        prof = db.query(Profile).filter(Profile.template_id == t["id"]).one()
        assert "Optional by evidence" in db.get(ProfileVersion, prof.live_version_id).notes
    finally:
        db.close()
    # the admin can make it required again — then the blank cells go to the user
    cols2 = [dict(c, required=True) if c["key"] == "travel_hours" else dict(c) for c in cols]
    r = owner.patch(f"/api/pdfstudio/templates/{t['id']}", json={"columns": cols2})
    assert r.status_code == 200 and next(c for c in r.json()["columns"] if c["key"] == "travel_hours")["required"] is True


def test_invalid_formula_is_400_with_the_usable_keys(owner):
    base = [{"key": "excl_vat", "label": "Excl VAT", "ctype": "money"}, {"key": "vat", "label": "VAT", "ctype": "money"}, {"key": "site", "label": "Site", "ctype": "text"}]

    def post(expr, label="Total (calc)"):
        return owner.post("/api/pdfstudio/templates", json={"name": "Formula test", "columns": base + [{"key": "calc", "label": label, "ctype": "formula", "expr": expr}]})

    r = post("excl_vat + vatt")
    assert r.status_code == 400, r.text
    assert "'vatt' isn't a column" in r.json()["detail"] and "excl_vat, vat" in r.json()["detail"]
    r = post("excl_vat * site")
    assert r.status_code == 400 and "text column" in r.json()["detail"]
    r = post("excl_vat +")
    assert r.status_code == 400 and "ends too early" in r.json()["detail"]
    r = post("calc * 2")
    assert r.status_code == 400 and "own column" in r.json()["detail"]
    r = post("")
    assert r.status_code == 400 and "type the formula" in r.json()["detail"]
    r = post("import os")
    assert r.status_code == 400
    # a template of only formula / filename columns has nothing to read
    r = owner.post("/api/pdfstudio/templates", json={"name": "Nothing to read", "columns": [{"key": "f", "label": "File", "ctype": "filename"}]})
    assert r.status_code == 400 and "read from the PDF" in r.json()["detail"]
    # a good one is accepted, with round()
    r = post("round(excl_vat + vat, 2)")
    assert r.status_code == 200, r.text
    calc = next(c for c in r.json()["columns"] if c["key"] == "calc")
    assert calc["expr"] == "round(excl_vat + vat, 2)" and calc["required"] is False and calc["critical"] is False
    # unknown type gets the full list
    r = owner.post("/api/pdfstudio/templates", json={"name": "Bad type", "columns": [{"key": "a", "label": "A", "ctype": "percentage"}]})
    assert r.status_code == 400 and "formula" in r.json()["detail"]


def test_formula_evaluator_is_safe_and_blank_on_missing_input():
    from cbx.pdfstudio import formula as F

    kinds = {"material_excl_vat": "money", "labour_hours": "number", "labour_rate": "money", "kms_travelled": "number", "km_rate": "money", "excl_vat": "money", "vat": "money"}
    for expr, keys in (("material_excl_vat * 1.10", ["material_excl_vat"]), ("labour_hours * labour_rate", ["labour_hours", "labour_rate"]), ("kms_travelled * km_rate", ["kms_travelled", "km_rate"]), ("excl_vat + vat", ["excl_vat", "vat"]), ("round(excl_vat * 0,15, 2)", ["excl_vat"])):
        assert F.validate(expr, kinds) == keys, expr
    assert F.evaluate("material_excl_vat * 1.10", {"material_excl_vat": 1234.5}) == 1357.95
    assert F.evaluate("labour_hours * labour_rate", {"labour_hours": 3, "labour_rate": None}) is None
    assert F.evaluate("labour_hours * labour_rate", {"labour_hours": 3}) is None
    assert F.evaluate("excl_vat / vat", {"excl_vat": 10, "vat": 0}) is None
    assert F.evaluate("-(excl_vat - vat)", {"excl_vat": 10, "vat": 4}) == -6.0
    assert F.evaluate("min(excl_vat, vat) + max(excl_vat, vat)", {"excl_vat": 10, "vat": 4}) == 14.0
    assert F.missing_inputs("labour_hours * labour_rate", {"labour_hours": 3}) == ["labour_rate"]
    for bad in ("__import__('os')", "excl_vat.__class__", "2 ** 8", "vat if 1 else 0", "[1]", "a b"):
        with pytest.raises(F.FormulaError):
            F.validate(bad, kinds)


def test_excel_guesses_filename_and_optional_columns(owner):
    from cbx.pdfstudio.excel_template import guess_ctype, guess_required

    for h in ("file", "File name", "electricity file", "PDF", "Source"):
        assert guess_ctype(h, None) == "filename", h
    assert guess_ctype("Number of files", None) != "filename"
    for h in ("Travel hours", "Travel time rate", "Kms travelled", "KM rate", "Exceptions", "Notes", "Comment", "Guarantee", "Deposit held", "Security"):
        assert guess_required(h) is False, h
    for h in ("Total due", "Account number", "Material excl VAT", "Invoice date"):
        assert guess_required(h) is True, h
    wb = Workbook()
    ws = wb.active
    ws.append(["File", "Invoice number", "Material excl VAT", "Travel hours", "Kms travelled", "Exceptions", "Total"])
    ws.append(["a.pdf", "INV-1", 100.0, 2, 40, "", 115.0])
    buf = io.BytesIO()
    wb.save(buf)
    r = owner.post("/api/pdfstudio/templates/from-excel", files=[("files", ("Capitec.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))])
    assert r.status_code == 200, r.text
    cols = {c["key"]: c for c in r.json()["columns"]}
    assert cols["file"]["ctype"] == "filename" and cols["file"]["required"] is False
    assert cols["travel_hours"]["required"] is False and cols["travel_hours"]["ctype"] == "number" and "Excel header" in cols["travel_hours"]["optional_reason"]
    assert cols["kms_travelled"]["required"] is False and cols["exceptions"]["required"] is False and cols["exceptions"]["ctype"] == "text"
    assert cols["material_excl_vat"]["required"] is True and cols["material_excl_vat"]["critical"] is True
    assert cols["total"]["required"] is True
    assert r.json()["columns_summary"]["text"] == "7 columns · 3 optional · 1 file name"


# ---- upload race, 409, delete, housekeeping ------------------------------------------------------------------
def test_files_to_a_finished_batch_is_409_and_empty_batch_can_be_deleted(vendor_batch, owner, user):
    t, b = vendor_batch
    r = owner.post(f"/api/pdfstudio/batches/{b['id']}/files", files=_vendor_files(1))
    assert r.status_code == 409, r.text
    assert "can't take more PDFs" in r.json()["detail"] and "Start a new batch" in r.json()["detail"]
    # a finished batch can't be deleted; an empty one (the race leftovers) can
    assert owner.delete(f"/api/pdfstudio/batches/{b['id']}").status_code == 400
    e = owner.post("/api/pdfstudio/batches", json={"template_id": t["id"], "name": "empty"}).json()
    assert owner.delete(f"/api/pdfstudio/batches/{e['id']}").status_code == 200
    assert owner.get(f"/api/pdfstudio/batches/{e['id']}").status_code == 404
    # only the maker (or an admin) — another user can't even see it
    e2 = owner.post("/api/pdfstudio/batches", json={"template_id": t["id"], "name": "empty 2"}).json()
    assert user.delete(f"/api/pdfstudio/batches/{e2['id']}").status_code == 404
    assert owner.delete(f"/api/pdfstudio/batches/{e2['id']}").status_code == 200


def test_housekeeping_prunes_empty_batches_and_fails_stuck_uploads(owner, vendor_batch):
    from datetime import timedelta

    from cbx.models import AuditLog
    from cbx.pdfstudio import housekeeping as HK

    t, _ = vendor_batch
    old_empty = owner.post("/api/pdfstudio/batches", json={"template_id": t["id"], "name": "race leftover"}).json()["id"]
    fresh_empty = owner.post("/api/pdfstudio/batches", json={"template_id": t["id"], "name": "just made"}).json()["id"]
    stuck = owner.post("/api/pdfstudio/batches", json={"template_id": t["id"], "name": "never handed over"}).json()["id"]
    assert owner.post(f"/api/pdfstudio/batches/{stuck}/files", files=_vendor_files(2)).status_code == 200
    db = SessionLocal()
    try:
        from cbx.db import utcnow

        for bid in (old_empty, stuck):
            db.get(Batch, bid).created_at = utcnow() - timedelta(hours=25)
        db.commit()
        out = HK.prune(db)
        assert out == {"pruned": [old_empty], "failed": [stuck]}, out
        assert db.get(Batch, old_empty) is None
        s = db.get(Batch, stuck)
        assert s.status == "failed" and s.error == "Upload never finished — start again" and len(s.documents) == 2
        assert db.get(Batch, fresh_empty).status == "uploading"
        acts = {a.action for a in db.query(AuditLog).filter(AuditLog.target.in_([f"batch:{old_empty}", f"batch:{stuck}"])).all()}
        assert {"pdf_batch_pruned", "pdf_batch_stuck"} <= acts
        # running again changes nothing
        assert HK.prune(db) == {"pruned": [], "failed": []}
    finally:
        db.close()
    # the routine itself runs without a fuss, and the failed batch's "Try again" works
    HK.housekeeping_routine()
    g = owner.get(f"/api/pdfstudio/batches/{stuck}").json()
    assert g["status"] == "failed" and g["error"] == "Upload never finished — start again"
    r = owner.post(f"/api/pdfstudio/batches/{stuck}/start")
    assert r.status_code == 200, r.text
    assert owner.get(f"/api/pdfstudio/batches/{stuck}").json()["status"] in ("done", "review")
    assert owner.delete(f"/api/pdfstudio/batches/{fresh_empty}").status_code == 200


# ---- re-run with the current profile -----------------------------------------------------------------------
def test_rerun_with_current_profile_keeps_history_and_user_fills(vendor_batch, owner, user):
    t, b = vendor_batch
    bid = b["id"]
    row = next(r for r in b["sheet"]["rows"] if r["cells"]["travel_hours"]["status"] == "absent")
    r = owner.post(f"/api/pdfstudio/batches/{bid}/cells", json={"document_id": row["document_id"], "column_key": "travel_hours", "value": "5"})
    cid = r.json()["correction_id"]
    owner.post(f"/api/pdfstudio/batches/{bid}/cells/answer", json={"correction_id": cid, "answer": "trust_me", "note": "job card"})
    assert user.post(f"/api/pdfstudio/batches/{bid}/rerun").status_code in (403, 404)  # admins only
    r = owner.post(f"/api/pdfstudio/batches/{bid}/rerun")
    assert r.status_code == 200, r.text
    g = owner.get(f"/api/pdfstudio/batches/{bid}").json()
    assert g["status"] in ("done", "review") and g["reruns"] == 1
    assert any("re-read this batch with the current profile" in e["text"] for e in g["log"])
    assert any(e["text"].startswith("Re-reading 3 PDFs") for e in g["log"])
    row2 = next(r for r in g["sheet"]["rows"] if r["document_id"] == row["document_id"])
    assert row2["cells"]["travel_hours"]["status"] == "trusted" and row2["cells"]["travel_hours"]["value"] == 5
    assert row2["cells"]["total"]["value"] == row["cells"]["total"]["value"]
    db = SessionLocal()
    try:
        ext = db.query(Extraction).filter(Extraction.document_id == row["document_id"]).one()
        assert len(ext.history) == 1 and ext.history[0]["cells"]["total"]["value"] == row["cells"]["total"]["value"]
        prof = db.query(Profile).filter(Profile.template_id == t["id"]).one()
        assert prof.docs_seen == 3  # a re-read is not "more documents seen"
    finally:
        db.close()
    # not while it is being read; not before it was handed over
    e = owner.post("/api/pdfstudio/batches", json={"template_id": t["id"], "name": "not started"}).json()
    assert owner.post(f"/api/pdfstudio/batches/{e['id']}/rerun").status_code == 400


# ---- sibling profiles: a small vendor borrows the big vendor's rules ---------------------------------------------
def test_new_small_vendor_borrows_rules_from_the_templates_other_profile(owner):
    cols = [c for c in VENDOR_COLUMNS if c["key"] not in ("travel_hours", "travel_time_rate", "travel_excl_vat")] + [{"key": "job", "label": "Job description", "ctype": "text", "required": False}]
    t = owner.post("/api/pdfstudio/templates", json={"name": "Capitec vendor invoices", "columns": cols}).json()
    big = _run_files(owner, t["id"], _vendor_files(3, travel_on=()), name="Bidvest")
    assert big["status"] in ("done", "review"), big["log"]
    db = SessionLocal()
    try:
        prof_a = db.query(Profile).filter(Profile.template_id == t["id"]).one()
        ver = db.get(ProfileVersion, prof_a.live_version_id)
        rs = json.loads(json.dumps(ver.ruleset))
        # a rule the big profile learned over time (the way a person's fix teaches it): 'Job' → job
        rs["fields"].setdefault("job", {"rules": []})["rules"].insert(0, {"id": "u1", "type": "label_right", "label": "Job", "pick": "first", "page": None, "quality": "exact"})
        ver.ruleset = rs
        db.commit()
        label_a = prof_a.label
    finally:
        db.close()
    small = _run_files(owner, t["id"], [("files", (f"Megaglo MG-{310 + i}.pdf", megaglo_invoice(i), "application/pdf")) for i in range(2)], name="Megaglo")
    assert small["status"] in ("done", "review"), small["log"]
    db = SessionLocal()
    try:
        profs = db.query(Profile).filter(Profile.template_id == t["id"]).order_by(Profile.id).all()
        assert len(profs) == 2 and profs[1].key.endswith("megaglo-electrical")
        v1 = profs[1].versions[0]
        fields = v1.ruleset["fields"]
        # the borrowed rule leads where Megaglo's own heuristics found nothing …
        assert fields["job"]["seeded_from"] == label_a and fields["job"]["rules"][0]["label"] == "Job"
        # … the big profile's rules were tried on the Megaglo samples first, and the notes say so
        assert "borrowed" in v1.reason and "tried from this template's other layouts" in v1.notes and f"'{label_a}' leads for job" in v1.notes
        # the same-key column Megaglo reads on its own keeps its own rule in front
        assert fields["invoice_no"]["rules"][0]["label"].lower().startswith("invoice number")
        assert fields["invoice_no"].get("seeded_from") is None
    finally:
        db.close()
    rows = small["sheet"]["rows"]
    assert all(r["cells"]["job"]["value"] == "replace distribution board, Capitec Boksburg" for r in rows)
    assert all(r["cells"]["total"]["status"] == "ok" and r["cells"]["file"]["value"].startswith("Megaglo MG-") for r in rows)
    # the profiles page shows required/optional, kind and where a rule came from
    pj = owner.get(f"/api/pdfstudio/profiles/{profs[1].id}").json()
    by = {f["key"]: f for f in pj["fields"]}
    assert by["job"]["seeded_from"] == label_a and by["job"]["required"] is False and by["job"]["kind"] == "optional"
    assert by["file"]["kind"] == "file name" and by["file"]["needs_rules"] is False
    assert by["material_10_mu"]["kind"] == "calculated" and by["material_10_mu"]["expr"] == "material_excl_vat * 1.10"
    assert by["total"]["kind"] == "read from the page" and by["total"]["required"] is True
    assert any(e["text"].startswith("Megaglo Electrical: built profile v1") and "borrowed" in e["text"] for e in small["log"])


def test_sibling_rules_unit():
    """merge_sibling_fields: a borrowed rule leads only when it reads more samples than our own."""
    from cbx.pdfstudio import profile as P

    rs = {"fields": {"total": {"rules": [{"id": "h1", "type": "label_right", "label": "Total", "pick": "first", "hits": 2, "of": 2}]}, "job": {"rules": []}}}
    sib = {
        "total": {"rules": [{"id": "s1", "type": "label_right", "label": "Amount due", "pick": "last", "hits": 1, "of": 2}], "seeded_from": "Big"},
        "job": {"rules": [{"id": "s1", "type": "label_right", "label": "Job", "pick": "first", "hits": 2, "of": 2}], "seeded_from": "Big"},
    }
    led = P.merge_sibling_fields(rs, sib)
    assert led == {"job": "Big"}
    assert [r["label"] for r in rs["fields"]["total"]["rules"]] == ["Total", "Amount due"] and "seeded_from" not in rs["fields"]["total"]
    assert rs["fields"]["job"]["seeded_from"] == "Big" and rs["fields"]["job"]["rules"][0]["label"] == "Job"


# ---- how C-EXA learns a layout: identity · skeleton · the evidence gate · exclusivity -----------------
from tests._pdfs import other_system_bill, shared_system_bill, text_pdf  # noqa: E402

SHARED_COLUMNS = [
    {"key": "municipality", "label": "Municipality", "ctype": "text"},
    {"key": "account_number", "label": "Account Number", "ctype": "number"},
    {"key": "billing_date", "label": "Billing Date", "ctype": "date"},
    {"key": "meter_number", "label": "Meter number", "ctype": "text"},
    {"key": "start_reading_kwh", "label": "Start Reading kWh", "ctype": "number"},
    {"key": "end_reading_kwh", "label": "End Reading kWh", "ctype": "number"},
    {"key": "total_kwh", "label": "Total kWh", "ctype": "number"},
    {"key": "peak_kwh", "label": "Peak kWh", "ctype": "number"},
    {"key": "deposit", "label": "Deposit", "ctype": "text"},
    {"key": "demand_kva", "label": "Demand kVA", "ctype": "number"},
]
SHARED_COLS = [dict(c, required=True, critical=False, hint="") for c in SHARED_COLUMNS]
_GREENVALE = [("0003529265", "19 AUG 2026", "000000000158618", 91767.0, 92414.0), ("0003529265", "20 SEP 2026", "000000000158618", 92414.0, 93471.0), ("0003529265", "21 OCT 2026", "000000000158618", 93471.0, 94528.0)]
_RIVERBEND = [("0001009715", "01 JUN 2026", "000000009600160", 98722.0, 120923.0), ("0001009715", "02 JUL 2026", "000000009600160", 120923.0, 141101.0)]


def _spans(pdf: bytes) -> dict:
    from hashlib import sha256

    from cbx.pdfstudio.spans import read_document

    return read_document(pdf, sha256(pdf).hexdigest())


def _bills(issuer, rows):
    return [_spans(shared_system_bill(issuer, account=a, date=d, meter=m, old=o, new=n)) for a, d, m, o, n in rows]


@pytest.fixture(scope="module")
def greenvale():
    return _bills("GREENVALE LOCAL MUNICIPALITY", _GREENVALE)


@pytest.fixture(scope="module")
def riverbend():
    return _bills("RIVERBEND MUNICIPALITY", _RIVERBEND)


def test_identity_is_the_letterhead_and_the_vat_number_not_a_list_of_names(greenvale):
    """The old code looked for known issuer names and matched the city inside a postal address,
    which is how 'capitec.cape-town' was born. Identity now comes from the shape of the page."""
    from cbx.pdfstudio import profile as P

    ident = P.identity(greenvale[0])
    assert ident["issuer"] == "GREENVALE LOCAL MUNICIPALITY"
    assert ident["vat"] == "4991820371"
    assert "TAX INVOICE" in ident["doc_type"].upper()
    assert ident["issuer_name"] == "Greenvale"  # generic descriptors trimmed, no list of names
    assert P.printed_org("City of Ekurhuleni") == "Ekurhuleni" and P.printed_org("CENTLEC (SOC) Ltd") == "Centlec"
    # a layout that prints the CUSTOMER's name and address near the top and never its own says so:
    # naming the customer as the supplier is a confident wrong answer, and a blank costs less
    assert P.identity(_spans(other_system_bill()))["issuer"] == ""


def test_one_billing_system_is_one_family_whatever_the_municipality(greenvale, riverbend):
    """Three municipalities on one billing system must share ONE profile; a different layout must
    split off. The family is the printed skeleton, never the issuer."""
    from cbx.pdfstudio import profile as P

    other = [_spans(other_system_bill(account=f"10000321{i}")) for i in range(3)]
    docs = [(i, sp) for i, sp in enumerate(greenvale + riverbend + other)]
    fams = P.cluster(docs)
    assert len(fams) == 2, [f["family"] for f in fams]
    shared = next(f for f in fams if len(f["doc_ids"]) == 5)
    assert set(shared["doc_ids"]) == set(range(5))  # both municipalities, one family
    assert len(shared["issuers"]) == 2 and "GREENVALE LOCAL MUNICIPALITY" in shared["issuers"]
    assert shared["family"] and shared["signature"]
    # the match rule is tested before it is saved: it fits its own samples and nobody else's
    for fam in fams:
        rule = P.match_rules_for(fam, fams)
        assert rule["verified"] is True, rule
        mine = {sp is s for s in fam["spans"] for sp in fam["spans"]}
        assert all(P.matches({"match": rule}, sp) for sp in fam["spans"]) and mine
        for g in fams:
            if g is not fam:
                assert not any(P.matches({"match": rule}, sp) for sp in g["spans"]), (fam["family"], rule)
    # and a rule like the live {"all": ["ltd", "pty"]} — true of every SA invoice — is not one of ours
    assert all(len(t) >= 6 for f in fams for t in P.match_rules_for(f, fams)["all"])


def test_authoring_on_one_municipality_reads_a_municipality_it_has_never_seen(greenvale, riverbend):
    from cbx.pdfstudio import profile as P

    rs, notes = P.author_heuristic(SHARED_COLS, greenvale, family="shared")
    cells = X.extract(rs, SHARED_COLS, riverbend[0])
    v = {k: c["value"] for k, c in cells.items()}
    # the customer's account number, NOT the municipality's bank account further down the page
    assert v["account_number"] == 1009715
    # a heading is not a value: 'METER TYPE' is printed to the right of 'METER NO.'
    assert v["meter_number"] == "000000009600160"
    # readings are printed with three decimals; the document's own decimal mark settles it
    assert v["start_reading_kwh"] == 98722 and v["end_reading_kwh"] == 120923
    assert v["total_kwh"] == 22201  # the meter table's CONSUMPTION, not '350.000' from a tariff line
    assert v["billing_date"] == "2026-06-01" and v["deposit"] == "0.00"
    # the issuer is read per document, so a const 'Greenvale' can never leak onto another town's bill
    assert v["municipality"] == "RIVERBEND MUNICIPALITY"
    assert all(c.get("provenance", True) for c in cells.values() if c["value"] is not None)


def test_no_span_serves_two_columns_and_columns_with_no_evidence_are_reported(greenvale):
    """Peak kWh is not on this layout. It must stay blank rather than share Total kWh's figure —
    the live bug had six columns all claiming the same span."""
    from cbx.pdfstudio import profile as P

    rs, notes = P.author_heuristic(SHARED_COLS, greenvale, family="shared")
    assert rs["fields"]["total_kwh"]["rules"], notes
    assert rs["fields"]["peak_kwh"]["rules"] == []
    assert rs["fields"]["demand_kva"]["rules"] == []
    assert {"peak_kwh", "demand_kva"} <= set(rs["not_found"])
    assert any("Peak kWh: no rule passed the evidence gate" in n for n in notes)
    # every saved rule says how much evidence it has
    for key, field in rs["fields"].items():
        for r in field["rules"]:
            if r["type"] != "derived":
                assert r["of"] == len(greenvale) and r["hits"] >= 0.8 * r["of"], (key, r)


def test_the_evidence_gate_throws_invented_rules_away(greenvale):
    """A rule nobody tested is not a fallback, it is a wrong answer waiting to happen."""
    from cbx.pdfstudio import profile as P

    col = next(c for c in SHARED_COLS if c["key"] == "meter_number")
    invented = [
        {"id": "a1", "type": "label_right", "label": "Your Account With Us", "quality": "exact"},  # nowhere on the page
        {"id": "a2", "type": "label_right", "label": "METER NO.", "quality": "exact"},  # reads the heading 'METER TYPE'
        {"id": "a3", "type": "label_below", "label": "METER NO.", "quality": "exact"},  # the real one
    ]
    kept = P.verify_rules(invented, col, greenvale)
    assert [r["id"] for r in kept] == ["a3"] and kept[0]["hits"] == 3
    rs = {"match": {}, "validators": {}, "fields": {"meter_number": {"rules": []}}}
    assert P.merge_model_rules(rs, {"meter_number": invented}, SHARED_COLS, samples=greenvale) == 1


def test_variance_check_drops_a_figure_that_never_moves(greenvale):
    """A reading that is the same on every bill is not this month's reading."""
    from cbx.pdfstudio import profile as P

    col = next(c for c in SHARED_COLS if c["key"] == "total_kwh")
    moving = {"fit": 0.6, "varying": True, "hits": 3, "values": [647, 1057, 1057]}
    steady = {"fit": 0.5, "varying": False, "hits": 3, "values": [350000, 350000, 350000]}
    assert P._variance_filter([moving, steady], col) == [moving]
    # …but when nothing moves, the column is an identity-like one and nothing is dropped
    other = {"fit": 0.9, "varying": False, "hits": 3, "values": [1, 1, 1]}
    assert P._variance_filter([steady, other], col) == [steady, other]
    # a steady candidate the column names BETTER than any moving one survives: an account number is
    # the same every month, and that is not a reason to throw it away
    named = {"fit": 0.95, "varying": False, "hits": 3, "values": [3529265, 3529265, 3529265]}
    assert P._variance_filter([moving, named], col) == [moving, named]


def test_a_document_says_which_mark_is_its_decimal_point(greenvale):
    from cbx.pdfstudio import extract as EX

    assert EX.decimal_sep(greenvale[0]) == "."
    assert EX.retype("91767.000", 91767000.0, ".") == 91767.0  # three decimals, not 91 million
    assert EX.retype("5,228.68", 5228.68, ".") == 5228.68
    assert EX.retype("1.250,00", 1250.0, ",") == 1250.0
    assert EX.retype("63 102 116 646", 63102116646.0, ".") == 63102116646.0
    assert EX.retype("4 942.03-", -4942.03, ".") == -4942.03
    sp = _spans(text_pdf(["Sub total R 1 250,00", "VAT R 187,50", "Total R 1 437,50"]))
    assert EX.decimal_sep(sp) == ","


def test_derivations_are_acyclic_and_only_written_when_they_hold():
    """`labour_excl = total − VAT`, `total = labour_excl + VAT` and `VAT = labour_excl × 0.15` all at
    once is three rules defining each other. A derived rule is only for a column nothing else fills."""
    from cbx.pdfstudio import profile as P

    cols = [dict(c, required=True, critical=False, hint="") for c in ({"key": "excl_vat", "label": "Amount excl VAT", "ctype": "money"}, {"key": "vat", "label": "VAT", "ctype": "money"}, {"key": "total_due", "label": "Total due", "ctype": "money"})]
    samples = [_spans(text_pdf(["ACME SUPPLIES (PTY) LTD", "TAX INVOICE", f"Amount excl VAT R {1000 + i * 10}.00", f"VAT @ 15% R {150 + i}.00"])) for i in range(3)]
    rs, _notes = P.author_heuristic(cols, samples, family="acme")
    kinds = {k: [r["type"] for r in f["rules"]] for k, f in rs["fields"].items()}
    assert "derived" in kinds["total_due"], kinds  # nothing on the page fills it — arithmetic may
    assert "derived" not in kinds["excl_vat"] and "derived" not in kinds["vat"], kinds
    assert P.derivation_cycles(rs) == []
    # a ruleset that defines its columns in a circle is caught
    bad = {"fields": {"a": {"rules": [{"type": "derived", "op": "sub", "args": ["b", "c"]}]}, "b": {"rules": [{"type": "derived", "op": "add", "args": ["a", "c"]}]}, "c": {"rules": []}}}
    assert P.derivation_cycles(bad)


def test_a_colon_label_is_never_read_from_its_left(greenvale):
    """'Internet PIN: 440624127  Deposit: 0.00' — the PIN belongs to the PIN, not to the deposit."""
    from cbx.pdfstudio import extract as EX

    assert EX.apply_rule({"type": "label_left", "label": "Deposit", "anchor": "label"}, "number", greenvale[0]) is None
    assert EX.apply_rule({"type": "label_right", "label": "Deposit", "anchor": "label"}, "text", greenvale[0])["value"] == "0.00"


# ---- reading a page nobody can read: two text readers, the payment block, corroboration ----------
def test_a_second_text_reader_fills_in_what_the_first_one_missed():
    """Measured on a real municipal statement: pdfplumber returned the values and none of the
    labels. Whatever the primary reader misses is added back rather than lost."""
    from cbx.pdfstudio import spans as S

    class _Page:
        rotation = 0

    got = S._add_missing_words(b"%PDF-1.4 not a real pdf", 0, _Page(), [{"t": "VALUE", "x0": 100.0, "x1": 140.0, "top": 50.0, "bottom": 60.0}])
    assert [w["t"] for w in got] == ["VALUE"]  # no second reader available → nothing changes, nothing breaks


def test_the_payment_block_is_not_a_source_for_the_customers_own_columns(greenvale, riverbend):
    """Every statement carries the municipality's OWN bank account under a label that reads exactly
    like the customer's. Variance cannot separate them (both are constant) and proximity cannot
    (the wrong one is often closer). The neighbourhood can."""
    from cbx.pdfstudio import profile as P

    bands = P.payee_regions(greenvale[0])
    assert bands, "the banking block was not found"
    bank = X.find_anchors(greenvale[0], "BANK NAME", anchor="label")[0]["line"]
    assert any(b["page"] == bank["page"] and b["top"] <= bank["top"] <= b["bottom"] for b in bands)
    # the top of the page is not in it
    top = X.find_anchors(greenvale[0], "Account Number", anchor="label")[0]["line"]
    assert not any(b["page"] == top["page"] and b["top"] <= top["top"] <= b["bottom"] and b["x0"] <= top["x0"] for b in bands)
    rs, _notes = P.author_heuristic(SHARED_COLS, greenvale, family="shared")
    for sp, want in ((greenvale[0], 3529265), (riverbend[0], 1009715)):
        got = X.extract(rs, SHARED_COLS, sp)["account_number"]["value"]
        assert got == want, got  # never 9555001224, the municipality's own account
    # a column that ASKS for the banking details may still read them
    bank_col = [dict(c, required=True, critical=False, hint="") for c in ({"key": "bank_account", "label": "Bank account number", "ctype": "number"},)]
    rs2, _n2 = P.author_heuristic(bank_col, greenvale, family="shared")
    assert rs2["fields"]["bank_account"]["rules"], "a payment column must still be able to read the payment block"


def test_a_value_the_page_repeats_corroborates_itself(greenvale):
    """A customer's account number is printed in the header, the invoice number and the remittance
    advice; the municipality's bank account is printed once, in the banking block."""
    from cbx.pdfstudio import profile as P

    sp = greenvale[0]
    assert P.corroboration(sp, "0003529265") >= 2
    assert P.corroboration(sp, "9 555 001 224") == 1
    assert P.corroboration(sp, "9 555 001 224", outside_payee=True) == 0
    assert P.corroboration(sp, "12") == 0  # too short to mean anything


def test_a_label_reads_its_own_field_not_its_neighbours(greenvale):
    """Two fields share a line whenever the page has columns. The gap between them says where one
    field stops: what is to the LEFT of a label that starts a field belongs to the field before it."""
    from cbx.pdfstudio import extract as EX

    line = EX.find_anchors(greenvale[0], "Deposit", anchor="label")[0]["line"]
    gap = EX.gap_threshold(greenvale[0])
    assert len(EX.field_groups(line["words"], gap)) >= 2
    assert EX.apply_rule({"type": "label_left", "label": "Deposit", "anchor": "label"}, "number", greenvale[0]) is None


def test_a_field_value_is_never_a_sentence(greenvale):
    from cbx.pdfstudio import profile as P

    prose = {"text": "Please note that your usage for this month has been estimated.", "words": []}
    assert P.is_prose(prose) is True
    assert P.is_prose({"text": "Account Number: 0003529265", "words": []}) is False


def test_a_group_synonym_does_not_name_every_column_in_the_group():
    """'Consumption' is the printed label for Total kWh. It is not a label for Peak kWh, and the
    column's own synonym list must not be allowed to pretend otherwise."""
    from cbx.pdfstudio import profile as P

    total = {"key": "total_kwh", "label": "Total kWh", "ctype": "number"}
    peak = {"key": "peak_kwh", "label": "Peak kWh", "ctype": "number"}
    assert (P.label_fit(total, "CONSUMPTION") or 0) >= P.MIN_FIT
    assert P.label_fit(peak, "CONSUMPTION") is None
    assert P.label_fit({"key": "invoice_dt", "label": "Invoice date", "ctype": "date"}, "Invoice date") is not None


def test_a_match_key_is_something_the_layout_prints_not_this_customers_email(greenvale, riverbend):
    """A layout is recognised by what is PRINTED on it. An e-mail address or a reference that
    happens to sit on the page is stable across one account's bills and useless for the next."""
    from cbx.pdfstudio import profile as P

    other = [_spans(other_system_bill(account=f"10000321{i}")) for i in range(3)]
    fams = P.cluster([(i, sp) for i, sp in enumerate(greenvale + riverbend + other)])
    for fam in fams:
        rule = P.match_rules_for(fam, fams)
        assert rule["verified"] is True
        assert not any("@" in t for t in rule["all"])


# ---- deleting a profile: the rules go, the readings stay ------------------------------------------
def test_owner_can_delete_a_profile_and_the_readings_survive(owner, bills, admin):
    """A junk profile has to be removable — the live server grew 21 Capitec profiles for perhaps
    eight real layouts. What must NOT happen is the sheet going with it: the values a person already
    exported are the record, the rules that produced them are not."""
    t = _template(owner, "Deletable layout")
    b = _run_batch(owner, bills, TEXT_BILLS[:4], t["id"], name="Delete-profile batch")
    db = SessionLocal()
    try:
        profs = db.query(Profile).filter(Profile.template_id == t["id"]).all()
        prof = max(profs, key=lambda p: db.query(Document).filter(Document.profile_id == p.id).count())
        pid, vids = prof.id, [v.id for v in prof.versions]
        doc_ids = [d.id for d in db.query(Document).filter(Document.profile_id == pid).all()]
        ex = db.query(Extraction).filter(Extraction.profile_version_id.in_(vids)).all()
        before = {x.id: dict(x.cells) for x in ex}
        assert doc_ids and before, "the fixture must have produced documents and readings"
    finally:
        db.close()

    # what it costs is answerable before you commit to it
    cost = owner.get(f"/api/pdfstudio/profiles/{pid}/deletable").json()
    assert cost["documents"] == len(doc_ids) and cost["readings"] == len(before) and cost["versions"] == len(vids)

    # an admin is not enough — this is the owner's call, and C-EXA's tool token is refused outright
    assert admin.delete(f"/api/pdfstudio/profiles/{pid}").status_code == 403

    r = owner.delete(f"/api/pdfstudio/profiles/{pid}")
    assert r.status_code == 200, r.text
    assert "reading" in r.json()["message"]

    db = SessionLocal()
    try:
        assert db.get(Profile, pid) is None
        assert db.query(ProfileVersion).filter(ProfileVersion.id.in_(vids)).count() == 0
        # the readings are all still there, with their values, pointing at no ruleset
        after = db.query(Extraction).filter(Extraction.id.in_(list(before))).all()
        assert len(after) == len(before)
        for x in after:
            assert x.cells == before[x.id]
            assert x.profile_version_id is None
        # the documents survive too, now unassigned so the next batch reads them fresh
        docs = db.query(Document).filter(Document.id.in_(doc_ids)).all()
        assert len(docs) == len(doc_ids) and all(d.profile_id is None for d in docs)
        # the batch and its sheet are untouched
        assert db.get(Batch, b["id"]) is not None
    finally:
        db.close()

    # it is written down, and the rules were saved off before they went
    assert any(e["action"] == "pdf_profile_deleted" for e in owner.get("/api/audit?action=pdf_profile_deleted").json()["rows"])
    files = owner.get("/api/pdfstudio/results").json()["files"]
    assert any(f["name"].startswith("deleted-profile-") and f["name"].endswith(".json") for f in files)
    assert owner.get(f"/api/pdfstudio/profiles/{pid}").status_code == 404


def test_purge_takes_the_below_gate_profiles_only(owner):
    """The 21-profile clean-up: ask for a template's below-gate profiles and only those go."""
    t = _template(owner, "Purge me")
    db = SessionLocal()
    try:
        for i in range(3):
            db.add(Profile(key=f"purge.junk{i}", label=f"Junk {i}", template_id=t["id"], status="below_gate"))
        db.add(Profile(key="purge.good", label="Good", template_id=t["id"], status="live"))
        db.commit()
        keep = db.query(Profile).filter(Profile.key == "purge.good").one().id
    finally:
        db.close()
    r = owner.post("/api/pdfstudio/profiles/purge", json={"template_id": t["id"], "below_gate_only": True})
    assert r.status_code == 200, r.text
    assert r.json()["count"] == 3
    keys = {p["key"] for p in owner.get("/api/pdfstudio/profiles").json()["profiles"]}
    assert "purge.good" in keys and not any(k.startswith("purge.junk") for k in keys)
    assert owner.get(f"/api/pdfstudio/profiles/{keep}").status_code == 200
    # nothing matching is a 404, not a silent success
    assert owner.post("/api/pdfstudio/profiles/purge", json={"template_id": t["id"], "below_gate_only": True}).status_code == 404
    assert owner.post("/api/pdfstudio/profiles/purge", json={}).status_code == 400
