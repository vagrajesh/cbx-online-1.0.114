"""1.0.104 (SPEC-1104 Phase 2, Builder U): V07 consumption and tariff checks on utility bills, and V08 the estimate
register with its true-up. Bills are made straight in the database as Uma would have read them (account, read dates,
kWh, services excl. VAT) and go through Mia, Pip and Tess. Account numbers, amounts and tariffs are made up."""
from __future__ import annotations

import io
import random
from datetime import date

import openpyxl
import pytest

from cbx.ap import bill_history as BH
from cbx.ap import capture as C
from cbx.ap import estimates as ES
from cbx.ap import masters as M
from cbx.ap import variance as VAR
from cbx.db import SessionLocal
from cbx.models import AuditLog, User
from cbx.models.ap import (ApBatch, ApBillHistory, ApCapture, ApCheckMeter, ApDispute, ApEstimate, ApItem, ApRole, ApTariff,
                           ApUtilityAccount)

from .test_ap_capture_1098 import _admin, _app, _chart_of_accounts, _user, signed_in  # noqa: F401

MARK = "v1104u test"


@pytest.fixture(autouse=True)
def _clean():
    yield
    with SessionLocal() as s:
        keys = [k for (k,) in s.query(ApUtilityAccount.account_key).filter(ApUtilityAccount.updated_by == MARK).all()]
        s.query(ApBillHistory).filter(ApBillHistory.account_key.in_(keys or [""])).delete()
        s.query(ApEstimate).filter(ApEstimate.account_key.in_(keys or [""])).delete()
        s.query(ApDispute).filter(ApDispute.account_key.in_(keys or [""])).delete()
        s.query(ApCheckMeter).delete()
        s.query(ApTariff).delete()
        s.query(ApRole).delete()
        bids = [b for (b,) in s.query(ApBatch.id).filter(ApBatch.name.like("v1104u%") | ApBatch.name.like("Estimates %")).all()]
        s.query(ApCapture).filter(ApCapture.batch_id.in_(bids or [0])).delete()
        s.query(ApItem).filter(ApItem.batch_id.in_(bids or [0])).delete()
        s.query(ApBatch).filter(ApBatch.id.in_(bids or [0])).delete()
        s.query(ApUtilityAccount).filter(ApUtilityAccount.updated_by == MARK).delete()
        s.commit()


def _acct(site: str = "") -> str:
    no = "77" + "".join(random.choice("0123456789") for _ in range(8))
    with SessionLocal() as s:
        s.add(ApUtilityAccount(account_key=M.acct_key(no), account=no, vendor="700999", cost_centre="01306", site=site or ("Site " + no), municipality="ESKOM",
                               company_code="", source="excel", updated_by=MARK))
        s.commit()
    return no


def _uid(email: str) -> int:
    with SessionLocal() as s:
        return s.query(User).filter(User.email == email).one().id


def _batch(s, name: str = "v1104u bills") -> ApBatch:
    uid = s.query(User.id).order_by(User.id).first()[0]
    b = ApBatch(user_id=uid, name=name, status="ready", counts={})
    s.add(b)
    s.flush()
    return b


def _month_end(p: str) -> str:
    import calendar

    y, m = int(p[:4]), int(p[5:7])
    return f"{p}-{calendar.monthrange(y, m)[1]:02d}"


def _bill(s, b: ApBatch, acct: str, period: str, kwh: float, ex: float, *, rt: str = "Actual", start: str = "", water: float = 0, extras=None) -> ApCapture:
    end = _month_end(period)
    services = [{"service": "electricity", "ex": "%.2f" % ex, "vat": "%.2f" % (ex * 0.15), "how": "test", "sure": True}]
    if water:
        services.append({"service": "water", "ex": "%.2f" % water, "vat": "%.2f" % (water * 0.15), "how": "test", "sure": True})
    inv = f"INV{acct[-4:]}{period.replace('-', '')}{random.randint(100, 999)}"
    f = {"kind": "utility", "municipality": "ESKOM", "profile": "eskom", "account": acct, "site": "", "billing_date": end, "due_date": end,
         "statement_month": period, "invoice_no": inv, "read_start": start or f"{period}-01", "read_end": end, "reading_type": rt, "kwh": "%g" % kwh,
         "services": services, "adjustments": {}, "vat_printed": None, "current_ex": "", "current_incl": "", "checks": [], "ocr": False,
         "bill_extras": extras or {"kl": None, "tariff": "Business LV", "opening": None, "arrears": None, "closing": None}}
    it = ApItem(batch_id=b.id, name=f"{inv}.pdf", kind="utility", status="ok", fields={}, findings=[], history=[])
    s.add(it)
    s.flush()
    c = ApCapture(item_id=it.id, batch_id=b.id, bucket="utility", stage="match", fields=f, read={}, match={}, sap_rows=[], sap_issues=[], read_by="uma", profile="eskom")
    s.add(c)
    s.flush()
    return c


def _run(s, b: ApBatch) -> None:
    C.advance(s, b, live=False)
    s.commit()


def _history(s, acct: str, periods: list[str], kwh=1000.0, ex=1000.0) -> ApBatch:
    b = _batch(s)
    for p in periods:
        _bill(s, b, acct, p, kwh, ex)
    _run(s, b)
    return b


SIX = ["2026-03", "2026-04", "2026-05", "2026-06", "2026-07", "2026-08"]


# ---- the bill history ------------------------------------------------------------------------------------------------------
def test_every_utility_bill_goes_into_the_history_oldest_first():
    acct = _acct()
    with SessionLocal() as s:
        b = _batch(s)
        for p in reversed(SIX):  # read newest first: the history is still oldest → newest
            _bill(s, b, acct, p, 1000, 1000, water=200, extras={"kl": 31.5, "tariff": "Business LV", "opening": 500.0, "arrears": 0.0, "closing": 2610.0})
        _run(s, b)
        got = BH.periods(s, acct)
    assert [g["period"] for g in got] == SIX
    one = got[0]
    assert set(one) >= {"period", "read_start", "read_end", "kwh", "kl", "amounts", "reading_type", "tariff", "opening", "arrears", "closing", "capture_id"}
    assert one["amounts"] == {"electricity": 1000.0, "water": 200.0} and one["kwh"] == 1000 and one["kl"] == 31.5
    assert (one["reading_type"], one["tariff"], one["opening"], one["closing"]) == ("A", "Business LV", 500.0, 2610.0)
    assert one["read_start"] == "2026-03-01" and one["read_end"] == "2026-03-31"
    with SessionLocal() as s:  # the account as printed (leading zeros, spaces) finds the same bills
        assert len(BH.periods(s, "0" + acct)) == 6


def test_extras_reads_kl_tariff_and_balances_from_the_profile_rows():
    got = {"profile": {"fields": [{"field": "bal_opening", "type": "money"}, {"field": "bal_closing", "type": "money"}, {"field": "water_kl", "type": "number"},
                                  {"field": "start_reading_kl", "type": "number"}]},
           "canon": {"tariff_name": "Megaflex"},
           "rows": [{"service": "electricity", "bal_opening": "1 200,50", "bal_closing": "3 400.00", "x_txt": "x"},
                    {"service": "water", "water_kl": "42", "start_reading_kl": "9000"}]}
    x = BH.extras(got)
    assert x == {"kl": 42.0, "tariff": "Megaflex", "opening": 1200.5, "arrears": None, "closing": 3400.0}
    assert BH.extras({"profile": None}) in ({}, {"kl": None, "tariff": "", "opening": None, "arrears": None, "closing": None})


def test_the_backfill_fills_the_history_from_bills_read_before():
    acct = _acct()
    with SessionLocal() as s:
        _history(s, acct, SIX[:3])
        s.query(ApBillHistory).filter(ApBillHistory.account_key == M.acct_key(acct)).delete()
        s.commit()
        assert BH.periods(s, acct) == []
    n = BH.run_backfill(force=True)
    assert n >= 3
    with SessionLocal() as s:
        assert [g["source"] for g in BH.periods(s, acct)] == ["backfill"] * 3
        assert s.query(AuditLog).filter(AuditLog.action == "ap_bill_history_backfill").count() >= 1


# ---- V07: the checks -------------------------------------------------------------------------------------------------------
def test_a_bill_out_of_line_is_yellow_with_the_figures():
    acct, acct2 = _acct(), _acct()
    with SessionLocal() as s:
        _history(s, acct, SIX)
        b = _batch(s)
        c = _bill(s, b, acct, "2026-09", 2000, 2000)
        _run(s, b)
        s.refresh(c)
        why = "; ".join(x["text"] for x in c.sap_issues)
        assert c.sap_status == "yellow", c.sap_issues
        assert "amount excl. VAT R 2 000.00 is 100% above the average of the last 6 bills (R 1 000.00; allowed 20%)" in why  # TELKOM (T000): 20 %
        assert "kWh 2 000 kWh is 100% above the average of the last 6 bills (1 000 kWh" in why
        assert "rate" not in why  # the same R/kWh as before
        h = BH.row_for(s, c.id)
        assert {f["code"] for f in h.findings} == {"VAR_R", "VAR_KWH"} and h.flags == 2
        fig = next(f for f in h.findings if f["code"] == "VAR_R")["figures"]
        assert fig["expected"] == 1000 and fig["periods"] == SIX and fig["pct"] == 100
        # in line: green
        c2 = _bill(s, b, acct2, "2026-09", 1000, 1000)
        _run(s, b)
        s.refresh(c2)
        assert BH.row_for(s, c2.id).findings == []


def test_the_rate_is_checked_against_the_tariff_table_then_the_history():
    acct = _acct()
    with SessionLocal() as s:
        _history(s, acct, SIX)  # R 1.00 / kWh
        b = _batch(s)
        c = _bill(s, b, acct, "2026-09", 1000, 1300)  # R 1.30 / kWh: 30 % above, the amount itself 30 % above too
        _run(s, b)
        s.refresh(c)
        why = "; ".join(x["text"] for x in c.sap_issues)
        assert "rate R 1.3000/kWh is 30% above the account's last 6 bills (R 1.0000/kWh; allowed 15%)" in why
    a = _admin()
    r = a.post("/api/ap/tariffs", json={"profile": "eskom", "tariff": "Business LV", "effective_from": "2026-01-01", "c_per_kwh": "100", "fixed_r": "0"})
    assert r.status_code == 200, r.text
    assert a.post("/api/ap/tariffs", json={"municipality": "x", "tariff": "y"}).status_code == 400  # needs a rate
    assert _user().post("/api/ap/tariffs", json={"profile": "eskom", "c_per_kwh": 1}).status_code == 403  # admins only
    with SessionLocal() as s:
        C.test_batch(s, s.get(ApBatch, b.id))
        s.commit()
        h = BH.row_for(s, c.id)
        t = next(f for f in h.findings if f["code"] == "TARIFF")
        assert "electricity R 1 300.00 for 1 000 kWh is 30% above the tariff Business LV (R 1 000.00 at 100 c/kWh; allowed 15%)" == t["text"]
        assert not any(f["code"] == "RATE" for f in h.findings)  # the table wins over the history
    # blocks: 500 at 80 c + the rest at 120 c + R 100 fixed = 400 + 600 + 100 = 1 100
    with SessionLocal() as s:
        from cbx.ap import tariffs as T

        tr = s.query(ApTariff).one()
        tr.blocks, tr.c_per_kwh, tr.fixed_r = T.blocks_in("500=80; =120"), None, 100.0
        s.commit()
        assert T.expected(tr, 1000)["total"] == 1100.0
    # upload / download round trip (replaces the table)
    data = a.get("/api/ap/tariffs.xlsx").content
    r = a.post("/api/ap/tariffs/upload", files={"file": ("t.xlsx", data, "application/octet-stream")})
    assert r.status_code == 200 and r.json() == {"rows": 1, "replaced": 1}, r.text
    with SessionLocal() as s:
        assert s.query(ApTariff).one().blocks == [{"upto_kwh": 500.0, "c_per_kwh": 80.0}, {"upto_kwh": None, "c_per_kwh": 120.0}]
        assert s.query(AuditLog).filter(AuditLog.action == "ap_tariff").count() >= 2  # added, uploaded


def test_estimated_readings_running_and_a_gap_are_yellow():
    acct = _acct()
    with SessionLocal() as s:
        _history(s, acct, SIX[:5])
        b = _batch(s)
        _bill(s, b, acct, "2026-08", 1000, 1000, rt="Estimated")
        c = _bill(s, b, acct, "2026-09", 1000, 1000, rt="E", start="2026-09-12")  # the last reading ended 31 Aug
        _run(s, b)
        s.refresh(c)
        why = "; ".join(x["text"] for x in c.sap_issues)
        assert c.sap_status == "yellow"
        assert "estimated reading 2 bills running: ask for an actual reading" in why
        assert "reading starts 2026-09-12, 12 days after the last bill's reading ended (2026-08-31): a gap" in why


def test_check_meter_readings_and_the_sites_without_one():
    acct, other = _acct(site="TKM SITE A"), _acct()
    with SessionLocal() as s:
        _history(s, acct, SIX)
        _history(s, other, SIX, kwh=5000, ex=5000)
    u = _user()
    csv = b"site,period,kwh\nTKM SITE A,2026-09,800\nnowhere,2026-09,1\nbad,,\n"
    r = u.post("/api/ap/check-meters", files={"file": ("bcx.csv", csv, "text/csv")})
    assert r.status_code == 200 and r.json()["readings"] == 2 and r.json()["skipped_rows"] == [4], r.text
    with SessionLocal() as s:
        b = _batch(s)
        c = _bill(s, b, acct, "2026-09", 1000, 1000)
        _run(s, b)
        s.refresh(c)
        assert "billed 1 000 kWh is 25% above the check meter's 800 kWh (allowed 10%)" in "; ".join(x["text"] for x in c.sap_issues)
    rank = u.get("/api/ap/check-meters/ranking").json()
    keys = [x["account_key"] for x in rank["sites"]]
    assert M.acct_key(other) in keys and M.acct_key(acct) not in keys  # acct has a check meter now
    assert keys.index(M.acct_key(other)) < len(keys)
    wb = openpyxl.load_workbook(io.BytesIO(u.get("/api/ap/check-meters/ranking.xlsx").content))
    assert wb.active.title == "No check meter" and wb.active["A1"].value == "Rank"


def test_accept_and_dispute_with_evidence_and_the_hold_on_the_upload():
    acct, acct2 = _acct(), _acct()
    with SessionLocal() as s:
        _history(s, acct, SIX)
        b = _batch(s)
        c = _bill(s, b, acct, "2026-09", 2000, 2000)
        c2 = _bill(s, b, acct2, "2026-09", 1000, 1000)
        _run(s, b)
        hid, bid, cid, cid2 = BH.row_for(s, c.id).id, b.id, c.id, c2.id
    rev = _user("rev.u@test.local", "Rita Reviewer")
    lst = rev.get("/api/ap/variances").json()
    assert hid in [x["id"] for x in lst["bills"]] and lst["open"] >= 1
    assert any(x["key"] == "variances" and x["sub"] == "variances" for x in rev.get("/api/ap/needs").json()["controls"])
    # accept: a reason is needed; the decision is kept, audited and the bill turns green
    assert rev.post(f"/api/ap/variances/{hid}/accept", json={"reason": ""}).status_code == 400
    r = rev.post(f"/api/ap/variances/{hid}/accept", json={"reason": "New air-conditioning plant since September"})
    assert r.status_code == 200 and r.json()["decision"] == "accepted", r.text
    with SessionLocal() as s:
        assert s.get(ApCapture, cid).sap_status == "green"
        assert s.query(AuditLog).filter(AuditLog.action == "ap_variance_accepted", AuditLog.actor_id == _uid("rev.u@test.local")).count() == 1
        it = s.get(ApItem, s.get(ApCapture, cid).item_id)
        assert any("Variance accepted: New air-conditioning" in h["text"] for h in it.history)
    # the review workbook carries the comparison: periods, expected, actual, who accepted
    wb = openpyxl.load_workbook(io.BytesIO(rev.get(f"/api/ap/batches/{bid}/sap/review").content))
    ws = wb["Bill comparison"]
    head = [x.value for x in ws[1]]
    row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[1] == acct)
    got = dict(zip(head, row))
    assert got["Compared with (periods)"] == ", ".join(SIX) and got["Average R"] == 1000 and got["R excl. VAT"] == 2000
    assert got["Decision"] == "accepted" and got["Decided by"] == "Rita Reviewer" and "100% above" in got["Findings"]

    # dispute: held out of the upload (red) until paid under dispute (yellow) or resolved
    hid2 = None
    with SessionLocal() as s:
        h = BH.row_for(s, cid)
        h.decision, h.decision_reason = "", ""  # (as if the acceptance had not happened, to dispute it)
        s.commit()
        hid2 = h.id
    capt = _user("cap.u@test.local", "Carl Capturer")
    assert capt.post(f"/api/ap/variances/{hid2}/dispute", json={"reason": ""}).status_code == 400
    r = capt.post(f"/api/ap/variances/{hid2}/dispute", json={"reason": "Meter read twice", "municipal_ref": "COJ-123", "owner": "Lindi (Council Liaison)"})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["amount"] == 1000.0 and d["status"] == "open" and not d["pay_under_dispute"]  # the part above the average
    with SessionLocal() as s:
        cap = s.get(ApCapture, cid)
        assert cap.sap_status == "red" and "disputed with the municipality (dispute #%d, municipal ref COJ-123): held out of the upload" % d["id"] in cap.sap_issues[-1]["text"]
    r = capt.post(f"/api/ap/batches/{bid}/sap/upload", json={"include_yellow": True})
    assert r.status_code == 200
    with SessionLocal() as s:
        assert s.get(ApCapture, cid).uploaded_at is None and s.get(ApCapture, cid2).uploaded_at is not None  # the disputed bill stayed out
    assert capt.put(f"/api/ap/disputes/{d['id']}", json={"pay_under_dispute": True}).status_code == 200
    with SessionLocal() as s:
        cap = s.get(ApCapture, cid)
        assert cap.sap_status == "yellow" and any("paid under dispute" in x["text"] for x in cap.sap_issues)
    assert capt.put(f"/api/ap/disputes/{d['id']}", json={"status": "resolved", "resolution": ""}).status_code == 400
    r = capt.put(f"/api/ap/disputes/{d['id']}", json={"status": "resolved", "resolution": "Credit of R 1 000 on the next bill"})
    assert r.status_code == 200 and r.json()["status"] == "resolved"
    assert capt.put(f"/api/ap/disputes/{d['id']}", json={"pay_under_dispute": False}).status_code == 400  # closed
    with SessionLocal() as s:
        assert s.query(AuditLog).filter(AuditLog.action == "ap_dispute").count() >= 3
        assert s.get(ApCapture, cid).sap_status == "green"
    x = openpyxl.load_workbook(io.BytesIO(capt.get("/api/ap/disputes.xlsx").content)).active
    assert x["A1"].value == "Dispute" and x.max_row >= 2


def test_a_reviewer_who_worked_on_the_bill_cannot_accept_its_variance():
    acct = _acct()
    a = _admin()
    with SessionLocal() as s:
        _history(s, acct, SIX)
        b = _batch(s)
        c = _bill(s, b, acct, "2026-09", 2000, 2000)
        _run(s, b)
        hid = BH.row_for(s, c.id).id
        c.decided_by_id = _uid("cap.admin@test.local")  # the admin typed this bill: someone else reviews it
        s.commit()
    r = a.post(f"/api/ap/variances/{hid}/accept", json={"reason": "fine"})
    assert r.status_code == 403 and "different person" in r.json()["detail"]


# ---- V08: estimates -----------------------------------------------------------------------------------------------------------
def test_proposals_by_basis():
    acct = _acct()
    with SessionLocal() as s:
        b = _batch(s)
        for n, p in enumerate(SIX):
            _bill(s, b, acct, p, 100 * (n + 1), 100 * (n + 1))  # 100 … 600
        _run(s, b)
        k = M.acct_key(acct)
        assert ES.propose(s, k, "2026-09", "last")["value"] == 600
        assert ES.propose(s, k, "2026-09", "avg3")["value"] == 500
        p6 = ES.propose(s, k, "2026-09", "avg6")
        assert p6["value"] == 350 and p6["workings"]["periods"] == SIX and p6["vat"] == 52.5
        daily = ES.propose(s, k, "2026-09", "daily")  # 2 100 over 178 days read (1–31 Mar is 30 days … ) × 30 days in September
        assert daily["value"] == round(2100 / 178 * 30, 2) and "178 days read" in daily["workings"]["how"]
        assert ES.propose(s, k, "2026-09")["basis"] == "avg6"  # the portfolio's default


def test_missing_bill_prepare_approve_post_and_true_up(monkeypatch):
    monkeypatch.setattr(ES, "_today", lambda: date(2026, 10, 19))
    acct = _acct()
    with SessionLocal() as s:
        _history(s, acct, SIX)  # last bill August; September's due 30 Sep + 10 days grace → missing on 19 Oct
        miss = {m["account_key"]: m for m in ES.missing(s, today=date(2026, 10, 19))}
        m = miss[M.acct_key(acct)]
        assert m["period"] == "2026-09" and m["missing"] == 1 and m["proposal"]["value"] == 1000 and m["due"] == "2026-10-10"
        assert M.acct_key(acct) not in {x["account_key"] for x in ES.missing(s, today=date(2026, 10, 5))}  # still within the grace
    prep, appr = _user("prep.u@test.local", "Pat Preparer"), _user("appr.u@test.local", "Abe Approver")
    k = M.acct_key(acct)
    assert prep.post("/api/ap/estimates", json={"account_key": k, "period": "2026-09", "value": "1500"}).status_code == 400  # a changed value says why
    r = prep.post("/api/ap/estimates", json={"account_key": k, "period": "2026-09"})
    assert r.status_code == 200, r.text
    e = r.json()
    assert e["status"] == "prepared" and e["value"] == 1000 and e["reference"] == f"EST-{k}-202609"
    assert prep.post("/api/ap/estimates", json={"account_key": k, "period": "2026-09"}).status_code == 400  # one per period
    assert prep.post(f"/api/ap/estimates/{e['id']}/approve", json={}).status_code == 403  # never the preparer
    assert any(x["key"] == "estimates_approve" for x in appr.get("/api/ap/needs").json()["controls"])
    assert not any(x["key"] == "estimates_approve" for x in prep.get("/api/ap/needs").json()["controls"])
    r = appr.post(f"/api/ap/estimates/{e['id']}/approve", json={})
    assert r.status_code == 200, r.text
    e = r.json()
    assert e["status"] == "approved" and e["approved_by"] == "Abe Approver" and e["capture_id"]
    with SessionLocal() as s:
        ec = s.get(ApCapture, e["capture_id"])
        b = s.get(ApBatch, ec.batch_id)
        assert b.name.startswith("Estimates 2026-09") and ec.stage == "ready" and ec.sap_status in ("green", "yellow"), ec.sap_issues
        assert ec.sap_rows[0]["refdocnr"] == f"EST-{k}-202609" and "Estimate" in ec.sap_rows[0]["docheadertext"]
        assert ec.sap_rows[0]["amount"] == "1150.00"  # R 1 000 + 15 % VAT to the vendor
        assert BH.row_for(s, ec.id) is None  # an estimate is not a bill
        assert {a for (a,) in s.query(AuditLog.action).filter(AuditLog.target == f"ap_estimate:{e['id']}").all()} >= {"ap_estimate_created", "ap_estimate_approved"}
        ebid = b.id
    assert appr.get(f"/api/ap/captures/{e['capture_id']}").status_code == 200 and appr.get(f"/api/ap/batches/{ebid}").status_code == 200  # no file: still opens
    # the estimate goes out in an upload, then the real bill comes: reversed by a credit memo EST-…-R
    assert appr.post(f"/api/ap/batches/{ebid}/sap/upload", json={"include_yellow": True}).status_code == 200
    with SessionLocal() as s:
        b = _batch(s)
        real = _bill(s, b, acct, "2026-09", 1100, 1100)
        _run(s, b)
        est = s.get(ApEstimate, e["id"])
        assert (est.status, est.actual_capture_id, est.actual_value, est.difference, est.true_up) == ("trued_up", real.id, 1100.0, 100.0, "reverse")
        rc = s.get(ApCapture, est.reversal_capture_id)
        assert rc.fields["invoice_no"] == f"EST-{k}-202609-R"
        it = s.get(ApItem, real.item_id)
        assert any(f"Trues up estimate EST-{k}-202609" in h["text"] for h in it.history)
        assert s.query(AuditLog).filter(AuditLog.action == "ap_estimate_trued_up", AuditLog.target == f"ap_estimate:{e['id']}").count() == 1
        rcid = rc.id
    with SessionLocal() as s:  # the reversal went through Mia, Pip and Tess (after the commit): a KG credit memo
        rc = s.get(ApCapture, rcid)
        assert rc.stage == "ready" and rc.sap_rows and rc.sap_rows[0]["pkey"] == "21" and rc.sap_rows[0]["amount"] == "1150.00", (rc.stage, rc.sap_rows[:1])
    wb = openpyxl.load_workbook(io.BytesIO(appr.get("/api/ap/estimates/report.xlsx?month=2026-09").content))
    assert wb["Per portfolio"]["A2"].value == "Portfolio"


def test_an_estimate_not_yet_uploaded_is_taken_out_when_the_bill_comes(monkeypatch):
    monkeypatch.setattr(ES, "_today", lambda: date(2026, 10, 19))
    acct = _acct()
    with SessionLocal() as s:
        _history(s, acct, SIX)
        k = M.acct_key(acct)
        p, a = s.query(User).filter(User.email == "prep.u@test.local").first(), s.query(User).filter(User.email == "appr.u@test.local").first()
    if p is None or a is None:
        _user("prep.u@test.local", "Pat Preparer"), _user("appr.u@test.local", "Abe Approver")
    with SessionLocal() as s:
        p, a = s.query(User).filter(User.email == "prep.u@test.local").one(), s.query(User).filter(User.email == "appr.u@test.local").one()
        e = ES.prepare(s, p, {"account_key": k, "period": "2026-09", "basis": "last"})
        ES.approve(s, a, e.id)
        s.commit()
        assert ES.report_rows(s, "2026-09")[0]  # an open estimate counts in the month's report
        b = _batch(s)
        _bill(s, b, acct, "2026-09", 900, 900)
        _run(s, b)
        e = s.get(ApEstimate, e.id)
        ec = s.get(ApCapture, e.capture_id)
        assert e.status == "trued_up" and e.true_up == "cancelled" and e.difference == -100.0 and e.reversal_capture_id is None
        assert ec.stage == "skipped" and ec.uploaded_at is None


def test_estimates_running_escalate_to_council_liaison(monkeypatch):
    monkeypatch.setattr(ES, "_today", lambda: date(2026, 12, 19))
    acct = _acct()
    u = _user("prep.u@test.local", "Pat Preparer")
    _user("appr.u@test.local", "Abe Approver")
    with SessionLocal() as s:
        _history(s, acct, SIX)
        k = M.acct_key(acct)
        p, a = s.query(User).filter(User.email == "prep.u@test.local").one(), s.query(User).filter(User.email == "appr.u@test.local").one()
        for per in ("2026-09", "2026-10"):
            ES.approve(s, a, ES.prepare(s, p, {"account_key": k, "period": per}).id)
        s.commit()
        assert k not in [r["account_key"] for r in ES.escalations(s)]  # 2 < 3
        ES.approve(s, a, ES.prepare(s, p, {"account_key": k, "period": "2026-11"}).id)
        s.commit()
        esc = {r["account_key"]: r for r in ES.escalations(s)}
        assert esc[k]["count"] == 3 and esc[k]["periods"] == ["2026-09", "2026-10", "2026-11"] and esc[k]["value"] == 3000
    assert any(x["key"] == "estimates_escalate" for x in u.get("/api/ap/needs").json()["controls"])
    r = u.post("/api/ap/estimates/escalate", json={"account_key": k, "note": "Mailed Lindi"})
    assert r.status_code == 200 and r.json()["escalated_by"] == "Pat Preparer"
    assert not any(x["key"] == "estimates_escalate" and "1 account" in x["text"] and k in str(x) for x in u.get("/api/ap/needs").json()["controls"])
    with SessionLocal() as s:
        assert s.query(AuditLog).filter(AuditLog.action == "ap_estimate_escalated").count() >= 1
    got = u.get("/api/ap/estimates?status=").json()
    assert {"missing", "estimates", "escalations", "bases"} <= set(got)


def test_the_new_settings_are_offered_in_settings_portfolios():
    from cbx.ap import portfolio as PF

    keys = PF.known_keys()
    for k in ("variance_min_bills", "estimate_run_n", "period_gap_days", "check_meter_pct", "estimate_basis_by", "estimate_true_up"):
        assert k in keys, k


def test_read_again_leaves_estimate_documents_as_they_are(monkeypatch):
    monkeypatch.setattr(ES, "_today", lambda: date(2026, 10, 19))
    acct, other = _acct(), _acct()
    prep, appr = _user("prep.u@test.local", "Pat Preparer"), _user("appr.u@test.local", "Abe Approver")
    with SessionLocal() as s:
        _history(s, acct, SIX)
    k = M.acct_key(acct)
    e = prep.post("/api/ap/estimates", json={"account_key": k, "period": "2026-09"}).json()
    e = appr.post(f"/api/ap/estimates/{e['id']}/approve", json={}).json()
    with SessionLocal() as s:
        ec = s.get(ApCapture, e["capture_id"])
        ebid, before = ec.batch_id, (ec.stage, ec.sap_status, dict(ec.fields), list(ec.sap_rows))
    # a folder of estimates only: refused, in plain words, nothing changed
    r = appr.post(f"/api/ap/batches/{ebid}/read")
    assert r.status_code == 400 and r.json()["detail"] == "Estimates are not read again: change or reject the estimate instead."
    with SessionLocal() as s:
        ec = s.get(ApCapture, e["capture_id"])
        assert (ec.stage, ec.sap_status, dict(ec.fields), list(ec.sap_rows)) == before
        # a folder with a real bill beside it: the bill goes back to Sol, the estimate stays, its file says why
        b = s.get(ApBatch, ebid)
        bill = _bill(s, b, other, "2026-09", 1000, 1000)
        _run(s, b)
        from cbx.ap import service as S

        S.read_again(s, s.query(User).filter(User.email == "appr.u@test.local").one(), b)
        s.flush()
        ec, bill = s.get(ApCapture, e["capture_id"]), s.get(ApCapture, bill.id)
        assert (ec.stage, ec.sap_status, dict(ec.fields), list(ec.sap_rows)) == before
        assert bill.stage == "sort" and bill.fields == {}
        assert any(h["text"] == ES.NOT_READ_AGAIN for h in s.get(ApItem, ec.item_id).history)
        assert ES.item_ids(s, ebid) == {ec.item_id}
        s.rollback()


# ---- 1.0.104 review fixes ------------------------------------------------------------------------------------------------------
def test_the_person_who_dropped_the_folder_cannot_accept_its_variance():
    acct = _acct()
    rev = _user("rev.u@test.local", "Rita Reviewer")
    with SessionLocal() as s:
        _history(s, acct, SIX)
        b = _batch(s)
        b.user_id = _uid("rev.u@test.local")  # Rita dropped this folder
        c = _bill(s, b, acct, "2026-09", 2000, 2000)
        _run(s, b)
        hid = BH.row_for(s, c.id).id
    r = rev.post(f"/api/ap/variances/{hid}/accept", json={"reason": "fine by me"})
    assert r.status_code == 403 and "dropped its folder" in r.json()["detail"]


def test_the_preparer_and_approver_of_an_estimate_touched_its_document(monkeypatch):
    from cbx.ap import roles as R

    monkeypatch.setattr(ES, "_today", lambda: date(2026, 10, 19))
    acct = _acct()
    prep, appr = _user("prep.u@test.local", "Pat Preparer"), _user("appr.u@test.local", "Abe Approver")
    with SessionLocal() as s:
        _history(s, acct, SIX)
    e = prep.post("/api/ap/estimates", json={"account_key": M.acct_key(acct), "period": "2026-09"}).json()
    e = appr.post(f"/api/ap/estimates/{e['id']}/approve", json={}).json()
    with SessionLocal() as s:
        ec = s.get(ApCapture, e["capture_id"])
        assert {_uid("prep.u@test.local"), _uid("appr.u@test.local")} <= R.touched_by(s, ec)
        assert not R.independent(s, s.query(User).filter(User.email == "prep.u@test.local").one(), ec)


def test_an_accepted_variance_needs_a_new_decision_when_the_figures_change():
    acct = _acct()
    rev = _user("rev2.u@test.local", "Rob Reviewer")  # (item ids are reused after the clean-up: a fresh reviewer has no old audit rows)
    with SessionLocal() as s:
        _history(s, acct, SIX)
        b = _batch(s)
        c = _bill(s, b, acct, "2026-09", 2000, 2000)
        _run(s, b)
        hid, cid, bid = BH.row_for(s, c.id).id, c.id, b.id
    assert rev.post(f"/api/ap/variances/{hid}/accept", json={"reason": "New plant"}).status_code == 200
    with SessionLocal() as s:
        c = s.get(ApCapture, cid)
        assert c.sap_status == "green"
        f = dict(c.fields)
        f["services"] = [{**f["services"][0], "ex": "2500.00", "vat": "375.00"}]  # a person corrected the bill's figures
        c.fields, c.stage = f, "sap"
        s.commit()
        C.advance(s, s.get(ApBatch, bid), live=False)
        s.commit()
        h, c = s.get(ApBillHistory, hid), s.get(ApCapture, cid)
        assert h.decision == "" and h.total_excl == 2500 and c.sap_status == "yellow"
        assert any(x["text"] == BH.CHANGED_LINE for x in s.get(ApItem, c.item_id).history)
        # the same figures again: nothing more to decide than before (no second line)
        C.test_batch(s, s.get(ApBatch, bid))
        s.commit()
        assert sum(x["text"] == BH.CHANGED_LINE for x in s.get(ApItem, c.item_id).history) == 1


def test_the_backfill_commits_in_chunks_and_is_safe_to_run_twice(monkeypatch):
    acct = _acct()
    with SessionLocal() as s:
        _history(s, acct, SIX[:5])
        s.query(ApBillHistory).filter(ApBillHistory.account_key == M.acct_key(acct)).delete()
        s.commit()
    monkeypatch.setattr(BH, "BACKFILL_CHUNK", 2)
    calls = []
    real = BH.backfill
    monkeypatch.setattr(BH, "backfill", lambda db, after_id=0, limit=None: calls.append(limit) or real(db, after_id, limit))
    assert BH.run_backfill(force=True) >= 5 and len(calls) >= 3 and set(calls) == {2}
    assert BH.run_backfill(force=True) == 0  # the second run adds nothing
    with SessionLocal() as s:
        assert [g["period"] for g in BH.periods(s, acct)] == SIX[:5]
