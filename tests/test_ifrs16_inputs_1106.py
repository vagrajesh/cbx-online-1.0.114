"""1.0.106: IFRS 16 reads the files the team really feeds it (Edwin, 19 Oct: "IFRS is not working, I'm getting errors").
Made-up files with the real exports' shapes — every name and number is invented:
  - the Sequentra 'Validation - Record Info Tab with Custom Fields' export: a title block above the headings, no Address 1/2
    (Building Number + Street Name instead), discount rates as fractions (0.0745) and as percentages (7.51), postal codes as text;
  - the 'PIW Step Payments' report: a title block with 'Total Records', an amount stored as text;
  - the RR06 contract register: many tabs; the Contract Movement tab's move type column has NO heading (a merged 'Movement'
    banner over the date and the move type); merged group banners over Rent/Parking/Other; 'Current Lease Term (Years)';
    headings and values with trailing spaces; numbers stored as text;
  - the LA 'Portfolio Transactions' report: a title row, text numbers, dates as Excel serial numbers, a Parameters tab — saved
    as an .xlsx with an .xls name here (the real one is a BIFF .xls, which xlrd reads; there is no .xls writer to make one).
"""
import io
import tempfile
import time
import zipfile
from datetime import datetime
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from cbx.ifrs16 import ENGINE_DIR  # noqa: F401 — puts the engine on sys.path
from cbx.main import app

from .conftest import _make, login

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
D = datetime


def _serial(d: datetime) -> float:
    return float((d - datetime(1899, 12, 30)).days)


def make_files(out: Path) -> dict:
    out.mkdir(parents=True, exist_ok=True)
    # ---- the LA Portfolio Transactions report ------------------------------------------------------------------
    la = Workbook()
    ws = la.active
    ws.title = "Portfolio Transactions"
    ws.append(["Portfolio Transactions"])
    ws.append(["Lessee", "Funder", "Transaction #", "Cost Center", "Lease Type", "BU", "Entity", "Asset Location(s)", "Start Date",
               "Effective End Date", "Term (mos.)", "Adv/Arr", "Current Payment Freq.", "Original Amount Financed (Reporting Currency) ",
               "Current Payment (Transactional Currency)", "Trans. Curr.", "Remaining (Reporting Currency)", "Deal Status", "IBR"])
    for ref, start, end, pay in (("T002P001", D(2023, 9, 1), D(2026, 8, 31), "21 000"), ("T003P001", D(2021, 3, 1), D(2026, 12, 31), "9000"),
                                 ("T007P001-M001", D(2024, 1, 1), D(2027, 12, 31), "5 000.50")):
        ws.append(["ABSA Bank Limited", "Invented Landlords (Pty) Ltd", ref, ref[:4], "FMV", "ABSA Bank Limited", "ABSA Bank Limited",
                   "South Africa", _serial(start), _serial(end), 36, "Advance", "Monthly", pay, float(pay.replace(" ", "")), "ZAR", "100000",
                   "Booked", 7.25])
    ws.append(["ABSA Bank Limited", "Old Landlord", "T008P001", "T008", "FMV", "ABSA Bank Limited", "ABSA Bank Limited", "South Africa",
               _serial(D(2020, 1, 1)), _serial(D(2022, 1, 1)), 24, "Advance", "Monthly", "1", 1.0, "ZAR", "0", "Disposed (Modification)", 7.0])
    p = la.create_sheet("Parameters")
    for r in (["Portfolio Transactions", ""], ["Requested By:", "Made Up"], [], ["Parameter", "Value Specified"], ["As At", "09/30/2026"]):
        p.append(r)

    # ---- the RR06 contract register -----------------------------------------------------------------------------
    rr = Workbook()
    s2 = rr.active
    s2.title = "Sheet2"
    s2.append(["Lease Reference", "Current Rentable Area", "Annual Rent"])
    s2.append(["T002P001", 100, 252000])
    rr.create_sheet("Control Sheet").append(["Lease Actions"])
    cover = rr.create_sheet("Absa CRES Contracts Cover")
    cover.append(["Absa CRES Contracts"])
    cover.append([])
    cover.append(["Report Description"])
    con = rr.create_sheet("Contracts")
    banner = ["Property"] + [None] * 8 + ["Current Charges", None, None, "Charges Old", None, None]
    heads = ["Property Ref", "Building Name ", "Lease Type", "Lease Marker", "Lease Reference", "Initial lease Start Date", "Lease Expiry Date",
             "Current Lease Term (Years)", "Legal Entity", "Rent", "Parking", "Other", "Rent", "Parking", "Other"]
    con.append(banner)
    con.append(heads)
    con.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
    con.merge_cells(start_row=1, start_column=10, end_row=1, end_column=12)
    con.merge_cells(start_row=1, start_column=13, end_row=1, end_column=15)
    for r in (["T001", "Made Up Mall", "Retail", "Payable   ", "T001P001", D(2026, 9, 1), D(2029, 8, 31), 3.0, "Absa Bank Limited", 10000, 500, 0, 0, 0, 0],
              ["T002", "Invented Plaza", "Retail", "Payable", "T002P001", D(2023, 9, 1), D(2029, 8, 31), 3.0, "Absa Bank Limited", 21000, "1 500,00", 0, 0, 0, 0],
              ["T003", "Pretend Kiosk", "ATM", "Payable", "T003P001", D(2021, 3, 1), D(2026, 12, 31), 5.0, "Absa Bank Limited", 9000, 0, 0, 0, 0, 0],
              ["T004", "Short Stay", "Retail", "Payable", "T004P001", D(2026, 1, 1), D(2026, 12, 31), 1.0, "Absa Bank Limited", 4000, 0, 0, 0, 0, 0],
              ["T005", "Free Corner", "ATM", "Payable", "T005P001", D(2024, 1, 1), D(2028, 12, 31), 5.0, "Absa Bank Limited", 0, 0, 0, 900, 0, 0],
              ["T006", "Month Parking", "Parking", "Payable", "T006P001", D(2015, 1, 1), "Monthly", "Monthly", "Absa Bank Limited", 0, 850, 0, 0, 0, 0],
              ["T009", "Not Booked", "Retail", "Payable", "T009P001", D(2024, 1, 1), D(2029, 12, 31), 5.0, "Absa Bank Limited", 7000, 0, 0, 0, 0, 0]):
        con.append(r)
    cm = rr.create_sheet("Contract Movement")
    cm.append(["Contract Movement"])
    cm.append(["Movement do not include rent escalations as per lease agreements."])
    cm.append([])
    cm.append(["Movement", None, "Property", None, None, "Contracts", None, None, None, None, None, None, None])
    cm.merge_cells("A4:B4")
    cm.append([None, None, "Property Ref", "One View Ref", "Building Name ", "Lease Reference", "Contract Name", "Landlord / Tenant Name",
               "Initial Lease Start Date", "Current Start Date", "Lease Expiry Date", "Termination Date", "Comments"])
    for r in ([D(2026, 8, 1), "Renewal", "T010", "90010X", "Earlier Mall", "T010P001", "Earlier", "Earlier Ltd", D(2020, 8, 1), D(2026, 8, 1),
               D(2029, 7, 31), None, "Renewal loaded August 2026"],
              [D(2026, 9, 1), "Addition ", "T001", "90001X", "Made Up Mall", "T001P001", "Made Up Mall", "Made-Up Properties (Pty) Ltd", D(2026, 9, 1),
               D(2026, 9, 1), D(2029, 8, 31), None, "New Lease loaded September 2026"],
              [D(2026, 9, 1), "Renewal", "T002", "90002X", "Invented Plaza", "T002P001", "Invented Plaza", "Invented Landlords (Pty) Ltd", D(2023, 9, 1),
               D(2026, 9, 1), D(2029, 8, 31), None, "Renewal loaded September 2026"],
              [D(2026, 9, 1), "Removal", "T003", "90003X", "Pretend Kiosk", "T003P001", "Pretend Kiosk", "Pretend Trust", D(2021, 3, 1), D(2021, 3, 1),
               D(2026, 12, 31), D(2026, 9, 30), "Termination loaded September 2026"]):
        cm.append(r)
    rr.create_sheet("AuditLogs").append(["Full Name", "Sheet Name", "Column Name"])

    # ---- the Sequentra Record Info export -----------------------------------------------------------------------
    sq = Workbook()
    ws = sq.active
    ws.title = "Sheet0"
    ws.append(["2. Validation - Record Info Tab with Custom Fields"])
    ws.append([])
    ws.append(["Created by someone@example.com on Monday, October 19, 2026 05:15:23 EDT GMT-0400."])
    ws.append([])
    ws.append(["Level 1", "Record ID", "Record Name", "Building Number", "Street Name", "Town", "Suite", "City", "State", "Postal Code", "Country",
               "Lease Status", "ABSA Lease Type", "Legal Entity", "Move-In Date", "Current Term Start Date", "Current Lease End Date", "Currency",
               "Contract Cost Centre", "Discount Rate", "Signature Date", "New Discount Rate", "External Reference"])
    for r in (["ABSAHP", "T001P001", "Made Up Mall", "15", "Made Up Rd", "Sandton", "", "Johannesburg", "Gauteng", "0200", "South Africa", "Active",
               "Retail", "Absa Bank Limited", D(2026, 9, 1), D(2026, 9, 1), D(2029, 8, 31), "ZAR", "SA600001", 0.0745, D(2026, 8, 20), None, "90001X"],
              ["ABSAHP", "T002P001", "Invented Plaza", "", "Pretend Ave", "", "", "Durban", "KwaZulu-Natal", "4001", "South Africa", "Active",
               "Retail", "Absa Bank Limited", D(2023, 9, 1), D(2026, 9, 1), D(2029, 8, 31), "ZAR", "SA600002", 0.0749, D(2026, 8, 15), 7.51, "90002X"],
              ["ABSAHP", "T003P001", "Pretend Kiosk", "1", "Nowhere St", "", "", "Pretoria", "Gauteng", "0002", "South Africa", "Terminated",
               "ATM", "Absa Bank Limited", D(2021, 3, 1), D(2021, 3, 1), D(2026, 12, 31), "ZAR", "SA600003", None, D(2021, 2, 1), 7.0, "90003X"]):
        ws.append(r)

    # ---- the PIW Step Payments report ---------------------------------------------------------------------------
    st = Workbook()
    ws = st.active
    ws.title = "Sheet0"
    ws.append(["DRAFT PIW Step Payments - 3/23/26", None, None, None, None, None, "ABSA"])
    ws.append([])
    ws.append(["Created by someone@example.com on Monday, October 19, 2026.", None, None, None, None, None, "Total Records 7"])
    ws.append([])
    ws.append(["Schedule Number", "Expense Type Name", "StartDate", "EndDate", "Number of Payments", "Payment Amount", "Payment Frequency"])
    for r in (["T001P001", "Base Rent", D(2026, 9, 1), D(2027, 8, 31), 12, 10000, "Monthly"],
              ["T001P001", "Base Rent", D(2027, 9, 1), D(2028, 8, 31), 12, 10600, "Monthly"],
              ["T001P001", "Base Rent", D(2028, 9, 1), D(2029, 8, 31), 12, 11236, "Monthly"],
              ["T001P001", "Parking", D(2026, 9, 1), D(2029, 8, 31), 36, 500, "Monthly"],
              ["T002P001", "Base Rent", D(2026, 9, 1), D(2029, 8, 31), 36, 21000, "Monthly"],
              ["T002P001", "Parking", D(2026, 9, 1), D(2029, 8, 31), 36, "1 500,00", "Monthly"],
              ["T003P001", "Base Rent", D(2021, 3, 1), D(2026, 12, 31), 70, 9000, "Monthly"]):
        ws.append(r)

    paths = {"la": out / "Transaction-Made.Up-10621.xls", "movements": out / "RR06_Made_Up_Contracts.xlsx",
             "sequentra": out / "2._Validation_-_Record_Info_Tab_with_Custom_Fields.xlsx", "steps": out / "DRAFT_PIW_Step_Payments.xlsx"}
    for key, wb in (("la", la), ("movements", rr), ("sequentra", sq), ("steps", st)):
        wb.save(paths[key])
    return paths


@pytest.fixture()
def files(tmp_path):
    return make_files(tmp_path / "in")


def _person(owner, email, name):
    ids = [t["id"] for t in owner.get("/api/users").json()["tools"] if t["slug"] == "ifrs16"]
    pw = _make(owner, email, name, "user", tools=ids)
    c = TestClient(app)
    login(c, email, pw)
    return c


def _upload(c, paths: dict, slots=None):
    slots = slots or {k: k for k in paths}
    r = c.post("/api/ifrs16/uploads", files={slot: (Path(paths[src]).name, Path(paths[src]).read_bytes(), XLSX) for slot, src in slots.items()})
    assert r.status_code == 200, r.text
    return next(u for u in c.get("/api/ifrs16").json()["uploads"] if u["id"] == r.json()["id"])


def _sheet(z, prefix, tab):
    name = next(n for n in z.namelist() if n.startswith(prefix))
    ws = load_workbook(io.BytesIO(z.read(name)))[tab]
    rows = list(ws.iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:]]


def _ciw_tab(z):
    name = next(n for n in z.namelist() if n.startswith("CIW"))
    return next(t for t in load_workbook(io.BytesIO(z.read(name))).sheetnames if "ddress" in t)


# ---- the whole month, as the page does it --------------------------------------------------------------------------
def test_the_real_shapes_are_read_and_a_month_comes_out(owner, files):
    c = _person(owner, "nomsa.ifrs1106@test.local", "Nomsa Leases")
    u = _upload(c, files)
    assert u["status"] == "ready", u["note"]
    assert [m["value"] for m in u["months"]][:2] == ["2026-09", "2026-08"], u["months"]
    assert "7 step payments" in u["note"] and "3 Sequentra rows" in u["note"], u["note"]
    r = c.post(f"/api/ifrs16/uploads/{u['id']}/month", json={"month": "2026-09"})
    assert r.status_code == 200, r.text
    m = next(x for x in c.get("/api/ifrs16").json()["months"] if x["id"] == r.json()["id"])
    assert m["status"] == "done", m["note"]
    assert "Movement tab used: Contract Movement" in m["recon"] and "Move type: Addition: 1" in m["recon"], m["recon"]
    z = zipfile.ZipFile(io.BytesIO(c.get(f"/api/ifrs16/files/{m['workbooks']['id']}").content))
    # the addition (not on LA): its payment schedule is the step payments report's, rent + parking added per month
    steps = [(x["Starting Payment Number"], x["Number Payments"], x["Payment Amount"]) for x in _sheet(z, "PIW", "Step Payments") if x["Schedule Number"] == "T001P001"]
    assert steps == [(1, 12, 10500), (13, 12, 11100), (25, 12, 11736)], steps
    piw = next(x for x in _sheet(z, "PIW", "Schedule") if x["Schedule Number"] == "T001P001")
    assert float(piw["IBR"]) == pytest.approx(7.45), piw["IBR"]  # Sequentra's 0.0745 is 7.45 %, not 0.0745 %
    assert piw["Duration"] == 36
    # the renewal (on LA): the parking amount stored as '1 500,00' is R1 500, not nil
    mdiw = [(x["Starting Payment Number"], x["Number Payments"], x["Payment Amount"]) for x in _sheet(z, "MDIW", "Step Payments") if x["Schedule Number"] == "T002P001"]
    assert mdiw == [(1, 36, 22500)], mdiw
    # the CIW address of the new lease: Building Number + Street Name, not the lease name
    ciw = _sheet(z, "CIW", _ciw_tab(z))
    assert any(str(x.get("Address Line 1") or "").startswith("15 Made Up Rd") for x in ciw), ciw
    # the removal (on LA) is a termination
    assert any(n.startswith("Terminations") for n in z.namelist()), z.namelist()


def test_a_file_in_the_wrong_slot_says_which_slot_it_belongs_in(owner, files):
    c = _person(owner, "pule.ifrs1106@test.local", "Pule Wrongslot")
    for slots, words in (({"la": "steps", "movements": "movements"}, "belongs in 'The step payments report'"),
                         ({"la": "la", "movements": "sequentra"}, "belongs in 'The Sequentra export'"),
                         ({"la": "la", "movements": "movements", "steps": "sequentra"}, "belongs in 'The Sequentra export'"),
                         ({"la": "la", "movements": "movements", "sequentra": "steps"}, "belongs in 'The step payments report'")):
        u = _upload(c, files, slots)
        assert u["status"] == "failed" and words in u["note"], (slots, u["note"])
        assert "Traceback" not in u["note"] and "KeyError" not in u["note"], u["note"]
    rows = c.get("/api/ifrs16").json()["uploads"]
    assert rows, rows


def test_la_vs_rr06_reads_the_charges_under_their_banner_and_the_term_in_years(owner, files):
    c = _person(owner, "thabo.ifrs1106@test.local", "Thabo Recon")
    r = c.post("/api/ifrs16/recon", data={"period": "2026-09"},
               files={"la": (files["la"].name, files["la"].read_bytes(), XLSX), "rr06": (files["movements"].name, files["movements"].read_bytes(), XLSX)})
    assert r.status_code == 200, r.text
    rc = c.get(f"/api/ifrs16/recon/{r.json()['id']}").json()
    assert rc["status"] == "done", rc["note"]
    assert not any("can't be used" in n for n in rc["notes"]), rc["notes"]
    cat = {x["Lease Reference"]: x["Category"] for x in rc["rr06_not_la"]}
    assert cat["T004P001"] == "A year or less" and cat["T005P001"] == "Zero rental" and cat["T006P001"] == "Monthly", cat
    assert cat["T001P001"] == "Exception" and cat["T009P001"] == "Exception", cat  # T001 starts in the period; T009 is simply missing from LA
    charges = {x["Lease Reference"]: x["Current Charges"] for x in rc["rr06_not_la"]}
    assert charges["T001P001"] == 10500 and charges["T005P001"] == 0, charges  # Charges Old (900) is not a current charge
    assert any(x["Lease Reference"] == "T007P001" for x in rc["la_not_rr06"])


# ---- the readers on their own ------------------------------------------------------------------------------------
def test_text_numbers_dates_and_serials():
    import lease_engine as le

    for text, want in (("1 234 567.89", 1234567.89), ("1,234,567.89", 1234567.89), ("1.234.567,89", 1234567.89), ("1 234,56", 1234.56),
                       ("R 12 500", 12500), ("(1 500.00)", -1500), ("7.5%", 7.5), ("12,500", 12500), ("12,5", 12.5), ("  3 000 ", 3000),
                       ("3006P001", None), ("abc", None), ("", None), ("1.2.3", None)):
        assert le._num(text) == (pytest.approx(want) if want is not None else None), text
    from datetime import date

    for v, want in ((45809, date(2025, 6, 1)), (45809.0, date(2025, 6, 1)), ("45809", date(2025, 6, 1)), ("2026-07-01 00:00:00", date(2026, 7, 1)),
                    ("01/07/2026", date(2026, 7, 1)), ("2026/07/01", date(2026, 7, 1)), ("12/25/2026", date(2026, 12, 25)), ("Monthly", None),
                    (9716, None), (None, None)):
        assert le._to_date(v) == want, v
    import pandas as pd

    assert le._to_date(pd.NaT) is None and le.month_token(pd.NaT) is None
    assert le.month_token(datetime(2026, 9, 14)) == date(2026, 9, 1)


def test_title_blocks_csv_and_an_xlsx_with_an_xls_name(tmp_path, files):
    import lease_engine as le

    steps = le.load_step_payments(str(files["steps"]))
    assert len(steps) == 7 and list(steps.columns)[:2] == ["Schedule Number", "Expense Type Name"]
    seq = le.load_sequentra(str(files["sequentra"]))
    row = seq.set_index("Record ID").loc["T001P001"]
    assert row["Address 1"] == "15 Made Up Rd" and row["Address 2"] == "Sandton" and str(row["Postal Code"]) == "0200"
    la = le.load_la_register(str(files["la"]))  # an .xlsx saved with an .xls name
    assert "T002P001" in set(la["Transaction #"]) and le._to_date(la.set_index("Transaction #").loc["T002P001", "Start Date"]).year == 2023
    mv, _ = le.load_movements(str(files["movements"]))
    assert "Move Type" in mv.columns and set(mv["Move Type"].str.strip()) == {"Addition", "Renewal", "Removal"}
    assert mv.attrs["sheet_name"] == "Contract Movement" and mv.attrs["month_source"] == "Movement"
    csv = tmp_path / "steps.csv"
    csv.write_text("Made up step payments\n\nSchedule Number;Expense Type Name;StartDate;EndDate;Number of Payments;Payment Amount;Payment Frequency\n"
                   "T001P001;Base Rent;2026-09-01;2027-08-31;12;10 000,00;Monthly\n", encoding="utf-8")
    got = le.load_step_payments(str(csv))
    assert len(got) == 1 and le._num(got.iloc[0]["Payment Amount"]) == 10000


def test_a_ledger_export_with_a_title_block_and_text_amounts(tmp_path, files):
    import ledger_check as lc

    wb = Workbook()
    ws = wb.active
    ws.title = "Ledger Entries"
    ws.append(["Ledger Export - made up"])
    ws.append([])
    ws.append(["Ledger Date", "Segment2", "Segment3", "Segment4", "DR", "CR", "Comments", "JEShortDesc", "Transactional Currency"])
    ws.append(["2026-06-30", "BC000", "P1", "CC1", "12 000,00", None, "Schedule T001P001 initial", "Initial Recognition", "ZAR"])
    ws.append(["2026-06-30", "HA101", "P1", "CC1", None, "12 000.00", "Schedule T001P001 initial", "Initial Recognition", "ZAR"])
    p = tmp_path / "ledger.xlsx"
    wb.save(p)
    led = lc.load_ledger_export(p)
    assert list(led["DR"]) == [12000, 0] and list(led["CR"]) == [0, 12000] and set(led["Lease Reference"]) == {"T001P001"}
    with pytest.raises(ValueError) as e:
        lc.load_ledger_export(files["sequentra"])
    assert "belongs in 'The Sequentra export'" in str(e.value) and "Worksheet named" not in str(e.value)


def test_a_book_is_read_once_per_file_and_a_big_report_is_quick(tmp_path, monkeypatch):
    import lease_engine as le

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Sheet0")
    ws.append(["DRAFT PIW Step Payments"])
    ws.append([])
    ws.append(["Schedule Number", "Expense Type Name", "StartDate", "EndDate", "Number of Payments", "Payment Amount", "Payment Frequency"])
    for i in range(20000):
        ws.append([f"B{i % 4000:04d}P001", "Base Rent", D(2026, 1, 1), D(2026, 12, 31), 12, 1000 + i, "Monthly"])
    p = tmp_path / "big_steps.xlsx"
    wb.save(p)
    t = time.time()
    assert len(le.load_step_payments(str(p))) == 20000
    assert time.time() - t < 30
    calls = []
    real = le._read_book
    monkeypatch.setattr(le, "_read_book", lambda *a, **k: calls.append(a) or real(*a, **k))
    le.load_step_payments(str(p))
    le.load_step_payments(str(p))
    assert calls == [], "the second and third reads come from the kept book"
    other = tmp_path / "other.xlsx"
    other.write_bytes(p.read_bytes())
    le.load_step_payments(str(other))
    assert len(calls) == 1
