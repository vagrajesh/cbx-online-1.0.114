"""1.0.104 (SPEC-1104 Phase 2, Builder D): the Valuence AP controls V01 and V09.

V01 — new SAP vendors are reviewed before their first invoice posts: the LFA1 reader keeps ERDAT / ERNAM / AEDAT; an
upload of SAP's vendor list opens a review for each vendor created or changed from the go-live day (the shipped seed has
no dates and opens none); Tess stops an invoice to an unreviewed vendor (red, or yellow by the portfolio's setting; a
failed review always red); the reviewer is never SAP's ERNAM nor who attached the creation request; Pass needs the
checklist and the files; the report counts the errors by type per month.

V09 — the duplicates register: every duplicate Tess (batch, upload, SAP) and Bo raise is a case; valid lifts the red
(Tess says who cleared it), genuine never posts; a valid release at or above the review threshold waits for a second
person; the near-duplicate utility rule; the register's age bands and KPI; the backlog import with proposals, bulk
confirm and the conclusions' export. Vendors, accounts, invoices and amounts are made up."""
from __future__ import annotations

import io
import random
import uuid

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func

from cbx import knowledge as K
from cbx.ap import capture as C
from cbx.ap import dupcases as DC
from cbx.ap import history as H
from cbx.ap import masters as M
from cbx.ap import vendor_review as VR
from cbx.db import SessionLocal
from cbx.models import AuditLog, User
from cbx.models.ap import (ApBatch, ApCapture, ApDupCase, ApDupImport, ApEntity, ApItem, ApPortfolio, ApPosted, ApRole, ApSapVendor, ApUtilityAccount,
                           ApVendorCode, ApVendorReview)

from .fixtures import make_ap
from .test_ap_1102 import FULL, _answer, _code, _run, _vat
from .test_ap_capture_1098 import _admin, _app, _chart_of_accounts, _drop, _sheet, _user, eskom_pdf, signed_in  # noqa: F401

TAG = "v1104 D test"
MINE: dict[str, set] = {"codes": set(), "posted": set(), "sapv": set()}


@pytest.fixture(autouse=True)
def _clean():
    """X143's full name taught (a sure bill-to match). Afterwards: no review, case, listing, role, test account, SAP
    posting or vendor coding of these tests is left, the portfolios' settings are as they were, and nothing these tests
    dropped waits in the readers' queues."""
    with SessionLocal() as s:
        e = s.query(ApEntity).filter(ApEntity.code == "X143").one()
        if FULL not in (e.aliases or ""):
            e.aliases = (e.aliases + "\n" if e.aliases else "") + FULL
        s.commit()
        first = {"rev": (s.query(func.max(ApVendorReview.id)).scalar() or 0), "case": (s.query(func.max(ApDupCase.id)).scalar() or 0),
                 "imp": (s.query(func.max(ApDupImport.id)).scalar() or 0), "cap": (s.query(func.max(ApCapture.id)).scalar() or 0) + 1}
        before = {p.code: dict(p.settings or {}) for p in s.query(ApPortfolio).all()}
    K.forget()
    yield
    with SessionLocal() as s:
        s.query(ApCapture).filter(ApCapture.id >= first["cap"], ApCapture.stage.in_(C.READING)).update(
            {"bucket": "skipped", "stage": "skipped", "needs": "", "claimed_at": None}, synchronize_session=False)
        s.query(ApVendorReview).filter(ApVendorReview.id > first["rev"]).delete(synchronize_session=False)
        s.query(ApDupCase).filter(ApDupCase.id > first["case"]).delete(synchronize_session=False)
        s.query(ApDupImport).filter(ApDupImport.id > first["imp"]).delete(synchronize_session=False)
        s.query(ApRole).delete()
        for p in s.query(ApPortfolio).all():
            if p.code in before:
                p.settings = before[p.code]
        s.query(ApUtilityAccount).filter(ApUtilityAccount.updated_by == TAG).delete(synchronize_session=False)
        s.query(ApVendorCode).filter(ApVendorCode.vendor_no.in_(sorted(MINE["codes"]) or ["-"])).delete(synchronize_session=False)
        s.query(ApSapVendor).filter(ApSapVendor.vendor.in_(sorted(MINE["sapv"]) or ["-"])).delete(synchronize_session=False)
        for co, doc in MINE["posted"]:
            s.query(ApPosted).filter(ApPosted.company_code == co, ApPosted.doc_number == doc).delete(synchronize_session=False)
        s.commit()
    for k in MINE:
        MINE[k].clear()
    K.forget()


def _xlsx(rows: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _acct() -> str:
    no = "51" + "".join(random.choice("0123456789") for _ in range(8))
    with SessionLocal() as s:
        s.add(ApUtilityAccount(account_key=M.acct_key(no), account=no, vendor="700999", cost_centre="01306", site="Test Site " + no, municipality="ESKOM",
                               company_code="", source="excel", updated_by=TAG))
        s.commit()
    return no


def _rev(email="dvr.rev@test.local", name="Rita Reviewer"):
    return signed_in(email, "user", name=name, tools=["ap"])


def _appr():
    return signed_in("dvr.appr@test.local", "user", name="Adam Approver", tools=["ap"])


def _uid(email: str) -> int:
    with SessionLocal() as s:
        return s.query(User).filter(User.email == email).one().id


def _review(vendor: str, *, kind: str = "new", created_by: str = "JSMITH", day: str = "2026-10-20", portfolio: str = "EFM-SA") -> int:
    with SessionLocal() as s:
        VR.open_reviews(s, None, [{"vendor": vendor, "name": f"Review Test {vendor}", "kind": kind, "stamp": day, "created_on": day if kind == "new" else "2019-01-01",
                                   "created_by": created_by, "changed_on": day if kind == "changed" else "", "portfolio": portfolio, "company_codes": "X143",
                                   "snapshot": {"name": f"Review Test {vendor}", "vat": "", "payment_terms": "Z030"}, "changes": []}], source="excel")
        s.commit()
        return s.query(ApVendorReview.id).filter(ApVendorReview.vendor == vendor).order_by(ApVendorReview.id.desc()).first()[0]


def _setting(code: str, **kw) -> None:
    with SessionLocal() as s:
        p = s.query(ApPortfolio).filter(ApPortfolio.code == code).one()
        p.settings = {**(p.settings or {}), **kw}
        s.commit()


def _retest(bid: int) -> None:
    with SessionLocal() as s:
        C.test_batch(s, s.get(ApBatch, bid))
        s.commit()


def _texts(doc: dict, level: str = "") -> list[str]:
    return [x["text"] for x in doc["issues"] if not level or x["level"] == level]


def _hist(item_id: int) -> list[str]:
    with SessionLocal() as s:
        return [h["text"] for h in s.get(ApItem, item_id).history or []]


# ---- V01 --------------------------------------------------------------------------------------------------------------------------
def test_the_lfa1_reader_keeps_created_and_changed_and_the_seed_opens_no_review():
    rows = M.read_sap_vendors(_xlsx([["LIFNR", "NAME1", "STCEG", "ERDAT", "ERNAM", "AEDAT"],
                                     ["991001", "New One (Pty) Ltd", "4123412341", "20261020", "jsmith", ""],
                                     ["991002", "Old One CC", "", "2019-03-01", "OLDUSER", "21.10.2026"],
                                     ["991003", "Older Ltd", "", "01.01.2015", "X", "00000000"]]), "lfa1.xlsx")
    assert rows.mapping["created_on"] == "ERDAT" and rows.mapping["created_by"] == "ERNAM" and rows.mapping["changed_on"] == "AEDAT"
    recs = {r["vendor"]: r for r in (M._sap_vendor_row(x) for x in rows)}
    assert (recs["991001"]["created_on"], recs["991001"]["created_by"], recs["991001"]["changed_on"]) == ("2026-10-20", "JSMITH", None)
    assert recs["991002"]["changed_on"] == "2026-10-21" and recs["991003"]["created_on"] == "2015-01-01" and recs["991003"]["changed_on"] is None
    # a list without the columns says nothing about them (the table keeps what it has); SAP's English labels are read too
    assert M._sap_vendor_row({"vendor": "1", "name": "x"})["created_on"] is None
    lab = M.read_sap_vendors(_xlsx([["Supplier", "Name", "Created On", "Created By", "Changed On"], ["991004", "Lab", "2026-10-22", "AB", ""]]), "x.xlsx")
    assert M._sap_vendor_row(lab[0])["created_on"] == "2026-10-22"
    # the shipped seed has no dates: the seeded master opens no review
    with SessionLocal() as s:
        assert s.query(ApSapVendor).filter(ApSapVendor.source == "seed", ApSapVendor.created_on.isnot(None)).count() == 0
        assert s.query(ApVendorReview).filter(ApVendorReview.source == "seed").count() == 0


def test_an_upload_opens_reviews_for_vendors_created_or_changed_from_the_go_live_day():
    _admin()
    sheet = _xlsx([["LIFNR", "NAME1", "STCEG", "ERDAT", "ERNAM", "AEDAT"],
                   ["991001", "New One (Pty) Ltd", "4123412341", "20261020", "JSMITH", ""],
                   ["991002", "Old One CC", "", "20190301", "OLDUSER", "20261021"],
                   ["991003", "Older Ltd", "", "20150101", "X", ""],
                   ["991005", "Before Go-live", "", "20261018", "X", ""]])
    with SessionLocal() as s:
        admin = s.query(User).filter(User.email == "cap.admin@test.local").one()
        rows = M.read_sap_vendors(sheet, "lfa1.xlsx")
        out = M.replace_sap_vendors(s, admin, rows, source="excel", confirm=False)
        assert out["reviews"]["due"] == 2 and out["reviews"]["new"] == 1 and out["reviews"]["changed"] == 1 and out["reviews"]["dated"] is True
        assert s.query(ApVendorReview).filter(ApVendorReview.vendor.in_(["991001", "991002"])).count() == 0  # the preview opens nothing
        # saved (inside a transaction rolled back below: the shared list stays as it is for the other tests)
        out = M.replace_sap_vendors(s, admin, rows, source="excel", confirm=True)
        assert out["saved"] and out["reviews"]["opened"] == 2
        got = {r.vendor: r for r in s.query(ApVendorReview).filter(ApVendorReview.vendor.in_(["991001", "991002", "991003", "991005"])).all()}
        assert set(got) == {"991001", "991002"}
        assert (got["991001"].kind, got["991001"].created_by, got["991001"].stamp, got["991001"].state) == ("new", "JSMITH", "2026-10-20", "open")
        assert (got["991002"].kind, got["991002"].stamp) == ("changed", "2026-10-21") and got["991001"].portfolio
        assert s.query(ApSapVendor).filter(ApSapVendor.vendor == "991001").one().created_on == "2026-10-20"
        # the same list again: nothing new to review
        again = M.replace_sap_vendors(s, admin, rows, source="excel", confirm=False)
        assert again["reviews"]["due"] == 0
        # a list with no created-on date says so
        plain = M.replace_sap_vendors(s, admin, M.read_sap_vendors(_xlsx([["LIFNR", "NAME1"], ["991001", "New One"]]), "x.xlsx"), source="excel", confirm=False)
        assert plain["reviews"]["due"] == 0 and "ERDAT" in plain["reviews"]["note"]
        s.rollback()
    K.forget()


def test_tess_stops_an_unreviewed_vendor_and_an_independent_reviewer_passes_it(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    vat = _vat()
    name = f"Review Me {run} (Pty) Ltd"
    vendor = _code(name, vat)
    MINE["codes"].add(vendor)
    rid = _review(vendor, created_by="JSMITH")
    u = _user()
    b = _run(monkeypatch, u, {f"rv_{run}.pdf": _answer(f"RV-{run}", name, vat)}, f"Rev {run}")
    doc = _sheet(u, b["id"])[f"rv_{run}.pdf"]
    want = f"vendor {vendor} not yet independently reviewed (created 20.10.2026)"
    assert doc["status"] == "red" and want in _texts(doc, "red"), doc["issues"]
    # the portfolio may let it through yellow
    _setting("EFM-SA", unreviewed_vendor_severity="yellow")
    _retest(b["id"])
    doc = _sheet(u, b["id"])[f"rv_{run}.pdf"]
    assert want in _texts(doc, "yellow") and want not in _texts(doc, "red")
    _setting("EFM-SA", unreviewed_vendor_severity="red")
    _retest(b["id"])
    # Needs you: one line for a reviewer, into Controls → New vendors
    rev, rev2 = _rev(), _rev("dvr.rev2@test.local", "Ravi Second")
    line = next(x for x in rev.get("/api/ap/needs").json()["controls"] if x["key"] == "new_vendors")
    assert line["sub"] == "new_vendors" and "independent review" in line["text"]
    assert any(r["id"] == rid for r in rev.get("/api/ap/vendor-reviews").json()["rows"])
    # SAP's creator (ERNAM JSMITH) never reviews it
    js = signed_in("jsmith@test.local", "user", name="John Smith", tools=["ap"])
    full = {"name": "ok", "vat": "ok", "company": "ok", "terms": "ok", "method": "ok", "house_bank": "ok", "bank": "ok"}
    r = js.post(f"/api/ap/vendor-reviews/{rid}/decide", json={"outcome": "pass", "checklist": full})
    assert r.status_code == 403 and "SAP says you created this vendor" in r.json()["detail"]
    assert "SAP says you created" in js.get(f"/api/ap/vendor-reviews/{rid}").json()["blockers"][0]
    # nor who attached the creation request
    up = rev.post(f"/api/ap/vendor-reviews/{rid}/files", data={"kind": "request"}, files={"file": ("request.pdf", make_ap.invoice(), "application/pdf")})
    assert up.status_code == 200, up.text
    r = rev.post(f"/api/ap/vendor-reviews/{rid}/decide", json={"outcome": "pass", "checklist": full})
    assert r.status_code == 403 and "attached the creation request" in r.json()["detail"]
    # Pass needs the bank letter (unless the bank is not applicable) and every item ticked
    r = rev2.post(f"/api/ap/vendor-reviews/{rid}/decide", json={"outcome": "pass", "checklist": full})
    assert r.status_code == 400 and "bank confirmation letter" in r.json()["detail"]
    assert rev2.post(f"/api/ap/vendor-reviews/{rid}/files", data={"kind": "letter"}, files={"file": ("bank.pdf", make_ap.bank_letter(), "application/pdf")}).status_code == 200
    r = rev2.post(f"/api/ap/vendor-reviews/{rid}/decide", json={"outcome": "pass", "checklist": {**full, "terms": ""}})
    assert r.status_code == 400 and "Tick every item" in r.json()["detail"]
    r = rev2.post(f"/api/ap/vendor-reviews/{rid}/decide", json={"outcome": "pass", "checklist": full})
    assert r.status_code == 200 and r.json()["state"] == "passed" and r.json()["reviewed_by"] == "Ravi Second"
    got = rev2.get(f"/api/ap/vendor-reviews/{rid}").json()
    assert got["request"]["name"].endswith("request.pdf") and got["letter"] and rev2.get(f"/api/ap/vendor-reviews/{rid}/files/{got['letter']['id']}").status_code == 200
    doc = _sheet(u, b["id"])[f"rv_{run}.pdf"]
    assert want not in _texts(doc)  # Tess tested it again
    assert any(f"Vendor {vendor} passed its independent review" in h for h in _hist(doc["item_id"]))
    with SessionLocal() as s:
        assert s.query(AuditLog).filter(AuditLog.action == "ap_vendor_review", AuditLog.target == f"ap_vendor:{vendor}").count() >= 3  # 2 files + the decision


def test_a_failed_review_is_always_red_and_the_report_counts_errors_by_type(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    vat = _vat()
    name = f"Fail Me {run} (Pty) Ltd"
    vendor = _code(name, vat)
    MINE["codes"].add(vendor)
    rid = _review(vendor, kind="changed", created_by="SOMEONE", day="2026-10-21")
    _setting("EFM-SA", unreviewed_vendor_severity="yellow")  # a failed review stays red whatever the portfolio says
    rev = _rev()
    r = rev.post(f"/api/ap/vendor-reviews/{rid}/decide", json={"outcome": "fail", "checklist": {"bank": "wrong"}})
    assert r.status_code == 400 and "Say why" in r.json()["detail"]
    r = rev.post(f"/api/ap/vendor-reviews/{rid}/decide", json={"outcome": "fail", "checklist": {"name": "ok", "bank": "wrong", "terms": "wrong"}, "reason": "Bank account differs from the letter"})
    assert r.status_code == 200 and r.json()["state"] == "failed" and set(r.json()["errors"]) == {"bank", "terms"}
    u = _user()
    b = _run(monkeypatch, u, {f"fl_{run}.pdf": _answer(f"FL-{run}", name, vat)}, f"Fail {run}")
    doc = _sheet(u, b["id"])[f"fl_{run}.pdf"]
    assert doc["status"] == "red" and any(t.startswith(f"vendor {vendor} failed its independent review on ") and t.endswith("Bank account differs from the letter") for t in _texts(doc, "red"))
    wb = load_workbook(io.BytesIO(rev.get("/api/ap/vendor-reviews/report").content))
    assert wb.sheetnames == ["Review log", "Errors by month"]
    log_rows = [r for r in wb["Review log"].iter_rows(values_only=True) if r[0] == vendor]
    assert log_rows and log_rows[0][2] == "changed" and log_rows[0][10] == "failed"
    heads = [c.value for c in wb["Errors by month"][1]]
    last = [c.value for c in list(wb["Errors by month"].iter_rows())[-1]]
    assert last[heads.index("Failed")] >= 1 and last[heads.index("Errors: bank")] >= 1 and last[heads.index("Errors: terms")] >= 1
    # the registers are for reviewers and admins
    viewer = signed_in("dvr.viewer@test.local", "viewer", name="Vic Viewer", tools=["ap"])
    assert viewer.get("/api/ap/vendor-reviews").status_code == 403 and viewer.get("/api/ap/dups").status_code == 403


# ---- V09 --------------------------------------------------------------------------------------------------------------------------
def _case(**where) -> ApDupCase:
    with SessionLocal() as s:
        q = s.query(ApDupCase)
        for k, v in where.items():
            q = q.filter(getattr(ApDupCase, k) == v)
        c = q.order_by(ApDupCase.id.desc()).first()
        s.expunge_all()
        return c


def _cap_of(u, bid: int, name: str) -> dict:
    return _sheet(u, bid)[name]


def test_a_duplicate_tess_raises_is_a_case_valid_lifts_the_red_and_genuine_never_posts():
    u, rev = _user(), _rev()
    acct, inv = _acct(), str(random.randint(700000000, 799999999))
    b = _drop(u, {"first.pdf": eskom_pdf(acct, inv), "second.pdf": eskom_pdf(acct, inv)}, f"Dup {inv}")
    first, second = _cap_of(u, b["id"], "first.pdf"), _cap_of(u, b["id"], "second.pdf")
    assert any("in this batch" in t for t in _texts(second, "red")) and not any("in this batch" in t for t in _texts(first))
    case = _case(capture_id=second["id"], source="tess")
    assert case.code == "batch" and case.twin_id == first["id"] and case.twin_ref == "first.pdf" and case.state == "open" and case.amount == pytest.approx(11512.34)
    # Tess again: the same case, not a second one
    _retest(b["id"])
    with SessionLocal() as s:
        assert s.query(ApDupCase).filter(ApDupCase.capture_id == second["id"], ApDupCase.source == "tess").count() == 1
    assert rev.post(f"/api/ap/dups/{case.id}/decide", json={"outcome": "valid", "reason": ""}).status_code == 400
    r = rev.post(f"/api/ap/dups/{case.id}/decide", json={"outcome": "valid", "reason": "Eskom re-issued the bill under a new cover"})
    assert r.status_code == 200 and r.json()["state"] == "valid"  # R 11 512.34 is under the R 50 000 review threshold
    second = _cap_of(u, b["id"], "second.pdf")
    assert not any("in this batch" in t for t in _texts(second, "red"))
    assert any(t.startswith("cleared as not a duplicate by Rita Reviewer") and t.endswith("Eskom re-issued the bill under a new cover") for t in _texts(second, "note"))
    assert any("Cleared as not a duplicate (Rita Reviewer)" in h for h in _hist(second["item_id"]))
    with SessionLocal() as s:
        assert s.query(AuditLog).filter(AuditLog.action == "ap_dup_classified", AuditLog.target == f"ap_item:{second['item_id']}").count() == 1
    # decided cases are closed
    assert rev.post(f"/api/ap/dups/{case.id}/decide", json={"outcome": "genuine", "reason": "x"}).status_code == 400
    # genuine: the document never posts
    acct2, inv2 = _acct(), str(random.randint(700000000, 799999999))
    b2 = _drop(u, {"one.pdf": eskom_pdf(acct2, inv2), "two.pdf": eskom_pdf(acct2, inv2)}, f"Dup2 {inv2}")
    two = _cap_of(u, b2["id"], "two.pdf")
    c2 = _case(capture_id=two["id"], source="tess")
    r = rev.post(f"/api/ap/dups/{c2.id}/decide", json={"outcome": "genuine", "reason": "The same bill twice"})
    assert r.status_code == 200 and r.json()["state"] == "genuine"
    with SessionLocal() as s:
        assert s.get(ApCapture, two["id"]).stage == "skipped"


def test_a_valid_release_over_the_threshold_needs_a_second_person():
    u, rev, appr = _user(), _rev(), _appr()
    acct, inv = _acct(), str(random.randint(700000000, 799999999))
    b = _drop(u, {"a.pdf": eskom_pdf(acct, inv), "b.pdf": eskom_pdf(acct, inv)}, f"Thr {inv}")
    bdoc = _cap_of(u, b["id"], "b.pdf")
    case = _case(capture_id=bdoc["id"], source="tess")
    _setting(case.portfolio, review_threshold_r=1000)
    r = rev.post(f"/api/ap/dups/{case.id}/decide", json={"outcome": "valid", "reason": "Different period, same number reused"})
    assert r.status_code == 200 and r.json()["state"] == "pending"
    assert any("a second person confirms it" in t for t in _texts(_cap_of(u, b["id"], "b.pdf"), "red"))  # still red until then
    assert any(x["key"] == "duplicates" for x in appr.get("/api/ap/needs").json()["controls"])
    r = rev.post(f"/api/ap/dups/{case.id}/second", json={"agree": True})
    assert r.status_code == 403 and "different person" in r.json()["detail"]
    r = appr.post(f"/api/ap/dups/{case.id}/second", json={"agree": False})
    assert r.status_code == 400  # sending it back needs a reason
    r = appr.post(f"/api/ap/dups/{case.id}/second", json={"agree": True, "reason": "Checked the period"})
    assert r.status_code == 200 and r.json()["state"] == "valid" and r.json()["second_by"] == "Adam Approver"
    bdoc = _cap_of(u, b["id"], "b.pdf")
    assert not any("in this batch" in t for t in _texts(bdoc, "red")) and any("Rita Reviewer and Adam Approver" in t for t in _texts(bdoc, "note"))


def test_a_near_duplicate_utility_bill_is_yellow_with_a_case():
    u, rev = _user(), _rev()
    acct = _acct()
    inv = random.randint(700000000, 799999999)
    b = _drop(u, {"bill.pdf": eskom_pdf(acct, str(inv)), "reissued.pdf": eskom_pdf(acct, str(inv + 7))}, f"Near {inv}")
    doc = _cap_of(u, b["id"], "reissued.pdf")
    near = [t for t in _texts(doc, "yellow") if t.startswith("possible duplicate: account")]
    assert near and "for SEPTEMBER 2026, R 11 512.34" in near[0] and "like bill.pdf" in near[0], doc["issues"]
    assert not any(t.startswith("possible duplicate") for t in _texts(_cap_of(u, b["id"], "bill.pdf")))  # only the later one
    case = _case(capture_id=doc["id"], source="near")
    assert case.code == "near" and case.twin_ref == "bill.pdf"
    assert rev.post(f"/api/ap/dups/{case.id}/decide", json={"outcome": "valid", "reason": "Two meters on one account"}).status_code == 200
    doc = _cap_of(u, b["id"], "reissued.pdf")
    assert not any(t.startswith("possible duplicate") for t in _texts(doc, "yellow")) and any("Two meters on one account" in t for t in _texts(doc, "note"))


def test_bo_s_twice_in_a_run_is_a_case_and_a_valid_one_no_longer_stops():
    u, rev = _user(), _rev()
    b = _drop(u, {f"run_{uuid.uuid4().hex[:6]}.xlsx": make_ap.payment_run()}, "Run")
    it = b["items"][0]
    case = _case(item_id=it["id"], code="duplicate_in_run")
    assert case.source == "bo" and case.state == "open" and case.row is not None and case.amount == pytest.approx(11500.0)
    with SessionLocal() as s:
        stops = s.get(ApItem, it["id"]).stops
    r = rev.post(f"/api/ap/dups/{case.id}/decide", json={"outcome": "valid", "reason": "Two invoices share the reference"})
    assert r.status_code == 200 and r.json()["state"] == "valid"
    with SessionLocal() as s:
        x = s.get(ApItem, it["id"])
        f = next(f for f in x.findings if f.get("code") == "duplicate_in_run")
        assert f["severity"] == "info" and "Cleared as not a duplicate by Rita Reviewer" in f["text"] and x.stops == stops - 1


def test_the_register_ages_cases_counts_them_on_the_board_and_needs_the_role():
    u, rev = _user(), _rev()
    acct, inv = _acct(), str(random.randint(700000000, 799999999))
    b = _drop(u, {"x.pdf": eskom_pdf(acct, inv), "y.pdf": eskom_pdf(acct, inv)}, f"Reg {inv}")
    case = _case(capture_id=_cap_of(u, b["id"], "y.pdf")["id"], source="tess")
    got = rev.get("/api/ap/dups").json()
    assert got["ages"] == ["0–30 days", "31–60 days", "60+ days"] and any(r["id"] == case.id and r["age"] == "0–30 days" for r in got["rows"])
    band = next(x for x in got["bands"] if x["portfolio"] == case.portfolio)
    assert band["0–30 days"] >= 1 and band["total"] >= 1
    assert u.get("/api/ap").json()["board"]["dup_cases_open"] >= 1
    # an owner per case
    assert rev.post(f"/api/ap/dups/{case.id}/take", json={}).json()["owner"] == "Rita Reviewer"
    assert any(o["owner"] == "Rita Reviewer" for o in rev.get("/api/ap/dups").json()["owners"])
    # once roles are given, only a reviewer decides; the person who typed the document never does
    _appr()
    with SessionLocal() as s:
        s.add(ApRole(user_id=_uid("dvr.appr@test.local"), portfolio="", role="approver", granted_by="test"))
        s.commit()
    r = rev.post(f"/api/ap/dups/{case.id}/decide", json={"outcome": "valid", "reason": "x"})
    assert r.status_code == 403 and "reviewer role" in r.json()["detail"]
    with SessionLocal() as s:
        s.query(ApRole).delete()
        s.commit()
    with SessionLocal() as s:  # someone who corrected the document
        s.add(AuditLog(action="ap_capture_save", actor_id=_uid("dvr.rev@test.local"), target=f"ap_item:{case.item_id}", detail="typed"))
        s.commit()
    r = rev.post(f"/api/ap/dups/{case.id}/decide", json={"outcome": "valid", "reason": "x"})
    assert r.status_code == 403 and "worked on this document" in r.json()["detail"]


def test_the_backlog_listing_is_matched_proposed_confirmed_in_bulk_and_exported():
    rev = _rev()
    run = uuid.uuid4().hex[:6].upper()
    vendor = "88" + str(uuid.uuid4().int)[:5]
    doc = "97" + str(uuid.uuid4().int)[:7]
    with SessionLocal() as s:
        s.add(ApPosted(company_code="X143", vendor=vendor, reference=f"MOL-{run}", ref_key=H.ref_key(f"MOL-{run}"), doc_type="KR", doc_number=doc,
                       posting_date="2026-09-01", doc_date="2026-09-01", amount=1000.0, currency="ZAR", source="excel"))
        s.commit()
    MINE["posted"].add(("X143", doc))
    listing = _xlsx([["Workflow Basket - possible duplicates"], [],
                     ["Vendor", "Vendor name", "Reference", "Amount", "Document date", "Company code"],
                     [vendor, "Listing Supplier", f"MOL {run}", 1000.0, "2026-09-02", "X143"],
                     [vendor, "Listing Supplier", f"NEW-{run}", 500.0, "2026-09-03", "X143"],
                     [vendor, "Listing Supplier", f"NEW-{run}", 500.0, "2026-09-04", "X143"],
                     ["", "No Vendor", "X-1", 5.0, "2026-09-05", "X143"]])
    r = rev.post("/api/ap/dups/import", files={"file": ("basket.xlsx", listing, "application/octet-stream")}, data={"portfolio": "EFM-SA", "confirm": "false"})
    assert r.status_code == 200, r.text
    p = r.json()
    assert (p["lines"], p["genuine"], p["valid"], p["look"], p["saved"]) == (4, 2, 1, 1, False)
    assert p["mapping"]["reference"] == "Reference" and p["sample"][0]["why"].startswith(f"already posted in SAP: doc {doc}")
    assert "line 2 of this listing" in p["sample"][2]["why"]
    r = rev.post("/api/ap/dups/import", files={"file": ("basket.xlsx", listing, "application/octet-stream")}, data={"portfolio": "EFM-SA", "confirm": "true"})
    imp = r.json()["import_id"]
    rows = rev.get(f"/api/ap/dups?import_id={imp}").json()["rows"]
    assert len(rows) == 4 and {x["proposed"] for x in rows} == {"genuine", "valid", ""} and rows[0]["twin_kind"] == "sap"
    out = rev.post("/api/ap/dups/confirm", json={"ids": [x["id"] for x in rows]}).json()
    assert out["done"] == 3 and out["skipped_n"] == 1
    states = {x["line"]["reference"]: x["state"] for x in rev.get(f"/api/ap/dups?import_id={imp}&state=all").json()["rows"]}
    assert states[f"MOL {run}"] == "genuine" and states["X-1"] == "open"
    wb = load_workbook(io.BytesIO(rev.get(f"/api/ap/dups/export?import_id={imp}").content))
    vals = [r for r in wb.active.iter_rows(values_only=True)]
    heads = list(vals[0])
    first = next(r for r in vals[1:] if r[heads.index("Reference")] == f"MOL {run}")
    assert first[heads.index("Outcome")] == "Genuine duplicate: do not pay" and first[heads.index("Reason")].startswith("As proposed: already posted in SAP")
    with SessionLocal() as s:
        assert s.query(AuditLog).filter(AuditLog.action == "ap_dup_classified", AuditLog.target.like("ap_dup_case:%")).count() >= 3


def test_a_new_vendor_s_first_invoice_is_known_for_the_review_selection():
    vendor = "93" + str(uuid.uuid4().int)[:5]
    with SessionLocal() as s:
        assert VR.first_invoice(s, vendor) is False  # no review: not a new vendor
    _review(vendor)
    with SessionLocal() as s:
        assert VR.first_invoice(s, vendor) is True and VR.first_invoice(s, "00" + vendor) is True
    doc = "96" + str(uuid.uuid4().int)[:7]
    with SessionLocal() as s:
        s.add(ApPosted(company_code="X143", vendor=vendor, reference="FIRST-1", ref_key="FIRST1", doc_type="KR", doc_number=doc, posting_date="2026-10-20", amount=10.0, source="excel"))
        s.commit()
    MINE["posted"].add(("X143", doc))
    with SessionLocal() as s:
        assert VR.first_invoice(s, vendor) is False  # SAP posted one already


# ---- 1.0.104 review findings ---------------------------------------------------------------------------------------------------
def test_the_person_who_dropped_the_folder_never_settles_its_duplicate():
    u, rev = _user(), _rev()
    acct, inv = _acct(), str(random.randint(700000000, 799999999))
    b = _drop(u, {"p.pdf": eskom_pdf(acct, inv), "q.pdf": eskom_pdf(acct, inv)}, f"Own {inv}")
    case = _case(capture_id=_cap_of(u, b["id"], "q.pdf")["id"], source="tess")
    r = u.post(f"/api/ap/dups/{case.id}/decide", json={"outcome": "valid", "reason": "mine, honestly"})
    assert r.status_code == 403 and "dropped its folder" in r.json()["detail"]
    _setting(case.portfolio, review_threshold_r=1000)  # nor confirms it as the second person
    assert rev.post(f"/api/ap/dups/{case.id}/decide", json={"outcome": "valid", "reason": "Checked"}).json()["state"] == "pending"
    r = u.post(f"/api/ap/dups/{case.id}/second", json={"agree": True})
    assert r.status_code == 403 and "dropped its folder" in r.json()["detail"]


def test_the_near_duplicate_rule_reads_the_account_s_bill_history():
    from cbx.models.ap import ApBillHistory

    u = _user()
    acct = _acct()
    inv = random.randint(700000000, 799999999)
    b = _drop(u, {"h1.pdf": eskom_pdf(acct, str(inv)), "h2.pdf": eskom_pdf(acct, str(inv + 3))}, f"Hist {inv}")
    h1, h2 = _cap_of(u, b["id"], "h1.pdf"), _cap_of(u, b["id"], "h2.pdf")
    assert any(t.startswith("possible duplicate") for t in _texts(h2, "yellow"))
    with SessionLocal() as s:  # the twin is found through its history row (no history row: no near-duplicate)
        assert s.query(ApBillHistory).filter(ApBillHistory.capture_id == h1["id"], ApBillHistory.account_key == M.acct_key(acct)).count() == 1
        g = s.get(ApCapture, h2["id"])
        assert DC.near_twin(s, g)[0].id == h1["id"]
        s.query(ApBillHistory).filter(ApBillHistory.capture_id == h1["id"]).update({"account_key": "X" + M.acct_key(acct)}, synchronize_session=False)
        assert DC.near_twin(s, g) is None
        s.rollback()


def test_a_review_s_files_need_a_role_in_its_own_portfolio():
    rev, rev2 = _rev(), _rev("dvr.rev2@test.local", "Ravi Second")
    rid = _review("94" + str(uuid.uuid4().int)[:5], portfolio="EFM-SA")
    up = rev.post(f"/api/ap/vendor-reviews/{rid}/files", data={"kind": "letter"}, files={"file": ("bank.pdf", make_ap.bank_letter(), "application/pdf")})
    assert up.status_code == 200, up.text
    with SessionLocal() as s:
        s.add(ApRole(user_id=_uid("dvr.rev@test.local"), portfolio="EFM-SA", role="reviewer", granted_by="test"))
        s.add(ApRole(user_id=_uid("dvr.rev2@test.local"), portfolio="SHELL", role="reviewer", granted_by="test"))
        s.commit()
    url = f"/api/ap/vendor-reviews/{rid}/files/{up.json()['id']}"
    assert rev.get(url).status_code == 200
    assert rev2.get(url).status_code == 404  # a reviewer in Shell doesn't see EFM-SA's bank letters
    # nor the review itself, nor it in the list
    assert rev.get(f"/api/ap/vendor-reviews/{rid}").status_code == 200 and rev2.get(f"/api/ap/vendor-reviews/{rid}").status_code == 404
    assert any(r["id"] == rid for r in rev.get("/api/ap/vendor-reviews").json()["rows"])
    assert not any(r["id"] == rid for r in rev2.get("/api/ap/vendor-reviews?state=all").json()["rows"])
    assert _admin().get(f"/api/ap/vendor-reviews/{rid}").status_code == 200  # admins see every portfolio's
