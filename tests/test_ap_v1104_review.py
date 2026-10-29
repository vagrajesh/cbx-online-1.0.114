"""1.0.104 (SPEC-1104 Phase 2, Builder R): the Valuence AP controls V05, V06 and V12.

V05 — documents are picked for an independent review at Tess time (amount, typed or coded by hand, yellow, tax code, not
a valid tax invoice, first invoice to a vendor, a repeatable random sample); a reviewer who did not drop or touch the
document agrees, corrects or rejects it; the SAP upload needs the poster role and waits while a picked document is not
reviewed; the review workbook and the monthly coverage report carry the outcome.
V06 — an invoice that is not a valid tax invoice (s20) is red until a reviewer decides: Claim VAT, or Don't claim (the
no-claim tax code, the VAT folded into the expense lines, no JA600 line); the supplier report.
V12 — every download files the upload and the review workbook with their SHA-256; evidence is attached on a folder, a
payment run or an account (.eml too); the checklist and Needs you; the zip export with its index; retention.
Invoices, accounts and people are made up."""
from __future__ import annotations

import hashlib
import io
import random
import uuid
import zipfile
from datetime import timedelta

import pytest
from openpyxl import load_workbook

from cbx.ap import capture as C
from cbx.ap import masters as M
from cbx.ap import portfolio as PF
from cbx.ap import review as RV
from cbx.ap import sap as P
from cbx.ap import vatcheck as VC
from cbx.db import SessionLocal, utcnow
from cbx.models import AuditLog, StoredFile, User
from cbx.models.ap import ApBatch, ApCapture, ApEntity, ApEvidence, ApItem, ApPortfolio, ApPosted, ApRole, ApUtilityAccount, ApVendorCode
from cbx.warehouse import storage

from .fixtures import make_ap
from .test_ap_capture_1098 import _admin, _app, _chart_of_accounts, _drop, _invoice_answer, _sheet, _user, _vision, eskom_pdf, signed_in  # noqa: F401

VENDOR_VAT, VENDOR_NO, SUPPLIER = "4555000111", "300555", "Reviewco Services (Pty) Ltd"


@pytest.fixture(autouse=True)
def _clean():
    """No AP role, test vendor, posted line or account is left behind (other tests run in bootstrap)."""
    yield
    with SessionLocal() as s:
        s.query(ApRole).delete()
        s.query(ApPosted).filter(ApPosted.doc_number.like("RVTEST%")).delete(synchronize_session=False)
        s.query(ApUtilityAccount).filter(ApUtilityAccount.updated_by == "v1104 review test").delete()
        s.commit()


def _uid(email: str) -> int:
    with SessionLocal() as s:
        return s.query(User).filter(User.email == email).one().id


def _rene():
    return signed_in("v1104.rev@test.local", "user", name="Rene Reviewer", tools=["ap"])


def _vendor():
    with SessionLocal() as s:
        s.query(ApVendorCode).filter(ApVendorCode.vat_no == VENDOR_VAT).delete()
        s.add(ApVendorCode(vendor_no=VENDOR_NO, name=SUPPLIER, name_key=M.name_key(SUPPLIER), vat_no=VENDOR_VAT, gl="6100100", cost_centre="CC77", tax_code="I9", source="excel"))
        x143 = s.query(ApEntity).filter(ApEntity.code == "X143").one()
        if "CBRE Excellerate Facilities Management" not in x143.aliases:
            x143.aliases = (x143.aliases + "\n" if x143.aliases else "") + "CBRE Excellerate Facilities Management"
        s.commit()


def _answer(no, excl="10000.00", vat="1500.00", incl="11500.00", bill_to_vat=None, words=True):
    a = _invoice_answer(no, supplier=SUPPLIER, vat_no=VENDOR_VAT, excl=excl, vat=vat, incl=incl)
    a["bill_to_vat"], a["tax_invoice_words"] = bill_to_vat, words
    return a


def _drop_invoices(monkeypatch, u, answers: dict, name: str) -> dict:
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision(answers))
    return _drop(u, {n: make_ap.invoice(no=a["invoice_no"], supplier=SUPPLIER, vat_no=VENDOR_VAT) for n, a in answers.items()}, name)


def _acct(vendor: str) -> str:
    no = "51" + "".join(random.choice("0123456789") for _ in range(8))
    with SessionLocal() as s:
        s.add(ApUtilityAccount(account_key=M.acct_key(no), account=no, vendor=vendor, cost_centre="01306", site="Review Site " + no, municipality="ESKOM",
                               company_code="", source="excel", updated_by="v1104 review test"))
        s.commit()
    return no


def _posted(vendor: str) -> None:
    """SAP posted this vendor before: its bills are not a first invoice."""
    with SessionLocal() as s:
        s.add(ApPosted(company_code="T000", vendor=vendor, reference="OLD-1", ref_key="OLD1", doc_type="KR", doc_number=f"RVTEST{random.randint(10**8, 10**9)}",
                       posting_date="2026-09-01", amount=100.0, source="excel"))
        s.commit()


def _cap(c, cid: int) -> dict:
    r = c.get(f"/api/ap/captures/{cid}")
    assert r.status_code == 200, r.text
    return r.json()


# ---- V06: the rules on their own ---------------------------------------------------------------------------------------------
def test_s20_checks_and_folding_the_vat_in():
    wants = {k for k, _w, _s in VC.REQUIREMENTS}
    f = {"kind": "manual", "supplier": "Acme", "supplier_vat": "4123456789", "bill_to": "CBRE", "bill_to_vat": "", "invoice_no": "1", "invoice_date": "2026-10-01",
         "lines": [{"description": "Cleaning", "excl": "10000.00", "vat_rate": 15, "vat": "1500.00"}], "excl": "10000.00", "vat": "1500.00", "incl": "11500.00",
         "currency": "ZAR", "tax_invoice_words": True}
    assert VC.check(f, wants=wants, threshold=5000) == (False, ["no recipient VAT number (needed above R 5 000)"])
    assert VC.check({**f, "bill_to_vat": "4999888777"}, wants=wants, threshold=5000) == (True, [])
    assert VC.check(f, wants=wants - {"recipient_vat"}, threshold=5000) == (True, [])  # a portfolio may drop a requirement
    assert VC.check({**f, "incl": "4600.00"}, wants=wants, threshold=5000)[0] is True  # below the threshold the recipient's number is not needed
    bad = VC.check({**f, "bill_to_vat": "4999888777", "supplier_vat": "123", "tax_invoice_words": False, "lines": [{"description": "", "excl": "10000.00"}]}, wants=wants)
    assert bad[0] is False and bad[1] == ['the words "Tax Invoice" are not on it', "supplier VAT number 123 is not 10 digits starting with 4", "no description of the goods or services"]
    assert VC.check({**f, "tax_invoice_words": None, "bill_to_vat": "4999888777"}, wants=wants)[1] == ['the words "Tax Invoice" are not confirmed']
    # nothing to claim: not checked
    assert VC.check({**f, "vat": "0.00", "lines": []}, wants=wants) == (None, []) and VC.check({**f, "credit": True}, wants=wants) == (None, [])
    assert VC.check({**f, "currency": "EGP"}, wants=wants) == (None, [])
    assert VC.vat_no_ok("4380190845") and not VC.vat_no_ok("438019084") and not VC.vat_no_ok("5380190845")
    # Don't claim: the VAT line goes, the VAT is spread over the lines that carried it, the vendor line stays the total
    tables = {"fixed": dict(P.SAP_DEFAULT_FIXED), "gl": {c: g for c, g, _ in P.DEFAULT_GL}}
    rows, st, _ = P.manual_doc({**f, "lines": [{"description": "A", "excl": "6000.00", "vat_rate": 15, "vat": "900.00"}, {"description": "B", "excl": "4000.00", "vat_rate": 15, "vat": "600.00"}]},
                               {"vendor_no": "300555", "name": "Acme", "gl": "6100100", "cost_centre": "CC77", "tax_code": "I9",
                                "line_rules": [{"words": "B", "gl": "6100200"}]}, "X143", tables, file="x.pdf")
    assert [r["account"] for r in rows] == ["300555", "6100100", "6100200", "JA600"]
    out = VC.fold(rows, tables, "I0")
    assert [(r["pkey"], r["account"], r["amount"], r["tc"]) for r in out] == [("31", "300555", "11500.00", ""), ("40", "6100100", "6900.00", "I0"), ("40", "6100200", "4600.00", "I0")]
    assert P.balance_fails([{**r, "_file": "x"} for r in out]) == (1, set())


def test_the_settings_are_offered_and_the_sample_repeats():
    keys = {k["key"]: k for k in PF.keys_out()}
    assert keys["vat_invalid"]["choices"] == ["red", "no_claim", "off"] and keys["review_before_upload"]["kind"] == "bool"
    assert keys["vat_number_profiles"]["kind"] == "lines" and keys["vat_requirements"]["kind"] == "lines" and keys["evidence_missing_days"]["default"] == 7
    picks = [RV.sampled("abc", i, 10) for i in range(2000)]
    assert picks == [RV.sampled("abc", i, 10) for i in range(2000)] and 120 < sum(picks) < 280  # the same seed: the same sample; about 10 %
    assert not any(RV.sampled("abc", i, 0) for i in range(200)) and all(RV.sampled("abc", i, 100) for i in range(200))


# ---- V06 in the flow -----------------------------------------------------------------------------------------------------------
def test_not_a_valid_tax_invoice_is_red_until_a_reviewer_decides(monkeypatch, ap_review_gates):
    _vendor()
    u, rene, adm = _user(), _rene(), _admin()
    run = uuid.uuid4().hex[:6].upper()
    bad, good = f"RV-{run}-1", f"RV-{run}-2"
    b = _drop_invoices(monkeypatch, u, {"bad.pdf": _answer(bad), "good.pdf": _answer(good, bill_to_vat="4999888777")}, f"VAT {run}")
    docs = _sheet(u, b["id"])
    assert docs["bad.pdf"]["status"] == "red" and docs["bad.pdf"]["vat_valid"] is False, docs["bad.pdf"]
    assert any("not a valid tax invoice: no recipient VAT number (needed above R 5 000)" in x["text"] for x in docs["bad.pdf"]["issues"])
    assert docs["good.pdf"]["status"] == "green" and docs["good.pdf"]["vat_valid"] is True
    cid = docs["bad.pdf"]["id"]
    # the person who dropped the folder doesn't decide; a reason is needed; a viewer-free reviewer decides
    r = u.post(f"/api/ap/captures/{cid}/vat", json={"decision": "no_claim", "reason": "supplier to send a valid tax invoice"})
    assert r.status_code == 403 and "dropped or worked on" in r.json()["detail"]
    assert rene.post(f"/api/ap/captures/{cid}/vat", json={"decision": "no_claim", "reason": ""}).status_code == 400
    assert rene.post(f"/api/ap/captures/{docs['good.pdf']['id']}/vat", json={"decision": "claim", "reason": "x" * 5}).status_code == 400  # valid: nothing to decide
    lines = rene.get("/api/ap/needs").json()["controls"]
    assert any(x["key"] == "vat" and x["sub"] == "review" for x in lines)
    assert any(d["id"] == cid and d["can"] for d in rene.get("/api/ap/review?show=vat").json()["docs"])
    r = rene.post(f"/api/ap/captures/{cid}/vat", json={"decision": "dont claim", "reason": "no recipient VAT number; supplier asked for a new one"})
    assert r.status_code == 200, r.text
    c = _cap(u, cid)
    assert c["sap_status"] == "yellow" and c["vat"]["vat_decision"] == "no_claim" and c["vat"]["vat_decided_by"] == "Rene Reviewer"
    assert [(x["pkey"], x["account"], x["amount"], x["tc"]) for x in c["sap_rows"]] == [("31", VENDOR_NO, "11500.00", ""), ("40", "6100100", "11500.00", "I0")]
    assert any("VAT not claimable: posted with I0" in x["text"] for x in c["sap_issues"])
    assert c["review"]["state"] == "selected" and "not a valid tax invoice" in c["review"]["why"]  # V05 picks it
    assert any("VAT: Don't claim" in h["text"] for h in c["history"])
    with SessionLocal() as s:
        assert s.query(AuditLog).filter(AuditLog.action == "ap_vat_decision", AuditLog.target == f"ap_item:{c['item_id']}").count() == 1
    # changed mind: claim → the VAT line is back, still yellow with who decided
    assert rene.post(f"/api/ap/captures/{cid}/vat", json={"decision": "claim", "reason": "SARS accepts the supplier's letter"}).status_code == 200
    c = _cap(u, cid)
    assert [x["account"] for x in c["sap_rows"]] == [VENDOR_NO, "6100100", "JA600"] and c["sap_rows"][1]["tc"] == "I9"
    assert any("VAT claimed on an invoice that is not a valid tax invoice" in x["text"] and "Rene Reviewer" in x["text"] for x in c["sap_issues"])
    assert rene.post(f"/api/ap/captures/{cid}/vat", json={"decision": "no_claim", "reason": "back to no claim"}).status_code == 200
    # it goes out once reviewed (Rene agrees both); the report lists the supplier with the VAT not claimed
    for d in rene.get(f"/api/ap/review?batch={b['id']}").json()["docs"]:
        assert rene.post(f"/api/ap/review/{d['id']}", json={"action": "agree"}).status_code == 200
    r = u.post(f"/api/ap/batches/{b['id']}/sap/upload", json={"include_yellow": True})
    assert r.status_code == 200, r.text
    month = f"{utcnow():%Y-%m}"
    viewer = signed_in("v1104.plain@test.local", "viewer", name="Vic Viewer", tools=["ap"])
    assert viewer.get(f"/api/ap/reports/vat?month={month}").status_code == 403  # registers and reports are for reviewers (bootstrap: everyone who writes)
    ws = load_workbook(io.BytesIO(adm.get(f"/api/ap/reports/vat?month={month}").content))["Suppliers"]
    row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[1] == VENDOR_VAT)
    assert row[0] == SUPPLIER and row[3] >= 1 and row[4] >= 1500 and "no recipient VAT number" in row[5]
    assert adm.get("/api/ap/reports/vat?month=soon").status_code == 400


# ---- V05: selection, the gate, the reviewer, the evidence at download ------------------------------------------------------------
def test_the_sample_the_gate_a_reviewer_and_the_evidence_filed_at_download(monkeypatch, ap_review_gates):
    u, rene = _user(), _rene()
    vendor = str(random.randint(710000, 719999))
    _posted(vendor)
    run = random.randint(100000000, 199999999)
    # no reason to review it (green, under R50 000, SAP posted the vendor before) and a 0 % sample: not picked; it goes out
    monkeypatch.setitem(PF.DEFAULTS, "review_sample_pct", 0)
    b1 = _drop(u, {"plain.pdf": eskom_pdf(_acct(vendor), str(run))}, f"Plain {run}")
    d1 = _sheet(u, b1["id"])["plain.pdf"]
    assert d1["status"] == "green" and d1["review"] == "", d1
    up = u.post(f"/api/ap/batches/{b1['id']}/sap/upload", json={"include_yellow": False})
    assert up.status_code == 200, up.text
    with SessionLocal() as s:
        ev = {e.type: e for e in s.query(ApEvidence).filter(ApEvidence.batch_id == b1["id"]).all()}
        assert set(ev) == {"SAP upload file", "Review workbook"} and ev["SAP upload file"].source == "download" and ev["SAP upload file"].by_name == "Cara Person"
        assert ev["SAP upload file"].sha256 == hashlib.sha256(up.content).hexdigest()  # exactly the file that was downloaded
        f = s.get(StoredFile, ev["Review workbook"].file_id)
        assert hashlib.sha256(storage.abs_path(f).read_bytes()).hexdigest() == ev["Review workbook"].sha256 and f.rel_path.startswith("Tools/AP evidence/")
        assert ev["SAP upload file"].keep_until[:4] == str(utcnow().year + 5) and ev["SAP upload file"].portfolio and ev["SAP upload file"].stage == "upload"
        assert s.query(AuditLog).filter(AuditLog.action == "ap_evidence_added", AuditLog.target == f"ap_batch:{b1['id']}").count() == 1
        eid = ev["SAP upload file"].id
    got = rene.get(f"/api/ap/evidence/{eid}/file")
    assert got.status_code == 200 and got.content == up.content
    # the whole sample: picked, and the upload waits for it
    monkeypatch.setitem(PF.DEFAULTS, "review_sample_pct", 100)
    b2 = _drop(u, {"sampled.pdf": eskom_pdf(_acct(vendor), str(run + 1))}, f"Sampled {run}")
    d2 = _sheet(u, b2["id"])["sampled.pdf"]
    assert d2["status"] == "green" and d2["review"] == "selected" and d2["review_why"] == "random sample (100 %)"
    with SessionLocal() as s:
        seed = s.get(ApBatch, b2["id"]).review_seed
        assert seed and RV.sampled(seed, d2["id"], 100)
    r = u.post(f"/api/ap/batches/{b2['id']}/sap/upload", json={"include_yellow": False})
    assert r.status_code == 400 and r.json()["detail"].startswith("1 document waits for review")
    # Needs you: Rene has a line into Controls → Review; Cara (who dropped it) has none for it
    assert any(x["key"] == "review" and x["sub"] == "review" for x in rene.get("/api/ap/needs").json()["controls"])
    q = rene.get("/api/ap/review").json()
    mine = next(d for d in q["docs"] if d["id"] == d2["id"])
    assert mine["can"] is True and mine["why"] == "random sample (100 %)" and "No twin found" in mine["duplicate"]
    one = rene.get(f"/api/ap/review/{d2['id']}").json()
    assert one["pages"] == 1 and one["is_pdf"] and one["editable"]["account"] and "invoice_no" in one["error_fields"]
    r = u.post(f"/api/ap/review/{d2['id']}", json={"action": "agree"})
    assert r.status_code == 403 and "dropped or worked on" in r.json()["detail"]
    assert rene.post(f"/api/ap/review/{d2['id']}", json={"action": "approve"}).status_code == 400
    r = rene.post(f"/api/ap/review/{d2['id']}", json={"action": "agree", "note": "checked against the bill"})
    assert r.status_code == 200 and r.json()["state"] == "agreed" and r.json()["reviewed_by"] == "Rene Reviewer"
    assert rene.post(f"/api/ap/review/{d2['id']}", json={"action": "agree"}).status_code == 400  # not waiting any more
    up2 = u.post(f"/api/ap/batches/{b2['id']}/sap/upload", json={"include_yellow": False})
    assert up2.status_code == 200, up2.text
    # the review workbook says why, who, when and the outcome
    ws = load_workbook(io.BytesIO(u.get(f"/api/ap/batches/{b2['id']}/sap/review").content)).active
    heads = [c.value for c in ws[1]]
    row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[heads.index("pkey")] in ("31", "21"))
    assert row[heads.index("Selected why")] == "random sample (100 %)" and row[heads.index("Reviewer")] == "Rene Reviewer" and row[heads.index("Outcome")] == "agreed"
    assert row[heads.index("Reviewed")]
    with SessionLocal() as s:
        it = s.get(ApItem, d2["item_id"])
        assert any("Review: Agreed" in h["text"] for h in it.history)
        assert s.query(AuditLog).filter(AuditLog.action == "ap_capture_reviewed", AuditLog.target == f"ap_item:{it.id}").count() == 1


def test_correct_reject_changes_after_the_review_roles_and_the_coverage_report(monkeypatch, ap_review_gates):
    _vendor()
    u, rene, adm = _user(), _rene(), _admin()
    run = uuid.uuid4().hex[:6].upper()
    big = f"RV-{run}-9"
    b = _drop_invoices(monkeypatch, u, {"big.pdf": _answer(big, excl="60000.00", vat="9000.00", incl="69000.00", bill_to_vat="4999888777")}, f"Big {run}")
    d = _sheet(u, b["id"])["big.pdf"]
    assert d["status"] == "green" and d["review"] == "selected" and d["review_why"].startswith("R 69 000.00 is R 50 000 or more"), d
    cid = d["id"]
    # Reject: a note is needed; then red until it changes
    assert rene.post(f"/api/ap/review/{cid}", json={"action": "reject"}).status_code == 400
    r = rene.post(f"/api/ap/review/{cid}", json={"action": "reject", "note": "the invoice number is the PO number", "errors": ["invoice_no", "nonsense"]})
    assert r.status_code == 200 and r.json()["state"] == "rejected" and r.json()["errors"] == ["invoice_no"]
    c = _cap(u, cid)
    assert c["sap_status"] == "red" and any("rejected in the review by Rene Reviewer: the invoice number is the PO number" in x["text"] for x in c["sap_issues"])
    # Cara fixes the number: changed after the review → picked again, no longer red
    fixed = f"RV-{run}-10"
    r = u.post(f"/api/ap/captures/{cid}/decide", json={"action": "save", "data": {"fields": {"invoice_no": fixed}}})
    assert r.status_code == 200, r.text
    c = _cap(u, cid)
    assert c["sap_status"] == "green" and c["review"]["state"] == "selected" and c["review"]["why"].startswith("changed after the review") and "typed by a person" in c["review"]["why"]
    assert any("Changed after the review (rejected)" in h["text"] for h in c["history"])
    # Correct: nothing changed is refused; a corrected number goes back to Pip and Tess
    assert rene.post(f"/api/ap/review/{cid}", json={"action": "correct", "fields": {"invoice_no": fixed}}).status_code == 400
    final = f"RV-{run}-11"
    r = rene.post(f"/api/ap/review/{cid}", json={"action": "correct", "fields": {"invoice_no": final}, "errors": ["supplier"], "note": "read from the page"})
    assert r.status_code == 200, r.text
    assert r.json()["state"] == "corrected" and r.json()["errors"] == ["invoice_no", "supplier"]
    with SessionLocal() as s:
        cap = s.get(ApCapture, cid)
        assert cap.sap_rows[0]["refdocnr"] == final and cap.review_hash == RV.fingerprint(cap) and cap.review_state == "corrected"
    # roles: once they exist, the upload needs the poster role (Rene reviews, Cara doesn't post) — then Cara gets it
    adm.put(f"/api/ap/roles/{_uid('v1104.rev@test.local')}", json={"portfolio": "", "roles": ["reviewer"]})
    r = u.post(f"/api/ap/batches/{b['id']}/sap/upload", json={"include_yellow": True})
    assert r.status_code == 403 and "poster role" in r.json()["detail"]
    adm.put(f"/api/ap/roles/{_uid('cap.user@test.local')}", json={"portfolio": "", "roles": ["capturer", "poster"]})
    r = u.post(f"/api/ap/batches/{b['id']}/sap/upload", json={"include_yellow": True})
    assert r.status_code == 200, r.text
    # the coverage report: selected, reviewed, corrected, and the fields found wrong
    rep = rene.get(f"/api/ap/reports/review?month={utcnow():%Y-%m}")
    assert rep.status_code == 200
    wb = load_workbook(io.BytesIO(rep.content))
    ws = wb.worksheets[0]
    heads = [c.value for c in ws[1]]
    row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == "EFM-SA")
    assert row[heads.index("Selected")] >= 1 and row[heads.index("Corrected")] >= 1 and row[heads.index("Errors: Invoice no")] >= 1
    rv = [r for r in wb["Reviews"].iter_rows(min_row=2, values_only=True) if r[4] == final]
    assert rv and rv[0][8] == "corrected" and rv[0][9] == "Rene Reviewer" and "Invoice no" in rv[0][11]
    assert u.get(f"/api/ap/reports/review?month={utcnow():%Y-%m}").status_code == 200  # Cara holds the poster role now
    assert rene.get("/api/ap/reports/review?month=13-2026").status_code == 400


# ---- V12: attaching, the checklist, Needs you, the export, retention -----------------------------------------------------------------
def _uploaded_shell_batch(days_ago: int) -> int:
    with SessionLocal() as s:
        b = ApBatch(user_id=_uid("cap.user@test.local"), name=f"Shell run {uuid.uuid4().hex[:4]}", status="ready", portfolio="SHELL", portfolio_how="picked")
        s.add(b)
        s.flush()
        it = ApItem(batch_id=b.id, name="bill.pdf", kind="utility", status="ok")
        s.add(it)
        s.flush()
        s.add(ApCapture(item_id=it.id, batch_id=b.id, bucket="utility", stage="ready", sap_status="green", portfolio="SHELL", uploaded_at=utcnow() - timedelta(days=days_ago)))
        s.commit()
        return b.id


EML = b"From: sfm@shell.example\r\nTo: ap@cbre.example\r\nSubject: Approved: payment schedule run 12\r\n\r\nApproved.\r\n"


def test_evidence_attached_checklist_needs_you_export_and_retention():
    u, rene, adm = _user(), _rene(), _admin()
    bid = _uploaded_shell_batch(10)
    st = rene.get(f"/api/ap/batches/{bid}/evidence").json()
    assert st["portfolio"] == "SHELL" and st["late"] is True and "Signed payment schedule" in st["missing"] and "Proof of payment" in st["missing"]
    assert any(x["key"] == "evidence" and x["sub"] == "evidence" for x in rene.get("/api/ap/needs").json()["controls"])
    # attach an approval mail (.eml) on the folder: its subject goes in the note; audited
    r = rene.post("/api/ap/evidence", files={"file": ("approval.eml", EML, "message/rfc822")}, data={"type": "signed payment schedule", "batch_id": str(bid), "note": "SFM"})
    assert r.status_code == 200, r.text
    e = r.json()
    assert e["type"] == "Signed payment schedule" and e["stage"] == "approval" and e["portfolio"] == "SHELL" and "Mail: Approved: payment schedule run 12" in e["note"]
    assert e["sha256"] == hashlib.sha256(EML).hexdigest() and e["deletable"] is False
    st = rene.get(f"/api/ap/batches/{bid}/evidence").json()
    assert "Signed payment schedule" not in st["missing"] and next(x for x in st["items"] if x["type"] == "Signed payment schedule")["n"] == 1
    with SessionLocal() as s:
        assert s.query(AuditLog).filter(AuditLog.action == "ap_evidence_added", AuditLog.target == f"ap_batch:{bid}").count() == 1
    # plain-words refusals
    bad = (({"file": ("x.exe", b"MZ", "application/octet-stream")}, {"type": "Other", "batch_id": str(bid)}, 400, "can't be kept"),
           ({"file": ("x.pdf", b"%PDF-1.4", "application/pdf")}, {"type": "A poem", "batch_id": str(bid)}, 400, "what kind of evidence"),
           ({"file": ("x.pdf", b"%PDF-1.4", "application/pdf")}, {"type": "Other", "batch_id": str(bid), "account": "123"}, 400, "one folder"),
           ({"file": ("x.pdf", b"%PDF-1.4", "application/pdf")}, {"type": "Other", "account": "999000999000"}, 404, "not in the utilities list"),
           ({"file": ("x.pdf", b"", "application/pdf")}, {"type": "Other", "batch_id": str(bid)}, 400, "empty"))
    for files, data, code, words in bad:
        r = rene.post("/api/ap/evidence", files=files, data=data)
        assert r.status_code == code and words in r.json()["detail"], (data, r.text)
    # on an account: a dispute reference
    acct = _acct("700999")
    r = rene.post("/api/ap/evidence", files={"file": ("dispute.pdf", b"%PDF-1.4 dispute", "application/pdf")}, data={"type": "Dispute reference", "account": acct})
    assert r.status_code == 200 and r.json()["account"] == M.acct_key(acct) and r.json()["stage"] == "account"
    listed = rene.get("/api/ap/evidence?portfolio=SHELL").json()
    assert any(x["id"] == e["id"] for x in listed["evidence"]) and "Proof of payment" in listed["checklist"]
    # the export: the naming convention, the files, the index with every hash checked
    month = e["period"]
    z = rene.get(f"/api/ap/evidence/export?portfolio=SHELL&month={month}")
    assert z.status_code == 200 and z.headers["content-type"] == "application/zip"
    with zipfile.ZipFile(io.BytesIO(z.content)) as zf:
        names = zf.namelist()
        mail = next(n for n in names if n.endswith(".eml"))
        assert mail.startswith(f"SHELL/{month}/approval/batch{bid}_") and "_Signed_payment_schedule_" in mail and zf.read(mail) == EML
        ws = load_workbook(io.BytesIO(zf.read(f"SHELL/{month}/index.xlsx"))).active
        heads = [c.value for c in ws[1]]
        row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == mail)
        assert row[heads.index("Check")] == "ok" and row[heads.index("SHA-256")] == e["sha256"]
    assert rene.get("/api/ap/evidence/export?portfolio=NOPE").status_code == 404
    # retention: nobody deletes inside it; an admin deletes after it
    assert rene.delete(f"/api/ap/evidence/{e['id']}").status_code == 403
    r = adm.delete(f"/api/ap/evidence/{e['id']}")
    assert r.status_code == 400 and "Kept until" in r.json()["detail"]
    with SessionLocal() as s:
        row = s.get(ApEvidence, e["id"])
        path = storage.abs_path(s.get(StoredFile, row.file_id))
        row.keep_until = "2001-01-01"
        s.commit()
    assert path.is_file() and adm.delete(f"/api/ap/evidence/{e['id']}").status_code == 200 and not path.exists()
    with SessionLocal() as s:
        assert s.get(ApEvidence, e["id"]) is None and s.query(AuditLog).filter(AuditLog.action == "ap_evidence_deleted", AuditLog.target == f"ap_evidence:{e['id']}").count() == 1
    # the warehouse never bins a tool's file (the evidence can't be deleted around the retention)
    with SessionLocal() as s:
        fid = s.query(ApEvidence).filter(ApEvidence.account == M.acct_key(acct)).one().file_id
    assert adm.delete(f"/api/warehouse/files/{fid}").status_code in (400, 403, 404)
    # everyday users don't see the register
    plain = signed_in("v1104.plain@test.local", "viewer", name="Vic Viewer", tools=["ap"])
    assert plain.get("/api/ap/evidence").status_code == 403 and plain.get(f"/api/ap/evidence/{e['id'] + 1}/file").status_code in (403, 404)


def test_the_pages_are_served():
    u = _user()
    idx = u.get("/static/views/ap_controls/index.js").text
    assert "export { default as review } from './review.js';" in idx and "export { default as evidence } from './evidence.js';" in idx
    for name in ("review", "evidence"):
        r = u.get(f"/static/views/ap_controls/{name}.js")
        assert r.status_code == 200 and f"key: '{name}'" in r.text
    page = u.get("/static/views/ap.js").text
    assert 'value=\\"review\\">To review' in page and "tax_invoice_words" in page and "/evidence`" in page


# ---- review findings (coordinator, 19 Oct) ----------------------------------------------------------------------------------------
def test_changed_after_the_review_stays_picked_until_a_reviewer_acts(monkeypatch, ap_review_gates):
    """HIGH: a reviewed document whose rows change (here: an admin changes the portfolio's G/L table) is picked again, and
    stays picked through every later Tess pass — even when its first reason (the sample) is gone — so the download refuses."""
    u, rene, adm = _user(), _rene(), _admin()
    vendor = str(random.randint(720000, 729999))
    _posted(vendor)
    run = random.randint(200000000, 299999999)
    monkeypatch.setitem(PF.DEFAULTS, "review_sample_pct", 100)
    b = _drop(u, {"bill.pdf": eskom_pdf(_acct(vendor), str(run))}, f"Sticky {run}")
    d = _sheet(u, b["id"])["bill.pdf"]
    assert d["review"] == "selected" and d["review_why"] == "random sample (100 %)"
    assert rene.post(f"/api/ap/review/{d['id']}", json={"action": "agree"}).status_code == 200
    monkeypatch.setitem(PF.DEFAULTS, "review_sample_pct", 0)  # its only reason is gone
    with SessionLocal() as s:
        pf = s.get(ApCapture, d["id"]).portfolio
        before = dict(s.query(ApPortfolio).filter(ApPortfolio.code == pf).one().gl or {})
    try:
        r = adm.put(f"/api/ap/portfolios/{pf}", json={"gl": {**before, "electricity": "SA010"}})  # Pip builds it again, Tess tests it (pass 1)
        assert r.status_code == 200, r.text
        with SessionLocal() as s:  # pass 2
            C.test_batch(s, s.get(ApBatch, b["id"]))
            s.commit()
            cap = s.get(ApCapture, d["id"])
            assert "SA010" in [x["account"] for x in cap.sap_rows]
            assert cap.review_state == "selected" and cap.review_why.startswith("changed after the review"), (cap.review_state, cap.review_why)
        r = u.post(f"/api/ap/batches/{b['id']}/sap/upload", json={"include_yellow": True})  # download() runs Tess again (pass 3)
        assert r.status_code == 400 and "waits for review" in r.json()["detail"], r.text
        assert _sheet(u, b["id"])["bill.pdf"]["review"] == "selected"
        # a reviewer acts: it goes out
        assert rene.post(f"/api/ap/review/{d['id']}", json={"action": "agree", "note": "new G/L checked"}).status_code == 200
        assert u.post(f"/api/ap/batches/{b['id']}/sap/upload", json={"include_yellow": True}).status_code == 200
    finally:
        with SessionLocal() as s:
            s.query(ApPortfolio).filter(ApPortfolio.code == pf).one().gl = before
            s.commit()


def test_who_may_review_is_worked_out_in_bulk(monkeypatch, ap_review_gates):
    """Perf: Needs you's counts read who may review many documents in a few queries, with the same answer as roles.require."""
    from sqlalchemy import event

    from cbx.ap import roles as R
    from cbx.db import engine

    u, rene = _user(), _rene()
    vendor = str(random.randint(730000, 739999))
    _posted(vendor)
    run = random.randint(300000000, 399999999)
    monkeypatch.setitem(PF.DEFAULTS, "review_sample_pct", 100)
    _drop(u, {f"b{i}.pdf": eskom_pdf(_acct(vendor), str(run + i)) for i in range(4)}, f"Bulk {run}")
    with SessionLocal() as s:
        caps = s.query(ApCapture).filter(ApCapture.review_state == "selected", ApCapture.uploaded_at.is_(None), ApCapture.stage == "ready").all()
        assert len(caps) >= 4
        for email in ("cap.user@test.local", "v1104.rev@test.local"):
            who = s.query(User).filter(User.email == email).one()
            bulk = RV._can_many(s, who, caps)
            for c in caps:  # the same answer as one document at a time
                ref = R.may(s, who, "reviewer", c.portfolio or R.portfolio_of(s, c)) and who.id not in R.touched_by(s, c, include_dropper=True, exclude_actions=RV.OWN_ACTIONS)
                assert bulk[c.id][0] == ref, (email, c.id)
        n = {"q": 0}

        def count(*_a, **_k):
            n["q"] += 1

        event.listen(engine, "before_cursor_execute", count)
        try:
            out = RV.counts(s, s.query(User).filter(User.email == "v1104.rev@test.local").one())
        finally:
            event.remove(engine, "before_cursor_execute", count)
        assert out["review_mine"] >= 4 and n["q"] < 40, n
