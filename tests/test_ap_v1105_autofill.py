"""1.0.105 (Builder G): never stop, guess well, learn (cbx/ap/autofill.py).
Edwin (19 Oct): "Can't we make it a simple, single click — it does everything and gets me results." Where CBX can't be sure:
the best guess, filled in, marked yellow — never a blank, never a stop. The group's own entities (EPS and Nervada → Z999,
Excellerate Holdings → Z102) are in the company list and never land on X143; a supplier with several SAP vendors gets one;
the model's coding is checked against the lists it was given; dates fall back in order; the corrected download teaches
CBX; a one-time repair takes what waited in Needs you out by itself. Suppliers, VAT numbers and amounts are made up
(the group's own VAT numbers are the real ones Edwin gave)."""
from __future__ import annotations

import io
import uuid
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from cbx import knowledge as K
from cbx.ap import autofill as AF
from cbx.ap import capture as C
from cbx.ap import masters as M
from cbx.ap import sap as P
from cbx.db import SessionLocal
from cbx.models import AppState, AuditLog, User
from cbx.models.ap import ApCapture, ApEntity, ApItem, ApSapVendor, ApVendorCode, ApVendorCompany

from .test_ap_capture_1098 import _admin, _app, _drop, _sheet, _user, _vision  # noqa: F401 — _app: the app started, the AP lists seeded
from .test_ap_scans_1098 import brightline_pdf

EPS, EPS_VAT = "Excellerate Property Services (Pty) Ltd", "4720256405"
NERVADA, NERVADA_VAT = "NERVADA TRADING 81 (PTY) LIMITED", "4910267154"
EHL, EHL_VAT = "Excellerate Holdings Limited", "4050175522"
INV_DAY = date.today() - timedelta(days=12)
MINE: dict[str, set] = {"sapv": set(), "vc": set(), "codes": set()}


@pytest.fixture(autouse=True)
def _tidy():
    """The group's entities as 1.0.105 adds them (an earlier test may have replaced the company list by an upload)."""
    with SessionLocal() as s:
        st = s.get(AppState, AF.GROUP_KEY)
        if st is not None:
            s.delete(st)
            s.flush()
        AF.seed_group_entities(s)
        s.commit()
    K.forget()
    yield
    with SessionLocal() as s:
        s.query(ApSapVendor).filter(ApSapVendor.vendor.in_(sorted(MINE["sapv"]) or ["-"])).delete(synchronize_session=False)
        for v, co in MINE["vc"]:
            s.query(ApVendorCompany).filter(ApVendorCompany.vendor == v, ApVendorCompany.company_code == co).delete(synchronize_session=False)
        s.query(ApVendorCode).filter(ApVendorCode.vendor_no.in_(sorted(MINE["codes"]) or ["-"])).delete(synchronize_session=False)
        for g in AF.GROUP:  # names a test taught the group's entities (a company corrected in a file) go again
            e = s.query(ApEntity).filter(ApEntity.code == g["code"]).first()
            if e is not None and e.source != "person":
                e.aliases = "\n".join(g["aliases"])
        s.commit()
    for k in MINE:
        MINE[k].clear()
    K.forget()


def _vat() -> str:
    return "46" + str(uuid.uuid4().int)[:8]


def _no(prefix: str) -> str:
    return prefix + str(uuid.uuid4().int)[:6]


def _sapv(vendor: str, name: str, vat: str, *, created: str = "", terms: str = "", opened=()) -> str:
    with SessionLocal() as s:
        s.add(ApSapVendor(vendor=vendor, name=name, name_key=M.name_key(name), vat=vat, tax1="", city="Sandton", country="ZA", blocked=False,
                          payment_terms=terms, source="excel", account_group="ZMGN", created_on=created or None))
        for co in opened:
            s.add(ApVendorCompany(vendor=vendor, company_code=co))
            MINE["vc"].add((vendor, co))
        s.commit()
    MINE["sapv"].add(vendor)
    MINE["codes"].add(vendor)
    K.forget()
    return vendor


def _answer(no: str, supplier: str, vat: str, *, bill_to: str = EPS, bill_vat: str = EPS_VAT, due: bool = True, inv: bool = True, incl: str = "1150.00") -> dict:
    conf = {k: 93 for k in ("supplier", "bill_to", "invoice_no", "invoice_date", "total_incl", "total_excl", "vat")}
    excl = f"{float(incl) / 1.15:.2f}"
    vat_amt = f"{float(incl) - float(excl):.2f}"
    if not inv:
        conf.pop("invoice_date")
    return {"is_invoice": True, "credit_note": False, "supplier": supplier, "supplier_vat": vat, "bill_to": bill_to, "bill_to_vat": bill_vat, "invoice_no": no,
            "invoice_date": INV_DAY.isoformat() if inv else None, "due_date": (INV_DAY + timedelta(days=20)).isoformat() if due else None, "po": "", "our_account": "",
            "currency": "ZAR", "lines": [{"description": "Internet services September", "amount_excl": excl, "vat_rate": 15, "vat": vat_amt}],
            "total_excl": excl, "vat": vat_amt, "total_incl": incl, "tax_invoice_words": True, "confidence": conf}


def _run(monkeypatch, u, answers: dict, name: str, *, model=None, date_reader=None) -> dict:
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision(answers))
    monkeypatch.setattr(AF, "ask_model", model or (lambda *a, **k: {}))  # the AI's coding: nothing unless a test says
    monkeypatch.setattr(AF, "ask_date", date_reader or (lambda db, c: ""))
    return _drop(u, {n: brightline_pdf(n) for n in answers}, name)


def _cap(cid: int) -> ApCapture:
    with SessionLocal() as s:
        c = s.get(ApCapture, cid)
        s.expunge(c)
        return c


def _texts(d: dict) -> set[str]:
    return {x["text"] for x in d["issues"]}


# ---- 1. the group's own entities -------------------------------------------------------------------------------------------------
def test_eps_nervada_and_holdings_are_in_the_company_list_by_vat_and_by_name():
    with SessionLocal() as s:
        z999, z102 = (s.query(ApEntity).filter(ApEntity.code == c).one() for c in ("Z999", "Z102"))
        assert (z999.name, z999.vat_no, z999.source) == (EPS, EPS_VAT, "seed") and "Nervada" in z999.aliases.splitlines() and "EPS" in z999.aliases.splitlines()
        assert (z102.name, z102.vat_no) == ("Excellerate Holdings Ltd", EHL_VAT) and "EHL" in z102.aliases.splitlines()
        assert AF.extra_vats(z999) == {NERVADA_VAT}
        # VAT first (Nervada's own number is Z999's too), then the name — never X143, never the Congo company X135
        for bill_to, vat, code, how in ((EPS, EPS_VAT, "Z999", "VAT number"), (NERVADA, NERVADA_VAT, "Z999", "VAT number"), (EHL, EHL_VAT, "Z102", "VAT number"),
                                        ("EXCELLERATE PROPERTY SERVICES PTY LTD", "", "Z999", "name"), (NERVADA, "", "Z999", "name"),
                                        ("Excellerate Holdings Ltd", "", "Z102", "name")):
            got = C.match_entity(s, bill_to, vat)
            assert (got["code"], got["how"]) == (code, how), (bill_to, got)
        # a word of the name: "EPS Group" is Z999 (a guess); one shared word ("Holdings") names nobody
        assert AF.by_words(s, "EPS Group", "")["code"] == "Z999"
        assert AF.by_words(s, "Somebody Else Holdings", "") is None
    # the one-time add never touches a row a person keeps
    with SessionLocal() as s:
        e = s.query(ApEntity).filter(ApEntity.code == "Z102").one()
        keep = (e.aliases, e.source)
        e.source, e.aliases = "person", "Holdings As Typed"
        s.delete(s.get(AppState, AF.GROUP_KEY))
        s.commit()
    with SessionLocal() as s:
        AF.seed_group_entities(s)
        s.commit()
        e = s.query(ApEntity).filter(ApEntity.code == "Z102").one()
        assert e.aliases == "Holdings As Typed" and s.get(AppState, AF.GROUP_KEY) is not None
        e.aliases, e.source = keep
        s.commit()
    K.forget()


def test_a_model_guess_never_puts_a_bill_to_that_names_another_company_on_x143(monkeypatch):
    with SessionLocal() as s:
        x143 = s.query(ApEntity).filter(ApEntity.code == "X143").one()
        # how X143 came onto EPS invoices before 1.0.105: EPS was in no list, its SA VAT number ruled out the Congo company
        # X135 by name, and Mia's model was asked — X143 "CBRE Excellerate Faciliti" was the nearest "Excellerate" left
        assert not AF.model_pick_ok(x143, EPS, EPS_VAT) and not AF.model_pick_ok(x143, "EPS Head Office", "")
        assert AF.model_pick_ok(x143, "CBRE Excellerate Facilities Management (Pty) Ltd - RIO2X21", "")

        class Fake:
            def __init__(self, db):
                pass

            def json(self, *a, **k):
                return {"code": "X143", "reason": "the nearest Excellerate company"}

        import cbx.cexa.gateway as G

        monkeypatch.setattr(G, "Gateway", Fake)
        assert C._ask_entity(s, "Excellerate Property Services Head Office", EPS_VAT) is None
        assert C._ask_entity(s, "CBRE Excellerate Facilities Mgmt", "")["code"] == "X143"


def test_eps_nervada_and_holdings_invoices_post_in_z999_and_z102_through_the_line(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    eps_vat, nv_vat, ehl_vat = _vat(), _vat(), _vat()
    eps_v = _sapv(_no("81"), f"Netup {run} (Pty) Ltd", eps_vat, opened=("Z999",))
    nv_v = _sapv(_no("82"), f"Cleaners {run} (Pty) Ltd", nv_vat, opened=("Z999",))
    ehl_v = _sapv(_no("83"), f"Secretarial {run} (Pty) Ltd", ehl_vat, opened=("Z102",))
    u = _user()
    b = _run(monkeypatch, u, {f"eps_{run}.pdf": _answer(f"E-{run}", f"Netup {run} (Pty) Ltd", eps_vat),
                              f"nv_{run}.pdf": _answer(f"N-{run}", f"Cleaners {run} (Pty) Ltd", nv_vat, bill_to=NERVADA, bill_vat=NERVADA_VAT),
                              f"ehl_{run}.pdf": _answer(f"H-{run}", f"Secretarial {run} (Pty) Ltd", ehl_vat, bill_to=EHL, bill_vat=EHL_VAT),
                              f"eps2_{run}.pdf": _answer(f"E2-{run}", f"Netup {run} (Pty) Ltd", eps_vat, bill_to="EXCELLERATE PROPERTY SERVICES PTY LTD", bill_vat="")},
             f"Group {run}")
    docs = _sheet(u, b["id"])
    assert {k: (d["entity"], d["vendor"], d["needs"]) for k, d in docs.items()} == {
        f"eps_{run}.pdf": ("Z999", eps_v, ""), f"nv_{run}.pdf": ("Z999", nv_v, ""), f"ehl_{run}.pdf": ("Z102", ehl_v, ""), f"eps2_{run}.pdf": ("Z999", eps_v, "")}
    assert all(d["status"] == "yellow" for d in docs.values()), {k: d["issues"] for k, d in docs.items()}  # guessed coding: never green
    # Nervada's cost centre from its own name; EPS (no word for it) the company's generic one; no G/L anywhere (no history, no
    # AI answer): a yellow "G/L needed", the cell left for Edwin
    nv, eps = _cap(docs[f"nv_{run}.pdf"]["id"]), _cap(docs[f"eps_{run}.pdf"]["id"])
    gl_lines = lambda c: [r for r in c.sap_rows if r["pkey"] == "40" and r.get("dt") != "X"]  # noqa: E731
    assert {r["costcentre"] for r in gl_lines(nv)} == {"NV0001"} and nv.guesses["cost_centre"] == "the bill-to names it (Nervada)"
    assert {r["costcentre"] for r in gl_lines(eps)} == {"NV0009"} and eps.guesses["cost_centre"].startswith("the company's generic cost centre")
    assert {r["account"] for r in gl_lines(eps)} == {""} and eps.guesses["gl"] == AF.GL_NEEDED
    assert "G/L needed (the G/L cell is empty)" in _texts(docs[f"eps_{run}.pdf"])
    assert not any(x["level"] == "red" for x in docs[f"eps_{run}.pdf"]["issues"])
    # the guessed cells, for the styled download
    vl = next(r for r in eps.sap_rows if r["pkey"] == "31")
    assert AF.row_guesses(eps, vl) == {} or set(AF.row_guesses(eps, vl)) <= {"comp", "account", "DueDate", "docdate"}
    assert set(AF.row_guesses(eps, gl_lines(eps)[0])) >= {"account", "costcentre"}


# ---- 2. the vendor, the coding, the dates ---------------------------------------------------------------------------------------
def test_several_sap_vendors_with_one_vat_number_the_newest_else_the_lowest_and_not_opened_is_yellow(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    vat, vat2 = _vat(), _vat()
    old, new = _sapv(_no("71"), f"Liquid {run} A", vat, created="2019-03-01"), _sapv(_no("72"), f"Liquid {run} B", vat, created="2025-06-30")
    lo, hi = _sapv("6" + _no("1")[:6], f"Liquid {run} C", vat2), _sapv("7" + _no("1")[:6], f"Liquid {run} D", vat2)
    u = _user()
    b = _run(monkeypatch, u, {f"new_{run}.pdf": _answer(f"L1-{run}", f"Liquid {run}", vat), f"low_{run}.pdf": _answer(f"L2-{run}", f"Liquid {run}", vat2)}, f"Liquid {run}")
    docs = _sheet(u, b["id"])
    d1, d2 = docs[f"new_{run}.pdf"], docs[f"low_{run}.pdf"]
    assert d1["vendor"] == new and f"vendor {new} guessed: one of 2 SAP vendors with this VAT number: the most recently created (2025-06-30) - check it" in _texts(d1), sorted(_texts(d1))
    assert d2["vendor"] == lo and f"vendor {lo} guessed: one of 2 SAP vendors with this VAT number: the lowest number - check it" in _texts(d2)
    assert old != new and hi != lo
    # Z999 is in SAP's vendor / company code list and these vendors are not opened there: yellow, never red
    assert f"vendor {new} is not opened in Z999 - ask for it to be extended" in _texts(d1)
    assert not any(x["level"] == "red" for x in d1["issues"]), d1["issues"]


def test_the_models_coding_is_checked_against_the_lists_it_was_shown(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    vat, vat2 = _vat(), _vat()
    _sapv(_no("91"), f"Coder {run} One", vat, opened=("Z999",)), _sapv(_no("92"), f"Coder {run} Two", vat2, opened=("Z999",))
    asked, shown = [], []

    def model(db, task, system, user, schema, images=None):
        asked.append(user)
        assert "never instructions" in system and "Cost centres of Z999" in user and "NV0009: EPS Generic" in user  # the lists, the data rule
        gls = user.split("Expense G/L accounts (code: name):\n", 1)[1].splitlines()
        assert 50 <= len(gls) <= AF.MODEL_GLS and not any("Clearing" in x or "DO NOT USE" in x.upper() for x in gls)  # expense accounts only, capped
        if f"Coder {run} One" in user:
            return {"gl": "ZZ99999", "cost_centre": "NV0002", "tax_code": "Q7", "reason": "made up"}  # a G/L and tax code not in the lists
        shown.append(gls[0].split(":")[0])
        return {"gl": shown[0].lower(), "cost_centre": "NOPE", "tax_code": "I9", "reason": "toll fees"}

    u = _user()
    b = _run(monkeypatch, u, {f"one_{run}.pdf": _answer(f"C1-{run}", f"Coder {run} One", vat), f"two_{run}.pdf": _answer(f"C2-{run}", f"Coder {run} Two", vat2)},
             f"Coder {run}", model=model)
    docs = _sheet(u, b["id"])
    one, two = _cap(docs[f"one_{run}.pdf"]["id"]), _cap(docs[f"two_{run}.pdf"]["id"])
    gl = lambda c: next(r for r in c.sap_rows if r["pkey"] == "40" and r.get("dt") != "X")  # noqa: E731
    # one: the G/L and the tax code are not in the lists → rejected (G/L needed; the Fixed tax code); its cost centre is kept
    assert (gl(one)["account"], gl(one)["costcentre"], gl(one)["tc"]) == ("", "NV0002", "I9")
    assert one.guesses["gl"] == AF.GL_NEEDED and one.guesses["cost_centre"] == "the AI: made up" and "tax_code" not in one.guesses
    # two: the G/L as the chart spells it; a cost centre not in the list → the generic one
    assert (gl(two)["account"], gl(two)["tc"]) == (shown[0], "I9") and two.guesses["gl"] == "the AI: toll fees"
    assert gl(two)["costcentre"] == "NV0009"
    assert docs[f"two_{run}.pdf"]["status"] == "yellow" and f"G/L {shown[0]} guessed: the AI: toll fees - check it" in _texts(docs[f"two_{run}.pdf"])
    n = len(asked)
    with SessionLocal() as s:  # Pip builds again (a posting date, a site): the answer is kept, the model is not asked again
        c = s.get(ApCapture, two.id)
        C.build_one(s, c, s.get(ApItem, c.item_id), M.tables(s))
        s.rollback()
    assert len(asked) == n


def test_the_due_date_and_the_invoice_date_fall_back_in_order(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    vat_terms, vat_none = _vat(), _vat()
    _sapv(_no("51"), f"Terms {run}", vat_terms, terms="N14", opened=("Z999",))
    _sapv(_no("52"), f"Plain {run}", vat_none, opened=("Z999",))
    coded_vat = _vat()
    coded = _no("53")
    with SessionLocal() as s:
        s.add(ApVendorCode(vendor_no=coded, name=f"Coded {run}", name_key=M.name_key(f"Coded {run}"), vat_no=coded_vat, gl="SC205", cost_centre="NV0002", tax_code="I9",
                           payment_days=45, source="excel"))
        s.commit()
    MINE["codes"].add(coded)
    u = _user()
    reads = []

    def date_reader(db, c):
        reads.append(c.id)
        return (INV_DAY - timedelta(days=1)).isoformat() if len(reads) == 1 else ""

    b = _run(monkeypatch, u, {f"print_{run}.pdf": _answer(f"D0-{run}", f"Terms {run}", vat_terms),
                              f"coded_{run}.pdf": _answer(f"D1-{run}", f"Coded {run}", coded_vat, due=False),
                              f"terms_{run}.pdf": _answer(f"D2-{run}", f"Terms {run}", vat_terms, due=False),
                              f"plus30_{run}.pdf": _answer(f"D3-{run}", f"Plain {run}", vat_none, due=False),
                              f"nodate_{run}.pdf": _answer(f"D4-{run}", f"Plain {run}", vat_none, due=False, inv=False)},
             f"Dates {run}", date_reader=date_reader)
    docs = _sheet(u, b["id"])
    due = {k.split("_")[0]: next(r for r in _cap(d["id"]).sap_rows if r["pkey"] == "31")["DueDate"] for k, d in docs.items()}
    assert due["print"] == f"{INV_DAY + timedelta(days=20):%d%m%Y}" and "due_date" not in (_cap(docs[f"print_{run}.pdf"]["id"]).guesses or {})
    assert due["coded"] == f"{INV_DAY + timedelta(days=45):%d%m%Y}"
    assert _cap(docs[f"coded_{run}.pdf"]["id"]).guesses["due_date"] == "the invoice date + 45 days (the vendor's payment terms)"
    assert due["terms"] == f"{INV_DAY + timedelta(days=14):%d%m%Y}"
    assert _cap(docs[f"terms_{run}.pdf"]["id"]).guesses["due_date"] == "the invoice date + 14 days (SAP's payment terms N14)"
    assert due["plus30"] == f"{INV_DAY + timedelta(days=30):%d%m%Y}"
    # no invoice date on the page: the senior's reading goes on (not Needs you); the reader is asked once more for the date
    nd = docs[f"nodate_{run}.pdf"]
    c = _cap(nd["id"])
    assert nd["needs"] == "" and nd["status"] == "yellow" and c.read.get("autofill_date_asked")
    assert c.guesses["invoice_date"] == "read again from the page image" and next(r for r in c.sap_rows if r["pkey"] == "31")["docdate"] == f"{INV_DAY - timedelta(days=1):%d%m%Y}"
    assert c.guesses["due_date"] == "the invoice date + 30 days"
    assert reads == [c.id]  # once
    # the reader can't say either: the posting date (unit, no AI)
    with SessionLocal() as s:
        x = ApCapture(fields={"supplier": "X", "invoice_no": "1", "incl": "1.00"}, match={}, read={"autofill_date_asked": "earlier"}, bucket="manual")
        af: dict = {}
        AF.dates(s, x, {}, af)
        assert af["invoice_date"] == {**af["invoice_date"], "value": C.posting_day(None).isoformat(), "how": "the posting date", "sure": False}
    assert AF.reading_goes_on({"supplier": "A", "invoice_no": "1", "incl": "1.00"}, ["no invoice date", "not sure enough (60%)"])
    assert not AF.reading_goes_on({"supplier": "A", "invoice_no": "1", "incl": "1.00"}, ["not sure enough (40%)"])  # only doubt: a person
    assert not AF.reading_goes_on({"supplier": "A", "invoice_no": "", "incl": "1.00"}, ["no invoice date"])


# ---- 3. learn from the corrected Excel ------------------------------------------------------------------------------------------
def _upload_of(caps: list[ApCapture], change) -> bytes:
    rows = []
    for c in caps:
        for r in c.sap_rows:
            r = dict(r)
            change(c, r)
            rows.append(r)
    return P.upload_file(rows)


def test_the_corrected_download_teaches_cbx_and_the_next_invoice_uses_it(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    vat, vat2 = _vat(), _vat()
    v1 = _sapv(_no("41"), f"Teach {run} (Pty) Ltd", vat, opened=("Z999",))
    v2 = _sapv(_no("42"), f"Sent {run} (Pty) Ltd", vat2, opened=("Z999",))
    u = _user()
    b = _run(monkeypatch, u, {f"t1_{run}.pdf": _answer(f"T1-{run}", f"Teach {run} (Pty) Ltd", vat), f"s1_{run}.pdf": _answer(f"S1-{run}", f"Sent {run} (Pty) Ltd", vat2)},
             f"Teach {run}")
    docs = _sheet(u, b["id"])
    t1, s1 = _cap(docs[f"t1_{run}.pdf"]["id"]), _cap(docs[f"s1_{run}.pdf"]["id"])
    assert t1.guesses.get("gl") == AF.GL_NEEDED
    with SessionLocal() as s:  # s1 went out in an upload already
        s.get(ApCapture, s1.id).uploaded_at = s.get(ApCapture, s1.id).built_at
        s.commit()

    def fix(c, r):  # Edwin fills the yellow cells: the G/L and the cost centre of each G/L line
        if r["pkey"] == "40" and r.get("dt") != "X":
            r["account"], r["costcentre"] = ("SC205", "NV0003") if c.id == t1.id else ("SB000", "NV0004")

    data = _upload_of([t1, s1], fix)
    with SessionLocal() as s:
        me = s.query(User).filter_by(email="cap.user@test.local").one()
        out = AF.learn_from_upload(s, me, data)
        s.commit()
    assert out["learnt"] == 2 and out["text"] == "Learnt 2 codings from your file." and out["matched"] == 2
    MINE["codes"] |= {v1, v2}
    with SessionLocal() as s:
        rows = {r.vendor_no: r for r in s.query(ApVendorCode).filter(ApVendorCode.vendor_no.in_([v1, v2])).all()}
        assert (rows[v1].gl, rows[v1].cost_centre, rows[v1].company_code, rows[v1].learned, rows[v1].learned_by) == ("SC205", "NV0003", "Z999", True, "Cara Person")
        a = s.query(AuditLog).filter(AuditLog.action == "ap_learned_from_upload").order_by(AuditLog.id.desc()).first()
        assert a is not None and a.detail.startswith("2 codings learnt")
    # t1 was not out yet: built again with the person's coding — still yellow (learned: a second person confirms it)
    d = _sheet(u, b["id"])[f"t1_{run}.pdf"]
    c = _cap(d["id"])
    gl = next(r for r in c.sap_rows if r["pkey"] == "40" and r.get("dt") != "X")
    assert (gl["account"], gl["costcentre"]) == ("SC205", "NV0003") and d["status"] == "yellow"
    assert c.guesses["gl"].startswith("as coded by Cara Person on ")
    # s1 went out: what was sent stays exactly as sent (1.0.105 review); the person's correction is kept beside it, audited
    s1b = _cap(s1.id)
    assert s1b.sap_rows == s1.sap_rows and (s1b.company_code, s1b.vendor_no) == (s1.company_code, s1.vendor_no)
    corr = s1b.match["corrected_after_upload"]
    assert corr["by"] == "Cara Person" and corr["values"] == {"cost_centre": "NV0004", "gl": "SB000"}
    with SessionLocal() as s:
        assert s.query(AuditLog).filter(AuditLog.action == "ap_corrected_after_upload", AuditLog.target == f"ap_item:{s1.item_id}").first() is not None
    # the vendor's next invoice uses the coding at once
    b2 = _run(monkeypatch, u, {f"t2_{run}.pdf": _answer(f"T2-{run}", f"Teach {run} (Pty) Ltd", vat)}, f"Teach2 {run}")
    d2 = _sheet(u, b2["id"])[f"t2_{run}.pdf"]
    c2 = _cap(d2["id"])
    gl2 = next(r for r in c2.sap_rows if r["pkey"] == "40" and r.get("dt") != "X")
    assert (gl2["account"], gl2["costcentre"], d2["status"]) == ("SC205", "NV0003", "yellow")
    assert f"vendor {v1} is learned: an admin confirms it once" in _texts(d2) and c2.guesses["gl"].startswith("as coded by Cara Person")
    # nothing changed and nothing guessed: nothing to learn; a file that is not the upload is refused in plain words
    with SessionLocal() as s:
        me = s.query(User).filter_by(email="cap.user@test.local").one()
        assert AF.learn_from_upload(s, me, _upload_of([_cap(s1.id)], lambda c, r: None))["text"] == "Nothing new to learn from your file."
        with pytest.raises(AF.Refused):
            AF.learn_from_upload(s, me, b"not a workbook")
        s.rollback()
    # the yellow cells in the styled file carry the how (Builder S's download marks them with row_guesses)
    ws = load_workbook(io.BytesIO(P.upload_file([{**r, "_guess": AF.row_guesses(c2, r)} for r in c2.sap_rows]))).active
    heads = [x.value for x in ws[1]]
    gl_row = next(i for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), 2) if r[heads.index("pkey")] == "40" and r[heads.index("account")] == "SC205")
    note = ws.cell(row=gl_row, column=heads.index("account") + 1).comment
    assert note is not None and "as coded by Cara Person" in note.text


# ---- 4. the one-time repair ---------------------------------------------------------------------------------------------------------
def test_the_repair_takes_what_waited_in_needs_you_out_and_x143_guesses_back_to_mia(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    vat = _vat()
    v = _sapv(_no("21"), f"Stuck {run} (Pty) Ltd", vat, opened=("Z999",))
    u = _user()
    names = [f"ent_{run}.pdf", f"ven_{run}.pdf", f"x143_{run}.pdf", f"unread_{run}.pdf"]
    b = _run(monkeypatch, u, {n: _answer(f"R{i}-{run}", f"Stuck {run} (Pty) Ltd", vat) for i, n in enumerate(names)}, f"Repair {run}")
    docs = _sheet(u, b["id"])
    with SessionLocal() as s:  # as 1.0.104 left them on 19 Oct
        for n in names:
            c = s.get(ApCapture, docs[n]["id"])
            m = dict(c.match or {})
            if n.startswith("ent_"):
                c.bucket, c.stage, c.needs, c.match, c.company_code = "other", "needs", "entity", {"entity": {"found": False, "bill_to": EPS}}, ""
            elif n.startswith("ven_"):
                m["vendor"] = {"found": False, "supplier": f"Stuck {run} (Pty) Ltd", "vat": vat}
                c.stage, c.needs, c.match, c.vendor_no = "needs", "vendor", m, ""
            elif n.startswith("x143_"):
                m["entity"] = {"code": "X143", "name": "CBRE Excellerate Faciliti", "how": "Mia's reading (closest Excellerate)", "unsure": True}
                c.match, c.company_code = m, "X143"
            else:
                f = dict(c.fields)
                f["invoice_date"] = ""
                c.fields, c.stage, c.needs = f, "needs", "unread"
                c.read = {**(c.read or {}), "seth": {"problems": ["no invoice date", "not sure enough (60%)"], "confidence": 60}}
        s.commit()
    assert len([x for x in u.get("/api/ap/needs").json()["vendors"] if x["vat"] == vat]) == 1
    n = AF.run_repair(force=True)
    assert n >= 4
    docs = _sheet(u, b["id"])
    assert {k: (d["entity"], d["vendor"], d["needs"], d["stage"]) for k, d in docs.items()} == {k: ("Z999", v, "", "ready") for k in names}
    assert all(d["status"] == "yellow" for d in docs.values())
    assert _cap(docs[f"unread_{run}.pdf"]["id"]).guesses["invoice_date"] == "the posting date"
    with SessionLocal() as s:
        a = s.query(AuditLog).filter(AuditLog.action == "ap_entity_rematched").order_by(AuditLog.id.desc()).first()
        assert a is not None and "back to Mia" in a.detail
        assert any("Back to Mia: CBX now fills in" in h["text"] for h in s.get(ApItem, docs[f"ent_{run}.pdf"]["item_id"]).history)
    assert AF.run_repair() is None  # once
    assert not [x for x in u.get("/api/ap/needs").json()["vendors"] if x["vat"] == vat]


# ---- 5. 1.0.105 review ---------------------------------------------------------------------------------------------------------------
def test_review_learning_never_changes_what_went_out_nor_amounts_and_only_from_the_persons_own_documents(monkeypatch):
    from cbx.models.ap import ApRole

    run = uuid.uuid4().hex[:6].upper()
    vat = _vat()
    v = _sapv(_no("43"), f"Guard {run} (Pty) Ltd", vat, opened=("Z999",))
    u = _user()
    b = _run(monkeypatch, u, {f"g1_{run}.pdf": _answer(f"G1-{run}", f"Guard {run} (Pty) Ltd", vat), f"g2_{run}.pdf": _answer(f"G2-{run}", f"Guard {run} (Pty) Ltd", vat)},
             f"Guard {run}")
    docs = _sheet(u, b["id"])
    g1, g2 = _cap(docs[f"g1_{run}.pdf"]["id"]), _cap(docs[f"g2_{run}.pdf"]["id"])
    with SessionLocal() as s:
        s.get(ApCapture, g1.id).uploaded_at = s.get(ApCapture, g1.id).built_at
        s.commit()
    g1 = _cap(g1.id)

    def tamper(c, r):  # amounts ×10, the company and the vendor changed, a G/L filled in
        r["amount"] = f"{float(r['amount']) * 10:.2f}"
        r["comp"] = "Z102"
        if r["pkey"] == "31":
            r["account"] = "999999"
        elif r["pkey"] == "40" and r.get("dt") != "X":
            r["account"] = "SC205"

    # another person (not the dropper, no capturer role in GENSEC once roles exist): nothing of hers is used
    mal = __import__("tests.test_ap_capture_1098", fromlist=["signed_in"]).signed_in(f"mal.{run.lower()}@test.local", "user", name="Mal Other", tools=["ap"])
    assert mal.get("/api/ap/needs").status_code == 200
    with SessionLocal() as s:
        cara = s.query(User).filter_by(email="cap.user@test.local").one()
        mallory = s.query(User).filter_by(email=f"mal.{run.lower()}@test.local").one()
        s.add(ApRole(user_id=cara.id, portfolio="GENSEC", role="capturer"))
        s.commit()
    try:
        def keep_vendor(c, r):
            tamper(c, r)
            if r["pkey"] == "31":
                r["account"] = c.vendor_no
            r["comp"] = c.company_code

        with SessionLocal() as s:
            out = AF.learn_from_upload(s, s.get(User, mallory.id), _upload_of([g1, g2], keep_vendor))
        assert (out["learnt"], out["not_yours"]) == (0, 2) and "2 invoices in it are not yours" in out["text"]
    finally:
        with SessionLocal() as s:
            s.query(ApRole).filter(ApRole.portfolio == "GENSEC", ApRole.role == "capturer").delete()
            s.commit()
    # the vendor in the file is not in SAP's vendor list: not learnt, said; and a reference whose vendor is another one is
    # never matched by the reference alone
    with SessionLocal() as s:
        out = AF.learn_from_upload(s, s.get(User, cara.id), _upload_of([g1, g2], tamper))
    assert out["learnt"] == 0 and out["matched"] == 0
    # the right vendor, the company changed on the one that went out: no match (its company is as sent); on the other: learnt
    def fix(c, r):
        if r["pkey"] == "40" and r.get("dt") != "X":
            r["account"], r["amount"] = "SC205", "99999.00"
        r["comp"] = "Z102"

    with SessionLocal() as s:
        out = AF.learn_from_upload(s, s.get(User, cara.id), _upload_of([g1, g2], fix))
    assert (out["matched"], out["learnt"]) == (1, 1)
    MINE["codes"].add(v)
    g1b, g2b = _cap(g1.id), _cap(g2.id)
    assert g1b.sap_rows == g1.sap_rows and g1b.company_code == "Z999" and "corrected_after_upload" not in g1b.match
    assert g2b.company_code == "Z102" and g2b.match["entity"]["how"] == "a person"
    assert all(r["amount"] != "99999.00" for r in g2b.sap_rows)  # never an amount from the file
    # a vendor SAP's vendor list doesn't have: the line says so
    def other_vendor(c, r):
        if r["pkey"] == "31":
            r["account"] = "999999"
        elif r["pkey"] == "40" and r.get("dt") != "X":
            r["account"] = "SC206"

    with SessionLocal() as s:
        x = s.get(ApCapture, g2.id)
        x.vendor_no = "999999"  # as if the document had that vendor
        s.commit()
        out = AF.learn_from_upload(s, s.get(User, cara.id), _upload_of([_cap(g2.id)], other_vendor))
        x = s.get(ApCapture, g2.id)
        x.vendor_no = v
        s.commit()
    assert out["learnt"] == 0 and out["skipped"] == [f"G2-{run}: vendor 999999 is not in SAP's vendor list - not learnt"]


def test_review_a_workbook_claiming_a_million_rows_is_refused_at_once_and_the_route_is_not_async():
    import inspect
    import re
    import time
    import zipfile

    from openpyxl import Workbook

    from cbx.api import ap_simple_api

    wb = Workbook()
    wb.active.append(P.SAP_HEADERS)
    buf = io.BytesIO()
    wb.save(buf)
    zin, out = zipfile.ZipFile(io.BytesIO(buf.getvalue())), io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for i in zin.infolist():
            d = zin.read(i.filename)
            if i.filename == "xl/worksheets/sheet1.xml":
                x = re.sub(r'<dimension ref="[^"]+"/>', '<dimension ref="A1:XFD1048576"/>', d.decode("utf-8"))
                d = x.replace("</sheetData>", '<row r="1048576"><c r="A1048576" t="inlineStr"><is><t>x</t></is></c></row></sheetData>').encode("utf-8")
            zout.writestr(i, d)
    zin.close()
    t = time.time()
    with pytest.raises(AF.Refused, match="more than 20 000 rows"):
        AF.read_upload(out.getvalue())
    assert time.time() - t < 5
    with pytest.raises(AF.Refused, match="larger than 10 MB"):
        AF.read_upload(b"x" * (AF.MAX_MB * 1024 * 1024 + 1))
    assert not inspect.iscoroutinefunction(ap_simple_api.simple_learn)


def test_review_a_bill_to_naming_another_company_is_a_group_entity_or_no_company_never_x143(monkeypatch):
    from types import SimpleNamespace

    monkeypatch.setattr(C, "model_on", lambda: False)
    with SessionLocal() as s:
        doc = lambda bt: SimpleNamespace(batch_id=None, fields={"bill_to": bt, "bill_to_vat": ""})  # noqa: E731
        got = AF.company(s, doc("Excellerate Property"), mia_tried=True)
        assert (got["code"], got["unsure"]) == ("Z999", True) and "property" in got["how"]
        assert AF.company(s, doc("Somebody Else Holdings"), mia_tried=True)["code"] == "Z102"
        assert AF.company(s, doc("Nobody We Know Ltd"), mia_tried=True)["code"] == ""
        assert AF.company(s, doc("CBRE Excellerate"), mia_tried=True)["code"] == "X143"  # names nobody else: the default
    run = uuid.uuid4().hex[:6].upper()
    u = _user()
    b = _run(monkeypatch, u, {f"nobody_{run}.pdf": _answer(f"NB-{run}", f"Nobody Supplier {run}", _vat(), bill_to="Nobody We Know Ltd", bill_vat="")}, f"Nobody {run}")
    with SessionLocal() as s:
        c = s.query(ApCapture).filter(ApCapture.batch_id == b["id"]).one()
        assert (c.company_code, c.bucket, c.stage, c.needs, c.sap_status) == ("", "other", "needs", "entity", "red")
        assert c.sap_issues == [{"level": "red", "text": AF.UNKNOWN_COMPANY}]
