"""1.0.89: the Site Schedules page as CBX-Connect 2 has it — the site's PDFs in, an agreement to check (fields, the
charge tables with every comment its own line, who signs), Word out; every read, change, Word file and download
on the agreement's history."""
import io

from fastapi.testclient import TestClient

from cbx.main import app

from ._swiftnet import REF, ntp, quote
from .conftest import _make, login


def _docx_text(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(" | ".join(c.text for c in r.cells) for t in d.tables for r in t.rows)


def _person(owner, email, name, tools):
    ids = [t["id"] for t in owner.get("/api/users").json()["tools"] if t["slug"] in tools]
    pw = _make(owner, email, name, "user", tools=ids)
    c = TestClient(app)
    login(c, email, pw)
    return c


def test_pdfs_in_agreement_out_and_every_step_on_its_history(owner):
    nobody = _person(owner, "noss@test.local", "No SS", [])
    assert nobody.get("/api/schedules").status_code == 403
    c = _person(owner, "sipho.ss@test.local", "Sipho SS", ["site-schedules"])
    assert c.post("/api/schedules", files=[("files", ("q.txt", b"hello", "text/plain"))]).status_code == 400
    a = c.post("/api/schedules", files=[("files", (REF + ".pdf", quote(), "application/pdf")), ("files", ("NTP Tiervlei.pdf", ntp(), "application/pdf"))]).json()
    f = a["fields"]
    assert f["site_name"] == "TIERVLEI EXCHANGE" and f["company"] == "Vodacom (Pty) Ltd" and f["contract_no"] == "70204797"
    assert f["co_site_name"] == "Beaconvale" and f["start_date"] == "01 August 2026" and f["co_code"] == "VDC"
    assert [x["kind"] for x in a["files"]] == ["quote", "ntp"] and not a["missing"]
    assert [r["item"] for r in a["tables"]["t1"]] == ["Mast Space", "Mast Space", "Ground Space", "Ground Space", "Ground Space"]
    assert a["values"]["t1_total"] == "R 6 305.34" and a["file_name"].startswith("9130 - 5183 - 70204797 - TIERVLEI EXCHANGE")
    # fix a line, pick who signs, make the Word file
    a["tables"]["t1"][0]["desc"] = "Three panel antennas, north leg"
    a["signing"]["swift"]["people"] = [{"name": "Gert Vorster", "title": "CFO for Swiftnet"}]
    r = c.post(f"/api/schedules/{a['id']}", json={"tables": a["tables"], "signing": a["signing"], "fields": {"swiftnet_rep": "Gert Vorster"}}).json()
    assert r["tables"]["t1"][0]["desc"] == "Three panel antennas, north leg"
    b = c.post(f"/api/schedules/{a['id']}/build").json()
    assert b["status"] == "built" and b["agreement"]["name"] == a["file_name"]
    doc = c.get(f"/api/schedules/files/{b['agreement']['id']}")
    assert doc.status_code == 200
    text = _docx_text(doc.content)
    for want in ("SITE SCHEDULE AGREEMENT", "Renewal", "TIERVLEI EXCHANGE", "Three panel antennas, north leg", "Gert Vorster", "CFO: Swiftnet",  # 1.0.107: the fixed three
                 "Total Rental Charges Payable (Excl. VAT) per Month | R 6 305.34"):
        assert want in text, want
    c.get(f"/api/schedules/files/{b['agreement']['id']}")
    d = c.get(f"/api/schedules/{a['id']}").json()
    whats = [h["what"] for h in d["history"]]
    assert d["downloads"] == 2 and sum(w == "Downloaded the Word file" for w in whats) == 2
    assert any(w.startswith("Made the Word file") for w in whats) and any(w.startswith("Changed") for w in whats)
    assert any(w.startswith("Added") and "read 2 PDFs" in w for w in whats)
    assert c.get("/api/schedules").json()["agreements"][0]["downloads"] == 2
    # someone else with the tool sees only their own; an admin sees all
    other = _person(owner, "thandi.ss@test.local", "Thandi SS", ["site-schedules"])
    assert other.get("/api/schedules").json()["agreements"] == []
    assert other.get(f"/api/schedules/{a['id']}").status_code == 404
    assert other.get(f"/api/schedules/files/{b['agreement']['id']}").status_code == 404
    assert any(x["id"] == a["id"] for x in owner.get("/api/schedules").json()["agreements"])
    # put away: gone from the list
    assert c.post(f"/api/schedules/{a['id']}/archive").json()["ok"]
    assert c.get(f"/api/schedules/{a['id']}").status_code == 404


def test_template_versions_and_people_who_sign(owner):
    from cbx.schedules import template as T

    c = _person(owner, "lebo.ss@test.local", "Lebo SS", ["site-schedules"])
    t = c.get("/api/schedules/template").json()
    assert t["current"] == 0 and t["template"]["blocks"]
    tpl = T.default()
    tpl["blocks"][0] = {**tpl["blocks"][0], "text": "SITE SCHEDULE AGREEMENT (TEST WORDING)"}
    assert c.post("/api/schedules/template", json={"template": tpl, "note": "test"}).status_code == 403  # admins change it
    assert owner.post("/api/schedules/template", json={"template": tpl, "note": "test"}).json()["current"] == 1
    assert c.post("/api/schedules/template/try", json={"template": tpl}).content.startswith(b"PK")
    assert owner.post("/api/schedules/template/use", json={"n": 0}).json()["current"] == 2
    p = c.get("/api/schedules/people").json()
    assert any(x["name"] == "Gert Vorster" for x in p["swift"])
    assert c.post("/api/schedules/people", json={"swift": [{"name": "Bad", "sig": "javascript:x"}], "couser": []}).status_code == 400
    assert c.post("/api/schedules/people", json={"swift": [{"name": "New Person", "title": "CFO"}], "couser": []}).json()["swift"][0]["name"] == "New Person"


def test_the_menu_opens_the_page():
    from cbx.tools import REGISTRY

    ss = next(t for t in REGISTRY if t["slug"] == "site-schedules")
    assert ss["admin_only"] is False  # people are given it, like the other tools
    js = open("web/views/index.js", encoding="utf-8").read()
    assert "'site-schedules': siteSchedules" in js


def test_the_tool_folders_rows_come_to_the_page_once(admin, db):
    from cbx import state
    from cbx.models import AppState
    from cbx.models.schedules import SsAgreement
    from cbx.schedules import copy_in

    base = "/api/engine/tools/site-schedules/data"
    r = admin.post(f"{base}/quotes", json={"values": {}}).json()
    admin.post(f"{base}/quotes/{r['id']}/files/quote", files=[("files", (REF + ".pdf", quote(), "application/pdf"))])
    admin.patch(f"{base}/quotes/{r['id']}", json={"values": {"co_site_name": "Typed by hand", "swift_signers": "Gert Vorster, CFO for Swiftnet"}})
    db.query(AppState).filter(AppState.key == copy_in.KEY_PAGE).delete()
    db.commit()
    got = copy_in.folder_to_page(db)
    db.commit()
    assert got["agreements"] >= 1
    a = db.query(SsAgreement).order_by(SsAgreement.id.desc()).first()
    assert a.fields["site_name"] == "TIERVLEI EXCHANGE" and a.fields["co_site_name"] == "Typed by hand"
    assert a.signing["swift"]["people"][0]["name"] == "Gert Vorster" and len(a.tables["t1"]) == 5
    assert any("Copied from the old Site Schedules table" in h["what"] for h in a.history)
    assert copy_in.folder_to_page(db) == {} and state.get(db, copy_in.KEY_PAGE)


def test_witnesses_left_and_the_people_who_sign_right_as_the_real_agreement():
    """12 Oct: witnesses on the left; on the right each person's line to sign on, the name, the title, 'For Swiftnet
    duly authorized' — side by side, as the real Site Schedule Agreement."""
    import docx

    from cbx.schedules import template as T

    ag = T.sample_agreement()
    ag["signing"]["swift"]["people"] = [{"name": "Gert Vorster", "title": "CFO for Swiftnet"}, {"name": "Paul Jansen van Rensburg", "title": "CEO for Swiftnet"}]
    ag["signing"]["swift"]["witnesses"] = ["Thandi Witness", ""]
    d = docx.Document(io.BytesIO(T.to_docx(T.default(), ag)))
    t = next(t for t in d.tables if any("Gert Vorster" in c.text for c in t.rows[0].cells))
    left, right = t.rows[0].cells
    assert "Thandi Witness" in left.text and "Gert Vorster" not in left.text
    r = [p.text for p in right.paragraphs if p.text]
    i = r.index("Gert Vorster")
    # 1.0.107: Swiftnet's signatories are the template's fixed three (the agreement's own list is not used for Swiftnet)
    assert r[i - 1].startswith("…") and r[i + 1] == "CFO: Swiftnet" and r[i + 2] == "For Swiftnet duly authorized"
    assert "Paul Jansen Van Rensburg" in r and r[1] == "Eddy Kgomongwe"
    co = next(t for t in d.tables if any("For the Co-User duly authorized" in c.text for c in t.rows[0].cells))
    assert "For the Co-User duly authorized" in co.rows[0].cells[1].text


# ---- 1.0.107: the Swiftnet team's test feedback (20 Oct) ---------------------------------------------------------
def _bold_values(doc, number: str) -> tuple[str, str]:
    """The paragraph numbered `number` in the Word file: (its text, the text of its bold runs)."""
    p = next(p for p in doc.paragraphs if p.text.startswith(number + "\t"))
    return p.text, "".join(r.text for r in p.runs if r.bold)


def test_first_escalation_and_bold_values(owner):
    from cbx.schedules import template as T

    def esc(start, **f):
        return T.values({"fields": {"start_date": start, "term_years": 5, **f}})["esc_date"]

    assert esc("15 January 2027") == "01 April 2027"  # 01 Jan–31 Mar: 01 April the same year
    assert esc("31 March 2027") == "01 April 2027"
    assert esc("01 April 2027") == "01 April 2028"  # after 31 March: as before, the next year
    assert esc("01 August 2026") == "01 April 2027"
    assert esc("15 January 2027", esc_year=2029) == "01 April 2029"  # a person's own date wins
    # 4.1–4.3 print their values in bold, in Word and on the page
    import docx

    d = docx.Document(io.BytesIO(T.to_docx(T.default(), {"fields": {"start_date": "15 January 2027", "term_years": 5}})))
    assert _bold_values(d, "4.1.") == ("4.1.\tEscalation Frequency: Annually", "Annually")
    assert _bold_values(d, "4.2.") == ("4.2.\tRate: CPI", "CPI")
    assert _bold_values(d, "4.3.") == ("4.3.\tFirst Escalation: 01 April 2027", "01 April 2027")
    assert "**{esc_date}**" in next(b["text"] for b in owner.get("/api/schedules/template").json()["template"]["blocks"] if b.get("number") == "4.3.")
    # read from a quote, then the start date changed: the first escalation follows, unless a person set their own
    from ._swiftnet import quote

    a = owner.post("/api/schedules", files=[("files", ("q.pdf", quote(start="15 January 2027"), "application/pdf"))]).json()
    assert a["fields"]["start_date"] == "15 January 2027" and a["fields"]["esc_year"] == 2027 and a["values"]["esc_date"] == "01 April 2027"
    r = owner.post(f"/api/schedules/{a['id']}", json={"fields": {**a["fields"], "start_date": "01 May 2027"}}).json()
    assert r["values"]["esc_date"] == "01 April 2028"
    r = owner.post(f"/api/schedules/{a['id']}", json={"fields": {"esc_year": "2031"}}).json()
    r = owner.post(f"/api/schedules/{a['id']}", json={"fields": {"start_date": "10 February 2027"}}).json()
    assert r["values"]["esc_date"] == "01 April 2031"


def test_saved_template_gets_the_new_standard_words(owner):
    """A template saved before 1.0.107 with the standard 4.1–4.3 and signing words gets the bold values and Pretoria."""
    from cbx.schedules import template as T

    old = T.default()
    old.pop("swift_sign")
    for b in old["blocks"]:
        for new, was in ((T.UPGRADES[k], k) for k in T.UPGRADES):
            if b.get("text") == new:
                b["text"] = was
    old["blocks"][0]["text"] = "SITE SCHEDULE AGREEMENT (OURS)"
    assert owner.post("/api/schedules/template", json={"template": old, "note": "before 1.0.107"}).status_code == 200
    tpl = owner.get("/api/schedules/template").json()["template"]
    texts = [b.get("text") for b in tpl["blocks"]]
    assert "Rate: **CPI**" in texts and texts[0] == "SITE SCHEDULE AGREEMENT (OURS)" and any("{sign_city_swift|Pretoria}" in (t or "") for t in texts)
    assert tpl["swift_sign"]["city"] == "Pretoria" and len(tpl["swift_sign"]["people"]) == 3
    owner.post("/api/schedules/template/use", json={"n": 0})


def test_swiftnet_signs_in_pretoria_by_its_three_signatories(owner):
    import docx

    from ._swiftnet import REF, ntp, quote

    a = owner.post("/api/schedules", files=[("files", (REF + ".pdf", quote(), "application/pdf")), ("files", ("NTP.pdf", ntp(), "application/pdf"))]).json()
    assert a["values"]["sign_city_swift"] == "Pretoria"
    owner.post(f"/api/schedules/{a['id']}", json={"signing": {"swift": {"city": "Centurion", "people": [{"name": "Someone Else", "title": "X"}]},
                                                              "couser": {"city": "Midrand", "people": [{"name": "Eugene Roets", "title": "Financial Director"}]}}})
    b = owner.post(f"/api/schedules/{a['id']}/build").json()
    d = docx.Document(io.BytesIO(owner.get(f"/api/schedules/files/{b['agreement']['id']}").content))
    paras = [p.text for p in d.paragraphs]
    assert any(p.startswith("Signed at Pretoria on") for p in paras) and not any("Centurion" in p for p in paras)
    sign = [t for t in d.tables if any("For Swiftnet duly authorized" in c.text for r in t.rows for c in r.cells)][0]
    left, right = sign.rows[0].cells
    lines = [p.text for p in right.paragraphs if p.text.strip()]
    assert lines == ["…………………………………", "Eddy Kgomongwe", "CCO: Swiftnet", "For Swiftnet duly authorized",
                     "…………………………………", "Gert Vorster", "CFO: Swiftnet", "For Swiftnet duly authorized",
                     "…………………………………", "Paul Jansen Van Rensburg", "CEO: Swiftnet", "For Swiftnet duly authorized"]
    assert "Someone Else" not in "\n".join(paras) + _docx_text(owner.get(f"/api/schedules/files/{b['agreement']['id']}").content)
    assert [p.text for p in left.paragraphs if p.text.strip()][:1] == ["…………………………………"]  # the witness line on the left
    text = _docx_text(owner.get(f"/api/schedules/files/{b['agreement']['id']}").content)
    assert "Signed at Midrand on" in text and "Eugene Roets" in text and "For the Co-User duly authorized" in text and "As witness:" in text
    # an admin changes them in the template: no code
    tpl = owner.get("/api/schedules/template").json()["template"]
    tpl["swift_sign"] = {"city": "Centurion", "people": [{"name": "New Chief", "title": "CEO: Swiftnet"}]}
    c = _person(owner, "naledi.ss@test.local", "Naledi SS", ["site-schedules"])
    assert c.post("/api/schedules/template", json={"template": tpl}).status_code == 403
    assert owner.post("/api/schedules/template", json={"template": {**tpl, "swift_sign": {"city": "X", "people": [{"name": str(i)} for i in range(9)]}}}).status_code == 400
    assert owner.post("/api/schedules/template", json={"template": tpl, "note": "board change"}).status_code == 200
    b = owner.post(f"/api/schedules/{a['id']}/build").json()
    text = _docx_text(owner.get(f"/api/schedules/files/{b['agreement']['id']}").content)
    assert "Signed at Centurion on" in text and "New Chief" in text and "Eddy Kgomongwe" not in text
    assert owner.get(f"/api/schedules/{a['id']}").json()["values"]["sign_city_swift"] == "Centurion"
    owner.post("/api/schedules/template/use", json={"n": 0})


def test_standby_at_nothing_is_0_amp_and_na():
    from cbx.schedules import reader as R

    from ._swiftnet import quote

    t2 = R.read_pdfs([("q.pdf", quote(standby="R 0.00"))])["tables"]["t2"]
    assert t2[1] == {"item": "Standby Electricity", "req": "0 Amp", "desc": "N/A", "amt": "R 0.00"}
    assert t2[0]["item"] == "Electricity" and t2[0]["req"] == "40 Amp" and t2[0]["desc"] != "N/A"  # the metered line keeps its own
    assert R.read_pdfs([("q.pdf", quote())])["tables"]["t2"][1]["req"] == "40 Amp"  # a standby charge: as before


def test_odd_spacing_on_quotes_still_gives_the_fields():
    from cbx.schedules import reader as R

    from ._swiftnet import _pdf, quote

    for odd in ("cells", "tight", "split"):
        f = R.read_pdfs([("q.pdf", quote(odd=odd))])["fields"]
        assert (f["lat"], f["lon"], f["pid"], f["region"]) == ("-26.13194", "28.203140", "7295", "Northern"), (odd, f)
        assert f["site_name"] == "TIERVLEI EXCHANGE" and f["struct_cat"] == "Sectional Mast"
    f = R.read_pdfs([("q.pdf", quote())])["fields"]
    assert (f["lat"], f["lon"], f["pid"], f["region"]) == ("33 55 37", "18 35 19", "5431", "Western")  # the usual layout: as before

    def only_structure_id(c):
        c.drawString(40, 800, "Quotation for Co-location")
        c.drawString(40, 780, "Structure ID   :")
        c.drawString(200, 780, "7295")
        c.drawString(40, 767, "Latitude")
        c.drawString(40, 755, "-26.5")
        c.drawString(120, 767, "Structure Name : TVI: Vodacom Mast")
    r = R.read_pdfs([("q.pdf", _pdf(only_structure_id))])
    assert r["fields"]["pid"] == "7295" and any("Structure ID" in n for n in r["notes"])
    assert r["fields"]["lat"] == ""  # no colon after the label: not taken as a label (and nothing below is taken by guess)


def test_wrapped_descriptions_stay_on_their_own_line():
    """The screenshot: "N/A", "&" and "Rf cable trenching." on three charge lines. Each line keeps its own description."""
    from cbx.schedules import reader as R

    from ._swiftnet import agreement, quote

    rows = [("Mast Space", "1.20m²", ["Three panel antennas at 36m"], "R 1 500.00"),
            ("Ground Space", "25.00m²", ["Power cable", "&", "Rf cable trenching."], "R 2 485.25"),
            ("Ground Space", "12.00m²", ["Container plinth inside the", "Swiftnet yard"], "R 1 192.92"),
            ("Ground Space", "4.80m²", ["N/A"], "R 477.17")]
    for align in ("top", "middle"):  # an earlier agreement's table, cells at the top of the row (Word's default) or in the middle
        t1 = R.read_pdfs([("a.pdf", agreement(rows, align=align))])["tables"]["t1"]
        assert [(r["item"], r["req"], r["desc"], r["amt"]) for r in t1] == [
            ("Mast Space", "1.20m²", "Three panel antennas at 36m", "R 1 500.00"), ("Ground Space", "25.00m²", "Power cable & Rf cable trenching.", "R 2 485.25"),
            ("Ground Space", "12.00m²", "Container plinth inside the Swiftnet yard", "R 1 192.92"), ("Ground Space", "4.80m²", "N/A", "R 477.17")], align

    def ground(comments):
        return [r["desc"] for r in R.read_pdfs([("q.pdf", quote(ground_comments=comments))])["tables"]["t1"] if r["item"] == "Ground Space"]

    # a quote's comments: "N/A" holds its item's place (the item keeps its own name), "&" and a fragment join their comment
    assert ground(["N/A", "70205157 Power cable &", "Rf cable trenching.", "70204797 Container plinth"]) == [
        "Ground Space", "70205157 Power cable & Rf cable trenching.", "70204797 Container plinth"]
    assert ground(["1) N/A", "2) 70205157* Power cable (40m)", "&", "Rf cable trenching.", "3) 70204797* Container plinth"]) == [
        "Ground Space", "70205157* Power cable (40m) & Rf cable trenching.", "70204797* Container plinth"]
    # still: every numbered comment its own line (Edwin, 9 Oct), and N/A never a line of its own
    assert ground(["1) A for item one", "2) B for item two", "3) C for item three", "4) D an extra note", "5) N/A"]) == [
        "A for item one", "B for item two", "C for item three", "D an extra note"]
    assert R.glue_fragments(["Power cable", "&", "Rf cable trenching.", "2) Next one", "and its tail"]) == ["Power cable & Rf cable trenching.", "2) Next one and its tail"]


def test_scope_from_the_sow_then_the_ntp_then_the_quote(owner):
    from cbx.schedules import reader as R

    from ._swiftnet import REF, ntp, quote, sow

    q, n, s = (REF + ".pdf", quote()), ("NTP Tiervlei.pdf", ntp(scope="Install 2 x RRU and one 0.6m dish at 30m")), ("Tiervlei SOW.pdf", sow())
    r = R.read_pdfs([q, n, s])
    assert r["fields"]["scope"] == ("Install 3 x panel antennas at 36m on the north leg and one 0.6m dish at 30m. "
                                    "Run fibre & RF cables in the existing trench.") and r["fields"]["scope_src"] == "sow"
    assert [x["kind"] for x in r["files"]] == ["quote", "ntp", "sow"] and not any("Scope" in x for x in r["notes"])
    r = R.read_pdfs([q, n])
    assert r["fields"]["scope"] == "Install 2 x RRU and one 0.6m dish at 30m" and r["fields"]["scope_src"] == "ntp"
    r = R.read_pdfs([q, ("NTP.pdf", ntp())])
    assert r["fields"]["scope_src"] == "quote" and "Three panel antennas" in r["fields"]["scope"]
    assert any(x.startswith("Scope of approved work: no SOW or NTP") for x in r["notes"])  # marked for checking
    assert R.classify(R.model_from_pdf(sow(), "scope.pdf")) == "sow"
    # the page takes a SOW beside the quote and the NTP; clause 7 prints its scope
    a = owner.post("/api/schedules", files=[("files", (n[0], n[1], "application/pdf")), ("files", (q[0], q[1], "application/pdf")),
                                            ("files", (s[0], s[1], "application/pdf"))]).json()
    assert [x["kind"] for x in a["files"]] == ["ntp", "quote", "sow"] and a["fields"]["scope"].startswith("Install 3 x panel")
    b = owner.post(f"/api/schedules/{a['id']}/build").json()
    assert "Run fibre & RF cables in the existing trench." in _docx_text(owner.get(f"/api/schedules/files/{b['agreement']['id']}").content)


# ---- 1.0.108: the T100/T000 lines; clause 7 sub-headings only with text under them ----------------------------------
def _paras(data: bytes) -> list[str]:
    import docx

    return [p.text for p in docx.Document(io.BytesIO(data)).paragraphs]


def _ag(**fields):
    from cbx.schedules import template as T

    ag = T.sample_agreement()
    ag["fields"].update(fields)
    return ag


def test_reference_lines_always_read_t100_and_t000_slash_number():
    from cbx.schedules import reader as R
    from cbx.schedules import template as T

    for v in ("T100/9130", "T100 9130", "t100-9130", "9130", " T100 / 9130 ", "T000/9130", "T100:9130"):
        assert R.ref_text("t100_ref", v) == "T100/9130", v
        assert R.ref_text("t000_ref", v) == "T000/9130", v  # a value with the other prefix gets its own field's prefix
    for v in ("", None, "T100/", "T000 /", "abc"):
        assert R.ref_text("t100_ref", v) == "" and R.ref_number(v) == "", v
    assert R.fix_refs({"t100_ref": "T000 5183", "t000_ref": "t000-9130", "x": "y"}) == {"t100_ref": "T100/5183", "t000_ref": "T000/9130", "x": "y"}
    # the Word lines under Contract no., from every shape a person or an older version saved
    for t100, t000 in (("T100 9130", "5183"), ("t100-9130", "T000/5183"), ("9130", "T100/5183"), ("T100/9130", "T000 - 5183")):
        paras = _paras(T.to_docx(T.default(), _ag(t100_ref=t100, t000_ref=t000)))
        i = paras.index("Contract no.: 70204797")
        assert paras[i + 1:i + 3] == ["T100/9130", "T000/5183"], (t100, t000, paras[i:i + 3])
        assert T.values(_ag(t100_ref=t100, t000_ref=t000))["t100_ref"] == "T100/9130"
    # no number (a new application): the lines are left out
    paras = _paras(T.to_docx(T.default(), _ag(lease_type="New Application", contract_no="", t100_ref="", t000_ref="T000/")))
    i = paras.index("Contract no.: New Application")
    assert not any(p.startswith(("T100", "T000")) for p in paras) and "CO-USER SITE DETAILS" in paras[i + 1]
    # only one of them: that one, in its own place
    paras = _paras(T.to_docx(T.default(), _ag(t100_ref="", t000_ref="5183")))
    i = paras.index("Contract no.: 70204797")
    assert paras[i + 1] == "T000/5183" and "T100/" not in "\n".join(paras)


def test_the_reader_never_swaps_t100_and_t000():
    from cbx.schedules import reader as R

    from ._swiftnet import REF, _pdf, quote

    # Swiftnet's file name: T100 - T000 - contract
    f = R.read_pdfs([(REF + ".pdf", quote())])["fields"]
    assert (f["t100_ref"], f["t000_ref"]) == ("T100/9130", "T000/5183")
    # the cause: an agreement with a T000 number only was named "5183 - 70204797 - …", and reading that file back made
    # 5183 the T100. One number alone now keeps its prefix in the file name, and a prefixed number is read as that one.
    assert R.parse_ref("T000-5183 - 70204797 - TIERVLEI EXCHANGE - 01.08.2026 - 31.07.2031 - 7708 - R - VDC") == {
        "t000": "5183", "start": [1, 8, 2026], "end": [31, 7, 2031], "col_id": "7708", "lease_code": "R", "co_code": "VDC"}
    assert R.parse_ref("T100-9130 - 70204797 - SITE")["t100"] == "9130" and "t000" not in R.parse_ref("T100-9130 - 70204797 - SITE")

    # an earlier agreement: the lines under Contract no., however they were typed
    def earlier(lines):
        def d(c):
            c.drawString(200, 800, "SITE SCHEDULE AGREEMENT")
            c.drawString(250, 786, "Renewal")
            c.drawString(40, 760, "Contract no.: 70204797")
            for k, s in enumerate(lines):
                c.drawString(60, 748 - k * 12, s)
        return _pdf(d)
    for lines in (["T100/9130", "T000/5183"], ["T100 / 9130", "T000-5183"], ["t100 9130", "T000 / 5183"]):
        f = R.read_pdfs([("old agreement.pdf", earlier(lines))])["fields"]
        assert (f["t100_ref"], f["t000_ref"], f["contract_no"]) == ("T100/9130", "T000/5183", "70204797"), lines
    f = R.read_pdfs([("old agreement.pdf", earlier(["T000/5183"]))])["fields"]
    assert (f["t100_ref"], f["t000_ref"]) == ("", "T000/5183")  # never the T000 number in T100

    # a quote whose Contract No is the T000 reference: it is the T000 line, never the T100
    def q(c):
        c.drawString(40, 800, "Quotation for Co-location")
        c.drawString(40, 780, "Contract No :")
        c.drawString(120, 780, "T000-5183")
        c.drawString(40, 767, "Site Name :")
        c.drawString(120, 767, "TIERVLEI EXCHANGE")
    f = R.read_pdfs([("q.pdf", _pdf(q))])["fields"]
    assert (f["t100_ref"], f["t000_ref"]) == ("", "T000/5183")


def test_the_form_keeps_the_prefixes_fixed(owner):
    from cbx.db import SessionLocal
    from cbx.models.schedules import SsAgreement

    from ._swiftnet import REF, quote

    a = owner.post("/api/schedules", files=[("files", (REF + ".pdf", quote(), "application/pdf"))]).json()
    assert [f for s in owner.get("/api/schedules").json()["form"] for f in s["fields"] if f["key"] in ("t100_ref", "t000_ref")] == [
        {"key": "t100_ref", "label": "T100 reference", "kind": "ref"}, {"key": "t000_ref", "label": "T000 reference", "kind": "ref"}]
    for t100, t000 in (("9130", "5183"), ("t100-9130", "T000 5183"), ("T000/9130", "T100/5183"), ("T100 9130", "T000/5183")):
        r = owner.post(f"/api/schedules/{a['id']}", json={"fields": {"t100_ref": t100, "t000_ref": t000}}).json()
        assert (r["fields"]["t100_ref"], r["fields"]["t000_ref"]) == ("T100/9130", "T000/5183"), (t100, t000)
        assert r["file_name"].startswith("9130 - 5183 - 70204797 - TIERVLEI EXCHANGE")
    # a value saved before (as typed, or by an older version) is shown, named and printed the fixed way
    with SessionLocal() as s:
        x = s.get(SsAgreement, a["id"])
        x.fields = {**x.fields, "t100_ref": "T100 9130", "t000_ref": "t000-5183"}
        s.commit()
    r = owner.get(f"/api/schedules/{a['id']}").json()
    assert (r["fields"]["t100_ref"], r["fields"]["t000_ref"], r["values"]["t000_ref"]) == ("T100/9130", "T000/5183", "T000/5183")
    b = owner.post(f"/api/schedules/{a['id']}/build").json()
    assert b["agreement"]["name"].startswith("9130 - 5183 - 70204797")
    paras = _paras(owner.get(f"/api/schedules/files/{b['agreement']['id']}").content)
    i = paras.index("Contract no.: 70204797")
    assert paras[i + 1:i + 3] == ["T100/9130", "T000/5183"]
    # one alone: its prefix in the file name, so reading the file back never makes it the other one
    r = owner.post(f"/api/schedules/{a['id']}", json={"fields": {"t100_ref": "", "t000_ref": "5183"}}).json()
    assert r["file_name"].startswith("T000-5183 - 70204797 - ") and r["fields"]["t100_ref"] == "" and "t100_ref" in r["missing"]
    # "T100/" without a number is not filled in; a new application needs neither, and prints neither
    r = owner.post(f"/api/schedules/{a['id']}", json={"fields": {"t100_ref": "T100/", "t000_ref": "T000/", "lease_type": "New Application"}}).json()
    assert (r["fields"]["t100_ref"], r["fields"]["t000_ref"]) == ("", "") and not {"t100_ref", "t000_ref"} & set(r["missing"])
    b = owner.post(f"/api/schedules/{a['id']}/build").json()
    assert not any(p.startswith(("T100", "T000")) for p in _paras(owner.get(f"/api/schedules/files/{b['agreement']['id']}").content))


CLAUSE7 = [("current_install", "Client current installation on site:"), ("scope", "Scope of work:"), ("equip_comments", "Equipment accommodation:"),
           ("scope", "Antenna and feeder work:")]


def _clause7(paras: list[str]) -> list[str]:
    i = next(k for k, p in enumerate(paras) if "SCOPE OF APPROVED WORK" in p)
    j = next(k for k, p in enumerate(paras) if "SPECIAL REQUESTS" in p)
    return paras[i:j]


def test_clause_7_sub_headings_only_with_text_under_them():
    import itertools

    from cbx.schedules import template as T

    texts = {"current_install": "Two panel antennas at 30m, one 0.6m dish.", "scope": "Install 3 x panel antennas at 36m.",
             "equip_comments": "Container plinth inside the yard."}
    for on in itertools.product((False, True), repeat=3):
        given = {k: (texts[k] if o else "") for k, o in zip(("current_install", "scope", "equip_comments"), on)}
        c7 = _clause7(_paras(T.to_docx(T.default(), _ag(lease_type="Renewal", **given))))
        assert c7[0] == "7.\tSCOPE OF APPROVED WORK:" and c7[1] == "Renewal:"  # the clause heading and the lease line stay
        want = []
        for key, head in CLAUSE7:
            if given[key]:
                want += [head, given[key]]
        assert [p for p in c7[2:] if p.strip()] == want, given
        assert "" not in c7[:2]


def test_clause_7_rule_in_a_saved_template_and_the_editor(owner):
    from cbx.schedules import template as T

    # a template saved before 1.0.108: the sub-headings had no "only when"; "Client current installation" had no line
    old = T.default()
    blocks = []
    for b in old["blocks"]:
        if b.get("text") == "{current_install}":
            continue
        if b.get("text") in T.SUBHEADS:
            b = {k: v for k, v in b.items() if k != "if"}
        blocks.append(b)
    old["blocks"] = blocks
    up = T.upgrade(old)
    heads = [b for b in up["blocks"] if b.get("text") in T.SUBHEADS]
    assert [b["if"] for b in heads] == ["current_install", "scope", "equip_comments", "scope"]
    assert T.upgrade(up)["blocks"] == up["blocks"]  # once only
    assert _clause7(_paras(T.to_docx(old, _ag(scope="", equip_comments="Plinth.", current_install=""))))[2:] == ["Equipment accommodation:", "Plinth."]
    # an admin's own choice stays theirs: words of their own under a heading, or their own "only when"
    mine = T.default()
    for b in mine["blocks"]:
        if b.get("text") == "**Scope of work:**":
            b.pop("if")
        if b.get("text") == "{scope}":
            b["text"] = "As per the attached SOW."
            b.pop("if")
            break
    assert "if" not in next(b for b in T.upgrade(mine)["blocks"] if b.get("text") == "**Scope of work:**")
    # the editor saves the rule and it holds; the field is on the form and is a token the editor offers
    tpl = owner.get("/api/schedules/template").json()
    assert any(t["key"] == "current_install" for t in tpl["tokens"])
    assert next(b for b in tpl["template"]["blocks"] if b.get("text") == "**Equipment accommodation:**")["if"] == "equip_comments"
    form = [f["key"] for s in owner.get("/api/schedules").json()["form"] for f in s["fields"]]
    assert "current_install" in form
    t2 = tpl["template"]
    next(b for b in t2["blocks"] if b.get("text") == "**Equipment accommodation:**")["if"] = "scope"  # an admin points it elsewhere
    assert owner.post("/api/schedules/template", json={"template": t2, "note": "mine"}).status_code == 200
    got = owner.get("/api/schedules/template").json()["template"]
    assert next(b for b in got["blocks"] if b.get("text") == "**Equipment accommodation:**")["if"] == "scope"
    owner.post("/api/schedules/template/use", json={"n": 0})


def test_file_name_lease_code_follows_the_lease_type(owner):
    """1.0.108: not a code read from the PDFs (the quote's name said R; the agreement is a New Application)."""
    from cbx.schedules import reader as R

    from ._swiftnet import REF, quote

    assert [R.lease_code_for(t) for t in ("New Application", "Renewal", "Re-Assessment", "New Re-Assessment", "MOU", "renewal")] == ["N", "R", "RA", "NA", "MOU", "R"]
    assert R.lease_code_for("Something else", "RV") == "RV" and R.lease_code_for("", "R") == "R"
    a = owner.post("/api/schedules", files=[("files", (REF + ".pdf", quote(), "application/pdf"))]).json()
    assert a["fields"]["lease_code"] == "R" and a["file_name"].endswith(" - R - VDC.docx")
    for lease, code in (("New Application", "N"), ("Renewal", "R"), ("Re-Assessment", "RA"), ("New Re-Assessment", "NA"), ("MOU", "MOU")):
        r = owner.post(f"/api/schedules/{a['id']}", json={"fields": {"lease_type": lease}}).json()
        assert r["file_name"].endswith(f" - {code} - VDC.docx"), (lease, r["file_name"])
    b = owner.post(f"/api/schedules/{a['id']}/build").json()
    assert b["agreement"]["name"].endswith(" - MOU - VDC.docx")
