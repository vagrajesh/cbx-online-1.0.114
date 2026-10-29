"""1.0.92: IFRS 16 is a page again, the same page as CBX-Connect 2 — the source files kept, a month worked out, the
ledger check (with the discount rates from the Sequentra validation report and the net asset check), LA vs RR06
(ABSA) with a drill-down and its own Excel, the Rates table; the tool folder's rows come to the page once."""
import io
import tempfile
import zipfile
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from cbx.main import app

from .conftest import _make, login
from .fixtures import make_ifrs16

FIX = Path(__file__).parent / "fixtures" / "ifrs16"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _person(owner, email, name, tools, role="user"):
    ids = [t["id"] for t in owner.get("/api/users").json()["tools"] if t["slug"] in tools]
    pw = _make(owner, email, name, role, tools=ids)
    c = TestClient(app)
    login(c, email, pw)
    return c


def _upload(c, la=None, mv=None, seq=None):
    files = {"la": (Path(la).name if la else "la_register.xlsx", Path(la or FIX / "la_register.xlsx").read_bytes(), XLSX),
             "movements": ("movements_rr06.xlsx", Path(mv or FIX / "movements_rr06.xlsx").read_bytes(), XLSX)}
    if seq:
        files["sequentra"] = ("sequentra.xlsx", Path(seq).read_bytes(), XLSX)
    r = c.post("/api/ifrs16/uploads", files=files)
    assert r.status_code == 200, r.text
    return next(u for u in c.get("/api/ifrs16").json()["uploads"] if u["id"] == r.json()["id"])


def _month(c, uid, month):
    r = c.post(f"/api/ifrs16/uploads/{uid}/month", json={"month": month})
    assert r.status_code == 200, r.text
    return next(m for m in c.get("/api/ifrs16").json()["months"] if m["id"] == r.json()["id"])


def test_files_in_a_month_out_and_only_your_own(owner):
    assert _person(owner, "noifrs@test.local", "No IFRS", []).get("/api/ifrs16").status_code == 403
    c = _person(owner, "lerato.ifrs@test.local", "Lerato Leases", ["ifrs16"])
    rich = make_ifrs16.make(Path(tempfile.mkdtemp()))
    u = _upload(c, rich["la"], rich["movements"], rich["sequentra"])
    assert u["status"] == "ready" and "2026-08" in [m["value"] for m in u["months"]], u
    assert c.post(f"/api/ifrs16/uploads/{u['id']}/month", json={"month": "1999-01"}).status_code == 400
    m = _month(c, u["id"], "2026-08")
    assert m["status"] == "done" and m["imports"] == 4 and m["report"]["name"] == "IFRS16_Report_2026-08.xlsx", m
    z = zipfile.ZipFile(io.BytesIO(c.get(f"/api/ifrs16/files/{m['workbooks']['id']}").content))
    piw = next(n for n in z.namelist() if n.startswith("PIW"))
    ws = load_workbook(io.BytesIO(z.read(piw)))["Schedule"]
    col = next(x.column for x in ws[1] if x.value == "Lease Start Date")
    assert not isinstance(ws.cell(2, col).value, str) and ws.cell(2, col).number_format == "m/d/yyyy;@"
    other = _person(owner, "bongi.ifrs@test.local", "Bongi B", ["ifrs16"])
    assert all(x["id"] != u["id"] for x in other.get("/api/ifrs16").json()["uploads"])
    assert other.get(f"/api/ifrs16/files/{m['report']['id']}").status_code == 404
    viewer = _person(owner, "vee.ifrs@test.local", "Vee Viewer", ["ifrs16"], role="viewer")
    assert viewer.post("/api/ifrs16/rates", json={"month": "2026-08", "a": 10}).status_code == 403


def _ledger_book():
    wb = Workbook()
    ws = wb.active
    ws.title = "Ledger Entries"
    ws.append(["Ledger Date", "Segment2", "Segment3", "Segment4", "DR", "CR", "Comments", "JEShortDesc", "Transactional Currency"])
    for r in [("2026-06-30", "BC000", "P1", "CC1", 12000, 0, "Schedule TXN00001 initial", "Initial Recognition", "ZAR"),
              ("2026-06-30", "HA101", "P1", "CC1", 0, 12000, "Schedule TXN00001 initial", "Initial Recognition", "ZAR"),
              ("2026-07-31", "HA101", "P1", "CC1", 1000, 0, "Schedule TXN00001 payment", "Payment", "ZAR"),
              ("2026-07-31", "SB424", "P1", "CC1", 0, 1000, "Schedule TXN00001 payment", "Payment", "ZAR")]:
        ws.append(list(r))
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_the_ledger_check_with_sequentra_rates_and_the_net_asset_check(owner):
    c = _person(owner, "ledger.ifrs@test.local", "Lebo Ledger", ["ifrs16"])
    seq = Workbook()
    seq.active.append(["Record ID", "Discount Rate"])
    seq.active.append(["TXN00001", 0.1125])
    sb = io.BytesIO()
    seq.save(sb)
    r = c.post("/api/ifrs16/ledger", files={"file": ("Ledger Export.xlsx", _ledger_book(), XLSX), "sequentra": ("Sequentra Validation.xlsx", sb.getvalue(), XLSX)},
               data={"as_at": "2026-08", "months_back": "12"})
    assert r.status_code == 200, r.text
    row = next(x for x in c.get("/api/ifrs16").json()["ledger"] if x["id"] == r.json()["id"])
    assert row["status"] == "done", row
    assert "Discount rates from: Sequentra Validation.xlsx" in row["overview"] and "Net asset check" in row["overview"]
    wb = load_workbook(io.BytesIO(c.get(f"/api/ifrs16/files/{row['report']['id']}").content))
    assert row["report"]["name"] == "Consolidated_Report_FULL_SA_August_2026.xlsx" and "Check Net Asset" not in wb.sheetnames
    assert row["simple"]["name"] == "Lease_Report_SIMPLE_August_2026.xlsx"
    heads = [x.value for r_ in wb["Data"].iter_rows(max_row=6) for x in r_ if isinstance(x.value, str)]
    assert any("NET ASSET" in h and "DEPRECIATION" in h for h in heads) and not any("NET ASSET_OPENNING" in h for h in heads)


def test_la_vs_rr06_with_a_drill_down_and_its_own_excel(owner):
    from datetime import datetime

    c = _person(owner, "recon.ifrs@test.local", "Rina Recon", ["ifrs16"])
    la = Workbook()
    ws = la.active
    ws.title = "Portfolio Transactions"
    ws.append(["Transaction #", "Deal Status", "Lessee", "Asset Location(s)"])
    for r in [("4041P001-01", "Booked", "ABSA Bank Limited", "South Africa"), ("4041P002-01", "Disposed (Modification)", "ABSA Bank Limited", "South Africa"),
              ("4041P003-01", "Booked", "ABSA Bank Limited", "South Africa")]:
        ws.append(list(r))
    rr = Workbook()
    ws = rr.active
    ws.title = "Contract List"
    ws.append(["Lease Reference", "Lease Marker", "Initial Lease Date", "Lease Expiry Date", "Current Lease Term", "Current Charges"])
    for r in [("4041P001", "Payable", "2020-01-01", "2030-12-31", 10, 1000), ("4041P002", "Payable", "2020-01-01", "2030-12-31", 10, 1000),
              ("4041R007", "Receivable", "2020-01-01", "2030-12-31", 10, 500), ("4041P006", "Payable", datetime(2027, 1, 1), "2030-12-31", 5, 500)]:
        ws.append(list(r))
    bl, br = io.BytesIO(), io.BytesIO()
    la.save(bl)
    rr.save(br)
    r = c.post("/api/ifrs16/recon", data={"period": "2026-09"}, files={"la": ("LA.xlsx", bl.getvalue(), XLSX), "rr06": ("RR06.xlsx", br.getvalue(), XLSX)})
    assert r.status_code == 200, r.text
    row = next(x for x in c.get("/api/ifrs16").json()["recons"] if x["id"] == r.json()["id"])
    assert row["status"] == "done", row
    d = c.get(f"/api/ifrs16/recon/{row['id']}").json()
    assert {x["Lease Reference"]: x["Category"] for x in d["rr06_not_la"]} == {"4041P002": "Exception", "4041R007": "Receivable", "4041P006": "Future acquisition"}
    assert [x["Lease Reference"] for x in d["la_not_rr06"]] == ["4041P003"]
    assert c.get(f"/api/ifrs16/files/{row['report']['id']}").content.startswith(b"PK")


def test_the_menu_opens_the_page_and_the_folder_is_hidden():
    import json

    js = open("web/app.js", encoding="utf-8").read()
    assert "ifrs16: 'ifrs16'" in js and "new Set(['todo'])" in js
    assert "ifrs16," in open("web/views/index.js", encoding="utf-8").read()
    assert json.load(open("tools/ifrs16/tool.json", encoding="utf-8"))["menu"]["hidden"] is True


def test_the_folders_rows_come_to_the_page_once(admin, db, tmp_path, monkeypatch):
    from cbx import state
    from cbx.config import ROOT
    from cbx.engine import files as FILES
    from cbx.engine import loader
    from cbx.engine import store as STORE
    from cbx.ifrs16 import copy_in
    from cbx.models import AppState
    from cbx.models.ifrs16_page import IfrsLedgerCheck, IfrsMonth, IfrsRate, IfrsUpload

    loader.reload(ROOT / "tools")
    monkeypatch.setattr(STORE, "path_for", lambda slug: tmp_path / f"{slug}.sqlite")
    monkeypatch.setattr(FILES, "root_for", lambda slug: tmp_path / "_files" / slug)
    base = "/api/engine/tools/ifrs16/data"
    up = admin.post(f"{base}/uploads", json={"values": {"months": "2026-08", "movement": "RR06", "status": "Ready", "summary": "1 LA row"}})
    assert up.status_code == 200, up.text
    uid = up.json()["id"]
    admin.post(f"{base}/uploads/{uid}/files/la", files=[("files", ("la_register.xlsx", (FIX / "la_register.xlsx").read_bytes(), XLSX))])
    admin.post(f"{base}/rates", json={"values": {"month": "2026-08", "a": 9.5, "b": 9.75}})
    for m in (IfrsMonth, IfrsLedgerCheck, IfrsUpload):
        db.query(m).delete()
    db.query(IfrsRate).delete()
    db.query(AppState).filter(AppState.key == copy_in.KEY_PAGE).delete()
    db.commit()
    got = copy_in.folder_to_page(db)
    db.commit()
    assert got["uploads"] >= 1 and got["rates"] >= 1, got
    u = db.query(IfrsUpload).order_by(IfrsUpload.id.desc()).first()
    assert u.status == "ready" and u.scan["months"][0]["value"] == "2026-08" and u.la_file_id
    assert db.query(IfrsRate).filter(IfrsRate.month == "2026-08").one().a == 9.5
    assert copy_in.folder_to_page(db) == {} and state.get(db, copy_in.KEY_PAGE)


# ---- 1.0.93: as the desktop — the ledger report's inputs and outputs, the master report, the Sequentra test --------
def test_the_ledger_report_and_the_master_report_as_the_desktop(owner):
    c = _person(owner, "desk.ifrs@test.local", "Desk Person", ["ifrs16"])
    rich = make_ifrs16.make(Path(tempfile.mkdtemp()))
    u = _upload(c, rich["la"], rich["movements"], rich["sequentra"])
    assert _month(c, u["id"], "2026-08")["status"] == "done"
    r = c.post("/api/ifrs16/ledger", files={"file": ("Ledger Export.xlsx", _ledger_book(), XLSX)},
               data={"as_at": "2026-06", "months_back": "12", "upload_id": str(u["id"]), "entity": "ABSA Bank Limited", "span": "from"})
    assert r.status_code == 200, r.text
    row = next(x for x in c.get("/api/ifrs16").json()["ledger"] if x["id"] == r.json()["id"])
    assert row["status"] == "done", row
    assert "Periods: from the as-at month on" in row["overview"] and "Discount rates from: sequentra.xlsx" in row["overview"]
    full = load_workbook(io.BytesIO(c.get(f"/api/ifrs16/files/{row['report']['id']}").content))
    assert "Recon" in full.sheetnames and "ABSA Bank Limited" in " ".join(str(x.value) for r_ in full["Recon"].iter_rows(max_row=6) for x in r_ if x.value)
    assert load_workbook(io.BytesIO(c.get(f"/api/ifrs16/files/{row['simple']['id']}").content)).sheetnames[0] == "Start here"
    m = c.post(f"/api/ifrs16/uploads/{u['id']}/master", data={"month": "2026-08", "as_at": "2026-08", "entity": "ABSA Bank Limited"},
               files={"ledger": ("Ledger Export.xlsx", _ledger_book(), XLSX)})
    assert m.status_code == 200, m.text
    mr = next(x for x in c.get("/api/ifrs16").json()["masters"] if x["id"] == m.json()["id"])
    assert mr["status"] == "done" and mr["report"]["name"] == "IFRS16Workflow_Master_Report_August2026.xlsx", mr
    names = load_workbook(io.BytesIO(c.get(f"/api/ifrs16/files/{mr['report']['id']}").content), read_only=True).sheetnames
    assert names[0] == "Contents" and "Report Recon" in names and "Recon Summary" in names


def test_the_sequentra_test_calls_a_report(owner, monkeypatch):
    import json

    from cbx.ifrs16 import ENGINE_DIR  # noqa: F401 — puts the engine on the path
    import sequentra_api as sq

    asked = []

    def fake_get(url, header, key, timeout):
        asked.append(url)
        return 200, "application/json", json.dumps({"value": [{"Property_ID": "9713P001", "Name": "A lease"}]}).encode()

    monkeypatch.setattr(sq, "_http_get", fake_get)
    key = "TEST-CLIENTKEY-0093"
    assert owner.post("/api/ifrs16/connection", json={"base_url": "https://api.excellerate.tech/sequentra/", "header_name": "Clientkey", "secret": key, "is_active": True}).status_code == 200
    r = owner.post("/api/ifrs16/connection/test").json()
    assert r["ok"] and "Connected" in r["note"] and asked[-1].rstrip("/").endswith("buildingrecordtabs") and key not in json.dumps(r), r


# ---- the AI check, and .xls exports that are really web pages ------------------------------------------------
def test_the_ai_check_names_the_problem_and_what_to_do():
    from cbx import aicheck

    c = aicheck.check("IFRS 16: reading", "The LA register could not be read: `Import xlrd` failed. Install xlrd >= 2.0.1 for xls Excel support")
    assert c["fix"] == "install:xlrd" and "xlrd" in c["words"]  # here xlrd is there already, so it counts as fixed
    n = aicheck.note("IFRS 16: reading", "File is locked: Permission denied")
    assert n.startswith("File is locked") and "\nAI check: The file is open somewhere else" in n
    assert aicheck.check("x", "something nobody has seen")["words"] == ""  # stub mode: no model, no words
    calls = []
    def flaky():
        calls.append(1)
        if len(calls) == 1:
            raise ValueError("No module named 'xlrd'")
        return "ok"
    assert aicheck.with_retry(flaky, "x") == "ok" and len(calls) == 2  # the reader is there now: tried once more


def test_an_xls_that_is_really_a_web_page_is_read_and_a_failed_upload_can_be_read_again(owner):
    CLIENT = _person(owner, "xls.ifrs@test.local", "Xavier Xls", ["ifrs16"])
    import tempfile as _t
    d = Path(_t.mkdtemp())
    rows = "".join(f"<tr><td>TXN000{i}</td><td>Company {i}</td><td>ZAR</td><td>{i}000</td></tr>" for i in (1, 2, 3))
    (d / "Transaction-export.xls").write_text("<html><body><table><tr><td>Landlord Advantage - Portfolio Transactions</td></tr>"
                                              "<tr><th>Transaction #</th><th>Company Name</th><th>Currency</th><th>Current Payment (Transactional)</th></tr>" + rows + "</table></body></html>", encoding="utf-8")
    u = _upload(CLIENT, d / "Transaction-export.xls")
    assert u["status"] == "ready", u
    (d / "broken.xlsx").write_bytes(b"not a workbook")
    bad = _upload(CLIENT, d / "broken.xlsx")
    assert bad["status"] == "failed" and "AI check:" in bad["note"], bad
    r = CLIENT.post(f"/api/ifrs16/uploads/{bad['id']}/read")
    assert r.status_code == 200, r.text
