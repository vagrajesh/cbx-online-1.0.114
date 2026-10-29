"""1.0.98: from CBX-Connect 2.10.3 tests/test_fix210_ap.py, on the old conftest.

2.10.0 review fixes, AP invoice capture (group ap210): one test per risk the review found, each the flipped probe."""
from __future__ import annotations

import io
import json
import os
import random
import re
import subprocess
import sys
import threading
import uuid
import zipfile
from datetime import timedelta
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import httpx
import pymupdf
import pytest
from openpyxl import Workbook, load_workbook

from cbx.ap import capture as C
from cbx.ap import masters as M
from cbx.ap import sap as P
from cbx.config import settings
from cbx.db import SessionLocal, utcnow
from cbx.models import User
from cbx.models.ap import ApCapture, ApEntity, ApItem, ApUtilityAccount, ApVendorCode
from cbx.models.connections import ApiConnection

from .fixtures import make_ap
from .test_ap_capture_1098 import ESKOM, _admin, _drop, _invoice_answer, _sheet, _tables, _user, _vision, _xlsx, eskom_pdf, make_person, utilities, _app  # noqa: F401
from .test_ap_capture_1098 import signed_in  # 1.0.104: a second admin

FULL = "CBRE Excellerate Facilities Management"


@pytest.fixture(autouse=True)
def _no_connections_left():
    """1.0.98: the suite shares one database: the SAP connections these tests make are removed after each test."""
    yield
    from cbx.models.connections import ApiConnection as _Conn

    with SessionLocal() as s:
        s.query(_Conn).filter(_Conn.slug.like("fix-sap-%")).delete(synchronize_session=False)
        s.commit()


@pytest.fixture(autouse=True)
def _chart_of_accounts():
    """2.10.3: every G/L is checked against the chart of accounts: the made-up G/Ls here are put in it first."""
    make_ap.chart()


@pytest.fixture(autouse=True)
def _bill_to_taught():
    """The invoices below are addressed to X143's full name: taught once, so the bill-to is a sure match (SAP's cut name
    alone is only the start of it, which is yellow since 2.10.0)."""
    with SessionLocal() as s:
        e = s.query(ApEntity).filter(ApEntity.code == "X143").one()
        if FULL not in (e.aliases or ""):
            e.aliases = (e.aliases + "\n" if e.aliases else "") + FULL
            s.commit()


def _vat(lead="4") -> str:
    return lead + "".join(random.choice("0123456789") for _ in range(9))


def _vendor(adm, no, name, vat):
    r = adm.post("/api/ap/masters/vendors", json={"vendor_no": no, "name": name, "vat_no": vat, "gl": "6100100", "cost_centre": "CC1", "tax_code": "I9"})
    assert r.status_code == 200, r.text
    return r.json()["id"]


def _rows(content: bytes) -> list:
    return [[c.value for c in row] for row in load_workbook(io.BytesIO(content)).active.iter_rows(min_row=2)]


def _pdf(text: str) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_textbox(pymupdf.Rect(20, 20, 592, 820), text, fontsize=7.0, fontname="cour")
    data = doc.tobytes()
    doc.close()
    return data


# ---- 1. a red document never reaches the upload file ---------------------------------------------------------------
def test_the_same_invoice_in_two_folders_goes_out_only_once(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    vat, no = _vat(), f"TW-{run}"
    _vendor(_admin(), "300600", f"Twice {run}", vat)
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision({f"a_{run}.pdf": _invoice_answer(no, supplier=f"Twice {run}", vat_no=vat),
                                                  f"b_{run}.pdf": _invoice_answer(no, supplier=f"Twice {run}", vat_no=vat)}))
    u = _user(f"fix1.{run.lower()}@test.local", "Fix One")
    ba = _drop(u, {f"a_{run}.pdf": make_ap.invoice(no=no)}, f"F1a {run}")
    bb = _drop(u, {f"b_{run}.pdf": make_ap.invoice(no=no)}, f"F1b {run}")
    assert _sheet(u, ba["id"])[f"a_{run}.pdf"]["status"] == "green"
    assert u.post(f"/api/ap/batches/{ba['id']}/sap/upload", json={"include_yellow": False}).status_code == 200
    rb = u.post(f"/api/ap/batches/{bb['id']}/sap/upload", json={"include_yellow": True})
    assert rb.status_code == 400 and "Nothing here is ready" in rb.json()["detail"]  # Tess turned it red just now: it stays out
    d = _sheet(u, bb["id"])[f"b_{run}.pdf"]
    assert d["status"] == "red" and not d["uploaded"] and any("already went out" in x["text"] for x in d["issues"])
    assert not any("In the SAP upload" in h["text"] for h in u.get(f"/api/ap/items/{d['item_id']}").json()["history"])


def test_a_vendor_removed_since_the_last_test_stays_out_of_the_upload(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    vat, adm = _vat(), _admin()
    vid = _vendor(adm, "VGONE" + run, f"Gone {run}", vat)
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision({f"g_{run}.pdf": _invoice_answer(f"GO-{run}", supplier=f"Gone {run}", vat_no=vat)}))
    u = _user(f"fix1b.{run.lower()}@test.local", "Fix One B")
    b = _drop(u, {f"g_{run}.pdf": make_ap.invoice(no=f"GO-{run}")}, f"F1c {run}")
    assert _sheet(u, b["id"])[f"g_{run}.pdf"]["status"] == "green"
    assert adm.delete(f"/api/ap/masters/vendors/{vid}").status_code == 200
    assert u.post(f"/api/ap/batches/{b['id']}/sap/upload", json={"include_yellow": True}).status_code == 400
    d = _sheet(u, b["id"])[f"g_{run}.pdf"]
    assert d["status"] == "red" and not d["uploaded"] and any("not in the vendor tables" in x["text"] for x in d["issues"])


# ---- 2. people's coding never changes what admins keep; Tess checks against what admins keep -----------------------------
def test_a_plain_users_coding_never_changes_an_admin_vendor_row(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    adm = _admin()
    acme, acme_vat = f"Acme {run} Supplies", _vat()
    vid = _vendor(adm, "300100", acme, acme_vat)
    newco, new_no, acme_no, other, other_no = f"Newco {run}", f"NC-{run}", f"AC-{run}", f"Otherco {run}", f"OC-{run}"
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision({f"newco_{run}.pdf": _invoice_answer(new_no, supplier=newco, vat_no=_vat()),
                                                  f"other_{run}.pdf": _invoice_answer(other_no, supplier=other, vat_no=_vat()),
                                                  f"acme_{run}.pdf": _invoice_answer(acme_no, supplier=acme, vat_no=acme_vat)}))
    u = _user(f"fix2.{run.lower()}@test.local", "Plain User")
    b1 = _drop(u, {f"newco_{run}.pdf": make_ap.invoice(no=new_no, supplier=newco), f"other_{run}.pdf": make_ap.invoice(no=other_no, supplier=other)}, f"F2 {run}")
    groups = {g["supplier"]: g for g in u.get("/api/ap/needs").json()["vendors"]}
    assert u.get("/api/ap/needs").json()["confirm"] == []  # only admins see what to confirm
    # a plain user "codes" Newco but names the admin's vendor (name and VAT are free fields): a new learned row, the admin's stays
    r = u.post("/api/ap/vendors/code", json={"ids": groups[newco]["ids"], "vendor_no": "999001", "gl": "999999", "cost_centre": "CCX", "tax_code": "ZZ", "name": acme, "vat_no": acme_vat})
    assert r.status_code == 200 and r.json()["vendor"]["learned"] and r.json()["vendor"]["id"] != vid
    with SessionLocal() as s:
        v = s.get(ApVendorCode, vid)
        assert (v.vendor_no, v.gl, v.tax_code, v.learned) == ("300100", "6100100", "I9", False)
    b2 = _drop(u, {f"acme_{run}.pdf": make_ap.invoice(no=acme_no, supplier=acme, vat_no=acme_vat)}, f"F2b {run}")
    d2 = _sheet(u, b2["id"])[f"acme_{run}.pdf"]
    assert d2["status"] == "green" and d2["vendor"] == "300100" and "999999" not in d2["gl"]  # the admin's coding wins
    # a person's own coding is used, but yellow (learned) - its GL and tax code, in no admin table, are yellow too, not green
    g = groups[other]
    r = u.post("/api/ap/vendors/code", json={"ids": g["ids"], "vendor_no": "999002", "gl": "999999", "cost_centre": "CCX", "tax_code": "ZZ", "name": g["supplier"], "vat_no": g["vat"]})
    d = _sheet(u, b1["id"])[f"other_{run}.pdf"]
    texts = {x["text"] for x in d["issues"]}
    assert d["status"] == "yellow" and {"vendor 999002 is learned: an admin confirms it once", "GL 999999 is learned: an admin confirms it once",
                                        "tax code ZZ is learned: an admin confirms it once"} <= texts
    # only an admin confirms; then Tess tests it again
    lid = r.json()["vendor"]["id"]
    assert u.post(f"/api/ap/masters/vendors/{lid}/confirm").status_code == 403
    assert any(x["id"] == lid for x in adm.get("/api/ap/needs").json()["confirm"] if x["kind"] == "vendors")
    assert adm.post(f"/api/ap/masters/vendors/{lid}/confirm").status_code == 200
    assert _sheet(u, b1["id"])[f"other_{run}.pdf"]["status"] == "green"


def test_a_person_added_utility_account_is_yellow_until_an_admin_confirms_it(utilities):  # noqa: F811 (the fixture)
    adm = utilities
    run = uuid.uuid4().hex[:6].upper()
    acct = "50" + "".join(random.choice("0123456789") for _ in range(8))
    u = _user(f"fix2b.{run.lower()}@test.local", "Plain Two")
    b = _drop(u, {f"eskom_{run}.pdf": eskom_pdf(account=acct, invoice=f"7{random.randint(10**7, 10**8 - 1)}")}, f"F2c {run}")
    other = next(x for x in u.get("/api/ap/needs").json()["other"] if x["batch_id"] == b["id"])
    # 1.0.101: only admins add accounts from Other now; a row a person added before (source "person") still waits for an admin
    r = u.post(f"/api/ap/captures/{other['id']}/decide", json={"action": "manual", "data": {"vendor": "NOT-A-BP-" + run, "cost_centre": "1", "site": "x"}})
    assert r.status_code == 403
    with SessionLocal() as s:
        M.add_account(s, s.query(User).filter(User.email == f"fix2b.{run.lower()}@test.local").one(), {"account": acct, "vendor": "NOT-A-BP-" + run, "cost_centre": "1", "site": "x"})
        assert C.rematch(s) == 1
        s.commit()
    r = u.get(f"/api/ap/captures/{other['id']}")
    assert r.json()["vendor_no"] == "NOT-A-BP-" + run
    assert r.json()["sap_status"] == "yellow" and any("learned: an admin confirms it once" in x["text"] for x in r.json()["sap_issues"])
    row = next(x for x in adm.get("/api/ap/needs").json()["confirm"] if x["kind"] == "utilities" and x["vendor"] == "NOT-A-BP-" + run)
    assert u.post(f"/api/ap/masters/utilities/{row['id']}/confirm").status_code == 403
    assert adm.post(f"/api/ap/masters/utilities/{row['id']}/confirm").status_code == 200
    assert u.get(f"/api/ap/captures/{other['id']}").json()["sap_status"] == "green"


def test_an_admins_own_coding_waits_for_a_second_admin(monkeypatch):
    """1.0.104 (Valuence F0b): was "an admin's own coding is confirmed at once" — that was a self-approval path. Now an
    admin's coding is learned (yellow) like anyone's, she can't confirm it herself, and a second admin does."""
    run = uuid.uuid4().hex[:6].upper()
    adm, sup = _admin(), f"Adminco {run}"
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision({f"ad_{run}.pdf": _invoice_answer(f"AD-{run}", supplier=sup, vat_no=_vat())}))
    b = _drop(adm, {f"ad_{run}.pdf": make_ap.invoice(no=f"AD-{run}", supplier=sup)}, f"F2d {run}")
    grp = next(g for g in adm.get("/api/ap/needs").json()["vendors"] if g["supplier"] == sup)
    r = adm.post("/api/ap/vendors/code", json={"ids": grp["ids"], "vendor_no": "301" + run[:3], "gl": "6100100", "cost_centre": "CC1", "tax_code": "I9"})
    assert r.status_code == 200 and r.json()["vendor"]["learned"] is True
    assert _sheet(adm, b["id"])[f"ad_{run}.pdf"]["status"] == "yellow"
    lid = r.json()["vendor"]["id"]
    own = adm.post(f"/api/ap/masters/vendors/{lid}/confirm")
    assert own.status_code == 403 and "a second admin confirms it" in own.json()["detail"]
    adm2 = signed_in("cap.admin2@test.local", "admin", name="Bongani Admin")
    assert adm2.post(f"/api/ap/masters/vendors/{lid}/confirm").status_code == 200
    assert _sheet(adm, b["id"])[f"ad_{run}.pdf"]["status"] == "green"


# ---- 3. a model answer that can't be used never strands an invoice -------------------------------------------------
def test_an_infinite_confidence_counts_a_try_and_the_batch_goes_on(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    poison = _invoice_answer(f"PZ-{run}")
    poison["confidence"]["supplier"] = "1e999"
    vat = _vat()
    _vendor(_admin(), "300200", f"Fine {run}", vat)
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision({f"a_poison_{run}.pdf": poison, f"b_fine_{run}.pdf": _invoice_answer(f"OK-{run}", supplier=f"Fine {run}", vat_no=vat)}))
    u = _user(f"fix3.{run.lower()}@test.local", "Fix Three")
    b = _drop(u, {f"a_poison_{run}.pdf": make_ap.invoice(no=f"PZ-{run}"), f"b_fine_{run}.pdf": make_ap.invoice(no=f"OK-{run}")}, f"F3 {run}")
    with SessionLocal() as s:
        caps = {s.get(ApItem, c.item_id).name: c for c in s.query(ApCapture).filter(ApCapture.batch_id == b["id"]).all()}
        bad, fine = caps[f"a_poison_{run}.pdf"], caps[f"b_fine_{run}.pdf"]
        assert (bad.stage, bad.tries, bad.needs) == ("needs", 2, "unread")  # Cara/Dex/Eli, then Seth, then a person
        assert fine.stage == "ready" and fine.sap_status == "green"
    r = u.post(f"/api/ap/batches/{b['id']}/sap/upload", json={"include_yellow": False})
    assert r.status_code == 200 and f"OK-{run}" in {row[6] for row in _rows(r.content)}
    assert C.check_read({"supplier": "x"}, {"supplier": float("inf"), "bill_to": float("nan")}, 80)[0] == 60  # no confidence given, not an OverflowError


def test_a_reading_that_cannot_be_saved_moves_on(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()

    def boom(*a, **k):
        raise OverflowError("cannot convert float infinity to integer")

    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision({f"s_{run}.pdf": _invoice_answer(f"SV-{run}")}))
    monkeypatch.setattr(C, "to_item", boom)
    u = _user(f"fix3b.{run.lower()}@test.local", "Fix Three B")
    b = _drop(u, {f"s_{run}.pdf": make_ap.invoice(no=f"SV-{run}")}, f"F3b {run}")
    with SessionLocal() as s:
        c = s.query(ApCapture).filter(ApCapture.batch_id == b["id"]).one()
        assert (c.stage, c.needs, c.tries, c.claimed_at) == ("needs", "unread", 2, None)
    row = next(x for x in u.get("/api/ap/needs").json()["unread"] if x["batch_id"] == b["id"])
    assert any("could not be saved" in p for p in row["problems"])


# ---- 4. Read again after a sign-off --------------------------------------------------------------------------------
def test_read_again_after_a_sign_off_leaves_no_file_at_sorting(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    vat = _vat()
    _vendor(_admin(), "300900", f"Signed {run}", vat)
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision({f"a_{run}.pdf": _invoice_answer(f"SA-{run}", supplier=f"Signed {run}", vat_no=vat),
                                                  f"b_{run}.pdf": _invoice_answer(f"SB-{run}", supplier=f"Signed {run}", vat_no=vat)}))
    u = _user(f"fix4.{run.lower()}@test.local", "Fix Four")
    b = _drop(u, {f"a_{run}.pdf": make_ap.invoice(no=f"SA-{run}"), f"b_{run}.pdf": make_ap.invoice(no=f"SB-{run}")}, f"F4 {run}")
    docs = _sheet(u, b["id"])
    assert u.post(f"/api/ap/items/{docs[f'a_{run}.pdf']['item_id']}/sign", json={"action": "clear", "note": "ok"}).json()["status"] == "cleared"
    assert u.post(f"/api/ap/batches/{b['id']}/read").status_code == 200
    with SessionLocal() as s:
        assert s.query(ApCapture).filter(ApCapture.batch_id == b["id"], ApCapture.stage.in_(C.READING)).count() == 0
    bk = u.get(f"/api/ap/batches/{b['id']}/sap").json()["capture"]
    assert bk["busy"] is False and bk["sap"]["green"] == 2
    assert {d["status"] for d in _sheet(u, b["id"]).values()} == {"green"}


# ---- 5. the AP team's own decisions count ---------------------------------------------------------------------------
def test_an_invoice_the_ap_team_rejected_is_red_and_stays_out(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    vat = _vat()
    _vendor(_admin(), "300500", f"Shady {run}", vat)
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision({f"shady_{run}.pdf": _invoice_answer(f"SH-{run}", supplier=f"Shady {run}", vat_no=vat)}))
    u = _user(f"fix5.{run.lower()}@test.local", "Fix Five")
    b = _drop(u, {f"shady_{run}.pdf": make_ap.invoice(no=f"SH-{run}", supplier=f"Shady {run}", vat_no=vat)}, f"F5 {run}")
    d = _sheet(u, b["id"])[f"shady_{run}.pdf"]
    assert d["status"] == "green"
    assert u.post(f"/api/ap/items/{d['item_id']}/sign", json={"action": "reject", "note": "Fraud: not our supplier"}).json()["status"] == "rejected"
    d = _sheet(u, b["id"])[f"shady_{run}.pdf"]  # red at once, before any download
    assert d["status"] == "red" and any("rejected by the AP team" in x["text"] and "Fraud" in x["text"] for x in d["issues"])
    assert u.post(f"/api/ap/batches/{b['id']}/sap/upload", json={"include_yellow": True}).status_code == 400


def test_an_invoice_ivy_stopped_as_a_duplicate_is_yellow_with_why(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    vat, no = _vat(), f"SN-{run}"
    _vendor(_admin(), "300510", f"Seen {run}", vat)
    u = _user(f"fix5b.{run.lower()}@test.local", "Fix Five B")
    monkeypatch.setattr(C, "model_on", lambda: False)  # the first folder is read by the AP team's rules only
    _drop(u, {f"first_{run}.pdf": make_ap.invoice(no=no, supplier=f"Seen {run}", vat_no=vat)}, f"F5b1 {run}")
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision({f"again_{run}.pdf": _invoice_answer(no, supplier=f"Seen {run}", vat_no=vat)}))
    b = _drop(u, {f"again_{run}.pdf": make_ap.invoice(no=no, supplier=f"Seen {run}", vat_no=vat)}, f"F5b2 {run}")
    d = _sheet(u, b["id"])[f"again_{run}.pdf"]
    stop = next(f for f in u.get(f"/api/ap/items/{d['item_id']}").json()["findings"] if f["code"] == "duplicate_seen")
    assert d["status"] == "yellow" and any(x["text"] == stop["text"] for x in d["issues"])
    assert u.post(f"/api/ap/batches/{b['id']}/sap/upload", json={"include_yellow": False}).status_code == 400
    # a person clears it (the stop was checked): then it is green
    assert u.post(f"/api/ap/items/{d['item_id']}/sign", json={"action": "clear", "note": "Checked: the first copy was never posted"}).json()["status"] == "cleared"
    assert _sheet(u, b["id"])[f"again_{run}.pdf"]["status"] == "green"


# ---- 6. a scan goes to the capture bots, OCR only finds a profile -----------------------------------------------------
def _scan(pdf: bytes) -> bytes:
    src = pymupdf.open(stream=pdf, filetype="pdf")
    out = pymupdf.open()
    for p in src:
        pix = p.get_pixmap(dpi=100)
        page = out.new_page(width=p.rect.width, height=p.rect.height)
        page.insert_image(page.rect, stream=pix.tobytes("png"))
    return out.tobytes()


def _fake_ocr(monkeypatch, text: str):
    """Tesseract or not: with OCR allowed, a page with no text layer 'reads' as `text`."""
    from cbx.pdfstudio.municipal import engine as ME

    real = ME.read_pdf

    def read_pdf(data, ocr="auto"):
        rd = real(data, ocr="never")
        if ocr != "never" and not rd["full"].strip():
            return {**rd, "full": text, "pages": [text], "ocr": True}
        return rd

    monkeypatch.setattr(ME, "read_pdf", read_pdf)


def test_a_scanned_invoice_goes_to_manual_capture_even_with_ocr(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    _fake_ocr(monkeypatch, f"Acme Cleaning (Pty) Ltd\nTAX INVOICE\nInvoice no SC-{run}\nVAT 1 500.00\nTotal 11 500.00")
    u = _user(f"fix6.{run.lower()}@test.local", "Fix Six")
    b = _drop(u, {f"scan_{run}.pdf": _scan(make_ap.invoice(no=f"SC-{run}"))}, f"F6 {run}")
    with SessionLocal() as s:
        c = s.query(ApCapture).filter(ApCapture.batch_id == b["id"]).one()
        assert c.bucket == "manual" and c.needs != "odd"


def test_ocr_still_lets_a_profile_know_a_scanned_utility_bill(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    _fake_ocr(monkeypatch, ESKOM.format(account="5012345678", invoice=f"7{random.randint(10**7, 10**8 - 1)}"))
    u = _user(f"fix6b.{run.lower()}@test.local", "Fix Six B")
    b = _drop(u, {f"eskom_scan_{run}.pdf": _scan(eskom_pdf())}, f"F6b {run}")
    hist = u.get(f"/api/ap/items/{b['items'][0]['id']}").json()["history"]
    assert any(h["who"] == "Sol" and h["text"].startswith("A utility bill") for h in hist)


# ---- 7. amounts and currency ---------------------------------------------------------------------------------------------
def test_a_negative_eskom_adjustment_posts_on_the_other_side(utilities):  # noqa: F811 (the fixture)
    run = uuid.uuid4().hex[:6].upper()
    inv = str(random.randint(10**8, 10**9 - 1))
    text = ESKOM.format(account="5012345678", invoice=inv).replace("Interest on overdue account R 12.34", "Interest on overdue account R -12.34")
    u = _user(f"fix7.{run.lower()}@test.local", "Fix Seven")
    b = _drop(u, {f"eskom_{run}.pdf": _pdf(text)}, f"F7 {run}")
    d = _sheet(u, b["id"])[f"eskom_{run}.pdf"]
    assert d["status"] == "green", d["issues"]
    rows = u.get(f"/api/ap/captures/{d['id']}").json()["sap_rows"]
    assert all(float(r["amount"]) > 0 for r in rows)
    assert [(r["pkey"], r["account"], r["amount"]) for r in rows if r["account"] == "SF101"] == [("50", "SF101", "12.34")]
    assert P.balance_fails([{**r, "_file": "x"} for r in rows]) == (1, set())


def test_a_credit_bill_with_interest_balances():
    fields = {"kind": "utility", "account": "5012345678", "billing_date": "2026-09-15", "due_date": "2026-10-05", "invoice_no": "700999111",
              "statement_month": "SEPTEMBER 2026", "read_start": "2026-08-01", "read_end": "2026-08-31", "kwh": "0", "reading_type": "ACTUAL",
              "services": [{"service": "electricity", "ex": "-1000.00", "vat": "-150.00", "how": "x", "sure": True}], "adjustments": {"interest": "12.34"}}
    acct = {"vendor": "700999", "cc": "01306", "name1": "Badplaas ETE", "muni": "ESKOM"}
    rows, st, issues = P.utility_doc(fields, None, acct, _tables(), file="credit.pdf")
    assert [(r["pkey"], r["account"], r["amount"]) for r in rows] == [("21", "700999", "1137.66"), ("50", "SA110", "1000.00"), ("40", "SF101", "12.34"), ("50", "JA600", "150.00")]
    known = {"vendors": {"700999"}, "gls": set(_tables()["gl"].values()), "companies": {"T000"}, "doc_types": {"KR", "KG"}, "tax_codes": {"I9", "I0"}}
    assert P.test_doc(rows, known, pip_status=st, pip_issues=issues)[0] == "yellow"  # a credit memo: confirm with finance


def test_tess_tests_signs_sizes_and_currency():
    rows, _, _ = P.manual_doc({"supplier": "Acme", "invoice_no": "T-1", "invoice_date": "2026-10-01", "due_date": "2026-10-31", "excl": "100.00", "vat": "15.00", "incl": "115.00"},
                              {"vendor_no": "300100", "gl": "6100100", "cost_centre": "CC1"}, "X143", _tables(), file="t.pdf")
    known = {"vendors": {"300100"}, "gls": {"6100100", "JA600"}, "companies": {"X143"}, "doc_types": {"KR", "KG"}, "tax_codes": {"I9", "I0"}}
    assert P.test_doc(rows, known) == ("green", [])
    for amt in ("-15.00", "0.00", "12345678901234.00"):
        bad = [dict(r) for r in rows]
        bad[2]["amount"] = amt
        st, why = P.test_doc(bad, known)
        assert st == "red" and any(f"amount {amt}" in x["text"] for x in why), (amt, why)
    assert [P.currency_code(x) for x in ("R", "zar", "Rand", "US$", "EUR", "€")] == ["ZAR", "ZAR", "ZAR", "USD", "EUR", "EUR"]
    st, why = P.test_doc([{**r, "curr": "USD"} for r in rows], known)
    assert st == "yellow" and why == [{"level": "yellow", "text": "in USD, not rand: confirm before upload"}]
    assert P.test_doc([{**r, "curr": "XYZ"} for r in rows], known)[0] == "red"


def test_an_invoice_in_r_is_rand_and_a_foreign_one_is_not_green(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    vat = _vat()
    _vendor(_admin(), "300300", f"Rand {run}", vat)
    a1 = _invoice_answer(f"CR-{run}", supplier=f"Rand {run}", vat_no=vat)
    a1["currency"] = "R"
    a2 = _invoice_answer(f"CU-{run}", supplier=f"Rand {run}", vat_no=vat, excl="100.00", vat="0.00", incl="100.00",
                         lines=[{"description": "Licence", "amount_excl": "100.00", "vat_rate": 0, "vat": "0.00"}])
    a2["currency"] = "US$"
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision({f"r_{run}.pdf": a1, f"usd_{run}.pdf": a2}))
    u = _user(f"fix7c.{run.lower()}@test.local", "Fix Seven C")
    b = _drop(u, {f"r_{run}.pdf": make_ap.invoice(no=f"CR-{run}"), f"usd_{run}.pdf": make_ap.invoice(no=f"CU-{run}")}, f"F7c {run}")
    docs = _sheet(u, b["id"])
    assert docs[f"r_{run}.pdf"]["status"] == "green" and docs[f"usd_{run}.pdf"]["status"] == "yellow"
    r = u.post(f"/api/ap/batches/{b['id']}/sap/upload", json={"include_yellow": False})
    cur = P.SAP_HEADERS.index("curr")
    assert {row[cur] for row in _rows(r.content)} == {"ZAR"}


# ---- 8. a cut SAP name is only the start of a name ---------------------------------------------------------------------
def test_the_start_of_a_cut_company_name_is_unsure_and_a_za_vat_number_never_matches_abroad(monkeypatch):
    vat = _vat()
    with SessionLocal() as s:
        m = C.match_entity(s, "Excellerate Facilities Management (Pty) Ltd", vat)
        assert m is None or (m.get("unsure") and m["code"] != "X136")  # not Angola
        m = C.match_entity(s, "CBRE Excellerate Property Services (Pty) Ltd", vat)
        assert m is None or (m.get("unsure") and m["code"] != "X110")  # not Ivory Coast
        m = C.match_entity(s, "Excellerate Facilities Management (Pty) Ltd", "")
        assert m["code"] == "X136" and m["unsure"] and m["how"] == "the start of the name"
    run = uuid.uuid4().hex[:6].upper()
    sup_vat = _vat()
    _vendor(_admin(), "300400", f"Clean {run}", sup_vat)
    ans = _invoice_answer(f"EF-{run}", supplier=f"Clean {run}", vat_no=sup_vat, bill_to="Excellerate Facilities Management (Pty) Ltd")
    ans["bill_to_vat"] = vat
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision({f"efm_{run}.pdf": ans}))
    u = _user(f"fix8.{run.lower()}@test.local", "Fix Eight")
    b = _drop(u, {f"efm_{run}.pdf": make_ap.invoice(no=f"EF-{run}")}, f"F8 {run}")
    d = _sheet(u, b["id"])[f"efm_{run}.pdf"]
    assert d["status"] != "green" and d["entity"] != "X136"


# ---- 9. re-uploading a list keeps what the sheet doesn't carry -----------------------------------------------------------
def _stat_codes(rename: str = "") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["Company Code", "Company Name", "Entity Name"])
    for code, name, entity in M.COMPANY_CODES:
        ws.append([code, rename if (rename and code == "X121") else name, entity])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_reuploading_the_company_codes_keeps_vat_aliases_and_doc_types():
    adm, vat = _admin(), _vat("9")  # not a South African VAT number: nothing else matches on it
    x = next(r for r in adm.get("/api/ap/masters/entities", params={"q": "X121"}).json()["rows"] if r["code"] == "X121")
    seed = {k: x[k] for k in ("code", "name", "entity", "vat_no", "aliases", "doc_type", "doc_type_credit")}
    try:
        r = adm.put(f"/api/ap/masters/entities/{x['id']}", json={**seed, "vat_no": vat, "aliases": "EPS Eswatini Facilities", "doc_type": "YR", "doc_type_credit": "YG"})
        assert r.status_code == 200, r.text
        pre = adm.post("/api/ap/masters/entities/import", files={"file": ("Stat Company Codes.xlsx", _stat_codes())}, data={"confirm": "false"}).json()
        assert not any(s.startswith("X121 ") for s in pre["sample"]["changed"])
        pre = adm.post("/api/ap/masters/entities/import", files={"file": ("Stat Company Codes.xlsx", _stat_codes("EPS Eswatini Ltd"))}, data={"confirm": "false"}).json()
        ch = next(c for c in pre["changes"] if c["row"].startswith("X121 "))
        assert ch["fields"] == [{"field": "name", "from": seed["name"], "to": "EPS Eswatini Ltd"}]  # per field, and only what the sheet carries
        r = adm.post("/api/ap/masters/entities/import", files={"file": ("Stat Company Codes.xlsx", _stat_codes())}, data={"confirm": "true"})
        assert r.status_code == 200 and r.json()["saved"]
        with SessionLocal() as s:
            e = s.query(ApEntity).filter(ApEntity.code == "X121").one()
            assert (e.vat_no, e.aliases, e.doc_type, e.doc_type_credit) == (vat, "EPS Eswatini Facilities", "YR", "YG")
    finally:
        adm.put(f"/api/ap/masters/entities/{x['id']}", json=seed)


def test_a_vendor_sheet_without_aliases_keeps_the_aliases():
    adm = _admin()
    run = uuid.uuid4().hex[:6].upper()
    vat = _vat()
    with SessionLocal() as s:
        s.query(ApVendorCode).filter(ApVendorCode.vendor_no == "377" + run[:3]).delete()
        s.add(ApVendorCode(vendor_no="377" + run[:3], name=f"Alias {run}", name_key=M.name_key(f"Alias {run}"), vat_no=vat, gl="6100100", aliases=f"Alias Trading {run}", source="excel"))
        s.commit()
    sheet = _xlsx([["Vendor", "Name", "VAT No", "GL", "Cost centre", "Tax code"], ["377" + run[:3], f"Alias {run}", vat, "6100200", "CC9", "I9"]])
    r = adm.post("/api/ap/masters/vendors/import", files={"file": ("vendors.xlsx", sheet)}, data={"confirm": "false"}).json()
    assert {f["field"] for c in r["changes"] if c["row"].startswith("377" + run[:3]) for f in c["fields"]} == {"gl", "cost_centre", "tax_code"}
    with SessionLocal() as s:
        old = [(v.vendor_no, v.vat_no, v.name, v.gl, v.cost_centre, v.tax_code, v.company_code, v.aliases) for v in s.query(ApVendorCode).filter(ApVendorCode.learned.is_(False)).all()]
    sheet = _xlsx([["Vendor", "Name", "VAT No", "GL", "Cost centre", "Tax code", "Company code"]] + [[a, c, b, d, e, f, g] for a, b, c, d, e, f, g, _h in old if a != "377" + run[:3]]
                  + [["377" + run[:3], f"Alias {run}", vat, "6100200", "CC9", "I9", ""]])
    assert adm.post("/api/ap/masters/vendors/import", files={"file": ("vendors.xlsx", sheet)}, data={"confirm": "true"}).json()["saved"]
    with SessionLocal() as s:
        v = s.query(ApVendorCode).filter(ApVendorCode.vendor_no == "377" + run[:3]).one()
        assert (v.gl, v.aliases) == ("6100200", f"Alias Trading {run}")


# ---- 10. hostile workbooks -----------------------------------------------------------------------------------------------------
def _claiming(ref: str, far: bool = True) -> bytes:
    """A two-row company-code sheet whose <dimension> claims `ref` (with one more cell in the last row it claims)."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Company Code", "Company Name"])
    ws.append(["X999", "Bomb Co"])
    buf = io.BytesIO()
    wb.save(buf)
    last = re.match(r"A1:([A-Z]+)(\d+)", ref).group(2)
    zin, out = zipfile.ZipFile(io.BytesIO(buf.getvalue())), io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for i in zin.infolist():
            d = zin.read(i.filename)
            if i.filename == "xl/worksheets/sheet1.xml":
                d = re.sub(rb'<dimension ref="[^"]+"/>', f'<dimension ref="{ref}"/>'.encode(), d)
                if far:
                    d = d.replace(b"</sheetData>", f'<row r="{last}"><c r="A{last}" t="inlineStr"><is><t>x</t></is></c></row></sheetData>'.encode())
            z.writestr(i, d)
    return out.getvalue()


def test_a_sheet_claiming_a_million_rows_is_read_lazily():
    import tracemalloc

    tracemalloc.start()
    try:
        with pytest.raises(M.Refused, match="More than 60000 rows"):
            M.read_sheet("entities", _claiming("A1:J1048576"), "codes.xlsx")  # a value far down: too long, said plainly
        peak = tracemalloc.get_traced_memory()[1]
    finally:
        tracemalloc.stop()
    assert peak < 64 * 1024 * 1024, peak  # not a million rows of ten cells built first (~150 MB)
    assert M.read_sheet("entities", _claiming("A1:XFD1048576", far=False), "codes.xlsx") == [{"code": "X999", "name": "Bomb Co"}]  # only the claim: the real rows


def test_the_full_width_claim_does_not_run_out_of_memory(tmp_path):
    p = str(tmp_path / "bomb210.xlsx")  # 1.0.98: the old conftest has no CBX_HOME
    with open(p, "wb") as f:
        f.write(_claiming("A1:XFD1048576"))
    code = ("import sys\nfrom cbx.ap import masters as M\n"
            "try:\n    M.read_sheet('entities', open(sys.argv[1], 'rb').read(), 'codes.xlsx')\n    print('READ')\n"
            "except MemoryError:\n    print('MEMORYERROR')\nexcept M.Refused as e:\n    print('REFUSED', e)\n")
    lim = 3 * 1024 ** 3
    cap = None
    if sys.platform != "win32":  # 2.10.1: `resource` is Linux/mac only; on the Windows CBX PC the refusal is checked without the cap
        import resource

        def cap():
            resource.setrlimit(resource.RLIMIT_AS, (lim, lim))

    r = subprocess.run([sys.executable, "-c", code, p], capture_output=True, text=True, timeout=300, preexec_fn=cap,
                       cwd=os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    assert "REFUSED" in r.stdout, (r.stdout, r.stderr[-2000:])


def test_a_one_column_sap_config_is_read_not_a_500():
    wb = Workbook()
    ws = wb.active
    ws.title = "Fixed"
    ws.append(["Key"])
    ws.append(["Company"])
    g = wb.create_sheet("GL Codes")
    g.append(["Charge"])
    g.append(["electricity"])
    buf = io.BytesIO()
    wb.save(buf)
    fixed, gl = M.read_sap_config(buf.getvalue())
    assert fixed == [{"key": "Company", "value": ""}] and gl == [{"key": "electricity", "value": "", "description": ""}]
    r = _admin().post("/api/ap/sap-values/import", files={"file": ("sap_config.xlsx", buf.getvalue())}, data={"confirm": "false"})
    assert r.status_code == 200 and r.json()["saved"] is False


# ---- 11. Pull from SAP: every page, said plainly when cut, saved as previewed ----------------------------------------------
def _row(i: int) -> dict:
    return {"Vendor": str(700000 + i), "PID": "%05d" % (1000 + i % 900), "Name1": f"Site number {i}", "Name3": f"6000/{i:05d}/X1",
            "AccountReference": "%012d" % (6600 + i), "Municipality": "CHIEF ALBERT LUTHULI MUNICIPALITY"}


class Fake(BaseHTTPRequestHandler):
    calls = 0

    def log_message(self, *a):
        pass

    def do_GET(self):
        p = self.path
        if p.startswith("/odata/Big"):
            body = {"d": {"results": [_row(i) for i in range(10632)]}}
        elif p.startswith("/odata/Shifty"):
            Fake.calls += 1
            body = {"value": [_row(i) for i in range(10 if Fake.calls == 1 else 3)]}
        elif p.startswith("/odata/Paged?$skiptoken=1000"):
            body = {"d": {"results": [_row(i) for i in range(1000, 1500)]}}
        elif p.startswith("/odata/Paged"):
            body = {"d": {"results": [_row(i) for i in range(1000)], "__next": "/odata/Paged?$skiptoken=1000"}}
        elif p.startswith("/odata/V4"):
            n = int(p.split("skip=")[1]) if "skip=" in p else 0
            body = {"value": [_row(i) for i in range(n, n + 100)], **({"@odata.nextLink": f"V4?skip={n + 100}"} if n < 200 else {})}
        elif p.startswith("/odata/Endless"):
            body = {"d": {"results": [_row(1)], "__next": "/odata/Endless?$skiptoken=1"}}
        elif p.startswith("/odata/Away"):
            body = {"d": {"results": [_row(1)], "__next": "https://elsewhere.example.com/odata/Away?$skiptoken=1"}}
        else:
            body = {}
        b = json.dumps(body).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(b)))
        self.end_headers()
        self.wfile.write(b)


@pytest.fixture()
def conn(monkeypatch):
    monkeypatch.setattr(settings, "api_allow_private", True)
    srv = ThreadingHTTPServer(("127.0.0.1", 0), Fake)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    with SessionLocal() as s:
        c = ApiConnection(slug="fix-sap-" + uuid.uuid4().hex[:6], name="SAP master data", base_url=f"http://127.0.0.1:{srv.server_address[1]}",
                          auth_kind="none", tools=["ap"], is_active=True)
        s.add(c)
        s.commit()
        cid = c.id
    yield cid
    srv.shutdown()


COLS = {"vendor": "Vendor", "cost_centre": "PID", "site": "Name1", "name3": "Name3", "account": "AccountReference", "municipality": "Municipality"}


def _pull(adm, cid, path, **kw):
    return adm.post("/api/ap/masters/utilities/pull", json={"connection_id": cid, "path": path, "columns": COLS, "confirm": False, **kw})


def test_a_real_sized_utilities_list_is_pulled_and_a_cut_answer_is_said_plainly(conn, monkeypatch):
    adm = _admin()
    r = _pull(adm, conn, "/odata/Big")
    assert r.status_code == 200 and r.json()["rows"] == 10632, r.text
    monkeypatch.setattr(M, "PULL_MB", 1)  # an answer over the cap
    r = _pull(adm, conn, "/odata/Big")
    assert r.status_code == 400 and "was cut" in r.json()["detail"]


def test_every_page_is_pulled_and_nothing_else_is_followed(conn):
    adm = _admin()
    sheet = [["VENDOR/BP NUMBER", "PID/CC", "NAME1", "NAME3", "ACCOUNT REFERENCE", "MUNICIPALITY/PAYEE"]] + \
        [[r["Vendor"], r["PID"], r["Name1"], r["Name3"], r["AccountReference"], r["Municipality"]] for r in (_row(i) for i in range(1500))]
    assert adm.post("/api/ap/masters/utilities/import", files={"file": ("vendor_mappings.xlsx", _xlsx(sheet))}, data={"confirm": "true"}).json()["saved"]
    pre = _pull(adm, conn, "/odata/Paged").json()
    assert pre["rows"] == 1500 and pre["removed"] == 0 and not pre["saved"]
    assert _pull(adm, conn, "/odata/V4").json()["rows"] == 300  # OData v4 nextLink, relative to the page
    r = _pull(adm, conn, "/odata/Endless")
    assert r.status_code == 400 and "keeps sending more pages" in r.json()["detail"]
    r = _pull(adm, conn, "/odata/Away")
    assert r.status_code == 400 and "not on the connection's address" in r.json()["detail"]
    with SessionLocal() as s:
        assert s.query(ApUtilityAccount).count() == 1500


def test_save_stores_exactly_the_rows_that_were_previewed(conn):
    adm = _admin()
    Fake.calls = 0
    pre = _pull(adm, conn, "/odata/Shifty").json()
    assert pre["rows"] == 10 and not pre["saved"] and pre["preview_id"]
    r = adm.post("/api/ap/masters/utilities/pull", json={"confirm": True})
    assert r.status_code == 400 and "preview is gone" in r.json()["detail"]
    done = adm.post("/api/ap/masters/utilities/pull", json={"connection_id": conn, "path": "/odata/Shifty", "columns": COLS, "confirm": True, "preview_id": pre["preview_id"]}).json()
    assert done["saved"] and done["rows"] == 10 and Fake.calls == 1  # not pulled a second time
    with SessionLocal() as s:
        assert s.query(ApUtilityAccount).count() == 10
    again = adm.post("/api/ap/masters/utilities/pull", json={"confirm": True, "preview_id": pre["preview_id"]})
    assert again.status_code == 400  # used once


# ---- 12. a short AI outage --------------------------------------------------------------------------------------------------
def test_a_short_ai_outage_puts_invoices_back_in_the_queue(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    vat = _vat()
    _vendor(_admin(), "301200", f"Out {run}", vat)

    def down(*a, **k):
        raise httpx.ConnectError("[Errno 111] Connection refused")

    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", down)
    u = _user(f"fix12.{run.lower()}@test.local", "Fix Twelve")
    files = {f"inv{i}_{run}.pdf": make_ap.invoice(no=f"OUT{i}-{run}") for i in range(3)}
    b = _drop(u, files, f"F12 {run}")
    with SessionLocal() as s:
        caps = s.query(ApCapture).filter(ApCapture.batch_id == b["id"]).all()
        assert len(caps) == 3 and all((c.stage, c.needs, c.tries) == ("queued", "", 0) and c.claimed_at > utcnow() for c in caps)
    assert not [x for x in u.get("/api/ap/needs").json()["unread"] if x["batch_id"] == b["id"]]
    # the AI is back and the wait is over: the wake routine sends the readers
    monkeypatch.setattr(C, "ask_vision", _vision({n: _invoice_answer(f"OUT{i}-{run}", supplier=f"Out {run}", vat_no=vat) for i, n in enumerate(files)}))
    with SessionLocal() as s:
        for c in s.query(ApCapture).filter(ApCapture.batch_id == b["id"]).all():
            c.claimed_at = utcnow() - timedelta(minutes=1)
        s.commit()
    C.wake_routine()
    with SessionLocal() as s:
        caps = s.query(ApCapture).filter(ApCapture.batch_id == b["id"]).all()
        assert all((c.stage, c.sap_status, c.tries) == ("ready", "green", 1) for c in caps)
        hist = [h["text"] for h in s.get(ApItem, caps[0].item_id).history]
        assert sum("The AI did not answer" in h for h in hist) == 1


def test_an_invoice_the_ai_never_answers_moves_on_after_a_while(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()

    def slow(*a, **k):
        raise httpx.ReadTimeout("timed out")

    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", slow)
    u = _user(f"fix12b.{run.lower()}@test.local", "Fix Twelve B")
    b = _drop(u, {f"slow_{run}.pdf": make_ap.invoice(no=f"SL-{run}")}, f"F12b {run}")
    with SessionLocal() as s:
        c = s.query(ApCapture).filter(ApCapture.batch_id == b["id"]).one()
        assert (c.stage, c.tries) == ("queued", 0)
        c.claimed_at, c.read = None, {**c.read, "ai_down": (utcnow() - C.AI_PATIENCE - timedelta(minutes=1)).isoformat()}
        s.commit()
        cid = c.id
    C.wake_routine()  # waited too long: the reader counts a try and it goes to Seth, who waits for the AI with his own patience
    with SessionLocal() as s:
        c = s.get(ApCapture, cid)
        assert (c.stage, c.tries) == ("senior", 1) and c.claimed_at > utcnow() and "ai_down" in c.read


def test_what_counts_as_an_outage():
    from cbx.cexa.gateway import ModelError  # 1.0.98: the old gateway

    try:
        try:
            raise httpx.ReadTimeout("slow")
        except httpx.HTTPError as e:
            raise ModelError("The AI box is not answering at http://ai:11434: slow") from e
    except ModelError as e:
        assert C._outage(e)
    assert C._outage(ModelError("AI box error 404: {\"error\":\"model 'qwen3-vl:32b' not found\"}"))
    assert not C._outage(ModelError("the model did not return valid JSON: Expecting value"))
    assert not C._outage(ValueError("bad"))
    # 1.0.98: the old gateway's words
    assert C._outage(ModelError("Ollama not reachable at http://ai:11434: [Errno 111] Connection refused"))
    assert C._outage(ModelError("Ollama error 503: busy")) and not C._outage(ModelError("Ollama error 400: bad request"))
    assert C._outage(ModelError("Ollama error: model 'qwen3-vl:32b' not found, try pulling it first"))


# ---- 13. "already went out" is not only the vendor BP --------------------------------------------------------------------
def test_the_same_invoice_after_a_vendor_bp_correction_is_red(monkeypatch):
    run = uuid.uuid4().hex[:6].upper()
    vat, adm = _vat(), _admin()
    body = {"vendor_no": "300800", "name": f"Fix {run}", "vat_no": vat, "gl": "6100100", "cost_centre": "CC1", "tax_code": "I9"}
    vid = adm.post("/api/ap/masters/vendors", json=body).json()["id"]
    no = f"FX-{run}"
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision({f"a_{run}.pdf": _invoice_answer(no, supplier=f"Fix {run}", vat_no=vat),
                                                  f"b_{run}.pdf": _invoice_answer(no, supplier=f"Fix {run}", vat_no=vat)}))
    u = _user(f"fix13.{run.lower()}@test.local", "Fix Thirteen")
    ba = _drop(u, {f"a_{run}.pdf": make_ap.invoice(no=no)}, f"F13a {run}")
    assert u.post(f"/api/ap/batches/{ba['id']}/sap/upload", json={"include_yellow": False}).status_code == 200
    assert adm.put(f"/api/ap/masters/vendors/{vid}", json={**body, "vendor_no": "300801"}).status_code == 200
    bb = _drop(u, {f"b_{run}.pdf": make_ap.invoice(no=no)}, f"F13b {run}")
    d = _sheet(u, bb["id"])[f"b_{run}.pdf"]
    assert d["vendor"] == "300801" and d["status"] == "red" and any("already went out" in x["text"] for x in d["issues"])
    assert u.post(f"/api/ap/batches/{bb['id']}/sap/upload", json={"include_yellow": True}).status_code == 400


# ---- 14. odd input -------------------------------------------------------------------------------------------------------------
def test_odd_input_is_refused_not_a_500(monkeypatch):
    from fastapi.testclient import TestClient

    from cbx.main import app

    from .conftest import STRONG

    def client(email, role="user", tools=("ap",)):
        make_person(email, role, tools=list(tools))
        c = TestClient(app, base_url="http://localhost", raise_server_exceptions=False)
        assert c.post("/api/auth/login", json={"email": email, "password": STRONG}).status_code == 200  # 1.0.98: the old sign-in
        return c

    run = uuid.uuid4().hex[:6].upper()
    monkeypatch.setattr(C, "model_on", lambda: True)
    monkeypatch.setattr(C, "ask_vision", _vision({f"x_{run}.pdf": _invoice_answer(f"X-{run}", conf=10)}))
    u = _user(f"fix14.{run.lower()}@test.local", "Fix Fourteen")
    b = _drop(u, {f"x_{run}.pdf": make_ap.invoice(no=f"X-{run}")}, f"F14 {run}")
    with SessionLocal() as s:
        cid = s.query(ApCapture.id).filter(ApCapture.batch_id == b["id"]).scalar()
    c = client(f"fix14.{run.lower()}@test.local")
    for bad in ("supplier=Acme", ["supplier"], 5):
        r = c.post(f"/api/ap/captures/{cid}/decide", json={"action": "save", "data": {"fields": bad}})
        assert r.status_code == 400, (bad, r.status_code)
    r = c.post(f"/api/ap/captures/{cid}/decide", json={"action": "save", "data": {"fields": {"lines": {"a": 1}, "supplier": "Acme", "invoice_no": f"X-{run}"}}})
    assert r.status_code in (200, 400)
    adm = client(f"fix14a.{run.lower()}@test.local", "admin", ())
    for v in ("1e999", "nan", "inf"):
        assert adm.put("/api/ap/capture-settings", json={"min_confidence": v}).status_code == 400


def test_a_huge_page_is_rasterised_at_most_3000_pixels():
    doc = pymupdf.open()
    p = doc.new_page(width=5000, height=2500)
    p.insert_text((100, 100), "TAX INVOICE", fontsize=40)
    imgs, _ = C.render(doc.tobytes(), "big.pdf", 200, 1)
    pix = pymupdf.Pixmap(imgs[0])
    assert max(pix.width, pix.height) <= 3000 and pix.width > pix.height
    small, _ = C.render(make_ap.invoice(), "a.pdf", 200, 1)  # an A4 page keeps its dpi
    assert pymupdf.Pixmap(small[0]).width == round(595 * 200 / 72)


def test_an_emptied_gl_table_is_red_never_guessed():
    with SessionLocal() as s:
        adm = s.get(User, make_person("fix14gl@test.local", "admin"))
        fixed = [{"key": k, "value": v} for k, v in M.tables(s)["fixed"].items()]
        assert M.save_sap_values(s, adm, fixed, [], confirm=True)["saved"]
        t = M.tables(s)
        assert t["gl"] == {}
        f = {"account": "5012345678", "billing_date": "2026-09-15", "due_date": "2026-10-05", "invoice_no": "INV-1", "statement_month": "SEPTEMBER 2026",
             "services": [{"service": "electricity", "ex": "1000.00", "vat": "150.00"}], "adjustments": {"interest": "10.00"}}
        rows, st, issues = P.utility_doc(f, None, {"vendor": "700999", "cc": "01306", "name1": "B", "muni": "E"}, t, file="e.pdf")
        assert st == "RED" and {"no GL for electricity in the GL Codes table", "no GL for VAT in the GL Codes table", "no GL for interest in the GL Codes table"} <= set(issues)
        rows, st, issues = P.manual_doc({"supplier": "A", "invoice_no": "1", "invoice_date": "2026-10-01", "due_date": "2026-10-31", "excl": "100.00", "vat": "15.00", "incl": "115.00"},
                                        {"vendor_no": "1", "gl": "6100100", "cost_centre": "C"}, "X143", t, file="m.pdf")
        assert st == "RED" and "no GL for VAT in the GL Codes table" in issues
        s.rollback()


# ---- 1.0.98 review: never two jobs with one ref at once; capture keeps a thread free ---------------------------------------------
def _pool(fn):
    from cbx import jobs

    was = jobs._inline
    jobs.set_inline(False)  # the live server: jobs on the pool
    try:
        return fn()
    finally:
        jobs.set_inline(was)


def _wait(cond, secs=5.0):
    import time

    for _ in range(int(secs / 0.05)):
        if cond():
            return True
        time.sleep(0.05)
    return cond()


def test_a_second_job_with_the_same_ref_waits_for_the_running_one_and_then_runs_once_more():
    from cbx import compat2

    started, release, lock = [], threading.Event(), threading.Lock()
    peak = {"now": 0, "max": 0}

    @compat2.handler("reg_1098_same_ref")
    def h(job_id, payload):
        with lock:
            peak["now"] += 1
            peak["max"] = max(peak["max"], peak["now"])
            started.append(payload["n"])
        release.wait(10)
        with lock:
            peak["now"] -= 1

    def go():
        try:
            with SessionLocal() as s:
                assert not compat2.enqueue(s, tool="ap", kind="reg_1098_same_ref", ref="ap_capture:reg", payload={"n": 1}).dropped
                s.commit()
            assert _wait(lambda: started == [1])  # the first is running now
            with SessionLocal() as s:  # the 5-minute wake, or the reader re-queuing itself: not beside it
                j = compat2.enqueue(s, tool="ap", kind="reg_1098_same_ref", ref="ap_capture:reg", payload={"n": 2})
                s.commit()
            with SessionLocal() as s:  # asked again: still only once more
                j2 = compat2.enqueue(s, tool="ap", kind="reg_1098_same_ref", ref="ap_capture:reg", payload={"n": 3})
                s.commit()
            assert j.dropped and j.again and j2.dropped
            import time

            time.sleep(0.3)
            assert started == [1] and peak["max"] == 1  # never two at once
            release.set()
            assert _wait(lambda: len(started) == 2 and not compat2.running("ap_capture:reg"))
            time.sleep(0.2)
            assert started[0] == 1 and len(started) == 2 and peak["max"] == 1  # it ran once more, after
        finally:
            release.set()
            compat2.HANDLERS.pop("reg_1098_same_ref", None)

    _pool(go)


def test_capture_jobs_never_take_more_than_three_threads():
    from cbx import compat2

    assert compat2.MOST.get("ap_capture") == 3
    release, lock = threading.Event(), threading.Lock()
    peak, ran = {"now": 0, "max": 0}, []

    @compat2.handler("reg_1098_lane", most=3)
    def h(job_id, payload):
        with lock:
            peak["now"] += 1
            peak["max"] = max(peak["max"], peak["now"])
        release.wait(10)
        with lock:
            peak["now"] -= 1
            ran.append(payload["w"])

    other = threading.Event()

    def go():
        try:
            with SessionLocal() as s:
                for w in ("cara", "dex", "eli", "seth"):
                    compat2.enqueue(s, tool="ap", kind="reg_1098_lane", ref=f"reg_lane:{w}", payload={"w": w})
                s.commit()
            assert _wait(lambda: peak["now"] == 3)
            from cbx import jobs

            jobs.submit(other.set)  # another tool's job still gets the fourth thread
            assert other.wait(3)
            assert peak["max"] == 3
            release.set()
            assert _wait(lambda: len(ran) == 4)  # the held-back one ran once a thread came free
            assert peak["max"] == 3
        finally:
            release.set()
            compat2.HANDLERS.pop("reg_1098_lane", None)
            compat2.MOST.pop("reg_1098_lane", None)

    _pool(go)
