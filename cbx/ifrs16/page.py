"""IFRS 16 as a page (1.0.92) — CBX-Connect 2's service, line for line (cbx/ifrs16/page_jobs.py stands in for 2.x's jobs
registry and office). IFRS 16: keep the source files, read them for their months, work out a month, check a ledger, and talk to
Sequentra — every slow step a job on the one jobs registry (progress, ETA, the office shows who is on it).

The engine (cbx/ifrs16/engine — the finance team's desktop IFRS 16 Workflow 5.16, unmodified) is called with
exactly the arguments the old CBX used (1.0.80 tool folder, cbx/engine/actions_core.py):
    run_workflow_from_files(la, movements, sequentra, workflow_month=, steps_path=, codes_path=, rates_path=,
                            assets_path=, ciw_master_path=, movement_label=, steps_df=, sequentra_df=)
    → write_import_workbooks + write_monthly_report (+ every tab as the working papers)
    run_ledger_check(ledger, as_at=, months_back=) → write_consolidated_workbook
"""
from __future__ import annotations

import calendar
import io
import logging
import shutil
import tempfile
import zipfile
from datetime import date
from pathlib import Path
from typing import Optional

from sqlalchemy.orm import Session

from .. import aicheck
from . import page_jobs as jobs
from ..audit import log
from ..config import settings
from ..db import session_scope, utcnow
from ..models import StoredFile, User
from ..models.ifrs16_page import IfrsApiRun, IfrsLedgerCheck, IfrsMaster, IfrsMonth, IfrsRate, IfrsRecon, IfrsUpload
from ..warehouse import storage
from . import CODES_PATH, scanning, template_info

logger = logging.getLogger("cbx.ifrs16")
SLUG = "ifrs16"
MAX_MB = 100
REPORT_NAME = "IFRS16_Report_%s.xlsx"
LEDGER_REPORT_NAME = "IFRS16_Consolidated_Report.xlsx"
RECON_REPORT_NAME = "LA_vs_RR06_Reconciliation_%s.xlsx"
RECON_LIST_CAP = 5000
FILE_KEYS = ("la", "movements", "sequentra", "steps", "assets", "ciw")
FILE_LABELS = {"la": "The LA register", "movements": "The RR06 / ARO06 movements", "sequentra": "The Sequentra export", "steps": "The step payments report",
               "assets": "The LA BU asset report", "ciw": "The CIW master"}


class Refused(ValueError):
    pass


def _log_failed(what: str, ident, e: Exception) -> None:
    """1.0.106: a file problem (a refusal, or the engine's plain-words ValueError: the wrong file, a missing column) is one
    line in the log - the same words the person sees on the row; anything else keeps its traceback for Edwin."""
    if isinstance(e, ValueError):
        logger.warning("ifrs16 %s %s: %s", what, ident, str(e)[:500])
    else:
        logger.error("ifrs16 %s %s failed", what, ident, exc_info=e)


# ---- small helpers ------------------------------------------------------------------------------------------
def month_date(month: str) -> date:
    y, m = int(month[:4]), int(month[5:7])
    if len(month) != 7 or month[4] != "-" or not (1 <= m <= 12):
        raise ValueError(month)
    return date(y, m, 1)


def parse_as_at(text: str) -> Optional[date]:
    s = (text or "").strip()
    if not s:
        return None
    if len(s) == 7:
        d = month_date(s)
        return date(d.year, d.month, calendar.monthrange(d.year, d.month)[1])
    return date.fromisoformat(s[:10])


def _scratch() -> Path:
    return Path(tempfile.mkdtemp(prefix="ifrs16_", dir=str(settings.path("tmp"))))


def _path(db: Session, fid: Optional[int]) -> Optional[str]:
    if not fid:
        return None
    f = db.get(StoredFile, fid)
    if f is None:
        return None
    p = storage.abs_path(f)
    return str(p) if p.is_file() else None


def _name(db: Session, fid: Optional[int]) -> str:
    f = db.get(StoredFile, fid) if fid else None
    return f.name if f else ""


def _keep(db: Session, user: User, name: str, data: bytes, meta: dict) -> StoredFile:
    return storage.store(db, user, name, data, sub="Results", source="result", mirror=False, meta={"purpose": SLUG, **meta})


def can_see(me: User, row) -> bool:
    return me.is_admin or row.user_id == me.id


# ---- uploads ----------------------------------------------------------------------------------------------
def create_upload(db: Session, user: User, files: dict[str, tuple[str, bytes]]) -> IfrsUpload:
    """Keep the files in the person's Inbox; read them in the background when both required files are there."""
    for k, (name, data) in files.items():
        if k not in FILE_KEYS:
            raise Refused(f"Unknown file '{k}'.")
        if len(data) > MAX_MB * 1024 * 1024:
            raise Refused(f"{FILE_LABELS[k]} is over {MAX_MB} MB.")
        if not name.lower().endswith((".xlsx", ".xls", ".xlsm", ".csv")):
            raise Refused(f"{FILE_LABELS[k]} must be an Excel file.")
    if "la" not in files or "movements" not in files:
        raise Refused("Both the LA register and the RR06 / ARO06 movements are needed.")
    u = IfrsUpload(user_id=user.id, status="reading")
    for k, (name, data) in files.items():
        f = storage.store(db, user, name, data, sub="Inbox", source=SLUG, mirror=False, meta={"purpose": SLUG, "ifrs16": k})
        setattr(u, f"{k}_file_id", f.id)
    db.add(u)
    db.flush()
    jobs.enqueue(db, tool=SLUG, kind="ifrs16_scan", title="Read the IFRS 16 source files", user=user, ref=f"ifrs_upload:{u.id}", total=1, payload={"upload_id": u.id})
    log(db, "ifrs16_uploaded", actor=user, target=f"ifrs_upload:{u.id}", detail=", ".join(f"{k}={n}" for k, (n, _) in files.items())[:400])
    return u


@jobs.handler("ifrs16_scan")
def scan_job(job_id: int, payload: dict) -> dict:
    from .page_jobs import company as CO

    with session_scope() as db:
        u = db.get(IfrsUpload, payload["upload_id"])
        if u is None:
            return {}
        paths = {k: _path(db, getattr(u, f"{k}_file_id")) for k in ("la", "movements", "sequentra", "steps")}
        with CO.working(SLUG, "calculator", task=f"Reading IFRS 16 upload #{u.id}", ref=f"ifrs_upload:{u.id}", db=db) as a:
            try:
                if not paths["la"] or not paths["movements"]:
                    scan, err = {}, "The stored source files are no longer available — upload them again."
                else:
                    scan, err = scanning.scan_sources(paths["la"], paths["movements"], paths["sequentra"], paths["steps"])
                    if err and aicheck.check("IFRS 16: reading the uploaded files", err)["fixed"]:  # e.g. a reader library installed
                        scan, err = scanning.scan_sources(paths["la"], paths["movements"], paths["sequentra"], paths["steps"])
            except Exception as e:  # noqa: BLE001
                _log_failed("scan", u.id, e)
                scan, err = {}, f"Reading the files failed: {str(e)[:300]}"
            u.scan = scan or {}
            u.status = "failed" if err else "ready"
            if err:
                u.note = aicheck.note("IFRS 16: reading the uploaded files", err, ", ".join(filter(None, (_name(db, getattr(u, f"{k}_file_id")) for k in FILE_KEYS))))
                a.say(f"Could not read upload #{u.id}: {err[:200]}")
            else:
                months = scan.get("months", [])
                parts = [f"{scan['la_rows']} LA register rows", f"{scan['move_rows']} movement rows in {len(months)} month{'s' if len(months) != 1 else ''}"]
                if scan.get("sequentra_rows"):
                    parts.append(f"{scan['sequentra_rows']} Sequentra rows")
                if scan.get("steps_rows"):
                    parts.append(f"{scan['steps_rows']} step payments")
                u.note = " · ".join(parts)
                a.say(f"Upload #{u.id}: {u.note}.")
            # 1.0.106: the read is in the audit trail too (a file that could not be read left no trace there)
            log(db, "ifrs16_read", actor=db.get(User, u.user_id), target=f"ifrs_upload:{u.id}", detail=(u.note or "")[:400],
                outcome="failed" if err else "ok")
        db.commit()
        jobs.progress(job_id, done=1, total=1, note=u.note[:200])
        return {"status": u.status, "link": "#/ifrs16"}


def read_again(db: Session, user: User, u: IfrsUpload) -> IfrsUpload:
    """A failed upload, read again with the same files (after a fix on the server, or the AI check's)."""
    if not can_see(user, u):
        raise PermissionError("That upload isn't yours.")
    if u.status == "reading":
        raise Refused("It's being read now.")
    u.status, u.note = "reading", ""
    jobs.enqueue(db, tool=SLUG, kind="ifrs16_scan", title="Read the IFRS 16 source files", user=user, ref=f"ifrs_upload:{u.id}", total=1, payload={"upload_id": u.id})
    log(db, "ifrs16_read_again", actor=user, target=f"ifrs_upload:{u.id}")
    return u


# ---- a month ----------------------------------------------------------------------------------------------
def start_month(db: Session, user: User, u: IfrsUpload, month: str, *, source: str = "files", use_charges: bool = False) -> IfrsMonth:
    if not can_see(user, u):
        raise PermissionError("That upload isn't yours.")
    if u.status != "ready":
        raise Refused("These files are still being read, or could not be read — see the note on the upload.")
    try:
        month_date(month)
    except (ValueError, IndexError):
        raise Refused(f"'{month}' is not a month — write it like 2026-08.") from None
    held = [m["value"] for m in (u.scan or {}).get("months", [])]
    if month not in held:
        raise Refused(f"That month is not in the movements file. It holds: {', '.join(held) or 'nothing'}.")
    m = IfrsMonth(upload_id=u.id, user_id=user.id, month=month, source="api" if source == "api" else "files", use_charges=bool(use_charges), status="working")
    db.add(m)
    db.flush()
    j = jobs.enqueue(db, tool=SLUG, kind="ifrs16_month", title=f"IFRS 16 · {month}" + (" · Sequentra API" if m.source == "api" else ""), user=user,
                     ref=f"ifrs_month:{m.id}", total=4, payload={"month_id": m.id})
    m.job_id = j.id
    log(db, "ifrs16_month_started", actor=user, target=f"ifrs_month:{m.id}", detail=f"{month} from {'the Sequentra API' if m.source == 'api' else 'the files'}")
    return m


def rates_workbook(db: Session, folder: Path) -> Optional[Path]:
    """The Rates table written the way the engine reads its rates workbook; None when the table is empty."""
    rows = db.query(IfrsRate).order_by(IfrsRate.month).all()
    if not rows:
        return None
    import lease_engine  # noqa: the engine module, unmodified

    path = folder / "discount_rates.xlsx"
    for r in rows:
        try:
            m = month_date(r.month)
        except (ValueError, IndexError):
            continue
        bands = {b.upper(): getattr(r, b) for b in "abcdef" if getattr(r, b) is not None}
        if bands:
            lease_engine.save_rate_month(str(path), m, bands)
    return path if path.exists() else None


def write_papers(results: dict, path: Path) -> None:
    """Every tab the run produced in one workbook — what the desktop shows as its grid tabs."""
    import pandas as pd

    used: set[str] = set()
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in results.items():
            if not isinstance(df, pd.DataFrame):
                continue
            sheet = "".join(ch for ch in str(name) if ch not in "[]:*?/\\")[:31] or "Sheet"
            n = 2
            while sheet in used:
                sheet = f"{sheet[:28]} {n}"
                n += 1
            used.add(sheet)
            df.to_excel(xw, sheet_name=sheet, index=False)
        if not used:
            pd.DataFrame({"Note": ["The run produced no tabs."]}).to_excel(xw, sheet_name="Note", index=False)


def run_month(db: Session, m: IfrsMonth, tmp: Path, progress=lambda *a, **k: None) -> dict:
    """The engine run itself (also used by the golden test). Returns {workbooks, report, papers, imports, recon}
    as (name, bytes) pairs and text."""
    from . import api as SQ

    import lease_engine  # noqa: the engine module, unmodified

    u = db.get(IfrsUpload, m.upload_id) if m.upload_id else None
    if u is None:
        raise Refused("The upload this month came from is gone — upload the files again.")
    wm = month_date(m.month)
    s = {k: _path(db, getattr(u, f"{k}_file_id")) for k in FILE_KEYS}
    if not s["la"] or not s["movements"]:
        raise Refused("The stored source files for that upload are no longer available — upload them again.")
    out = tmp / "out"
    out.mkdir()
    label = "ARO06" if "aro" in (_name(db, u.movements_file_id) or s["movements"]).lower() else "RR06"
    kind = "ARO" if label == "ARO06" else "SA"
    seq_df = steps_df = cfg = None
    notes: list[str] = []
    if m.source == "api":
        progress(1, "asking Sequentra for the lease facts")
        try:
            cfg = SQ.config(db, use_charges=m.use_charges)
            seq_df, steps_df, notes = SQ.facts_and_steps(cfg, kind, wm)
        except SQ.ApiProblem as e:
            raise Refused(f"{e} — or work the month out from the files instead.") from None
        if steps_df is None:
            notes.append(f"Step payments taken from the stored file '{_name(db, u.steps_file_id)}'" if s["steps"] else
                         "No step payments: add the step payments file until the API charge reports work - payment schedules are flagged, not drafted")
    rates = rates_workbook(db, tmp)
    progress(2, "the lease engine is working the month out")
    results = lease_engine.run_workflow_from_files(s["la"], s["movements"], s["sequentra"], workflow_month=wm, steps_path=s["steps"], codes_path=str(CODES_PATH),
                                                   rates_path=str(rates) if rates else None, assets_path=s["assets"], ciw_master_path=s["ciw"],
                                                   movement_label=label, steps_df=steps_df, sequentra_df=seq_df)
    if m.source == "api":
        SQ.cross_check(results, seq_df, cfg, kind, wm, notes)
    progress(3, "writing the import workbooks and the report")
    written = lease_engine.write_import_workbooks(results, str(out))
    report = out / (REPORT_NAME % m.month)
    lease_engine.write_monthly_report(results, str(report))
    summary = template_info.run_summary(results, written)
    zbuf = io.BytesIO()
    with zipfile.ZipFile(zbuf, "w", zipfile.ZIP_DEFLATED) as z:
        for p in sorted(out.iterdir()):
            if p.is_file() and p != report:
                z.write(p, p.name)
    papers = tmp / f"IFRS16_Working_papers_{m.month}.xlsx"
    write_papers(results, papers)
    recon = "\n".join(f"{x.get('item')}: {x.get('value')}" for x in summary.get("recon") or [])
    if m.source == "api":
        recon += "".join(f"\nAPI: {n}" for n in notes if f"API: {n}" not in recon)
    if not rates and kind == "SA":
        recon += "\nDiscount rates: none on the Rates page — SA rates came from Sequentra (flagged for review)."
    return {"workbooks": (f"IFRS16_Imports_{m.month}.zip", zbuf.getvalue()) if written else None, "report": (report.name, report.read_bytes()),
            "papers": (papers.name, papers.read_bytes()), "imports": len(written),
            "recon": recon or (f"{m.month} ran, but the engine found nothing to import for that month — only the report was written." if not written else "")}


@jobs.handler("ifrs16_month")
def month_job(job_id: int, payload: dict) -> dict:
    from .page_jobs import company as CO

    tmp = _scratch()
    try:
        with session_scope() as db:
            m = db.get(IfrsMonth, payload["month_id"])
            if m is None:
                return {}
            user = db.get(User, m.user_id)
            with CO.working(SLUG, "calculator", task=f"Working out {m.month}", ref=f"ifrs_month:{m.id}", db=db) as a:
                def prog(n, note):
                    a.progress(n / 4, note)
                    db.commit()  # one writer at a time on SQLite: the job's row is written in its own session
                    jobs.progress(job_id, done=n, total=4, note=note)
                try:
                    r = aicheck.with_retry(lambda: run_month(db, m, tmp, prog), f"IFRS 16: working out {m.month}")
                    if r["workbooks"]:
                        m.workbooks_file_id = _keep(db, user, *r["workbooks"], {"month": m.month, "what": "workbooks"}).id
                    m.report_file_id = _keep(db, user, *r["report"], {"month": m.month, "what": "report"}).id
                    m.papers_file_id = _keep(db, user, *r["papers"], {"month": m.month, "what": "papers"}).id
                    m.imports, m.recon, m.status, m.note = r["imports"], r["recon"][:20000], "done", ""
                    a.say(f"{m.month}: {m.imports} import workbook{'s' if m.imports != 1 else ''} and the report are ready.")
                    log(db, "ifrs16_month_done", actor=user, target=f"ifrs_month:{m.id}", detail=f"{m.month}: {m.imports} import workbook(s)")
                except Exception as e:  # noqa: BLE001 — a month that fails says why on its row
                    _log_failed("month", m.id, e)
                    m.status = "failed"
                    m.note = aicheck.note(f"IFRS 16: working out {m.month}", str(e) if isinstance(e, Refused) else f"Working out {m.month} failed: {str(e)[:400]}")
                    a.say(f"{m.month} failed: {m.note[:200]}")
                    log(db, "ifrs16_month_done", actor=user, target=f"ifrs_month:{m.id}", detail=m.note[:400], outcome="failed")
            m.finished_at = utcnow()
            db.commit()
            jobs.progress(job_id, done=4, total=4, note="done" if m.status == "done" else "failed")
            return {"status": m.status, "link": "#/ifrs16", "note": m.note}
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ---- the ledger check (as the desktop's Report) -------------------------------------------------------------
SPANS = {"life": (True, False), "from": (True, True), "back": (False, False)}  # (full_life, from_selected)
SPAN_WORDS = {"life": "full lease life", "from": "from the as-at month on", "back": "only the months back"}


def _latest_upload_file(db: Session, user: User, key: str) -> Optional[int]:
    """The person's newest kept upload that has this file (the Sequentra export, the LA register)."""
    col = getattr(IfrsUpload, f"{key}_file_id")
    u = db.query(IfrsUpload).filter(IfrsUpload.user_id == user.id, col.isnot(None)).order_by(IfrsUpload.id.desc()).first()
    return getattr(u, f"{key}_file_id") if u else None


def _source_upload(db: Session, user: User, upload_id: Optional[int]) -> Optional[IfrsUpload]:
    """The kept upload whose files fill the report's lease names, dates and entities: the one asked for, else the
    person's newest."""
    if upload_id:
        u = db.get(IfrsUpload, upload_id)
        return u if u is not None and can_see(user, u) else None
    return db.query(IfrsUpload).filter(IfrsUpload.user_id == user.id).order_by(IfrsUpload.id.desc()).first()


def _from_papers(db: Session, u: Optional[IfrsUpload]) -> tuple[dict, object]:
    """What the desktop takes from the loaded run: the leases it ruled out (Not Applicable) and the RR06/ARO06
    movements — from the working papers of the upload's newest finished month."""
    import pandas as pd

    if u is None:
        return {}, None
    m = db.query(IfrsMonth).filter(IfrsMonth.upload_id == u.id, IfrsMonth.status == "done", IfrsMonth.papers_file_id.isnot(None)).order_by(IfrsMonth.id.desc()).first()
    p = _path(db, m.papers_file_id) if m else None
    if not p:
        return {}, None
    try:
        book = pd.read_excel(p, sheet_name=None)
    except Exception:  # noqa: BLE001
        return {}, None
    na = {}
    nadf = book.get("Not Applicable")
    if nadf is not None and len(nadf) and "Lease Reference" in nadf.columns:
        for _, r in nadf.iterrows():
            na[str(r.get("Lease Reference", "")).strip().upper()] = str(r.get("Reason", ""))
    contracts = next((book[k] for k in ("RR06 Movements", "ARO06 Movements") if k in book), None)
    return na, contracts


def _as_at_or_today(as_at: Optional[date]) -> date:
    d = as_at or date.today()
    return date(d.year, d.month, calendar.monthrange(d.year, d.month)[1])


def start_ledger(db: Session, user: User, name: str, data: bytes, as_at: str, months_back: int,
                 sequentra: Optional[tuple[str, bytes]] = None, upload_id: Optional[int] = None, entity: str = "", span: str = "life") -> IfrsLedgerCheck:
    try:
        parse_as_at(as_at)
    except ValueError:
        raise Refused("The 'as at' month must look like 2026-08 (year-month).") from None
    if not 1 <= int(months_back) <= 120:
        raise Refused("Months back must be between 1 and 120.")
    if not name.lower().endswith((".xlsx", ".xls")):
        raise Refused("The ledger export must be an Excel file.")
    if sequentra and not sequentra[0].lower().endswith((".xlsx", ".xls", ".csv")):
        raise Refused("The Sequentra validation report must be an Excel file.")
    for what, blob in (("The ledger export", data), ("The Sequentra validation report", sequentra[1] if sequentra else b"")):
        if len(blob) > MAX_MB * 1024 * 1024:  # 1.0.106: was read cut short at the limit and failed as 'not a ledger'
            raise Refused(f"{what} is over {MAX_MB} MB.")
    if span not in SPANS:
        raise Refused("Choose full lease life, from the as-at month, or only the months back.")
    u = _source_upload(db, user, upload_id)
    if upload_id and u is None:
        raise Refused("That upload isn't yours.")
    f = storage.store(db, user, name, data, sub="Inbox", source=SLUG, mirror=False, meta={"purpose": SLUG, "ifrs16": "ledger"})
    if sequentra:
        seq_id = storage.store(db, user, sequentra[0], sequentra[1], sub="Inbox", source=SLUG, mirror=False, meta={"purpose": SLUG, "ifrs16": "sequentra"}).id
    else:  # the Sequentra export kept with that upload, else the newest one that has one
        seq_id = (u.sequentra_file_id if u is not None and u.sequentra_file_id else None) or _latest_upload_file(db, user, "sequentra")
    r = IfrsLedgerCheck(user_id=user.id, ledger_file_id=f.id, sequentra_file_id=seq_id, as_at=(as_at or "").strip()[:10], months_back=int(months_back),
                        upload_id=u.id if u is not None else None, entity=(entity or "").strip()[:160], span=span)
    db.add(r)
    db.flush()
    j = jobs.enqueue(db, tool=SLUG, kind="ifrs16_ledger", title="IFRS 16 ledger check", user=user, ref=f"ifrs_ledger:{r.id}", total=1, payload={"check_id": r.id})
    r.job_id = j.id
    return r


def run_ledger(db: Session, r: IfrsLedgerCheck, tmp: Path) -> dict:
    """The desktop's Report, with the same inputs as its web interface: the ledger export, the leases the last run
    ruled out, the Sequentra validation export, the BU asset report, the entity data (the RR06), the LA report.
    Writes the full Consolidated Report (pivot, Data, period sheet, Recon) and the simple report."""
    import ledger_check  # noqa: the engine module

    p = _path(db, r.ledger_file_id)
    if not p:
        raise Refused("The ledger export is no longer on the server — upload it again.")
    user = db.get(User, r.user_id)
    u = db.get(IfrsUpload, r.upload_id) if r.upload_id else None
    na, _contracts = _from_papers(db, u)  # the movements list is not used to filter, as the desktop's web interface
    full_life, from_selected = SPANS.get(r.span or "life", SPANS["life"])
    as_at = _as_at_or_today(parse_as_at(r.as_at))
    res = ledger_check.run_ledger_check(
        p, not_applicable=na, as_at=as_at, months_back=r.months_back or 12, full_life=full_life, from_selected=from_selected,
        sequentra_path=_path(db, r.sequentra_file_id), assets_path=_path(db, u.assets_file_id) if u else None,
        entity_path=_path(db, u.movements_file_id) if u else None, la_path=_path(db, u.la_file_id) if u else _path(db, _latest_upload_file(db, user, "la")))
    con = res.get("Report Consolidated")
    if con is None or not len(con):
        led = res.get("_ledger")
        have = sorted(led["Period"].dropna().unique()) if led is not None and "Period" in getattr(led, "columns", []) and len(led) else []
        raise Refused("The report has no rows for those periods" + (f" (the ledger runs {have[0]} to {have[-1]})" if have else "")
                      + ". Pick an earlier as-at month, or Full lease life.")
    full = tmp / f"Consolidated_Report_FULL_SA_{as_at:%B_%Y}.xlsx"
    ledger_check.write_consolidated_workbook(res, str(full), ledger=res.get("_ledger"), refs=res.get("_refs"), as_at=as_at, entity=r.entity or "")
    simple = tmp / f"Lease_Report_SIMPLE_{as_at:%B_%Y}.xlsx"
    try:
        ledger_check.write_simple_report(res, str(simple), as_at=as_at)
    except Exception:  # noqa: BLE001 — as the desktop: the full report stands on its own
        simple = None
    lines = []
    try:
        for _, row in res["Report Overview"].iterrows():
            if row["Item"] != "Net asset roll-forward breaks":  # said once, below, with the leases
                lines.append(f"{row['Item']}: {row['Value']}")
    except Exception:  # noqa: BLE001
        pass
    lines.append(f"Periods: {SPAN_WORDS.get(r.span or 'life')}")
    lines.append("Lease names, dates and entities from: " + (", ".join(n for n in (_name(db, u.la_file_id), _name(db, u.movements_file_id), _name(db, u.assets_file_id)) if n) if u else "no kept upload (add one on Monthly workflow)"))
    lines.append("Discount rates from: " + (_name(db, r.sequentra_file_id) or "no Sequentra validation report (add it to fill the rates)"))
    if na:
        lines.append(f"Leases the last run ruled out (Not Applicable): {len(na)}")
    con = res.get("Report Consolidated")
    if con is not None and len(con) and con["Lease Name"].astype(str).str.strip().isin(("", "nan", "None")).all():
        lines.append("NOTE: Lease Name, the term dates and Entity are blank — keep the Sequentra validation export with your upload (or add it here) and run it again.")
    chk = res.get("Check Net Asset")
    if chk is not None and len(chk):
        bad = chk[chk["Status"] != "OK"]
        lines.append(f"Net asset check (opening + movements = closing): {len(chk) - len(bad)} of {len(chk)} rows agree")
        for _, b in bad.head(10).iterrows():
            lines.append(f"  {b['Lease Reference']} {b['Period']}: opening {b['Opening']:,.2f} + movements {b['Movements']:,.2f} vs closing {b['Closing']:,.2f} (difference {b['Difference']:,.2f})")
        if len(bad) > 10:
            lines.append(f"  and {len(bad) - 10} more")
    return {"full": (full.name, full.read_bytes()), "simple": (simple.name, simple.read_bytes()) if simple else None, "overview": "\n".join(lines[:50])}


@jobs.handler("ifrs16_ledger")
def ledger_job(job_id: int, payload: dict) -> dict:
    from .page_jobs import company as CO

    tmp = _scratch()
    try:
        with session_scope() as db:
            r = db.get(IfrsLedgerCheck, payload["check_id"])
            if r is None:
                return {}
            user = db.get(User, r.user_id)
            with CO.working(SLUG, "reconciler", task="Checking a ledger", ref=f"ifrs_ledger:{r.id}", db=db):
                try:
                    out = aicheck.with_retry(lambda: run_ledger(db, r, tmp), "IFRS 16: the ledger check")
                    r.report_file_id = _keep(db, user, *out["full"], {"what": "ledger"}).id
                    if out["simple"]:
                        r.simple_file_id = _keep(db, user, *out["simple"], {"what": "ledger simple"}).id
                    r.overview, r.status, r.note = out["overview"], "done", ""
                except Exception as e:  # noqa: BLE001
                    _log_failed("ledger", r.id, e)
                    r.status, r.note = "failed", aicheck.note("IFRS 16: the ledger check", str(e) if isinstance(e, Refused) else f"Ledger check failed: {str(e)[:400]}")
            log(db, "ifrs16_ledger_checked", actor=user, target=f"ifrs_ledger:{r.id}", detail=r.note[:400] or "report written", outcome="ok" if r.status == "done" else "failed")
            db.commit()
            jobs.progress(job_id, done=1, total=1, note=r.status)
            return {"status": r.status, "link": "#/ifrs16?tab=ledger"}
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ---- the Master report (as the desktop) ---------------------------------------------------------------------
SHEET_GROUPS = [
    ("Report", ("Dashboard", "Report ")),
    ("Recon & Review", ("Recon Summary", "Routing", "Exceptions", "Not Applicable", "IBR Check", "Monthly Report")),
    ("RR06", ("RR06 ",)), ("ARO06", ("ARO06 ",)), ("PIW", ("PIW ",)), ("CIW", ("CIW ",)), ("MDIW", ("MDIW ",)),
    ("Rent Review", ("Rent Review",)), ("Terminations", ("Terminations",)),
]


def _group_of(name: str) -> str:
    for group, keys in SHEET_GROUPS:
        if any(name == k or name.startswith(k) for k in keys):
            return group
    return "Other"


def write_master_workbook(sheets: dict, out_path: Path) -> int:
    """Every sheet in one workbook, in the desktop's group order, with a Contents page in front (finance_harvester
    _write_master_workbook). Returns the number of sheets after Contents."""
    import pandas as pd
    from openpyxl.styles import Font, PatternFill

    order = [g for g, _ in SHEET_GROUPS] + ["Other"]
    groups: dict[str, list[str]] = {g: [] for g in order}
    for name in sheets:
        if not str(name).startswith("_") and isinstance(sheets[name], pd.DataFrame):
            groups[_group_of(str(name))].append(str(name))
    used = {"Contents"}
    tabs = []
    for g in order:
        for n in groups[g]:
            t = "".join(ch for ch in n if ch not in "[]:*?/\\")[:31] or "Sheet"
            k = 2
            while t in used:
                t = f"{t[:28]} {k}"
                k += 1
            used.add(t)
            tabs.append((g, n, t))
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="003F2D")
    with pd.ExcelWriter(out_path, engine="openpyxl") as xl:
        pd.DataFrame([{"Section": g, "Sheet": t, "Rows": int(len(sheets[n]))} for g, n, t in tabs]).to_excel(xl, sheet_name="Contents", index=False)
        for _g, n, t in tabs:
            sheets[n].to_excel(xl, sheet_name=t, index=False)
        for ws in xl.book.worksheets:
            for cell in ws[1]:
                if cell.value is not None:
                    cell.font, cell.fill = head_font, head_fill
            ws.freeze_panes = "A2"
    return len(tabs)


def start_master(db: Session, user: User, u: IfrsUpload, month: str, as_at: str = "", entity: str = "",
                 ledger: Optional[tuple[str, bytes]] = None) -> IfrsMaster:
    if not can_see(user, u):
        raise PermissionError("That upload isn't yours.")
    if u.status != "ready":
        raise Refused("These files are still being read, or could not be read — see the note on the upload.")
    held = [m["value"] for m in (u.scan or {}).get("months", [])]
    if month not in held:
        raise Refused(f"That month is not in the movements file. It holds: {', '.join(held) or 'nothing'}.")
    try:
        parse_as_at(as_at)
    except ValueError:
        raise Refused("The 'as at' month must look like 2026-08 (year-month).") from None
    lid = None
    if ledger:
        if not ledger[0].lower().endswith((".xlsx", ".xls")):
            raise Refused("The ledger export must be an Excel file.")
        if len(ledger[1]) > MAX_MB * 1024 * 1024:  # 1.0.106
            raise Refused(f"The ledger export is over {MAX_MB} MB.")
        lid = storage.store(db, user, ledger[0], ledger[1], sub="Inbox", source=SLUG, mirror=False, meta={"purpose": SLUG, "ifrs16": "ledger"}).id
    r = IfrsMaster(user_id=user.id, upload_id=u.id, month=month, ledger_file_id=lid, as_at=(as_at or "").strip()[:10], entity=(entity or "").strip()[:160])
    db.add(r)
    db.flush()
    j = jobs.enqueue(db, tool=SLUG, kind="ifrs16_master", title=f"IFRS 16 · master report {month}", user=user, ref=f"ifrs_master:{r.id}", total=1, payload={"master_id": r.id})
    r.job_id = j.id
    log(db, "ifrs16_master_started", actor=user, target=f"ifrs_master:{r.id}", detail=f"{month}" + (" with the ledger check" if lid else ""))
    return r


def run_master(db: Session, r: IfrsMaster, tmp: Path) -> tuple[str, bytes, int]:
    import ledger_check  # noqa: the engine module
    import lease_engine  # noqa: the engine module

    u = db.get(IfrsUpload, r.upload_id) if r.upload_id else None
    if u is None:
        raise Refused("The upload this came from is gone — upload the files again.")
    s = {k: _path(db, getattr(u, f"{k}_file_id")) for k in FILE_KEYS}
    if not s["la"] or not s["movements"]:
        raise Refused("The stored source files for that upload are no longer available — upload them again.")
    wm = month_date(r.month)
    label = "ARO06" if "aro" in (_name(db, u.movements_file_id) or s["movements"]).lower() else "RR06"
    rates = rates_workbook(db, tmp)
    results = lease_engine.run_workflow_from_files(s["la"], s["movements"], s["sequentra"], workflow_month=wm, steps_path=s["steps"], codes_path=str(CODES_PATH),
                                                   rates_path=str(rates) if rates else None, assets_path=s["assets"], ciw_master_path=s["ciw"], movement_label=label)
    master = {str(k): v for k, v in results.items() if not str(k).startswith("_")}
    tag = None
    recon = results.get("Recon Summary")
    if recon is not None:
        for _, x in recon.iterrows():
            if str(x["Item"]).strip().lower() == "workflow month":
                tag = str(x["Value"]).replace(" ", "")
    led = _path(db, r.ledger_file_id)
    if led:
        import pandas as pd

        as_at = _as_at_or_today(parse_as_at(r.as_at) or wm)
        na = {}
        nadf = results.get("Not Applicable")
        if nadf is not None and len(nadf):
            for _, x in nadf.iterrows():
                na[str(x.get("Lease Reference", "")).strip().upper()] = str(x.get("Reason", ""))
        res = ledger_check.run_ledger_check(led, not_applicable=na, contracts=results.get(label + " Movements"), as_at=as_at, sequentra_path=s["sequentra"],
                                            assets_path=s["assets"], entity_path=s["movements"])
        master.update({str(k): v for k, v in res.items() if not str(k).startswith("_")})
        try:
            rows = ledger_check._recon_rows(res.get("_ledger"), res.get("_refs"), as_at, r.entity or "")
            width = max(len(x) for x in rows)
            master["Report Recon"] = pd.DataFrame([list(x) + [""] * (width - len(x)) for x in rows],
                                                  columns=["Recon", "Line", "LEASE LIABILITY", "ROU ASSET", "ACC DEP", "Impairment"][:width])
        except Exception:  # noqa: BLE001 — as the desktop
            pass
        tag = tag or as_at.strftime("%B%Y")
    out = tmp / f"IFRS16Workflow_Master_Report_{tag or r.month}.xlsx"
    n = write_master_workbook(master, out)
    return out.name, out.read_bytes(), n


@jobs.handler("ifrs16_master")
def master_job(job_id: int, payload: dict) -> dict:
    from .page_jobs import company as CO

    tmp = _scratch()
    try:
        with session_scope() as db:
            r = db.get(IfrsMaster, payload["master_id"])
            if r is None:
                return {}
            user = db.get(User, r.user_id)
            with CO.working(SLUG, "calculator", task=f"Master report {r.month}", ref=f"ifrs_master:{r.id}", db=db):
                try:
                    name, data, n = aicheck.with_retry(lambda: run_master(db, r, tmp), "IFRS 16: the master report")
                    r.report_file_id = _keep(db, user, name, data, {"month": r.month, "what": "master"}).id
                    r.sheets, r.status, r.note = n, "done", ""
                except Exception as e:  # noqa: BLE001
                    _log_failed("master", r.id, e)
                    r.status, r.note = "failed", aicheck.note("IFRS 16: the master report", str(e) if isinstance(e, Refused) else f"The master report failed: {str(e)[:400]}")
            log(db, "ifrs16_master_done", actor=user, target=f"ifrs_master:{r.id}", detail=r.note[:400] or f"{r.sheets} sheets", outcome="ok" if r.status == "done" else "failed")
            db.commit()
            jobs.progress(job_id, done=1, total=1, note=r.status)
            return {"status": r.status, "link": "#/ifrs16"}
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ---- LA vs RR06 (ABSA) -------------------------------------------------------------------------------------
def start_recon(db: Session, user: User, period: str, la: Optional[tuple[str, bytes]] = None, rr06: Optional[tuple[str, bytes]] = None,
                upload: Optional[IfrsUpload] = None) -> IfrsRecon:
    """The LA transaction report against the RR06 contract list, for a period (YYYY-MM). Files uploaded here, or
    the ones kept with an upload (its LA register and its RR06)."""
    try:
        month_date((period or "").strip())
    except (ValueError, IndexError):
        raise Refused("The period must look like 2026-09 (year-month).") from None
    ids = {}
    for key, given, kept in (("la", la, "la_file_id"), ("rr06", rr06, "movements_file_id")):
        if given:
            name, data = given
            if len(data) > MAX_MB * 1024 * 1024:
                raise Refused(f"{name} is over {MAX_MB} MB.")
            if not name.lower().endswith((".xlsx", ".xls", ".xlsm")):
                raise Refused(("The LA report" if key == "la" else "The RR06") + " must be an Excel file.")
            ids[key] = storage.store(db, user, name, data, sub="Inbox", source=SLUG, mirror=False, meta={"purpose": SLUG, "ifrs16": key}).id
        elif upload is not None and getattr(upload, kept):
            ids[key] = getattr(upload, kept)
    if "la" not in ids or "rr06" not in ids:
        raise Refused("Both the LA transaction report and the RR06 contract list are needed.")
    r = IfrsRecon(user_id=user.id, la_file_id=ids["la"], rr06_file_id=ids["rr06"], period=period.strip()[:7], result={})
    db.add(r)
    db.flush()
    j = jobs.enqueue(db, tool=SLUG, kind="ifrs16_recon", title="IFRS 16 · LA vs RR06", user=user, ref=f"ifrs_recon:{r.id}", total=1, payload={"recon_id": r.id})
    r.job_id = j.id
    log(db, "ifrs16_recon_started", actor=user, target=f"ifrs_recon:{r.id}", detail=f"{r.period}: {_name(db, r.la_file_id)} vs {_name(db, r.rr06_file_id)}"[:400])
    return r


def _records(df, cap: int) -> list[dict]:
    import math
    from datetime import date as _d, datetime as _dt

    out = []
    for row in df.head(cap).to_dict("records"):
        clean = {}
        for k, v in row.items():
            if isinstance(v, float) and math.isnan(v):
                v = None
            elif isinstance(v, (_dt, _d)):
                v = v.strftime("%d/%m/%Y")
            elif v is not None and not isinstance(v, (int, float, str, bool)):
                v = str(v)
                v = "" if v in ("NaT", "nan") else v
            clean[str(k)] = v
        out.append(clean)
    return out


@jobs.handler("ifrs16_recon")
def recon_job(job_id: int, payload: dict) -> dict:
    import rr06_recon  # noqa: the engine module

    tmp = _scratch()
    try:
        with session_scope() as db:
            r = db.get(IfrsRecon, payload["recon_id"])
            if r is None:
                return {}
            user = db.get(User, r.user_id)
            try:
                la, rr = _path(db, r.la_file_id), _path(db, r.rr06_file_id)
                if not la or not rr:
                    raise Refused("A file is no longer on the server — upload it again.")
                res = aicheck.with_retry(lambda: rr06_recon.run_recon(la, rr, r.period), "IFRS 16: LA vs RR06")
                out = tmp / (RECON_REPORT_NAME % r.period)
                rr06_recon.write_recon_workbook(res, str(out))
                r.report_file_id = _keep(db, user, out.name, out.read_bytes(), {"what": "recon"}).id
                a, b = res["RR06 not on LA"], res["LA not on RR06"]
                r.result = {"overview": _records(res["Overview"], 50), "summary": _records(res["Summary"], 20), "notes": list(res["Notes"]),
                            "rr06_not_la": _records(a, RECON_LIST_CAP), "la_not_rr06": _records(b, RECON_LIST_CAP),
                            "capped": len(a) > RECON_LIST_CAP or len(b) > RECON_LIST_CAP}
                r.status, r.note = "done", ""
            except Exception as e:  # noqa: BLE001
                _log_failed("recon", r.id, e)
                r.status, r.note = "failed", aicheck.note("IFRS 16: LA vs RR06", str(e) if isinstance(e, (Refused, ValueError)) else f"Reconciliation failed: {str(e)[:400]}")
            log(db, "ifrs16_recon_done", actor=user, target=f"ifrs_recon:{r.id}", detail=r.note[:400] or "reconciliation written", outcome="ok" if r.status == "done" else "failed")
            db.commit()
            jobs.progress(job_id, done=1, total=1, note=r.status)
            return {"status": r.status, "link": "#/ifrs16?tab=recon"}
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ---- Sequentra: check and download ---------------------------------------------------------------------------
def start_api(db: Session, user: User, kind: str, month: str = "", to_month: str = "") -> IfrsApiRun:
    if kind not in ("check", "download"):
        raise Refused("Check or download.")
    for v in (month, to_month):
        if v:
            try:
                month_date(v)
            except (ValueError, IndexError):
                raise Refused("Months look like 2026-08.") from None
    if kind == "download" and not month:
        raise Refused("Which month? (YYYY-MM)")
    if month and to_month and to_month < month:
        raise Refused("The 'to' month is before the 'from' month.")
    r = IfrsApiRun(user_id=user.id, kind=kind, month=month[:7], to_month=to_month[:7])
    db.add(r)
    db.flush()
    j = jobs.enqueue(db, tool=SLUG, kind="ifrs16_api", title="Sequentra · " + ("check every report" if kind == "check" else f"download {month}"), user=user,
                     ref=f"ifrs_api:{r.id}", total=1, payload={"run_id": r.id})
    r.job_id = j.id
    return r


@jobs.handler("ifrs16_api")
def api_job(job_id: int, payload: dict) -> dict:
    from . import api as SQ

    tmp = _scratch()
    try:
        with session_scope() as db:
            r = db.get(IfrsApiRun, payload["run_id"])
            if r is None:
                return {}
            user = db.get(User, r.user_id)
            try:
                cfg = SQ.config(db)
                if r.kind == "check":
                    lines = SQ.check(cfg, month_date(r.month) if r.month else None)
                    r.result = "\n".join(lines)[:8000]
                    r.ok = "no" if lines and lines[0].startswith("Not connected") else "yes"
                else:
                    a, b = month_date(r.month), (month_date(r.to_month) if r.to_month else None)
                    lines, data = SQ.download(cfg, a, b, tmp / f"Sequentra_{a:%Y-%m}")
                    name = f"Sequentra_{a:%Y-%m}{'_to_' + format(b, '%Y-%m') if b else ''}.zip"
                    r.download_file_id = _keep(db, user, name, data, {"what": "sequentra"}).id
                    r.result, r.ok = "\n".join(lines)[:8000], "yes"
                r.status = "done"
            except SQ.ApiProblem as e:
                r.status, r.note, r.ok = "failed", str(e)[:1000], "no"
            except Exception as e:  # noqa: BLE001
                _log_failed("api", r.id, e)
                r.status, r.note, r.ok = "failed", aicheck.note("IFRS 16: asking Sequentra", f"Sequentra: {str(e)[:400]}"), "no"
            log(db, "ifrs16_api", actor=user, target=f"ifrs_api:{r.id}", detail=f"{r.kind}: {r.ok or r.status}")
            db.commit()
            jobs.progress(job_id, done=1, total=1, note=r.status)
            return {"status": r.status, "link": "#/ifrs16?tab=sequentra"}
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ---- output -----------------------------------------------------------------------------------------------
def _f(db: Session, fid: Optional[int]) -> Optional[dict]:
    f = db.get(StoredFile, fid) if fid else None
    return {"id": f.id, "name": f.name, "size": f.size} if f else None


def upload_out(db: Session, u: IfrsUpload, owner: Optional[User] = None) -> dict:
    owner = owner or db.get(User, u.user_id)
    return {"id": u.id, "owner": owner.name if owner else "", "status": u.status, "note": u.note or "", "months": (u.scan or {}).get("months", []),
            "movement": (u.scan or {}).get("movement_label", ""), "files": {k: _f(db, getattr(u, f"{k}_file_id")) for k in FILE_KEYS},
            "at": u.created_at.isoformat() if u.created_at else None}


def month_out(db: Session, m: IfrsMonth, owner: Optional[User] = None) -> dict:
    owner = owner or db.get(User, m.user_id)
    return {"id": m.id, "upload_id": m.upload_id, "owner": owner.name if owner else "", "month": m.month, "source": m.source, "status": m.status,
            "imports": m.imports, "recon": m.recon or "", "note": m.note or "", "job_id": m.job_id,
            "workbooks": _f(db, m.workbooks_file_id), "report": _f(db, m.report_file_id), "papers": _f(db, m.papers_file_id),
            "at": m.created_at.isoformat() if m.created_at else None, "finished_at": m.finished_at.isoformat() if m.finished_at else None}


def ledger_out(db: Session, r: IfrsLedgerCheck) -> dict:
    return {"id": r.id, "status": r.status, "as_at": r.as_at, "months_back": r.months_back, "overview": r.overview or "", "note": r.note or "",
            "ledger": _f(db, r.ledger_file_id), "sequentra": _f(db, r.sequentra_file_id), "report": _f(db, r.report_file_id), "simple": _f(db, r.simple_file_id),
            "upload_id": r.upload_id, "entity": r.entity or "", "span": r.span or "life", "span_words": SPAN_WORDS.get(r.span or "life", ""),
            "at": r.created_at.isoformat() if r.created_at else None}


def master_out(db: Session, r: IfrsMaster) -> dict:
    return {"id": r.id, "upload_id": r.upload_id, "month": r.month, "as_at": r.as_at, "entity": r.entity, "status": r.status, "sheets": r.sheets,
            "note": r.note or "", "ledger": _f(db, r.ledger_file_id), "report": _f(db, r.report_file_id), "at": r.created_at.isoformat() if r.created_at else None}


def recon_out(db: Session, r: IfrsRecon, full: bool = False) -> dict:
    res = r.result or {}
    out = {"id": r.id, "status": r.status, "period": r.period, "note": r.note or "", "la": _f(db, r.la_file_id), "rr06": _f(db, r.rr06_file_id),
           "report": _f(db, r.report_file_id), "overview": res.get("overview", []), "summary": res.get("summary", []), "notes": res.get("notes", []),
           "at": r.created_at.isoformat() if r.created_at else None}
    if full:
        out.update(rr06_not_la=res.get("rr06_not_la", []), la_not_rr06=res.get("la_not_rr06", []), capped=res.get("capped", False))
    return out


def api_out(db: Session, r: IfrsApiRun) -> dict:
    return {"id": r.id, "kind": r.kind, "month": r.month, "to_month": r.to_month, "status": r.status, "ok": r.ok, "result": r.result or "", "note": r.note or "",
            "download": _f(db, r.download_file_id), "at": r.created_at.isoformat() if r.created_at else None}


def rate_out(r: IfrsRate) -> dict:
    return {"id": r.id, "month": r.month, **{b: getattr(r, b) for b in "abcdef"}, "note": r.note or ""}
