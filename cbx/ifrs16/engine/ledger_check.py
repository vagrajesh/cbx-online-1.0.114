"""
Ledger check for IFRS 16 Workflow.

Reads the Lease Accelerator "Ledger Export" (IFRS 16 set of books) and
builds the balance check: a movement report per lease and period, a
roll-forward proof per balance category, and a full lease list showing
which leases were included and why the rest were filtered out.

Fields used from the export (Ledger Entries tab), as specified by the
finance team:

    Ledger Date      -> reporting period (YYYYMM)
    Transactional
      Currency       -> currency
    Segment2         -> GL account code, which gives the balance category
    DR, CR           -> movement value, taken as DR - CR
    Comments         -> the lease reference ("Schedule A5344P01-M001")
    JEShortDesc      -> the movement type

The lease reference is the eight-character schedule number; the "-M001"
modification suffix and any trailing "[...]" annotation are stripped, so
A5344P01-M001 and A5344P01 roll up to the same lease.
"""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

import pandas as pd

# ---------------------------------------------------------------------
# Segment2 (GL account) -> balance category
# ---------------------------------------------------------------------
SEG_CATEGORY = {
    "BC000": "1. ROU Asset",
    "BC100": "2. Accumulated Depreciation",
    "BC101": "3. Impairment",
    "HA101": "4. Lease Liability",
}
# accounts that are profit and loss / clearing rather than the four
# IFRS 16 balance categories
SEG_PL = {
    "SE301": "Depreciation expense",
    "SF117": "Interest expense",
    "SB424": "Payments and rental expense",
    "SE651": "Gain (loss) on modification",
    "SE652": "Gain (loss) on return",
}

BALANCE_CATEGORIES = ["1. ROU Asset", "2. Accumulated Depreciation",
                      "3. Impairment", "4. Lease Liability"]

# ---------------------------------------------------------------------
# JEShortDesc -> movement column
# ---------------------------------------------------------------------
MOVEMENT_MAP = {
    "capital lease modification takedown": "Addition",
    "closeout asset and liability balances upon termination":
        "Termination",
    "capital lease depreciation expense recognition": "Depreciation",
    "interest accrual": "Interest",
    "payment": "Payment",
    "external adjusting entry -- do not export": "Adjustment",
    "cumulative catchup for penny rounding": "Rounding",
    "capital long term to short term reclass": "Reclass",
    "capital long term from short term reversal": "Reclass",
    "operating lease monthly expense recognition": "Operating",
    "operating lease periodic lease payment": "Operating",
    "return fees and grv": "Termination",
}
MOVEMENT_ORDER = ["Addition", "Adjustment", "Depreciation", "Interest",
                  "Payment", "Termination", "Rounding", "Reclass",
                  "Operating", "Other"]

SCHEDULE_RE = re.compile(r"Schedule\s+([A-Za-z0-9]+)", re.I)


def schedule_ref(comment):
    """The 8-character schedule number from a ledger comment, with the
    -M001 modification suffix and any [...] annotation removed."""
    text = str(comment or "")
    m = SCHEDULE_RE.search(text)
    if not m:
        return ""
    ref = m.group(1).strip().upper()
    ref = re.split(r"-M", ref, maxsplit=1)[0]      # drop -M001 / -M002
    return ref


def _num(v):
    """Blank DR/CR cells arrive as NaN; they must read as nil, otherwise
    every movement becomes NaN because one side of each entry is blank."""
    if v is None:
        return 0.0
    try:
        f = float(str(v).replace(",", "").strip() or 0)
    except (TypeError, ValueError):
        # 1.0.106 / 2.10.9: '1 234.56', 'R 1 234,56', '(500.00)' are
        # numbers stored as text, not nil (they used to read as 0.0)
        import lease_engine
        f = lease_engine._text_num(v)
        return 0.0 if f is None else f
    return 0.0 if f != f else f          # NaN -> 0.0


LEDGER_NEEDED = ["Ledger Date", "Segment2", "DR", "CR", "Comments",
                 "JEShortDesc"]


def _read_sheet(path, sheet, header=0):
    """Read one sheet, preferring calamine (LA exports carry style XML
    that openpyxl rejects)."""
    try:
        return pd.read_excel(path, sheet_name=sheet, engine="calamine",
                             header=header)
    except Exception:
        return pd.read_excel(path, sheet_name=sheet, header=header)


def _ledger_place(path, preferred="Ledger Entries"):
    """1.0.106 / 2.10.9: (sheet, header row, probe) - the tab and the row
    holding the ledger headings, found by content (a re-saved export can
    carry a title block above them, or rename the tab); (None, None,
    probe) when no tab has them."""
    try:
        probe = pd.read_excel(path, sheet_name=None, header=None, nrows=40,
                              engine="calamine")
    except Exception:
        try:
            probe = pd.read_excel(path, sheet_name=None, header=None,
                                  nrows=40)
        except Exception:
            return None, None, {}
    names = sorted(probe, key=lambda n: n != preferred)
    for name in names:
        df = probe[name]
        for i in range(min(40, len(df))):
            cells = {re.sub(r"\s+", " ", str(v)).strip()
                     for v in df.iloc[i].tolist()}
            if all(c in cells for c in LEDGER_NEEDED):
                return name, i, probe
    return None, None, probe


def _ledger_sheet_name(path, preferred="Ledger Entries"):
    """The sheet holding the ledger entries. The named tab is used when
    it carries the required columns; otherwise every sheet is probed, so
    a re-saved or renamed export still loads."""
    try:
        probe = pd.read_excel(path, sheet_name=None, nrows=3,
                              engine="calamine")
    except Exception:
        try:
            probe = pd.read_excel(path, sheet_name=None, nrows=3)
        except Exception:
            return preferred
    def fits(df):
        cols = [str(c).strip() for c in (df.columns if df is not None
                                         else [])]
        return all(c in cols for c in LEDGER_NEEDED)
    if preferred in probe and fits(probe[preferred]):
        return preferred
    for name, df in probe.items():
        if fits(df):
            return name
    return preferred


def load_ledger_export(path, sheet="Ledger Entries"):
    """Read a Ledger Export into the normalised frame the check uses."""
    path = Path(path)
    if not path.exists():
        raise ValueError("Ledger export not found: %s" % path.name)
    found, hr, probe = _ledger_place(path, sheet)
    if found is None:
        try:
            import lease_engine
            hint = lease_engine.what_file(probe, "ledger")
        except Exception:
            hint = ""
        raise ValueError(
            "'%s' does not look like the LA ledger export - no tab has the "
            "columns %s.%s" % (path.name, ", ".join(LEDGER_NEEDED), hint))
    raw = _read_sheet(path, found, header=hr)
    raw.columns = [re.sub(r"\s+", " ", str(c)).strip() for c in raw.columns]

    needed = list(LEDGER_NEEDED)
    missing = [c for c in needed if c not in raw.columns]
    if missing:
        raise ValueError(
            "'%s' does not look like a Ledger Export - missing column(s): "
            "%s. Expected the 'Ledger Entries' tab of the Lease "
            "Accelerator ledger export." % (path.name, ", ".join(missing)))

    cur_col = next((c for c in ("Transactional Currency",
                                "Functional Currency", "Currency")
                    if c in raw.columns), None)

    out = pd.DataFrame({
        "Lease Reference": raw["Comments"].map(schedule_ref),
        "Ledger Date": pd.to_datetime(raw["Ledger Date"],
                                      errors="coerce", dayfirst=False),
        "Currency": raw[cur_col].astype(str).str.strip()
                    if cur_col else "",
        "Segment2": raw["Segment2"].astype(str).str.strip(),
        "Journal Type": raw["JEShortDesc"].astype(str).str.strip(),
        "DR": raw["DR"].map(_num),
        "CR": raw["CR"].map(_num),
        "Prop Ref": (raw["Segment3"].astype(str).str.strip()
                     if "Segment3" in raw.columns else ""),
        "Cost Centre": (raw["Segment4"].astype(str).str.strip()
                        if "Segment4" in raw.columns else ""),
    })
    out["Movement"] = out["DR"] - out["CR"]
    out["Period"] = out["Ledger Date"].dt.strftime("%Y%m")
    out["Category"] = out["Segment2"].map(SEG_CATEGORY).fillna(
        out["Segment2"].map(SEG_PL)).fillna("Other")
    out["Movement Type"] = (out["Journal Type"].str.lower()
                            .map(MOVEMENT_MAP).fillna("Other"))
    out["Source File"] = path.name
    out = out[out["Lease Reference"] != ""]
    return out.reset_index(drop=True)


# ---------------------------------------------------------------------
# Combining several ledger exports into one report
# ---------------------------------------------------------------------
def ledger_files_in(folder):
    """Every Excel file in a folder, in name order (Excel lock files and
    files already produced by this app are ignored)."""
    folder = Path(folder)
    if not folder.is_dir():
        return []
    skip = ("consolidated_report", "master_report", "_extract_",
            "secure_report")
    out = []
    for p in sorted(folder.iterdir()):
        if not p.is_file() or p.suffix.lower() not in (".xlsx", ".xlsm",
                                                      ".xls"):
            continue
        if p.name.startswith("~$"):
            continue
        if any(s in p.name.lower() for s in skip):
            continue
        out.append(p)
    return out


def combine_ledger_exports(paths):
    """Read several ledger exports and stack them into one frame.

    Returns (ledger, used, skipped, dropped) where
      used    = [{File, Rows, Leases, First Period, Last Period}, ...]
      skipped = [{File, Reason}, ...]      (not a ledger export, etc.)
      dropped = rows that repeated a line already read from an earlier
                file (overlapping exports), so balances are not doubled.

    De-duplication is deliberately ACROSS files only: a line repeated
    inside one export is a genuine second posting and is kept."""
    frames, used, skipped = [], [], []
    for p in paths:
        p = Path(p)
        try:
            df = load_ledger_export(p)
        except Exception as exc:
            skipped.append({"File": p.name, "Reason": str(exc)[:300]})
            continue
        if not len(df):
            skipped.append({"File": p.name,
                            "Reason": "no ledger lines with a schedule "
                                      "reference"})
            continue
        periods = sorted(df["Period"].dropna().unique())
        used.append({"File": p.name, "Rows": len(df),
                     "Leases": int(df["Lease Reference"].nunique()),
                     "First Period": periods[0] if periods else "",
                     "Last Period": periods[-1] if periods else ""})
        frames.append(df)
    if not frames:
        raise ValueError(
            "No ledger export could be read from the folder. %s"
            % ("Files tried: " + "; ".join(
                "%s (%s)" % (s["File"], s["Reason"]) for s in skipped)
               if skipped else "The folder holds no Excel files."))
    led = pd.concat(frames, ignore_index=True)
    dropped = 0
    if len(frames) > 1:
        key = ["Lease Reference", "Ledger Date", "Segment2",
               "Journal Type", "DR", "CR"]
        # occurrence number within each file, so repeats inside one
        # export survive while cross-file repeats are removed
        led["_occ"] = led.groupby(["Source File"] + key).cumcount()
        before = len(led)
        led = led.drop_duplicates(subset=key + ["_occ"], keep="first")
        dropped = before - len(led)
        led = led.drop(columns=["_occ"])
    return led.reset_index(drop=True), used, skipped, dropped


def load_ledger_folder(folder):
    """Combine every ledger export in a folder. Returns the same tuple
    as combine_ledger_exports."""
    files = ledger_files_in(folder)
    if not files:
        raise ValueError("No Excel files found in %s" % folder)
    return combine_ledger_exports(files)


# ---------------------------------------------------------------------
# lease selection
# ---------------------------------------------------------------------
def _period_floor(as_at, months_back):
    """The earliest period to include: months_back months before as_at."""
    y, m = as_at.year, as_at.month
    total = y * 12 + (m - 1) - int(months_back)
    return "%04d%02d" % (total // 12, total % 12 + 1)


def _period_of(as_at):
    """The reporting period itself - the ledger carries future-dated
    entries for the whole lease life, which must not be counted as
    activity in the reporting window."""
    return as_at.strftime("%Y%m")


def select_leases(ledger, not_applicable=None, as_at=None, months_back=12,
                  contracts=None, full_life=False):
    """Decide which leases the check covers.

    Returns (kept_refs, listing) where listing is a row per lease with
    the decision and the reason - so every exclusion is visible.

    not_applicable: {reference: reason} from the workflow's IFRS 16 tests
    contracts:      optional frame from the movement file's Contracts tab
    """
    as_at = as_at or date.today()
    if full_life:
        have = ledger["Period"].dropna()
        floor = have.min() if len(have) else _period_of(as_at)
        ceiling = have.max() if len(have) else _period_of(as_at)
    else:
        floor = _period_floor(as_at, months_back)
        ceiling = _period_of(as_at)
    na = {str(k).strip().upper(): v for k, v in (not_applicable or {}).items()}

    contract_refs = set()
    if contracts is not None and len(contracts):
        col = next((c for c in contracts.columns
                    if "lease reference" in str(c).strip().lower()
                    or str(c).strip().lower() in ("schedule number",
                                                  "schedule")), None)
        if col:
            contract_refs = {str(v).strip().upper()
                             for v in contracts[col].dropna()}

    rows, keep = [], []
    for ref, grp in ledger.groupby("Lease Reference", sort=True):
        in_window = grp[(grp["Period"] >= floor)
                        & (grp["Period"] <= ceiling)]
        periods = sorted(grp["Period"].dropna().unique())
        cats = set(grp["Category"])
        reason, decision = "", "Included"

        if ref in na:
            decision, reason = "Excluded", "Not IFRS 16 applicable - %s" \
                % na[ref]
        elif not len(in_window):
            decision = "Excluded"
            reason = ("No ledger activity in the %d months to %s (last "
                      "period %s)" % (months_back, ceiling,
                                      periods[-1] if periods else "none"))
        elif not (cats & set(BALANCE_CATEGORIES)):
            decision, reason = "Excluded", \
                "No IFRS 16 balance-sheet entries (profit and loss only)"
        elif contract_refs and ref not in contract_refs:
            decision, reason = "Excluded", \
                "Not present on the Contracts tab of the movement file"
        else:
            reason = "IFRS 16 balances present in the reporting window"
            keep.append(ref)

        rows.append({
            "Lease Reference": ref,
            "Decision": decision,
            "Reason": reason,
            "Currency": next((c for c in grp["Currency"] if c), ""),
            "Ledger Lines": len(grp),
            "Lines In Window": len(in_window),
            "First Period": periods[0] if periods else "",
            "Last Period": periods[-1] if periods else "",
            "Categories": ", ".join(sorted(
                c for c in cats if c in BALANCE_CATEGORIES)),
        })
    listing = pd.DataFrame(rows, columns=[
        "Lease Reference", "Decision", "Reason", "Currency",
        "Ledger Lines", "Lines In Window", "First Period", "Last Period",
        "Categories"])
    return keep, listing


# ---------------------------------------------------------------------
# the report and the check
# ---------------------------------------------------------------------
def movement_report(ledger, refs=None, as_at=None, months_back=12,
                    full_life=False):
    """Per lease, period and category: the movement columns plus the
    opening and closing balance. Opening is the cumulative position
    before the period, so opening + movements = closing by construction
    and any break is a genuine data problem."""
    as_at = as_at or date.today()
    if full_life:
        have = ledger["Period"].dropna()
        floor = have.min() if len(have) else _period_of(as_at)
        ceiling = have.max() if len(have) else _period_of(as_at)
    else:
        floor = _period_floor(as_at, months_back)
        ceiling = _period_of(as_at)
    led = ledger[ledger["Category"].isin(BALANCE_CATEGORIES)].copy()
    if refs is not None:
        led = led[led["Lease Reference"].isin(set(refs))]
    if led.empty:
        return pd.DataFrame(columns=(
            ["Lease Reference", "Period", "Category", "Currency",
             "Opening Balance"] + MOVEMENT_ORDER + ["Closing Balance"]))

    # movement columns per lease / period / category
    piv = (led.pivot_table(index=["Lease Reference", "Period", "Category"],
                           columns="Movement Type", values="Movement",
                           aggfunc="sum", fill_value=0.0)
           .reset_index())
    for c in MOVEMENT_ORDER:
        if c not in piv.columns:
            piv[c] = 0.0

    # cumulative closing balance across ALL periods, then opening
    led = led[led["Period"] <= ceiling]
    tot = (led.groupby(["Lease Reference", "Category", "Period"])
           ["Movement"].sum().reset_index()
           .sort_values(["Lease Reference", "Category", "Period"]))
    tot["Closing Balance"] = (tot.groupby(["Lease Reference", "Category"])
                              ["Movement"].cumsum())
    tot["Opening Balance"] = tot["Closing Balance"] - tot["Movement"]

    rep = piv.merge(tot[["Lease Reference", "Category", "Period",
                         "Opening Balance", "Closing Balance"]],
                    on=["Lease Reference", "Category", "Period"],
                    how="left")
    cur = (led.groupby("Lease Reference")["Currency"]
           .agg(lambda s: next((x for x in s if x), "")))
    rep["Currency"] = rep["Lease Reference"].map(cur)
    rep = rep[(rep["Period"] >= floor) & (rep["Period"] <= ceiling)]
    cols = (["Lease Reference", "Period", "Category", "Currency",
             "Opening Balance"] + MOVEMENT_ORDER + ["Closing Balance"])
    rep = rep[cols].sort_values(["Lease Reference", "Category", "Period"])
    return rep.reset_index(drop=True)


def balance_check(report):
    """Roll-forward proof per period and category:
       opening + movements = closing, and the difference must be nil."""
    if report is None or not len(report):
        return pd.DataFrame(columns=[
            "Period", "Category", "Opening Balance", "Movements",
            "Closing Balance", "Expected Closing", "Difference",
            "Status"])
    move_cols = [c for c in MOVEMENT_ORDER if c in report.columns]
    g = (report.groupby(["Period", "Category"])
         .agg({**{"Opening Balance": "sum", "Closing Balance": "sum"},
               **{c: "sum" for c in move_cols}})
         .reset_index())
    g["Movements"] = g[move_cols].sum(axis=1)
    g["Expected Closing"] = g["Opening Balance"] + g["Movements"]
    g["Difference"] = g["Closing Balance"] - g["Expected Closing"]
    g["Status"] = g["Difference"].abs().map(
        lambda d: "OK" if d < 0.01 else "CHECK")
    cols = ["Period", "Category", "Opening Balance", "Movements",
            "Closing Balance", "Expected Closing", "Difference", "Status"]
    return g[cols].sort_values(["Period", "Category"]).reset_index(
        drop=True)


def ledger_integrity(ledger, as_at=None, months_back=12):
    """Double-entry proof straight off the export: every period's debits
    must equal its credits, so the net movement is nil."""
    as_at = as_at or date.today()
    floor = _period_floor(as_at, months_back)
    led = ledger[(ledger["Period"] >= floor)
                 & (ledger["Period"] <= _period_of(as_at))]
    if led.empty:
        return pd.DataFrame(columns=["Period", "Debits", "Credits",
                                     "Net", "Lines", "Status"])
    g = (led.groupby("Period")
         .agg(Debits=("DR", "sum"), Credits=("CR", "sum"),
              Lines=("Movement", "size"))
         .reset_index())
    g["Net"] = g["Debits"] - g["Credits"]
    g["Status"] = g["Net"].abs().map(
        lambda d: "OK" if d < 0.01 else "CHECK")
    return g[["Period", "Debits", "Credits", "Net", "Lines",
              "Status"]].sort_values("Period").reset_index(drop=True)


def category_summary(ledger, refs=None, as_at=None):
    """Portfolio closing balance by category as at the reporting period.

    This is the cumulative position of every lease up to and including
    the period - not only the leases that moved in that month - so it is
    the figure that ties to the trial balance and to the consolidated
    report's closing balances."""
    as_at = as_at or date.today()
    period = _period_of(as_at)
    led = ledger[(ledger["Period"] <= period)
                 & ledger["Category"].isin(BALANCE_CATEGORIES)]
    if refs is not None:
        led = led[led["Lease Reference"].isin(set(refs))]
    if led.empty:
        return pd.DataFrame(columns=["Category", "Period", "Leases",
                                     "Closing Balance"])
    g = (led.groupby("Category")
         .agg(Leases=("Lease Reference", "nunique"),
              **{"Closing Balance": ("Movement", "sum")})
         .reset_index())
    g["Period"] = period
    g["Closing Balance"] = g["Closing Balance"].round(2)
    return g[["Category", "Period", "Leases",
              "Closing Balance"]].sort_values("Category").reset_index(
        drop=True)



CONSOLIDATED_COLS = [
    "Lease Reference", "Lease Name", "Prop Ref", "Period", "Currency",
    "Current lease Start Date", "Current lease End Date", "Entity",
    "Country", "Discount Rate", "Cost Centre",
    "1.ASSET_OPENING BALANCE", "1.ASSET_ADDITION", "1.ASSET_ADJUSTMENT",
    "1.ASSET_TERMINATION", "1.ASSET_CLOSING BALANCE",
    "2.ACCUMULATED DEPRECIATION_OPENING BALANCE",
    "2.ACCUMULATED DEPRECIATION_ADDITION",
    "2.ACCUMULATED DEPRECIATION_ADJUSTMENT",
    "2.ACCUMULATED DEPRECIATION_DEPRECIATION",
    "2.ACCUMULATED DEPRECIATION_TERMINATION",
    "2.ACCUMULATED DEPRECIATION_CLOSING BALANCE",
    "3.IMPAIRMENT_OPENING BALANCE", "3.IMPAIRMENT_ADDITION",
    "3.IMPAIRMENT_ADJUSTMENT", "3.IMPAIRMENT_TERMINATION",
    "3.IMPAIRMENT_CLOSING BALANCE",
    "4.LEASE LIABILITY_OPENING BALANCE", "4.LEASE LIABILITY_ADDITION",
    "4.LEASE LIABILITY_INTEREST", "4.LEASE LIABILITY_PAYMENT",
    "4.LEASE LIABILITY_TERMINATION", "4.LEASE LIABILITY_CLOSING BALANCE",
    # finance, Oct 2026: the net asset rolled forward like the other blocks,
    # after the lease liability closing balance - each column the sum of the
    # ROU asset, accumulated depreciation and impairment columns of the same
    # name (depreciation: the accumulated depreciation charge)
    "5.NET ASSET_OPENING BALANCE", "5.NET ASSET_ADDITION",
    "5.NET ASSET_ADJUSTMENT", "5.NET ASSET_TERMINATION",
    "5.NET ASSET_DEPRECIATION", "5.NET ASSET_CLOSING BALANCE",
    "DEPRECIATION", "INTEREST EXPENSE", "PAYMENT",
]
NET_ASSET_PARTS = {
    "OPENING BALANCE": ("1.ASSET_OPENING BALANCE",
                        "2.ACCUMULATED DEPRECIATION_OPENING BALANCE",
                        "3.IMPAIRMENT_OPENING BALANCE"),
    "ADDITION": ("1.ASSET_ADDITION", "2.ACCUMULATED DEPRECIATION_ADDITION",
                 "3.IMPAIRMENT_ADDITION"),
    "ADJUSTMENT": ("1.ASSET_ADJUSTMENT",
                   "2.ACCUMULATED DEPRECIATION_ADJUSTMENT",
                   "3.IMPAIRMENT_ADJUSTMENT"),
    "TERMINATION": ("1.ASSET_TERMINATION",
                    "2.ACCUMULATED DEPRECIATION_TERMINATION",
                    "3.IMPAIRMENT_TERMINATION"),
    "DEPRECIATION": ("2.ACCUMULATED DEPRECIATION_DEPRECIATION",),
    "CLOSING BALANCE": ("1.ASSET_CLOSING BALANCE",
                        "2.ACCUMULATED DEPRECIATION_CLOSING BALANCE",
                        "3.IMPAIRMENT_CLOSING BALANCE"),
}
NET_ASSET_MOVES = ("ADDITION", "ADJUSTMENT", "TERMINATION", "DEPRECIATION")

# category -> the movement columns it carries in the consolidated layout
_CONSOLIDATED_BLOCKS = {
    "1. ROU Asset": ("1.ASSET", ["ADDITION", "ADJUSTMENT",
                                 "TERMINATION"]),
    "2. Accumulated Depreciation": ("2.ACCUMULATED DEPRECIATION",
                                    ["ADDITION", "ADJUSTMENT",
                                     "DEPRECIATION", "TERMINATION"]),
    "3. Impairment": ("3.IMPAIRMENT", ["ADDITION", "ADJUSTMENT",
                                       "TERMINATION"]),
    "4. Lease Liability": ("4.LEASE LIABILITY", ["ADDITION", "INTEREST",
                                                 "PAYMENT",
                                                 "TERMINATION"]),
}
# our movement names -> the consolidated column suffix
_MOVE_TO_COL = {"Addition": "ADDITION", "Adjustment": "ADJUSTMENT",
                "Depreciation": "DEPRECIATION", "Interest": "INTEREST",
                "Payment": "PAYMENT", "Termination": "TERMINATION"}


def _month_range(first, last):
    """Inclusive list of YYYYMM period strings."""
    fy, fm = int(first[:4]), int(first[4:6])
    ly, lm = int(last[:4]), int(last[4:6])
    out, y, m = [], fy, fm
    while (y, m) <= (ly, lm):
        out.append("%04d%02d" % (y, m))
        m += 1
        if m > 12:
            y, m = y + 1, 1
    return out


def consolidated_report(ledger, refs=None, as_at=None, months_back=12,
                        lease_names=None, lease_terms=None,
                        full_life=False, entities=None, period_from=None,
                        discount_rates=None):
    """The consolidated movement report in the team's column layout: one
    row per lease and period, with the four balance blocks rolled
    opening -> movements -> closing, plus the profit and loss columns.

    Balances carry forward: a lease with no movement in a month still
    reports its opening and closing position for that month, which is
    what makes the roll-forward tie."""
    as_at = as_at or date.today()
    if full_life:
        # every period the ledger carries, so each lease is shown from
        # its first entry through to the end of its schedule
        have = ledger["Period"].dropna()
        floor = have.min() if len(have) else _period_of(as_at)
        ceiling = have.max() if len(have) else _period_of(as_at)
    else:
        floor = _period_floor(as_at, months_back)
        ceiling = _period_of(as_at)
    led = ledger[ledger["Period"].notna() & (ledger["Period"] <= ceiling)]
    if refs is not None:
        led = led[led["Lease Reference"].isin(set(refs))]
    if led.empty:
        return pd.DataFrame(columns=CONSOLIDATED_COLS)
    led = led.copy()

    bal = led[led["Category"].isin(BALANCE_CATEGORIES)]
    if bal.empty:
        return pd.DataFrame(columns=CONSOLIDATED_COLS)

    all_periods = _month_range(bal["Period"].min(), ceiling)
    window = [p for p in all_periods if p >= floor]

    # cumulative closing balance for every lease/category at every period
    piv = bal.pivot_table(index=["Lease Reference", "Category"],
                          columns="Period", values="Movement",
                          aggfunc="sum", fill_value=0.0)
    piv = piv.reindex(columns=all_periods, fill_value=0.0)
    closing = piv.cumsum(axis=1)
    opening = closing.shift(axis=1).fillna(0.0)

    clo = (closing[window].stack().rename("closing").reset_index()
           .rename(columns={"level_2": "Period"}))
    ope = (opening[window].stack().rename("opening").reset_index()
           .rename(columns={"level_2": "Period"}))
    bals = clo.merge(ope, on=["Lease Reference", "Category", "Period"])

    # movements actually posted in each period
    mv = (bal[bal["Period"].isin(window)]
          .groupby(["Lease Reference", "Period", "Category",
                    "Movement Type"])["Movement"].sum().reset_index())
    mv_idx = {(r[0], r[1], r[2], r[3]): r[4]
              for r in mv.itertuples(index=False)}

    # drop rows with nothing to show: no balance either side, no movement
    moved = {(r[0], r[1], r[2]) for r in mv.itertuples(index=False)}
    has_move = pd.Series(
        [(r, p, c) in moved for r, p, c in
         zip(bals["Lease Reference"], bals["Period"], bals["Category"])],
        index=bals.index)
    bals = bals[(bals["closing"].abs() > 0.004)
                | (bals["opening"].abs() > 0.004)
                | has_move]
    # bound each lease by its own posted periods: the schedule runs from
    # the lease's first entry to its last, not to the last period that
    # happens to exist in the file for other leases
    span = (bal.groupby("Lease Reference")["Period"]
            .agg(["min", "max"]).to_dict("index"))
    asat_period = _period_of(as_at)
    lo = bals["Lease Reference"].map(lambda r: span.get(r, {}).get("min"))
    # a lease still carrying a balance is shown up to the reporting month
    # even if nothing was posted in between; a lease whose schedule runs
    # into future periods is shown to the end of that schedule
    hi = bals["Lease Reference"].map(
        lambda r: max(span.get(r, {}).get("max") or asat_period,
                      asat_period))
    bals = bals[(bals["Period"] >= lo) & (bals["Period"] <= hi)]

    # profit and loss per lease / period
    pl_map = {"Depreciation expense": "DEPRECIATION",
              "Interest expense": "INTEREST EXPENSE",
              "Payments and rental expense": "PAYMENT"}
    pl = (led[led["Category"].isin(pl_map) & led["Period"].isin(window)]
          .groupby(["Lease Reference", "Period", "Category"])
          ["Movement"].sum().reset_index())
    pl_idx = {(r[0], r[1], pl_map[r[2]]): r[3]
              for r in pl.itertuples(index=False)}

    meta = (led.groupby("Lease Reference")
            .agg(Currency=("Currency", lambda s: next((x for x in s if x),
                                                      "")),
                 **{"Prop Ref": ("Prop Ref",
                                 lambda s: next((x for x in s if x), "")),
                    "Cost Centre": ("Cost Centre",
                                    lambda s: next((x for x in s if x),
                                                   ""))}))
    names = lease_names or {}
    terms = lease_terms or {}
    ents = entities or {}
    rates = discount_rates or {}

    # index the balances for quick lookup, then build one row per
    # lease/period across the four blocks
    bal_idx = {}
    for r in bals.itertuples(index=False):
        bal_idx[(r[0], r[1], r[2])] = (r[4], r[3])   # (opening, closing)
    keys = sorted({(r, p) for r, _c, p in
                   zip(bals["Lease Reference"], bals["Category"],
                       bals["Period"])})

    rows = []
    for ref, period in keys:
        m = meta.loc[ref] if ref in meta.index else None
        t = terms.get(ref, {})
        row = {c: "" for c in CONSOLIDATED_COLS}
        row.update({
            "Lease Reference": ref,
            "Lease Name": names.get(ref, ""),
            "Prop Ref": m["Prop Ref"] if m is not None else "",
            "Period": period,
            # currency per the Sequentra data, falling back to the ledger
            "Currency": (t.get("currency")
                         or (m["Currency"] if m is not None else "")),
            "Current lease Start Date": t.get("start", ""),
            "Current lease End Date": t.get("end", ""),
            "Entity": ents.get(ref, t.get("entity", "")),
            "Country": t.get("country", ""),
            "Discount Rate": rates.get(ref, ""),
            "Cost Centre": m["Cost Centre"] if m is not None else "",
        })
        for cat, (prefix, cols) in _CONSOLIDATED_BLOCKS.items():
            got = bal_idx.get((ref, cat, period))
            if got is None:
                continue
            op, cl = got
            row["%s_OPENING BALANCE" % prefix] = round(op, 6)
            row["%s_CLOSING BALANCE" % prefix] = round(cl, 6)
            for suffix in cols:
                total, seen = 0.0, False
                for our, col in _MOVE_TO_COL.items():
                    if col != suffix:
                        continue
                    v = mv_idx.get((ref, period, cat, our))
                    if v is not None:
                        total += v
                        seen = True
                if suffix == "ADJUSTMENT":
                    v = mv_idx.get((ref, period, cat, "Rounding"))
                    if v is not None:
                        total += v
                        seen = True
                if seen:
                    row["%s_%s" % (prefix, suffix)] = round(total, 6)
        # Net asset: ROU asset less accumulated depreciation and
        # impairment (the three are already signed, so a plain sum), for
        # the opening, every movement and the closing.
        for suffix, parts in NET_ASSET_PARTS.items():
            vals = [row.get(c) for c in parts]
            vals = [v for v in vals if isinstance(v, (int, float))]
            if vals:
                row["5.NET ASSET_%s" % suffix] = round(sum(vals), 6)
        for col in ("DEPRECIATION", "INTEREST EXPENSE", "PAYMENT"):
            v = pl_idx.get((ref, period, col))
            if v is not None:
                row[col] = round(v, 6)
        rows.append(row)
    out = pd.DataFrame(rows, columns=CONSOLIDATED_COLS)
    if period_from:
        out = out[out["Period"] >= str(period_from)]
    return out.sort_values(["Lease Reference", "Period"]).reset_index(
        drop=True)



# --------------------------------------------------------------------------
# the Consolidated Report workbook, in the layout the team produces today
# --------------------------------------------------------------------------
CR_PIVOT_COLS = ["Lease Reference", "Lease Name", "Prop Ref",
                 "Cost Centre", "Period"] + [
    "Sum of " + c for c in CONSOLIDATED_COLS
    if c.startswith(("1.", "2.", "3.", "4.", "5."))]

# the two banner rows above the Data header, as in the team's workbook
_DATA_GROUP = {c: (c.split("_")[0] if "_" in c and c[0].isdigit() else "")
               for c in CONSOLIDATED_COLS}
_DATA_SUB = {c: (c.split("_", 1)[1] if "_" in c and c[0].isdigit() else "")
             for c in CONSOLIDATED_COLS}
_DATA_SOURCE = {
    "Lease Reference": "LeaseData[Lease Reference]",
    "Lease Name": "Asset Report[Address]",
    "Prop Ref": "LeaseData[Property Reference]",
    "Period": "LeaseData[Period]",
    "Currency": "Sequentra lease data[Currency]",
    "Current lease Start Date": "Sequentra lease data[Start]",
    "Current lease End Date": "Sequentra lease data[End]",
    "Entity": "Sequentra lease data[Entity]",
    "Country": "Sequentra lease data[Country]",
    "Discount Rate": "LA transaction report[Discount Rate]",
}


def _recon_rows(ledger, refs, as_at, entity=""):
    """The roll-forward proof laid out as the team's Recon tab: prior
    month closing, this month's movements by type, computed closing and
    the difference against the reported closing."""
    period = _period_of(as_at)
    prev = _period_floor(as_at, 1)
    cats = [("LEASE LIABILITY", "4. Lease Liability"),
            ("ROU ASSET", "1. ROU Asset"),
            ("ACC DEP", "2. Accumulated Depreciation"),
            ("Impairment", "3. Impairment")]
    led = ledger[ledger["Category"].isin(BALANCE_CATEGORIES)]
    if refs is not None:
        led = led[led["Lease Reference"].isin(set(refs))]

    def total(cat, upto):
        return float(led[(led["Category"] == cat)
                         & (led["Period"] <= upto)]["Movement"].sum())

    def moved(cat, mtype):
        return float(led[(led["Category"] == cat)
                         & (led["Period"] == period)
                         & (led["Movement Type"] == mtype)]["Movement"]
                     .sum())

    rows = [[entity, period, "", "", "", ""], []]
    rows.append(["", "", cats[0][0], cats[1][0], cats[2][0], cats[3][0]])
    rows.append(["Balance per previous report", "Prior Month C/B"]
                + [total(c, prev) for _l, c in cats])
    rows.append(["", "Current Month O/B"] + [total(c, prev)
                                             for _l, c in cats])
    rows.append([])
    for label, mtype in (("Depreciation", "Depreciation"),
                         ("Additions", "Addition"),
                         ("Disposals", "Termination"),
                         ("Interest", "Interest"),
                         ("Payment", "Payment"),
                         ("Adjustments", "Adjustment"),
                         ("Rounding", "Rounding")):
        vals = [moved(c, mtype) for _l, c in cats]
        if any(abs(v) > 0.005 for v in vals):
            rows.append(["Grand Total", label] + vals)
    computed = [total(c, prev) + float(
        led[(led["Category"] == c) & (led["Period"] == period)]
        ["Movement"].sum()) for _l, c in cats]
    rows.append([])
    rows.append(["", "Computed closing"] + computed)
    rows.append(["Grand Total", "Current Month C/B"]
                + [total(c, period) for _l, c in cats])
    rows.append(["", "Difference"]
                + [round(total(c, period) - comp, 6)
                   for (_l, c), comp in zip(cats, computed)])
    return rows


def workbook_layout(results, as_at=None):
    """What the Excel file will contain: each tab, its columns and how
    many rows - so the layout can be shown before the file is opened."""
    period = _period_of(as_at) if as_at is not None else ""
    con = results.get("Report Consolidated")
    n_all = 0 if con is None else len(con)
    n_period = 0
    if con is not None and period:
        n_period = int((con["Period"] == period).sum())
    tabs = [
        {"tab": "Consolidated Report",
         "what": "Pivot view - one row per lease and period",
         "rows": n_all, "columns": CR_PIVOT_COLS,
         "header_row": 3},
        {"tab": "Data",
         "what": "The full table, with the two banner rows above the "
                 "header", "rows": n_all,
         "columns": list(CONSOLIDATED_COLS), "header_row": 3},
    ]
    if period:
        tabs.append({"tab": period,
                     "what": "Only %s, with a total row at the bottom"
                             % period,
                     "rows": n_period + 2, "columns": CR_PIVOT_COLS,
                     "header_row": 3})
    tabs.append({"tab": "Recon",
                 "what": "Roll-forward proof: opening + movements = "
                         "closing, per category",
                 "rows": 16,
                 "columns": ["", "", "LEASE LIABILITY", "ROU ASSET",
                             "ACC DEP", "Impairment"],
                 "header_row": 3})
    for key, tab, what in (
            ("Report Category Summary", "Category Summary",
             "Portfolio closing balance by category"),
            ("Report Ledger Integrity", "Ledger Integrity",
             "Debits equal credits in every period"),
            ("Report Lease Selection", "Lease Selection",
             "Every lease, included or excluded, with the reason"),
            ("Report Source Files", "Source Files",
             "Each ledger export that went into this report")):
        df = results.get(key)
        if df is None:
            continue
        tabs.append({"tab": tab, "what": what, "rows": len(df),
                     "columns": [str(c) for c in df.columns],
                     "header_row": 1})
    return tabs


def net_asset_check(consolidated, tolerance=0.01):
    """Finance's check (Oct 2026), kept off the report: for every lease and
    period, net asset opening + additions + adjustment + termination +
    depreciation must equal the net asset closing balance."""
    cols = ["Lease Reference", "Period", "Opening", "Movements",
            "Expected Closing", "Closing", "Difference", "Status"]
    if consolidated is None or not len(consolidated):
        return pd.DataFrame(columns=cols)

    def n(r, c):
        v = r.get(c)
        return float(v) if isinstance(v, (int, float)) and v == v else 0.0
    rows = []
    for _i, r in consolidated.iterrows():
        op = n(r, "5.NET ASSET_OPENING BALANCE")
        mv = sum(n(r, "5.NET ASSET_%s" % m) for m in NET_ASSET_MOVES)
        cl = n(r, "5.NET ASSET_CLOSING BALANCE")
        diff = round(cl - (op + mv), 2)
        rows.append({"Lease Reference": r["Lease Reference"],
                     "Period": r["Period"], "Opening": round(op, 2),
                     "Movements": round(mv, 2),
                     "Expected Closing": round(op + mv, 2),
                     "Closing": round(cl, 2), "Difference": diff,
                     "Status": "OK" if abs(diff) <= tolerance else "CHECK"})
    return pd.DataFrame(rows, columns=cols)


def discount_rates_from_sequentra(path):
    """{lease reference: discount rate (%)} from the Sequentra validation
    report (finance, Oct 2026: the report's Discount Rate comes from it).
    'Discount Rate' first, else 'New Discount Rate'; a fraction (0.105) is
    read as a percentage (10.5)."""
    out = {}
    if not path or not Path(path).exists():
        return out
    try:
        import lease_engine
        seq = lease_engine.load_sequentra(path)
    except Exception:
        return out
    cols = [c for c in ("Discount Rate", "New Discount Rate", "IBR")
            if c in seq.columns]
    if not cols or "Record ID" not in seq.columns:
        return out
    for _i, r in seq.iterrows():
        ref = str(r.get("Record ID", "")).strip().upper()
        if not ref or ref == "NAN":
            continue
        for c in cols:
            try:
                v = float(str(r.get(c)).replace("%", "").strip())
            except (TypeError, ValueError):
                continue
            if v != v or v <= 0:
                continue
            if v < 1:
                v = v * 100
            out.setdefault(ref, round(v, 4))
            break
    return out


def write_consolidated_workbook(results, out_path, ledger=None,
                                refs=None, as_at=None, entity=""):
    """Write the Consolidated Report workbook: the pivot view, the Data
    table with its banner rows, and the Recon proof - the same shape the
    team produces by hand today."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    con = results.get("Report Consolidated")
    if con is None or not len(con):
        raise ValueError("No consolidated report to write - run the "
                         "ledger check first.")
    out_path = Path(out_path)
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="003F2D")
    band_font = Font(bold=True, color="003F2D")
    band_fill = PatternFill("solid", fgColor="DBE7E1")

    # ---- 1. Consolidated Report (the pivot view) --------------------
    ws = wb.active
    ws.title = "Consolidated Report"
    ws.append([])
    ws.append([])
    ws.append(CR_PIVOT_COLS)
    for cell in ws[3]:
        if cell.value:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
    money = [c for c in CONSOLIDATED_COLS
             if c.startswith(("1.", "2.", "3.", "4.", "5."))]
    for _i, r in con.iterrows():
        ws.append([r["Lease Reference"], r["Lease Name"], r["Prop Ref"],
                   r["Cost Centre"], r["Period"]]
                  + [r[c] if r[c] != "" else None for c in money])
    ws.freeze_panes = "A4"

    # ---- 2. Data (the full table with its banner rows) --------------
    ws2 = wb.create_sheet("Data")
    ws2.append([_DATA_GROUP.get(c, "") for c in CONSOLIDATED_COLS])
    ws2.append([_DATA_SUB.get(c) or _DATA_SOURCE.get(c, "")
                for c in CONSOLIDATED_COLS])
    ws2.append(list(CONSOLIDATED_COLS))
    for cell in ws2[1]:
        cell.font = band_font
        cell.fill = band_fill
    for cell in ws2[3]:
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for _i, r in con.iterrows():
        ws2.append([r[c] if r[c] != "" else None
                    for c in CONSOLIDATED_COLS])
    ws2.freeze_panes = "A4"

    # ---- 2b. the period sheet, named after the reporting month ------
    period = _period_of(as_at) if as_at is not None else ""
    if period:
        ws2b = wb.create_sheet(period)
        ws2b.append([])
        ws2b.append([])
        ws2b.append(CR_PIVOT_COLS)
        for cell in ws2b[3]:
            if cell.value:
                cell.font = head_font
                cell.fill = head_fill
                cell.alignment = Alignment(wrap_text=True,
                                           vertical="center")
        # only the selected period on this sheet
        just = con[con["Period"] == period]
        for _i, r in just.iterrows():
            ws2b.append([r["Lease Reference"], r["Lease Name"],
                         r["Prop Ref"], r["Cost Centre"], r["Period"]]
                        + [r[c] if r[c] != "" else None for c in money])
        num = just[money].apply(pd.to_numeric, errors="coerce").fillna(0.0)
        ws2b.append([])
        ws2b.append(["Total for " + period, "", "", "",
                     "%d lease(s)" % len(just)]
                    # an empty month's object columns can sum to '' on newer pandas
                    + [round(float(pd.to_numeric(num[c], errors="coerce")
                                   .fillna(0.0).sum() or 0.0), 6)
                       for c in money])
        for cell in ws2b[ws2b.max_row]:
            cell.font = band_font
            cell.fill = band_fill
        ws2b.freeze_panes = "A4"

    # ---- 3. Recon (the roll-forward proof) --------------------------
    ws3 = wb.create_sheet("Recon")
    if ledger is not None and as_at is not None:
        for row in _recon_rows(ledger, refs, as_at, entity):
            ws3.append(row)
        for cell in ws3[3]:
            if cell.value:
                cell.font = band_font
                cell.fill = band_fill
    else:
        chk = results.get("Report Balance Check")
        if chk is not None and len(chk):
            ws3.append(list(chk.columns))
            for cell in ws3[1]:
                cell.font = head_font
                cell.fill = head_fill
            for _i, r in chk.iterrows():
                ws3.append(list(r.values))

    # ---- 4. the supporting check sheets ----------------------------
    for key, tab in (("Report Category Summary", "Category Summary"),
                     ("Report Ledger Integrity", "Ledger Integrity"),
                     ("Report Lease Selection", "Lease Selection"),
                     ("Report Source Files", "Source Files")):
        df = results.get(key)
        if df is None or not len(df):
            continue
        w = wb.create_sheet(tab)
        w.append([str(c) for c in df.columns])
        for cell in w[1]:
            cell.font = head_font
            cell.fill = head_fill
        for _i, r in df.iterrows():
            w.append([None if (v is None or (isinstance(v, float)
                                             and v != v)) else v
                      for v in r.values])
        w.freeze_panes = "A2"

    for sheet in wb.worksheets:
        for col in list(sheet.columns)[:8]:
            letter = col[0].column_letter
            width = max((len(str(c.value)) for c in col[:60]
                         if c.value is not None), default=10)
            sheet.column_dimensions[letter].width = min(max(width + 2,
                                                            11), 34)
    wb.save(out_path)
    return out_path


# ---------------------------------------------------------------------
# The plain-language report: one row per lease, one column per month
# ---------------------------------------------------------------------
SIMPLE_MEASURES = [
    ("Net asset", "5.NET ASSET_CLOSING BALANCE",
     "What the leased assets are worth at each month end, after "
     "depreciation and impairment."),
    ("ROU asset", "1.ASSET_CLOSING BALANCE",
     "The right-of-use asset at each month end, before depreciation."),
    ("Depreciation to date", "2.ACCUMULATED DEPRECIATION_CLOSING BALANCE",
     "Total depreciation charged against the asset so far."),
    ("Impairment", "3.IMPAIRMENT_CLOSING BALANCE",
     "Impairment held against the asset."),
    ("Lease liability", "4.LEASE LIABILITY_CLOSING BALANCE",
     "What is still owed on the lease at each month end."),
    ("Depreciation charge", "DEPRECIATION",
     "The depreciation charged in the month itself."),
    ("Interest", "INTEREST EXPENSE",
     "The interest charged in the month itself."),
    ("Payments", "PAYMENT", "What was paid in the month itself."),
]
SIMPLE_KEYS = ["Lease Reference", "Lease Name", "Entity", "Country",
               "Currency"]


def _period_label(p):
    """202607 -> Jul 2026."""
    p = str(p).strip()
    if len(p) != 6 or not p.isdigit():
        return p
    months = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug",
              "Sep", "Oct", "Nov", "Dec")
    return "%s %s" % (months[int(p[4:6]) - 1], p[:4])


def simple_tables(consolidated):
    """{measure name: frame} - one row per lease, one column per month,
    months oldest to newest left to right."""
    out = {}
    if consolidated is None or not len(consolidated):
        return out
    con = consolidated.copy()
    con["Period"] = con["Period"].astype(str)
    periods = sorted(p for p in con["Period"].unique() if p)
    for name, col, _blurb in SIMPLE_MEASURES:
        if col not in con.columns:
            continue
        vals = pd.to_numeric(con[col], errors="coerce")
        work = con[SIMPLE_KEYS + ["Period"]].copy()
        work["_v"] = vals
        if not work["_v"].notna().any():
            continue
        piv = (work.pivot_table(index=SIMPLE_KEYS, columns="Period",
                                values="_v", aggfunc="sum")
               .reindex(columns=periods))
        piv = piv.reset_index()
        piv.columns = SIMPLE_KEYS + [_period_label(c) for c in periods]
        out[name] = piv.sort_values("Lease Reference").reset_index(
            drop=True)
    return out


def write_simple_report(results, out_path, as_at=None):
    """The easy-to-read version: a short explanation, then one sheet per
    figure with the months running across the page. The lease columns
    are frozen, so the months scroll while you keep your place."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    con = results.get("Report Consolidated")
    tables = simple_tables(con)
    if not tables:
        raise ValueError("No consolidated report to summarise.")
    out_path = Path(out_path)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="003F2D")
    title_font = Font(bold=True, size=14, color="003F2D")
    band = PatternFill("solid", fgColor="DBE7E1")

    wb = Workbook()
    ws = wb.active
    ws.title = "Start here"
    ws["A1"] = "IFRS 16 lease report"
    ws["A1"].font = title_font
    period = _period_label(_period_of(as_at)) if as_at is not None else ""
    lines = [
        "",
        "Reporting month: %s" % (period or "(not set)"),
        "",
        "Every sheet below shows one lease per row and one month per "
        "column, oldest month first.",
        "The lease columns stay in place while you scroll sideways "
        "through the months.",
        "",
        "What each sheet shows:",
    ]
    for i, line in enumerate(lines, start=2):
        ws.cell(i, 1, line)
    row = 2 + len(lines)
    for name, col, blurb in SIMPLE_MEASURES:
        if name not in tables:
            continue
        ws.cell(row, 1, name).font = Font(bold=True)
        ws.cell(row, 2, blurb)
        row += 1
    row += 1
    ws.cell(row, 1, "Figures are as posted in the Lease Accelerator "
                    "ledger. Balances are shown at each month end.")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 96

    for name, _col, _blurb in SIMPLE_MEASURES:
        df = tables.get(name)
        if df is None:
            continue
        w = wb.create_sheet(name[:31])
        w.append([str(c) for c in df.columns])
        for cell in w[1]:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center",
                                       vertical="center")
        for _i, r in df.iterrows():
            w.append([None if (v is None or (isinstance(v, float)
                                             and v != v)) else v
                      for v in r.values])
        last = w.max_row
        if last > 1:
            w.append([])
            total = ["Total", "", "", "", ""]
            for ci in range(len(SIMPLE_KEYS) + 1, len(df.columns) + 1):
                L = get_column_letter(ci)
                total.append("=SUM(%s2:%s%d)" % (L, L, last))
            w.append(total)
            for cell in w[w.max_row]:
                cell.font = Font(bold=True, color="003F2D")
                cell.fill = band
        for ci in range(len(SIMPLE_KEYS) + 1, len(df.columns) + 1):
            L = get_column_letter(ci)
            w.column_dimensions[L].width = 15
            for r in range(2, w.max_row + 1):
                w.cell(r, ci).number_format = "#,##0.00"
        for ci, width in enumerate((16, 34, 26, 16, 10), start=1):
            w.column_dimensions[get_column_letter(ci)].width = width
        # keep the lease columns visible while the months scroll
        w.freeze_panes = "%s2" % get_column_letter(len(SIMPLE_KEYS) + 1)
    wb.save(out_path)
    return out_path


def lease_names_from_sequentra(path):
    """{lease reference: contract name} from the Sequentra export.

    Sequentra carries the contract / building name in 'Record Name'
    (for example 'Meeg Bank Centre'), which is what the report shows as
    the Lease Name."""
    out = {}
    if not path or not Path(path).exists():
        return out
    try:
        import lease_engine
        seq = lease_engine.load_sequentra(path)
    except Exception:
        return out
    name_col = next((c for c in ("Contract Name", "Record Name",
                                 "Expanded Name")
                     if c in seq.columns), None)
    if not name_col:
        return out
    for ref, nm in zip(seq["Record ID"], seq[name_col]):
        ref = str(ref).strip().upper()
        nm = "" if nm is None else str(nm).strip()
        if ref and nm and nm.lower() != "nan":
            out.setdefault(ref, nm)
    return out


def entity_map(path, settings=None):
    """{lease reference: entity} from an entity data report.

    The report is whatever extract carries the legal entity per lease.
    The columns are found by name, and the names can be overridden in
    settings under 'entity_report' -> {'reference': ..., 'entity': ...}
    so a new extract can be pointed at without changing code.
    """
    out = {}
    if not path or not Path(path).exists():
        return out
    cfg = ((settings or {}).get("entity_report") or {})
    want_ref = str(cfg.get("reference", "")).strip().lower()
    want_ent = str(cfg.get("entity", "")).strip().lower()
    ref_names = ([want_ref] if want_ref else []) + [
        "lease reference", "schedule number", "schedule #", "record id",
        "schedule", "contract reference"]
    ent_names = ([want_ent] if want_ent else []) + [
        "entity", "legal entity", "company", "company name", "lessee",
        "bu", "business unit"]
    try:
        book = pd.read_excel(path, sheet_name=None, header=None,
                             engine="calamine")
    except Exception:
        try:
            book = pd.read_excel(path, sheet_name=None, header=None)
        except Exception:
            return out
    for _sheet, raw in book.items():
        if raw is None or not len(raw):
            continue
        for i in range(min(12, len(raw))):
            cells = [str(v).strip().lower() for v in raw.iloc[i].tolist()]
            rcol = next((j for j, c in enumerate(cells)
                         if c in ref_names), None)
            ecol = next((j for j, c in enumerate(cells)
                         if c in ent_names), None)
            if rcol is None or ecol is None:
                continue
            body = raw.iloc[i + 1:]
            for ref, ent in zip(body.iloc[:, rcol], body.iloc[:, ecol]):
                ref = str(ref).strip().upper()
                ent = "" if ent is None else str(ent).strip()
                if ref and ref != "NAN" and ent and ent.lower() != "nan":
                    out.setdefault(ref, ent)
            if out:
                return out
    return out


def lease_names_from_assets(path):
    """{lease reference: building / address} from the LA asset report."""
    out = {}
    if not path or not Path(path).exists():
        return out
    try:
        raw = pd.read_excel(path, sheet_name=0, header=1,
                            engine="calamine")
    except Exception:
        return out
    raw.columns = [str(c).strip() for c in raw.columns]
    key = next((c for c in raw.columns
                if str(c).strip().lower() in ("schedule #", "schedule number",
                                              "schedule")), None)
    name = next((c for c in raw.columns
                 if str(c).strip().lower() in ("address", "description",
                                               "building", "location")), None)
    if not key or not name:
        return out
    for k, v in zip(raw[key], raw[name]):
        k = str(k).strip().upper()
        if k and k != "NAN" and str(v).strip().lower() != "nan":
            out.setdefault(k, str(v).strip())
    return out


def discount_rates_from_la(path):
    """{lease reference: discount rate} from the LA transaction report.

    Keyed on 'Transaction #'; the rate column is matched by name so a
    'Discount Rate', 'IBR' or 'Current Discount Rate' heading all work."""
    out = {}
    if not path or not Path(path).exists():
        return out
    try:
        import lease_engine
        la = lease_engine.load_la_register(path)
    except Exception:
        return out
    key = next((c for c in la.columns
                if str(c).strip().lower() == "transaction #"), None)
    if key is None:
        return out
    rate_col = None
    for c in la.columns:
        n = re.sub(r"[^a-z0-9]+", " ", str(c).strip().lower()).strip()
        if n in ("discount rate", "ibr", "current discount rate",
                 "incremental borrowing rate", "interest rate"):
            rate_col = c
            break
    if rate_col is None:
        rate_col = next((c for c in la.columns
                         if "discount rate" in str(c).strip().lower()
                         or str(c).strip().lower().startswith("ibr")), None)
    if rate_col is None:
        return out
    for ref, val in zip(la[key], la[rate_col]):
        ref = str(ref).strip().upper()
        if not ref or ref == "NAN" or val is None or pd.isna(val):
            continue
        try:
            rate = float(str(val).replace("%", "").strip())
        except (TypeError, ValueError):
            continue
        if 0 < rate < 1:            # 0.105 -> 10.5
            rate = round(rate * 100, 4)
        out.setdefault(ref, rate)
    return out


def lease_terms_from_sequentra(path):
    """{lease reference: {start, end, entity}} from the Sequentra export."""
    out = {}
    if not path or not Path(path).exists():
        return out
    try:
        import lease_engine
        seq = lease_engine.load_sequentra(path)
    except Exception:
        return out

    def fmt(v):
        try:
            d = pd.to_datetime(v, errors="coerce")
            return "" if pd.isna(d) else d.strftime("%Y-%m-%d")
        except Exception:
            return ""
    def first(r, names):
        for n in names:
            v = r.get(n)
            if v is not None and str(v).strip() \
                    and str(v).strip().lower() != "nan":
                return str(v).strip()
        return ""
    for _i, r in seq.iterrows():
        ref = str(r.get("Record ID", "")).strip().upper()
        if not ref:
            continue
        out[ref] = {
            "start": fmt(r.get("Current Term Start Date")),
            "end": fmt(r.get("Current Lease End Date")),
            "entity": first(r, ("Lessee", "Legal Entity", "Entity")),
            "country": first(r, ("Country", "Asset Location(s)")),
            "currency": first(r, ("Currency", "Currency Code",
                                  "Transactional Currency")),
        }
    return out


def run_ledger_check(ledger_path, not_applicable=None, contracts=None,
                     as_at=None, months_back=12, assets_path=None,
                     sequentra_path=None, full_life=True,
                     entity_path=None, settings=None, period_from=None,
                     from_selected=False, la_path=None):
    """Read a ledger export and produce the report sheets.

    ledger_path may be a single export, a FOLDER (every ledger export in
    it is combined into one report), or a list of files.

    By default every lease runs from its first posted period to the end
    of its own schedule - the report is not truncated at the reporting
    month. Set from_selected=True to start at the selected month instead
    and still run to the end of each lease.
    """
    used, skipped, dupes = [], [], 0
    if isinstance(ledger_path, (list, tuple, set)):
        paths = [Path(x) for x in ledger_path]
        if len(paths) == 1:
            ledger = load_ledger_export(paths[0])
            source_label = paths[0].name
        else:
            ledger, used, skipped, dupes = combine_ledger_exports(paths)
            source_label = "%d files combined" % len(used)
    elif Path(ledger_path).is_dir():
        ledger, used, skipped, dupes = load_ledger_folder(ledger_path)
        source_label = "%d file(s) combined from %s" % (
            len(used), Path(ledger_path).name)
    else:
        ledger = load_ledger_export(ledger_path)
        source_label = Path(ledger_path).name
    as_at = as_at or date.today()
    if from_selected and not period_from:
        period_from = _period_of(as_at)
    refs, listing = select_leases(ledger, not_applicable=not_applicable,
                                  as_at=as_at, months_back=months_back,
                                  contracts=contracts,
                                  full_life=full_life)
    report = movement_report(ledger, refs=refs, as_at=as_at,
                             months_back=months_back,
                             full_life=full_life)
    check = balance_check(report)
    integrity = ledger_integrity(ledger, as_at=as_at,
                                 months_back=months_back)
    summary = category_summary(ledger, refs=refs, as_at=as_at)
    breaks = int((check["Status"] == "CHECK").sum()) if len(check) else 0
    unbal = int((integrity["Status"] == "CHECK").sum()) \
        if len(integrity) else 0
    overview = pd.DataFrame([
        {"Item": "Ledger file", "Value": source_label},
        {"Item": "As at", "Value": as_at.strftime("%d/%m/%Y")},
        {"Item": "Months included", "Value": months_back},
        {"Item": "Periods included",
         "Value": ("full lease life (%s to %s)"
                   % (ledger["Period"].dropna().min(),
                      ledger["Period"].dropna().max())) if full_life
         else ("%s to %s" % (_period_floor(as_at, months_back),
                             _period_of(as_at)))},
        {"Item": "Ledger lines read", "Value": len(ledger)},
        {"Item": "Leases on the ledger",
         "Value": int(ledger["Lease Reference"].nunique())},
        {"Item": "Leases included in the check", "Value": len(refs)},
        {"Item": "Leases filtered out",
         "Value": int((listing["Decision"] == "Excluded").sum())},
        {"Item": "Report rows", "Value": len(report)},
        {"Item": "Roll-forward breaks", "Value": breaks},
        {"Item": "Periods where debits <> credits", "Value": unbal},
    ], columns=["Item", "Value"])
    if used:
        extra = [{"Item": "Ledger files combined", "Value": len(used)}]
        if dupes:
            extra.append({"Item": "Overlapping lines removed",
                          "Value": dupes})
        if skipped:
            extra.append({"Item": "Files skipped (not a ledger export)",
                          "Value": len(skipped)})
        overview = pd.concat([overview, pd.DataFrame(extra)],
                             ignore_index=True)
    names = lease_names_from_sequentra(sequentra_path)
    if not names:                          # fall back to the asset report
        names = lease_names_from_assets(assets_path)
    consolidated = consolidated_report(
        ledger, refs=refs, as_at=as_at, months_back=months_back,
        lease_names=names,
        lease_terms=lease_terms_from_sequentra(sequentra_path),
        entities=entity_map(entity_path, settings),
        discount_rates={**discount_rates_from_la(la_path),
                        **discount_rates_from_sequentra(sequentra_path)},
        full_life=full_life, period_from=period_from)
    na_check = net_asset_check(consolidated)
    na_breaks = int((na_check["Status"] == "CHECK").sum()) \
        if len(na_check) else 0
    with_rate = int((consolidated["Discount Rate"].astype(str).str.strip()
                     != "").sum()) if len(consolidated) else 0
    overview = pd.concat([overview, pd.DataFrame([
        {"Item": "Net asset roll-forward breaks", "Value": na_breaks},
        {"Item": "Rows with a discount rate",
         "Value": "%d of %d" % (with_rate, len(consolidated))},
    ])], ignore_index=True)
    out = {
        "Report Overview": overview,
        "Report Consolidated": consolidated,
        # a check, not a report sheet (finance: "does not have to be saved
        # on the report")
        "Check Net Asset": na_check,
        "Report Balance Check": check,
        "Report Category Summary": summary,
        "Report Ledger Integrity": integrity,
        "Report Lease Movements": report,
        "Report Lease Selection": listing,
    }
    if used or skipped:
        rows = [dict(u, Status="Combined") for u in used]
        rows += [{"File": s["File"], "Rows": "", "Leases": "",
                  "First Period": "", "Last Period": "",
                  "Status": "Skipped - " + s["Reason"]}
                 for s in skipped]
        out["Report Source Files"] = pd.DataFrame(
            rows, columns=["File", "Rows", "Leases", "First Period",
                           "Last Period", "Status"])
    # kept for the workbook writer; not a displayed sheet
    out["_ledger"] = ledger
    out["_refs"] = refs
    out["_as_at"] = as_at
    return out
