# 1.0.98: from CBX-Connect 2.10.3 cbx/knowledge.py
"""What CBX knows about SAP (2.10.3, Edwin 16 Oct: "here is some knowledge … it needs to read images, and learn, and know
what to use where; if it sees it's SAP, it needs to put it into the summary").

The one place the app asks about SAP. The AP Tool (Tess, Pip, Mia, Settings) and the Support desk (the SAP block in a
ticket's summary) both use it:

    doc_type(db, code)                    an SAP document type: what SAP says it is, which account types it allows, and
                                          what it means for AP (Edwin's table)
    gl_account(db, account)               a G/L account and its name, from the chart of accounts
    company(db, code_or_name_or_vat)      a company code from the AP Entities
    sap_words()                           SAP transaction codes and screen words (a precompiled pattern)
    vendor_in_company(db, vendor, cc)     is the vendor opened in that company code in SAP (None: CBX can't tell)
    companies_of_vendor(db, vendor)       the company codes a vendor is opened in
    cost_centre(db, cc, company_code)     a cost centre in SAP (its company code, currency, profit centre, description)
    company_currency(db, company_code)    the currency a company code books in (most of its cost centres)
    sap_vendor / sap_vendors_by_vat       1.0.102: SAP's vendor general data (LFA1): name, VAT number, blocked, payment
    vendor_blocked / terms_days           terms ("N30" → 30 days)

The tables (ap_doc_types, ap_gl_accounts, ap_vendor_companies, ap_cost_centres) start from the seed files in cbx/knowledge_data/ (SAP's
own exports, converted by tools/make_knowledge_data.py) and admins replace them in AP → Settings (upload or Pull from
SAP). Lookups are cached in this process for a minute; a save in Settings forgets the cache at once.

Everything here is read-only for callers: what SAP's lists say is data, never instructions.
"""
from __future__ import annotations

import csv
import io
import re
import threading
import time
from pathlib import Path
from typing import Any, Callable, Iterable, Optional

DATA = Path(__file__).resolve().parent / "knowledge_data"
DOC_TYPES_CSV = DATA / "sap_doc_types.csv"
CHART_CSV = DATA / "chart_of_accounts.csv"
VENDOR_CC_CSV = DATA / "vendor_company_codes.csv"
COST_CENTRES_CSV = DATA / "cost_centres.csv"
DUMMY_CC = "SAP-DUMMY"  # SAP's placeholder cost centre: never a real one
TTL = 60.0  # seconds a lookup table is kept in this process

# SAP's account-type letters (Document Type List, column "Account Types")
LETTERS = {"A": "asset", "D": "customer", "K": "vendor", "M": "material", "S": "gl"}
LETTER_WORDS = {"A": "assets", "D": "customers", "K": "vendors", "M": "material", "S": "G/L accounts"}
KINDS = ("invoice", "credit", "payment", "clearing", "journal", "migration", "other")


class Refused(Exception):
    pass


# ---- the cache ----------------------------------------------------------------------------------------------------------
_cache: dict[str, tuple[float, Any]] = {}
_lock = threading.Lock()


def _cached(key: str, build: Callable[[], Any]) -> Any:
    now = time.monotonic()
    with _lock:
        hit = _cache.get(key)
        if hit is not None and now - hit[0] < TTL:
            return hit[1]
    value = build()
    with _lock:
        _cache[key] = (time.monotonic(), value)
    return value


def forget(*keys: str) -> None:
    """Drop cached lookups (all of them, or the ones named): a save in Settings calls this."""
    with _lock:
        if not keys:
            _cache.clear()
        for k in keys:
            _cache.pop(k, None)


# ---- small helpers --------------------------------------------------------------------------------------------------------
def text(v: Any, n: int = 200) -> str:
    """A cell as text: an Excel number 120000.0 is "120000"; inner spaces stay, new lines go."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"[\r\n\t]+", " ", str(v)).strip()[:n]


def gl_key(v: Any) -> str:
    """A G/L account compared ignoring case, spaces and leading zeros: "0006123" = "6123", "ja600" = "JA600"."""
    s = re.sub(r"\s+", "", text(v, 40)).upper()
    return s.lstrip("0") or ("0" if s else "")


def vendor_variants(v: Any) -> set[str]:
    """A vendor number as SAP may show it: as given, without leading zeros, and padded to SAP's 10 digits."""
    s = re.sub(r"\s+", "", text(v, 20))
    if not s:
        return set()
    bare = s.lstrip("0") or "0"
    out = {s, bare}
    if bare.isdigit():  # SAP pads vendor numbers with zeros to 10 digits; an export may keep some, all or none of them
        out |= {bare.zfill(n) for n in range(len(bare) + 1, 11)}
    return out


def letters_of(account_types: str) -> str:
    return "".join(ch for ch in "ADKMS" if ch in (account_types or "").upper())


def allows_of(account_types: str) -> set[str]:
    return {LETTERS[ch] for ch in letters_of(account_types)}


def letters_said(account_types: str) -> str:
    """"DS" → "D, S" (the way Tess and the Settings page say it)."""
    return ", ".join(letters_of(account_types)) or "none"


# ---- reading SAP's exports (the seed tool and the Settings uploads use the same readers) ----------------------------------
MAX_MB = 25


def _rows(data: bytes, filename: str, most: int, cols: int = 60) -> Iterable[list]:
    """Rows of an .xlsx (the first sheet) or .csv, lazily; a list longer than `most` rows is refused. 1.0.103: `cols` —
    SAP's vendor general data has 200+ columns (its payment block is the 71st)."""
    if len(data) > MAX_MB * 1024 * 1024:
        raise Refused(f"The file is over {MAX_MB} MB.")
    n = 0
    if filename.lower().endswith(".csv"):
        rd = csv.reader(io.StringIO(data.decode("utf-8-sig", "replace")))
        try:
            for r in rd:
                n += 1
                if n > most + 20:
                    raise Refused(f"More than {most} rows: split the list.")
                yield list(r[:cols])
        except csv.Error as e:
            raise Refused(f"The file could not be read as CSV ({type(e).__name__}).") from None
        return
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise Refused(f"The file could not be read as Excel ({type(e).__name__}).") from None
    try:
        ws = wb["Data"] if "Data" in wb.sheetnames else wb.worksheets[0]
        for r in ws.iter_rows(values_only=True, max_col=cols):
            n += 1
            if n > most + 20:
                raise Refused(f"More than {most} rows: split the list.")
            yield list(r)
    except Refused:
        raise
    except Exception as e:  # noqa: BLE001 — a broken sheet is "could not be read", not a 500
        raise Refused(f"The file could not be read as Excel ({type(e).__name__}).") from None
    finally:
        wb.close()  # 2.10.3: closed before anything else happens (Windows)


def _head(v: Any) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip().lower()


def _find(heads: list, names: tuple[str, ...]) -> Optional[int]:
    norm = [_head(h) for h in heads]
    for nm in names:
        if nm in norm:
            return norm.index(nm)
    return None


DOC_NEEDED = ("document type", "account types", "description")


def read_doc_type_list(data: bytes, filename: str = "") -> list[dict]:
    """SAP's Document Type List (sheet Data: Document Type, Number Range, Account Types, Reverse DocumentType, …,
    Description) → [{code, description, account_types, number_range, reverse, sap: {column: value}}]."""
    it = iter(_rows(data, filename, 2000))
    heads: list = []
    for _ in range(10):
        r = next(it, None)
        if r is None:
            break
        if all(_find(r, (nm,)) is not None for nm in DOC_NEEDED):
            heads = [text(h, 60) for h in r]
            break
    if not heads:
        raise Refused("No header row with Document Type, Account Types and Description in the first 10 rows.")
    while heads and not heads[-1]:
        heads.pop()
    ic, ia, idesc = (_find(heads, (nm,)) for nm in DOC_NEEDED)
    inr, irev = _find(heads, ("number range",)), _find(heads, ("reverse documenttype", "reverse document type", "reverse doc type"))
    out, seen = [], set()
    for r in it:
        cell = lambda i: text(r[i], 80) if i is not None and i < len(r) else ""  # noqa: E731
        code = cell(ic).upper()
        if not re.fullmatch(r"[A-Z0-9]{1,4}", code) or code in seen:
            continue
        seen.add(code)
        out.append({"code": code, "description": cell(idesc), "account_types": letters_of(cell(ia)), "number_range": cell(inr)[:8],
                    "reverse": cell(irev).upper()[:4], "sap": {h: cell(i) for i, h in enumerate(heads) if h}})
    if not out:
        raise Refused("No document types in it.")
    return out


def read_chart(data: bytes, filename: str = "") -> list[dict]:
    """The chart of accounts (G/L Account | G/L Acct Long Text) → [{account, name}] (the same account once)."""
    it = iter(_rows(data, filename, 60000))
    acc = name = None
    for _ in range(10):
        r = next(it, None)
        if r is None:
            break
        acc = _find(r, ("g/l account", "gl account", "g/l acct", "account", "g/l"))
        name = _find(r, ("g/l acct long text", "long text", "g/l account long text", "g/l acct short text", "short text", "name", "description"))
        if acc is not None:
            break
    if acc is None:
        raise Refused("No header row with G/L Account in the first 10 rows.")
    out, seen = [], set()
    for r in it:
        a = re.sub(r"\s+", "", text(r[acc] if acc < len(r) else "", 40))
        if not a or len(a) > 20 or gl_key(a) in seen:
            continue
        seen.add(gl_key(a))
        out.append({"account": a, "name": text(r[name], 200) if name is not None and name < len(r) else ""})
    if not out:
        raise Refused("No G/L accounts in it.")
    return out


VENDOR_PAIRS_MAX = 500000


def read_vendor_companies(data: bytes, filename: str = "") -> list[tuple[str, str]]:
    """SAP's vendor ↔ company code export (Vendor | Company Code) → [(vendor, company_code)], each pair once. Vendor
    numbers stay text as SAP gives them (leading zeros kept)."""
    it = iter(_rows(data, filename, VENDOR_PAIRS_MAX))
    iv = ic = None
    for _ in range(10):
        r = next(it, None)
        if r is None:
            break
        iv = _find(r, ("vendor", "supplier", "vendor number", "vendor no", "account", "lifnr"))
        ic = _find(r, ("company code", "cocd", "company", "bukrs"))
        if iv is not None and ic is not None:
            break
    if iv is None or ic is None:
        raise Refused("No header row with Vendor and Company Code in the first 10 rows.")
    out: dict[tuple[str, str], None] = {}
    for r in it:
        v = re.sub(r"\s+", "", text(r[iv] if iv < len(r) else "", 20))
        c = re.sub(r"\s+", "", text(r[ic] if ic < len(r) else "", 12)).upper()
        if v and c:
            out[(v, c)] = None
    if not out:
        raise Refused("No vendor and company code pairs in it.")
    return list(out)


COST_CENTRES_MAX = 300000
CC_FIELDS = ("currency", "profit_centre", "description")


def read_cost_centres(data: bytes, filename: str = "") -> list[dict]:
    """SAP's Cost Center List (sheet Data: Cost Center | Company Code | Currency | Profit Center | Description) →
    [{cost_centre, company_code, currency, profit_centre, description}], one per (cost centre, company code): the first
    row wins, a later row of the same pair only fills what the first left blank. Codes stay text as SAP gives them."""
    it = iter(_rows(data, filename, COST_CENTRES_MAX))
    cols: dict[str, Optional[int]] = {}
    for _ in range(10):
        r = next(it, None)
        if r is None:
            break
        cols = {"cost_centre": _find(r, ("cost center", "cost centre", "kostl")), "company_code": _find(r, ("company code", "bukrs", "cocd")),
                "currency": _find(r, ("currency", "waers")), "profit_centre": _find(r, ("profit center", "profit centre", "prctr")),
                "description": _find(r, ("description", "name", "ktext", "ltext"))}
        if cols["cost_centre"] is not None and cols["company_code"] is not None:
            break
    if cols.get("cost_centre") is None or cols.get("company_code") is None:
        raise Refused("No header row with Cost Center and Company Code in the first 10 rows.")
    out: dict[tuple[str, str], dict] = {}
    for r in it:
        cell = lambda k, n: re.sub(r"\s+", " ", text(r[cols[k]], n)) if cols.get(k) is not None and cols[k] < len(r) else ""  # noqa: E731
        cc, co = cell("cost_centre", 20).replace(" ", ""), cell("company_code", 12).replace(" ", "").upper()
        if not cc or not co:
            continue
        row = {"cost_centre": cc, "company_code": co, "currency": cell("currency", 5).upper(), "profit_centre": cell("profit_centre", 20).replace(" ", ""),
               "description": cell("description", 120)}
        old = out.get((cc, co))
        if old is None:
            out[(cc, co)] = row
        else:
            for f in CC_FIELDS:
                old[f] = old[f] or row[f]
    if not out:
        raise Refused("No cost centres in it.")
    return list(out.values())


# ---- the seed files ---------------------------------------------------------------------------------------------------------
def _read_csv(path: Path) -> list[list[str]]:
    if not path.is_file():
        return []
    with path.open(encoding="utf-8", newline="") as fh:  # 2.10.3: closed at once (Windows)
        return list(csv.reader(fh))


def seed_doc_types() -> list[dict]:
    rows = _read_csv(DOC_TYPES_CSV)
    if not rows:
        return []
    return read_doc_type_list(_csv_bytes(rows), "seed.csv")


def seed_chart() -> list[dict]:
    rows = _read_csv(CHART_CSV)
    return [{"account": r[0], "name": r[1] if len(r) > 1 else ""} for r in rows[1:] if r and r[0]]


def seed_vendor_companies() -> list[tuple[str, str]]:
    rows = _read_csv(VENDOR_CC_CSV)
    return [(r[0], r[1]) for r in rows[1:] if len(r) > 1 and r[0] and r[1]]


def seed_cost_centres() -> list[dict]:
    rows = _read_csv(COST_CENTRES_CSV)
    return [{"cost_centre": r[0], "company_code": r[1], "currency": r[2], "profit_centre": r[3], "description": r[4]} for r in rows[1:] if len(r) >= 5 and r[0] and r[1]]


def _csv_bytes(rows: list[list]) -> bytes:
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\n").writerows(rows)
    return buf.getvalue().encode("utf-8")


def write_csv(path: Path, header: list[str], rows: Iterable[Iterable]) -> int:
    """The seed files: utf-8, one header row (tools/make_knowledge_data.py). 1.0.103: a path ending .gz is gzipped."""
    n = 0
    if path.name.endswith(".gz"):
        import gzip

        with gzip.open(path, "wt", encoding="utf-8", newline="") as fh:
            w = csv.writer(fh, lineterminator="\n")
            w.writerow(header)
            for r in rows:
                w.writerow(r)
                n += 1
        return n
    with path.open("w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(header)
        for r in rows:
            w.writerow(r)
            n += 1
    return n


# ---- document types ---------------------------------------------------------------------------------------------------------
def _doc_index(db) -> dict[str, dict]:
    def build() -> dict[str, dict]:
        from .models.ap import ApDocType

        seed = {d["code"]: d for d in seed_doc_types()}
        out: dict[str, dict] = {}
        rows = db.query(ApDocType).all()
        for d in rows:
            s = seed.get(d.code.upper(), {})
            at = d.account_types or ("" if d.sap else s.get("account_types", ""))  # a row an admin added by hand: SAP's seed fills in what it knows
            out[d.code.upper()] = {"code": d.code.upper(), "description": d.description or s.get("description", ""), "account_types": at,
                                   "allows": allows_of(at), "reverse": d.reverse or s.get("reverse", ""), "number_range": d.number_range or s.get("number_range", ""),
                                   "kind": d.meaning if d.meaning in KINDS else "other", "label": d.label or "", "needs_check": bool(d.needs_check)}
        if not rows:  # an empty table: SAP's list as shipped, Edwin's meanings unknown
            for code, s in seed.items():
                out[code] = {"code": code, "description": s["description"], "account_types": s["account_types"], "allows": allows_of(s["account_types"]),
                             "reverse": s["reverse"], "number_range": s["number_range"], "kind": "other", "label": "", "needs_check": False}
        return out

    return _cached("doc_types", build)


def doc_type(db, code: Any) -> Optional[dict]:
    """{code, description, account_types: "AKMS", allows: {"asset","customer","vendor","material","gl"}, reverse,
    number_range, kind} — kind is CBX's meaning for AP (invoice / credit / payment / clearing / journal / migration /
    other). Also label (Edwin's) and needs_check. None: not a document type CBX knows."""
    c = text(code, 8).upper()
    d = _doc_index(db).get(c) if c else None
    return {**d, "allows": set(d["allows"])} if d else None


def doc_types_all(db) -> dict[str, dict]:
    return _doc_index(db)


# ---- G/L accounts -------------------------------------------------------------------------------------------------------------
def _gl_index(db) -> dict[str, tuple[str, str]]:
    def build() -> dict[str, tuple[str, str]]:
        from .models.ap import ApGlAccount

        rows = db.query(ApGlAccount.account, ApGlAccount.name).filter(ApGlAccount.is_active.is_(True)).all()
        if not rows and not db.query(ApGlAccount.id).first():
            rows = [(r["account"], r["name"]) for r in seed_chart()]  # an empty table: the chart as shipped
        return {gl_key(a): (a, n or "") for a, n in rows if a}

    return _cached("gl", build)


def gl_account(db, account: Any) -> Optional[dict]:
    """{account, name} from the chart of accounts, matching ignoring case, spaces and leading zeros. None: not in it."""
    k = gl_key(account)
    hit = _gl_index(db).get(k) if k else None
    return {"account": hit[0], "name": hit[1]} if hit else None


def chart_keys(db) -> Optional[set[str]]:
    """The chart's accounts (as gl_key), or None when there is no chart at all (then nothing is checked against it)."""
    idx = _gl_index(db)
    return set(idx) if idx else None


def gl_search(db, q: str, limit: int = 20) -> list[dict]:
    """G/L accounts by number (start first) or by words in the name — for the coding forms' search."""
    q = text(q, 80)
    if not q:
        return []
    idx = _gl_index(db)
    k, words = gl_key(q), [w for w in re.split(r"\W+", q.lower()) if w]
    starts, has, named = [], [], []
    for key, (a, n) in idx.items():
        if k and key.startswith(k):
            starts.append((a, n))
        elif k and k in key:
            has.append((a, n))
        elif words and all(w in n.lower() for w in words):
            named.append((a, n))
        if len(starts) >= limit:
            break
    out = sorted(starts)[:limit] + sorted(has)[:limit] + sorted(named, key=lambda x: (len(x[1]), x[0]))[:limit]
    return [{"account": a, "name": n} for a, n in out[: max(1, min(int(limit or 20), 100))]]


# ---- company codes ------------------------------------------------------------------------------------------------------------
def split_bill_to(bill_to: Any) -> tuple[str, str]:
    """"CBRE Excellerate Facilities Management (Pty) Ltd - RIO2X21" → (the name, "RIO2X21"): the part after the last
    " - " is a client or site code when it is short (at most 4 words, 40 characters)."""
    s = re.sub(r"\s+", " ", text(bill_to, 300))
    m = re.match(r"^(.*\S)\s+[-–]\s+(\S.{0,39})$", s)
    if not m or len(m.group(2).split()) > 4 or len(m.group(1)) < 4:
        return s, ""
    return m.group(1).strip(), m.group(2).strip()


def company(db, code_or_name_or_vat: Any) -> Optional[dict]:
    """{code, name, entity} from AP Entities: by company code, by VAT number (10 digits), or by name (the name before
    " - " too; a name only part of which fits is not taken). None: no company code CBX knows."""
    from .ap import masters as M
    from .models.ap import ApEntity

    raw = text(code_or_name_or_vat, 300)
    if not raw:
        return None
    ents = _cached("entities", lambda: [{"code": e.code, "name": e.name or "", "entity": e.entity or "", "vat": e.vat_no or "", "aliases": e.aliases or ""}
                                        for e in db.query(ApEntity).filter(ApEntity.is_active.is_(True)).all()])
    up = raw.upper()
    e = next((x for x in ents if x["code"].upper() == up), None)
    if e is None:
        v = M.vat_key(raw)
        if len(v) == 10 and re.fullmatch(r"[\d\s/]+", raw):
            e = next((x for x in ents if x["vat"] == v), None)
    if e is None:
        for nm in dict.fromkeys([raw, split_bill_to(raw)[0]]):
            k = M.name_key(nm)
            if len(k) < 4:
                continue
            e = next((x for x in ents if k in {M.name_key(n) for n in [x["name"], x["entity"], *x["aliases"].splitlines()] if n}), None)
            if e is not None:
                break
    return {"code": e["code"], "name": e["name"], "entity": e["entity"]} if e else None


# ---- vendors per company code -------------------------------------------------------------------------------------------------
def _vc_state(db) -> dict:
    def build() -> dict:
        from sqlalchemy import func

        from .models.ap import ApVendorCompany

        n = db.query(func.count(ApVendorCompany.id)).scalar() or 0
        codes = {c for (c,) in db.query(ApVendorCompany.company_code).distinct().all()} if n else set()
        return {"pairs": n, "codes": codes}

    return _cached("vendor_companies", build)


def vendor_in_company(db, vendor: Any, company_code: Any) -> Optional[bool]:
    """True / False: is the vendor opened in that company code in SAP. None when CBX can't tell: the table is empty, or
    the company code is not in SAP's export at all (a part of the export, or a company code SAP's list doesn't carry) —
    never a guess that turns a document red."""
    from .models.ap import ApVendorCompany

    vs, cc = vendor_variants(vendor), re.sub(r"\s+", "", text(company_code, 12)).upper()
    if not vs or not cc:
        return None
    st = _vc_state(db)
    if not st["pairs"] or cc not in st["codes"]:
        return None
    return db.query(ApVendorCompany.id).filter(ApVendorCompany.company_code == cc, ApVendorCompany.vendor.in_(sorted(vs))).first() is not None


def companies_of_vendor(db, vendor: Any) -> list[str]:
    """The company codes the vendor is opened in (sorted); [] when none or the table is empty."""
    from .models.ap import ApVendorCompany

    vs = vendor_variants(vendor)
    if not vs or not _vc_state(db)["pairs"]:
        return []
    return sorted({c for (c,) in db.query(ApVendorCompany.company_code).filter(ApVendorCompany.vendor.in_(sorted(vs))).all()})


# ---- cost centres -------------------------------------------------------------------------------------------------------------
def cc_variants(v: Any) -> set[str]:
    """A cost centre as SAP may show it: as given, without leading zeros, padded to SAP's 10 digits (numbers only)."""
    s = re.sub(r"\s+", "", text(v, 20)).upper()
    if not s:
        return set()
    bare = s.lstrip("0") or "0"
    out = {s, bare}
    if bare.isdigit():
        out |= {bare.zfill(n) for n in range(len(bare) + 1, 11)}
    return out


def _cc_state(db) -> dict:
    def build() -> dict:
        from sqlalchemy import func

        from .models.ap import ApCostCentre

        n = db.query(func.count(ApCostCentre.id)).scalar() or 0
        cur: dict[str, dict[str, int]] = {}
        if n:
            for co, c, k in db.query(ApCostCentre.company_code, ApCostCentre.currency, func.count(ApCostCentre.id)).group_by(ApCostCentre.company_code, ApCostCentre.currency).all():
                if c:
                    cur.setdefault(co, {})[c] = k
        codes = {co for (co,) in db.query(ApCostCentre.company_code).distinct().all()} if n else set()
        return {"rows": n, "codes": codes, "currency": {co: max(sorted(d.items()), key=lambda x: x[1])[0] for co, d in cur.items()}}

    return _cached("cost_centres", build)


def _cc_out(r) -> dict:
    return {"cost_centre": r.cost_centre, "company_code": r.company_code, "currency": r.currency, "profit_centre": r.profit_centre, "description": r.description}


def cost_centre(db, cc: Any, company_code: Any = None) -> Optional[dict]:
    """{cost_centre, company_code, currency, profit_centre, description} from SAP's Cost Center List; with no company
    code, the first match (by company code). None: not in the list (or the list is empty)."""
    from .models.ap import ApCostCentre

    vs = cc_variants(cc)
    if not vs or not _cc_state(db)["rows"]:
        return None
    q = db.query(ApCostCentre).filter(ApCostCentre.cost_centre.in_(sorted(vs)))
    co = re.sub(r"\s+", "", text(company_code, 12)).upper()
    if co:
        q = q.filter(ApCostCentre.company_code == co)
    r = q.order_by(ApCostCentre.company_code, ApCostCentre.id).first()
    return _cc_out(r) if r is not None else None


def cost_centre_in_company(db, cc: Any, company_code: Any) -> Optional[str]:
    """For Tess: "dummy" (SAP-DUMMY), "yes", "no" — or None when CBX can't tell (the list is empty, or SAP's list has no
    cost centre at all for that company code)."""
    s = re.sub(r"\s+", "", text(cc, 20)).upper()
    co = re.sub(r"\s+", "", text(company_code, 12)).upper()
    if not s or not co:
        return None
    if s == DUMMY_CC:
        return "dummy"
    st = _cc_state(db)
    if not st["rows"] or co not in st["codes"]:
        return None
    return "yes" if cost_centre(db, s, co) else "no"


def cost_centres_cover(db, company_code: Any = None) -> bool:
    """Is SAP's Cost Center List loaded (and, with a company code, does it have cost centres for that company)?"""
    st = _cc_state(db)
    co = re.sub(r"\s+", "", text(company_code, 12)).upper()
    return bool(st["rows"]) and (not co or co in st["codes"])


def company_currency(db, company_code: Any) -> Optional[str]:
    """The currency a company code books in: the most common currency of its cost centres. None: not in the list."""
    co = re.sub(r"\s+", "", text(company_code, 12)).upper()
    return _cc_state(db)["currency"].get(co) if co else None


# ---- 1.0.102: SAP's vendor general data (LFA1) ------------------------------------------------------------------------------
SAP_VENDORS_MAX = 300000


def sap_vat(v: Any) -> str:
    """A VAT number as CBX compares it: digits only, South Africa's leading 4 kept ("ZA 4123 456 789" → "4123456789")."""
    return re.sub(r"\D", "", text(v, 40))[:20]


def terms_days(zterm: Any) -> Optional[int]:
    """SAP's terms of payment → days, only for a plain days term: "N30" / "Z030" → 30; anything else (e.g. "0001",
    "EOM") → None (a person types the days)."""
    m = re.fullmatch(r"[NZ]0*(\d{1,3})", re.sub(r"\s+", "", text(zterm, 8)).upper())
    n = int(m.group(1)) if m else None
    return n if n is not None and n <= 365 else None


def _sv_count(db) -> int:
    def build() -> int:
        from sqlalchemy import func

        from .models.ap import ApSapVendor

        return db.query(func.count(ApSapVendor.id)).scalar() or 0

    return _cached("sap_vendors", build)


def _sv_out(r) -> dict:
    return {"vendor": r.vendor, "name": r.name, "vat": r.vat, "tax1": r.tax1, "city": r.city, "country": r.country, "blocked": bool(r.blocked),
            "payment_terms": r.payment_terms, "payment_days": terms_days(r.payment_terms), "account_group": r.account_group or ""}  # 1.0.103


# 1.0.103: SAP's vendor general data shipped in the code (Edwin's Supplier Master - Gen Sec - EFM: vendor, name, VAT
# number, tax number 1, city, country, blocked, account group); gzipped when it is big
SAP_VENDORS_CSV = DATA / "sap_vendors.csv"
SAP_VENDOR_HEADS = ["Vendor", "Name", "VAT Registration No.", "Tax Number 1", "City", "Country", "Blocked", "Account group"]


def sap_vendors_seed_path() -> Optional[Path]:
    for p in (SAP_VENDORS_CSV, SAP_VENDORS_CSV.with_name(SAP_VENDORS_CSV.name + ".gz")):
        if p.is_file():
            return p
    return None


def read_csv_any(path: Path) -> list[list[str]]:
    """A seed file, plain or gzipped (read as utf-8 text), closed at once."""
    if path.name.endswith(".gz"):
        import gzip

        with gzip.open(path, "rt", encoding="utf-8", newline="") as fh:
            return list(csv.reader(fh))
    return _read_csv(path)


def seed_sap_vendors() -> list[dict]:
    """The shipped vendor list as rows of {vendor, name, vat, tax1, city, country, blocked, account_group} (vat already
    normalised the way the 2.10.7 reader does it)."""
    p = sap_vendors_seed_path()
    rows = read_csv_any(p) if p else []
    out = []
    for r in rows[1:]:
        if len(r) >= 8 and r[0]:
            out.append({"vendor": r[0], "name": r[1], "vat": r[2], "tax1": r[3], "city": r[4], "country": r[5], "blocked": r[6].strip().upper() in ("X", "1", "TRUE", "YES"),
                        "account_group": r[7]})
    return out


def sap_vendor(db, vendor: Any) -> Optional[dict]:
    """A vendor number in SAP's vendor general data (leading zeros as SAP may give them). None: not in it (or no list)."""
    from .models.ap import ApSapVendor

    vs = vendor_variants(vendor)
    if not vs or not _sv_count(db):
        return None
    r = db.query(ApSapVendor).filter(ApSapVendor.vendor.in_(sorted(vs))).order_by(ApSapVendor.id).first()
    return _sv_out(r) if r is not None else None


def sap_vendors_by_vat(db, vat: Any, limit: int = 20) -> list[dict]:
    """Every SAP vendor with this VAT number (a supplier may have more than one vendor number). 1.0.103: `limit` — one VAT
    number can belong to many vendors (4380190845: 54 "Excellerate Facilities …" vendors; normal in SAP)."""
    from .models.ap import ApSapVendor

    v = sap_vat(vat)
    if len(v) < 6 or not _sv_count(db):
        return []
    return [_sv_out(r) for r in db.query(ApSapVendor).filter(ApSapVendor.vat == v).order_by(ApSapVendor.blocked, ApSapVendor.vendor).limit(max(1, int(limit))).all()]


def sap_vendors_by_name(db, name: Any, name_key: Callable[[Any], str]) -> list[dict]:
    """SAP vendors whose name (squeezed the way CBX compares names) is this one's."""
    from .models.ap import ApSapVendor

    k = name_key(name)
    if len(k) < 4 or not _sv_count(db):
        return []
    return [_sv_out(r) for r in db.query(ApSapVendor).filter(ApSapVendor.name_key == k).order_by(ApSapVendor.blocked, ApSapVendor.vendor).limit(20).all()]


def vendor_blocked(db, vendor: Any) -> Optional[bool]:
    """True when SAP has the vendor blocked or marked for deletion; None when CBX can't tell (no list, or not in it)."""
    hit = sap_vendor(db, vendor)
    return None if hit is None else hit["blocked"]


# ---- SAP words --------------------------------------------------------------------------------------------------------------
TCODES = (
    # posting and display (FI)
    "FB01", "FB01L", "FB02", "FB03", "FB08", "FB50", "FB60", "FB65", "FB70", "FB75", "FBV0", "FBV2", "FBV3", "FBRA", "FB1K", "FB1D", "FB1S",
    "F-02", "F-03", "F-04", "F-28", "F-32", "F-43", "F-44", "F-47", "F-48", "F-53", "F-58", "F110", "F111", "F.13", "FBL1N", "FBL3N", "FBL5N",
    "FAGLL03", "FAGLB03", "FS10N", "FS00", "FK01", "FK02", "FK03", "FK05", "FK08", "FD01", "FD02", "FD03", "XK01", "XK02", "XK03", "XK05",
    "XD01", "XD02", "XD03", "OB52", "FBZP", "FCH1", "FCHN", "F-90", "AS01", "AS02", "AS03", "AW01N",
    # logistics, purchasing, sales
    "MIRO", "MIGO", "MIR4", "MIR6", "MIR7", "MRBR", "MR8M", "ME21N", "ME22N", "ME23N", "ME51N", "ME52N", "ME53N", "ME54N", "ME2N", "ME2L",
    "ME2M", "ME29N", "ME28", "ML81N", "MM01", "MM02", "MM03", "MK01", "MK02", "MK03", "VA01", "VA02", "VA03", "VF01", "VF02", "VF03", "VF04",
    "VL01N", "VL02N", "VL03N",
    # controlling, tables, admin
    "KSB1", "KSB5", "KS01", "KS02", "KS03", "KE53", "KO88", "CJ20N", "CJI3", "SE16N", "SE16", "SE11", "SM35", "SM37", "SU01", "SU53", "SBWP",
)
SCREEN = (
    r"SAP\s+Easy\s+Access", r"SAP\s+GUI", r"SAP\s+Logon", r"SAP\s+Fiori", r"Fiori\s+Launchpad",
    r"Document\s+\d{6,10}\s+was\s+posted\s+in\s+company\s+code", r"was\s+posted\s+in\s+company\s+code",
    r"Balance\s+in\s+transaction\s+currency", r"Message\s+no\.?\s*[A-Z0-9_/]{1,20}\s?\d{3}\b",
    # 2.10.3: a message id ("F5 702") only with "message" / "msg" just before it, or in SAP's own "(F5 702)" — "office F2 105"
    # and "lift M3 101" are rooms and lifts, not SAP
    r"\b(?:message|msg)\b[^\n]{0,20}?\b(?:F5|F6|F2|FS|M7|M8|M3|KI|KM|FF|RW|GLT2|FICORE)\s?\d{3}\b",
    r"\((?:F5|F6|F2|FS|M7|M8|M3|KI|KM|FF|RW|GLT2|FICORE)\s?\d{3}\)",
    r"\bS_ALR_\d{6,}\w*", r"(?:transaction|t-?code)\s*:?\s*\"?BP\b", r"(?<=/[nNoO])BP\b", r"\bBP\s+transaction\b",
    r"Manage\s+Supplier\s+Line\s+Items", r"Manage\s+Journal\s+Entries", r"Post\s+General\s+Journal\s+Entries", r"Create\s+Supplier\s+Invoice",
    r"Supplier\s+Invoices\s+List", r"Display\s+Supplier\s+Balances", r"Manage\s+Purchase\s+Orders", r"Approve\s+Purchase\s+Orders",
    r"Manage\s+Supplier\s+Master\s+Data", r"Display\s+Line\s+Items\s+in\s+General\s+Ledger", r"(?-i:\bMy Inbox\b)",
    r"Enter\s+Vendor\s+Invoice", r"Enter\s+Incoming\s+Invoice", r"Display\s+Document:\s+(?:Data\s+Entry\s+View|Line\s+Item)", r"Vendor\s+Line\s+Item\s+Display",
    r"G/L\s+Account\s+Line\s+Item\s+Display", r"Posting\s+period\s+\d{1,3}\s+\d{4}\s+is\s+not\s+open", r"Company\s+code\s+[A-Z0-9]{4}\s+is\s+not\s+defined",
)
_SAP: Optional[re.Pattern] = None


def sap_words() -> re.Pattern:
    """SAP transaction codes (FB60, MIRO, ME21N, F-53, S_ALR_…, "/nFB03") and SAP screen words ("SAP Easy Access",
    "Document … was posted in company code", "Balance in transaction currency", message ids like "F5 702" or
    "Message no. M8081", Fiori apps). Precompiled, case-insensitive; `.search(text)` says whether a text is about SAP."""
    global _SAP
    if _SAP is None:
        codes = "|".join(re.escape(c) for c in sorted(_long_codes(), key=len, reverse=True))
        tcode = rf"(?:(?<=/[nNoO])|(?<![A-Za-z0-9_\-.]))(?:{codes})(?![A-Za-z0-9_\-])"
        _SAP = re.compile("|".join([tcode, _short_rx(), *SCREEN]), re.I)
    return _SAP


# 2.10.3: the short codes (F-03, F-53, F.13) are also unit and room numbers ("unit F-03"): they count only with
# transaction / t-code / SAP / posting next to them, or typed as SAP's "/nF-53"
SHORT = re.compile(r"^F[-.]\d\d$")
NEAR = r"\b(?:transactions?|t-?codes?|sap|posting|posted|post)\b"


def _long_codes() -> list[str]:
    return [c for c in TCODES if not SHORT.match(c)]


def _short_rx(group: bool = False) -> str:
    codes = "|".join(re.escape(c) for c in TCODES if SHORT.match(c))
    code = f"({codes})" if group else f"(?:{codes})"
    tail = r"(?![A-Za-z0-9_\-])"
    return (rf"(?<=/[nNoO]){code}{tail}|{NEAR}[^\n]{{0,25}}?(?<![A-Za-z0-9_\-.]){code}{tail}"
            rf"|(?<![A-Za-z0-9_\-.]){code}{tail}[^\n]{{0,25}}?{NEAR}")


def tcodes_in(s: str) -> list[str]:
    """The SAP transaction codes in a text, upper case, in the order they first appear (each once)."""
    rx = re.compile(rf"(?:(?<=/[nNoO])|(?<![A-Za-z0-9_\-.]))({'|'.join(re.escape(c) for c in sorted(_long_codes(), key=len, reverse=True))}|S_ALR_\d{{6,}}\w*)(?![A-Za-z0-9_\-])", re.I)
    found = [(m.start(1), m.group(1).upper()) for m in rx.finditer(s or "")]
    short = re.compile(_short_rx(group=True), re.I)  # 2.10.3: F-53 only where the words say it is a transaction
    for m in short.finditer(s or ""):
        i = next((k for k in range(1, (m.lastindex or 0) + 1) if m.group(k)), 0)
        if i:
            found.append((m.start(i), m.group(i).upper()))
    out = [c for _, c in sorted(found)]
    if re.search(r"(?:transaction|t-?code)\s*:?\s*\"?BP\b|(?<=/[nNoO])BP\b|\bBP\s+transaction\b", s or "", re.I):
        out.append("BP")  # the business partner transaction, only where the words say it is one ("BP number" is not)
    return list(dict.fromkeys(out))
