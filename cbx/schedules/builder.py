"""Fill the Site Schedule Agreement Word template. Placeholders look like {{site}}; a default template
is generated on first use if the shipped one is missing, so the tool always works."""
from __future__ import annotations

import io
import re
from datetime import date
from pathlib import Path

from .parser import FIELDS, MONEY_FIELDS, fmt_money
from ..util import human_date

TEMPLATE_DIR = Path(__file__).resolve().parent / "templates"
DEFAULT_TEMPLATE = TEMPLATE_DIR / "site_schedule_default.docx"
PLACEHOLDER = re.compile(r"\{\{\s*([a-z_]+)\s*\}\}")


def ensure_default_template() -> Path:
    if DEFAULT_TEMPLATE.exists():
        return DEFAULT_TEMPLATE
    import docx
    from docx.shared import Pt

    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    d.add_heading("SITE SCHEDULE AGREEMENT", level=0)
    d.add_paragraph("Schedule to the Master Site Access and Co-location Agreement")
    d.add_paragraph("Site: {{site}}    Province: {{province}}")
    d.add_paragraph("Swiftnet quotation reference: {{quote_no}}    Date: {{date}}")
    d.add_heading("1. Site particulars", level=1)
    t = d.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in (("Site name", "{{site}}"), ("Province", "{{province}}"), ("Quotation number", "{{quote_no}}"), ("Equipment accommodation", "{{accommodation_desc}}"), ("Antennas", "{{antennas}}")):
        r = t.add_row().cells
        r[0].text, r[1].text = k, v
    d.add_heading("2. Fees", level=1)
    t2 = d.add_table(rows=1, cols=3)
    t2.style = "Table Grid"
    h = t2.rows[0].cells
    h[0].text, h[1].text, h[2].text = "Item", "Per month (excl. VAT)", "Once-off (excl. VAT)"
    for k, m, o in (("Accommodation", "{{accommodation}}", ""), ("Electricity", "{{electricity}}", ""), ("Access road", "{{access_road}}", ""), ("Installation / establishment", "", "{{once_off}}"), ("Total", "{{total_monthly}}", "{{total_once_off}}")):
        r = t2.add_row().cells
        r[0].text, r[1].text, r[2].text = k, m, o
    d.add_heading("3. Term", level=1)
    d.add_paragraph("This Site Schedule commences on the Commencement Date and continues for the Initial Term set out in the Master Agreement, subject to the terms of that agreement. All amounts exclude VAT and escalate annually as provided in the Master Agreement.")
    d.add_heading("4. Signatures", level=1)
    d.add_paragraph("For CBRE Excellerate: ______________________    Date: __________")
    d.add_paragraph("For Swiftnet: ______________________    Date: __________")
    d.add_paragraph("Prepared by CBX-Online on {{date}} from quotation {{quote_no}}.")
    d.save(DEFAULT_TEMPLATE)
    return DEFAULT_TEMPLATE


def context_from(fields: dict, site: str, quote_no: str) -> dict:
    ctx = {"site": site or "", "quote_no": quote_no or "", "date": human_date(date.today())}
    for key, _label in FIELDS:
        f = fields.get(key) or {}
        v = f.get("value")
        if key in MONEY_FIELDS:
            ctx[key] = fmt_money(v) if v is not None else "—"
            if v == 0 and f.get("note"):
                ctx[key] = f["note"]
        elif key not in ctx:
            ctx[key] = "" if v is None else str(v)
    ctx["accommodation_desc"] = (fields.get("accommodation") or {}).get("quote", "").split(" R")[0].strip() or "As per quotation"
    return ctx


def _replace_in_paragraph(p, ctx: dict) -> None:
    if "{{" not in p.text:
        return
    new = PLACEHOLDER.sub(lambda m: str(ctx.get(m.group(1), m.group(0))), p.text)
    if p.runs:
        p.runs[0].text = new
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(new)


def build_docx(fields: dict, site: str, quote_no: str, template: Path | None = None) -> bytes:
    import docx

    tpl = template or ensure_default_template()
    d = docx.Document(str(tpl))
    ctx = context_from(fields, site, quote_no)
    for p in d.paragraphs:
        _replace_in_paragraph(p, ctx)
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _replace_in_paragraph(p, ctx)
    for s in d.sections:
        for p in list(s.header.paragraphs) + list(s.footer.paragraphs):
            _replace_in_paragraph(p, ctx)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
