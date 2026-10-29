"""The Site Schedule Agreement template: the wording, as blocks an admin edits in the app (Site Schedules →
Template), and the Word document made from it.

A block is one paragraph-like thing on the page:
  para     text with {tokens}; **bold**, __underline__; align left|center|justify; indent 0–3; an optional
           number ("6.2.1.") with a tab; an optional "right" text at the middle tab stop (two columns);
           size (pt), space_before / space_after (pt); "if": a token — the block only shows when it has a value
  heading  number + text, bold, no line under it (Edwin, 9 Oct)
  table    one of the charge tables t1–t4 (or "extra": the extra tables), with its total label; "if" works too
  witnesses / signatories   the signing people of one side ("swift" or "couser")
  pagebreak
{token|fallback} prints the fallback when the token is empty, e.g. {swiftnet_rep|____________________}.
"""
from __future__ import annotations

import base64
import copy
import io
import re
from typing import Any, Optional

from .reader import fix_refs, fmt_zar, num

TOKENS: list[tuple[str, str]] = [
    ("lease_type", "Lease type"), ("lease_code", "Lease code (N, R, RA…)"), ("company", "Co-user company"), ("co_code", "Co-user code (VDC…)"),
    ("col_id", "Co-locate ID"), ("soc_no", "Swiftnet SOC No"), ("contract_no", "Contract No"), ("t100_ref", "T100 reference"), ("t000_ref", "T000 reference"),
    ("swiftnet_rep", "Swiftnet representative"), ("couser_rep", "Co-user representative"),
    ("site_name", "Site name"), ("stac", "STAC code"), ("pid", "Site PID"), ("region", "Region"), ("lat", "Latitude"), ("lon", "Longitude"),
    ("struct_name", "Structure name"), ("struct_cat", "Structure category"), ("site_owner", "Site owner"),
    ("co_site_name", "Co-user site name"), ("bs_site_no", "BS / site no"), ("co_region", "Co-user region"),
    ("term", "Term (5 Years)"), ("start_date", "Commencement date"), ("end_date", "Expiry date"), ("esc_date", "First escalation (01 April 2027)"),
    ("esc_freq", "Escalation frequency"), ("mast_cat", "Mast category"), ("mast_hagl", "Mast height (m)"), ("inst_hagl", "Installation height (m)"),
    ("scope", "Scope of approved work (SOW / NTP)"), ("equip_comments", "Equipment comments"),
    ("current_install", "Client current installation on site"),
    ("sign_city_swift", "Swiftnet: signed at"), ("sign_date_swift", "Swiftnet: signed on"), ("sign_city_couser", "Co-user: signed at"),
    ("sign_date_couser", "Co-user: signed on"), ("today", "Today"), ("t1_total", "6.1 total"), ("t2_total", "6.2 total"), ("t3_total", "6.3 total"), ("t4_total", "6.4 total"),
]
TABLES = {"t1": "Rental charges", "t2": "Electricity charges", "t3": "Escort / road charges", "t4": "Once-off charges", "extra": "The extra tables"}
COLUMNS = ["Leased Item", "Requirement", "Comment/Equipment Description", "Amount (excl VAT)"]
TOKEN_RE = re.compile(r"\{([a-z0-9_]+)(?:\|([^}]*))?\}")

LINE = "\u00a0" * 24  # underlined in the template: an empty signing line, as Word shows it
DOTS = "…………………………………"


def _p(text: str = "", **o) -> dict:
    return {"type": "para", "text": text, **o}


def _h(number: str, text: str) -> dict:
    return {"type": "heading", "number": number, "text": text}


# 1.0.107: Swiftnet signs at one place, by the same three people, on every agreement — an admin setting kept with the
# template (Site Schedules → Template → The page), so a change on Swiftnet's board needs no code
SWIFT_SIGN: dict = {"city": "Pretoria", "people": [{"name": "Eddy Kgomongwe", "title": "CCO: Swiftnet"}, {"name": "Gert Vorster", "title": "CFO: Swiftnet"},
                                                   {"name": "Paul Jansen Van Rensburg", "title": "CEO: Swiftnet"}]}

DEFAULT: dict = {
    "font": "Arial", "size": 11, "margin_cm": 2.54, "swift_sign": copy.deepcopy(SWIFT_SIGN),
    "footer": "{lease_type} Site Schedule - Site Name: {site_name}   Contract no: {contract_no}   Co-user: {company}",
    "blocks": [
        _p("SITE SCHEDULE AGREEMENT", align="center", bold=True, size=14),
        _p("{lease_type}", align="center", bold=True, size=12, space_after=10),
        _p("Entered into by and between", align="center", space_after=8),
        _p("**__Swiftnet (Pty) Ltd,__** herein referred to as “Swiftnet”, represented by __{swiftnet_rep|" + LINE + "}__, who by his/her signature hereto "
           "warrants that he/she is duly authorised hereto to enter into this Site Schedule Agreement;", align="justify"),
        _p("and", align="center"),
        _p("**__{company},__** herein referred to as the “Co-User”, represented by __{couser_rep|" + LINE + "}__ who by his/her signature hereto warrants "
           "that he/she is duly authorised hereto to enter into this Site Schedule Agreement.", align="justify"),
        _p("Jointly the “Parties”", align="center"),
        _p("The Parties have entered into a Master Site Lease Agreement for Mast Space (hereinafter referred to as “the Master Site Lease Agreement” for "
           "purposes of this Site Schedule Agreement) and the terms and conditions of such agreement shall apply to the lease of the Designated Area "
           "(as defined in the Master Site Lease Agreement) described herein.", align="justify"),
        _h("1.", "SWIFTNET SITE DETAILS:"),
        _p("Site name: **{site_name}**", indent=1, right="STAC Code: **{stac}**"),
        _p("Site PID: **{pid}**", indent=1, right="Region: **{region}**"),
        _p("Latitude: **{lat}**", indent=1, right="Longitude: **{lon}**"),
        _p("Co-locate ID: **{col_id}**", indent=1),
        _p("Swiftnet (Pty) Ltd: **{soc_no}**", indent=1),
        _p("Contract no.: **{contract_no|New Application}**", indent=1),
        _p("**{t100_ref}**", indent=2, **{"if": "t100_ref"}),
        _p("**{t000_ref}**", indent=2, **{"if": "t000_ref"}),
        _h("2.", "CO-USER SITE DETAILS:"),
        _p("Site Name: **{co_site_name}**", indent=1, right="BS/Site. No: **{bs_site_no}**"),
        _p("Region: **{co_region}**", indent=1),
        _h("3.", "SITE SCHEDULE AGREEMENT DURATION:"),
        _p("Term: **{term}**", indent=1, number="3.1."),
        _p("Commencement Date: **{start_date}**", indent=1, number="3.2."),
        _p("Expiry Date: **{end_date}**", indent=1, number="3.3."),
        _h("4.", "ESCALATION:"),
        _p("Rental escalation, as contemplated in the Site Lease Agreement, shall apply in respect of Rental Charges.", indent=1, align="justify"),
        _p("Escalation Frequency: **{esc_freq|Annually}**", indent=1, number="4.1."),  # 1.0.107: the values in bold
        _p("Rate: **CPI**", indent=1, number="4.2."),
        _p("First Escalation: **{esc_date}**", indent=1, number="4.3."),
        _h("5.", "SITE/INFRASTRUCTURE INFORMATION:"),
        _p("Mast Category: **{mast_cat}**", indent=1, number="5.1."),
        _p("Mast Height Above Ground Level (HAGL): **{mast_hagl}**", indent=1, number="5.2."),
        _p("Installation Height on Mast (HAGL): **{inst_hagl}**", indent=1, number="5.3."),
        _h("6.", "RENTALS AND OTHER CHARGES:"),
        _p("**Rental Charges**", indent=1, number="6.1."),
        _p("In respect of rental of the Designated Area (as fully indicated in Appendix 1 of this Site Schedule Agreement and attached quotation), the "
           "following rentals shall be payable by The Co-User:", indent=1, align="justify"),
        {"type": "table", "table": "t1", "total": "Total Rental Charges Payable (Excl. VAT) per Month"},
        {"type": "table", "table": "extra", "total": "Sub-Total"},
        _p("**Other Charges**", indent=1, number="6.2.", space_before=8),
        _p("Other Charges, for electricity and other services or products provided to the Co-User Swiftnet Premises shall be as follows:", indent=1, align="justify"),
        {"type": "table", "table": "t2", "total": "Total Electricity Charges Payable (Excl. VAT) per Month"},
        _p("**Note:**", indent=1),
        _p("Charges reflected above are as at date of signature hereto and Swiftnet shall be entitled to amend the above electricity charges as and when "
           "contemplated in the Electricity Agreement.", indent=2, number="6.2.1.", align="justify"),
        _p("The following further charges shall be payable by the Co-User:", indent=1, number="6.3", space_before=6),
        {"type": "table", "table": "t3", "total": "Total Other Charges Payable (Excl. VAT) per Month"},
        _p("Charges reflected above are as at date of signature hereto and Swiftnet shall be entitled to amend the above charges from time to time upon "
           "written notice to this effect to the Co-User.", indent=2, number="6.3.1", align="justify"),
        _p("Escort fees shall be payable for both an hourly and a kilometre rate, as determined by the relevant Swiftnet charge-out rates from time to "
           "time, provided that such fees shall be based on the necessary and reasonable costs associated with such activity.", indent=2, number="6.3.2", align="justify"),
        _p("Any increase in respect of Access Road Usage fees, as contemplated in note 1 above, shall be subject to an annual review such that the "
           "adjusted rate pursuant to such annual review shall not exceed a rate equal to the Consumer Price Index (“CPI”) as published by Statistics "
           "South Africa (P0141 CPI Headline, or replacement, as per the website of Statistics South Africa), calculated over the twelve (12) months prior "
           "to the first (1st) day of May of the previous year.", indent=2, number="6.3.3", align="justify"),
        _p("The following further charges shall be payable by the Co-user:", indent=1, number="6.4", space_before=6, **{"if": "t4"}),
        _p("Application fees/Structural fees shall be payable for both as determined by the relevant Swiftnet charge-out rates from time to time.",
           indent=2, align="justify", **{"if": "t4"}),
        {"type": "table", "table": "t4", "total": "Total Other Charges Payable (Excl. VAT)", "if": "t4"},
        _h("7.", "SCOPE OF APPROVED WORK:"),
        _p("**{lease_type}:**", indent=1),
        # 1.0.108: a sub-heading only when there is text under it (the Swiftnet team, 21 Oct: no empty headings)
        _p("**Client current installation on site:**", indent=1, **{"if": "current_install"}),
        _p("{current_install}", indent=1, align="justify", **{"if": "current_install"}),
        _p("**Scope of work:**", indent=1, **{"if": "scope"}),
        _p("{scope}", indent=1, align="justify", **{"if": "scope"}),
        _p("**Equipment accommodation:**", indent=1, **{"if": "equip_comments"}),
        _p("{equip_comments}", indent=1, align="justify", **{"if": "equip_comments"}),
        _p("**Antenna and feeder work:**", indent=1, **{"if": "scope"}),
        _p("{scope}", indent=1, align="justify", **{"if": "scope"}),
        _h("8.", "SPECIAL REQUESTS/CONDITIONS"),
        _p("For the processing of all Swiftnet applications in terms of the Master Site Lease Agreement. Email :(Colocation@swiftnet.co.za).", indent=1, number="8.1.", align="justify"),
        _p("Any person working on Masts on a Swiftnet Site shall be duly certified and in possession of a valid mast climbing certificate and must wear the "
           "appropriate Personal Protective Equipment (PPE) whilst climbing or working on the above-mentioned structures at all times. Proof of the "
           "aforementioned certificate shall be submitted to the Swiftnet site manager or his authorised representative before any work on Masts may "
           "commence.", indent=1, number="8.2.", align="justify"),
        _p("The Co-User shall be liable for any costs to provide feeder/cable access (and related ducts) between the Co-Users accommodation and the "
           "Swiftnet equipment room or frames for services required from Swiftnet's network. The Co-User shall have a Swiftnet's Requirement for "
           "Occupation (RFO) specification for the above-mentioned in hand before any build work may commence on site.", indent=1, number="8.3.", align="justify"),
        _p("The approved installation and/or work shall be completed within 6 (six) months from date of authorisation to commence with installation, or as "
           "agreed between the Parties in writing. Non-compliance with the aforementioned may result in unavailability of approved facilities, space "
           "and/or electricity for which Swiftnet shall not be held liable.", indent=1, number="8.4.", align="justify"),
        _p("**Swiftnet reserves the right to amend the Rental and Other Charges as determined herein should it be found that the Co-User’s use of the "
           "station is different to the detail as depicted herein.**", number="9.", align="justify", space_before=12),
        {"type": "pagebreak"},
        _p("Signed at __{sign_city_swift|Pretoria}__ on {sign_date_swift|" + DOTS + "} 20", space_before=10),
        _p("As witness:"),
        {"type": "witnesses", "side": "swift"},
        {"type": "signatories", "side": "swift", "caption": "For Swiftnet duly authorized"},
        _p("Signed at __{sign_city_couser|" + DOTS + "}__ on __{sign_date_couser|" + DOTS + "}__ 20", space_before=16),
        _p("As witness:"),
        {"type": "witnesses", "side": "couser"},
        {"type": "signatories", "side": "couser", "caption": "For the Co-User duly authorized"},
    ],
}

BLOCK_TYPES = ("para", "heading", "table", "witnesses", "signatories", "pagebreak")

# 1.0.107: standard words that changed; a saved template that still has the old words gets the new ones (one an admin
# rewrote keeps its own)
UPGRADES = {
    "Escalation Frequency: {esc_freq|Annually}": "Escalation Frequency: **{esc_freq|Annually}**",
    "Rate: CPI": "Rate: **CPI**",
    "First Escalation: {esc_date}": "First Escalation: **{esc_date}**",
    "Signed at __{sign_city_swift|Centurion}__ on {sign_date_swift|" + DOTS + "} 20": "Signed at __{sign_city_swift|Pretoria}__ on {sign_date_swift|" + DOTS + "} 20",
}


def swift_sign(tpl: Optional[dict]) -> dict:
    """Where Swiftnet signs and who signs for it (the template's setting, else the standard three)."""
    x = (tpl or {}).get("swift_sign") if isinstance(tpl, dict) else None
    if not isinstance(x, dict):
        return copy.deepcopy(SWIFT_SIGN)
    ppl = [{"name": str(p.get("name") or "").strip()[:120], "title": str(p.get("title") or "").strip()[:160]}
           for p in (x.get("people") or []) if isinstance(p, dict) and str(p.get("name") or "").strip()][:6]
    return {"city": str(x.get("city") or "").strip()[:80], "people": ppl}


# 1.0.108: clause 7's sub-headings and the field under each; a saved template whose sub-heading has no "only when" gets it
SUBHEADS = {"**Client current installation on site:**": "current_install", "**Scope of work:**": "scope",
            "**Equipment accommodation:**": "equip_comments", "**Antenna and feeder work:**": "scope"}


def _subheads(blocks: list) -> list:
    """# 1.0.108: in a template saved before: each standard clause 7 sub-heading shows only when its text does. Only the standard
    shape is touched — the heading followed by its {field} line (or, for "Client current installation on site", by
    nothing of its own: its line is added); a heading an admin gave other words under, or an "only when", stays theirs."""
    out: list = []
    for i, b in enumerate(blocks):
        out.append(b)
        if not (isinstance(b, dict) and b.get("type") == "para" and b.get("text") in SUBHEADS and "if" not in b):
            continue
        key = SUBHEADS[b["text"]]
        nxt = blocks[i + 1] if i + 1 < len(blocks) else {}
        nt = str(nxt.get("text") or "") if isinstance(nxt, dict) else ""
        if nt == "{" + key + "}":
            b["if"] = key
        elif key == "current_install" and (not nt or nt in SUBHEADS or (isinstance(nxt, dict) and nxt.get("type") != "para")):
            b["if"] = key
            out.append(_p("{current_install}", indent=b.get("indent", 1), align="justify", **{"if": key}))
    return out


def upgrade(tpl: dict) -> dict:
    """The template as it is used now: old standard words replaced, the Swiftnet signing setting filled in."""
    t = copy.deepcopy(tpl or {})
    for b in t.get("blocks") or []:
        if isinstance(b, dict) and b.get("type") == "para" and b.get("text") in UPGRADES:
            b["text"] = UPGRADES[b["text"]]
    if isinstance(t.get("blocks"), list):
        t["blocks"] = _subheads(t["blocks"])
    t["swift_sign"] = swift_sign(t)
    return t


def check(tpl: dict) -> list[str]:
    """What is wrong with a template an admin saved (empty list = fine)."""
    probs: list[str] = []
    if not isinstance(tpl, dict) or not isinstance(tpl.get("blocks"), list) or not tpl["blocks"]:
        return ["The template has no blocks."]
    if len(tpl["blocks"]) > 400:
        probs.append("At most 400 blocks.")
    known = {t for t, _ in TOKENS}
    for i, b in enumerate(tpl["blocks"], 1):
        if not isinstance(b, dict) or b.get("type") not in BLOCK_TYPES:
            probs.append(f"Block {i}: unknown kind.")
            continue
        for key in ("text", "right", "number"):
            for tk in TOKEN_RE.findall(str(b.get(key) or "")):
                if tk[0] not in known:
                    probs.append(f"Block {i}: {{{tk[0]}}} is not a field.")
        if b["type"] == "table" and b.get("table") not in TABLES:
            probs.append(f"Block {i}: pick one of the charge tables.")
        if b.get("if") and b["if"] not in known and b["if"] not in TABLES:
            probs.append(f"Block {i}: 'only when' needs a field or a table.")
    ss = tpl.get("swift_sign")
    if ss is not None and (not isinstance(ss, dict) or not isinstance(ss.get("people", []), list)):
        probs.append("Swiftnet signing: a place and a list of people.")
    elif isinstance(ss, dict) and len(ss.get("people") or []) > 6:
        probs.append("Swiftnet signing: at most 6 people.")
    return probs


# ---- values ----------------------------------------------------------------------------------------------
def values(ag: dict, today: str = "", tpl: Optional[dict] = None) -> dict:
    """Everything a token can print, from an agreement's fields and tables (and the template's Swiftnet signing place)."""
    f = fix_refs(ag.get("fields"))  # 1.0.108: the reference lines always "T100/<number>", "T000/<number>" (or left out)
    t = ag.get("tables") or {}
    out = {k: str(v).strip() for k, v in f.items() if v is not None}
    from datetime import date, timedelta

    from .reader import MONTHS, first_escalation_year, long_date, parse_long_date

    start = parse_long_date(out.get("start_date", ""))
    try:
        y, mo = int(f.get("term_years") or 0), int(f.get("term_months") or 0)
    except (TypeError, ValueError):
        y, mo = 5, 0
    parts = ([f"{y} Year{'s' if y != 1 else ''}"] if y else []) + ([f"{mo} Month{'s' if mo != 1 else ''}"] if mo else [])
    out["term"] = " and ".join(parts) or "0 Months"
    if start:
        total = start.month - 1 + y * 12 + mo
        first = date(start.year + total // 12, total % 12 + 1, 1)
        end = first + timedelta(days=start.day - 2)
        out["end_date"] = long_date(end.day, end.month, end.year)
        out.setdefault("esc_year", str(first_escalation_year(start)))  # 1.0.107: Jan–Mar starts escalate on 01 April that year
    esc = str(f.get("esc_year") or out.get("esc_year") or "")
    out["esc_date"] = f"01 April {esc}" if esc else ""
    out.setdefault("esc_freq", "Annually")
    out.setdefault("soc_no", "4320151972")
    for k in ("mast_hagl", "inst_hagl"):
        v = out.get(k, "")
        if v and not v.rstrip().lower().endswith("m"):
            out[k] = v + "m"
    out["today"] = today or long_date(date.today().day, date.today().month, date.today().year)
    for k in ("t1", "t2", "t3", "t4"):
        tot = sum(num(r.get("amt")) for r in (t.get(k) or []))
        if k == "t1":
            tot += sum(num(r[-1] if isinstance(r, list) else 0) for x in (t.get("extra") or []) for r in (x.get("rows") or []))
        out[f"{k}_total"] = fmt_zar(tot)
    sg = ag.get("signing") or {}
    for side in ("swift", "couser"):
        s = sg.get(side) or {}
        out[f"sign_city_{side}"] = str(s.get("city") or "").strip()
        out[f"sign_date_{side}"] = str(s.get("date") or "").strip()
    out["sign_city_swift"] = swift_sign(tpl if tpl is not None else DEFAULT)["city"]  # 1.0.107: the admin setting, not per agreement
    _ = MONTHS
    return out


def fill(text: str, vals: dict) -> str:
    return TOKEN_RE.sub(lambda m: vals.get(m.group(1)) or (m.group(2) or ""), text or "")


def shown(b: dict, vals: dict, tables: dict) -> bool:
    cond = b.get("if")
    if not cond:
        return True
    if cond in TABLES:
        return bool(tables.get(cond))
    return bool(vals.get(cond))


def runs(text: str) -> list[tuple[str, bool, bool]]:
    """'**a** __b__ c' → [(text, bold, underline)]."""
    out: list[tuple[str, bool, bool]] = []
    bold = under = False
    for part in re.split(r"(\*\*|__)", text):
        if part == "**":
            bold = not bold
        elif part == "__":
            under = not under
        elif part:
            out.append((part, bold, under))
    return out


# ---- Word ------------------------------------------------------------------------------------------------
INDENT_TWIPS = [0, 426, 720, 1080]


def to_docx(tpl: dict, ag: dict, people: Optional[dict] = None, today: str = "") -> bytes:
    import docx
    from docx.enum.section import WD_SECTION  # noqa: F401
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT  # noqa: F401
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, Twips

    tpl = upgrade(tpl)
    vals = values(ag, today, tpl)
    tables = ag.get("tables") or {}
    font, size = tpl.get("font") or "Arial", float(tpl.get("size") or 11)
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = font, Pt(size)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.space_before = Pt(0)
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    m = Cm(float(tpl.get("margin_cm") or 2.54))
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = m
    align = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "right": WD_ALIGN_PARAGRAPH.RIGHT}
    text_width = sec.page_width - sec.left_margin - sec.right_margin

    def add_runs(p, text, *, bold=False, size_pt=None):
        for t, b, u in runs(text):
            r = p.add_run(t)
            r.bold = True if (bold or b) else None
            r.underline = True if u else None
            if size_pt:
                r.font.size = Pt(size_pt)

    def para(b):
        p = d.add_paragraph()
        pf = p.paragraph_format
        ind = INDENT_TWIPS[max(0, min(3, int(b.get("indent") or 0)))]
        number = fill(b.get("number") or "", vals)
        if number:
            pf.left_indent = Twips(ind + 720)
            pf.first_line_indent = Twips(-720)
            pf.tab_stops.add_tab_stop(Twips(ind + 720))
        else:
            pf.left_indent = Twips(ind)
        pf.alignment = align.get(b.get("align") or "left")
        if b.get("space_before") is not None:
            pf.space_before = Pt(float(b["space_before"]))
        if b.get("space_after") is not None:
            pf.space_after = Pt(float(b["space_after"]))
        bold, size_pt = bool(b.get("bold")), b.get("size")
        if number:
            add_runs(p, number + "\t", bold=bold or b.get("type") == "heading", size_pt=size_pt)
        add_runs(p, fill(b.get("text") or "", vals), bold=bold or b.get("type") == "heading", size_pt=size_pt)
        if b.get("right"):
            pf.tab_stops.add_tab_stop(Twips(5040))
            p.add_run("\t")
            add_runs(p, fill(b["right"], vals), bold=bold, size_pt=size_pt)
        return p

    def borders(table):
        tbl = table._tbl
        pr = tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), "000000")
            b.append(e)
        pr.append(b)
        ind = OxmlElement("w:tblInd")
        ind.set(qn("w:w"), "426")
        ind.set(qn("w:type"), "dxa")
        pr.append(ind)

    def cell(c, text, bold=False):
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        add_runs(p, text, bold=bold, size_pt=size - 1)

    def charge_table(cols, rows, total_label, total):
        widths = [0.18, 0.18, 0.42, 0.22] if len(cols) == 4 else [1 / len(cols)] * len(cols)
        tw = text_width - Twips(426)
        t = d.add_table(rows=1, cols=len(cols))
        t.autofit = False
        borders(t)
        for i, h in enumerate(cols):
            cell(t.rows[0].cells[i], h, bold=True)
        hdr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        hdr.append(th)
        body = rows or [["N/A"] * (len(cols) - 1) + ["R 0.00"]]
        for r in body:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cell(cells[i], str(v if v not in (None, "") else ("" if i == len(cols) - 1 else "N/A")))
        cells = t.add_row().cells
        merged = cells[0].merge(cells[len(cols) - 2]) if len(cols) > 1 else cells[0]
        cell(merged, total_label, bold=True)
        cell(cells[len(cols) - 1], total, bold=True)
        ws = [int(tw * widths[min(i, len(widths) - 1)]) for i in range(len(cols))]
        for i, col in enumerate(t.columns):  # the grid Word and LibreOffice lay the table out by
            col.width = ws[i]
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        t._tbl.tblPr.append(lay)
        for row in t.rows:
            for i, c in enumerate(row.cells):
                c.width = ws[min(i, len(ws) - 1)]
        d.add_paragraph().paragraph_format.space_after = Pt(4)

    sigs = people or {}
    sg = ag.get("signing") or {}

    def sig_of(name):
        for side in ("swift", "couser"):
            for p in sigs.get(side) or []:
                if (p.get("name") or "").strip().lower() == (name or "").strip().lower() and p.get("sig"):
                    return p["sig"]
        return None

    def side_people(side):
        """1.0.107: Swiftnet's signatories are the template's fixed list; the co-user's are picked per agreement."""
        if side == "swift":
            return swift_sign(tpl)["people"]
        return (sg.get(side) or {}).get("people") or []

    def witness_names(side):
        w = [x for x in ((sg.get(side) or {}).get("witnesses") or []) if str(x).strip()]
        return (w + ["", ""])[:max(2, len(w))]

    def sign_lines(cell_or_doc, person, caption, first):
        """One person who signs: the line (or drawn signature), the name, the title, the caption — as the real agreement."""
        add = cell_or_doc.add_paragraph
        img = sig_of(person.get("name"))
        p = add()
        p.paragraph_format.space_before = Pt(18 if first else 26)
        p.paragraph_format.space_after = Pt(0)
        if img:
            try:
                raw = base64.b64decode(img.split(",", 1)[1])
                p.add_run().add_picture(io.BytesIO(raw), width=Cm(3.7))
            except Exception:  # noqa: BLE001 — a broken signature prints the dotted line instead
                p.add_run(DOTS)
        else:
            p.add_run(DOTS)
        for text, bold in ((person.get("name") or "", True), (person.get("title") or "", True), (caption or "", False)):
            if text:
                q = add()
                q.paragraph_format.space_after = Pt(0)
                add_runs(q, f"**{text}**" if bold else text)
        return p

    def signing_table(side, people_, caption):
        """Witnesses on the left, the people who sign on the right, side by side (Edwin, 12 Oct: like the real agreement)."""
        t = d.add_table(rows=1, cols=2)
        t.autofit = False
        left, right = t.rows[0].cells
        lw = Twips(4320)
        left.width, right.width = lw, text_width - lw
        left.paragraphs[0].paragraph_format.space_after = Pt(0)
        right.paragraphs[0].paragraph_format.space_after = Pt(0)
        for i, w in enumerate(witness_names(side)):
            p = left.add_paragraph()
            p.paragraph_format.space_before = Pt(18 if i == 0 else 26)
            p.paragraph_format.space_after = Pt(0)
            add_runs(p, f"__{w}__" if w else DOTS)
        for i, person in enumerate(people_ or [{"name": "", "title": ""}]):
            sign_lines(right, person, caption, i == 0)
        d.add_paragraph().paragraph_format.space_after = Pt(4)

    blocks = [b for b in (tpl.get("blocks") or []) if shown(b, vals, tables)]
    skip = set()
    for bi, b in enumerate(blocks):
        if bi in skip:
            continue
        kind = b.get("type")
        nxt = blocks[bi + 1] if bi + 1 < len(blocks) else {}
        if kind == "witnesses" and nxt.get("type") == "signatories" and nxt.get("side") == b.get("side"):
            signing_table(b.get("side"), side_people(b.get("side")), nxt.get("caption") or "")
            skip.add(bi + 1)
            continue
        if kind in ("para", "heading"):
            if kind == "heading":
                b = {**b, "bold": True, "space_before": b.get("space_before", 12), "space_after": b.get("space_after", 6)}
            para(b)
        elif kind == "pagebreak":
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == "table":
            k = b.get("table")
            if k == "extra":
                for x in tables.get("extra") or []:
                    if x.get("title"):
                        para({"text": "**" + x["title"] + "**", "indent": 1})
                    rows = x.get("rows") or []
                    tot = fmt_zar(sum(num(r[-1]) for r in rows if r))
                    charge_table(x.get("cols") or COLUMNS, rows, b.get("total") or "Sub-Total", tot)
            else:
                rows = [[r.get("item", ""), r.get("req", ""), r.get("desc", ""), r.get("amt", "")] for r in tables.get(k) or []]
                charge_table(COLUMNS, rows, fill(b.get("total") or "", vals), vals.get(f"{k}_total", "R 0.00"))
        elif kind == "witnesses":
            w = [x for x in ((sg.get(b.get("side")) or {}).get("witnesses") or []) if str(x).strip()]
            w = (w + ["", ""])[:max(2, len(w))]
            for i in range(0, len(w), 2):
                p = d.add_paragraph()
                p.paragraph_format.space_before = Pt(18)
                p.paragraph_format.tab_stops.add_tab_stop(Twips(5040))
                add_runs(p, (f"__{w[i]}__" if w[i] else DOTS) + "\t" + ((f"__{w[i + 1]}__" if w[i + 1] else DOTS) if i + 1 < len(w) else ""))
        elif kind == "signatories":
            # on their own (no witnesses just above): still in the right-hand column, as the real agreement
            ppl = side_people(b.get("side")) or [{"name": "", "title": ""}]
            for i, person in enumerate(ppl):
                sign_lines(d, person, b.get("caption") or "", i == 0)
            # indent what was just written into the right-hand column
            n = sum(1 + bool(x.get("name")) + bool(x.get("title")) + bool(b.get("caption")) for x in ppl)
            for q in d.paragraphs[-n:]:
                q.paragraph_format.left_indent = Twips(4320)
    if tpl.get("footer"):
        fp = sec.footer.paragraphs[0]
        fp.paragraph_format.tab_stops.add_tab_stop(text_width, alignment=2)  # right
        for t, bo, u in runs(fill(tpl["footer"], vals)):
            r = fp.add_run(t)
            r.font.size = Pt(8)
        r = fp.add_run("\tPage ")
        r.font.size = Pt(8)
        _field(fp, "PAGE")
        r = fp.add_run(" of ")
        r.font.size = Pt(8)
        _field(fp, "NUMPAGES")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _field(p, code: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    for kind, text in (("begin", None), (None, code), ("separate", None), (None, "1"), ("end", None)):
        r = p.add_run()
        r.font.size = Pt(8)
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        elif text == code:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {code} "
            r._r.append(it)
        else:
            r.text = text


def default() -> dict:
    return copy.deepcopy(DEFAULT)


def sample_agreement() -> dict:
    """Made-up values to show a template with (Site Schedules → Template)."""
    return {"fields": {"lease_type": "Renewal", "company": "Vodacom (Pty) Ltd", "col_id": "7708", "contract_no": "70204797", "t100_ref": "T100/9130",
                       "t000_ref": "T000/5183", "site_name": "SAMPLE EXCHANGE", "stac": "SMP", "pid": "5431", "region": "Western", "lat": "33 55 37",
                       "lon": "18 35 19", "co_site_name": "Sample Vale", "bs_site_no": "2981", "co_region": "Western", "start_date": "01 August 2026",
                       "term_years": 5, "term_months": 0, "mast_cat": "Sectional Mast", "mast_hagl": "35", "inst_hagl": "0.00",
                       "scope": "Install three panel antennas at 36m on the north leg.", "equip_comments": "Container plinth inside the yard.",
                       "current_install": ""},  # 1.0.108: text under two sub-headings, so the editor shows them
            "tables": {"t1": [{"item": "Ground Space", "req": "25.00m²", "desc": "Ground space for the 35m mast inside the yard", "amt": "R 2 485.25"},
                              {"item": "Ground Space", "req": "12.00m²", "desc": "Container plinth inside the yard", "amt": "R 1 192.92"}],
                       "t2": [{"item": "Electricity", "req": "40 Amp", "desc": "3 Phase – AC Metered", "amt": "R 0.00"},
                              {"item": "Standby Electricity", "req": "40 Amp", "desc": "3 Phase", "amt": "R 5 615.70"}],
                       "t3": [{"item": "Escort Fees", "req": "N/A", "desc": "N/A", "amt": "R 0.00"}, {"item": "Access Road Usage", "req": "0.00 Km", "desc": "N/A", "amt": "R 0.00"}],
                       "t4": [], "extra": []},
            "signing": {"swift": {"witnesses": [], "people": []},
                        "couser": {"witnesses": [], "people": [{"name": "Co-user signatory", "title": "Authorised Representative"}]}}}
