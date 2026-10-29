# 1.0.104: new — Builder R, V06 (VALUENCE-AP-GAPS.md §2 V06; SPEC-1104 Phase 2)
"""VAT validity decides the tax code.

A supplier invoice with VAT on it is checked against s20 of the VAT Act (the portfolio's list `vat_requirements`): the
words "Tax Invoice", the supplier's name and VAT number (10 digits starting with 4), the recipient's name — and, above
`vat_recipient_threshold_r` (R5,000), the recipient's VAT number —, the invoice number, the date, a description and the
VAT amount. A utility bill counts only when its profile is listed in `vat_number_profiles` (those bills print the
municipality's VAT number: a bill without it is not valid). The result is kept on the document (`vat_valid`,
`vat_reasons`).

When it is not valid, the portfolio's `vat_invalid` decides (Edwin, 19 Oct: red until a person decides):
    red        Tess shows red "not a valid tax invoice: <reasons>" until a reviewer (never someone who touched the
               document) picks Claim VAT or Don't claim, with a reason (audited ap_vat_decision);
    no_claim   posted without the claim straight away (Shell's rule);
    off        not checked (outside SA VAT, e.g. Egypt).
Don't claim posts with the portfolio's `no_claim_tax_code` (blank = the Fixed TaxCodeCredit, I0) on the expense lines,
the VAT folded into them and no VAT (JA600) line; the vendor line stays the invoice's total. V05 picks every such
document for review.

Pip calls `apply` (capture.build_one); the Review view calls `decide`; the report lists the suppliers whose invoices went
out without the claim (to chase them for valid tax invoices). A credit note, an invoice with no VAT and one not in rand
are not checked (vat_valid None)."""
from __future__ import annotations

import io
import re
from datetime import datetime
from decimal import Decimal
from typing import Optional

from sqlalchemy.orm import Session

from ..audit import log
from ..db import utcnow
from ..models import User
from ..models.ap import ApBatch, ApCapture, ApItem
from . import portfolio as PF
from . import roles as R
from . import sap as P

# the review's own actions: never a "touch" when the same reviewer looks again (roles.touched_by exclude_actions)
OWN_ACTIONS = ("ap_capture_reviewed", "ap_vat_decision")
DECISIONS = {"claim": "Claim VAT", "no_claim": "Don't claim"}


class Refused(PF.Refused):
    """A plain-words refusal (400)."""


# ---- s20: what a tax invoice shows ------------------------------------------------------------------------------------------
# key · the words in Settings (portfolio vat_requirements) · what is said when it is missing
REQUIREMENTS = (
    ("words", "Tax Invoice words", 'the words "Tax Invoice" are not on it'),
    ("supplier", "Supplier name", "no supplier name"),
    ("supplier_vat", "Supplier VAT number", "no supplier VAT number"),
    ("recipient", "Recipient name", "no recipient (our company's) name"),
    ("recipient_vat", "Recipient VAT number", "no recipient VAT number"),
    ("invoice_no", "Invoice number", "no invoice number"),
    ("date", "Invoice date", "no invoice date"),
    ("description", "Description", "no description of the goods or services"),
    ("vat", "VAT amount", "no VAT amount or rate"),
)


def _squash(s) -> str:
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


def wanted(db: Session, portfolio) -> set[str]:
    """The checks the portfolio asks for (its `vat_requirements` lines, matched loosely to the keys above)."""
    lines = PF.setting(db, portfolio, "vat_requirements", [r[1] for r in REQUIREMENTS]) or []
    said = {_squash(x) for x in lines if isinstance(x, str)}
    return {k for k, words, _ in REQUIREMENTS if _squash(words) in said or _squash(k) in said}


def vat_no_ok(v) -> bool:
    """A South African VAT number: 10 digits starting with 4."""
    return bool(re.fullmatch(r"4\d{9}", re.sub(r"\D", "", str(v or ""))))


def _vat_total(f: dict) -> Optional[Decimal]:
    v = P.dec(f.get("vat"))
    if v is None:
        lv = [P.dec(x.get("vat")) for x in f.get("lines") or [] if isinstance(x, dict) and P.dec(x.get("vat")) is not None]
        v = sum(lv, Decimal(0)) if lv else None
    return v


def check(f: dict, *, wants: set[str], threshold: float = 5000) -> tuple[Optional[bool], list[str]]:
    """A manual invoice's fields → (valid, reasons). None = not checked: a credit note, no VAT on it, not in rand."""
    if f.get("credit") or str(f.get("currency") or "ZAR").upper() not in ("ZAR", "R", ""):
        return None, []
    vat = _vat_total(f)
    rated = any((P.dec(x.get("vat_rate")) or 0) > 0 for x in f.get("lines") or [] if isinstance(x, dict))
    if (vat is None or vat == 0) and not (vat is None and rated):
        return None, []  # nothing to claim
    words = {k: w for k, _, w in REQUIREMENTS}
    out = []
    if "words" in wants and f.get("tax_invoice_words") is not True:
        out.append(words["words"] if f.get("tax_invoice_words") is False else 'the words "Tax Invoice" are not confirmed')
    if "supplier" in wants and not str(f.get("supplier") or "").strip():
        out.append(words["supplier"])
    if "supplier_vat" in wants and not vat_no_ok(f.get("supplier_vat")):
        out.append(words["supplier_vat"] if not re.sub(r"\D", "", str(f.get("supplier_vat") or "")) else f"supplier VAT number {f.get('supplier_vat')} is not 10 digits starting with 4")
    if "recipient" in wants and not str(f.get("bill_to") or "").strip():
        out.append(words["recipient"])
    incl = P.dec(f.get("incl"))
    if "recipient_vat" in wants and incl is not None and abs(incl) > Decimal(str(threshold or 0)) and not vat_no_ok(f.get("bill_to_vat")):
        out.append(f"{words['recipient_vat']} (needed above R {float(threshold or 0):,.0f})".replace(",", " "))
    if "invoice_no" in wants and not str(f.get("invoice_no") or "").strip():
        out.append(words["invoice_no"])
    if "date" in wants and not str(f.get("invoice_date") or "").strip():
        out.append(words["date"])
    if "description" in wants and not any(str(x.get("description") or "").strip() for x in f.get("lines") or [] if isinstance(x, dict)):
        out.append(words["description"])
    if "vat" in wants and vat is None:
        out.append(words["vat"])
    return (not out), out


VAT_NO_KEY = re.compile(r"(^|_)(vat|tax)_?(reg_?)?(no|nr|number|num)$", re.I)


def _bill_vat_no(c: ApCapture) -> str:
    """The municipality's VAT number as the profile read it (a vat_no / vat_number field on the bill's rows)."""
    f, rd = c.fields or {}, c.read or {}
    for src in (f, rd.get("elec_row") or {}):
        for k, v in (src or {}).items():
            if isinstance(k, str) and VAT_NO_KEY.search(k) and not k.startswith(("customer", "bill_to", "our")) and vat_no_ok(v):
                return re.sub(r"\D", "", str(v))
    return ""


def check_utility(db: Session, c: ApCapture, portfolio) -> tuple[Optional[bool], list[str]]:
    """A utility bill: checked only when its profile is one whose bills print the municipality's VAT number."""
    listed = {_squash(x) for x in PF.setting(db, portfolio, "vat_number_profiles", []) or [] if isinstance(x, str) and x.strip()}
    f = c.fields or {}
    names = {_squash(x) for x in (c.profile, f.get("profile"), (c.read or {}).get("profile"), (c.read or {}).get("profile_name"), f.get("municipality")) if x}
    if not listed or not (names & listed):
        return None, []
    vat = sum(((P.dec(s.get("vat")) or Decimal(0)) for s in f.get("services") or [] if isinstance(s, dict)), Decimal(0))
    if not vat:
        return None, []
    return (True, []) if _bill_vat_no(c) else (False, ["the bill does not show the municipality's VAT number"])


# ---- Pip's hook ---------------------------------------------------------------------------------------------------------------
def no_claim_code(db: Session, portfolio, tables: dict) -> str:
    return (str(PF.setting(db, portfolio, "no_claim_tax_code", "") or "").strip().upper()
            or str(tables["fixed"].get("TaxCodeCredit") or "I0").strip().upper())


def fold(rows: list[dict], tables: dict, code: str) -> list[dict]:
    """Don't claim: the VAT line(s) go, the VAT is spread over the expense lines in proportion (to the cent: the last line
    takes the rounding), and every expense line takes the no-claim tax code. The vendor line stays the invoice's total."""
    if not rows:
        return rows
    vat_gl = str(tables["gl"].get("vat") or "").strip()

    def signed(r) -> Decimal:
        a = P.dec(r.get("amount")) or Decimal(0)
        return a if r.get("pkey") == "40" else -a

    is_vat = [r.get("pkey") in ("40", "50") and (str(r.get("dt") or "") == "X" or (vat_gl and r.get("account") == vat_gl)) for r in rows]
    lines = [i for i, r in enumerate(rows) if i and r.get("pkey") in ("40", "50") and not is_vat[i]]
    # the lines that carried VAT (not those on the zero-rated / exempt code: interest, deposits, a 0 % line) take it back
    zero = {str(tables["fixed"].get(k) or "").strip().upper() for k in ("TaxCodeCredit", "TaxCodeExempt")} - {""}
    exp = [i for i in lines if str(rows[i].get("tc") or "").strip().upper() not in zero] or lines
    vat = sum((signed(r) for i, r in enumerate(rows) if is_vat[i]), Decimal(0))
    if not exp:
        return rows
    net = sum((signed(rows[i]) for i in exp), Decimal(0))
    out = [dict(r) for r in rows]
    new = {}
    if vat and net:
        done = Decimal(0)
        for n, i in enumerate(exp):
            if n == len(exp) - 1:
                new[i] = net + vat - done
            else:
                new[i] = (signed(rows[i]) * (net + vat) / net).quantize(Decimal("0.01"))
                done += new[i]
    for i in exp:
        o = out[i]
        if i in new:
            o["pkey"], o["amount"] = ("40" if new[i] >= 0 else "50"), P.m2(abs(new[i]))
        o["tc"] = code
    return [r for i, r in enumerate(out) if not is_vat[i] or not vat]


def apply(db: Session, c: ApCapture, rows: list[dict], st: str, issues: list[str], tables: dict) -> tuple[list[dict], str, list[str]]:
    """Pip's hook (capture.build_one, after the rows are built): is it a valid tax invoice, and what follows when it is not."""
    portfolio = c.portfolio or None
    mode = str(PF.setting(db, portfolio, "vat_invalid", "red") or "red")
    if mode == "off" or c.bucket not in ("manual", "utility") or not rows:
        c.vat_valid, c.vat_reasons = None, []
        return rows, st, issues
    if c.bucket == "manual":
        valid, reasons = check(c.fields or {}, wants=wanted(db, portfolio), threshold=PF.setting(db, portfolio, "vat_recipient_threshold_r", 5000) or 0)
    else:
        valid, reasons = check_utility(db, c, portfolio)
    c.vat_valid, c.vat_reasons = valid, reasons
    if valid is not False:
        return rows, st, issues
    said = "; ".join(reasons)
    decision = c.vat_decision or ("no_claim" if mode == "no_claim" else "")
    if decision == "claim":
        who = _name(db, c.vat_decided_by_id)
        return rows, (P.YLW if st == P.GRN else st), issues + [f"VAT claimed on an invoice that is not a valid tax invoice ({said}): decided by {who}"]
    if decision == "no_claim":
        code = no_claim_code(db, portfolio, tables)
        rows = fold(rows, tables, code)
        why = "decided by " + _name(db, c.vat_decided_by_id) if c.vat_decision else "the portfolio posts without the claim"
        return rows, (P.YLW if st == P.GRN else st), issues + [f"VAT not claimable: posted with {code} ({said}; {why})"]
    return rows, P.RED, issues + [f"not a valid tax invoice: {said} - a reviewer decides: claim VAT or don't claim"]


def _name(db: Session, uid: Optional[int]) -> str:
    u = db.get(User, uid) if uid else None
    return (u.name or u.email) if u else "a person"


# ---- a person's decision -----------------------------------------------------------------------------------------------------
def decide(db: Session, user: User, c: ApCapture, decision: str, reason: str) -> ApCapture:
    """Claim VAT or Don't claim, with a reason — a reviewer in the document's portfolio who did not touch it. Audited
    ap_vat_decision, a history line on the file; Pip builds it again (and Tess tests it)."""
    from . import capture as C

    decision = str(decision or "").strip().lower().replace(" ", "_").replace("'", "")
    if decision in ("dont_claim", "do_not_claim"):
        decision = "no_claim"
    if decision not in DECISIONS:
        raise Refused("Pick Claim VAT or Don't claim.")
    reason = re.sub(r"\s+", " ", str(reason or "")).strip()[:500]
    if len(reason) < 3:
        raise Refused("Say why (it is kept with the decision).")
    if c.uploaded_at:
        raise Refused("This invoice already went out in an SAP upload.")
    if c.vat_valid is not False:
        raise Refused("This is a valid tax invoice: there is nothing to decide.")
    pf = R.portfolio_of(db, c)
    R.require(db, user, "reviewer", pf, what="Deciding on the VAT")
    if user.id in R.touched_by(db, c, include_dropper=True, exclude_actions=OWN_ACTIONS):
        raise R.Forbidden("You dropped or worked on this document, so a different person decides on its VAT.")
    old = c.vat_decision or ""
    c.vat_decision, c.vat_decided_by_id, c.vat_decided_at, c.vat_decision_note = decision, user.id, utcnow(), reason
    it = db.get(ApItem, c.item_id)
    text = f"{DECISIONS[decision]} ({'; '.join(c.vat_reasons or [])}): {reason}" + (f" (was {DECISIONS.get(old, old)})" if old and old != decision else "")
    C._hist(it, user.name or user.email, "VAT: " + text)
    log(db, "ap_vat_decision", actor=user, target=f"ap_item:{c.item_id}", detail=f"{it.name if it else c.id}: {text}"[:400])
    if c.stage in ("test", "ready"):
        c.stage = "sap"  # Pip builds it again with the decision
    db.flush()
    b = db.get(ApBatch, c.batch_id)
    if b is not None:
        C.advance(db, b, live=False)
    return c


# ---- reading it ------------------------------------------------------------------------------------------------------------------
def out(db: Session, c: ApCapture) -> dict:
    return {"vat_valid": c.vat_valid, "vat_reasons": list(c.vat_reasons or []), "vat_decision": c.vat_decision or "",
            "vat_decided_by": _name(db, c.vat_decided_by_id) if c.vat_decided_by_id else "",
            "vat_decided_at": c.vat_decided_at.isoformat(timespec="minutes") if c.vat_decided_at else "", "vat_decision_note": c.vat_decision_note or ""}


def vat_of(c: ApCapture) -> float:
    f = c.fields or {}
    if c.bucket == "utility":
        return float(sum(((P.dec(s.get("vat")) or Decimal(0)) for s in f.get("services") or [] if isinstance(s, dict)), Decimal(0)))
    return float(_vat_total(f) or 0)


def report(db: Session, month: str = "", portfolio: str = "") -> bytes:
    """V06.5: the suppliers whose invoices went out without the VAT claim (to chase them for valid tax invoices): one line
    per supplier (invoices, VAT not claimed, the reasons), then every invoice."""
    import openpyxl

    from .. import xlsafe

    month = month_ok(month)
    q = db.query(ApCapture).filter(ApCapture.uploaded_at.isnot(None), ApCapture.vat_valid.is_(False))
    if portfolio:
        q = q.filter(ApCapture.portfolio == portfolio.upper())
    caps = [c for c in q.order_by(ApCapture.uploaded_at).all() if (not month or f"{c.uploaded_at:%Y-%m}" == month)
            and any(str(r.get("tc") or "") and r.get("pkey") in ("40", "50") for r in c.sap_rows or [])
            and not any(str(r.get("dt") or "") == "X" for r in c.sap_rows or [])]  # posted without a VAT line: not claimed
    items = {i.id: i for i in db.query(ApItem).filter(ApItem.id.in_([c.item_id for c in caps] or [0])).all()}
    by: dict[str, dict] = {}
    for c in caps:
        f = c.fields or {}
        name = f.get("supplier") or f.get("municipality") or c.vendor_no or "?"
        key = re.sub(r"\D", "", str(f.get("supplier_vat") or "")) or _squash(name)
        g = by.setdefault(key, {"supplier": name, "vat_no": f.get("supplier_vat") or "", "vendor": c.vendor_no, "n": 0, "vat": 0.0, "reasons": {}})
        g["n"] += 1
        g["vat"] += vat_of(c)
        for r in c.vat_reasons or []:
            g["reasons"][r] = g["reasons"].get(r, 0) + 1
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Suppliers"
    ws.append(["Supplier", "Supplier VAT no", "Vendor", "Invoices", "VAT not claimed (R)", "What their tax invoices miss"])
    for g in sorted(by.values(), key=lambda x: -x["vat"]):
        ws.append([g["supplier"], g["vat_no"], g["vendor"], g["n"], round(g["vat"], 2), "; ".join(f"{k} ({v})" for k, v in sorted(g["reasons"].items(), key=lambda x: -x[1]))])
    ws2 = wb.create_sheet("Invoices")
    ws2.append(["Portfolio", "Supplier", "Invoice no", "Invoice date", "Total (R)", "VAT not claimed (R)", "Why", "Decision", "Decided by", "Reason", "Uploaded", "File"])
    for c in caps:
        f = c.fields or {}
        it = items.get(c.item_id)
        ws2.append([c.portfolio or "", f.get("supplier") or f.get("municipality") or "", f.get("invoice_no") or "", f.get("invoice_date") or f.get("billing_date") or "",
                    float(P.dec(f.get("incl")) or 0) if c.bucket == "manual" else "", round(vat_of(c), 2), "; ".join(c.vat_reasons or []),
                    DECISIONS.get(c.vat_decision or "", "posted without the claim (portfolio rule)"), _name(db, c.vat_decided_by_id) if c.vat_decided_by_id else "",
                    c.vat_decision_note or "", f"{c.uploaded_at:%Y-%m-%d %H:%M}", it.name if it else ""])
    for sh in (ws, ws2):
        P._head_style(sh, sh.max_column)
        sh.freeze_panes = "A2"
        P._fit(sh, 60)
    xlsafe.book(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def month_ok(month: str) -> str:
    """'' or a yyyy-mm; anything else is refused in plain words."""
    m = str(month or "").strip()
    if not m:
        return ""
    try:
        datetime.strptime(m, "%Y-%m")
    except ValueError:
        raise Refused("A month as yyyy-mm.") from None
    return m
