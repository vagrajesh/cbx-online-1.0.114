# 1.0.105: new — Builder S (SPEC-1105): the AP Tool's one screen, "Make SAP upload"
"""Make SAP upload: drop invoices, one click, one result (Edwin, 19 Oct: "a simple, single click — it does everything and
gets me results").

A run is an AP batch named "Run 19 Oct, 14:05" (no portfolio question: each document's own, from its company code). The
whole pipeline runs as before — the AP team's batch reader (Paula, Ivy, Bo, Rae), Sol, Uma, the capture bots, Mia, Pip
(with Builder G's guesses), Tess, the 1.0.104 controls — and this module only says what the screen shows:

    while it runs   "Reading 7 of 20…" (a document is done once it is green, yellow or left out)
    the result      "20 invoices · 14 ready · 4 with guesses (yellow) · 2 left out", one Download button for every
                    document that may go out (green and yellow), the review workbook
    Left out        one plain line each, and the one thing a person can do right there: a VAT decision (V06), a second
                    person's check (V05) or the company an invoice is billed to (Builder G's "company unknown"), for people
                    allowed to; everyone else is told who it waits for (once the run has finished reading)
    Recent runs     date, count, "What was sent": the run's record — never a file SAP could take again (1.0.105 review:
                    a second copy of an upload is a double posting one click away)

The controls are those of 1.0.104: the poster role for the download, the reviewer role and separation of duties for a
check or a VAT decision, red never goes out, a document waiting for its review is left out of the file (never sent).
The old batch reader's findings are notes on the invoice's row (Tess's yellow), never a second list."""
from __future__ import annotations

import re
from datetime import datetime, timezone
from typing import Any, Optional

from sqlalchemy.orm import Session

from ..audit import log
from ..db import utcnow
from ..models import User
from ..models import AuditLog
from ..models.ap import ApBatch, ApCapture, ApEntity, ApItem, ApRole
from . import capture as C
from . import engine as E
from . import review as RV
from . import roles as R
from . import sap as P
from . import service as S

BUSY = C.READING + ("match", "sap", "test")
WORD = {"ready": "Ready", "guessed": "Guessed", "left": "Left out", "reading": "Reading", "sent": "Sent"}
RECENT = 8


class Refused(S.Refused):
    """A plain-words refusal (400)."""


def run_name(now: Optional[datetime] = None) -> str:
    """"Run 19 Oct, 14:05" — Johannesburg time."""
    t = (now or datetime.now(timezone.utc)).astimezone(P.SAST)
    return f"Run {t.day} {t:%b}, {t:%H:%M}"


def start(db: Session, user: User, files: list[tuple[str, bytes]]) -> ApBatch:
    """A run: the dropped files (PDFs or a zip) as one batch; its portfolio comes from each document (no question)."""
    b = S.create_batch(db, user, files, run_name())
    from . import portfolio as PF

    PF.set_batch(db, user, b, None, on_drop=True)
    return b


# ---- the guesses in the SAP file ----------------------------------------------------------------------------------------------
# Builder G's field → the SAP cell(s) it fills: (header, which rows) — "vendor" = the 31/21 line, "lines" = the 40/50 lines
GUESS_CELLS = {"company": (("comp", "vendor"),), "vendor": (("account", "vendor"),), "gl": (("account", "lines"),), "cost_centre": (("costcentre", "lines"),),
               "tax_code": (("tc", "lines"),), "due_date": (("DueDate", "vendor"),), "invoice_date": (("docdate", "vendor"),)}


def guesses_of(c: ApCapture) -> dict[str, str]:
    """{field: how} of what CBX guessed on this document (not what it was sure of). Builder G keeps `capture.guesses`
    as {field: how}; a list of fill() answers ({field, how, sure}) is read too."""
    g = getattr(c, "guesses", None)
    out: dict[str, str] = {}
    if isinstance(g, dict):
        for k, v in g.items():
            if isinstance(v, dict):
                if v.get("sure"):
                    continue
                v = v.get("how") or v.get("note") or ""
            if k and v:
                out[str(k)] = str(v)
    elif isinstance(g, list):
        for x in g:
            if isinstance(x, dict) and x.get("field") and not x.get("sure"):
                out[str(x["field"])] = str(x.get("how") or "guessed")
    return out


def cell_marks(c: ApCapture, r: dict) -> dict:
    """capture.download's hook: {"_guess": {SAP header: how}} for the cells of row r that CBX guessed — Builder G's
    autofill.row_guesses (SPEC-1105-interface.md) when it is there, else the same from `capture.guesses`."""
    try:
        from . import autofill as AF

        row_guesses = getattr(AF, "row_guesses", None)
    except ImportError:
        row_guesses = None
    if row_guesses is not None and isinstance(getattr(c, "match", None), dict) and c.match.get("autofill"):
        marks = {h: str(how) for h, how in (row_guesses(c, r) or {}).items() if h in P.SAP_HEADERS and how}
        return {"_guess": marks} if marks else {}
    g = guesses_of(c)
    if not g:
        return {}
    pk = str(r.get("pkey") or "")
    kind = "vendor" if pk in ("31", "21") else "lines" if pk in ("40", "50") and str(r.get("dt") or "") != "X" else ""
    marks = {}
    for field, how in g.items():
        for header, rows in GUESS_CELLS.get(field, ()):
            if rows == kind:
                marks[header] = how
    return {"_guess": marks} if marks else {}


# ---- what the screen shows -------------------------------------------------------------------------------------------------
def _money(v: Any) -> str:
    d = P.dec(v)
    return "" if d is None else f"R\u00a0{abs(d):,.2f}".replace(",", "\u00a0")  # no break inside an amount on a phone


def _red_text(c: ApCapture) -> str:
    reds = [str(x.get("text") or "") for x in c.sap_issues or [] if isinstance(x, dict) and x.get("level") == "red"]
    return reds[0] if reds else next((str(x.get("text") or "") for x in c.sap_issues or [] if isinstance(x, dict)), "")


NEEDS_WORDS = {"entity": "the account is not in the utilities list (an admin adds it under More)", "vendor": "the vendor needs its SAP coding",
               "profile": "PDF Studio needs a profile for this bill", "odd": "not an invoice CBX can place",
               "unread": "CBX could not read it"}


def reviewers_for(db: Session, c: ApCapture, limit: int = 3) -> list[str]:
    """Up to `limit` people who may check this document: the reviewer role in its portfolio, never someone who touched it
    or dropped its folder."""
    pf = c.portfolio or R.portfolio_of(db, c)
    out_ids = R.touched_by(db, c, include_dropper=True, exclude_actions=RV.OWN_ACTIONS)
    if R.has_rows(db, pf):
        ids = [uid for (uid,) in db.query(ApRole.user_id).filter(ApRole.role == "reviewer", ApRole.portfolio.in_((pf, ""))).distinct().all()]
        cands = db.query(User).filter(User.id.in_(ids or [0]), User.is_active.is_(True)).order_by(User.name).all()
    else:  # bootstrap: everyone who writes in the AP Tool reviews
        from ..models import Tool, ToolAccess

        tool = db.query(Tool).filter(Tool.slug == R.SLUG).one_or_none()
        ids = [uid for (uid,) in db.query(ToolAccess.user_id).filter(ToolAccess.tool_id == tool.id).limit(2000).all()] if tool else []
        cands = db.query(User).filter(User.id.in_(ids or [0]), User.is_active.is_(True), User.is_axl.is_(False)).order_by(User.name).limit(max(60, limit * 20)).all()
    names = []
    for u in cands:
        if u.id in out_ids or not R.may(db, u, "reviewer", pf):
            continue
        names.append(u.name or u.email)
        if len(names) >= limit:
            break
    return names


def _candidates(db: Session, pf: str, cache: dict) -> list[User]:
    """The people who may review in the portfolio, by name: its reviewer rows (or "" rows), or — bootstrap, no role rows yet —
    everyone who writes in the AP Tool (and the admins). One or two reads per portfolio."""
    if pf in cache:
        return cache[pf]
    active = (User.is_active.is_(True), User.is_axl.is_(False), User.role != "viewer")
    if R.has_rows(db, pf):
        ids = [uid for (uid,) in db.query(ApRole.user_id).filter(ApRole.role == "reviewer", ApRole.portfolio.in_((pf, ""))).distinct().all()]
        users = db.query(User).filter(User.id.in_(ids or [0]), *active).order_by(User.name).limit(200).all()
    else:
        if "" not in cache:  # the bootstrap list is the same for every portfolio
            from sqlalchemy import or_

            from ..models import Tool, ToolAccess

            tool = db.query(Tool).filter(Tool.slug == R.SLUG).one_or_none()
            has = db.query(ToolAccess.user_id).filter(ToolAccess.tool_id == (tool.id if tool else 0))
            cache[""] = db.query(User).filter(or_(User.id.in_(has), User.role.in_(("admin", "owner"))), *active).order_by(User.name).limit(200).all()
        users = cache[""]
    cache[pf] = users
    return users


def reviewer_names(db: Session, caps: list, limit: int = 3) -> dict[int, list[str]]:
    """1.0.105 review (perf): who may check each document, for many documents in a few reads — the reviewers of its
    portfolio, never someone who touched it, dropped its folder or coded its vendor (the same rule as review._can_many)."""
    if not caps:
        return {}
    pfs = {c.id: (c.portfolio or R.portfolio_of(db, c)) for c in caps}
    cache: dict = {}
    skip = set(R.NOT_A_TOUCH) | set(RV.OWN_ACTIONS)
    target = {f"ap_item:{c.item_id}": c.id for c in caps}
    touched: dict[int, set] = {c.id: set() for c in caps}
    keys = list(target)
    for i in range(0, len(keys), 500):
        for t, uid, a in db.query(AuditLog.target, AuditLog.actor_id, AuditLog.action).filter(AuditLog.target.in_(keys[i:i + 500]), AuditLog.actor_id.isnot(None),
                                                                                              AuditLog.action.like("ap\\_%", escape="\\")).all():
            if a not in skip:
                touched[target[t]].add(uid)
    droppers = dict(db.query(ApBatch.id, ApBatch.user_id).filter(ApBatch.id.in_({c.batch_id for c in caps})).all())
    coders: dict[str, set] = {}
    vends = sorted({c.vendor_no for c in caps if c.vendor_no})
    for t, uid in db.query(AuditLog.target, AuditLog.actor_id).filter(AuditLog.action == "ap_vendor_learned", AuditLog.actor_id.isnot(None),
                                                                      AuditLog.target.in_([f"ap_vendor:{v}" for v in vends] or ["-"])).all():
        coders.setdefault(t.split(":", 1)[1], set()).add(uid)
    out = {}
    for c in caps:
        no = touched[c.id] | {droppers.get(c.batch_id)} | coders.get(c.vendor_no or "", set())
        out[c.id] = [u.name or u.email for u in _candidates(db, pfs[c.id], cache) if u.id not in no][:limit]
    return out


def _who_checks(names: list[str]) -> str:
    if not names:
        return "waits for a reviewer to check it (an admin gives the role in AP Settings → AP people)"
    return "waits for " + (names[0] if len(names) == 1 else ", ".join(names[:-1]) + " or " + names[-1]) + " to check it"


def status(db: Session, user: User, b: ApBatch) -> dict:
    """The run as the screen shows it."""
    items = db.query(ApItem).filter(ApItem.batch_id == b.id).order_by(ApItem.id).all()
    caps = {c.item_id: c for c in db.query(ApCapture).filter(ApCapture.batch_id == b.id).all()}
    caps_by_id = {c.id: c for c in caps.values()}
    ai_on = C.model_on()
    cans = RV._can_many(db, user, [c for c in caps.values() if c.stage == "ready" and not c.uploaded_at and (c.vat_valid is False or c.review_state == "selected")])
    from . import portfolio as PF

    pf_cache: dict = {}

    def pf_of(c: ApCapture) -> str:
        return c.portfolio or R.portfolio_of(db, c)

    def waits(c: ApCapture) -> bool:  # review.waits, the portfolio's setting read once per portfolio
        if (c.review_state or "") != "selected":
            return False
        pf = pf_of(c)
        if ("wait", pf) not in pf_cache:
            pf_cache[("wait", pf)] = bool(PF.setting(db, pf or None, "review_before_upload", True))
        return pf_cache[("wait", pf)]

    def may_capture(c: ApCapture) -> bool:
        pf = pf_of(c)
        if ("cap", pf) not in pf_cache:
            pf_cache[("cap", pf)] = R.may(db, user, "capturer", pf or None)
        return pf_cache[("cap", pf)]

    docs, left, ai_off = [], [], 0
    named: dict[int, tuple] = {}  # capture id → (its line, "check" | "vat"): who it waits for, said once the reading is done
    total = len(items)
    for it in items:
        c = caps.get(it.id)
        f = (c.fields if c is not None else None) or {}
        rows = (c.sap_rows if c is not None else None) or []
        name = it.name.split("/")[-1]
        d = {"id": c.id if c is not None else None, "item_id": it.id, "file": name, "supplier": f.get("supplier") or f.get("municipality") or it.supplier or "",
             "amount": _money(rows[0].get("amount") if rows else (f.get("incl") or it.amount)), "word": "reading", "note": ""}
        if c is None:  # not sorted yet: reading
            docs.append(d)
            continue
        if c.uploaded_at:
            d["word"] = "sent"
            d["status"] = c.sap_status or ""
        elif c.stage in BUSY:
            if C.ai_waiting(c) or (not ai_on and c.stage in ("queued", "senior")):
                ai_off += 1
                d["word"], d["note"] = "left", "the AI is off: read when it is back"
                d["kind"] = "ai_off"
            # else: still reading
        elif c.bucket == "ap":
            d["word"], d["kind"] = "left", "not_invoice"
            d["note"] = f"not an invoice ({E.KINDS.get(it.kind, (it.kind,))[0].lower()}): the AP team's own checks"
        elif c.bucket == "skipped" or c.stage == "skipped":
            d["word"], d["kind"], d["note"] = "left", "skipped", "skipped: it stays out of the SAP upload"
        elif c.stage == "needs" and c.needs == "entity" and f.get("kind") == "manual":
            # 1.0.105 review (Builder G): no company could be guessed — a person picks it here (capture.decide "manual": Mia again)
            d["word"], d["kind"] = "left", "company"
            d["note"] = "company unknown: pick the company it is billed to" + (f" (billed to \"{str(f.get('bill_to'))[:80]}\")" if f.get("bill_to") else "")
            d["may"] = may_capture(c)
            if not d["may"]:
                d["waits"] = "waits for a capturer to pick the company"
        elif c.stage == "needs":
            d["word"], d["kind"] = "left", "needs"
            d["note"] = NEEDS_WORDS.get(c.needs, "needs a person") + (" (the AI is off)" if c.needs == "unread" and not ai_on else "")
            if c.needs == "vendor":
                d["note"] = (f"vendor {f['supplier']}" if f.get("supplier") else "the vendor") + " is not in SAP's vendor list"
        elif c.stage == "ready":
            if c.sap_status == "red":
                d["word"] = "left"
                text = _red_text(c)
                if c.vat_valid is False and not c.vat_decision:
                    d["kind"], d["note"] = "vat", "not a valid tax invoice: " + "; ".join(c.vat_reasons or []) if c.vat_reasons else "not a valid tax invoice"
                    ok, why = cans.get(c.id, (False, ""))
                    d["may"] = ok
                    if not ok:
                        named[c.id] = (d, "vat")
                elif RV.DUP_WORDS.search(text):
                    d["kind"], d["note"] = "duplicate", text
                elif "blocked in SAP" in text:
                    d["kind"], d["note"] = "blocked", text
                elif (c.review_state or "") == "rejected":
                    d["kind"], d["note"] = "rejected", text
                else:
                    d["kind"], d["note"] = "red", text or "Tess stopped it"
            elif waits(c):
                d["word"], d["kind"] = "left", "review"
                d["note"] = "waits for a second person's check" + (f" ({c.review_why})" if c.review_why else "")
                ok, _why = cans.get(c.id, (False, ""))
                d["may"] = ok
                if not ok:
                    named[c.id] = (d, "check")
            elif c.sap_status in ("green", "yellow"):
                d["word"], d["status"] = ("ready" if c.sap_status == "green" else "guessed"), c.sap_status
                g = guesses_of(c)
                if c.sap_status == "yellow":
                    d["note"] = "; ".join(f"{k.replace('_', ' ')}: {v}" for k, v in g.items()) or "; ".join(str(x.get("text") or "") for x in c.sap_issues or [] if isinstance(x, dict))[:300]
        docs.append(d)
        if d["word"] == "left" and d.get("kind") != "ai_off":
            left.append(d)
    reading = sum(1 for d in docs if d["word"] == "reading")
    phase = "failed" if b.status == "failed" else "reading" if (b.status == "reading" or reading) else "done"
    if phase == "done" and named:  # who each one waits for — worked out once, in bulk, when the screen shows the result
        names = reviewer_names(db, [caps_by_id[cid] for cid in named])
        for cid, (d, what) in named.items():
            said = _who_checks(names.get(cid) or [])
            d["waits"] = (said.replace(" to check it", " to decide on the VAT") if names.get(cid)
                          else "waits for a reviewer to decide on the VAT (an admin gives the role in AP Settings → AP people)") if what == "vat" else said
    sent = [d for d in docs if d["word"] == "sent"]
    out_now = [d for d in docs if d["word"] in ("ready", "guessed")]
    n_ready = sum(1 for d in docs if d["word"] == "ready")
    n_guess = sum(1 for d in docs if d["word"] == "guessed")
    n_left = sum(1 for d in docs if d["word"] == "left")
    parts = [f"{total} invoice{'s' if total != 1 else ''}"] + ([f"{len(sent)} sent"] if sent else [])  # what went out in an earlier download
    if n_ready or not sent:
        parts.append(f"{n_ready} ready")
    if n_guess:
        parts.append(f"{n_guess} with guesses (yellow)")
    if n_left:
        parts.append(f"{n_left} left out")
    label = (f"Download again ({len(out_now)} more)" if sent else "Download SAP upload") if out_now else ""
    users = {u.id: u for u in db.query(User).filter(User.id == b.user_id).all()}
    return {"run": {"id": b.id, "name": b.name, "at": b.created_at.isoformat(timespec="minutes") if b.created_at else "", "by": S._who(users.get(b.user_id)),
                    "files": total, "note": b.note if b.status == "failed" else ""},
            "phase": phase,
            "progress": {"done": total - reading, "total": total, "line": f"Reading {total - reading} of {total}…"},
            "summary": {"total": total, "ready": n_ready, "guessed": n_guess, "left_out": n_left, "sent": len(sent), "line": " · ".join(parts)},
            "download": {"n": len(out_now), "sent": len(sent), "label": label, "can": bool(out_now), "record": bool(sent)},
            "ai_off": ai_off, "ai_line": (f"The AI is off: {ai_off} invoice{'s' if ai_off != 1 else ''} will be read when it is back." if ai_off else ""),
            "left_out": left,
            "docs": [{k: d.get(k, "") for k in ("id", "file", "supplier", "amount", "word", "note")} | {"word": WORD.get(d["word"], d["word"])} for d in docs]}


def recent(db: Session, user: User, limit: int = RECENT) -> list[dict]:
    """The person's latest runs (every batch they dropped): date, how many, sent, whether a file can be fetched again."""
    from sqlalchemy import func

    bs = db.query(ApBatch).filter(ApBatch.user_id == user.id).order_by(ApBatch.id.desc()).limit(limit).all()
    ids = [b.id for b in bs] or [0]
    files = dict(db.query(ApItem.batch_id, func.count(ApItem.id)).filter(ApItem.batch_id.in_(ids)).group_by(ApItem.batch_id).all())
    sent = dict(db.query(ApCapture.batch_id, func.count(ApCapture.id)).filter(ApCapture.batch_id.in_(ids), ApCapture.uploaded_at.isnot(None)).group_by(ApCapture.batch_id).all())
    return [{"id": b.id, "name": b.name, "at": b.created_at.isoformat(timespec="minutes") if b.created_at else "", "files": files.get(b.id, 0), "sent": sent.get(b.id, 0),
             "reading": b.status == "reading"} for b in bs]


def download(db: Session, user: User, b: ApBatch) -> tuple[bytes, int]:
    """Every document that may go out — green and yellow, not sent before, not waiting for its review, never red."""
    return C.download(db, user, b, True)


SENT_SHEET = "ALREADY SENT - DO NOT UPLOAD"


def sent_record(db: Session, user: User, b: ApBatch) -> bytes:
    """Recent runs → "What was sent" (1.0.105 review: never a file SAP could take again): the run's review workbook (every
    row, its status and why, when it went out) behind a first sheet "ALREADY SENT - DO NOT UPLOAD" that lists each send —
    when, by whom, which invoices. SAP's uploader reads the first sheet, so this file can't post anything twice."""
    import io

    import openpyxl
    from openpyxl.styles import Font

    from .. import xlsafe

    caps = db.query(ApCapture).filter(ApCapture.batch_id == b.id, ApCapture.uploaded_at.isnot(None)).order_by(ApCapture.uploaded_at, ApCapture.item_id).all()
    if not caps:
        raise Refused("Nothing from this run went out yet.")
    items = {i.id: i for i in db.query(ApItem).filter(ApItem.batch_id == b.id).all()}
    users = {u.id: u for u in db.query(User).filter(User.id.in_({c.uploaded_by_id for c in caps if c.uploaded_by_id} or {0})).all()}
    wb = openpyxl.load_workbook(io.BytesIO(C.review(db, b)))
    ws = wb.create_sheet(SENT_SHEET, 0)
    ws.append([f"{b.name}: these invoices were sent to SAP already. Uploading them again posts them twice."])
    ws["A1"].font = Font(bold=True, color="C00000")
    ws.append([])
    ws.append(["Sent on", "Sent by", "Invoice no", "Supplier", "Amount", "Status", "File"])
    for c in caps:
        f, r0 = c.fields or {}, (c.sap_rows or [{}])[0]
        it = items.get(c.item_id)
        ws.append([f"{c.uploaded_at.astimezone(P.SAST) if c.uploaded_at.tzinfo else c.uploaded_at:%d %b %Y %H:%M}", S._who(users.get(c.uploaded_by_id)),
                   r0.get("refdocnr") or f.get("invoice_no") or "", f.get("supplier") or f.get("municipality") or "", float(P.dec(r0.get("amount")) or 0), c.sap_status or "",
                   it.name.split("/")[-1] if it else ""])
    for cell in ws[3]:
        cell.font = Font(bold=True)
    for col, w in zip("ABCDEFG", (18, 22, 18, 36, 12, 8, 36)):
        ws.column_dimensions[col].width = w
    wb.active = 0
    xlsafe.book(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def file_name(b: ApBatch, record: bool = False) -> str:
    base = re.sub(r"[^A-Za-z0-9]+", "_", b.name).strip("_")[:60] or f"run_{b.id}"
    return f"SAP_sent_{base}_record.xlsx" if record else f"SAP_upload_{base}.xlsx"


def companies(db: Session) -> dict:
    """The company picker (Left out → company unknown): the group's own entities first (the GENSEC portfolio's: EPS /
    Nervada Z999, Holdings Z102), then every other company in the list."""
    from . import portfolio as PF

    p = PF.get(db, "GENSEC")
    group = {x.strip().upper() for x in re.split(r"[,\s]+", (p.company_codes if p is not None else "") or "") if x.strip() and "-" not in x} or {"Z999", "Z102"}
    rows = db.query(ApEntity).filter(ApEntity.is_active.is_(True)).order_by(ApEntity.code).all()
    out = lambda e: {"code": e.code, "name": e.name or e.entity or ""}  # noqa: E731
    return {"group": [out(e) for e in rows if e.code.upper() in group], "other": [out(e) for e in rows if e.code.upper() not in group]}
