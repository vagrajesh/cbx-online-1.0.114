# 1.0.98: from CBX-Connect 2.10.3 cbx/ap/sap.py
"""The SAP upload (2.10.0): Pip builds the rows, Tess tests them, people download them.

Part 1 is the CBX Harvester's own SAP journal generator (pdf_harvester.py, "SAP journal generation (Edwin, July 2026)"),
copied as it is: SAP_HEADERS, the default Fixed values and GL Codes, generate_sap, export_sap_xlsx and the value check
the Harvester's SAP tab ran. It was verified line by line against finance's June 2025 pack (20 Eskom bills):
    31-line  = vendor, amount = elec incl VAT + posted adjustments
    SA110    = electricity excl VAT (canonical total_ex_vat_r)
    SA412    = (re/dis)connection work adjustments (no VAT split)
    SF101    = interest-on-overdue adjustments (no VAT split, tax code I0)
    JA600    = the bill's PRINTED VAT (never recomputed), dt = "X"
    text     = {CC}_{read start}_{read end}_Cons_{A|E}_Var_{kWh}
    header   = mapping NAME1 + " " + statement month
    credit bills → KG with pkeys 21/50, flagged yellow; the 40-lines must sum to the 31-line to the cent.
An electricity-only bill goes through generate_sap untouched.

Part 2 extends it (2.10.0, Edwin: "extend them to the other services the profiles give"): a bill with water, sanitation,
refuse, rates or sundry charges is one document — one 31-line for the whole bill, a 40-line per service to its GL from
the GL Codes table (a GL left blank there makes the document red, never a guess), the services' printed VAT on one JA600
line. Manual invoices: the 31-line is the invoice incl. VAT, a 40-line per GL and tax code (from the vendor's coding),
JA600 the printed VAT, credit notes as KG.

Part 3 is Tess: every document is tested before it can be downloaded. Part 4 writes the two files.

2.10.3 (SAP knowledge, cbx/knowledge.py): Tess also refuses a document type that doesn't allow the account types its
lines use (vendor lines need K, G/L lines need S), a vendor that is not opened in the posting company code and a cost
centre SAP doesn't have in that company code (or SAP's dummy); a G/L that is not in the chart of accounts is yellow (the
chart export may be incomplete: check with finance). A manual invoice's bill-to site code goes into the line text the
vendor coding says (or the header text); the review workbook names each G/L and cost centre.

1.0.102 (Edwin, 18 Oct): `postdate` is the batch's posting date (today unless a person set another; Tess is yellow when it
is far from today), not the document's date; a manual invoice's lines take their G/L (cost centre, tax code) from the
vendor coding's line rules, grouped by G/L, tax code and cost centre; a missing due date is the invoice date plus the
vendor's payment days; the header text is "<supplier short name> <MMM YY>" within SAP's 25 characters; a vendor SAP
has blocked is red.
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from typing import Optional

# =============================================================================================================
# Part 1 — the CBX Harvester's code, as written there
# =============================================================================================================
SAP_HEADERS = ["TransCode", "docdate", "postdate", "type", "comp", "curr",
               "refdocnr", "docheadertext", "pkey", "account", "sgl",
               "compn", "txi", "amount", "tc", "dt", "costcentre",
               "ordernr", "allocation", "text", "ProfitCentre", "DueDate"]

SAP_DEFAULT_FIXED = [("TransCode31", "FB01"), ("TransCode40", "ZB01"),
                     ("DocType", "KR"), ("Company", "T000"),
                     ("Currency", "ZAR"), ("TaxCode", "I9"),
                     ("TaxCodeCredit", "I0"), ("DocTypeCredit", "KG")]
SAP_DEFAULT_GL = [
    ("electricity", "SA110", "Electricity"),
    ("vat", "JA600", "VAT (bill's printed VAT)"),
    ("adjustment_work", "SA412", "Levies / connection-work adjustments"),
    ("adjustment_interest", "SF101", "Interest on overdue account"),
    ("adjustment_deposit", "DA201", "Sync-deposit credits (pkey 50)"),
    ("rates", "SA010", "Rates & Taxes (not generated yet)"),
    ("water", "SA210", "Water (not generated yet)"),
    ("refuse", "SA311", "Refuse (not generated yet)"),
    ("sanitation", "SA312", "Sanitation (not generated yet)"),
]


def _sap_norm_acct(v):
    s = re.sub(r"\D", "", str(v or ""))
    return s.lstrip("0")


def _sap_date(v):
    """yyyy-mm-dd / yyyy/mm/dd -> ddmmyyyy (SAP sheet format)."""
    m = re.match(r"(\d{4})[-/](\d{2})[-/](\d{2})", str(v or ""))
    return "%s%s%s" % (m.group(3), m.group(2), m.group(1)) if m else ""


def generate_sap(rows, tables):
    """rows: extracted rows (with _canon). Returns (sap_rows, issues,
    summary). Each sap_row is a dict over SAP_HEADERS plus _status
    (GRN/YLW/RED) and _issue."""
    fixed, gl, vendors = tables["fixed"], tables["gl"], tables["vendors"]
    out, issues = [], []
    n_ok = n_credit = n_unmapped = 0
    D_ = Decimal

    def dec(v):
        try:
            return D_(str(v).replace(",", "").replace(" ", ""))
        except Exception:
            return None

    for r in rows:
        if r.get("service") != "electricity" or r.get("_dup"):
            continue
        c = r.get("_canon") or {}
        acct = c.get("account_number") or r.get("account_number") or ""
        ex = dec(c.get("total_ex_vat_r"))
        vat = dec(r.get("elec_vat_r") or r.get("vat_out"))
        if ex is None:
            issues.append((r.get("file", ""), "no electricity total - not posted"))
            continue
        adj_i = dec(r.get("esk_adj_interest_r")) or D_(0)
        adj_w = dec(r.get("esk_adj_work_r")) or D_(0)
        adj_dep = dec(r.get("esk_adj_deposit_r")) or D_(0)
        if vat is None:
            vat = (ex * D_("0.15")).quantize(D_("0.01"))
            vat_flag = "VAT not printed on bill - calculated at 15%"
        else:
            vat_flag = ""
        total31 = ex + vat + adj_i + adj_w
        total31 = total31 + adj_dep
        credit = total31 < 0
        if credit:
            # Credit (negative) bills generate as KG vendor credit memos:
            # vendor line pkey 21, GL lines pkey 50, all amounts positive.
            # No worked finance example exists yet, so the doc is flagged
            # YLW for review before upload.
            n_credit += 1
            total31, ex, vat = abs(total31), abs(ex), abs(vat)
            adj_i, adj_w, adj_dep = abs(adj_i), abs(adj_w), -abs(adj_dep)
        m = vendors.get(_sap_norm_acct(acct))
        status, issue = "GRN", ""
        if credit:
            status = "YLW"
            issue = ("credit bill -> KG credit memo (pkey 21/50) - "
                     "confirm with finance before upload")
        if not m:
            n_unmapped += 1
            status, issue = "RED", "account %s not in vendor mapping" % acct
            m = {"vendor": "", "cc": "", "name1": "", "muni": ""}
        cc = m["cc"].lstrip("0")
        if vat_flag:
            status, issue = "YLW", (issue + "; " if issue else "") + vat_flag
        due = _sap_date(r.get("due_date"))
        if not due:
            status = "RED" if status == "RED" else "YLW"
            issue = (issue + "; " if issue else "") + "no due date on bill"
        bd = _sap_date(c.get("billing_date"))
        ref = str(r.get("invoice_number") or "").strip()
        if not ref:
            status, issue = "RED", (issue + "; " if issue else "") + "no invoice number extracted"
        hdr = (m["name1"] + " " + str(r.get("statement_month") or "")).strip()
        kwh = str(c.get("total_kwh") or "").strip()
        ae = ("A" if str(r.get("esk_reading_type", ""))
              .upper().startswith("A") else "E")
        rs = re.sub(r"/", "-", str(r.get("esk_read_start_date")
                                   or c.get("read_start") or ""))
        re_ = re.sub(r"/", "-", str(r.get("esk_read_end_date")
                                    or c.get("read_end") or ""))
        txt = ("%s_%s_%s_Cons_%s_Var_%s" % (cc, rs, re_, ae, kwh)
               if cc else "")

        doctype = (fixed.get("DocTypeCredit", "KG") if credit
                   else fixed.get("DocType", "KR"))

        def line(trans, pkey, account, amount, tc="", dt="", cc="",
                 alloc="", text="", pc=""):
            if credit:
                pkey = {"31": "21", "40": "50", "50": "40"}.get(pkey, pkey)
            return {"TransCode": trans, "docdate": bd, "postdate": bd,
                    "type": doctype,
                    "comp": fixed.get("Company", "T000"),
                    "curr": fixed.get("Currency", "ZAR"),
                    "refdocnr": ref, "docheadertext": hdr, "pkey": pkey,
                    "account": account, "sgl": "", "compn": "", "txi": "",
                    "amount": str(amount), "tc": tc, "dt": dt,
                    "costcentre": cc, "ordernr": "", "allocation": alloc,
                    "text": text, "ProfitCentre": pc, "DueDate": due,
                    "_status": status, "_issue": issue,
                    "_file": r.get("file", "")}

        t31 = fixed.get("TransCode31", "FB01")
        t40 = fixed.get("TransCode40", "ZB01")
        tc = fixed.get("TaxCode", "I9")
        out.append(line(t31, "31", m["vendor"], total31,
                        alloc=cc, text=txt, pc=cc))
        out.append(line(t40, "40", gl.get("electricity", "SA110"), ex,
                        tc=tc, cc=cc, alloc=cc, text=txt))
        if adj_w:
            out.append(line(t40, "40", gl.get("adjustment_work", "SA412"),
                            adj_w, tc=tc, cc=cc, alloc=cc,
                            text=str(r.get("esk_adj_work_desc")
                                     or "RECONNECTION WORK HRS")))
        if adj_i:
            # Interest posts VAT-exempt with the property's cost centre
            # (Mubeenah's sheet: tc I0, no dt, cc + allocation + wording).
            out.append(line(t40, "40", gl.get("adjustment_interest",
                            "SF101"), adj_i,
                            tc=fixed.get("TaxCodeCredit", "I0"),
                            cc=cc, alloc=cc,
                            text="Interest on overdue account"))
        out.append(line(t40, "40", gl.get("vat", "JA600"), vat,
                        tc=tc, dt="X", alloc=cc, pc=cc))
        if adj_dep:
            out.append(line(t40, "50",
                            gl.get("adjustment_deposit", "DA201"),
                            -adj_dep, tc=fixed.get("TaxCodeCredit", "I0"),
                            cc=cc, alloc=cc, text="Deposit"))
        n_ok += 1
    summary = ("%d document(s) generated (%d as KG credit memos) | "
               "%d account(s) unmapped"
               % (n_ok, n_credit, n_unmapped))
    return out, issues, summary


def export_sap_xlsx(path, sap_rows):
    """CBX: `path` may also be a file object (the page downloads it); the workbook is the Harvester's."""
    import openpyxl

    from .. import xlsafe

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(SAP_HEADERS)
    for r in sap_rows:
        ws.append([r.get(h, "") for h in SAP_HEADERS])
    xlsafe.book(wb)  # 2.10.0: a supplier's name starting with = stays text (nothing else in the file changes)
    wb.save(path)


def balance_fails(sap_rows) -> tuple[int, set]:
    """The Harvester's value check (_sap_balance_check): per document (grouped by source file) the vendor line must
    equal GL 40-lines minus 50-lines (mirrored for KG memos). Returns (n_docs, set of failing file keys)."""
    from decimal import Decimal as D
    docs = {}
    for r in sap_rows:
        docs.setdefault(r.get("_file", ""), []).append(r)
    fails = set()
    for f, rr in docs.items():
        try:
            vend = sum(D(r["amount"]) for r in rr
                       if r["pkey"] in ("31", "21"))
            kg = any(r["pkey"] == "21" for r in rr)
            gl = (sum(D(r["amount"]) for r in rr
                      if r["pkey"] == ("50" if kg else "40"))
                  - sum(D(r["amount"]) for r in rr
                        if r["pkey"] == ("40" if kg else "50")))
            if abs(vend - gl) > D("0.005"):
                fails.add(f)
        except Exception:
            fails.add(f)
    return len(docs), fails


# =============================================================================================================
# Part 2 — the other services, and manual invoices (2.10.0)
# =============================================================================================================
# a profile's service → its key in the GL Codes table (the Harvester called sewerage "sanitation")
SERVICE_GL = {"electricity": "electricity", "water": "water", "sewerage": "sanitation", "refuse": "refuse", "rates": "rates", "sundry": "sundry"}
SERVICE_LABEL = {"electricity": "Electricity", "water": "Water", "sewerage": "Sanitation", "refuse": "Refuse", "rates": "Rates", "sundry": "Sundry"}
# 2.10.0 seeds: the Harvester's GL Codes, now generated, plus sundry (no GL until an admin sets one: a bill with sundry
# charges stays red until then)
DEFAULT_GL = [(k, g, d.replace(" (not generated yet)", "")) for k, g, d in SAP_DEFAULT_GL] + [("sundry", "", "Sundry charges (set the GL before these post)")]
GRN, YLW, RED = "GRN", "YLW", "RED"
GL_LABEL = {"electricity": "electricity", "vat": "VAT", "adjustment_work": "work adjustments", "adjustment_interest": "interest", "adjustment_deposit": "deposits"}
# ISO 4217 (active codes): a currency Tess knows; anything else is red
CURRENCIES = set("""AED AFN ALL AMD ANG AOA ARS AUD AWG AZN BAM BBD BDT BGN BHD BIF BMD BND BOB BRL BSD BTN BWP BYN BZD CAD CDF CHF CLP CNY COP
CRC CUP CVE CZK DJF DKK DOP DZD EGP ERN ETB EUR FJD FKP GBP GEL GHS GIP GMD GNF GTQ GYD HKD HNL HTG HUF IDR ILS INR IQD IRR ISK JMD JOD JPY
KES KGS KHR KMF KPW KRW KWD KYD KZT LAK LBP LKR LRD LSL LYD MAD MDL MGA MKD MMK MNT MOP MRU MUR MVR MWK MXN MYR MZN NAD NGN NIO NOK NPR NZD
OMR PAB PEN PGK PHP PKR PLN PYG QAR RON RSD RUB RWF SAR SBD SCR SDG SEK SGD SHP SLE SOS SRD SSP STN SVC SYP SZL THB TJS TMT TND TOP TRY TTD
TWD TZS UAH UGX USD UYU UZS VES VND VUV WST XAF XCD XOF XPF YER ZAR ZMW ZWL""".split())
_CURRENCY_WORDS = {"R": "ZAR", "R.": "ZAR", "RAND": "ZAR", "RANDS": "ZAR", "US$": "USD", "USD$": "USD", "$US": "USD", "€": "EUR", "EURO": "EUR", "EUROS": "EUR",
                   "£": "GBP", "GB£": "GBP"}


def currency_code(v) -> str:
    """2.10.0: what an invoice prints → the ISO code SAP takes ("R" and "ZAR" are rand); unknown stays as printed (Tess: red)."""
    s = re.sub(r"\s+", "", str(v or "")).upper()
    return _CURRENCY_WORDS.get(s, s[:5])


def _need_gl(gl: dict, key: str, flag) -> str:
    """2.10.0: a GL from the GL Codes table, never the Harvester's default: a blank one makes the document red."""
    g = str(gl.get(key) or "").strip()
    if not g:
        flag(RED, f"no GL for {GL_LABEL.get(key, key)} in the GL Codes table")
    return g


def dec(v) -> Optional[Decimal]:
    if v is None or v == "":
        return None
    try:
        d = Decimal(str(v).replace(",", "").replace(" ", "").replace("R", ""))
    except Exception:  # noqa: BLE001
        return None
    return d if d.is_finite() else None


def m2(d: Decimal) -> str:
    return "%.2f" % d


def iso_to_sap(v) -> str:
    """An ISO date (or the bill's yyyy/mm/dd) → ddmmyyyy."""
    return _sap_date(v)


def month_words(iso: str) -> str:
    """'2026-09-15' → 'SEPTEMBER 2026' (the way Eskom prints its account month)."""
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", str(iso or ""))
    if not m:
        return ""
    try:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3))).strftime("%B %Y").upper()
    except ValueError:
        return ""


SAST = timezone(timedelta(hours=2))  # 1.0.102: Johannesburg (no daylight saving, so no tz database needed on Windows)
MONTHS = ("JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC")
HEADER_MAX = 25  # SAP keeps the first 25 characters of the header text (BKTXT)


def sap_today() -> date:
    """1.0.102: today in Johannesburg — the posting date a batch takes unless a person set another."""
    return datetime.now(SAST).date()


def sap_day(d: date) -> str:
    return f"{d:%d%m%Y}"


def far_from_today(d: date, today: Optional[date] = None) -> bool:
    """1.0.102: a posting date before the 1st of last month or more than 7 days ahead (Tess: yellow, not red)."""
    today = today or sap_today()
    first_last_month = (today.replace(day=1) - timedelta(days=1)).replace(day=1)
    return d < first_last_month or d > today + timedelta(days=7)


def due_from_terms(invoice_iso: str, days) -> str:
    """1.0.102: the invoice date + the vendor's payment days, as an ISO date ("" when either is missing)."""
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", str(invoice_iso or ""))
    try:
        n = int(days) if days not in (None, "") else None
    except (TypeError, ValueError):
        n = None
    if not m or n is None or not 0 <= n <= 365:
        return ""
    try:
        return (date(int(m.group(1)), int(m.group(2)), int(m.group(3))) + timedelta(days=n)).isoformat()
    except ValueError:
        return ""


_LEGAL = (re.compile(r"[\s,]*\(?\b(?:pty|proprietary)\b\)?\.?\s*(?:\b(?:ltd|limited)\b\.?)?\s*$", re.I),
          re.compile(r"[\s,]+\b(?:ltd|limited|cc|inc|llc|plc|sarl|lda|soc|w\.?l\.?l)\b\.?\s*$", re.I))


def short_name(name: str) -> str:
    """1.0.102: a supplier's name without its legal form: "Excellerate Services (Pty) Ltd" → "Excellerate Services"."""
    s = re.sub(r"\s+", " ", str(name or "")).strip()
    for _ in range(3):
        before = s
        for rx in _LEGAL:
            s = rx.sub("", s).strip(" ,-")
        if s == before:
            break
    return s or re.sub(r"\s+", " ", str(name or "")).strip()


def _fit_words(words: list[str], room: int) -> str:
    """Whole words while they fit; the next word shortened to four letters and a dot when that fits ("Serv."); a first
    word longer than the room is cut there."""
    out = ""
    for w in words:
        add = (" " if out else "") + w
        if len(out) + len(add) <= room:
            out += add
            continue
        if not out:
            out = w[:max(1, room - 1)] + "." if room > 1 else w[:room]
        elif len(w) > 5 and len(out) + 6 <= room:
            out += " " + w[:4] + "."
        break
    return out


def mmm_yy(iso: str) -> str:
    """'2026-08-15' → 'AUG 26' (fixed English month names: the same on every PC whatever its language)."""
    m = re.match(r"(\d{4})-(\d{2})-\d{2}", str(iso or ""))
    if not m or not 1 <= int(m.group(2)) <= 12:
        return ""
    return f"{MONTHS[int(m.group(2)) - 1]} {m.group(1)[2:]}"


def header_text(name: str, invoice_iso: str, site: str = "") -> str:
    """1.0.102: a manual invoice's header text, the useful part first and within SAP's 25 characters: "<supplier short
    name> <MMM YY>" cut at a word boundary ("Excellerate Serv. AUG 26"); a site code leads when the coding has no
    line-text rule ("DELL Shine Serv. SEP 26")."""
    tail = mmm_yy(invoice_iso)
    tail = (" " + tail) if tail else ""
    site = re.sub(r"\s+", " ", str(site or "")).strip()
    lead = (site[:max(0, HEADER_MAX - len(tail) - 6)] + " ") if site else ""
    room = HEADER_MAX - len(lead) - len(tail)
    body = _fit_words(short_name(name).split(), room)
    return (lead + body + tail).strip()[:HEADER_MAX]


def line_rule(rules: list, description: str) -> Optional[dict]:
    """1.0.102: the first of a vendor coding's line rules whose words are in the line's description (case does not
    matter; "a, b" = either; every word of one alternative must be there)."""
    desc = " " + re.sub(r"\s+", " ", str(description or "")).lower() + " "
    for r in rules or []:
        if not isinstance(r, dict):
            continue
        for alt in str(r.get("words") or "").lower().split(","):
            ws = alt.split()
            if ws and all(w in desc for w in ws):
                return r
    return None


class _Doc:
    """One document's lines with the Harvester's shape. Amounts are signed (+ = a debit on a normal invoice); when the
    whole document is negative it becomes a credit memo (KG, pkeys 21/50), as generate_sap does."""

    def __init__(self, *, fixed: dict, comp: str, curr: str, doctype: str, doctype_credit: str, bd: str, ref: str, hdr: str, due: str, file: str, pd: str = ""):
        self.fixed, self.comp, self.curr, self.bd, self.ref, self.hdr, self.due, self.file = fixed, comp, curr, bd, ref, hdr, due, file
        self.pd = pd or bd  # 1.0.102: the batch's posting date (the Harvester posted on the document's date)
        self.doctype, self.doctype_credit = doctype, doctype_credit
        self.lines: list[tuple] = []  # (account, signed amount, tc, dt, cc, alloc, text, pc)
        self.status, self.issues = GRN, []

    def add(self, account: str, amount: Decimal, *, tc="", dt="", cc="", alloc="", text="", pc="") -> None:
        if amount:
            self.lines.append((account, amount, tc, dt, cc, alloc, text, pc))

    def flag(self, level: str, text: str) -> None:
        if level == RED or (level == YLW and self.status == GRN):
            self.status = level
        if text and text not in self.issues:
            self.issues.append(text)

    def rows(self, vendor: str, total: Decimal, *, alloc="", text="", pc="") -> list[dict]:
        credit = total < 0
        if credit:
            self.flag(YLW, "credit -> KG credit memo (pkey 21/50) - confirm with finance before upload")
        t31, t40 = self.fixed.get("TransCode31", "FB01"), self.fixed.get("TransCode40", "ZB01")
        doctype = self.doctype_credit if credit else self.doctype

        def line(trans, pkey, account, amount, tc="", dt="", cc="", alloc="", text="", pc=""):
            return {"TransCode": trans, "docdate": self.bd, "postdate": self.pd, "type": doctype, "comp": self.comp, "curr": self.curr,
                    "refdocnr": self.ref, "docheadertext": self.hdr, "pkey": pkey, "account": account, "sgl": "", "compn": "", "txi": "",
                    "amount": m2(abs(amount)), "tc": tc, "dt": dt, "costcentre": cc, "ordernr": "", "allocation": alloc, "text": text,
                    "ProfitCentre": pc, "DueDate": self.due, "_status": self.status, "_issue": "; ".join(self.issues), "_file": self.file}

        out = [line(t31, "21" if credit else "31", vendor, total, alloc=alloc, text=text, pc=pc)]
        for account, amount, tc, dt, cc, al, tx, p in self.lines:
            # a GL debit is 40 and a GL credit 50 on an invoice and on a credit memo alike (generate_sap reaches the same
            # keys by making a credit bill's amounts positive and swapping 40 and 50); the vendor line is 31 or 21
            out.append(line(t40, "40" if amount >= 0 else "50", account, amount, tc=tc, dt=dt, cc=cc, alloc=al, text=tx, pc=p))
        for r in out:  # the status settles only once every line is known
            r["_status"], r["_issue"] = self.status, "; ".join(self.issues)
        return out


def _utility_header(fields: dict, acct: Optional[dict], tables: dict, company: str, doctypes: tuple[str, str], file: str) -> tuple[_Doc, dict]:
    fixed = tables["fixed"]
    m = acct or {"vendor": "", "cc": "", "name1": "", "muni": ""}
    d = _Doc(fixed=fixed, comp=company or fixed.get("Company", "T000"), curr=fixed.get("Currency", "ZAR"), doctype=doctypes[0], doctype_credit=doctypes[1],
             bd=iso_to_sap(fields.get("billing_date")), ref=str(fields.get("invoice_no") or "").strip(),
             hdr=(m["name1"] + " " + str(fields.get("statement_month") or "")).strip(), due=iso_to_sap(fields.get("due_date")), file=file)
    if not acct:
        d.flag(RED, "account %s not in vendor mapping" % (fields.get("account") or ""))
    if not d.due:
        d.flag(YLW, "no due date on bill")
    if not d.ref:
        d.flag(RED, "no invoice number extracted")
    if not d.bd:
        d.flag(RED, "no billing date on the bill")
    return d, m


def _harvester_row(fields: dict, elec: dict, elec_row: Optional[dict], file: str) -> dict:
    """The profile's electricity row as the Harvester saw it, with what the bill says now (a person may have corrected
    it) laid over it — for an untouched bill these are the same values."""
    row = dict(elec_row or {})
    row.update(service="electricity", file=file)
    canon = dict(row.get("_canon") or {})
    canon.update(account_number=fields.get("account") or "", total_ex_vat_r=elec.get("ex"), billing_date=fields.get("billing_date") or "",
                 read_start=fields.get("read_start") or "", read_end=fields.get("read_end") or "", total_kwh=fields.get("kwh") or "")
    row["_canon"] = canon
    adj = fields.get("adjustments") or {}
    row.update(elec_vat_r=elec.get("vat"), vat_out=None, esk_adj_interest_r=adj.get("interest"), esk_adj_work_r=adj.get("work"),
               esk_adj_work_desc=adj.get("work_desc"), esk_adj_deposit_r=adj.get("deposit"), due_date=fields.get("due_date") or "",
               invoice_number=fields.get("invoice_no") or "", statement_month=fields.get("statement_month") or "",
               esk_reading_type=fields.get("reading_type") or "", esk_read_start_date=fields.get("read_start") or "", esk_read_end_date=fields.get("read_end") or "")
    return row


def _one_sided(elec: dict, adj: dict) -> bool:
    """2.10.0: the Harvester's generator writes a negative adjustment as a negative 40-line and turns every amount positive
    on a credit bill (interest then lands on the wrong side). A bill whose lines don't all go one way goes through _Doc,
    which posts each negative line on the opposite posting key with a positive amount."""
    ex, vat = dec(elec.get("ex")), dec(elec.get("vat"))
    if vat is None and ex is not None:
        vat = (ex * Decimal("0.15")).quantize(Decimal("0.01"))
    i, w, dep = (dec(adj.get(k)) or Decimal(0) for k in ("interest", "work", "deposit"))
    parts = [x for x in (ex, vat, i, w) if x]
    if sum(parts, Decimal(0)) + dep >= 0:
        return all(x > 0 for x in parts) and dep <= 0
    return all(x < 0 for x in parts) and not dep


def utility_doc(fields: dict, elec_row: Optional[dict], acct: Optional[dict], tables: dict, *, file: str, company: str = "",
                doctypes: Optional[tuple[str, str]] = None, postdate: str = "") -> tuple[list[dict], str, list[str]]:
    """One utility bill → (SAP rows, GRN/YLW/RED, issues). `fields` is the bill as read (and corrected); `elec_row` the
    profile's electricity row; `acct` the utilities-list row {vendor, cc, name1, muni} (None: not in the list).
    1.0.102: `postdate` (ddmmyyyy) = the batch's posting date on every row; "" = the bill's date, as the Harvester did."""
    rows, st, issues = _utility_doc(fields, elec_row, acct, tables, file=file, company=company, doctypes=doctypes)
    if postdate:
        for r in rows:
            r["postdate"] = postdate
    return rows, st, issues


def _utility_doc(fields: dict, elec_row: Optional[dict], acct: Optional[dict], tables: dict, *, file: str, company: str = "",
                 doctypes: Optional[tuple[str, str]] = None) -> tuple[list[dict], str, list[str]]:
    fixed = dict(tables["fixed"])
    if company:
        fixed["Company"] = company
    doctypes = doctypes or (fixed.get("DocType", "KR"), fixed.get("DocTypeCredit", "KG"))
    fixed["DocType"], fixed["DocTypeCredit"] = doctypes
    services = [s for s in fields.get("services") or [] if s.get("service") in SERVICE_GL]
    elec = next((s for s in services if s["service"] == "electricity"), None)
    others = [s for s in services if s["service"] != "electricity" and dec(s.get("ex"))]
    if elec is not None and not others and _one_sided(elec, fields.get("adjustments") or {}):  # an electricity bill: the Harvester's own generator, untouched
        row = _harvester_row(fields, elec, elec_row, file)
        acct_key = _sap_norm_acct(fields.get("account"))
        gl_in = {**{k: "" for k in GL_LABEL}, **{k: str(v or "").strip() for k, v in tables["gl"].items()}}  # 2.10.0: a GL missing from the table is blank, never the Harvester's default
        rows, issues, _ = generate_sap([row], {"fixed": fixed, "gl": gl_in, "vendors": {acct_key: acct} if acct and acct_key else {}})
        if not rows:
            return [], RED, [t for _, t in issues] or ["no electricity total - not posted"]
        st, why = rows[0]["_status"], rows[0]["_issue"]
        issues = [x.strip() for x in why.split(";") if x.strip()]
        rows = [r for n, r in enumerate(rows) if n == 0 or dec(r["amount"])]  # 2.10.0: a zero line (VAT 0.00) is not posted
        adj = fields.get("adjustments") or {}
        ex0, vat0 = dec(elec.get("ex")), dec(elec.get("vat"))
        need = {"electricity": ex0, "vat": vat0 if vat0 is not None else ex0, "adjustment_work": dec(adj.get("work")), "adjustment_interest": dec(adj.get("interest")),
                "adjustment_deposit": dec(adj.get("deposit"))}
        for k, v in need.items():
            if v and not gl_in[k]:  # 2.10.0: blank GL = red with why
                st = RED
                issues.append(f"no GL for {GL_LABEL[k]} in the GL Codes table")
        return rows, st, issues
    if not services or (elec is None and not others):
        return [], RED, ["no charges read from the bill - not posted"]
    d, m = _utility_header(fields, acct, {**tables, "fixed": fixed}, fixed.get("Company", "T000"), doctypes, file)
    gl, tc, tc0 = tables["gl"], fixed.get("TaxCode", "I9"), fixed.get("TaxCodeCredit", "I0")
    cc = m["cc"].lstrip("0")
    rs, re_ = str(fields.get("read_start") or "").replace("/", "-"), str(fields.get("read_end") or "").replace("/", "-")
    total, vat_total, txt31 = Decimal(0), Decimal(0), ""
    for s in ([elec] if elec else []) + others:
        ex = dec(s.get("ex"))
        if ex is None:
            d.flag(RED, f"no {SERVICE_LABEL[s['service']].lower()} total - not posted")
            continue
        vat = dec(s.get("vat"))
        if vat is None:
            if s["service"] == "rates":
                vat = Decimal(0)  # property rates carry no VAT unless the bill prints it
            else:
                vat = (ex * Decimal("0.15")).quantize(Decimal("0.01"))
                d.flag(YLW, f"{SERVICE_LABEL[s['service']]} VAT not printed on bill - calculated at 15%")
        account = (gl.get(SERVICE_GL[s["service"]]) or "").strip()
        if not account:
            d.flag(RED, f"no GL for {SERVICE_LABEL[s['service']].lower()} in the GL Codes table")
        if s["service"] == "electricity":
            kwh = str(fields.get("kwh") or "").strip()
            ae = "A" if str(fields.get("reading_type") or "").upper().startswith("A") else "E"
            text = "%s_%s_%s_Cons_%s_Var_%s" % (cc, rs, re_, ae, kwh) if cc else ""
        else:
            text = "_".join(x for x in (cc, SERVICE_LABEL[s["service"]].upper(), rs, re_) if x) if cc else ""
        txt31 = txt31 or text
        d.add(account, ex, tc=tc if vat else tc0, cc=cc, alloc=cc, text=text)
        total += ex + vat
        vat_total += vat
    adj = fields.get("adjustments") or {}
    a_w, a_i, a_dep = dec(adj.get("work")) or Decimal(0), dec(adj.get("interest")) or Decimal(0), dec(adj.get("deposit")) or Decimal(0)
    if a_w:  # 2.10.0: the GLs come from the GL Codes table only (blank = red), as the services' do
        d.add(_need_gl(gl, "adjustment_work", d.flag), a_w, tc=tc, cc=cc, alloc=cc, text=str(adj.get("work_desc") or "RECONNECTION WORK HRS"))
    if a_i:
        d.add(_need_gl(gl, "adjustment_interest", d.flag), a_i, tc=tc0, cc=cc, alloc=cc, text="Interest on overdue account")
    if vat_total:
        d.add(_need_gl(gl, "vat", d.flag), vat_total, tc=tc, dt="X", alloc=cc, pc=cc)
    if a_dep:
        d.add(_need_gl(gl, "adjustment_deposit", d.flag), a_dep, tc=tc0, cc=cc, alloc=cc, text="Deposit")
    total += a_w + a_i + a_dep
    printed = dec(fields.get("vat_printed"))
    if printed is not None and abs(printed - vat_total) > Decimal("0.01") and vat_total:
        d.flag(YLW, f"the services' VAT R {m2(vat_total)} is not the bill's VAT total R {m2(printed)}")
    rows = d.rows(m["vendor"], total, alloc=cc, text=txt31, pc=cc)
    return rows, d.status, d.issues


def line_text(rule: str, *, site: str = "", description: str = "", supplier: str = "", invoice: str = "", po: str = "") -> str:
    """2.10.3: a vendor coding's line text (SGTXT), e.g. "{site} {description}" → "RIO2X21 Project Value" (50 characters)."""
    vals = {"site": site, "description": description, "supplier": supplier, "invoice": invoice, "po": po}
    out = re.sub(r"\{(\w+)\}", lambda m: str(vals.get(m.group(1).lower(), "")), str(rule or ""))
    return re.sub(r"\s+", " ", out).strip()[:50]


def manual_doc(fields: dict, coding: Optional[dict], company: str, tables: dict, *, file: str, doctypes: Optional[tuple[str, str]] = None,
               site: str = "", postdate: str = "") -> tuple[list[dict], str, list[str]]:
    """A manual invoice → (SAP rows, GRN/YLW/RED, issues). `coding` = the vendor's {vendor_no, name, gl, cost_centre, tax_code,
    item_text}. 2.10.3: `site` = the bill-to's site code (after " - "): into each G/L line's text the way the coding's
    item_text says, or at the start of the header text when the coding has no rule.
    1.0.102: `coding` may also carry line_rules ([{words, gl, cost_centre, tax_code}]), payment_days, and gl_fixed /
    cc_fixed (a person set the G/L / cost centre for this invoice: the rules don't change it); `postdate` (ddmmyyyy) is
    the batch's posting date ("" = the invoice date, as before)."""
    fixed = tables["fixed"]
    doctypes = doctypes or (fixed.get("DocType", "KR"), fixed.get("DocTypeCredit", "KG"))
    coding = coding or {}
    when = str(fields.get("invoice_date") or "")
    rule = str(coding.get("item_text") or "").strip()
    # 1.0.102: the useful part first, within the 25 characters SAP keeps; no line-text rule: the site leads (as 2.10.3)
    hdr = header_text(str(coding.get("name") or fields.get("supplier") or ""), when, site if not rule else "")
    due = iso_to_sap(fields.get("due_date")) or iso_to_sap(due_from_terms(when, coding.get("payment_days")))  # 1.0.102: the vendor's payment terms
    d = _Doc(fixed=fixed, comp=company, curr=currency_code(fields.get("currency") or fixed.get("Currency", "ZAR")), doctype=doctypes[0],  # 2.10.0: "R" is ZAR
             doctype_credit=doctypes[1], bd=iso_to_sap(when), ref=str(fields.get("invoice_no") or "").strip(),
             hdr=hdr, due=due, file=file, pd=postdate)
    if not coding.get("vendor_no"):
        d.flag(RED, "the vendor is not coded")
    if not company:
        d.flag(RED, "no company code")
    if not d.ref:
        d.flag(RED, "no invoice number")
    if not d.bd:
        d.flag(RED, "no invoice date")
    if not d.due:
        d.flag(YLW, "no due date on the invoice")
    tc, tc0 = (coding.get("tax_code") or fixed.get("TaxCode", "I9")), fixed.get("TaxCodeCredit", "I0")
    tcx = str(fixed.get("TaxCodeExempt") or "").strip().upper()  # 1.0.102: "" = an exempt line keeps its usual code
    cc, gl0 = str(coding.get("cost_centre") or "").strip(), str(coding.get("gl") or "").strip()
    rules = [r for r in coding.get("line_rules") or [] if isinstance(r, dict) and str(r.get("words") or "").strip()]
    excl, vat, incl = dec(fields.get("excl")), dec(fields.get("vat")) or Decimal(0), dec(fields.get("incl"))
    sign = Decimal(-1) if fields.get("credit") and (incl or 0) > 0 else Decimal(1)  # a credit note printed with positive figures
    if incl is None and excl is not None:
        incl = excl + vat
    if excl is None and incl is not None:
        excl = incl - vat
    if incl is None:
        d.flag(RED, "no total on the invoice")
        return [], RED, d.issues
    groups: dict[tuple, list] = {}
    lines = [ln for ln in fields.get("lines") or [] if dec(ln.get("excl")) is not None]
    if lines and sum((dec(ln["excl"]) for ln in lines), Decimal(0)) == excl:
        for ln in lines:
            lvat = dec(ln.get("vat"))
            rate = dec(ln.get("vat_rate"))
            t = tc if ((lvat or 0) != 0 or (rate or 0) > 0 or (lvat is None and rate is None and vat)) else tc0  # a 0 % line: the zero-rated code
            hit = line_rule(rules, ln.get("description"))  # 1.0.102: the coding's line rules
            g = str(ln.get("gl") or ((hit or {}).get("gl") if not coding.get("gl_fixed") else "") or gl0).strip()
            lcc = cc if coding.get("cc_fixed") else (str((hit or {}).get("cost_centre") or "").strip() or cc)
            rtc = str((hit or {}).get("tax_code") or "").strip().upper()
            if rtc == "EXEMPT":
                t = tcx or t  # the Fixed TaxCodeExempt; when it is empty nothing changes
            elif rtc:
                t = rtc
            groups.setdefault((g, t, lcc), []).append(ln)  # 1.0.102: grouped by G/L, tax code and cost centre
    else:
        if lines:
            d.flag(YLW, "the lines did not add up to the excl. total: posted as one line")
        groups[(gl0, tc if vat else tc0, cc)] = [{"description": "", "excl": str(excl)}]
    if any(not lcc for (_g, _t, lcc) in groups):
        d.flag(YLW, "no cost centre for this vendor")
    for (g, t, lcc), lns in groups.items():
        if not g:
            d.flag(RED, "no GL for this vendor")
        amount = sum((dec(ln["excl"]) for ln in lns), Decimal(0)) * sign
        text = (str(lns[0].get("description") or "").strip() or str(fields.get("supplier") or "")).replace("\n", " ")[:50]
        if rule:  # 2.10.3
            text = line_text(rule, site=site, description=text, supplier=str(fields.get("supplier") or ""), invoice=d.ref, po=str(fields.get("po") or ""))
        d.add(g, amount, tc=t, cc=lcc, alloc=lcc, text=text)
    if vat:
        d.add(_need_gl(tables["gl"], "vat", d.flag), vat * sign, tc=tc, dt="X", alloc=cc, pc=cc)  # 2.10.0: no guessed JA600
    if abs((excl + vat) - incl) > Decimal("0.005"):
        d.flag(RED, f"excl. R {m2(excl)} + VAT R {m2(vat)} is not the total R {m2(incl)}")
    rows = d.rows(str(coding.get("vendor_no") or ""), incl * sign, alloc=cc, text=(str(fields.get("supplier") or "") + " " + d.ref).strip()[:50], pc=cc)
    return rows, d.status, d.issues


# =============================================================================================================
# Part 3 — Tess
# =============================================================================================================
REQUIRED = ("TransCode", "docdate", "postdate", "type", "comp", "curr", "refdocnr", "pkey", "account", "amount")


def real_date(s: str) -> bool:
    m = re.fullmatch(r"(\d{2})(\d{2})(\d{4})", str(s or ""))
    if not m:
        return False
    try:
        date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        return True
    except ValueError:
        return False


def test_doc(rows: list[dict], known: dict, *, pip_status: str = GRN, pip_issues: Optional[list] = None, dup: str = "", warn: Optional[list] = None,
             stop: Optional[list] = None) -> tuple[str, list[dict]]:
    """Tess's tests on one document → ('green' | 'yellow' | 'red', [{level, text}]).
    known: {vendors, gls, companies, doc_types, tax_codes} (sets of what admins keep in the tables) and, optionally,
    learned_vendors / learned_gls / learned_tax_codes (what people added: yellow until an admin confirms it).
    2.10.3, optionally: chart (the chart of accounts as knowledge.gl_key values; None = no chart), doc_type_accounts
    ({code: SAP's account types}), vendor_in_company (vendor, company code → True / False / None), cost_centre_in_company
    (cost centre, company code → "yes" / "no" / "dummy" / None).
    1.0.102, optionally: today (a date: a posting date far from it is yellow) and vendor_blocked (vendor → True when SAP
    has it blocked or marked for deletion: red).
    warn / stop: more reasons for yellow / red (the AP team's own verdict on the file, a learned coding)."""
    out: list[dict] = []

    def say(level, text):
        if text and not any(x["text"] == text for x in out):
            out.append({"level": level, "text": text})

    def listed(what, value, admin, learned):
        """2.10.0: an admin-kept value passes; a learned one is yellow until an admin confirms it; anything else is red."""
        if value in known.get(admin, set()):
            return
        if value in known.get(learned, set()):
            say("yellow", f"{what} {value} is learned: an admin confirms it once")
        else:
            say("red", f"{what} {value} is not {'in the vendor tables' if what == 'vendor' else 'in the GL tables' if what == 'GL' else 'known'}")

    for t in pip_issues or []:
        say("red" if pip_status == RED else "yellow", t)
    for t in stop or []:
        say("red", t)
    for t in warn or []:
        say("yellow", t)
    if not rows:
        say("red", "no SAP rows")
        return "red", out
    _, fails = balance_fails([{**r, "_file": "x"} for r in rows])
    if fails:
        say("red", "the 40-lines do not add up to the 31-line")
    empty: dict[str, list] = {}
    for n, r in enumerate(rows, 1):
        for col in REQUIRED + (("tc",) if r.get("pkey") in ("40", "50") else ()):
            if not str(r.get(col) or "").strip():
                empty.setdefault(col, []).append(n)
    for col, lines in empty.items():  # once per column, not once per line
        say("red", f"{col} is empty" + (f" on line {lines[0]}" if len(lines) == 1 else f" on {len(lines)} lines"))
    for n, r in enumerate(rows, 1):
        for col in ("docdate", "postdate"):
            if r.get(col) and not real_date(r[col]):
                say("red", f"line {n}: {col} {r[col]} is not a real date (ddmmyyyy)")
        if r.get("DueDate") and not real_date(r["DueDate"]):
            say("red", f"line {n}: DueDate {r['DueDate']} is not a real date (ddmmyyyy)")
        if str(r.get("amount") or "").strip():  # 2.10.0: SAP takes a positive amount (the posting key gives the side), at most 13 digits
            a = dec(r["amount"])
            if a is None or a <= 0:
                say("red", f"line {n}: amount {r['amount']} is not above zero")
            elif len(re.sub(r"\D", "", m2(a))) > 13:
                say("red", f"line {n}: amount {m2(a)} is longer than SAP takes (13 digits)")
        acc = str(r.get("account") or "")
        if r.get("pkey") in ("31", "21"):
            if acc:
                listed("vendor", acc, "vendors", "learned_vendors")
        elif acc:
            listed("GL", acc, "gls", "learned_gls")
        if r.get("tc"):
            listed("tax code", r["tc"], "tax_codes", "learned_tax_codes")
    chart = known.get("chart")
    if chart is not None:  # 2.10.3: a G/L the chart of accounts doesn't have: yellow — the chart export may be incomplete
        from ..knowledge import gl_key

        for r in rows:
            acc = str(r.get("account") or "").strip()
            if r.get("pkey") in ("40", "50") and acc and gl_key(acc) not in chart:
                say("yellow", f"G/L {acc} is not in the chart of accounts — check with finance")
    r0 = rows[0]
    today = known.get("today")
    if isinstance(today, date) and real_date(r0.get("postdate")):  # 1.0.102: a posting date someone set far from today
        s = str(r0["postdate"])
        pd = date(int(s[4:]), int(s[2:4]), int(s[:2]))
        if far_from_today(pd, today):
            say("yellow", f"posting date {pd:%d.%m.%Y} is far from today")
    blocked = known.get("vendor_blocked")
    if callable(blocked):  # 1.0.102: SAP's vendor general data (LFA1 SPERR / LOEVM)
        for v in dict.fromkeys(str(r.get("account") or "").strip() for r in rows if r.get("pkey") in ("31", "21")):
            if v and blocked(v):
                say("red", f"vendor {v} is blocked in SAP")
    dt = str(r0.get("type") or "").strip().upper()
    kinds = (known.get("doc_type_accounts") or {}).get(dt)
    if kinds:  # 2.10.3: the document type must allow the account types its lines use (SAP's Document Type List)
        said = ", ".join(ch for ch in "ADKMS" if ch in kinds)
        if "K" not in kinds and any(r.get("pkey") in ("31", "21") for r in rows):
            say("red", f"{dt} doesn't allow vendor lines in SAP (account types {said})")
        if "S" not in kinds and any(r.get("pkey") in ("40", "50") for r in rows):
            say("red", f"{dt} doesn't allow G/L lines in SAP (account types {said})")
    opened = known.get("vendor_in_company")
    if callable(opened) and r0.get("comp"):  # 2.10.3: only when SAP's vendor / company code list knows the company code
        for v in dict.fromkeys(str(r.get("account") or "").strip() for r in rows if r.get("pkey") in ("31", "21")):
            if v and opened(v, r0["comp"]) is False:
                say("red", f"Vendor {v} is not opened in company code {r0['comp']} in SAP")
    in_co = known.get("cost_centre_in_company")
    if callable(in_co) and r0.get("comp"):  # 2.10.3: SAP's Cost Center List, only for a company code it has cost centres for
        for cc in dict.fromkeys(str(r.get("costcentre") or "").strip() for r in rows if r.get("pkey") in ("40", "50")):
            got = in_co(cc, r0["comp"]) if cc else None
            if got == "dummy":
                say("red", f"Cost centre {cc} is a dummy cost centre")
            elif got == "no":
                say("red", f"Cost centre {cc} doesn't exist in company code {r0['comp']} in SAP")
    if r0.get("comp") and r0["comp"] not in known.get("companies", set()):
        say("red", f"company code {r0['comp']} is not in the company list")
    if r0.get("type") and r0["type"] not in known.get("doc_types", set()):
        say("red", f"document type {r0['type']} is not in the document types")
    cur = str(r0.get("curr") or "").strip().upper()
    if cur and cur not in CURRENCIES:  # 2.10.0: the currency is tested too
        say("red", f"currency {cur} is not known")
    elif cur and cur != "ZAR":
        say("yellow", f"in {cur}, not rand: confirm before upload")
    if dup:
        say("red", dup)
    if any(x["level"] == "red" for x in out):
        return "red", out
    return ("yellow" if out else "green"), out


# =============================================================================================================
# Part 4 — the files
# =============================================================================================================
HARVESTER_GREEN = "003F2D"  # 1.0.99: the Harvester's grid header
VENDOR_ROW = "E3F1E8"  # each document's vendor line (FB01, pkey 31/21), as in the Harvester's grid
STATUS_FILL = {"green": "63BE7B", "yellow": "FFD966", "red": "F8696B"}  # Tess's verdict, on the vendor line's TransCode cell
TEXT_COLS = ("refdocnr", "account", "costcentre", "allocation", "ProfitCentre", "comp", "tc")  # leading zeros kept
STATUS_WORD = {"GRN": "green", "YLW": "yellow", "RED": "red"}


def _head_style(ws, ncols: int) -> None:
    """1.0.99: the Harvester's header: dark green, bold white, filters over the whole table, the header row frozen."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    for c in ws[1][:ncols]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HARVESTER_GREEN)
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{max(1, ws.max_row)}"


def _shown(v) -> str:
    return f"{v:.2f}" if isinstance(v, float) else str(v if v is not None else "")


def _fit(ws, cap: int = 40) -> None:
    """Column widths fitted to the content (the header plus room for the filter arrow), at most `cap`."""
    for col in ws.iter_cols(min_row=1, max_row=min(ws.max_row, 2000)):
        w = max([len(_shown(col[0].value)) + 3] + [len(_shown(c.value)) for c in col[1:]])
        ws.column_dimensions[col[0].column_letter].width = min(cap, max(6, w + 1))


def export_sap_xlsx_styled(path, sap_rows) -> None:
    """1.0.99 (Edwin's picture of the Harvester's grid): the same file as export_sap_xlsx — sheet "Sheet1", the 22
    SAP_HEADERS in order, the same values — styled: the green header with filters, frozen; each document's vendor line
    (pkey 31/21, FB01) shaded light green and its TransCode cell in Tess's colour (green / yellow / red); the 40/50 lines
    white. Amounts are numbers (format 0.00); refdocnr, account, cost centre, allocation, profit centre, company code and
    tax code are text (leading zeros kept); the dates stay ddmmyyyy text. No extra column: SAP's uploader reads it as is."""
    import openpyxl
    from openpyxl.styles import PatternFill

    from .. import xlsafe

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(SAP_HEADERS)
    a_col = SAP_HEADERS.index("amount") + 1
    shade = PatternFill("solid", fgColor=VENDOR_ROW)
    for r in sap_rows:
        vals = []
        for h in SAP_HEADERS:
            v = r.get(h, "")
            if h == "amount":
                d = dec(v) if str(v or "").strip() else None
                v = float(d) if d is not None else v  # a number SAP's uploader reads; what isn't one stays as it was
            elif h in TEXT_COLS and v is not None and not isinstance(v, str):
                v = str(v)
            vals.append(v)
        ws.append(vals)
        n = ws.max_row
        amt = ws.cell(row=n, column=a_col)
        if isinstance(amt.value, float):
            amt.number_format = "0.00"
        for h in TEXT_COLS:
            c = ws.cell(row=n, column=SAP_HEADERS.index(h) + 1)
            if isinstance(c.value, str):
                c.number_format = "@"  # Excel keeps it text when someone edits the cell
        if str(r.get("pkey") or "") in ("31", "21"):
            for c in ws[n][:len(SAP_HEADERS)]:
                c.fill = shade
            st = str(r.get("_tess") or STATUS_WORD.get(str(r.get("_status") or ""), "")).lower()
            if st in STATUS_FILL:
                ws.cell(row=n, column=1).fill = PatternFill("solid", fgColor=STATUS_FILL[st])
        _guess_cells(ws, n, r.get("_guess"))  # 1.0.105: a guessed cell is yellow with a note saying how it was guessed
    _head_style(ws, len(SAP_HEADERS))
    ws.freeze_panes = "A2"
    _fit(ws)
    xlsafe.book(wb)
    wb.save(path)


GUESS_FILL = "FFF2CC"  # 1.0.105: a cell CBX guessed (cbx/ap/simple.py cell_marks): light yellow, its note says how


def _guess_cells(ws, n: int, marks) -> None:
    """1.0.105: `marks` = {SAP header: how it was guessed} — each of those cells on row n light yellow with a note."""
    if not isinstance(marks, dict) or not marks:
        return
    from openpyxl.comments import Comment
    from openpyxl.styles import PatternFill

    for h, how in marks.items():
        if h in SAP_HEADERS and how:
            c = ws.cell(row=n, column=SAP_HEADERS.index(h) + 1)
            if h != "TransCode":  # the TransCode cell keeps Tess's colour
                c.fill = PatternFill("solid", fgColor=GUESS_FILL)
            c.comment = Comment(("Guessed by CBX: " + str(how))[:900], "CBX")


def upload_file(rows: list[dict]) -> bytes:
    """The Harvester's file: one sheet "Sheet1", the 22 SAP_HEADERS, the rows — 1.0.99: styled like the Harvester's grid
    (export_sap_xlsx_styled); each row's `_tess` (the document's green/yellow/red) colours its vendor line's TransCode."""
    buf = io.BytesIO()
    export_sap_xlsx_styled(buf, rows)
    return buf.getvalue()


def review_file(docs: list[dict]) -> bytes:
    """Every row with its document's status and why, and the source file. docs: [{file, status, issues, rows, uploaded}]."""
    import openpyxl
    from openpyxl.styles import PatternFill

    from .. import xlsafe

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SAP rows"
    head = ["Status", "Why", "Source file", "Uploaded"] + SAP_HEADERS + ["G/L name", "Cost centre name"]  # 2.10.3: SAP's names of each line's G/L and cost centre
    extra = list(dict.fromkeys(k for d in docs for k in (d.get("extra") or {})))  # 1.0.104 (Builder R, V05): a document's own columns (the review …)
    head = head + extra
    ws.append(head)
    fill = {"green": "E2F2E8", "yellow": "FFF2CC", "red": "F7D9D9"}
    for d in docs:
        why = "; ".join(x["text"] for x in d.get("issues") or [])
        for r in d.get("rows") or [{}]:
            ws.append([d.get("status", ""), why, d.get("file", ""), d.get("uploaded", "")] + [r.get(h, "") for h in SAP_HEADERS] + [r.get("_gl_name", ""), r.get("_cc_name", "")]
                      + [(d.get("extra") or {}).get(k, "") for k in extra])
            if d.get("status") in fill:
                for c in ws[ws.max_row]:
                    c.fill = PatternFill("solid", fgColor=fill[d["status"]])
    _head_style(ws, len(head))  # 1.0.99: the Harvester's green header and filters, like the upload file
    ws.freeze_panes = "E2"
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(48, max(8, *(len(str(c.value or "")) for c in col[:300])) + 2)
    xlsafe.book(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
