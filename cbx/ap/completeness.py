# 1.0.104: new (Builder C) — V10 account status, masterlists and a completeness view (VALUENCE-AP-GAPS.md §2 V10; SPEC-1104)
"""Which utility accounts should have a bill this month, and which did.

    Account fields (ApUtilityAccount): portfolio, status (active / suspended / terminated / disposed) with its effective
    date and source, Shell's CO/DO, Telkom's cluster, a check meter, the usual bill day (else learned from the history).

    Masterlist import per portfolio (`read_list`, `import_list`): Shell's monthly masterlist, Telkom's SAP Masterfile
    (a full list) and its Termination / Anomaly reports, Swiftnet's Utility Master List. The list is compared with the
    utilities list and with SAP's vendor list; what it says about an account (status, CO/DO, cluster, check meter, bill
    day, the portfolio) is taken; every difference is an exception a person resolves with a note:
        not_in_cbx         on the list, not in CBX's utilities list
        not_in_sap         its vendor is not in SAP's vendor list
        not_on_list        in CBX (this portfolio), not on the list (a full masterlist only)
        terminated_active  terminated or disposed on the list, still active in CBX until now
        terminated_billed  terminated, and a bill or a payment came after the date
        co_do              CO → DO (or DO → CO)
        anomaly            on the Anomaly report
    Audited ap_masterlist_import (the date and source of the refresh); ap_masterlist_exception (a resolution).

    Completeness (`month_view`) per portfolio, cluster and month: each account that should bill is posted (its bill went
    out in an SAP upload) · read (a bill read, not uploaded yet) · estimated (an approved estimate, Builder U's
    ap_estimates) · missing (with its age and the day it counts as late: the usual bill day + missing_grace_days). The
    month's payments show beside it once FBL1N payments are loaded (cbx/ap/payments.py). Missing and late → the Council
    Liaison list. `export` is the weekly workbook that replaces the offline Cluster Portfolio Tracker.

    Tess (`tess_note`): a bill on a terminated account is red "account terminated on dd.mm.yyyy" (a final bill whose
    reading ends by that day is yellow); a suspended account is yellow.

What a masterlist says is data, never instructions."""
from __future__ import annotations

import calendar
import copy
import io
import re
from datetime import date, datetime, timedelta
from statistics import median
from typing import Any, Iterable, Optional

from sqlalchemy.orm import Session

from .. import knowledge as K
from ..audit import log
from ..db import utcnow
from ..models import User
from ..models.ap import (ApBillHistory, ApCapture, ApEstimate, ApMasterlistException, ApMasterlistImport, ApPayment, ApSapVendor,
                         ApUtilityAccount)
from . import portfolio as PF

STATUSES = ("active", "suspended", "terminated", "disposed")
CLOSED = ("terminated", "disposed")
KINDS = {"masterlist": "Masterlist (the full list)", "terminations": "Termination report", "anomalies": "Anomaly report"}
EXCEPTIONS = {"not_in_cbx": "On the list, not in CBX's utilities list", "not_in_sap": "Its vendor is not in SAP's vendor list",
              "not_on_list": "In CBX, not on the list", "terminated_active": "Terminated on the list, active in CBX until now",
              "terminated_billed": "Terminated, still billed or paid", "co_do": "CO / DO changed", "anomaly": "On the anomaly report"}
STATES = {"posted": "Posted", "read": "Bill read", "estimated": "Estimated", "missing": "Missing"}
MOST_EXCEPTIONS = 2000  # one import keeps at most this many (the counts say how many there were)

# 1.0.104: this module's portfolio settings (Settings → Portfolios lists them with these words)
SETTINGS = {"default_bill_day": (20, "Usual bill day when unknown", "The day of the month a bill is expected when the account has no bill day and no history.")}
for _k, (_d, _l, _h) in SETTINGS.items():
    PF._SEEN.setdefault(_k, copy.deepcopy(_d))
    PF.DESCRIBE.setdefault(_k, (_l, _h, ()))


class Refused(Exception):
    """A plain-words refusal (400)."""


def _who(u: Optional[User]) -> str:
    return (u.name or u.email) if u else ""


def _dmy(iso_day: str) -> str:
    try:
        return date.fromisoformat(iso_day).strftime("%d.%m.%Y")
    except (TypeError, ValueError):
        return iso_day or "?"


def _iso(v: Any) -> str:
    from .history import iso

    return iso(v)


def _acct(v: Any) -> str:
    from . import masters as M

    return M.acct_key(v)


def this_month() -> str:
    return date.today().isoformat()[:7]


def period_ok(period: str) -> str:
    p = str(period or "").strip()[:7]
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", p):
        raise Refused("A month as yyyy-mm.")
    return p


# ---- who may change accounts and import lists -------------------------------------------------------------------------------
def may_manage(db: Session, user: Optional[User], portfolio: str) -> bool:
    """Admins, and reviewers of the portfolio (bootstrap: everyone with write access holds reviewer)."""
    from . import roles as R

    return bool(user is not None and not user.is_axl and (user.is_admin or R.may(db, user, "reviewer", portfolio)))


def _require_manage(db: Session, user: Optional[User], portfolio: str, what: str) -> None:
    if not may_manage(db, user, portfolio):
        from . import roles as R

        raise R.Forbidden(f"{what} needs the reviewer role in {portfolio or 'the portfolio'} (or an admin).")


# ---- the accounts of a portfolio ------------------------------------------------------------------------------------------
class Portfolios:
    """Which portfolio each utility account is in (its own, else its company code's, else the default) without a query per account."""

    def __init__(self, db: Session):
        self.ps = PF.portfolios(db)
        d = PF.default(db)
        self.default = d.code if d is not None else ""
        self.codes = {p.code for p in self.ps}
        self._co: dict[str, str] = {}

    def of(self, a: ApUtilityAccount) -> str:
        if a.portfolio and a.portfolio in self.codes:
            return a.portfolio
        co = (a.company_code or "").strip().upper()
        if co not in self._co:
            best = sorted(((PF.claims(p, co), -p.id, p.code) for p in self.ps), reverse=True)
            self._co[co] = best[0][2] if best and best[0][0] else self.default
        return self._co[co]


def accounts(db: Session, portfolio: str, cluster: str = "") -> list[ApUtilityAccount]:
    pf = Portfolios(db)
    out = [a for a in db.query(ApUtilityAccount).order_by(ApUtilityAccount.id).all() if pf.of(a) == portfolio]
    if cluster:
        out = [a for a in out if (a.cluster or "") == cluster]
    return out


def status_of(a: ApUtilityAccount) -> str:
    return (a.status or "").strip().lower() or "active"


def _closed_before(a: ApUtilityAccount, day: str) -> bool:
    """Terminated or disposed on or before this day (no date: closed)."""
    return status_of(a) in CLOSED and (not a.status_from or a.status_from <= day)


# ---- an account's fields (a person's change) ---------------------------------------------------------------------------------
def norm_status(v: Any) -> str:
    s = str(v or "").strip().lower()
    if not s:
        return ""
    if re.search(r"terminat|closed|cancel|deactivat|inactive", s):
        return "terminated"
    if re.search(r"dispos|sold|divest|transfer", s):
        return "disposed"
    if re.search(r"suspend|disconnect|cut", s):
        return "suspended"
    if re.search(r"active|live|open|current|operational|connected", s):
        return "active"
    return ""


def norm_co_do(v: Any) -> str:
    s = re.sub(r"[^a-z]", "", str(v or "").lower())
    if s in ("co", "coco", "companyowned", "companyoperated", "company"):
        return "CO"
    if s in ("do", "codo", "dealerowned", "dealeroperated", "dealer", "dodo"):
        return "DO"
    return ""


def norm_bool(v: Any) -> Optional[bool]:
    s = str(v if v is not None else "").strip().lower()
    if s in ("y", "yes", "true", "1", "x", "ja"):
        return True
    if s in ("n", "no", "false", "0", "nee"):
        return False
    return None


def norm_day(v: Any) -> Optional[int]:
    try:
        n = int(float(str(v).strip()))
    except (TypeError, ValueError):
        return None
    return n if 1 <= n <= 31 else None


def account_out(a: ApUtilityAccount, portfolio: str = "") -> dict:
    return {"id": a.id, "account": a.account, "account_key": a.account_key, "vendor": a.vendor, "site": a.site, "municipality": a.municipality,
            "company_code": a.company_code, "portfolio": portfolio or a.portfolio or "", "status": status_of(a), "status_from": a.status_from or "",
            "status_source": a.status_source or "", "co_do": a.co_do or "", "cluster": a.cluster or "", "check_meter": bool(a.check_meter), "bill_day": a.bill_day}


def save_account(db: Session, user: User, aid: int, data: dict) -> ApUtilityAccount:
    """A person sets an account's status (with its date), CO/DO, cluster, check meter, bill day or portfolio. Audited
    ap_account_status (what changed)."""
    a = db.get(ApUtilityAccount, int(aid or 0))
    if a is None:
        raise LookupError("No such account.")
    _require_manage(db, user, Portfolios(db).of(a), "Changing an account")
    said: list[str] = []
    data = data or {}
    if "status" in data:
        st = str(data.get("status") or "").strip().lower()
        if st not in STATUSES:
            raise Refused(f"Status is one of {', '.join(STATUSES)}.")
        frm = _iso(data.get("status_from")) if data.get("status_from") else ""
        if st in CLOSED and not frm:
            raise Refused("Give the day the account was terminated or disposed (yyyy-mm-dd).")
        if (st, frm) != (status_of(a), a.status_from or ""):
            said.append(f"status {status_of(a)}{' from ' + a.status_from if a.status_from else ''} → {st}{' from ' + frm if frm else ''}")
            a.status, a.status_from, a.status_source = st, frm, _who(user)[:80]
    if "co_do" in data:
        v = norm_co_do(data.get("co_do")) if str(data.get("co_do") or "").strip() else ""
        if str(data.get("co_do") or "").strip() and not v:
            raise Refused("CO/DO is CO or DO.")
        if v != (a.co_do or ""):
            said.append(f"CO/DO {a.co_do or '-'} → {v or '-'}")
            a.co_do = v
    if "cluster" in data:
        v = re.sub(r"\s+", " ", str(data.get("cluster") or "")).strip()[:20]
        if v != (a.cluster or ""):
            said.append(f"cluster {a.cluster or '-'} → {v or '-'}")
            a.cluster = v
    if "check_meter" in data:
        v = bool(data.get("check_meter"))
        if v != bool(a.check_meter):
            said.append(f"check meter {'yes' if v else 'no'}")
            a.check_meter = v
    if "bill_day" in data:
        raw = data.get("bill_day")
        v = norm_day(raw) if raw not in (None, "") else None
        if raw not in (None, "") and v is None:
            raise Refused("The bill day is a day of the month (1 to 31).")
        if v != a.bill_day:
            said.append(f"bill day {a.bill_day or 'learned'} → {v or 'learned'}")
            a.bill_day = v
    if "portfolio" in data:
        code = PF.pick(db, data.get("portfolio")) if str(data.get("portfolio") or "").strip() else ""
        if code != (a.portfolio or ""):
            _require_manage(db, user, code or Portfolios(db).of(a), "Moving an account")
            said.append(f"portfolio {a.portfolio or 'from its company code'} → {code or 'from its company code'}")
            a.portfolio = code
    if said:
        a.updated_by = _who(user)
        db.flush()
        log(db, "ap_account_status", actor=user, target=f"ap_account:{a.account_key}", detail=f"{a.account} {a.site}: " + "; ".join(said)[:900])
    return a


# ---- reading a masterlist -------------------------------------------------------------------------------------------------------
ML_COLS: dict[str, tuple[str, ...]] = {
    "account": ("account reference", "account number", "account no", "account no.", "account", "acc no", "municipal account", "utility account", "reference"),
    "vendor": ("vendor/bp number", "vendor / bp number", "vendor number", "bp number", "vendor", "supplier", "supplier number", "lifnr"),
    "site": ("site name", "name1", "site", "site code", "site id"),
    "municipality": ("municipality/payee", "municipality / payee", "municipality", "payee", "vendor name", "supplier name"),
    "status": ("status", "account status", "site status", "state"),
    "status_from": ("effective date", "termination date", "date terminated", "terminated on", "disposal date", "status date", "date of termination"),
    "co_do": ("co/do", "co / do", "co do", "co-do", "ownership", "operating model", "site type"),
    "cluster": ("cluster", "cluster no", "cluster number"),
    "check_meter": ("check meter", "check-meter", "check meter (y/n)", "bcx meter"),
    "bill_day": ("bill day", "usual bill day", "billing day"),
    "company_code": ("company code", "bukrs", "cocd"),
    "note": ("comment", "comments", "anomaly", "remarks", "note", "notes", "reason"),
}
ML_MOST = 60000


def _head(v: Any) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip().lower()


def read_list(data: bytes, filename: str = "") -> list[dict]:
    """A masterlist / Termination / Anomaly report (Excel or CSV): the header row (in the first 15 rows) needs an account
    column; the other columns are taken when the sheet has them. → [{field: text}] (account_key added)."""
    from . import masters as M

    try:
        rows = M._raw_rows(data, filename)
    except M.Refused as e:
        raise Refused(str(e)) from None
    for i, r in enumerate(rows[:15]):
        norm = [_head(h) for h in r]
        cols: dict[str, int] = {}
        for f, names in ML_COLS.items():
            for nm in names:
                if nm in norm and norm.index(nm) not in cols.values():
                    cols[f] = norm.index(nm)
                    break
        if "account" not in cols:
            continue
        out = []
        for raw in rows[i + 1:]:
            rec = {f: (raw[c] if c < len(raw) else None) for f, c in cols.items()}
            k = _acct(rec.get("account"))
            if not k:
                continue
            got = {f: K.text(v, 160) for f, v in rec.items() if f not in ("status_from",)}
            got["status_from"] = _iso(rec.get("status_from")) if rec.get("status_from") not in (None, "") else ""
            got["account_key"] = k[:40]
            out.append(got)
            if len(out) > ML_MOST:
                raise Refused(f"More than {ML_MOST} rows: split the list.")
        if not out:
            raise Refused("The list has an account column but no account numbers under it.")
        return out
    raise Refused("No header row with an account column (Account number, Account reference …) in the first 15 rows.")


# ---- comparing and taking a masterlist ---------------------------------------------------------------------------------------
def _sap_vendors(db: Session) -> Optional[set[str]]:
    """Every SAP vendor number (all the ways SAP writes it), or None when SAP's vendor list is not loaded."""
    if not db.query(ApSapVendor.id).first():
        return None
    out: set[str] = set()
    for (v,) in db.query(ApSapVendor.vendor).all():
        out |= K.vendor_variants(v)
    return out


def _after(db: Session, a: ApUtilityAccount, day: str) -> list[str]:
    """Bills read and payments made for the account after this day (yyyy-mm-dd): what says it is still billed or paid."""
    out = []
    for h in db.query(ApBillHistory).filter(ApBillHistory.account_key == a.account_key).all():
        d = h.read_start or h.bill_date
        if d and d > day:
            out.append(f"bill {h.invoice_no or h.period} of {_dmy(h.bill_date or d)}")
    for p in db.query(ApPayment).filter(ApPayment.posting_date > day).all():
        if a.account_key in (p.allocations or {}):
            out.append(f"payment {p.doc_number} of {_dmy(p.posting_date)}")
    return out


def _plan(db: Session, portfolio: str, kind: str, rows: list[dict]) -> dict:
    """What an import would change and which exceptions it raises (nothing is written)."""
    pf = Portfolios(db)
    by_key: dict[str, list[ApUtilityAccount]] = {}
    mine: dict[int, ApUtilityAccount] = {}
    for a in db.query(ApUtilityAccount).all():
        by_key.setdefault(a.account_key, []).append(a)
        if pf.of(a) == portfolio:
            mine[a.id] = a
    sap = _sap_vendors(db)
    changes: list[tuple[ApUtilityAccount, dict, list[str]]] = []
    exc: list[dict] = []
    seen: set[int] = set()
    listed: set[str] = set()
    today = date.today().isoformat()

    def ex(kind_: str, r: dict, text: str, a: Optional[ApUtilityAccount] = None) -> None:
        exc.append({"kind": kind_, "account_key": (a.account_key if a else r.get("account_key", ""))[:40], "account": (a.account if a else r.get("account", ""))[:60],
                    "vendor": (a.vendor if a else r.get("vendor", ""))[:40], "site": (a.site if a else r.get("site", ""))[:160], "text": text[:400]})

    for r in rows:
        k = r["account_key"]
        if k in listed:
            continue
        listed.add(k)
        hits = by_key.get(k) or []
        if r.get("vendor") and sap is not None and not (K.vendor_variants(r["vendor"]) & sap):
            ex("not_in_sap", r, f"Vendor {r['vendor']} of account {r.get('account') or k} is not in SAP's vendor list.")
        if not hits:
            ex("not_in_cbx", r, f"Account {r.get('account') or k} ({r.get('site') or r.get('municipality') or 'no site'}) is on the list but not in CBX's utilities list.")
            continue
        st = norm_status(r.get("status")) if r.get("status") else ""
        if kind == "terminations" and not st:
            st = "terminated"
        for a in hits:
            seen.add(a.id)
            want: dict[str, Any] = {}
            said: list[str] = []
            if st and (st, r.get("status_from") or "") != (status_of(a), a.status_from or ""):
                frm = r.get("status_from") or (a.status_from if st == status_of(a) else "") or (today if st in CLOSED else "")
                if (st, frm) != (status_of(a), a.status_from or ""):
                    want.update(status=st, status_from=frm)
                    said.append(f"status {status_of(a)} → {st}{' from ' + frm if frm else ''}")
                    if st in CLOSED and status_of(a) not in CLOSED:
                        ex("terminated_active", r, f"Account {a.account} ({a.site}) is {st} on the list{' from ' + _dmy(frm) if frm else ''} but was active in CBX.", a)
                    if st in CLOSED:
                        late = _after(db, a, frm)
                        if late:
                            ex("terminated_billed", r, f"Account {a.account} ({a.site}) is {st} from {_dmy(frm)} but still billed or paid: {', '.join(late[:4])}.", a)
            co = norm_co_do(r.get("co_do")) if r.get("co_do") else ""
            if co and co != (a.co_do or ""):
                want["co_do"] = co
                said.append(f"CO/DO {a.co_do or '-'} → {co}")
                if a.co_do:
                    ex("co_do", r, f"Account {a.account} ({a.site}): {a.co_do} → {co}.", a)
            cl = re.sub(r"\s+", " ", r.get("cluster") or "").strip()[:20]
            if cl and cl != (a.cluster or ""):
                want["cluster"] = cl
                said.append(f"cluster {a.cluster or '-'} → {cl}")
            cm = norm_bool(r.get("check_meter")) if r.get("check_meter") else None
            if cm is not None and cm != bool(a.check_meter):
                want["check_meter"] = cm
                said.append(f"check meter {'yes' if cm else 'no'}")
            bd = norm_day(r.get("bill_day")) if r.get("bill_day") else None
            if bd and bd != a.bill_day:
                want["bill_day"] = bd
                said.append(f"bill day {bd}")
            if not a.portfolio and pf.of(a) != portfolio:
                want["portfolio"] = portfolio
                said.append(f"portfolio {pf.of(a)} → {portfolio}")
            if kind == "anomalies":
                ex("anomaly", r, f"Account {a.account} ({a.site}) is on the anomaly report" + (f": {r['note']}" if r.get("note") else (f" ({r['status']})" if r.get("status") else "")) + ".", a)
            if want:
                changes.append((a, want, said))
    if kind == "masterlist":
        for a in mine.values():
            if a.id not in seen and status_of(a) not in CLOSED:
                ex("not_on_list", {}, f"Account {a.account} ({a.site}, {a.municipality}) is in CBX but not on the list.", a)
    counts = {"rows": len(rows), "accounts": len(listed), "matched": len(seen), "changed": len(changes)}
    for k in EXCEPTIONS:
        counts[k] = sum(1 for e in exc if e["kind"] == k)
    return {"changes": changes, "exceptions": exc, "counts": counts}


def import_list(db: Session, user: User, portfolio: str, kind: str, rows: list[dict], *, source: str = "", file_name: str = "", confirm: bool = False) -> dict:
    """A masterlist refresh for a portfolio. confirm=False: what would change and the exceptions it would raise; True:
    the accounts take what the list says, the exceptions are opened, audited ap_masterlist_import."""
    p = PF.get(db, portfolio)
    if p is None:
        raise Refused(f"{portfolio or '(none)'} is not a portfolio.")
    if kind not in KINDS:
        raise Refused(f"The list is one of: {', '.join(KINDS)}.")
    _require_manage(db, user, p.code, "Importing a masterlist")
    if not rows:
        raise Refused("No rows with an account number in the list.")
    plan = _plan(db, p.code, kind, rows)
    out = {"portfolio": p.code, "kind": kind, "counts": plan["counts"], "saved": False,
           "changes": [{"account": a.account, "site": a.site, "what": "; ".join(s)} for a, _w, s in plan["changes"][:20]],
           "exceptions": [{k: e[k] for k in ("kind", "account", "text")} for e in plan["exceptions"][:20]], "labels": EXCEPTIONS}
    if not confirm:
        return out
    who = _who(user)
    src = (source or KINDS[kind]).strip()[:120]
    stamp = f"{src} {date.today().isoformat()}"[:80]
    for a, want, _s in plan["changes"]:
        for f, v in want.items():
            setattr(a, f, v)
        if "status" in want:
            a.status_source = stamp
        a.updated_by = who
    imp = ApMasterlistImport(portfolio=p.code, kind=kind, source=src, file_name=(file_name or "")[:255], rows=len(rows), counts=plan["counts"], by=who, by_id=user.id)
    db.add(imp)
    db.flush()
    t = ApMasterlistException.__table__
    keep = plan["exceptions"][:MOST_EXCEPTIONS]
    if keep:
        db.execute(t.insert(), [{**e, "import_id": imp.id, "portfolio": p.code, "status": "open", "note": "", "resolved_by": "", "created_at": utcnow()} for e in keep])
    db.flush()
    log(db, "ap_masterlist_import", actor=user, target=f"ap_portfolio:{p.code}",
        detail=f"{KINDS[kind]} '{src}' ({file_name or 'no file name'}) refreshed on {date.today().isoformat()}: {len(rows)} rows, "
               f"{plan['counts']['changed']} accounts changed, {len(plan['exceptions'])} exceptions"
               + (f" ({len(keep)} kept)" if len(keep) < len(plan["exceptions"]) else ""))
    out["saved"], out["import_id"] = True, imp.id
    return out


def imports(db: Session, portfolio: str, limit: int = 12) -> list[dict]:
    return [{"id": i.id, "kind": i.kind, "kind_label": KINDS.get(i.kind, i.kind), "source": i.source, "file_name": i.file_name, "rows": i.rows,
             "counts": dict(i.counts or {}), "by": i.by, "at": i.created_at.isoformat(timespec="minutes") if i.created_at else ""}
            for i in db.query(ApMasterlistImport).filter(ApMasterlistImport.portfolio == portfolio).order_by(ApMasterlistImport.id.desc()).limit(limit).all()]


def exception_out(e: ApMasterlistException) -> dict:
    return {"id": e.id, "kind": e.kind, "label": EXCEPTIONS.get(e.kind, e.kind), "account": e.account, "account_key": e.account_key, "vendor": e.vendor,
            "site": e.site, "text": e.text, "status": e.status, "note": e.note, "resolved_by": e.resolved_by,
            "resolved_at": e.resolved_at.isoformat(timespec="minutes") if e.resolved_at else "", "at": e.created_at.isoformat(timespec="minutes") if e.created_at else ""}


def exceptions(db: Session, portfolio: str, status: str = "open", kind: str = "", limit: int = 300) -> dict:
    q = db.query(ApMasterlistException).filter(ApMasterlistException.portfolio == portfolio)
    if status in ("open", "resolved"):
        q = q.filter(ApMasterlistException.status == status)
    if kind:
        q = q.filter(ApMasterlistException.kind == kind)
    total = q.count()
    return {"total": total, "rows": [exception_out(e) for e in q.order_by(ApMasterlistException.id.desc()).limit(max(1, min(int(limit or 300), 2000))).all()]}


def resolve(db: Session, user: User, eid: int, note: str) -> ApMasterlistException:
    """A person resolves an exception with a note. Audited ap_masterlist_exception."""
    e = db.get(ApMasterlistException, int(eid or 0))
    if e is None:
        raise LookupError("No such exception.")
    _require_manage(db, user, e.portfolio, "Resolving a masterlist exception")
    note = re.sub(r"\s+", " ", str(note or "")).strip()[:1000]
    if len(note) < 3:
        raise Refused("Say how it was resolved.")
    if e.status == "resolved":
        raise Refused("It is resolved already.")
    e.status, e.note, e.resolved_by, e.resolved_by_id, e.resolved_at = "resolved", note, _who(user), user.id, utcnow()
    db.flush()
    log(db, "ap_masterlist_exception", actor=user, target=f"ap_account:{e.account_key}", detail=f"{EXCEPTIONS.get(e.kind, e.kind)} — {e.account}: resolved: {note}"[:900])
    return e


# ---- the completeness view ---------------------------------------------------------------------------------------------------
def _chunks(keys: list[str], n: int = 900) -> Iterable[list[str]]:
    for i in range(0, len(keys), n):
        yield keys[i:i + n]


def _learned_days(db: Session, keys: list[str]) -> dict[str, int]:
    """The usual bill day per account from its history (the median day of its bill dates)."""
    days: dict[str, list[int]] = {}
    for ch in _chunks(keys):
        for k, bd in db.query(ApBillHistory.account_key, ApBillHistory.bill_date).filter(ApBillHistory.account_key.in_(ch), ApBillHistory.bill_date != "").all():
            try:
                days.setdefault(k, []).append(int(bd[8:10]))
            except (TypeError, ValueError):
                continue
    return {k: int(median(v[-6:])) for k, v in days.items() if v}


def expected_day(period: str, day: int) -> date:
    y, m = int(period[:4]), int(period[5:7])
    return date(y, m, min(max(1, day), calendar.monthrange(y, m)[1]))


def month_view(db: Session, portfolio: str, period: str, *, cluster: str = "", state: str = "", q: str = "", limit: int = 300, today: Optional[date] = None) -> dict:
    """Every account of the portfolio that should bill in the month: its state (posted · read · estimated · missing), the
    bill, the payments of the month; missing ones with their age and the day they count as late."""
    from . import payments as PAY

    p = PF.get(db, portfolio)
    if p is None:
        raise LookupError("No such portfolio.")
    period = period_ok(period)
    today = today or date.today()
    grace = int(PF.setting(db, p, "missing_grace_days", 10) or 0)
    dflt_day = int(PF.setting(db, p, "default_bill_day", 20) or 20)
    month_start = f"{period}-01"
    accts = [a for a in accounts(db, p.code) if not _closed_before(a, month_start)]
    clusters = sorted({a.cluster or "" for a in accts})
    if cluster:
        accts = [a for a in accts if (a.cluster or "") == cluster]
    keys = sorted({a.account_key for a in accts if a.account_key})
    bills: dict[str, list[ApBillHistory]] = {}
    for ch in _chunks(keys):
        for h in db.query(ApBillHistory).filter(ApBillHistory.period == period, ApBillHistory.account_key.in_(ch)).all():
            bills.setdefault(h.account_key, []).append(h)
    cap_ids = sorted({h.capture_id for hs in bills.values() for h in hs if h.capture_id})
    uploaded: dict[int, Optional[datetime]] = {}
    for i in range(0, len(cap_ids), 900):
        uploaded.update({cid: up for cid, up in db.query(ApCapture.id, ApCapture.uploaded_at).filter(ApCapture.id.in_(cap_ids[i:i + 900])).all()})
    ests: dict[str, ApEstimate] = {}
    for ch in _chunks(keys):
        for e in db.query(ApEstimate).filter(ApEstimate.period == period, ApEstimate.account_key.in_(ch), ApEstimate.status.in_(("prepared", "approved", "trued_up"))).all():
            if e.account_key not in ests or e.status != "prepared":
                ests[e.account_key] = e
    paid = PAY.paid_in_month(db, keys, period)
    payments_loaded = PAY.loaded(db, period, period)
    learned = _learned_days(db, [a.account_key for a in accts if not a.bill_day])
    rows = []
    for a in accts:
        hs = sorted(bills.get(a.account_key, []), key=lambda h: (h.bill_date or "", h.id))
        e = ests.get(a.account_key)
        day = a.bill_day or learned.get(a.account_key) or dflt_day
        exp = expected_day(period, day)
        r = {**account_out(a, p.code), "state": "", "bill_day": day, "bill_day_how": "set" if a.bill_day else ("learned" if a.account_key in learned else "default"),
             "expected": exp.isoformat(), "late_from": (exp + timedelta(days=grace)).isoformat(), "age_days": None, "late": False, "bill": None, "estimate": None,
             "paid": paid.get(a.account_key) if payments_loaded else None, "after_termination": False}
        if hs:
            h = hs[-1]
            up = any(uploaded.get(x.capture_id) for x in hs if x.capture_id)
            r["state"] = "posted" if up else "read"
            r["bill"] = {"invoice_no": h.invoice_no, "bill_date": h.bill_date, "total_excl": h.total_excl, "vat": h.vat, "closing": h.closing,
                         "reading_type": h.reading_type, "capture_id": h.capture_id, "bills": len(hs)}
            r["after_termination"] = status_of(a) in CLOSED and bool(a.status_from) and (h.read_start or h.bill_date or "") > a.status_from
        elif e is not None and e.status != "prepared":
            r["state"] = "estimated"
            r["estimate"] = {"id": e.id, "value": e.value, "status": e.status, "basis": e.basis}
        else:
            r["state"] = "missing"
            r["age_days"] = max(0, (today - exp).days)
            r["late"] = today > exp + timedelta(days=grace)
            if e is not None:
                r["estimate"] = {"id": e.id, "value": e.value, "status": e.status, "basis": e.basis}
        rows.append(r)
    summary = {s: sum(1 for r in rows if r["state"] == s) for s in STATES}
    summary["accounts"] = len(rows)
    summary["late"] = sum(1 for r in rows if r["late"])
    summary["complete_pct"] = round(100.0 * (len(rows) - summary["missing"]) / len(rows), 1) if rows else 100.0
    by_cluster: dict[str, dict] = {}
    for r in rows:
        c = by_cluster.setdefault(r["cluster"] or "", {"cluster": r["cluster"] or "", "accounts": 0, **{s: 0 for s in STATES}, "late": 0})
        c["accounts"] += 1
        c[r["state"]] += 1
        c["late"] += r["late"]
    ql = re.sub(r"\s+", " ", str(q or "")).strip().lower()[:60]
    shown = [r for r in rows if (not state or r["state"] == state or (state == "late" and r["late"])) and
             (not ql or ql in f"{r['account']} {r['site']} {r['municipality']} {r['vendor']}".lower())]
    shown.sort(key=lambda r: (0 if r["late"] else 1, list(STATES).index(r["state"]) * -1, -(r["age_days"] or 0), r["account"]))
    return {"portfolio": p.code, "period": period, "cluster": cluster, "clusters": clusters, "summary": summary, "by_cluster": sorted(by_cluster.values(), key=lambda c: c["cluster"]),
            "payments_loaded": payments_loaded, "grace_days": grace, "total": len(shown), "rows": shown[:max(1, min(int(limit or 300), 20000))], "states": STATES}


def liaison(db: Session, portfolio: str, period: str, cluster: str = "", today: Optional[date] = None) -> list[dict]:
    """The Council Liaison list: missing bills past the usual day + the grace days, oldest first."""
    v = month_view(db, portfolio, period, cluster=cluster, state="late", limit=20000, today=today)
    return sorted(v["rows"], key=lambda r: -(r["age_days"] or 0))


# ---- the weekly workbook ----------------------------------------------------------------------------------------------------
def export(db: Session, portfolio: str, period: str, cluster: str = "") -> bytes:
    """The completeness workbook (weekly; replaces the offline Cluster Portfolio Tracker and Shell's manual tracker):
    Summary · Accounts · Council Liaison · Exceptions."""
    import openpyxl

    from .. import xlsafe

    v = month_view(db, portfolio, period, cluster=cluster, limit=20000)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    s = v["summary"]
    ws.append(["Portfolio", v["portfolio"]])
    ws.append(["Month", v["period"]])
    ws.append(["Cluster", cluster or "all"])
    ws.append(["Made", datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws.append([])
    ws.append(["Cluster", "Accounts", "Posted", "Bill read", "Estimated", "Missing", "Late (Council Liaison)"])
    for c in v["by_cluster"]:
        ws.append([c["cluster"] or "(none)", c["accounts"], c["posted"], c["read"], c["estimated"], c["missing"], c["late"]])
    ws.append(["All", s["accounts"], s["posted"], s["read"], s["estimated"], s["missing"], s["late"]])
    ws.append(["Complete %", s["complete_pct"]])
    a = wb.create_sheet("Accounts")
    a.append(["Account", "Site", "Municipality", "Vendor", "Company code", "Cluster", "CO/DO", "Status", "Status from", "Check meter", "State", "Bill date",
               "Invoice", "Amount excl. VAT", "Closing balance", "Reading", "Estimate", "Paid in the month", "Expected by", "Late from", "Age (days)"])
    for r in v["rows"]:
        b, e = r["bill"] or {}, r["estimate"] or {}
        a.append([r["account"], r["site"], r["municipality"], r["vendor"], r["company_code"], r["cluster"], r["co_do"], r["status"], r["status_from"],
                  "yes" if r["check_meter"] else "", STATES[r["state"]] + (" (late)" if r["late"] else ""), b.get("bill_date", ""), b.get("invoice_no", ""),
                  b.get("total_excl"), b.get("closing"), b.get("reading_type", ""), e.get("value"), r["paid"], r["expected"], r["late_from"], r["age_days"]])
    li = wb.create_sheet("Council Liaison")
    li.append(["Account", "Site", "Municipality", "Vendor", "Cluster", "Expected by", "Late from", "Age (days)", "Estimate"])
    for r in sorted((r for r in v["rows"] if r["late"]), key=lambda r: -(r["age_days"] or 0)):
        li.append([r["account"], r["site"], r["municipality"], r["vendor"], r["cluster"], r["expected"], r["late_from"], r["age_days"],
                   (r["estimate"] or {}).get("status", "")])
    ex = wb.create_sheet("Exceptions")
    ex.append(["Opened", "Kind", "Account", "Vendor", "Site", "What"])
    for e in exceptions(db, v["portfolio"], "open", limit=2000)["rows"]:
        ex.append([e["at"], e["label"], e["account"], e["vendor"], e["site"], e["text"]])
    xlsafe.book(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- Tess's hook -------------------------------------------------------------------------------------------------------------
def tess_note(db: Session, c: ApCapture) -> tuple[list[str], list[str]]:
    """(red reasons, yellow reasons) for a utility bill on a terminated, disposed or suspended account."""
    if c.bucket != "utility":
        return [], []
    aid = ((c.match or {}).get("account") or {}).get("id")
    a = db.get(ApUtilityAccount, int(aid)) if aid else None
    if a is None:
        return [], []
    st = status_of(a)
    if st in CLOSED:
        f = c.fields or {}
        end = _iso(f.get("read_end"))
        if a.status_from and end and end <= a.status_from:
            return [], [f"account {st} on {_dmy(a.status_from)}: final bill? (the reading ends {_dmy(end)})"]
        return [f"account {st} on {_dmy(a.status_from) if a.status_from else '(no date)'}"], []
    if st == "suspended":
        return [], [f"account suspended{' since ' + _dmy(a.status_from) if a.status_from else ''}"]
    return [], []


def portfolios_out(db: Session, user: Optional[User]) -> list[dict]:
    return [{"code": p.code, "name": p.name, "manage": may_manage(db, user, p.code)} for p in PF.portfolios(db)]
