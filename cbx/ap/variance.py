# 1.0.104: new (Builder U) — V07: consumption and tariff checks on utility bills, with evidence (VALUENCE-AP-GAPS.md §2 V07)
"""Tess compares every utility bill with the account's earlier bills (cbx/ap/bill_history.py) and the tariff table
(cbx/ap/tariffs.py). What she finds makes the bill yellow, with the figures used:

    VAR_R      the amount excl. VAT against the average of the last `compare_periods` bills, beyond `variance_pct`
    VAR_KWH    the kWh the same way
    RATE       the effective R/kWh against the account's own last 6 bills beyond `rate_pct` …
    TARIFF     … or against the tariff table when it has a row for the bill (the charge the tariff gives for the kWh)
    EST_RUN    an estimated reading (E) `estimate_run_n` or more bills running
    GAP        the reading starts more than `period_gap_days` after the last bill's reading ended (a gap) …
    OVERLAP    … or before it ended (an overlap)
    METER      a check-meter reading (the BCX CSV) this far (`check_meter_pct`) from the billed kWh

A person then decides (reviewer role, not someone who worked on the bill):
    Accept     with a reason (stored on the history row as evidence; audited ap_variance_accepted) → the yellow goes
    Dispute    (capturer role) an `ap_disputes` row: amount, reason, municipal reference, Council Liaison owner — the
               bill is red, held out of the SAP upload, until the dispute is resolved or "pay under dispute" is ticked
               (audited ap_dispute)

The comparison (periods used, expected, actual, who accepted) goes into the review workbook (`with_comparison`).
Check-meter readings come from a CSV upload (`upload_check_meters`); `no_meter_ranking` ranks the sites without one.
Settings (Settings → Portfolios): variance_pct, rate_pct, compare_periods (F0a), and variance_min_bills,
estimate_run_n, period_gap_days, check_meter_pct (added here through portfolio.setting)."""
from __future__ import annotations

import csv
import io
import re
from datetime import datetime
from typing import Any, Optional

from sqlalchemy.orm import Session

from ..audit import log
from ..db import utcnow
from ..models import User
from ..models.ap import ApBatch, ApBillHistory, ApCapture, ApCheckMeter, ApDispute, ApItem, ApUtilityAccount
from . import bill_history as BH
from . import portfolio as PF
from . import roles as R
from . import tariffs as T

RATE_PERIODS = 6  # "against the account's own last 6 months" (V07 item 2)
NEW_KEYS = {"variance_min_bills": 2,  # V07: earlier bills needed before the average means anything
            "estimate_run_n": 2,  # V07: estimated readings running before the bill is yellow
            "period_gap_days": 1,  # V07: days allowed between the last bill's reading end and this one's start
            "check_meter_pct": 10}  # V07: check-meter kWh against billed kWh (Telkom's BCX readings)
for _k, _v in NEW_KEYS.items():  # offered in Settings → Portfolios with these defaults (portfolio.setting registers them)
    PF.setting(None, None, _k, _v)
for _k, _d in {"variance_min_bills": ("Bills needed before comparing", "Earlier bills an account needs before its average is used.", ()),
               "estimate_run_n": ("Estimated readings running", "A bill whose reading is estimated this many bills running is checked.", ()),
               "period_gap_days": ("Reading gap allowed (days)", "Days between the last bill's reading end and this bill's start before it is a gap.", ()),
               "check_meter_pct": ("Check meter variance %", "Billed kWh this far from the check meter's reading is checked.", ())}.items():
    PF.DESCRIBE.setdefault(_k, _d)  # the words Settings → Portfolios shows beside the key


class Refused(Exception):
    pass


def _who(u: Optional[User]) -> str:
    return (u.name or u.email) if u else ""


def _set(db: Session, port: str, key: str) -> Any:
    v = PF.setting(db, port or None, key, NEW_KEYS.get(key))
    return v


def _n(db: Session, port: str, key: str) -> float:
    try:
        return float(_set(db, port, key))
    except (TypeError, ValueError):
        return float(PF.DEFAULTS.get(key, NEW_KEYS.get(key, 0)) or 0)


def money(v: Optional[float]) -> str:
    return "R {:,.2f}".format(v or 0).replace(",", " ")


def qty(v: Optional[float]) -> str:
    return "{:,.0f}".format(v or 0).replace(",", " ") if v is not None and abs(v) >= 100 else "{:g}".format(round(v or 0, 2))


def _pct(cur: float, base: float) -> float:
    return (cur - base) / abs(base) * 100


def _word(p: float) -> str:
    return f"{abs(p):.0f}% {'above' if p > 0 else 'below'}"


# ---- the checks --------------------------------------------------------------------------------------------------------------
def _account(db: Session, c: Optional[ApCapture], h: ApBillHistory) -> Optional[ApUtilityAccount]:
    aid = (((c.match or {}) if c is not None else {}).get("account") or {}).get("id")
    a = db.get(ApUtilityAccount, int(aid)) if aid else None
    return a or db.query(ApUtilityAccount).filter(ApUtilityAccount.account_key == h.account_key).order_by(ApUtilityAccount.id.desc()).first()


def meter_for(db: Session, h: ApBillHistory, a: Optional[ApUtilityAccount]) -> Optional[ApCheckMeter]:
    """The check-meter reading for the bill's account (or its site's name or cost centre) and period."""
    q = db.query(ApCheckMeter).filter(ApCheckMeter.period == h.period)
    m = q.filter(ApCheckMeter.account_key == h.account_key).first() if h.account_key else None
    if m is None and a is not None:
        keys = _site_keys(a)
        m = next((x for x in q.all() if T.key(x.site) in keys), None) if keys else None
    return m


def compute(db: Session, h: ApBillHistory, c: Optional[ApCapture] = None) -> list[dict]:
    """The findings for one history row: [{code, text, figures}] (empty when the bill is in line or nothing compares)."""
    port = h.portfolio or ""
    n_cmp = max(1, int(_n(db, port, "compare_periods")))
    vp, rp, minb = _n(db, port, "variance_pct"), _n(db, port, "rate_pct"), max(1, int(_n(db, port, "variance_min_bills")))
    prev = BH.rows_of(db, h.account_key, before=h)
    last = prev[-n_cmp:]
    out: list[dict] = []
    used = [p.period for p in last]

    # amount and kWh against the average of the last N bills
    for code, attr, show, label in (("VAR_R", "total_excl", money, "amount excl. VAT"), ("VAR_KWH", "kwh", lambda v: f"{qty(v)} kWh", "kWh")):
        cur = getattr(h, attr)
        vals = [getattr(p, attr) for p in last if getattr(p, attr) is not None]
        if cur is None or len(vals) < minb:
            continue
        avg = sum(vals) / len(vals)
        if not avg:
            continue
        p = _pct(cur, avg)
        if abs(p) > vp:
            out.append({"code": code, "text": f"{label} {show(cur)} is {_word(p)} the average of the last {len(vals)} bills ({show(avg)}; allowed {vp:g}%)",
                        "figures": {"actual": round(cur, 2), "expected": round(avg, 2), "pct": round(p, 1), "limit": vp, "periods": used[-len(vals):]}})

    # the rate: the tariff table first, else the account's own last 6 bills
    elec = (h.amounts or {}).get("electricity")
    if elec is not None and h.kwh and h.kwh > 0:
        rate = elec / h.kwh
        t = T.match(db, h.profile, h.municipality, h.tariff, h.read_end or h.bill_date)
        exp = T.expected(t, h.kwh) if t is not None else None
        if exp and exp["total"]:
            p = _pct(elec, exp["total"])
            if abs(p) > rp:
                out.append({"code": "TARIFF", "text": f"electricity {money(elec)} for {qty(h.kwh)} kWh is {_word(p)} the tariff {t.tariff or t.municipality or t.profile} "
                            f"({money(exp['total'])} at {exp['how']}; allowed {rp:g}%)",
                            "figures": {"actual": round(elec, 2), "expected": exp["total"], "pct": round(p, 1), "limit": rp, "tariff_id": t.id, "how": exp["how"]}})
        elif exp is None:
            hist = [(p.amounts or {}).get("electricity") / p.kwh for p in prev[-RATE_PERIODS:] if (p.amounts or {}).get("electricity") is not None and p.kwh and p.kwh > 0]
            if len(hist) >= minb:
                avg = sum(hist) / len(hist)
                p = _pct(rate, avg) if avg else 0
                if avg and abs(p) > rp:
                    out.append({"code": "RATE", "text": f"rate R {rate:.4f}/kWh is {_word(p)} the account's last {len(hist)} bills (R {avg:.4f}/kWh; allowed {rp:g}%)",
                                "figures": {"actual": round(rate, 4), "expected": round(avg, 4), "pct": round(p, 1), "limit": rp, "periods": [x.period for x in prev[-RATE_PERIODS:]]}})

    # estimated readings running
    if h.reading_type == "E":
        run = 1
        for p in reversed(prev):
            if p.reading_type != "E":
                break
            run += 1
        if run >= max(2, int(_n(db, port, "estimate_run_n"))):
            out.append({"code": "EST_RUN", "text": f"estimated reading {run} bills running: ask for an actual reading", "figures": {"run": run}})

    # a gap or an overlap with the last bill's reading
    before = next((p for p in reversed(prev) if (p.period or "") < (h.period or "") and p.read_end), None)
    if before is not None and h.read_start:
        d = BH.days_between(before.read_end, h.read_start)
        tol = _n(db, port, "period_gap_days")
        if d is not None and d > tol:
            out.append({"code": "GAP", "text": f"reading starts {h.read_start}, {d} days after the last bill's reading ended ({before.read_end}): a gap",
                        "figures": {"last_end": before.read_end, "start": h.read_start, "days": d}})
        elif d is not None and d < 0:
            out.append({"code": "OVERLAP", "text": f"reading starts {h.read_start}, {-d} days before the last bill's reading ended ({before.read_end}): an overlap",
                        "figures": {"last_end": before.read_end, "start": h.read_start, "days": d}})

    # the check meter (Telkom's BCX readings)
    if h.kwh is not None:
        m = meter_for(db, h, _account(db, c, h))
        if m is not None and m.kwh:
            p = _pct(h.kwh, m.kwh)
            lim = _n(db, port, "check_meter_pct")
            if abs(p) > lim:
                out.append({"code": "METER", "text": f"billed {qty(h.kwh)} kWh is {_word(p)} the check meter's {qty(m.kwh)} kWh (allowed {lim:g}%)",
                            "figures": {"actual": h.kwh, "expected": m.kwh, "pct": round(p, 1), "limit": lim, "meter_id": m.id}})
    return out


def tess(db: Session, c: ApCapture, it: Optional[ApItem] = None) -> tuple[list[str], list[str]]:
    """Tess's hook (capture.test_batch, utility bills not uploaded): the bill's history row, the true-up of an estimate it
    replaces, and the checks → (yellow lines, red lines)."""
    h = BH.record(db, c)
    if h is None:
        return [], []
    try:
        from . import estimates as ES

        ES.true_up(db, c, h, it)
    except Exception:  # noqa: BLE001 — the true-up never stops Tess; the Estimates view shows what is left open
        import logging

        logging.getLogger("cbx.ap.variance").exception("estimate true-up %s", c.id)
    found = compute(db, h, c)
    if found != (h.findings or []):
        h.findings = found
    if (h.flags or 0) != len(found):
        h.flags = len(found)
    warn, stop = [], []
    if h.decision == "disputed" and h.dispute_id:
        d = db.get(ApDispute, h.dispute_id)
        if d is not None and d.status == "open":
            ref = f"dispute #{d.id}" + (f", municipal ref {d.municipal_ref}" if d.municipal_ref else "")
            if d.pay_under_dispute:
                warn.append(f"paid under dispute ({ref})")
            else:
                stop.append(f"disputed with the municipality ({ref}): held out of the upload until it is resolved or paid under dispute")
        return warn, stop
    if h.decision == "accepted":
        return [], []
    return [f["text"] for f in found], []


# ---- a person's decision ------------------------------------------------------------------------------------------------------
def _hist_row(db: Session, hid: int) -> ApBillHistory:
    h = db.get(ApBillHistory, hid)
    if h is None:
        raise LookupError("No such bill.")
    return h


def _capture_of(db: Session, h: ApBillHistory) -> tuple[Optional[ApCapture], Optional[ApItem]]:
    c = db.get(ApCapture, h.capture_id) if h.capture_id else None
    return c, (db.get(ApItem, c.item_id) if c is not None else None)


def _line(it: Optional[ApItem], who: str, text: str) -> None:
    if it is not None:
        from . import service as S

        S._add(it, who, text)


def _retest(db: Session, c: Optional[ApCapture]) -> None:
    if c is None or c.uploaded_at or c.stage != "ready":
        return
    from . import capture as C

    db.flush()
    b = db.get(ApBatch, c.batch_id)
    if b is not None and not db.query(ApCapture.id).filter(ApCapture.batch_id == b.id, ApCapture.stage.in_(C.READING)).first():
        C.test_batch(db, b)


def accept(db: Session, user: User, hid: int, reason: str) -> ApBillHistory:
    h = _hist_row(db, hid)
    c, it = _capture_of(db, h)
    reason = str(reason or "").strip()
    if len(reason) < 3:
        raise Refused("Say why the variance is accepted (it is kept as evidence).")
    if not h.findings:
        raise Refused("This bill has no variance to accept.")
    if h.decision == "disputed":
        raise Refused("This bill is disputed: close the dispute first.")
    R.require(db, user, "reviewer", h.portfolio or None, capture=c, what="Accepting a variance", include_dropper=True)  # 1.0.104 (review): nor who dropped the folder
    h.decision, h.decision_reason, h.decided_by, h.decided_by_id, h.decided_at = "accepted", reason[:2000], _who(user), user.id, utcnow()
    _line(it, _who(user), f"Variance accepted: {reason[:300]}")
    log(db, "ap_variance_accepted", actor=user, target=f"ap_item:{it.id}" if it else f"ap_bill:{h.id}",
        detail=f"account {h.account_key} {h.period}: {'; '.join(f['text'] for f in h.findings)[:250]} — {reason[:120]}"[:400])
    _retest(db, c)
    return h


def dispute(db: Session, user: User, hid: int, data: dict) -> ApDispute:
    h = _hist_row(db, hid)
    c, it = _capture_of(db, h)
    reason = str((data or {}).get("reason") or "").strip()
    if len(reason) < 3:
        raise Refused("Say what is disputed.")
    if h.decision == "disputed" and h.dispute_id and (db.get(ApDispute, h.dispute_id) or ApDispute(status="")).status == "open":
        raise Refused("This bill is already disputed.")
    R.require(db, user, "capturer", h.portfolio or None, what="Disputing a bill")
    amt = BH.num((data or {}).get("amount"))
    if amt is None:
        exp = next((f["figures"].get("expected") for f in h.findings or [] if f.get("code") == "VAR_R"), None)
        amt = round((h.total_excl or 0) - exp, 2) if exp is not None and (h.total_excl or 0) > exp else h.total_excl
    d = ApDispute(account_key=h.account_key, municipality=h.municipality, portfolio=h.portfolio, period=h.period, capture_id=h.capture_id, history_id=h.id,
                  amount=amt, reason=reason[:4000], municipal_ref=str((data or {}).get("municipal_ref") or "").strip()[:80],
                  owner=str((data or {}).get("owner") or "").strip()[:160], pay_under_dispute=bool((data or {}).get("pay_under_dispute")),
                  created_by=_who(user), created_by_id=user.id)
    db.add(d)
    db.flush()
    h.decision, h.decision_reason, h.decided_by, h.decided_by_id, h.decided_at, h.dispute_id = "disputed", reason[:2000], _who(user), user.id, utcnow(), d.id
    _line(it, _who(user), f"Disputed with the municipality (dispute #{d.id}{', ref ' + d.municipal_ref if d.municipal_ref else ''}): {reason[:300]}"
          + (" Paid under dispute." if d.pay_under_dispute else " Held out of the SAP upload."))
    log(db, "ap_dispute", actor=user, target=f"ap_item:{it.id}" if it else f"ap_dispute:{d.id}",
        detail=f"opened #{d.id}: account {h.account_key} {h.period}, {money(amt)}{', pay under dispute' if d.pay_under_dispute else ''} — {reason[:160]}"[:400])
    _retest(db, c)
    return d


def dispute_update(db: Session, user: User, did: int, data: dict) -> ApDispute:
    """Change a dispute: municipal_ref, owner, amount, pay_under_dispute, or close it (status resolved / withdrawn with a
    resolution). Withdrawn puts the bill's variance back to decide."""
    d = db.get(ApDispute, did)
    if d is None:
        raise LookupError("No such dispute.")
    R.require(db, user, "capturer", d.portfolio or None, what="Changing a dispute")
    data = data or {}
    said = []
    for k in ("municipal_ref", "owner"):
        if k in data and str(data[k] or "").strip() != getattr(d, k):
            setattr(d, k, str(data[k] or "").strip()[:80 if k == "municipal_ref" else 160])
            said.append(f"{k.replace('_', ' ')} {getattr(d, k) or '(none)'}")
    if "amount" in data and BH.num(data["amount"]) is not None and BH.num(data["amount"]) != d.amount:
        d.amount = BH.num(data["amount"])
        said.append(f"amount {money(d.amount)}")
    if "pay_under_dispute" in data and bool(data["pay_under_dispute"]) != d.pay_under_dispute:
        if d.status != "open":
            raise Refused("This dispute is closed.")
        d.pay_under_dispute = bool(data["pay_under_dispute"])
        said.append("pay under dispute" if d.pay_under_dispute else "held out of the upload again")
    st = str(data.get("status") or "").strip()
    if st and st != d.status:
        if st not in ("resolved", "withdrawn"):
            raise Refused("A dispute is resolved or withdrawn.")
        if d.status != "open":
            raise Refused("This dispute is closed.")
        res = str(data.get("resolution") or "").strip()
        if len(res) < 3:
            raise Refused("Say how it ended.")
        d.status, d.resolution, d.closed_by, d.closed_at = st, res[:4000], _who(user), utcnow()
        said.append(f"{st}: {res[:160]}")
        h = db.get(ApBillHistory, d.history_id) if d.history_id else None
        if h is not None and st == "withdrawn" and h.dispute_id == d.id:
            h.decision, h.decision_reason, h.decided_by, h.decided_by_id, h.decided_at = "", "", "", None, None
    if not said:
        return d
    h = db.get(ApBillHistory, d.history_id) if d.history_id else None
    c, it = _capture_of(db, h) if h is not None else (None, None)
    _line(it, _who(user), f"Dispute #{d.id}: " + "; ".join(said)[:300])
    log(db, "ap_dispute", actor=user, target=f"ap_item:{it.id}" if it else f"ap_dispute:{d.id}", detail=f"#{d.id}: " + "; ".join(said)[:380])
    _retest(db, c)
    return d


# ---- the Variances & disputes view -----------------------------------------------------------------------------------------
def _ctx(db: Session, hs: list[ApBillHistory]) -> tuple[dict, dict, dict]:
    caps = {c.id: c for c in db.query(ApCapture).filter(ApCapture.id.in_([h.capture_id for h in hs if h.capture_id] or [0])).all()}
    items = {i.id: i for i in db.query(ApItem).filter(ApItem.id.in_([c.item_id for c in caps.values()] or [0])).all()}
    batches = {b.id: b for b in db.query(ApBatch).filter(ApBatch.id.in_({c.batch_id for c in caps.values()} or {0})).all()}
    return caps, items, batches


def bill_out(h: ApBillHistory, c: Optional[ApCapture], it: Optional[ApItem], b: Optional[ApBatch]) -> dict:
    f = (c.fields or {}) if c is not None else {}
    a = ((c.match or {}).get("account") or {}) if c is not None else {}
    return {"id": h.id, "account": f.get("account") or h.account_key, "account_key": h.account_key, "municipality": h.municipality, "site": a.get("name1") or f.get("site") or "",
            "portfolio": h.portfolio, "period": h.period, "read_start": h.read_start, "read_end": h.read_end, "kwh": h.kwh, "total_excl": h.total_excl, "amounts": dict(h.amounts or {}),
            "reading_type": h.reading_type, "tariff": h.tariff, "findings": h.findings or [], "decision": h.decision, "reason": h.decision_reason,
            "decided_by": h.decided_by, "decided_at": h.decided_at.isoformat() if h.decided_at else "", "dispute_id": h.dispute_id,
            "capture_id": h.capture_id, "item_id": c.item_id if c is not None else None, "file": it.name if it else "", "batch": b.name if b else "",
            "batch_id": b.id if b else None, "uploaded": bool(c is not None and c.uploaded_at), "status": c.sap_status if c is not None else ""}


def listing(db: Session, state: str = "open", portfolio: str = "", account: str = "", limit: int = 300) -> list[dict]:
    """Bills with a variance: open (nobody decided, not uploaded), accepted, disputed, or all."""
    q = db.query(ApBillHistory).filter(ApBillHistory.flags > 0)
    if portfolio:
        q = q.filter(ApBillHistory.portfolio == portfolio)
    if account:
        q = q.filter(ApBillHistory.account_key == BH.acct_key(account))
    if state in ("accepted", "disputed"):
        q = q.filter(ApBillHistory.decision == state)
    elif state == "open":
        q = q.filter(ApBillHistory.decision == "")
    hs = q.order_by(ApBillHistory.id.desc()).limit(4000).all()
    caps, items, batches = _ctx(db, hs)
    out = []
    for h in hs:
        c = caps.get(h.capture_id)
        if state == "open" and c is not None and c.uploaded_at:
            continue
        it = items.get(c.item_id) if c is not None else None
        out.append(bill_out(h, c, it, batches.get(c.batch_id) if c is not None else None))
        if len(out) >= limit:
            break
    return out


def open_count(db: Session) -> int:
    ids = [cid for (cid,) in db.query(ApBillHistory.capture_id).filter(ApBillHistory.decision == "", ApBillHistory.flags > 0, ApBillHistory.capture_id.isnot(None)).all()]
    if not ids:
        return 0
    up = {i for (i,) in db.query(ApCapture.id).filter(ApCapture.id.in_(ids), ApCapture.uploaded_at.isnot(None)).all()}
    return sum(1 for i in ids if i not in up)


def bill_detail(db: Session, hid: int) -> dict:
    """One bill with the account's earlier bills (the comparison a person looks at before deciding)."""
    h = _hist_row(db, hid)
    c, it = _capture_of(db, h)
    b = db.get(ApBatch, c.batch_id) if c is not None else None
    port = h.portfolio or ""
    earlier = BH.rows_of(db, h.account_key, before=h)[-max(int(_n(db, port, "compare_periods")), RATE_PERIODS):]
    d = db.get(ApDispute, h.dispute_id) if h.dispute_id else None
    return {**bill_out(h, c, it, b), "earlier": [BH._out(x) for x in earlier], "dispute": dispute_out(d) if d else None,
            "settings": {k: _set(db, port, k) for k in ("variance_pct", "rate_pct", "compare_periods", *NEW_KEYS)}}


def dispute_out(d: ApDispute) -> dict:
    return {"id": d.id, "account_key": d.account_key, "municipality": d.municipality, "portfolio": d.portfolio, "period": d.period, "capture_id": d.capture_id,
            "history_id": d.history_id, "amount": d.amount, "reason": d.reason, "municipal_ref": d.municipal_ref, "owner": d.owner, "status": d.status,
            "pay_under_dispute": d.pay_under_dispute, "resolution": d.resolution, "created_by": d.created_by,
            "created_at": d.created_at.isoformat() if d.created_at else "", "closed_by": d.closed_by, "closed_at": d.closed_at.isoformat() if d.closed_at else "",
            "age_days": ((utcnow() - d.created_at).days if d.created_at and d.status == "open" else None)}


def disputes(db: Session, status: str = "open", portfolio: str = "") -> list[dict]:
    q = db.query(ApDispute)
    if status in ("open", "resolved", "withdrawn"):
        q = q.filter(ApDispute.status == status)
    if portfolio:
        q = q.filter(ApDispute.portfolio == portfolio)
    return [dispute_out(d) for d in q.order_by(ApDispute.id.desc()).limit(500).all()]


# ---- the check meters (Telkom's BCX readings) ----------------------------------------------------------------------------------
def upload_check_meters(db: Session, user: User, data: bytes, name: str) -> dict:
    """A CSV (or .xlsx) with site, period, kWh (and optionally account): each reading is added or replaces the one for
    that site and period. Audited ap_check_meters."""
    R.require(db, user, "capturer", None, what="Loading check-meter readings")
    grid = T._table(data, name or "x.csv")
    def is_head(r) -> bool:
        ks = {T.key(c) for c in r or [] if c is not None}
        return bool(ks & {"site", "account", "accountnumber"}) and bool(ks & {"period", "month", "date"})

    head = next((i for i, r in enumerate(grid[:20]) if is_head(r)), None)
    if head is None:
        raise Refused("The file needs the columns site, period and kWh (account is optional).")
    names = {T.key(c): i for i, c in enumerate(grid[head]) if c is not None}
    col = {"site": names.get("site"), "period": names.get("period", names.get("month", names.get("date"))),
           "kwh": next((names[k] for k in ("kwh", "kwhs", "consumption", "units") if k in names), None), "account": names.get("account", names.get("accountnumber"))}
    if col["kwh"] is None:
        raise Refused("The file needs a kWh column.")
    done, bad = 0, []
    for n, r in enumerate(grid[head + 1:], head + 2):
        if not r or all(c in (None, "") for c in r):
            continue
        cell = lambda k: r[col[k]] if col[k] is not None and col[k] < len(r) else None  # noqa: E731
        site = str(cell("site") or "").strip()[:160]
        p = cell("period")
        per = p.strftime("%Y-%m") if isinstance(p, datetime) else (str(p).strip()[:7] if re.fullmatch(r"\d{4}-\d{2}.*", str(p or "").strip()) else BH.period_of("", str(p or ""), str(p or "")))
        kwh = BH.num(cell("kwh"))
        acct = BH.acct_key(cell("account")) if cell("account") not in (None, "") else ""
        if not (site or acct) or not per or kwh is None:
            bad.append(n)
            continue
        site = site or acct
        m = db.query(ApCheckMeter).filter(ApCheckMeter.site == site, ApCheckMeter.period == per).one_or_none()
        if m is None:
            m = ApCheckMeter(site=site, period=per)
            db.add(m)
        m.kwh, m.account_key, m.uploaded_by, m.uploaded_at = kwh, acct or m.account_key or "", _who(user), utcnow()
        done += 1
        if done > 50000:
            raise Refused("More than 50 000 readings in one file.")
    db.flush()
    log(db, "ap_check_meters", actor=user, target="ap_check_meters", detail=f"{name or 'a file'}: {done} readings" + (f", {len(bad)} rows skipped" if bad else ""), count=done)
    return {"readings": done, "skipped_rows": bad[:30]}


def _site_keys(a: Optional[ApUtilityAccount]) -> set:
    return {T.key(x) for x in (a.site, a.cost_centre, (a.cost_centre or "").lstrip("0")) if T.key(x)} if a is not None else set()


def no_meter_ranking(db: Session, portfolio: str = "") -> list[dict]:
    """Sites without a check meter (no "check meter" mark on the account, Builder C's V10 field, and no readings loaded
    for its account, site or cost centre), ranked by consumption (average kWh of their last 6 bills) and then by their
    largest variance (Telkom C: where a check meter would help most)."""
    q = db.query(ApBillHistory)
    if portfolio:
        q = q.filter(ApBillHistory.portfolio == portfolio)
    by: dict[str, list[ApBillHistory]] = {}
    for h in q.all():
        by.setdefault(h.account_key, []).append(h)
    meter_sites = {T.key(s) for (s,) in db.query(ApCheckMeter.site).distinct().all()}
    meter_accts = {k for (k,) in db.query(ApCheckMeter.account_key).filter(ApCheckMeter.account_key != "").distinct().all()}
    out = []
    for k, hs in by.items():
        a = db.query(ApUtilityAccount).filter(ApUtilityAccount.account_key == k).order_by(ApUtilityAccount.id.desc()).first()
        if (a is not None and getattr(a, "check_meter", None)) or k in meter_accts or _site_keys(a) & meter_sites:
            continue
        hs.sort(key=lambda h: (h.period or "", h.id))
        recent = hs[-6:]
        kwhs = [h.kwh for h in recent if h.kwh is not None]
        worst = max((abs(f["figures"].get("pct") or 0) for h in recent for f in h.findings or [] if f.get("code") in ("VAR_KWH", "VAR_R", "RATE", "TARIFF")), default=0.0)
        out.append({"account_key": k, "account": a.account if a else k, "site": a.site if a else "", "cost_centre": a.cost_centre if a else "",
                    "municipality": (a.municipality if a else "") or recent[-1].municipality, "portfolio": recent[-1].portfolio, "bills": len(hs),
                    "avg_kwh": round(sum(kwhs) / len(kwhs), 1) if kwhs else None, "last_kwh": recent[-1].kwh, "worst_variance_pct": round(worst, 1),
                    "variances": sum(1 for h in recent if h.findings)})
    out.sort(key=lambda r: (-(r["avg_kwh"] or 0), -r["worst_variance_pct"], r["account_key"]))
    for n, r in enumerate(out, 1):
        r["rank"] = n
    return out


def ranking_xlsx(db: Session, portfolio: str = "") -> bytes:
    rows = no_meter_ranking(db, portfolio)
    head = [("rank", "Rank"), ("account", "Account"), ("site", "Site"), ("cost_centre", "PID/CC"), ("municipality", "Municipality"), ("portfolio", "Portfolio"),
            ("avg_kwh", "Average kWh (last 6 bills)"), ("last_kwh", "Last kWh"), ("worst_variance_pct", "Largest variance %"), ("variances", "Bills with a variance"), ("bills", "Bills read")]
    return _xlsx("No check meter", head, rows)


def _xlsx(title: str, head: list[tuple[str, str]], rows: list[dict]) -> bytes:
    import openpyxl

    from .. import xlsafe

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([lbl for _k, lbl in head])
    for r in rows:
        ws.append([("; ".join(map(str, r.get(k))) if isinstance(r.get(k), list) else r.get(k)) for k, _l in head])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(50, max(8, *(len(str(c.value or "")) for c in col[:300])) + 2)
    ws.freeze_panes = "A2"
    xlsafe.book(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def disputes_xlsx(db: Session, status: str = "", portfolio: str = "") -> bytes:
    head = [("id", "Dispute"), ("status", "Status"), ("portfolio", "Portfolio"), ("municipality", "Municipality"), ("account_key", "Account"), ("period", "Period"),
            ("amount", "Amount (R)"), ("reason", "Reason"), ("municipal_ref", "Municipal reference"), ("owner", "Council Liaison owner"),
            ("pay_under_dispute", "Paid under dispute"), ("created_by", "Opened by"), ("created_at", "Opened"), ("age_days", "Days open"),
            ("resolution", "Resolution"), ("closed_by", "Closed by"), ("closed_at", "Closed")]
    return _xlsx("Disputes", head, disputes(db, status or "all", portfolio))


# ---- the review workbook: the comparison (V07 item 5) ---------------------------------------------------------------------------
COMPARE_HEAD = ["Source file", "Account", "Municipality", "Period", "Read start", "Read end", "Reading", "kWh", "R excl. VAT", "Tariff",
                "Compared with (periods)", "Average R", "Average kWh", "Rate R/kWh", "Expected (rate or tariff)", "Findings", "Decision", "Reason",
                "Decided by", "Decided at", "Dispute", "Check meter kWh"]


def comparison_rows(db: Session, caps: list[ApCapture], items: dict) -> list[list]:
    hs = {h.capture_id: h for h in db.query(ApBillHistory).filter(ApBillHistory.capture_id.in_([c.id for c in caps] or [0])).all()}
    out = []
    for c in caps:
        h = hs.get(c.id)
        if h is None:
            continue
        port = h.portfolio or ""
        earlier = BH.rows_of(db, h.account_key, before=h)[-max(1, int(_n(db, port, "compare_periods"))):]
        amts = [p.total_excl for p in earlier if p.total_excl is not None]
        kw = [p.kwh for p in earlier if p.kwh is not None]
        elec = (h.amounts or {}).get("electricity")
        rate = round(elec / h.kwh, 4) if elec is not None and h.kwh else None
        exp = next((f["figures"].get("expected") for f in h.findings or [] if f.get("code") in ("TARIFF", "RATE")), None)
        d = db.get(ApDispute, h.dispute_id) if h.dispute_id else None
        m = meter_for(db, h, _account(db, c, h))
        it = items.get(c.item_id)
        out.append([it.name if it else "", (c.fields or {}).get("account") or h.account_key, h.municipality, h.period, h.read_start, h.read_end, h.reading_type,
                    h.kwh, h.total_excl, h.tariff, ", ".join(p.period for p in earlier), round(sum(amts) / len(amts), 2) if amts else None,
                    round(sum(kw) / len(kw), 1) if kw else None, rate, exp, "; ".join(f["text"] for f in h.findings or []) or "in line",
                    h.decision or ("open" if h.findings else ""), h.decision_reason, h.decided_by, f"{h.decided_at:%Y-%m-%d %H:%M}" if h.decided_at else "",
                    (f"#{d.id} {d.status}" + (f" ref {d.municipal_ref}" if d.municipal_ref else "") + (" (paid under dispute)" if d.pay_under_dispute else "")) if d else "",
                    m.kwh if m is not None else None])
    return out


def with_comparison(db: Session, caps: list[ApCapture], items: dict, data: bytes) -> bytes:
    """The review workbook with a "Bill comparison" sheet (only when the batch has utility bills in the history)."""
    rows = comparison_rows(db, [c for c in caps if c.bucket == "utility"], items)
    if not rows:
        return data
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    from .. import xlsafe

    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.create_sheet("Bill comparison")
    ws.append(COMPARE_HEAD)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="003F2D")
    for r in rows:
        ws.append(r)
        if r[16] == "open":
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor="FFF2CC")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(60, max(8, *(len(str(c.value or "")) for c in col[:300])) + 2)
    xlsafe.book(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def csv_bytes(rows: list[dict], cols: list[str]) -> bytes:
    from .. import xlsafe

    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(cols)
    for r in rows:
        w.writerow([xlsafe.csv_text(r.get(k)) if isinstance(r.get(k), str) else r.get(k) for k in cols])
    return buf.getvalue().encode("utf-8")
