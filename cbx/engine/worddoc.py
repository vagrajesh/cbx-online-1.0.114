"""Make a Word document (1.0.56): a rule fills a .docx template from the tool's assets/ with the row's
values and puts the document into a file column of the row.

In the template, {{column}} is replaced by the row's value (a number shown as Rand when its column
says so, a date as 12 March 2026, a person or a linked row by its name), and {{today}} by today's date.
Placeholders are found in paragraphs, tables, headers and footers. A placeholder the row does not have
stays as it is, so a mistake shows in the document instead of vanishing.
"""
from __future__ import annotations

import io
import re
from pathlib import Path

PLACEHOLDER = re.compile(r"\{\{\s*([A-Za-z_][A-Za-z0-9_]*)\s*\}\}")


class DocError(ValueError):
    pass


def money(v) -> str:
    try:
        v = float(v)
    except (TypeError, ValueError):
        return ""
    whole, dec = f"{v:,.2f}".split(".")
    return "R " + whole.replace(",", " ") + "." + dec


def shown(col: dict, v, row: dict) -> str:
    """A value as a document shows it."""
    from datetime import date

    from ..util import human_date

    t = col.get("type", "text")
    if v in (None, ""):
        return ""
    if t in ("person", "link"):
        return str(((row.get("_labels") or {}).get(col["key"])) or v)
    if t in ("number", "int", "calc") and col.get("format") == "money":
        return money(v)
    if t == "number":
        return f"{float(v):,.2f}".replace(",", " ")
    if t == "date":
        try:
            return human_date(date.fromisoformat(str(v)[:10]))
        except ValueError:
            return str(v)
    if t == "bool":
        return "Yes" if v else "No"
    if t == "checklist":
        return "; ".join(str(x.get("text", "")) for x in (v or []) if isinstance(x, dict))
    if t == "file":
        return str((v or {}).get("name") or "")
    return str(v)


def _paragraphs(d):
    yield from d.paragraphs
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for s in d.sections:
        yield from s.header.paragraphs
        yield from s.footer.paragraphs


def fill(template: Path, values: dict) -> bytes:
    import docx

    d = docx.Document(str(template))
    for p in _paragraphs(d):
        if "{{" not in p.text:
            continue
        new = PLACEHOLDER.sub(lambda m: str(values[m.group(1)]) if m.group(1) in values else m.group(0), p.text)
        if new == p.text:
            continue
        if p.runs:
            p.runs[0].text = new
            for r in p.runs[1:]:
                r.text = ""
        else:
            p.add_run(new)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def check_step(t, at: str, s: dict) -> list[str]:
    errs = []
    tpl = str(s.get("template") or "")
    if not tpl.lower().endswith(".docx") or "/" in tpl or "\\" in tpl:
        errs.append(f"{at}: choose the Word template (a .docx uploaded to the tool)")
    elif not (t.path / "assets" / tpl).is_file():
        errs.append(f"{at}: assets/{tpl} is not in the folder — upload the template")
    b = t.data.get(s.get("data"))
    if b is None:
        errs.append(f"{at}: choose the table whose row fills the document")
    else:
        col = next((c for c in b.get("columns") or [] if c.get("key") == s.get("into")), None)
        if col is None or col.get("type") != "file":
            errs.append(f"{at}: say which file column of {s.get('data')} gets the document")
    if not s.get("row"):
        errs.append(f"{at}: say which row (its id, e.g. {{{{trigger.row.id}}}})")
    return errs


def run_step(db, t, s: dict, ctx: dict, user, *, dry: bool = False) -> str:
    from ..api.engine_api import label_people
    from ..db import utcnow
    from . import files as FILES
    from .flows import _num, _text, local_today, render
    from .store import Store

    b = t.data.get(s.get("data"))
    if b is None:
        raise DocError(f"The tool has no table '{s.get('data')}'.")
    cols = b.get("columns") or []
    tpl = t.path / "assets" / str(s.get("template") or "")
    if not tpl.is_file():
        raise DocError(f"The Word template {s.get('template')} is not in the tool.")
    rid = int(_num(render(s.get("row"), ctx)) or 0)
    st = Store(t.slug)
    row = st.get(s["data"], cols, rid)
    if row is None:
        raise DocError(f"Row {rid} is not in {s['data']}.")
    st.label_links([row], cols, t.data)
    label_people(db, [row], cols)
    values = {c["key"]: shown(c, row.get(c["key"]), row) for c in cols}
    from ..util import human_date

    values.update({"today": human_date(local_today()), "id": str(rid)})
    name = _text(render(s.get("name") or "{{trigger.row.id}}.docx", ctx)).strip() or f"document-{rid}.docx"
    name = FILES.clean_name(re.sub(r"[^\w\- .()]+", "", name))
    if not name.lower().endswith(".docx"):
        name += ".docx"
    try:
        data = fill(tpl, values)
    except Exception as e:  # noqa: BLE001 — a template Word cannot open
        raise DocError(f"The template could not be filled: {str(e)[:200]}")
    if dry:
        return name
    col = next(c for c in cols if c.get("key") == s.get("into"))
    try:
        meta = FILES.save(t.slug, s["data"], rid, {**col, "accept": ["docx"], "max_mb": 100}, name, data, user.id if user else None, utcnow().isoformat())
    except FILES.FileError as e:
        raise DocError(str(e))
    st.set_file(s["data"], cols, rid, col["key"], meta)
    return name

