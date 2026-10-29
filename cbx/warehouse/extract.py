"""Text out of uploaded files: pdfplumber first, PyMuPDF as fallback, tesseract for scans.
Every reader is optional — a missing library degrades to "no text", never to an error."""
from __future__ import annotations

import io
from typing import Tuple


def pdf_text(data: bytes, *, max_pages: int = 40) -> Tuple[str, int]:
    """Return (text, pages). Text is empty for image-only PDFs."""
    text, pages = "", 0
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = len(pdf.pages)
            parts = []
            for p in pdf.pages[:max_pages]:
                parts.append(p.extract_text() or "")
            text = "\n".join(parts)
    except Exception:  # noqa: BLE001
        text = ""
    if not text.strip():
        try:
            import fitz  # PyMuPDF

            doc = fitz.open(stream=data, filetype="pdf")
            pages = doc.page_count
            text = "\n".join(page.get_text() for page in list(doc)[:max_pages])
            doc.close()
        except Exception:  # noqa: BLE001
            pass
    return text.strip(), pages


def ocr_pdf(data: bytes, *, max_pages: int = 10, dpi: int = 200) -> str:
    """Render each page with PyMuPDF and read it with tesseract. Returns '' if either is missing."""
    try:
        import fitz
        import pytesseract
        from ..config import settings as _st
        if _st.tesseract_cmd:
            pytesseract.pytesseract.tesseract_cmd = _st.tesseract_cmd
        from PIL import Image
    except Exception:  # noqa: BLE001
        return ""
    out = []
    try:
        doc = fitz.open(stream=data, filetype="pdf")
        for page in list(doc)[:max_pages]:
            pix = page.get_pixmap(dpi=dpi)
            img = Image.open(io.BytesIO(pix.tobytes("png")))
            out.append(pytesseract.image_to_string(img))
        doc.close()
    except Exception:  # noqa: BLE001
        return ""
    return "\n".join(out).strip()


def text_of(filename: str, data: bytes) -> Tuple[str, int]:
    """Best-effort text for any upload kind (used for duplicate fingerprints and detection)."""
    name = filename.lower()
    if name.endswith(".pdf"):
        return pdf_text(data)
    if name.endswith((".csv", ".txt", ".json")):
        try:
            return data.decode("utf-8", errors="ignore")[:200_000], 0
        except Exception:  # noqa: BLE001
            return "", 0
    if name.endswith(".docx"):
        try:
            import docx

            d = docx.Document(io.BytesIO(data))
            return "\n".join(p.text for p in d.paragraphs), 0
        except Exception:  # noqa: BLE001
            return "", 0
    if name.endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets[:5]:
                for row in ws.iter_rows(max_row=200, values_only=True):
                    parts.append(" ".join("" if v is None else str(v) for v in row))
            return "\n".join(parts), 0
        except Exception:  # noqa: BLE001
            return "", 0
    return "", 0
