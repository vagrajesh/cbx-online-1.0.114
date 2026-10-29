"""Read the columns a user wants out of an Excel template.

The header is NOT simply the first non-empty row: real templates start with a title ("Electricity"),
a merged banner, a blank line, then the real header. We scan the first 20 rows of every sheet and
take the row with the most non-empty cells (earliest on ties), across the sheet that has the widest
such row. Group-header rows directly above (e.g. "Consumption" spanning kWh/kVA) are prefixed onto
the header text when the header cell alone would be ambiguous.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

SCAN_ROWS = 20


@dataclass
class ExcelHeader:
    sheet: str
    row_index: int  # 1-based
    headers: list[str]
    sample: list  # first data row values (aligned to headers), may be empty
    skipped: list[str] = field(default_factory=list)  # rows we passed over (titles etc.)


def _clean(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return re.sub(r"\s+", " ", s)


def read_headers(data: bytes) -> ExcelHeader:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    best: ExcelHeader | None = None
    for ws in wb.worksheets:
        rows = []
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=SCAN_ROWS, values_only=True), start=1):
            rows.append((i, [_clean(v) for v in row]))
        if not rows:
            continue
        counts = [(sum(1 for v in vals if v), i, vals) for i, vals in rows]
        n, idx, vals = max(counts, key=lambda t: (t[0], -t[1]))
        if n == 0:
            continue
        # group header directly above: a row with fewer cells whose text sits over blank header cells
        above = next((v for i, v in rows if i == idx - 1), None)
        headers = list(vals)
        if above and 0 < sum(1 for v in above if v) < n:
            current = ""
            for j in range(len(headers)):
                if j < len(above) and above[j]:
                    current = above[j]
                if headers[j] and current and current.lower() not in headers[j].lower() and len(headers[j]) <= 4:
                    headers[j] = f"{current} {headers[j]}"
        # first data row after the header
        sample: list = []
        for r in ws.iter_rows(min_row=idx + 1, max_row=idx + SCAN_ROWS, values_only=True):
            if any(_clean(v) for v in r):
                sample = list(r)
                break
        skipped = [" ".join(v for v in vals2 if v) for i2, vals2 in rows if i2 < idx and any(vals2)]
        cand = ExcelHeader(sheet=ws.title, row_index=idx, headers=headers, sample=sample, skipped=skipped)
        if best is None or n > sum(1 for h in best.headers if h):
            best = cand
    if best is None or not any(best.headers):
        raise ValueError("I couldn't find a header row in that workbook.")
    return best


FILENAME_WORDS = re.compile(r"\b(file ?name|filename|file|pdf|source ?file|source|document ?name|doc ?name)\b")
OPTIONAL_WORDS = re.compile(r"\b(travel|travelling|km|kms|kilomet\w*|exceptions?|notes?|comments?|remarks?|guarantees?|deposits?|security)\b")


def guess_ctype(label: str, sample) -> str:
    l = label.lower()
    # "file", "File name", "Electricity file", "PDF", "Source": the PDF's own name, not text on the page
    if FILENAME_WORDS.search(l) and not re.search(r"\b(date|amount|total|no|number|count|size|pages?)\b", l):
        return "filename"
    # "Due date" / "Date paid" are dates, not money — check the date words first
    if re.search(r"\b(date|period|month|expiry)\b", l):
        return "date"
    if re.search(r"\b(vat|total|amount|excl|incl|price|cost|charge|due|paid|r\b|rand|zar)\b", l):
        return "money"
    if re.search(r"\b(from|to)\b", l):
        return "date"
    if re.search(r"\b(kwh|kva|kw|units|qty|quantity|reading|count|number of|nmd|demand|consumption|litres|m3|m³)\b", l):
        return "number"
    if isinstance(sample, (int, float)) and not isinstance(sample, bool):
        return "number"
    return "text"


def guess_required(label: str, ctype: str = "") -> bool:
    """Headers for figures most documents don't have (travel, kms, exceptions, notes, guarantees,
    deposits) start out optional: a blank cell is the answer, not a miss. File-name and formula
    columns are never 'required' in the extraction sense."""
    if ctype in ("filename", "formula"):
        return False
    return not OPTIONAL_WORDS.search(label.lower())


def optional_reason(label: str) -> str:
    m = OPTIONAL_WORDS.search(label.lower())
    return f"from the Excel header: '{m.group(1)}' columns are usually not on every document" if m else ""
