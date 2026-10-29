"""Results out: an Excel workbook (rows · provenance · misses) and appending to a master sheet."""
from __future__ import annotations

import io
import re
import uuid
from typing import Optional

from sqlalchemy.orm import Session

from ..audit import log as audit
from ..db import utcnow
from ..models import Batch, Document, MasterSheet, StoredFile, Team, Template, User
from ..warehouse import storage
from . import scoring as SC

ACCEPT = ("ok", "filled", "trusted")


def columns_of(template: Template) -> list[dict]:
    """The column dicts the engine works with (one place, so every path sees hint/expr/required)."""
    return [column_dict(c) for c in template.columns]


def column_dict(c) -> dict:
    cfg = dict(c.config or {})
    return {"key": c.key, "label": c.label, "ctype": c.ctype, "required": bool(c.required), "critical": bool(c.critical), "hint": c.hint or "", "expr": (cfg.get("expr") or "") if c.ctype == "formula" else "", "config": cfg}


def column_kinds(columns: list[dict]) -> dict:
    """For the batch summary: '22 columns · 3 optional · 2 calculated · 1 file name'."""
    return {
        "total": len(columns),
        "optional": sum(1 for c in columns if not c.get("required", True) and c.get("ctype") not in ("filename", "formula")),
        "calculated": sum(1 for c in columns if c.get("ctype") == "formula"),
        "filename": sum(1 for c in columns if c.get("ctype") == "filename"),
        "read": sum(1 for c in columns if c.get("ctype") not in ("filename", "formula")),
    }


def rows_of(batch: Batch, columns: list[dict]) -> list[dict]:
    out = []
    for d in batch.documents:
        cells = (d.extraction.cells if d.extraction else {}) or {}
        _, k, n = SC.row_score(cells, columns)
        out.append({"document_id": d.id, "ordinal": d.ordinal, "filename": d.filename, "pages": d.pages, "scanned": d.scanned, "score": d.score, "correct": k, "of": n, "pct": round(k / n, 3) if n else None, "status": d.status, "cells": {c["key"]: cells.get(c["key"]) or {"value": None, "status": "miss"} for c in columns}, "validators": (d.extraction.validators if d.extraction else {}) or {}})
    return out


def workbook_bytes(batch: Optional[Batch], template: Template, rows: Optional[list[dict]] = None) -> bytes:
    """`rows` defaults to the batch's rows; Harvest passes rows built from HarvestRow (each with a
    'date', which becomes a Date column after File)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    columns = columns_of(template)
    rows = rows if rows is not None else rows_of(batch, columns)
    with_date = any("date" in r for r in rows)
    lead = ["#", "File"] + (["Date"] if with_date else [])
    wb = Workbook()
    ws = wb.active
    ws.title = "Rows"
    head = lead + [c["label"] for c in columns] + ["Correct", "Status"]
    ws.append(head)
    bold = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold
    miss_fill = PatternFill("solid", fgColor="FFF1DB")
    filled_fill = PatternFill("solid", fgColor="E3F4E8")
    for r in rows:
        vals = [r["ordinal"], r["filename"]] + ([r.get("date") or ""] if with_date else [])
        for c in columns:
            cell = r["cells"][c["key"]]
            vals.append(cell.get("value") if cell.get("status") in ACCEPT else None)
        vals += [f"{r['correct']} of {r['of']}", r["status"]]
        ws.append(vals)
        rr = ws.max_row
        for i, c in enumerate(columns, start=len(lead) + 1):
            cell = r["cells"][c["key"]]
            xc = ws.cell(row=rr, column=i)
            if cell.get("status") in ("miss", "flagged"):
                xc.fill = miss_fill
            elif cell.get("status") in ("filled", "trusted"):
                xc.fill = filled_fill
            if c["ctype"] == "money" and isinstance(xc.value, (int, float)):
                xc.number_format = '"R" #,##0.00'
            elif c["ctype"] == "formula" and isinstance(xc.value, (int, float)):
                xc.number_format = "#,##0.00"
            elif c["ctype"] == "number" and isinstance(xc.value, (int, float)):
                xc.number_format = "#,##0"
    for i in range(1, len(head) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16 if i > len(lead) else (6 if i == 1 else 30 if i == 2 else 12)

    wp = wb.create_sheet("Provenance")
    wp.append(["#", "File", "Column", "Value", "Page", "Quote on page", "Rule", "Quality", "Score", "Status"])
    for cell in wp[1]:
        cell.font = bold
    for r in rows:
        for c in columns:
            cell = r["cells"][c["key"]]
            wp.append([r["ordinal"], r["filename"], c["label"], cell.get("value"), cell.get("page"), cell.get("quote"), cell.get("method"), cell.get("quality"), cell.get("score"), cell.get("status")])
    for i, w in enumerate([6, 30, 18, 16, 6, 34, 12, 10, 8, 10], start=1):
        wp.column_dimensions[get_column_letter(i)].width = w

    wm = wb.create_sheet("Misses")
    wm.append(["#", "File", "Column", "Page", "What C-EXA saw", "Why it needs you"])
    for cell in wm[1]:
        cell.font = bold
    for r in rows:
        for c in columns:
            cell = r["cells"][c["key"]]
            if cell.get("status") in ("miss", "flagged"):
                wm.append([r["ordinal"], r["filename"], c["label"], cell.get("page"), cell.get("raw") or cell.get("value"), cell.get("note")])
    for i, w in enumerate([6, 30, 18, 6, 24, 50], start=1):
        wm.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_and_store(db: Session, batch: Batch, template: Template, user: User) -> tuple[bytes, Optional[StoredFile]]:
    data = workbook_bytes(batch, template)
    name = re.sub(r"[^A-Za-z0-9 ._()-]+", "", f"{(batch.name or template.name or 'PDF Studio').strip()} - results.xlsx").strip() or "results.xlsx"
    sha = storage.sha256_bytes(data)
    existing = db.query(StoredFile).filter(StoredFile.sha256 == sha, StoredFile.source == "result", StoredFile.uploaded_by_id == user.id).first()
    f = existing
    if f is None:
        f = storage.store(db, user, name, data, sub="Results", source="result", mirror=False, meta={"batch_id": batch.id, "template_id": template.id, "kind": "pdfstudio-export"})
        audit(db, "pdf_export", actor=user, target=f"batch:{batch.id}", detail=f"{f.name} ({len(data)} bytes)")
    return data, f


# ---- master sheet ---------------------------------------------------------------------------------
def _slug(s: str) -> str:
    import re

    return re.sub(r"[^a-z0-9]+", "_", str(s).lower()).strip("_")


def map_columns(sheet: MasterSheet, columns: list[dict]) -> tuple[dict[str, str], list[dict]]:
    """template key → sheet key by key or label match; unmatched template columns are added."""
    sheet_cols = list(sheet.columns or [])
    by_key = {c["key"]: c for c in sheet_cols}
    by_label = {_slug(c.get("label", "")): c for c in sheet_cols}
    mapping = {}
    added = []
    for c in columns:
        hit = by_key.get(c["key"]) or by_label.get(_slug(c["label"])) or by_key.get(_slug(c["label"]))
        if hit is None:
            hit = {"key": c["key"], "label": c["label"], "ctype": c["ctype"]}
            sheet_cols.append(hit)
            by_key[hit["key"]] = hit
            added.append(hit)
        mapping[c["key"]] = hit["key"]
    return mapping, sheet_cols if added else list(sheet.columns or [])


def append_to_master(db: Session, batch: Batch, template: Template, sheet: MasterSheet, user: User) -> dict:
    columns = columns_of(template)
    mapping, new_cols = map_columns(sheet, columns)
    if len(new_cols) != len(sheet.columns or []):
        sheet.columns = new_cols
    team = db.get(Team, user.team_id) if user.team_id else None
    now = utcnow().isoformat()
    rows = list(sheet.rows or [])
    exceptions = list(sheet.exceptions or [])
    existing_src = {(r.get("_src") or {}).get("doc") for r in rows} | {(r.get("_src") or {}).get("doc") for r in exceptions}
    added = skipped = excepted = 0
    for d in batch.documents:
        if d.id in existing_src:
            skipped += 1
            continue
        cells = (d.extraction.cells if d.extraction else {}) or {}
        validators = (d.extraction.validators if d.extraction else {}) or {}
        row = {
            "_id": uuid.uuid4().hex[:12],
            "_src": {"doc": d.id, "file": d.file_id, "batch": batch.id, "name": d.filename},
            "_by": user.id,
            "_by_name": user.name,
            "_at": now,
            "_client": _client_of(cells, columns),
            "_team": team.name if team else "",
        }
        incomplete = []
        for c in columns:
            cell = cells.get(c["key"]) or {}
            if cell.get("status") in ACCEPT and cell.get("value") is not None:
                row[mapping[c["key"]]] = cell["value"]
            else:
                row[mapping[c["key"]]] = None
                if c.get("required", True):
                    incomplete.append(c["label"])
        problems = []
        if validators.get("vat_ok") is False or validators.get("sum_ok") is False:
            problems.append(validators.get("note") or "figures don't add up")
        if incomplete:
            problems.append("missing: " + ", ".join(incomplete))
        if problems:
            row["_why"] = "; ".join(problems)
            exceptions.append(row)
            excepted += 1
        else:
            rows.append(row)
            added += 1
    sheet.rows = rows
    sheet.exceptions = exceptions
    sheet.updated_at = utcnow()
    audit(db, "master_sheet_append", actor=user, target=f"sheet:{sheet.id}", detail=f"batch {batch.id}: {added} rows, {excepted} exceptions, {skipped} already there")
    return {"added": added, "exceptions": excepted, "skipped": skipped, "sheet": {"id": sheet.id, "name": sheet.name, "rows": len(rows), "exceptions": len(exceptions)}}


def _client_of(cells: dict, columns: list[dict]) -> str:
    for c in columns:
        if any(w in (c["key"] + " " + c["label"]).lower() for w in ("client", "municipal", "vendor", "supplier", "site")):
            v = (cells.get(c["key"]) or {}).get("value")
            if v:
                return str(v)
    return ""
