"""PDF Studio API: templates, batches, cells (fill / verify / answer), export, master-sheet append,
profiles & versions (admin), spend and the model bake-off (admin), results."""
from __future__ import annotations

import io
import json
import re
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import Response, StreamingResponse
from pydantic import BaseModel
from sqlalchemy import or_
from sqlalchemy.orm import Session

from ..audit import log
from ..config import settings
from ..db import get_db, utcnow
from ..util import read_capped
from ..models.pdfstudio import COLUMN_TYPES, NO_RULE_TYPES
from ..models import Batch, BenchResult, Column, Correction, Document, Extraction, MasterSheet, Profile, ProfileVersion, StoredFile, Template, User
from ..pdfstudio import corrections as C
from ..pdfstudio import export as E
from ..pdfstudio import extract as X
from ..pdfstudio import formula as F
from ..pdfstudio import housekeeping as _housekeeping  # noqa: F401  (registers the hourly pdf-housekeeping routine)
from ..pdfstudio import lexicon as LX
from ..pdfstudio import profile as P
from ..pdfstudio import profiles as PR
from ..pdfstudio import scoring as SC
from ..pdfstudio.spans import cached, page_lines, quick_text, render_page_png
from ..security.deps import client_ip, require_admin, require_human, require_owner, require_tool
from ..security.rls import scope
from ..warehouse import storage
from .schemas import iso

router = APIRouter(prefix="/api/pdfstudio", tags=["pdfstudio"])
TOOL = require_tool("pdf-studio")
TYPE_LIST = "text, number, money, date, list, filename (the PDF's own name) or formula (worked out from other columns)"

from .pdfstudio_harvest import router as _harvest_router  # noqa: E402  (Harvest: folders watched continuously)

router.include_router(_harvest_router)


# ---- helpers ---------------------------------------------------------------------------------------
def _slug(s: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "_", str(s).lower()).strip("_")
    return s or "col"


def _template_q(db: Session, me: User):
    q = db.query(Template)
    if me.is_admin or me.is_axl:
        return q
    conds = [Template.owner_id == me.id]
    if me.team_id:
        conds.append((Template.team_id == me.team_id) & (Template.shared.is_(True)))
    return q.filter(or_(*conds))


def _get_template(db: Session, me: User, tid: int) -> Template:
    t = _template_q(db, me).filter(Template.id == tid).one_or_none()
    if t is None:
        raise HTTPException(404, "No such template.")
    return t


def _get_batch(db: Session, me: User, bid: int) -> Batch:
    b = scope(db.query(Batch), Batch, me).filter(Batch.id == bid).one_or_none()
    if b is None:
        raise HTTPException(404, "No such batch.")
    return b


def _get_doc(db: Session, me: User, did: int) -> tuple[Document, Batch]:
    d = db.get(Document, did)
    if d is None:
        raise HTTPException(404, "No such document.")
    b = _get_batch(db, me, d.batch_id)
    return d, b


def _col_out(c: Column) -> dict:
    cfg = dict(c.config or {})
    return {"id": c.id, "key": c.key, "label": c.label, "ctype": c.ctype, "required": c.required, "critical": c.critical, "hint": c.hint, "ordinal": c.ordinal, "expr": (cfg.get("expr") or "") if c.ctype == "formula" else "", "optional_reason": cfg.get("optional_reason") or "", "kind": _kind_label(c.ctype, c.required)}


def _kind_label(ctype: str, required: bool) -> str:
    """What an admin reads on a column: 'read from the page', 'optional', 'file name', 'calculated'."""
    if ctype == "filename":
        return "file name"
    if ctype == "formula":
        return "calculated"
    return "read from the page" if required else "optional"


def _columns_summary(columns: list[dict]) -> dict:
    k = E.column_kinds(columns)
    text = f"{k['total']} column{'s' if k['total'] != 1 else ''}"
    if k["optional"]:
        text += f" · {k['optional']} optional"
    if k["calculated"]:
        text += f" · {k['calculated']} calculated"
    if k["filename"]:
        text += f" · {k['filename']} file name"
    return dict(k, text=text)


def _template_out(t: Template, db: Session) -> dict:
    n_batches = db.query(Batch).filter(Batch.template_id == t.id).count()
    last = db.query(Batch).filter(Batch.template_id == t.id, Batch.accuracy.isnot(None)).order_by(Batch.created_at.desc()).first()
    return {"id": t.id, "name": t.name, "owner_id": t.owner_id, "team_id": t.team_id, "shared": t.shared, "master_sheet_id": t.master_sheet_id, "created_at": iso(t.created_at), "columns": [_col_out(c) for c in t.columns], "batches": n_batches, "accuracy": last.accuracy if last else None, "columns_summary": _columns_summary(E.columns_of(t))}


# ---- templates -------------------------------------------------------------------------------------
class ColumnIn(BaseModel):
    key: str = ""
    label: str
    ctype: str = "text"
    required: bool = True
    critical: Optional[bool] = None
    hint: str = ""
    expr: str = ""  # formula columns: e.g. "material_excl_vat * 1.10"
    optional_reason: Optional[str] = None


class TemplateIn(BaseModel):
    name: str
    columns: list[ColumnIn]
    shared: bool = True
    master_sheet_id: Optional[int] = None


def _apply_columns(t: Template, cols: list[ColumnIn]) -> None:
    if not cols:
        raise HTTPException(400, "Add at least one column.")
    if not any(c.ctype not in NO_RULE_TYPES for c in cols):
        raise HTTPException(400, "Add at least one column that is read from the PDF — file-name and formula columns need something to work from.")
    # keys first, so formulas can be checked against every column (including ones later in the list)
    seen: set[str] = set()
    keyed: list[tuple[str, ColumnIn]] = []
    for c in cols:
        if c.ctype not in COLUMN_TYPES:
            raise HTTPException(400, f"'{c.ctype}' is not a column type. Use {TYPE_LIST}.")
        key = _slug(c.key or c.label)
        base, n = key, 2
        while key in seen:
            key = f"{base}_{n}"
            n += 1
        seen.add(key)
        keyed.append((key, c))
    kinds = {k: c.ctype for k, c in keyed}
    old_cfg = {c.key: dict(c.config or {}) for c in t.columns}
    t.columns.clear()
    for i, (key, c) in enumerate(keyed, start=1):
        cfg = {}
        if c.ctype == "formula":
            expr = (c.expr or "").strip()
            if not expr:
                raise HTTPException(400, f"'{c.label}' is a formula column — type the formula, e.g. material_excl_vat * 1.10. Columns you can use: {', '.join(k for k, ct in kinds.items() if ct in F.NUMERIC_KINDS and k != key) or 'add a money or number column first'}.")
            try:
                F.validate(expr, kinds, own_key=key)
            except F.FormulaError as e:
                raise HTTPException(400, f"The formula for '{c.label}' isn't right: {e}") from e
            cfg["expr"] = expr
        reason = c.optional_reason if c.optional_reason is not None else old_cfg.get(key, {}).get("optional_reason")
        if reason and not c.required and c.ctype not in NO_RULE_TYPES:
            cfg["optional_reason"] = reason
        crit = c.critical if c.critical is not None else (c.ctype == "money" and c.required)
        if c.ctype in NO_RULE_TYPES:
            crit = False
        t.columns.append(Column(ordinal=i, key=key, label=c.label.strip() or key, ctype=c.ctype, required=c.required if c.ctype not in NO_RULE_TYPES else False, critical=crit, hint=c.hint or "", config=cfg))


@router.get("/templates")
def list_templates(db: Session = Depends(get_db), me: User = Depends(TOOL)):
    rows = _template_q(db, me).order_by(Template.created_at.desc()).all()
    return {"templates": [_template_out(t, db) for t in rows]}


@router.post("/templates")
def create_template(body: TemplateIn, db: Session = Depends(get_db), me: User = Depends(TOOL)):
    name = body.name.strip()
    if not name:
        raise HTTPException(400, "Give the template a name.")
    t = Template(name=name, owner_id=me.id, team_id=me.team_id, shared=body.shared, master_sheet_id=body.master_sheet_id)
    db.add(t)
    db.flush()
    _apply_columns(t, body.columns)
    db.flush()
    log(db, "pdf_template_created", actor=me, target=f"template:{t.id}", detail=f"{name}: {len(t.columns)} columns")
    return _template_out(t, db)


@router.get("/municipal")
def municipal_profiles(me: User = Depends(TOOL)):
    """1.0.75: the CBX Harvester's municipal profiles PDF Studio reads bills with."""
    from ..pdfstudio import municipal as MU

    return {"profiles": [{"id": p["profile_id"], "name": p.get("profile_name") or p["profile_id"], "short": p.get("short_name") or "", "custom": p["_custom"]} for p in MU.profiles()],
            "template": MU.TEMPLATE_NAME, "columns": [label for _k, label in MU.CANONICAL]}


@router.post("/templates/municipal")
def municipal_template(db: Session = Depends(get_db), me: User = Depends(TOOL)):
    """The Harvester's General template (28 columns), made once and shared; returned for use."""
    from ..models import Column as ColumnRow
    from ..pdfstudio import municipal as MU

    t = _template_q(db, me).filter(Template.name == MU.TEMPLATE_NAME).order_by(Template.id).first()
    if t is None:
        t = Template(name=MU.TEMPLATE_NAME, owner_id=me.id, team_id=me.team_id, shared=True)
        db.add(t)
        db.flush()
        for i, c in enumerate(MU.template_columns(), start=1):
            t.columns.append(ColumnRow(ordinal=i, key=c["key"], label=c["label"], ctype=c["ctype"], required=c["required"], critical=c["critical"], hint="", config={}))
        db.flush()
        log(db, "pdf_template_created", actor=me, target=f"template:{t.id}", detail=f"{t.name}: the CBX Harvester's 28 columns")
    return _template_out(t, db)


@router.post("/templates/from-excel")
async def template_from_excel(files: list[UploadFile] = File(...), db: Session = Depends(get_db), me: User = Depends(TOOL)):
    if not files:
        raise HTTPException(400, "Upload an Excel file.")
    up = files[0]
    data = await read_capped(up, settings.max_attachment_mb * 1024 * 1024, "template")
    try:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001
        raise HTTPException(400, "That file isn't an Excel workbook I can read (.xlsx).")
    from ..pdfstudio.excel_template import guess_ctype, guess_required, optional_reason, read_headers

    try:
        hdr = read_headers(data)
    except ValueError as e:
        raise HTTPException(400, str(e)) from e
    cols = []
    seen: set[str] = set()
    for i, h in enumerate(hdr.headers):
        if not h:
            continue
        key = _slug(h) or f"col_{i+1}"
        if key in seen:
            continue
        seen.add(key)
        sv = hdr.sample[i] if hdr.sample and i < len(hdr.sample) else None
        ct = guess_ctype(h, sv)
        req = guess_required(h, ct)
        cols.append(ColumnIn(key=key, label=h, ctype=ct, required=req, critical=(ct == "money" and req), optional_reason=optional_reason(h) if not req and ct not in NO_RULE_TYPES else None))
    if not cols:
        raise HTTPException(400, "The header row is empty.")
    name = re.sub(r"\.(xlsx|xlsm|xls)$", "", up.filename or "Template", flags=re.I).replace("_", " ").strip() or "Template"
    f = storage.store(db, me, up.filename or "template.xlsx", data, sub="Work", source="pdfstudio", mirror=False, meta={"kind": "pdfstudio-template"})
    t = Template(name=name, owner_id=me.id, team_id=me.team_id, shared=True, source_file_id=f.id)
    db.add(t)
    db.flush()
    _apply_columns(t, cols)
    db.flush()
    log(db, "pdf_template_created", actor=me, target=f"template:{t.id}", detail=f"from Excel '{up.filename}': {len(cols)} columns")
    return _template_out(t, db)


@router.delete("/templates/{tid}")
def delete_template(tid: int, db: Session = Depends(get_db), me: User = Depends(TOOL)):
    t = _get_template(db, me, tid)
    if t.owner_id != me.id and not me.is_admin:
        raise HTTPException(403, "Only the person who made this template (or an admin) can delete it.")
    n = db.query(Batch).filter(Batch.template_id == t.id).count()
    if n:
        # batches need their template (columns, labels) to show their sheet — the row can't go
        raise HTTPException(400, f"This template has {n} batch{'es' if n != 1 else ''} behind it, so it can't be deleted. Rename it instead.")
    name = t.name
    # profiles and bake-off rows point at the template; detach them so the delete doesn't fail
    for p in db.query(Profile).filter(Profile.template_id == t.id).all():
        p.template_id = None
    db.query(BenchResult).filter(BenchResult.template_id == t.id).delete()
    db.delete(t)
    db.flush()
    log(db, "pdf_template_deleted", actor=me, target=f"template:{tid}", detail=name)
    return {"ok": True}


class TemplatePatch(BaseModel):
    name: Optional[str] = None
    columns: Optional[list[ColumnIn]] = None
    shared: Optional[bool] = None
    master_sheet_id: Optional[int] = None


@router.patch("/templates/{tid}")
def patch_template(tid: int, body: TemplatePatch, db: Session = Depends(get_db), me: User = Depends(TOOL)):
    t = _get_template(db, me, tid)
    if t.owner_id != me.id and not me.is_admin:
        raise HTTPException(403, "Only the person who made this template (or an admin) can change it.")
    if body.name is not None:
        t.name = body.name.strip() or t.name
    if body.shared is not None:
        t.shared = body.shared
    if body.master_sheet_id is not None:
        t.master_sheet_id = body.master_sheet_id
    if body.columns is not None:
        _apply_columns(t, body.columns)
    db.flush()
    log(db, "pdf_template_changed", actor=me, target=f"template:{t.id}")
    return _template_out(t, db)


# ---- batches ---------------------------------------------------------------------------------------
class BatchIn(BaseModel):
    template_id: int
    name: str = ""


@router.post("/batches")
def create_batch(body: BatchIn, db: Session = Depends(get_db), me: User = Depends(TOOL)):
    t = _get_template(db, me, body.template_id)
    b = Batch(template_id=t.id, user_id=me.id, team_id=me.team_id, name=body.name.strip() or t.name, status="uploading", step="Waiting for PDFs", ceiling_zar=settings.model_batch_ceiling_zar)
    db.add(b)
    db.flush()
    log(db, "pdf_batch_created", actor=me, target=f"batch:{b.id}", detail=t.name)
    return {"id": b.id, "name": b.name, "status": b.status}


@router.get("/batches")
def list_batches(db: Session = Depends(get_db), me: User = Depends(TOOL)):
    rows = scope(db.query(Batch), Batch, me).order_by(Batch.created_at.desc()).limit(50).all()
    return {"batches": [_batch_brief(b) for b in rows]}


def _stalled(b: Batch) -> bool:
    """Marked as running but no job is alive for it (the server restarted mid-run)."""
    from ..pdfstudio import pipeline

    return b.status in ("learning", "running") and not pipeline.is_running(b.id)


def _batch_brief(b: Batch) -> dict:
    return {"id": b.id, "name": b.name, "status": b.status, "step": b.step, "progress": b.progress, "docs": len(b.documents), "cells_for_user": b.cells_for_user, "created_at": iso(b.created_at), "finished_at": iso(b.finished_at), "template_id": b.template_id, "stalled": _stalled(b)}


@router.delete("/batches/{bid}")
def delete_batch(bid: int, db: Session = Depends(get_db), me: User = Depends(TOOL)):
    """Only a batch nobody has read yet: an empty one a browser left behind, or one still uploading.
    Its PDFs stay in the person's warehouse folder."""
    b = _get_batch(db, me, bid)
    if b.user_id != me.id and not me.is_admin:
        raise HTTPException(403, "Only the person who made this batch (or an admin) can delete it.")
    n = len(b.documents)
    if b.status not in ("uploading", "queued") and n:
        raise HTTPException(400, "This batch has been read — it can't be deleted. The results stay for the record.")
    log(db, "pdf_batch_deleted", actor=me, target=f"batch:{b.id}", detail=f"'{b.name}' ({n} PDF{'s' if n != 1 else ''}, {b.status})")
    db.delete(b)
    db.flush()
    return {"ok": True, "deleted": bid}


@router.post("/batches/{bid}/files")
async def add_files(bid: int, files: list[UploadFile] = File(...), db: Session = Depends(get_db), me: User = Depends(TOOL)):
    b = _get_batch(db, me, bid)
    if b.status not in ("uploading", "queued"):
        state = {"learning": "C-EXA is reading it", "running": "C-EXA is reading it", "review": "it has been read and is waiting for you", "done": "it has been read", "failed": "it failed"}.get(b.status, f"it is {b.status}")
        raise HTTPException(409, f"This batch can't take more PDFs — {state}. Start a new batch for these files.")
    out = []
    rejected = []
    start = len(b.documents)
    # check every file before storing any, so one stray .xlsx in a drop doesn't sink the PDFs beside it
    loaded = []
    for up in files:
        data = await read_capped(up, 200 * 1024 * 1024)
        if data.startswith(b"%PDF"):
            loaded.append((up, data))
        else:
            rejected.append(up.filename or "file")
    if not loaded:
        raise HTTPException(400, f"'{rejected[0]}' isn't a PDF." if len(rejected) == 1 else "None of those files are PDFs.")
    for i, (up, data) in enumerate(loaded):
        text, pages = quick_text(data)
        f = storage.store(db, me, up.filename or f"document-{i+1}.pdf", data, sub="Inbox", source="pdfstudio", text=text, pages=pages, meta={"batch_id": b.id, "original_name": up.filename or f"document-{i+1}.pdf"})
        # PDF Studio handles these itself — keep them out of the inbox → master-sheet pipeline
        f.pipeline_status, f.detected_as, f.goes_to = "added", "PDF Studio batch", b.name or "PDF Studio"
        d = Document(batch_id=b.id, ordinal=start + i + 1, file_id=f.id, filename=f.name, sha256=f.sha256, pages=pages, status="queued")
        db.add(d)
        db.flush()
        out.append({"id": d.id, "ordinal": d.ordinal, "filename": d.filename, "pages": pages, "text_layer": bool(text.strip())})
    log(db, "pdf_files_added", actor=me, target=f"batch:{b.id}", detail=f"{len(out)} files" + (f", {len(rejected)} not PDFs" if rejected else ""))
    return {"documents": out, "count": start + len(out), "rejected": rejected}


@router.post("/batches/{bid}/start")
def start_batch(bid: int, db: Session = Depends(get_db), me: User = Depends(TOOL)):
    b = _get_batch(db, me, bid)
    from ..pdfstudio import pipeline

    if not b.documents:
        raise HTTPException(400, "Add at least one PDF first.")
    if b.status in ("review", "done"):
        raise HTTPException(400, "This batch is already finished.")
    if b.status in ("learning", "running") and not _stalled(b):
        raise HTTPException(400, "This batch is already running.")
    restart = b.status != "uploading"
    b.status = "learning"
    b.step = "Received"
    b.progress = 0.02
    b.error = ""
    db.commit()
    log(db, "pdf_batch_started", actor=me, target=f"batch:{b.id}", detail=f"{len(b.documents)} PDFs" + (" (started again)" if restart else ""))
    db.commit()
    pipeline.submit_batch(b.id)
    db.expire_all()
    b = db.get(Batch, bid)
    return {"id": b.id, "status": b.status, "step": b.step}


def _sheet_out(b: Batch, t: Template) -> dict:
    cols = E.columns_of(t)
    return {"columns": [dict(c, kind=_kind_label(c["ctype"], c["required"])) for c in cols], "rows": E.rows_of(b, cols)}


@router.post("/batches/{bid}/rerun")
def rerun_batch(bid: int, request: Request, db: Session = Depends(get_db), me: User = Depends(require_admin), _h: User = Depends(require_human)):
    """Re-read every PDF of a finished (or failed) batch with the template's CURRENT profile. Cells
    people filled or trusted stay; each document's earlier read is kept in Extraction.history."""
    from ..pdfstudio import pipeline

    b = _get_batch(db, me, bid)
    if not b.documents:
        raise HTTPException(400, "There are no PDFs in this batch to re-read.")
    if b.status in ("learning", "running") and not _stalled(b):
        raise HTTPException(400, "This batch is being read right now — wait for it to finish.")
    if b.status in ("uploading", "queued"):
        raise HTTPException(400, "This batch hasn't been handed to C-EXA yet — click 'Hand to C-EXA' instead.")
    b.status = "learning"
    b.step = "Received"
    b.progress = 0.02
    b.error = ""
    b.finished_at = None
    entries = list(b.log or [])
    who = me.name.split()[0] if me.name else "An admin"
    entries.append({"at": utcnow().isoformat(), "text": f"{who} asked me to re-read this batch with the current profile."})
    b.log = entries
    db.commit()
    log(db, "pdf_batch_rerun", actor=me, target=f"batch:{b.id}", detail=f"{len(b.documents)} PDFs, re-read with the current profile", ip=client_ip(request))
    db.commit()
    pipeline.submit_batch(b.id)
    db.expire_all()
    b = db.get(Batch, bid)
    return {"id": b.id, "status": b.status, "step": b.step}


@router.get("/batches/{bid}")
def get_batch(bid: int, db: Session = Depends(get_db), me: User = Depends(TOOL)):
    b = _get_batch(db, me, bid)
    t = db.get(Template, b.template_id)
    out = _batch_brief(b)
    cols = E.columns_of(t)
    out.update(
        {
            "template": {"id": t.id, "name": t.name, "columns": [_col_out(c) for c in t.columns]},
            "columns_summary": _columns_summary(cols),
            "cells_total": b.cells_total,
            "cells_filled": b.cells_filled,
            "cells_absent": SC.cell_counts(((d.extraction.cells if d.extraction else None) for d in b.documents if d.extraction), cols)["absent"] if b.status in ("review", "done") else 0,
            "runs_count": len(b.runs or []),
            "reruns": max((len(d.extraction.history or []) for d in b.documents if d.extraction), default=0),
            "documents": [{"id": d.id, "ordinal": d.ordinal, "filename": d.filename, "pages": d.pages, "scanned": d.scanned, "status": d.status, "is_sample": d.is_sample, "score": d.score if me.role != "user" else None} for d in b.documents],
            "accuracy_pct": round((b.accuracy or 0) * 100, 1) if b.accuracy is not None else None,
        }
    )
    if b.status in ("review", "done"):
        out["sheet"] = _sheet_out(b, t)
    if me.role != "user":
        out.update({"runs": b.runs or [], "by_column": b.by_column or {}, "log": b.log or [], "accuracy": b.accuracy, "spend_zar": b.spend_zar, "ceiling_zar": b.ceiling_zar, "error": b.error, "profiles": _batch_profiles(db, b)})
    return out


def _batch_profiles(db: Session, b: Batch) -> list[dict]:
    ids = {d.profile_id for d in b.documents if d.profile_id}
    out = []
    for pid in ids:
        p = db.get(Profile, pid)
        if p:
            v = PR.working_version(p)
            out.append({"id": p.id, "key": p.key, "label": p.label, "status": p.status, "version": v.number if v else None, "accuracy": p.accuracy, "docs": sum(1 for d in b.documents if d.profile_id == pid)})
    return out


# ---- cells -----------------------------------------------------------------------------------------
class CellIn(BaseModel):
    document_id: int
    column_key: str
    value: str


@router.post("/batches/{bid}/cells")
def fill_cell(bid: int, body: CellIn, db: Session = Depends(get_db), me: User = Depends(TOOL), _h: User = Depends(require_human)):
    b = _get_batch(db, me, bid)
    if b.status not in ("review", "done"):
        raise HTTPException(400, "The sheet isn't ready yet — C-EXA is still reading this batch.")
    d = db.get(Document, body.document_id)
    if d is None or d.batch_id != b.id:
        raise HTTPException(404, "That PDF isn't in this batch.")
    col = db.query(Column).filter(Column.template_id == b.template_id, Column.key == body.column_key).one_or_none()
    if col is None:
        raise HTTPException(404, "No such column.")
    if col.ctype == "filename":
        raise HTTPException(400, "This column is the PDF's file name — it can't be typed over.")
    if col.ctype == "formula":
        raise HTTPException(400, f"'{col.label}' is worked out from other columns ({F.pretty((col.config or {}).get('expr', ''))}). Fix the columns it uses and it follows.")
    value = body.value.strip()
    if not value:
        raise HTTPException(400, "Type a value first.")
    corr, cell = C.fill(db, b, d, col, value, me)
    db.flush()
    learned = None
    if corr.verified and d.profile_id:
        prof = db.get(Profile, d.profile_id)
        try:
            ver = C.learn(db, prof, d.extraction, col.key, value, me)
        except Exception:  # noqa: BLE001  learning must never block a save
            ver = None
        if ver is not None:
            corr.learned_version_id = ver.id
            learned = {"profile_id": prof.id, "version": ver.number, "reason": ver.reason, "backtest": ver.backtest}
    out = {"correction_id": corr.id, "verified": corr.verified, "status": cell["status"], "cell": cell, "learned": learned if me.role != "user" else (bool(learned)), "cells_for_user": b.cells_for_user, "batch_status": b.status}
    if not corr.verified:
        out["question"] = f"I can't find “{value}” on this bill — is it from another document, or should I trust you and note the source?"
        out["options"] = [{"id": "other_source", "label": "It's from another document", "needs_note": True}, {"id": "trust_me", "label": "Trust me", "needs_note": False}, {"id": "retract", "label": "Take it back", "needs_note": False}]
    return out


class AnswerIn(BaseModel):
    correction_id: int
    answer: str
    note: str = ""


@router.post("/batches/{bid}/cells/answer")
def answer_cell(bid: int, body: AnswerIn, db: Session = Depends(get_db), me: User = Depends(TOOL), _h: User = Depends(require_human)):
    b = _get_batch(db, me, bid)
    if b.status not in ("review", "done"):
        raise HTTPException(400, "The sheet isn't ready yet — C-EXA is still reading this batch.")
    corr = db.get(Correction, body.correction_id)
    if corr is None:
        raise HTTPException(404, "No such entry.")
    ext = db.get(Extraction, corr.extraction_id)
    d = db.get(Document, ext.document_id) if ext else None
    if d is None or d.batch_id != b.id:
        raise HTTPException(404, "That entry isn't in this batch.")
    try:
        cell = C.answer(db, corr, body.answer, body.note, me)
    except ValueError as e:
        raise HTTPException(400, str(e))
    return {"cell": cell, "status": cell["status"], "cells_for_user": b.cells_for_user, "batch_status": b.status}


# ---- documents: context and page images ------------------------------------------------------------------
@router.get("/documents/{did}/context")
def doc_context(did: int, col: str = Query(""), db: Session = Depends(get_db), me: User = Depends(TOOL)):
    d, b = _get_doc(db, me, did)
    spans = cached(d.sha256)
    if spans is None:
        return {"page": 1, "lines": [], "highlight": None, "label": "", "note": "This PDF hasn't been read yet."}
    cell = ((d.extraction.cells if d.extraction else {}) or {}).get(col) or {}
    column = db.query(Column).filter(Column.template_id == b.template_id, Column.key == col).one_or_none()
    page = cell.get("page")
    idx = None
    label_used = ""
    lines_all = None
    # 1. the line the cell came from
    if page and cell.get("quote"):
        lines_all = page_lines(spans, page)
        q = X.M.normalise_text(cell["quote"])
        for i, l in enumerate(lines_all):
            if q and q in X.M.normalise_text(l["text"]):
                idx = i
                label_used = cell.get("anchor") or ""
                break
    # 2. else the expected label from the profile's rules / column names
    if idx is None and column is not None:
        labels = []
        if d.profile_id:
            prof = db.get(Profile, d.profile_id)
            ver = PR.working_version(prof) if prof else None
            for r in ((ver.ruleset or {}).get("fields", {}).get(col, {}).get("rules", []) if ver else []):
                if r.get("label"):
                    labels.append(r["label"])
        labels += P.candidate_labels(_col_out(column))
        for lab in labels:
            anchors = X.find_anchors(spans, lab, fuzzy=True, limit=1)
            if anchors:
                line = anchors[0]["line"]
                page = line["page"]
                lines_all = page_lines(spans, page)
                idx = lines_all.index(line)
                label_used = lab
                break
    found = idx is not None
    if idx is None:
        page = page or 1
        lines_all = page_lines(spans, page)
        idx = 0 if lines_all else None
    lo, hi = (max(0, idx - 4), idx + 5) if idx is not None else (0, 8)
    window = lines_all[lo:hi] if lines_all else []
    return {
        "page": page,
        "pages": spans.get("pages", d.pages),
        "scanned": bool(spans.get("scanned")),
        "label": label_used,
        "found": found,
        "lines": [{"text": l["text"], "highlight": found and (lo + i) == idx} for i, l in enumerate(window)],
        "cell": {k: cell.get(k) for k in ("value", "raw", "quote", "status", "note", "page", "score") if k in cell},
        "note": ("Scanned page — the digits may be unclear. " if spans.get("scanned") else "") + ("" if found else f"I couldn't find a label for '{column.label if column else col}' on this PDF — this is the top of page {page}."),
    }


@router.get("/documents/{did}/page/{n}.png")
def doc_page_png(did: int, n: int, db: Session = Depends(get_db), me: User = Depends(TOOL)):
    d, _ = _get_doc(db, me, did)
    f = db.get(StoredFile, d.file_id)
    try:
        png = render_page_png(storage.read_bytes(f), n)
    except ValueError:
        raise HTTPException(404, "No such page.")
    except Exception:  # noqa: BLE001
        raise HTTPException(500, "I couldn't render that page.")
    return Response(content=png, media_type="image/png", headers={"Cache-Control": "private, max-age=600"})


# ---- export / master sheet -------------------------------------------------------------------------------
@router.get("/batches/{bid}/export.xlsx")
def export_xlsx(bid: int, db: Session = Depends(get_db), me: User = Depends(TOOL)):
    b = _get_batch(db, me, bid)
    if b.status not in ("review", "done"):
        raise HTTPException(400, "The sheet isn't ready yet.")
    t = db.get(Template, b.template_id)
    data, f = E.export_and_store(db, b, t, me)
    name = f.name if f else "results.xlsx"
    return StreamingResponse(io.BytesIO(data), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.post("/batches/{bid}/to-master/{sid}")
def to_master(bid: int, sid: int, db: Session = Depends(get_db), me: User = Depends(TOOL), _h: User = Depends(require_human)):
    b = _get_batch(db, me, bid)
    if b.status not in ("review", "done"):
        raise HTTPException(400, "The sheet isn't ready yet.")
    sheet = db.get(MasterSheet, sid)
    if sheet is None or not sheet.is_active:
        raise HTTPException(404, "No such master sheet.")
    t = db.get(Template, b.template_id)
    return E.append_to_master(db, b, t, sheet, me)


@router.get("/master-sheets")
def master_sheet_options(db: Session = Depends(get_db), me: User = Depends(TOOL)):
    rows = db.query(MasterSheet).filter(MasterSheet.is_active.is_(True)).order_by(MasterSheet.name).all()
    return {"sheets": [{"id": s.id, "name": s.name, "rows": len(s.rows or []), "columns": [c.get("label") for c in (s.columns or [])]} for s in rows]}


# ---- results ---------------------------------------------------------------------------------------------
@router.get("/results")
def results(db: Session = Depends(get_db), me: User = Depends(TOOL)):
    batches = scope(db.query(Batch), Batch, me).filter(Batch.status.in_(["review", "done"])).order_by(Batch.finished_at.desc()).limit(50).all()
    fq = db.query(StoredFile).filter(StoredFile.source == "result")
    if not me.is_admin:
        fq = fq.filter(StoredFile.uploaded_by_id == me.id)
    files = fq.order_by(StoredFile.created_at.desc()).limit(50).all()
    owners = {u.id: u.name for u in db.query(User).all()} if me.is_admin else {}
    return {
        "batches": [dict(_batch_brief(b), owner=owners.get(b.user_id, "")) for b in batches],
        "files": [{"id": f.id, "name": f.name, "kind": f.kind, "size": f.size, "at": iso(f.created_at), "owner": owners.get(f.uploaded_by_id, ""), "download": f"/api/warehouse/files/{f.id}/download"} for f in files],
    }


# ---- profiles (admin) --------------------------------------------------------------------------------------
def _version_out(v: ProfileVersion, db: Session) -> dict:
    return {"id": v.id, "number": v.number, "reason": v.reason, "author": v.author, "author_name": _author_name(v.author, db), "backtest": v.backtest or {}, "approved_at": iso(v.approved_at), "approved_by_id": v.approved_by_id, "created_at": iso(v.created_at), "pending": v.approved_at is None and v.author != "axl", "rules": sum(len(f.get("rules", [])) for f in (v.ruleset or {}).get("fields", {}).values()), "fields": len((v.ruleset or {}).get("fields", {}))}


def _author_name(author: str, db: Session) -> str:
    if author == "axl":
        return "C-EXA"
    m = re.match(r"user:(\d+)", author or "")
    if m:
        u = db.get(User, int(m.group(1)))
        return u.name if u else "a user"
    return author or ""


def _profile_out(p: Profile, db: Session, full: bool = False) -> dict:
    v = PR.working_version(p)
    pending = [x for x in p.versions if x.approved_at is None and x.author != "axl"]
    # `accuracy` is the HELD-OUT figure — what the profile scored on documents it was not built from
    # — or null when there were too few documents to hold any back. `honest` says which.
    h = PR.honest_accuracy(p)
    out = {"id": p.id, "key": p.key, "label": p.label, "status": p.status, "template_id": p.template_id, "docs_seen": p.docs_seen, "accuracy": h["accuracy"], "honest": h, "version": v.number if v else None, "versions_count": len(p.versions), "pending": len(pending), "created_at": iso(p.created_at), "shared": p.shared}
    rs = (v.ruleset if v else None) or {}
    out["seeded_from"] = rs.get("seeded_from") or None
    if full:
        out["notes"] = v.notes if v else ""
        out["versions"] = [_version_out(x, db) for x in reversed(p.versions)]
        out["ruleset"] = rs
        out["live_version_id"] = p.live_version_id
        out["fields"] = _fields_out(p, rs, db)
        out["scoring"] = {"base": SC.BASE, "ocr": SC.OCR_FACTOR, "arith": SC.ARITH_FACTOR, "parse": SC.PARSE_FACTOR, "gates": {"batch": SC.GATE_BATCH, "profile": SC.GATE_PROFILE, "critical": SC.GATE_CRITICAL}}
    return out


def _fields_out(p: Profile, rs: dict, db: Session) -> list[dict]:
    """Per-column rule list for the admin page: one entry per template column (or per field when the
    template is gone), with the rules in the order they are tried."""
    fields = rs.get("fields") or {}
    t = db.get(Template, p.template_id) if p.template_id else None
    cols = [_col_out(c) for c in t.columns] if t else [{"id": None, "key": k, "label": k, "ctype": "", "hint": "", "required": True, "expr": "", "optional_reason": "", "kind": "read from the page"} for k in fields]
    seeded = (rs.get("seeded_from") or {}).get("columns") or {}
    out = []
    for c in cols:
        f = fields.get(c["key"]) or {}
        rules = [{k: r.get(k) for k in ("id", "type", "label", "pick", "page", "pattern", "value", "op", "args", "after", "before", "quality", "hits", "of") if r.get(k) is not None} for r in (f.get("rules") or [])]
        out.append({**c, "rules": rules, "has_rules": bool(rules), "needs_rules": c.get("ctype") not in NO_RULE_TYPES, "seeded_from": seeded.get(c["key"]) or f.get("seeded_from"), "hint_used": f.get("hint")})
    return out


@router.post("/templates/{tid}/columns/{cid}/reauthor")
def reauthor_column(tid: int, cid: int, request: Request, db: Session = Depends(get_db), me: User = Depends(require_admin), _h: User = Depends(require_human)):
    """Re-author ONE column's rules in every live profile of this template. Every other field is
    copied unchanged into the new version (reason: "Re-authored column X")."""
    from ..pdfstudio import pipeline

    t = db.get(Template, tid)
    if t is None:
        raise HTTPException(404, "No such template.")
    col = next((c for c in t.columns if c.id == cid), None)
    if col is None:
        raise HTTPException(404, "That column isn't on this template.")
    profs = db.query(Profile).filter(Profile.template_id == t.id).all()
    if not profs:
        raise HTTPException(400, "This template has no profile yet — run a batch first and C-EXA will draft one.")
    done, skipped = [], []
    for p in profs:
        v = pipeline.author_missing_columns(db, t, p, force_keys=(col.key,))
        if v is None:
            skipped.append(p.label)
        else:
            done.append({"profile_id": p.id, "profile": p.label, "version": v.number, "reason": v.reason, "rules": len((v.ruleset.get("fields") or {}).get(col.key, {}).get("rules") or [])})
    db.flush()
    log(db, "pdf_column_reauthored", actor=me, target=f"template:{t.id}", detail=f"'{col.label}': {len(done)} profile(s) → new version" + (f"; no samples for {', '.join(skipped)}" if skipped else ""), ip=client_ip(request))
    return {"column": _col_out(col), "versions": done, "skipped": skipped, "message": (f"Re-authored '{col.label}' in {len(done)} profile{'s' if len(done) != 1 else ''}; the other columns' rules are unchanged." if done else "I have no sample pages for this template's profiles yet — run a batch first.")}


@router.get("/profiles")
def list_profiles(db: Session = Depends(get_db), me: User = Depends(require_admin)):
    rows = db.query(Profile).order_by(Profile.created_at.desc()).all()
    return {"profiles": [_profile_out(p, db) for p in rows], "pending": C.pending_count(db)}


@router.get("/profiles/{pid}")
def get_profile(pid: int, db: Session = Depends(get_db), me: User = Depends(require_admin)):
    p = db.get(Profile, pid)
    if p is None:
        raise HTTPException(404, "No such profile.")
    return _profile_out(p, db, full=True)


@router.get("/profiles/{pid}/deletable")
def profile_deletable(pid: int, db: Session = Depends(get_db), me: User = Depends(require_admin)):
    """What deleting this profile would cost, so the confirm can say it in plain words instead of
    'are you sure?'. Nothing is changed by asking."""
    p = db.get(Profile, pid)
    if p is None:
        raise HTTPException(404, "No such profile.")
    return _delete_cost(db, p)


def _delete_cost(db: Session, p: Profile) -> dict:
    vids = [v.id for v in p.versions]
    docs = db.query(Document).filter(Document.profile_id == p.id).count()
    reads = db.query(Extraction).filter(Extraction.profile_version_id.in_(vids)).count() if vids else 0
    return {"id": p.id, "key": p.key, "label": p.label, "status": p.status, "versions": len(vids), "documents": docs, "readings": reads}


def _export_profile(db: Session, me: User, p: Profile) -> Optional[dict]:
    """Write the whole profile — every version, every rule — to the person's Results folder before it
    goes, so 'delete' is never 'gone for good'. A failure here does not stop the delete; the audit
    row is the other record."""
    try:
        payload = {
            "exported_at": iso(utcnow()),
            "exported_by": me.name,
            "profile": {"id": p.id, "key": p.key, "label": p.label, "template_id": p.template_id, "status": p.status, "docs_seen": p.docs_seen, "accuracy": p.accuracy, "created_at": iso(p.created_at), "created_by": p.created_by, "live_version_id": p.live_version_id},
            "versions": [{"id": v.id, "number": v.number, "reason": v.reason, "author": v.author, "notes": v.notes, "backtest": v.backtest or {}, "created_at": iso(v.created_at), "approved_at": iso(v.approved_at), "ruleset": v.ruleset or {}} for v in p.versions],
        }
        data = json.dumps(payload, indent=2, ensure_ascii=False).encode("utf-8")
        f = storage.store(db, me, f"deleted-profile-{p.key}.json", data, sub="Results", source="result", mirror=False, meta={"kind": "pdfstudio-profile-export", "profile_key": p.key})
        return {"file_id": f.id, "name": f.name, "download": f"/api/warehouse/files/{f.id}/download"}
    except Exception:
        return None


def _delete_profile(db: Session, me: User, p: Profile) -> dict:
    """Delete a profile and its versions. The documents and the readings STAY — only the pointer at
    the rules that produced them goes null, because the values are the record and the rules are not."""
    cost = _delete_cost(db, p)
    cost["export"] = _export_profile(db, me, p)
    vids = [v.id for v in p.versions]
    db.query(Document).filter(Document.profile_id == p.id).update({Document.profile_id: None}, synchronize_session=False)
    if vids:
        db.query(Extraction).filter(Extraction.profile_version_id.in_(vids)).update({Extraction.profile_version_id: None}, synchronize_session=False)
        db.query(Correction).filter(Correction.learned_version_id.in_(vids)).update({Correction.learned_version_id: None}, synchronize_session=False)
    p.live_version_id = None  # self-reference: clear it before the versions go
    db.flush()
    db.delete(p)
    db.flush()
    return cost


@router.delete("/profiles/{pid}")
def delete_profile(pid: int, request: Request, db: Session = Depends(get_db), me: User = Depends(require_owner), _h: User = Depends(require_human)):
    """Owner only, and never C-EXA: a profile is how a layout gets read, so removing one is a
    production decision. Exports it first, then detaches and deletes."""
    p = db.get(Profile, pid)
    if p is None:
        raise HTTPException(404, "No such profile.")
    out = _delete_profile(db, me, p)
    log(db, "pdf_profile_deleted", actor=me, target=f"profile:{pid}", detail=f"'{out['label']}' ({out['key']}) · {out['versions']} version(s) · {out['documents']} document(s) detached · {out['readings']} reading(s) kept" + (f" · exported as {out['export']['name']}" if out.get("export") else " · export failed, audit is the record"), ip=client_ip(request))
    return {"ok": True, "deleted": out, "message": f"Deleted '{out['label']}'. {out['readings']} reading{'s' if out['readings'] != 1 else ''} kept — the values stay, only the link to the rules is gone."}


class ProfilePurge(BaseModel):
    ids: Optional[list[int]] = None
    template_id: Optional[int] = None
    below_gate_only: bool = True


@router.post("/profiles/purge")
def purge_profiles(body: ProfilePurge, request: Request, db: Session = Depends(get_db), me: User = Depends(require_owner), _h: User = Depends(require_human)):
    """Delete several at once — the case this exists for is a template whose profiles were all
    authored before the layout was identified properly, where one-at-a-time is 21 clicks."""
    q = db.query(Profile)
    if body.ids:
        q = q.filter(Profile.id.in_(body.ids))
    elif body.template_id:
        q = q.filter(Profile.template_id == body.template_id)
        if body.below_gate_only:
            q = q.filter(Profile.status == "below_gate")
    else:
        raise HTTPException(400, "Tell me which profiles: either a list of ids, or a template.")
    rows = q.order_by(Profile.id).all()
    if not rows:
        raise HTTPException(404, "Nothing matched — no profiles to delete.")
    done = [_delete_profile(db, me, p) for p in rows]
    docs = sum(d["documents"] for d in done)
    reads = sum(d["readings"] for d in done)
    log(db, "pdf_profiles_purged", actor=me, target=f"template:{body.template_id}" if body.template_id else "profiles", detail=f"{len(done)} profile(s): {', '.join(d['key'] for d in done)} · {docs} document(s) detached · {reads} reading(s) kept", ip=client_ip(request))
    return {"ok": True, "deleted": done, "count": len(done), "documents": docs, "readings": reads, "message": f"Deleted {len(done)} profile{'s' if len(done) != 1 else ''}. {reads} reading{'s' if reads != 1 else ''} kept, {docs} document{'s' if docs != 1 else ''} now unassigned — the next batch re-reads them from scratch."}


@router.post("/profiles/{pid}/versions/{vid}/approve")
def approve_version(pid: int, vid: int, request: Request, db: Session = Depends(get_db), me: User = Depends(require_admin), _h: User = Depends(require_human)):
    v = db.get(ProfileVersion, vid)
    if v is None or v.profile_id != pid:
        raise HTTPException(404, "No such version.")
    if v.approved_at is not None:
        raise HTTPException(400, "This version is already approved.")
    C.approve(db, v, me)
    log(db, "pdf_profile_approved", actor=me, target=f"profile:{pid}", detail=f"v{v.number}", ip=client_ip(request))
    return _profile_out(db.get(Profile, pid), db, full=True)


@router.get("/lexicon")
def get_lexicon(db: Session = Depends(get_db), me: User = Depends(require_admin)):
    """What C-EXA knows: the labels she has learned across every layout, and how she learned them."""
    return LX.stats(db, top=14)


@router.get("/spend")
def get_spend(db: Session = Depends(get_db), me: User = Depends(require_admin)):
    return PR.spend(db)


@router.get("/bench/{tid}")
def get_bench(tid: int, db: Session = Depends(get_db), me: User = Depends(require_admin)):
    rows = db.query(BenchResult).filter(BenchResult.template_id == tid).order_by(BenchResult.id).all()
    return {"rows": [_bench_out(b) for b in rows]}


@router.post("/bench/{tid}")
def run_bench(tid: int, db: Session = Depends(get_db), me: User = Depends(require_admin)):
    t = db.get(Template, tid)
    if t is None:
        raise HTTPException(404, "No such template.")
    rows = PR.run_bench(db, t)
    meta = getattr(rows[0], "_meta", {}) if rows else {}
    log(db, "pdf_bench_run", actor=me, target=f"template:{tid}", detail=f"{len(rows)} rows, gold cells {meta.get('gold_cells', 0)}")
    return {"rows": [_bench_out(b) for b in rows], "gold": meta}


def _bench_out(b: BenchResult) -> dict:
    return {"id": b.id, "model": b.model, "runs_on": b.runs_on, "docs": b.docs, "precision": b.precision, "recall": b.recall, "hallucination_rate": b.hallucination_rate, "cost_per_1000": b.cost_per_1000, "latency_s": b.latency_s, "role": b.role, "recommended": b.recommended, "at": iso(b.at)}
