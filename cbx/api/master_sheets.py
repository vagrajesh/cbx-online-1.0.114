"""Master Sheets: living workbooks. Users see the rows their row-level rules allow (and rows they
added); admins see everything. Downloads are built from the filtered rows only."""
from __future__ import annotations

import io
import re
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from ..audit import log
from ..db import get_db, utcnow
from ..models import MasterSheet, Team, User
from ..models.pdfstudio import COLUMN_TYPES
from ..security.deps import client_ip, require_admin, require_human, require_tool
from ..security.rls import filter_rows
from .schemas import iso

router = APIRouter(prefix="/api/master-sheets", tags=["master-sheets"])
TOOL = require_tool("master-sheets")


def _slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def _visible(db: Session, me: User, sheet: MasterSheet) -> tuple[list[dict], list[dict]]:
    rows = filter_rows(db, me, "master_rows", sheet.rows or [])
    exc = filter_rows(db, me, "master_rows", sheet.exceptions or [])
    return rows, exc


def _brief(db: Session, me: User, s: MasterSheet) -> dict:
    rows, exc = _visible(db, me, s)
    team = db.get(Team, s.team_id) if s.team_id else None
    return {"id": s.id, "name": s.name, "slug": s.slug, "description": s.description, "team": team.name if team else "", "columns": s.columns or [], "rows": len(rows), "rows_total": len(s.rows or []) if me.is_admin else None, "exceptions": len(exc), "updated_at": iso(s.updated_at), "created_at": iso(s.created_at), "is_active": s.is_active}


@router.get("")
def list_sheets(db: Session = Depends(get_db), me: User = Depends(TOOL)):
    q = db.query(MasterSheet).order_by(MasterSheet.name)
    if not me.is_admin:
        q = q.filter(MasterSheet.is_active.is_(True))
    return {"sheets": [_brief(db, me, s) for s in q.all()], "can_edit": me.is_admin}


def _get(db: Session, me: User, sid: int) -> MasterSheet:
    s = db.get(MasterSheet, sid)
    if s is None or (not s.is_active and not me.is_admin):
        raise HTTPException(404, "No such master sheet.")
    return s


@router.get("/{sid}")
def get_sheet(sid: int, db: Session = Depends(get_db), me: User = Depends(TOOL)):
    s = _get(db, me, sid)
    rows, exc = _visible(db, me, s)
    out = _brief(db, me, s)
    out["data"] = rows[-500:]
    out["exception_rows"] = exc[-200:]
    return out


@router.get("/{sid}/export.xlsx")
def export_sheet(sid: int, request: Request, db: Session = Depends(get_db), me: User = Depends(TOOL)):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    s = _get(db, me, sid)
    rows, exc = _visible(db, me, s)  # only the rows this person may see
    cols = s.columns or []
    wb = Workbook()
    ws = wb.active
    ws.title = "Rows"
    head = [c.get("label") or c.get("key") for c in cols] + ["Added by", "Added at", "Source"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r.get(c["key"]) for c in cols] + [r.get("_by_name") or r.get("_by"), r.get("_at"), (r.get("_src") or {}).get("name", "")])
    we = wb.create_sheet("Exceptions")
    we.append(head + ["Why"])
    for c in we[1]:
        c.font = Font(bold=True)
    for r in exc:
        we.append([r.get(c["key"]) for c in cols] + [r.get("_by_name") or r.get("_by"), r.get("_at"), (r.get("_src") or {}).get("name", ""), r.get("_why", "")])
    buf = io.BytesIO()
    wb.save(buf)
    log(db, "master_sheet_downloaded", actor=me, target=f"sheet:{s.id}", detail=f"{len(rows)} rows visible", ip=client_ip(request))
    return StreamingResponse(io.BytesIO(buf.getvalue()), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="%s.xlsx"' % re.sub(r"[^A-Za-z0-9 ._()-]+", "", s.name)})


# ---- admin ------------------------------------------------------------------------------------------
class ColIn(BaseModel):
    key: str = ""
    label: str
    ctype: str = "text"


class SheetIn(BaseModel):
    name: str
    description: str = ""
    team_id: Optional[int] = None
    columns: list[ColIn] = []


def _cols(cols: list[ColIn]) -> list[dict]:
    out, seen = [], set()
    for c in cols:
        if c.ctype not in COLUMN_TYPES:
            raise HTTPException(400, f"'{c.ctype}' is not a column type.")
        key = re.sub(r"[^a-z0-9]+", "_", (c.key or c.label).lower()).strip("_") or "col"
        base, n = key, 2
        while key in seen:
            key = f"{base}_{n}"
            n += 1
        seen.add(key)
        out.append({"key": key, "label": c.label.strip() or key, "ctype": c.ctype})
    return out


@router.post("")
def create_sheet(body: SheetIn, db: Session = Depends(get_db), me: User = Depends(require_admin), _h: User = Depends(require_human)):
    name = body.name.strip()
    if not name:
        raise HTTPException(400, "Give the sheet a name.")
    if db.query(MasterSheet).filter(MasterSheet.slug == _slug(name)).count():
        raise HTTPException(400, "There is already a sheet with that name.")
    s = MasterSheet(name=name, slug=_slug(name), description=body.description.strip(), team_id=body.team_id, columns=_cols(body.columns))
    db.add(s)
    db.flush()
    log(db, "master_sheet_created", actor=me, target=f"sheet:{s.id}", detail=f"{name}: {len(s.columns)} columns")
    return _brief(db, me, s)


class SheetPatch(BaseModel):
    name: Optional[str] = None
    description: Optional[str] = None
    columns: Optional[list[ColIn]] = None
    is_active: Optional[bool] = None


@router.patch("/{sid}")
def patch_sheet(sid: int, body: SheetPatch, db: Session = Depends(get_db), me: User = Depends(require_admin), _h: User = Depends(require_human)):
    s = db.get(MasterSheet, sid)
    if s is None:
        raise HTTPException(404, "No such master sheet.")
    changes = []
    if body.name is not None and body.name.strip() and body.name.strip() != s.name:
        new_name = body.name.strip()
        clash = db.query(MasterSheet).filter(MasterSheet.id != s.id, (MasterSheet.name == new_name) | (MasterSheet.slug == _slug(new_name))).count()
        if clash:
            raise HTTPException(400, "There is already a sheet with that name.")
        s.name, s.slug = new_name, _slug(new_name)
        changes.append("renamed")
    if body.description is not None:
        s.description = body.description.strip()
    if body.columns is not None:
        s.columns = _cols(body.columns)
        changes.append(f"{len(s.columns)} columns")
    if body.is_active is not None:
        s.is_active = body.is_active
        changes.append("active" if body.is_active else "archived")
    s.updated_at = utcnow()
    log(db, "master_sheet_changed", actor=me, target=f"sheet:{s.id}", detail="; ".join(changes))
    return _brief(db, me, s)
