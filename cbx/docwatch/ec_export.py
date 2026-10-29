"""1.0.99: Entity Compliance → one Excel workbook ("Export" on the tool's grid page) — CBX-Connect 2.10.4's book,
built from the tool folder's own tables (DATA_DIR/tools/entity-compliance.sqlite) and its own "who sees what" rules.

Only the companies the person may see (the tables' scopes, as the grid reads them: open companies, the ones
given to them, theirs; "only theirs" grants see no open ones; admins all); archived ones only on their own
sheet, and only for admins. Sheets: Overview, Documents, Expiring and expired, Entities, Archived entities
(admins), Document types, About. English only. Text cells go through xlsafe; dates are real Excel dates.

`workbook()` below is CBX-Connect 2.10.4's cbx/compliance/export.py, unchanged (it takes plain values only).
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime, timedelta, timezone
from typing import Callable, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .. import xlsafe

GREEN = "003F2D"  # the Harvester's header green
# the grid's colours as light fills, with a dark text of the same colour so the words stay readable
STATE_FILL = {
    "ok": ("DCEFE3", "0F5C36"), "no_expiry": ("DCEFE3", "0F5C36"), "expiring": ("FBE8C8", "7A4A00"), "expired": ("F5D3CC", "8A2A1B"),
    "needs_date": ("FFFFFF", "5F6B64"), "missing": ("FFFFFF", "5F6B64"), "not_required": ("E1E5E0", "5F6B64"),
}
OVERALL = {"expired": ("Expired", "expired"), "expiring": ("Expiring", "expiring"), "missing": ("Missing", "missing"), "ok": ("OK", "ok"), "empty": ("OK", "ok")}
WORST_ORDER = {"expired": 0, "expiring": 1, "missing": 2, "ok": 3, "empty": 4}
TODO = ("expired", "expiring", "missing", "needs_date")
DATE_FMT = "dd mmm yyyy"
TIME_FMT = "dd mmm yyyy hh:mm"
MAX_WIDTH = 40


def _en(text: str, **kw) -> str:
    return text.format(**kw) if kw else text


def _en_date(d: date) -> str:
    return f"{d.day} {d:%b} {d.year}"


def local(dt, tz: str = "Africa/Johannesburg") -> Optional[datetime]:
    """A stored time (UTC; naive means UTC) as the business's wall-clock time, without a zone (Excel has none)."""
    if dt in (None, ""):
        return None
    if isinstance(dt, str):
        try:
            dt = datetime.fromisoformat(dt.replace("Z", "+00:00"))
        except ValueError:
            return None
    if not isinstance(dt, datetime):
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    try:
        from zoneinfo import ZoneInfo

        zone = ZoneInfo(tz)
    except Exception:  # noqa: BLE001 — Windows without tzdata: SAST has no daylight saving
        zone = timezone(timedelta(hours=2))
    return dt.astimezone(zone).replace(tzinfo=None)


def state_word(status: str, expiry: Optional[date], today: date, T: Callable = _en, D: Callable = _en_date) -> str:
    """The cell's state in words: "Valid until 12 Mar 2027", "Expires in 23 days", "Expired 3 Feb 2026"…"""
    if status == "ok":
        return T("Valid until {date}", date=D(expiry)) if expiry else T("Valid")
    if status == "no_expiry":
        return T("Valid, doesn't expire")
    if status == "expiring" and expiry:
        n = (expiry - today).days
        return T("Expires today") if n <= 0 else T("Expires in 1 day") if n == 1 else T("Expires in {n} days", n=n)
    if status == "expired" and expiry:
        return T("Expired {date}", date=D(expiry))
    if status == "needs_date":
        return T("Filed, no expiry date")
    if status == "not_required":
        return T("Not required")
    return T("Missing")


def _paint(c, status: str) -> None:
    fill, ink = STATE_FILL.get(status, STATE_FILL["missing"])
    c.fill = PatternFill("solid", fgColor=fill)
    c.font = Font(color=ink, bold=status == "expired")


def _sheet_name(s: str) -> str:
    return re.sub(r"[\[\]:*?/\\]", " ", s).strip()[:31] or "Sheet"


def _sheet(wb: Workbook, title: str, headers: list, *, first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = _sheet_name(title)
    ws.append(headers)
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor=GREEN)
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    return ws


def _date(c, v, fmt: str = DATE_FMT) -> None:
    if v is not None:
        c.number_format = fmt


def _finish(ws, *, filters: bool = True) -> None:
    """Filters over the whole table and each column as wide as its longest value (at most 40)."""
    if filters and ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if v is None:
                continue
            n = 12 if isinstance(v, (date, datetime)) else len(str(v))
            widths[c.column] = max(widths.get(c.column, 0), n)
    for col, n in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(8, min(MAX_WIDTH, n + 2))


def summarise(statuses: list[str]) -> dict:
    n = {k: 0 for k in ("ok", "no_expiry", "expiring", "expired", "needs_date", "missing", "not_required")}
    for s in statuses:
        n[s] = n.get(s, 0) + 1
    out = {"valid": n["ok"] + n["no_expiry"], "expiring": n["expiring"], "expired": n["expired"], "missing": n["missing"] + n["needs_date"],
           "not_required": n["not_required"]}
    out["worst"] = "expired" if out["expired"] else "expiring" if out["expiring"] else "missing" if out["missing"] else "ok" if out["valid"] else "empty"
    return out


def workbook(data: dict, T: Callable = _en, D: Callable = _en_date) -> bytes:
    """The workbook from plain values:
    data = {"me": {"name", "email"}, "is_admin": bool, "at": local datetime, "tz": str, "today": date,
            "types": [{"id", "name", "short", "warn_days", "order", "shown"}]  (all, in the grid's order),
            "companies": [{"id", "name", "responsible", "responsible_email", "responsible_account", "open_to_all", "notes",
                           "created", "archived", "archived_at", "archived_reason", "people": [names] | None}],
            "documents": [{"id", "company_id", "type_id" | None, "title", "expiry", "no_expiry", "required", "status",
                           "file_name", "uploaded_by", "uploaded_at", "notes"}]}  (live documents only)."""
    today: date = data["today"]
    types = list(data.get("types") or [])
    shown = [t for t in types if t.get("shown", True)]
    tname = {t["id"]: t["name"] for t in types}
    live = [c for c in data.get("companies") or [] if not c.get("archived")]
    gone = [c for c in data.get("companies") or [] if c.get("archived")] if data.get("is_admin") else []
    docs = list(data.get("documents") or [])
    cell: dict[tuple, dict] = {}
    other: dict[int, list] = {}
    for d in sorted(docs, key=lambda x: x.get("id") or 0):  # the newest row of a cell is the cell
        if d.get("type_id"):
            cell[(d["company_id"], d["type_id"])] = d
        else:
            other.setdefault(d["company_id"], []).append(d)

    def rows_of(c: dict) -> list[dict]:
        """The company's documents as the grid has them: one per shown column (an empty one is missing), then the others."""
        out = []
        for t in shown:
            d = cell.get((c["id"], t["id"]))
            out.append(d if d is not None else {"id": None, "company_id": c["id"], "type_id": t["id"], "title": t["name"], "expiry": None,
                                                 "no_expiry": False, "required": True, "status": "missing", "file_name": "", "uploaded_by": "",
                                                 "uploaded_at": None, "notes": ""})
        return out + other.get(c["id"], [])

    per = {c["id"]: rows_of(c) for c in live + gone}
    summ = {cid: summarise([d["status"] for d in rows]) for cid, rows in per.items()}
    live.sort(key=lambda c: (WORST_ORDER.get(summ[c["id"]]["worst"], 9), (c.get("name") or "").lower()))
    gone.sort(key=lambda c: (c.get("name") or "").lower())
    email = lambda c: c.get("responsible_email") or c.get("responsible_account") or ""  # noqa: E731
    yes = lambda v: T("Yes") if v else T("No")  # noqa: E731

    wb = Workbook()
    # ---- Overview -------------------------------------------------------------------------------------
    ws = _sheet(wb, T("Overview"), [T("Entity"), T("Responsible person"), T("Responsible email"), T("Valid"), T("Expiring soon"), T("Expired"),
                                     T("Missing"), T("Not required"), T("Overall")] + [t["name"] for t in shown], first=True)
    ws.freeze_panes = "B2"
    for c in live:
        s = summ[c["id"]]
        word, paint = OVERALL[s["worst"]]
        ws.append([c["name"], c.get("responsible") or "", email(c), s["valid"], s["expiring"], s["expired"], s["missing"], s["not_required"], T(word)]
                  + [state_word(d["status"], d.get("expiry"), today, T, D) for d in per[c["id"]][:len(shown)]])
        r = ws.max_row
        _paint(ws.cell(r, 9), paint)
        for i, d in enumerate(per[c["id"]][:len(shown)]):
            _paint(ws.cell(r, 10 + i), d["status"])
    _finish(ws)

    # ---- Documents ------------------------------------------------------------------------------------
    ws = _sheet(wb, T("Documents"), [T("Entity"), T("Document type"), T("Title"), T("Expiry date"), T("Days left"), T("State"), T("Required"),
                                      T("No expiry"), T("File name"), T("Uploaded by"), T("Uploaded at"), T("Notes")])
    n_docs = 0
    for c in live:
        for d in per[c["id"]]:
            exp = d.get("expiry")
            left = (exp - today).days if exp and not d.get("no_expiry") else None
            ws.append([c["name"], tname.get(d.get("type_id"), T("Other")) if d.get("type_id") else T("Other"), d.get("title") or "", exp, left,
                       state_word(d["status"], exp, today, T, D), yes(d.get("required", True)), yes(d.get("no_expiry")), d.get("file_name") or "",
                       d.get("uploaded_by") or "", d.get("uploaded_at"), d.get("notes") or ""])
            r = ws.max_row
            _date(ws.cell(r, 4), exp)
            _date(ws.cell(r, 11), d.get("uploaded_at"), TIME_FMT)
            _paint(ws.cell(r, 6), d["status"])
            n_docs += 1 if d.get("id") else 0
    _finish(ws)

    # ---- Expiring and expired: the to-do list ---------------------------------------------------------
    ws = _sheet(wb, T("Expiring and expired"), [T("Entity"), T("Document type"), T("Title"), T("Expiry date"), T("Days left"), T("State"),
                                                 T("Responsible person"), T("Responsible email")])
    todo = [(c, d) for c in live for d in per[c["id"]] if d["status"] in TODO]
    todo.sort(key=lambda x: (x[1].get("expiry") is None, x[1].get("expiry") or date.max, (x[0].get("name") or "").lower()))
    for c, d in todo:
        exp = d.get("expiry")
        ws.append([c["name"], tname.get(d.get("type_id"), T("Other")) if d.get("type_id") else T("Other"), d.get("title") or "", exp,
                   (exp - today).days if exp else None, state_word(d["status"], exp, today, T, D), c.get("responsible") or "", email(c)])
        r = ws.max_row
        _date(ws.cell(r, 4), exp)
        _paint(ws.cell(r, 6), d["status"])
    _finish(ws)

    # ---- Entities -------------------------------------------------------------------------------------
    admin = bool(data.get("is_admin"))
    head = [T("Entity"), T("Responsible person"), T("Responsible account"), T("Responsible email"), T("Who sees it")]
    ws = _sheet(wb, T("Entities"), head + ([T("Given access")] if admin else []) + [T("Notes"), T("Created")])
    for c in sorted(live, key=lambda x: (x.get("name") or "").lower()):
        ws.append([c["name"], c.get("responsible") or "", c.get("responsible_account") or "", c.get("responsible_email") or "",
                   T("Everyone with the tool") if c.get("open_to_all") else T("Only its people")]
                  + ([", ".join(c.get("people") or [])] if admin else []) + [c.get("notes") or "", c.get("created")])
        _date(ws.cell(ws.max_row, ws.max_column), c.get("created"))
    _finish(ws)

    # ---- Archived entities (admins) ---------------------------------------------------------------------
    if gone:
        ws = _sheet(wb, T("Archived entities"), [T("Entity"), T("Responsible person"), T("Responsible email"), T("Archived on"), T("Why archived"),
                                                  T("Overall"), T("Documents"), T("Given access"), T("Notes"), T("Created")])
        for c in gone:
            s = summ[c["id"]]
            word, paint = OVERALL[s["worst"]]
            ws.append([c["name"], c.get("responsible") or "", email(c), c.get("archived_at"), c.get("archived_reason") or "", T(word),
                       sum(1 for d in per[c["id"]] if d.get("id")), ", ".join(c.get("people") or []), c.get("notes") or "", c.get("created")])
            r = ws.max_row
            _date(ws.cell(r, 4), c.get("archived_at"))
            _paint(ws.cell(r, 6), paint)
            _date(ws.cell(r, 10), c.get("created"))
        _finish(ws)

    # ---- Document types -----------------------------------------------------------------------------------
    ws = _sheet(wb, T("Document types"), [T("Name"), T("Short name"), T("Warn (days before)"), T("Order"), T("Active")])
    for t in types:
        ws.append([t["name"], t.get("short") or "", t.get("warn_days"), t.get("order"), yes(t.get("shown", True))])
    _finish(ws)

    # ---- About ----------------------------------------------------------------------------------------------
    ws = _sheet(wb, T("About"), [T("Entity Compliance"), ""])
    me = data.get("me") or {}
    about = [(T("Exported by"), f"{me.get('name') or ''} ({me.get('email') or ''})"), (T("Exported at"), data.get("at")),
             (T("Time zone"), data.get("tz") or "Africa/Johannesburg"), (T("Entities"), len(live))]
    if admin:
        about.append((T("Archived entities"), len(gone)))
    about += [(T("Documents"), n_docs), (T("Who is in it"), T("Only the entities you may see."))]
    for k, v in about:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    _date(ws.cell(3, 2), data.get("at"), TIME_FMT)
    ws.append([])
    ws.append([T("Colour key")])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    sample = today + timedelta(days=400)
    for status, exp, meaning in (("ok", sample, T("In date")), ("expiring", today + timedelta(days=23), T("Inside its warning days")),
                                 ("expired", today - timedelta(days=10), T("Past its expiry date")), ("missing", None, T("Nothing filed yet")),
                                 ("needs_date", None, T("Filed, but nobody has set its expiry date")),
                                 ("not_required", None, T("Not needed for this entity"))):
        ws.append([state_word(status, exp, today, T, D), meaning])
        _paint(ws.cell(ws.max_row, 1), status)
    _finish(ws, filters=False)

    xlsafe.book(wb)  # a company or file name starting with = is text, never a formula
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- the old CBX's side: the tool folder's tables, through the engine's scopes ---------------------------------
SLUG = "entity-compliance"
PAGE = 2000  # the most Store.rows returns at once


def _all_rows(db, me, t, block: str) -> list[dict]:
    """Every row of one of the tool's tables that `me` may see (the grid's own scope), page by page."""
    from ..api.engine_api import _mine_scope, label_people
    from ..engine.store import Store

    b = {**t.data[block], "_slug": SLUG, "_name": block}
    cols = b.get("columns") or []
    mine = _mine_scope(db, b, me, t)
    st = Store(SLUG)
    out: list[dict] = []
    while True:
        got = st.rows(block, cols, limit=PAGE, offset=len(out), mine=mine, sort="id", tables=t.data)["rows"]
        out += got
        if len(got) < PAGE:
            break
    label_people(db, out, cols)
    return out


def _day(v) -> Optional[date]:
    try:
        return date.fromisoformat(str(v)[:10]) if v not in (None, "") else None
    except ValueError:
        return None


def collect(db, me, t) -> dict:
    """The plain values for `workbook()` from the tool's tables."""
    from ..config import settings
    from ..models import User
    from ..util import local_today

    tz = getattr(settings, "timezone", "") or "Africa/Johannesburg"
    today = local_today()
    comps = [c for c in _all_rows(db, me, t, "companies") if me.is_admin or not c.get("archived")]
    ids = {c["id"] for c in comps}
    docs = [d for d in _all_rows(db, me, t, "documents") if d.get("company") in ids and not d.get("archived") and d.get("status") != "archived"]
    people: dict[int, list[str]] = {}
    if me.is_admin and "company_people" in t.data:
        for p in _all_rows(db, me, t, "company_people"):
            if p.get("company") in ids and isinstance(p.get("person"), int):
                people.setdefault(p["company"], []).append((p.get("_labels") or {}).get("person") or "")
    uids = {(d.get("file") or {}).get("by") for d in docs if isinstance(d.get("file"), dict)} | {c.get("responsible_account") for c in comps}
    uids = {u for u in uids if isinstance(u, int)}
    users = {u.id: u for u in db.query(User).filter(User.id.in_(uids)).all()} if uids else {}
    out_c = []
    for c in comps:
        acc = users.get(c.get("responsible_account")) if isinstance(c.get("responsible_account"), int) else None
        out_c.append({"id": c["id"], "name": c.get("name") or f"#{c['id']}", "responsible": c.get("responsible") or (acc.name if acc else ""),
                      "responsible_email": c.get("responsible_email") or "", "responsible_account": acc.email if acc else "",
                      "open_to_all": (c.get("visible") or "Everyone with the tool") == "Everyone with the tool", "notes": c.get("notes") or "",
                      "created": local(c.get("created_at"), tz), "archived": bool(c.get("archived")), "archived_at": None,
                      "archived_reason": c.get("archived_reason") or "",
                      "people": sorted((n for n in people.get(c["id"], []) if n), key=str.lower) if me.is_admin else None})
    out_d = []
    for d in docs:
        f = d.get("file") if isinstance(d.get("file"), dict) else None
        by = users.get(f.get("by")) if f and isinstance(f.get("by"), int) else None
        out_d.append({"id": d["id"], "company_id": d["company"], "type_id": d.get("doc_type") if isinstance(d.get("doc_type"), int) else None,
                      "title": d.get("title") or "", "expiry": _day(d.get("expiry")), "no_expiry": bool(d.get("no_expiry")),
                      "required": d.get("required") is not False, "status": str(d.get("status") or "missing"),
                      "file_name": (f or {}).get("name") or "", "uploaded_by": by.name if by else ((d.get("_labels") or {}).get("filed_by") or "" if f else ""),
                      "uploaded_at": local((f or {}).get("at"), tz) if f else None, "notes": d.get("notes") or ""})
    types = sorted(_all_rows(db, me, t, "doc_types"), key=lambda x: (x.get("order") if isinstance(x.get("order"), int) else 100, str(x.get("name") or "").lower()))
    out_t = [{"id": x["id"], "name": x.get("name") or f"#{x['id']}", "short": x.get("short") or "", "warn_days": x.get("warn_days"), "order": x.get("order"),
              "shown": x.get("shown") is not False} for x in types]
    return {"me": {"name": me.name, "email": me.email}, "is_admin": bool(me.is_admin), "at": local(datetime.now(timezone.utc), tz), "tz": tz,
            "today": today, "types": out_t, "companies": out_c, "documents": out_d}


def export(db, me, t) -> tuple[str, bytes]:
    """The file name and the workbook; one audit row with the counts."""
    from ..audit import log

    data = collect(db, me, t)
    body = workbook(data)
    live = {c["id"] for c in data["companies"] if not c["archived"]}
    n_docs = sum(1 for d in data["documents"] if d["company_id"] in live)
    log(db, "ec_export", actor=me, target=SLUG, detail=f"{len(live)} entities, {n_docs} documents")
    return f"Entity Compliance {data['today'].isoformat()}.xlsx", body
