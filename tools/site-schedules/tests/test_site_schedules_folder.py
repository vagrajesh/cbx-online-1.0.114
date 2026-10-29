"""Site Schedules as a tool folder (1.0.56; CBX-Connect 2's engine since 1.0.88): the Swiftnet PDFs on a row → the server
job reads them into the row → lines to check → Make the agreement (Word, like the real one) once they are ticked."""
import io
import json

from cbx.engine import loader

BASE = "/api/engine/tools/site-schedules/data"


def _row(c, rid):
    return next(r for r in c.get(f"{BASE}/quotes").json()["rows"] if r["id"] == rid)


def test_folder_is_sound():
    t = loader.registry().tools.get("site-schedules")
    assert t is not None and t.ok, (t.errors if t else "not loaded")
    assert set(t.pages) == {"quotes"} and set(t.flows) == {"read", "reread", "build"}
    for lang in ("en", "ar", "pt", "fr"):
        assert t.languages.get(lang), lang
    assert not any(t.missing().values()), t.missing()
    from cbx.engine.rules import why_not

    assert {n: why_not(f) for n, f in t.flows.items()} == {n: None for n in t.flows}


def test_a_quote_is_read_checked_and_built(admin, db):
    """1.0.88: CBX-Connect 2's engine — a real-shaped Swiftnet quote (and its NTP) in, each comment its own line,
    the Word agreement laid out like the real one, named like the real files."""
    from tests._swiftnet import REF, ntp, quote

    r = admin.post(f"{BASE}/quotes", json={"values": {}}).json()
    assert r["status"] == "New"
    up = admin.post(f"{BASE}/quotes/{r['id']}/files/quote", files=[("files", (REF + ".pdf", quote(), "application/pdf"))])
    assert up.status_code == 200, up.text
    row = _row(admin, r["id"])  # the upload is a change: the rule read it
    assert row["status"] == "Read" and row["site_name"] == "TIERVLEI EXCHANGE" and row["company"] == "Vodacom (Pty) Ltd", row
    assert row["contract_no"] == "70204797" and row["lease_type"] == "Renewal" and row["start_date"] == "01 August 2026"
    assert row["monthly"] == 6305.34 and "70205157* Ground space for 35m" in row["charges"]
    assert row["charges"].count("6.1 Rental") == 5  # every comment its own line
    assert any("Co-user site" in c["text"] or "Co-User site" in c["text"] for c in row["checks"]), row["checks"]
    # the NTP added later, read again: the co-user site comes in
    admin.post(f"{BASE}/quotes/{r['id']}/files/more", files=[("files", ("NTP Tiervlei.pdf", ntp(), "application/pdf"))])
    run = admin.post("/api/flows/site-schedules/reread/run", json={"values": {}, "block": "quotes", "row": r["id"]}).json()["run"]
    assert run["status"] == "done", run
    row = _row(admin, r["id"])
    assert row["co_site_name"] == "Beaconvale" and row["bs_site_no"] == "2981"
    # not before every line is ticked
    run = admin.post("/api/flows/site-schedules/build/run", json={"values": {}, "block": "quotes", "row": r["id"]}).json()["run"]
    if row["checks_n"]:
        assert run["status"] == "failed" and "Tick every line" in run["note"]
    # a person corrects a field, adds who signs, ticks the lines
    admin.patch(f"{BASE}/quotes/{r['id']}", json={"values": {"checks": json.dumps([{**c, "done": True} for c in row["checks"]]),
                                                             "swiftnet_rep": "Gert Vorster", "swift_signers": "Gert Vorster, CFO for Swiftnet"}})
    run = admin.post("/api/flows/site-schedules/build/run", json={"values": {}, "block": "quotes", "row": r["id"]}).json()["run"]
    assert run["status"] == "done", run
    row = _row(admin, r["id"])
    assert row["status"] == "Built" and row["agreement"]["name"] == "9130 - 5183 - 70204797 - TIERVLEI EXCHANGE - 01.08.2026 - 31.07.2031 - 7708 - R - VDC.docx"
    doc = admin.get(f"{BASE}/quotes/{r['id']}/files/agreement")
    assert doc.status_code == 200 and doc.content.startswith(b"PK")
    import docx

    d = docx.Document(io.BytesIO(doc.content))
    text = "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(" | ".join(c.text for c in rw.cells) for t in d.tables for rw in t.rows)
    for want in ("SITE SCHEDULE AGREEMENT", "Renewal", "TIERVLEI EXCHANGE", "T100/9130", "01 August 2026", "31 July 2031", "01 April 2027",
                 "R 2 485.25", "Total Rental Charges Payable (Excl. VAT) per Month | R 6 305.34", "Gert Vorster", "CFO: Swiftnet", "Eddy Kgomongwe", "Paul Jansen Van Rensburg", "Signed at Pretoria", "Beaconvale"):  # 1.0.107: the fixed signatories
        assert want in text, want
    assert "{" not in text and "**" not in text


def test_a_pdf_that_is_not_swiftnet_says_so(admin, db):
    from tests._pdfs import swiftnet_quote

    r = admin.post(f"{BASE}/quotes", json={"values": {}}).json()
    admin.post(f"{BASE}/quotes/{r['id']}/files/quote", files=[("files", ("odd.pdf", swiftnet_quote(), "application/pdf"))])
    run = admin.post("/api/flows/site-schedules/reread/run", json={"values": {}, "block": "quotes", "row": r["id"]}).json()["run"]
    assert run["status"] == "failed" and "Swiftnet quote" in run["note"], run


def test_someone_without_the_tool_cannot_read_quotes(owner, db):
    from fastapi.testclient import TestClient

    from cbx.main import app
    from tests.conftest import _make, login

    pw = _make(owner, "nosched@test.local", "No Schedules", "user", tools=[])
    c = TestClient(app)
    login(c, "nosched@test.local", pw)
    assert c.get(f"{BASE}/quotes").status_code == 403
    assert c.post("/api/flows/site-schedules/build/run", json={"values": {}, "block": "quotes", "row": 1}).status_code in (403, 404)


def test_the_pack_prints_t100_and_t000_with_their_prefix_and_no_empty_clause_7_headings():
    """1.0.108: the row's T numbers in any shape ("T100 9130", "9130", the other prefix) print as "T100/<n>" and
    "T000/<n>"; none, no line. A clause 7 sub-heading only with text under it."""
    import docx

    from cbx.schedules import agreement as A
    from cbx.schedules import template as T

    def paras(data):
        return [p.text for p in docx.Document(io.BytesIO(data)).paragraphs]

    details = {"fields": {k: v for k, v in T.sample_agreement()["fields"].items()}, "tables": T.sample_agreement()["tables"]}
    for t100, t000 in (("T100 9130", "5183"), ("9130", "T100/5183"), ("t100-9130", "T000 / 5183")):
        name, data = A.build(details, {"t100_ref": t100, "t000_ref": t000, "scope": "", "equip_comments": "Plinth."})
        p = paras(data)
        i = p.index("Contract no.: 70204797")
        assert p[i + 1:i + 3] == ["T100/9130", "T000/5183"] and name.startswith("9130 - 5183 - 70204797"), (t100, t000)
        c7 = p[next(k for k, x in enumerate(p) if "SCOPE OF APPROVED WORK" in x):next(k for k, x in enumerate(p) if "SPECIAL REQUESTS" in x)]
        assert [x for x in c7[2:] if x.strip()] == ["Equipment accommodation:", "Plinth."]
    name, data = A.build(details, {"t100_ref": "", "t000_ref": "T000/", "lease_type": "New Application", "contract_no": ""})
    assert not any(x.startswith(("T100", "T000")) for x in paras(data)) and name.startswith("SAMPLE EXCHANGE - ")
    name, _ = A.build(details, {"t100_ref": "", "t000_ref": "5183"})
    assert name.startswith("T000-5183 - 70204797 - ")


def test_the_pack_escalates_and_names_as_the_page():
    """1.0.108: the pack's read uses the reader's first-escalation rule (one function), and its file name's lease code
    follows the lease type."""
    from cbx.schedules import agreement as A
    from cbx.schedules import reader as R
    from cbx.schedules import template as T
    from tests._swiftnet import REF, quote

    assert A.R.first_escalation_year is R.first_escalation_year
    for start, year in (("15 January 2027", 2027), ("31 March 2027", 2027), ("1 April 2027", 2028), ("1 August 2026", 2027)):
        assert A.read([(REF + ".pdf", quote(start=start))])["esc_year"] == year, start
    details = {"fields": dict(T.sample_agreement()["fields"]), "tables": T.sample_agreement()["tables"]}
    for lease, code in (("New Application", "N"), ("Renewal", "R"), ("Re-Assessment", "RA"), ("New Re-Assessment", "NA"), ("MOU", "MOU")):
        name, _ = A.build(details, {"lease_type": lease, "lease_code": "R"})
        assert name.endswith(f" - {code}.docx"), (lease, name)
