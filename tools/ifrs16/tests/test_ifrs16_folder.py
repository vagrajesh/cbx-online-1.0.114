"""IFRS 16 as a tool folder (1.0.56): keep the source files → read for their months (a slow job) → work out
a month (the lease engine, in the background) → import workbooks and report on the month's row; the ledger check."""
import io
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

from cbx.engine import loader

BASE = "/api/engine/tools/ifrs16/data"
FIX = Path(__file__).resolve().parents[3] / "tests" / "fixtures" / "ifrs16"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _up(c, table, rid, col, name, data):
    r = c.post(f"{BASE}/{table}/{rid}/files/{col}", files=[("files", (name, data, XLSX))])
    assert r.status_code == 200, r.text


def _row(c, table, rid):
    return next(r for r in c.get(f"{BASE}/{table}").json()["rows"] if r["id"] == rid)


def _kept(c):
    u = c.post(f"{BASE}/uploads", json={"values": {}}).json()
    _up(c, "uploads", u["id"], "la", "la_register.xlsx", (FIX / "la_register.xlsx").read_bytes())
    _up(c, "uploads", u["id"], "movements", "movements_rr06.xlsx", (FIX / "movements_rr06.xlsx").read_bytes())
    return _row(c, "uploads", u["id"])


def test_folder_is_sound():
    t = loader.registry().tools.get("ifrs16")
    assert t is not None and t.ok, (t.errors if t else "not loaded")
    assert set(t.pages) == {"months", "ledger", "sequentra", "rates"} and set(t.flows) == {"read", "month", "ledger", "api_check", "api_download"}
    for lang in ("en", "ar", "pt", "fr"):
        assert t.languages.get(lang), lang
    assert not any(t.missing().values()), t.missing()
    from cbx.engine.rules import why_not

    assert {n: why_not(f) for n, f in t.flows.items()} == {n: None for n in t.flows}


def test_files_are_read_and_a_month_is_worked_out(user):
    u = _kept(user)
    assert u["status"] == "Ready", u  # the slow job ran (in line in the suite) once the upload was saved
    assert "2026-08" in u["months"] and "LA register rows" in u["summary"] and u["movement"] == "RR06"
    run = user.post("/api/flows/ifrs16/month/run", json={"values": {"answer": "1999-01"}, "block": "uploads", "row": u["id"]}).json()["run"]
    assert run["status"] == "failed" and "not in the movements file" in run["note"]
    run = user.post("/api/flows/ifrs16/month/run", json={"values": {"answer": "2026-08", "source": "Files"}, "block": "uploads", "row": u["id"]}).json()["run"]
    assert run["status"] == "waiting" and run["waiting"]["kind"] == "working"
    assert user.get(f"/api/flows/runs/{run['id']}").json()["run"]["status"] == "done"
    m = next(r for r in user.get(f"{BASE}/months").json()["rows"] if r["upload"] == u["id"] and r["month"] == "2026-08")
    assert m["status"] == "Done" and m["report"]["name"].startswith("IFRS16_Report_2026-08") and "Workflow month" in m["recon"] and m["source"] == "Files"
    assert "Discount rates: none on the Rates page" in m["recon"]  # an SA month with an empty Rates page says so
    rep = user.get(f"{BASE}/months/{m['id']}/files/report")
    assert rep.status_code == 200
    assert load_workbook(io.BytesIO(rep.content), read_only=True).sheetnames
    if m["imports"]:
        z = zipfile.ZipFile(io.BytesIO(user.get(f"{BASE}/months/{m['id']}/files/workbooks").content))
        assert len(z.namelist()) == m["imports"]


def test_unreadable_files_say_so_on_the_row(user):
    wb = Workbook()
    wb.active.append(["Nothing", "to", "see"])
    bio = io.BytesIO()
    wb.save(bio)
    u = user.post(f"{BASE}/uploads", json={"values": {}}).json()
    _up(user, "uploads", u["id"], "la", "junk.xlsx", bio.getvalue())
    _up(user, "uploads", u["id"], "movements", "junk2.xlsx", bio.getvalue())
    row = _row(user, "uploads", u["id"])
    assert row["status"] == "Could not read" and row["note"]


def test_the_ledger_check(user):
    wb = Workbook()
    ws = wb.active
    ws.title = "Ledger Entries"
    ws.append(["Ledger Date", "Segment2", "Segment3", "Segment4", "DR", "CR", "Comments", "JEShortDesc", "Transactional Currency"])
    for r in [("2026-06-30", "BC000", "P1", "CC1", 12000, 0, "Schedule TXN00001 initial", "Initial Recognition", "ZAR"),
              ("2026-06-30", "HA101", "P1", "CC1", 0, 12000, "Schedule TXN00001 initial", "Initial Recognition", "ZAR"),
              ("2026-07-31", "HA101", "P1", "CC1", 1000, 0, "Schedule TXN00001 payment", "Payment", "ZAR"),
              ("2026-07-31", "SB424", "P1", "CC1", 0, 1000, "Schedule TXN00001 payment", "Payment", "ZAR")]:
        ws.append(list(r))
    bio = io.BytesIO()
    wb.save(bio)
    r = user.post(f"{BASE}/ledger", json={"values": {"as_at": "2026-08"}}).json()
    assert r["months_back"] == 12 and r["status"] == "New"
    _up(user, "ledger", r["id"], "ledger", "Ledger Export.xlsx", bio.getvalue())
    row = _row(user, "ledger", r["id"])
    assert row["status"] == "Done", row
    assert row["report"]["name"] == "IFRS16_Consolidated_Report.xlsx" and "Leases included in the check" in row["overview"]


def test_people_see_their_own_and_nobody_adds_a_month_by_hand(user, admin):
    u = _kept(user)
    assert any(r["id"] == u["id"] for r in admin.get(f"{BASE}/uploads").json()["rows"])  # admins see all
    assert user.post(f"{BASE}/months", json={"values": {"month": "2026-08"}}).status_code == 403


# ---- 1.0.80: the desktop IFRS 16 Workflow 5.16 — the Sequentra API, the cross-check, the Rates page ---------------
KEY = "TEST-CLIENTKEY-0042"


def _fake_sequentra(monkeypatch, seen):
    """A stand-in for https://api.excellerate.tech/sequentra/: the engine's own client is used, only its one
    network call is replaced (sequentra_api._http_get — the seam the desktop's own tests use)."""
    import json

    from cbx.ifrs16 import ENGINE_DIR  # noqa: F401 — puts the engine on sys.path
    import sequentra_api as sq

    def get(url, header, key, timeout):
        seen.append((url, header, key))
        if key != KEY:
            raise sq.SequentraAPIError("HTTP 401: the key was refused", status=401, short="HTTP 401")
        if "buildingrecordtabs" in url or "keydateoptionstab" in url:
            rows = [{"recordId": "TXN00001", "recordName": "Test Branch", "country": "South Africa", "currency": "ZAR",
                     "currentTermStartDate": "2026-06-01T00:00:00", "currentLeaseEndDate": "2031-05-31T00:00:00"}]
        else:
            rows = []
        return 200, "application/json", json.dumps({"value": rows}).encode()

    monkeypatch.setattr(sq, "_http_get", get)


def _connect(admin, key=KEY, tools=("ifrs16",)):
    r = admin.post("/api/connections", json={"name": "Sequentra", "base_url": "https://api.excellerate.tech/sequentra/", "auth_kind": "header",
                                             "header_name": "Clientkey", "secret": key, "tools": list(tools)})
    assert r.status_code == 200, r.text
    return r.json()["connection"]


def _drop_connections(admin):
    for c in admin.get("/api/connections").json()["connections"]:
        admin.delete(f"/api/connections/{c['id']}")


def test_the_month_button_asks_where_the_lease_data_comes_from():
    t = loader.registry().tools["ifrs16"]
    btn = t.pages["months"]["layout"]["sections"][0]["columns"][1]["widgets"][0]["actions"][0]
    assert [q["type"] for q in btn["ask"]] == ["month", "choice"] and [o["value"] for o in btn["ask"][1]["options"]] == ["Files", "Sequentra API"]


def test_without_a_connection_the_api_month_says_how_to_add_it(user, admin):
    _drop_connections(admin)
    u = _kept(user)
    run = user.post("/api/flows/ifrs16/month/run", json={"values": {"answer": "2026-08", "source": "Sequentra API"}, "block": "uploads", "row": u["id"]}).json()["run"]
    m = next(r for r in user.get(f"{BASE}/months").json()["rows"] if r["upload"] == u["id"] and r["month"] == "2026-08")
    assert m["status"] == "Failed" and "no Sequentra connection" in m["note"] and "Settings → API connections" in m["note"], (run, m)


def test_a_month_from_the_sequentra_api_gets_the_cross_check_and_the_key_stays_put(user, admin, owner, monkeypatch):
    seen = []
    _drop_connections(admin)
    _fake_sequentra(monkeypatch, seen)
    _connect(admin)
    u = _kept(user)
    user.post("/api/flows/ifrs16/month/run", json={"values": {"answer": "2026-08", "source": "Sequentra API"}, "block": "uploads", "row": u["id"]})
    m = next(r for r in user.get(f"{BASE}/months").json()["rows"] if r["upload"] == u["id"] and r["month"] == "2026-08" and r["source"] == "Sequentra API")
    assert m["status"] == "Done", m
    wb = load_workbook(io.BytesIO(user.get(f"{BASE}/months/{m['id']}/files/papers").content), read_only=True)
    assert {"Recon Summary", "Routing", "Exceptions", "API Cross-Check"} <= set(wb.sheetnames), wb.sheetnames
    assert "API cross-check: " in m["recon"]  # the API notes are on the row too
    # the client asked Sequentra, with the key in the Clientkey header — and the key is nowhere a person can read
    assert seen and all(h == "Clientkey" and k == KEY for _, h, k in seen) and all(KEY not in url for url, _, _ in seen)
    import json

    for rows in (user.get(f"{BASE}/months").json()["rows"], owner.get("/api/audit?limit=200").json()["rows"], admin.get("/api/connections").json()["connections"]):
        assert KEY not in json.dumps(rows)


def test_check_and_download_the_sequentra_reports(user, admin, monkeypatch):
    seen = []
    _drop_connections(admin)
    _fake_sequentra(monkeypatch, seen)
    _connect(admin)
    r = user.post(f"{BASE}/api", json={"values": {"kind": "Check", "month": "2026-08"}}).json()
    row = _row(user, "api", r["id"])
    assert row["status"] == "Done" and row["ok"] == "yes" and "Connected" in row["result"] and "Building & Record" in row["result"], row
    r = user.post(f"{BASE}/api", json={"values": {"kind": "Download", "month": "2026-07", "to_month": "2026-08"}}).json()
    row = _row(user, "api", r["id"])
    assert row["status"] == "Done" and row["download"]["name"] == "Sequentra_2026-07_to_2026-08.zip", row
    z = zipfile.ZipFile(io.BytesIO(user.get(f"{BASE}/api/{row['id']}/files/download").content))
    assert any(n.endswith("00_Summary.csv") for n in z.namelist())
    assert KEY.encode() not in b"".join(z.read(n) for n in z.namelist())
    # a wrong key: said plainly on the row
    _drop_connections(admin)
    _connect(admin, key="WRONG-KEY-9999")
    r = user.post(f"{BASE}/api", json={"values": {"kind": "Check"}}).json()
    row = _row(user, "api", r["id"])
    assert row["ok"] == "no" and "401" in row["result"], row
    # not ticked for IFRS 16 → refused before any call
    _drop_connections(admin)
    _connect(admin, tools=())
    n = len(seen)
    row = _row(user, "api", user.post(f"{BASE}/api", json={"values": {"kind": "Check"}}).json()["id"])
    assert row["status"] == "Failed" and "not ticked for IFRS 16" in row["note"] and len(seen) == n
    _drop_connections(admin)


def test_sa_rates_come_from_the_rates_page(user, admin):
    r = user.post(f"{BASE}/rates", json={"values": {"month": "2026-06", "a": 10.5, "b": 10.75, "c": 11.0, "d": 11.25, "e": 11.5, "f": 11.75}})
    assert r.status_code == 200, r.text
    u = _kept(user)
    user.post("/api/flows/ifrs16/month/run", json={"values": {"answer": "2026-08", "source": "Files"}, "block": "uploads", "row": u["id"]})
    m = next(x for x in user.get(f"{BASE}/months").json()["rows"] if x["upload"] == u["id"] and x["month"] == "2026-08")
    assert m["status"] == "Done" and "Discount rates: none" not in m["recon"], m
    assert any(x["month"] == "2026-06" for x in admin.get(f"{BASE}/rates").json()["rows"])  # one table, shared
    admin.delete(f"{BASE}/rates/{r.json()['id']}")
